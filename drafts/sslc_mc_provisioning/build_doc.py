"""Build SSLC Mission Control + email provisioning analysis (.docx)."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/Users/exomac/.openclaw/workspace/drafts/sslc_mc_provisioning/SSLC_MC_Provisioning_Analysis_2026-10-19_r3.docx"

d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)


def link(par, text, url):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); par._p.append(h)


def p(text="", bold=False, style=None):
    para = d.add_paragraph(style=style)
    if text:
        run = para.add_run(text); run.bold = bold
    return para


def bullet(text, src=None):
    para = d.add_paragraph(style="List Bullet")
    para.add_run(text)
    if src:
        para.add_run(" (source: ")
        link(para, src[0], src[1])
        para.add_run(")")
    return para


def table(rows, widths=None):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j); cell.text = ""
            run = cell.paragraphs[0].add_run(str(val)); run.font.size = Pt(9.5)
            if i == 0: run.bold = True
    d.add_paragraph()
    return t


SRC_PROTON_WS = ("proton.me/business/plans", "https://proton.me/business/plans")
SRC_PROTON_ESS = ("Proton, A Practical Guide to Security for Growing Businesses", "https://proton.me/files/A_Practical_Guide_to_Security_for_Growing_Businesses.pdf")
SRC_MS = ("microsoft.com, compare Microsoft 365 business plans", "https://www.microsoft.com/en-us/microsoft-365/business/compare-all-microsoft-365-business-products")
SRC_CF = ("cloudflare.com, Access plans", "https://www.cloudflare.com/sase/products/access")
SRC_ITAR = ("Federal Register 84 FR 70887, 22 CFR 120.54(a)(5)", "https://www.federalregister.gov/documents/2019/12/26/2019-27438/international-traffic-in-arms-regulations-creation-of-definition-of-activities-that-are-not-exports")

h = d.add_heading("SSLC Mission Control and Email Provisioning: Analysis and Recommendation", level=1)
p("Prepared by Exo for John Quinn, CEO | 19 October 2026 | r3 | INTERNAL. r3 (Jon 05:15): records that Board Motion 2 books access is established through Mission Control and the director access emails were sent. Not for distribution to SSLC in this form. r1 changes (Jon 04:50-04:51): email cost is $0, and paid mailboxes are dropped. The evaluation uses the SSLC QuickBooks access we already have. The system is packaged as Docker containers, one per company, for reuse. The deliverable is a live demo for the SSLC board. r2 (Jon 04:55): adds the post-acquisition end state (Exos email for all SSLC staff; the full Mission Control suite in SSLC's own section), and the demo now covers every module.")

d.add_heading("1. Bottom line", level=2)
bullet("Deliverable: a live Mission Control demo for the SSLC board, running SSLC's real QuickBooks data in its own container on the Mac Studio. Only SSLC-authorized people and named Exos evaluators can see it.")
bullet("Containerize: package Mission Control as one Docker image, and run one container per company. Each container has its own data volume (database, users, documents), port, hostname and login policy. The SSLC container cannot read Exos files, because they are not in its container. Docker is already running on the Studio (colima, which also hosts Vaultwarden).")
bullet("Don't build it as a tab inside the Exos Mission Control. None of the 47 ERP tables has a company field, and the QuickBooks sync assumes one company, so SSLC data would mix with Exos data.")
bullet("Board Motion 2 is done: SSLC books access was established through Mission Control on 10/19. The QuickBooks (indirect, via Mission Control) access emails were sent from Proton to Markus Rufer and Robert Conger at 03:58 CDT on 10/19, cc John. The same Mission Control platform is what the Motion 4 ERP/MRP evaluation demonstrates.")
bullet("QuickBooks: use the SSLC access we already have, Jon's reports-only login, which already produced the GL detail through 10/12. No new Intuit app and no SSLC-side authorization step are needed for the evaluation.")
bullet("Email: $0. SSLC people log in with the addresses they already have, as Rufer and Conger do now. If SSLC wants branded addresses, Cloudflare Email Routing forwards them free. No paid mailboxes.")
bullet("Cash cost to Exos: $0. Build labor is Exo's time.")
bullet("Template: every future acquisition is a new container from the same image, created by one provisioning command.")

d.add_heading("1A. End state after acquisition (Jon, 19 Oct)", level=2)
bullet("Email: every SSLC employee gets an Exos (@exosaero.com) mailbox in the Exos Microsoft 365 GCC High tenant, the same as Exos staff. Seats would be added to the existing tenant. I couldn't verify a public GCC High per-seat price; it is quoted through the tenant provider (Atomus). Budget this as a post-close cost. The $0 email approach applies only to the pre-close evaluation.")
bullet("Mission Control: SSLC gets the full suite Exos runs, in its own section: work management (tasks, programs, Gantt, Decisions), ERP/MRP, inventory and receiving, suppliers and POs, accounting and finance, job costing, compliance sections (CMMC 2.0, AS9100 quality, export screening), the Process and Procedures library with Required Reading, opportunities, CRM, HR and assets.")
bullet("SSLC data stays in its own container even after close. The company is still one Exos company; the container is just how its records are kept separate. Exos leadership sees SSLC through the consolidated roll-up and named admin accounts.")
bullet("Compliance: putting SSLC staff on GCC High brings them into the Exos CMMC boundary. Scope it with Atomus before close (Jon is driving Atomus). SSLC's own process library starts as a copy of the Exos house template, then SSLC-specific procedures are added under the same CM rules.")
bullet("Keep this end state out of SSLC-facing documents. Post-close organization is Exos's decision alone; the SSLC demo shows the system, not the post-close structure.")

d.add_heading("2. What exists today", level=2)
bullet("Mission Control: server.js has 6,808 lines and 58 pages. It runs under launchd on port 8080, behind the exos-mc Cloudflare tunnel and an Access app restricted to email one-time PIN. The tunnel currently routes 4 hostnames (mc, vault, ssh, vnc).")
bullet("SSLC precedent, live as of today: the sslc_director role (Rufer, Conger) uses strict default-deny, reaches only /sslc-board.html, and sees only curated SSLC records. It works well for 2 read-only directors. It does not scale to 15 workers doing transactions, because every sandbox role is a hand-coded allowlist block in server.js and all of them share one database.")
bullet("ERP and MRP: Phase 5C MRP is live (demand from missions, purchase requests to PO drafts), along with Suppliers/POs, inventory, work orders, NCRs and assets. The QuickBooks link uses one credential set (one realm). SSLC QuickBooks is reached through Jon's reports-only login, not the API. That login is enough for the evaluation: the item, vendor, open PO, bill and inventory valuation reports can be exported into the SSLC database. The login session expires, so each data refresh needs a valid login.")
bullet("Users: 31 Mission Control accounts today. Cloudflare's free Zero Trust tier covers up to 50 users, then costs $7 per user per month on pay-as-you-go. A small SSLC demo audience fits in the Exos org. If SSLC adopts the system with about 15 users, move it to its own free org.", SRC_CF)
bullet("Docker/colima is already running on the Studio (Vaultwarden container), so the container runtime is in place.")
bullet("Containerizing work: server.js has 63 call sites to host tools (python3 scripts, the openclaw CLI, the Proton Bridge, /Users/exomac paths). In the company image these are either disabled (Exos-only integrations such as GCC High Graph, Exos Decisions and e-sign) or replaced with container-local equivalents.")
bullet("Hardware: the Mac Studio M3 Ultra (96 GB, 2 TB) is the primary host, and the Mac mini is the offsite DR node with daily backup and a failover drill. A second Node process plus a SQLite file is a negligible load; no new hardware is needed.")
bullet("SSLC staff currently use @smad.com addresses. Confirm who controls smad.com. If SSLC does not control it outright, SSLC needs its own domain for this.")

d.add_heading("3. Architecture options", level=2)
table([
    ["Option", "How", "Pros", "Cons", "Verdict"],
    ["A. SSLC role inside Exos MC", "Add an sslc_worker role, as with sslc_director; SSLC rows go in the shared mc.db", "Fastest (days)", "No entity column, so SSLC and Exos data mix in one database; leakage risk; Exos QB sync writes to the same DB; hard to separate on divestiture or a failed deal", "Reject for workers (keep for read-only board access)"],
    ["B. Entity column across the ERP", "Add entity_id to all 47 tables, filter every query", "One database, consolidated reporting", "Touches every module and query; a single missed filter is a breach; large regression risk to live Exos ops", "Reject now; revisit only for a full post-close merger"],
    ["C. Container per company (recommended)", "One Docker image of MC; the sslc container has its own volume (sslc.db, users, docs), port 8081, hostname sslc.exosaero.space and its own login policy", "Hard isolation by construction; Exos code improvements flow to every container; clean exit or merge; repeatable template", "Some up-front work to parameterize paths and the port; consolidated reporting needs a roll-up job", "Recommend"],
])

d.add_heading("4. Build plan: SSLC demo, then template", level=2)
table([
    ["Step", "Work", "Human step", "Elapsed"],
    ["1. MC image", "Dockerfile, plus ENTITY/PORT/DATA_DIR settings; Exos-only integrations switched off for company containers", "None", "2-3 days"],
    ["2. SSLC container", "Empty sslc.db from schema.sql (never migrate.js), SSLC users, volume, hostname sslc.exosaero.space, Access app and policy", "None", "1 day"],
    ["3. Load SSLC data", "Export items, vendors, open POs, bills and inventory valuation from SSLC QB through the existing login, then import into sslc.db", "Jon: refresh the SSLC QB login if the session has expired", "1-2 days"],
    ["4. Demo setup (full suite)", "All modules on SSLC data. ERP/MRP: tank BOMs from QB items, MRP run, PO-draft flow. Also inventory, receiving, suppliers, finance/AR-AP and job costing from QB reports, tasks and programs for the open tank orders, opportunities and CRM from SSLC customers, and starter compliance and process-library sections (house template, empty SSLC procedures). Plus a 1-page demo script", "None", "4-5 days"],
    ["5. Present", "Live walkthrough for the SSLC board, plus seats for whoever SSLC names", "Jon: schedule it", "-"],
    ["6. Template", "provision_entity command (new container, volume, hostname, Access policy, backup lane), plus a procedure on the house template", "Jon: approve the procedure", "1-2 days"],
])
p("Demo ready in about 1.5-2 weeks; Steps 1-4 need nothing from SSLC.")
bullet("Backups: the SSLC volume is added to the nightly asset backup and the Studio-to-Mini DR sync. Because the container is a self-contained unit, restore means copying the volume and starting the container.")

d.add_heading("5. Email: $0", level=2)
bullet("Logging in to MC does not need new mailboxes. The Cloudflare gate emails a one-time code to whichever address is on the approved list, so SSLC staff use the addresses they have today.")
bullet("SSLC-branded addresses, if wanted, are free with Cloudflare Email Routing, which forwards custom-domain addresses to existing inboxes. It is receive-only: replies go out from the person's current address.", ("cloudflare.com, Email Routing", "https://www.cloudflare.com/developer-platform/products/email-routing"))
bullet("Paid mailboxes are needed only if staff must send as the SSLC domain. Then buy seats only for those people. Proton business is billed per user, and free Proton accounts cannot use a company domain.", ("proton.me/mail", "https://proton.me/mail"))
bullet("Never give multiple people one shared login, and never self-host mail on the Studio (residential IPs land in spam).")
bullet("Export control: no SSLC USML technical data goes in email or the container until the Empowered Official determines its status. This also keeps the Studio out of CMMC scope.", SRC_ITAR)

d.add_heading("6. Cost summary", level=2)
table([
    ["Item", "Cash cost", "Basis"],
    ["Hardware and hosting", "$0", "Existing Mac Studio + Mini DR"],
    ["Container runtime", "$0", "colima/Docker already running"],
    ["Cloudflare Access", "$0", "Free up to 50 users"],
    ["QuickBooks", "$0", "Existing SSLC reports-only login"],
    ["Email (evaluation)", "$0", "Existing addresses + free forwarding"],
    ["Email (post-close)", "GCC High seats x 15", "Couldn't verify public price; quote via Atomus"],
    ["Labor", "Exo time", "About 1.5-2 weeks to demo"],
])

d.add_heading("7. Governance (because Exos is related-party and SSLC is not a subsidiary)", level=2)
bullet("Scope: the demo stays inside the approved ERP/MRP/inventory evaluation at no cost to SSLC. No cash goes to SSLC and no paid services are provisioned for it.")
bullet("Evaluation letter (one page, from SSLC to Exos): SSLC owns its data and its domain; Exos provides hosting and admin as a service; an end date or 90-day review; on termination, Exos exports the data (QuickBooks remains SSLC's system of record) and deletes the container after SSLC confirms the export.")
bullet("Disclosure: Exos-operated AI agents (Exo) will process SSLC data to run the system. State this in the letter.")
bullet("Exos-side visibility: only named Exos evaluators get accounts in the SSLC container, and their access is logged like the sslc_director activity log. No SSLC data flows into Exos MC, decks or analyses without SSLC board consent. This follows the same rule already applied to the director portal.")
bullet("Keep the letter limited to the evaluation. Post-transaction structure stays out of SSLC documents.")

d.add_heading("8. Acquisition template: what each new company gets", level=2)
for s in [
    "A container from the shared MC image, with its own volume (database, users, documents), port and ENTITY setting.",
    "A hostname on the Exos tunnel and its own Access policy. Move it to its own free Cloudflare org once its users plus Exos users approach 50.",
    "Email: before close, existing addresses plus free forwarding. After close, Exos GCC High mailboxes for all staff.",
    "QuickBooks: start with the access we have, then move to an API connection (namespaced credentials) once the company adopts the system.",
    "A backup lane in the nightly asset backup, the Studio-to-Mini DR sync and the weekly failover drill.",
    "A consolidated roll-up: a read-only nightly job that copies agreed KPIs (revenue, backlog, inventory value, open POs) from each container to an Exos-only portfolio page, with no raw records crossing over.",
    "The full MC suite in the company's own section, with compliance, process library, opportunities and work management, starting from the Exos house template.",
    "An exit or merge path: export or hand over the container to divest, or run a planned merge into the Exos database to integrate. Both are possible because the container was never mixed with Exos data.",
    "One command (provision_entity) plus a house-template procedure. Target: a new company's container is live within 1 day of QuickBooks access.",
]:
    bullet(s)

d.add_heading("9. Next actions", level=2)
bullet("Exo: build Steps 1-4 (image, SSLC container, data load, demo setup). Nothing changes for current Exos MC users.")
bullet("Jon: refresh the SSLC QB login if Exo reports the session has expired, and schedule the board demo when Exo reports it is ready.")

d.save(OUT)
print(OUT)
