#!/usr/bin/env python3
"""IDJ Oct 8 deck v7: Jon's v6_JQ (9/28) + Intro/Outro videos, condensed 8 -> 5 static slides.
Re-run with INTRO=/OUTRO= env to swap in new video cuts (e.g. v3 no-smoke)."""
import copy, os, re
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

ROOT = '/Users/exomac/.openclaw/workspace'
D = f'{ROOT}/drafts/idj_deck_v7'
SRC = f'{ROOT}/mission-control/uploads/1790638738245_Exos_JP_3min_InvestorPitch_IDJ_BRANDED_v6_2026-09-28_JQ.pptx'
OUT = os.environ.get('OUT', f'{ROOT}/drafts/Exos_JP_3min_InvestorPitch_IDJ_v7_VIDEO_2026-10-05.pptx')
INTRO = os.environ.get('INTRO', f'{D}/Intro_embed.mp4')
OUTRO = os.environ.get('OUTRO', f'{D}/Outro_embed.mp4')
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

prs = Presentation(SRC)
S = list(prs.slides)
title, who, apart, proof, why, invest, traction, invite = S


def set_text(slide, old_re, new):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if re.search(old_re, r.text):
                    r.text = re.sub(old_re, new, r.text)
                    return True
    return False


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


# --- renumber section labels + page numbers after removing Title / Proof / Invitation
lab = shape(why, 'TextBox 1').text_frame.paragraphs[0].runs   # Jon's edit split "04" into runs '0','4'
lab[1].text = '3' if lab[1].text == '4' else lab[1].text
set_text(invest, r'^05', '04'); set_text(traction, r'^06', '05')
for sl, n in ((why, '04'), (invest, '05'), (traction, '06')):
    for sh in sl.shapes:
        if sh.has_text_frame and re.fullmatch(r'0\d', sh.text_frame.text.strip()) and sh.left > Emu(11 * 914400):
            sh.text_frame.paragraphs[0].runs[0].text = n

# --- fold "The Proof" takeaway into "What sets us apart"
tk = copy.deepcopy(shape(who, 'TextBox 18')._element)   # same style as slide-2 takeaway line
apart.shapes._spTree.append(tk)
tkb = apart.shapes[-1]
tkb.top = Emu(int(5.80 * 914400))
tkb.text_frame.paragraphs[0].runs[0].text = 'Demonstrated hardware and mission heritage — not a promise of future technology.'
for r in tkb.text_frame.paragraphs[0].runs[1:]:
    r.text = ''

# --- carry the Invitation contacts onto Traction (outro video now carries "build it together / thank you")
for nm, top in (('TextBox 9', 6.30), ('TextBox 10', 6.68)):
    el = copy.deepcopy(shape(invite, nm)._element)
    traction.shapes._spTree.append(el)
    b = traction.shapes[-1]
    b.left, b.top, b.width = Emu(int(0.65 * 914400)), Emu(int(top * 914400)), Emu(int(12.0 * 914400))
b.text_frame.paragraphs[0].runs[0].text = 'Toshi Kumazawa  •  toshihisa@exosaero.com      |      John Quinn  •  Co-founder & CEO'
for r in b.text_frame.paragraphs[0].runs[1:]:
    r.text = ''


# --- video slides
def video_slide(path, poster, label):
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else title.slide_layout
    sl = prs.slides.add_slide(title.slide_layout)
    for ph in list(sl.placeholders):
        ph._element.getparent().remove(ph._element)
    bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(0, 0, 0)
    mv = sl.shapes.add_movie(path, 0, 0, prs.slide_width, prs.slide_height,
                             poster_frame_image=poster, mime_type='video/mp4')
    mv.name = label
    sp_id = mv.shape_id
    # replace python-pptx click-to-play timing with autoplay-on-slide-enter
    for t in sl._element.findall('p:timing', NS):
        sl._element.remove(t)
    timing = etree.fromstring(f'''
<p:timing xmlns:p="{NS['p']}">
 <p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
  <p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
   <p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
    <p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
     <p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
      <p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{sp_id}"/></p:tgtEl></p:cBhvr></p:cmd>
     </p:childTnLst></p:cTn></p:par>
    </p:childTnLst></p:cTn></p:par>
   </p:childTnLst></p:cTn></p:par>
  </p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
  <p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sp_id}"/></p:tgtEl></p:cMediaNode></p:video>
 </p:childTnLst></p:cTn></p:par></p:tnLst>
</p:timing>''')
    sl._element.append(timing)
    return sl


intro = video_slide(INTRO, f'{D}/intro_poster.jpg', 'Intro video')
outro = video_slide(OUTRO, f'{D}/outro_poster.jpg', 'Outro video')
intro.notes_slide.notes_text_frame.text = ('INTRO VIDEO (26 s, autoplays with sound). Carries the brand open + reusable / suborbital / flown / '
    'recovered / BLK3 story and the Innovation Dojo card - so the title slide and the "Proof" slide were cut. '
    'Speak over nothing; start talking on slide 2. Presenters: Toshi Kumazawa (Japan BD) · John Quinn (Co-founder & CEO).')
outro.notes_slide.notes_text_frame.text = ('OUTRO VIDEO (13 s, autoplays). "Let\'s build it together" + Thank you / ありがとうございました - '
    'replaces the Invitation slide; contacts moved to the Traction slide.')

# --- reorder / drop: Intro, Who, Apart, Why, Invest, Traction, Outro
lst = prs.slides._sldIdLst
el_of = {id(s.part): el for s, el in zip(prs.slides, list(lst))}
keep = [intro, who, apart, why, invest, traction, outro]
drop = [title, proof, invite]
keep_els = [el_of[id(s.part)] for s in keep]
drop_rids = [el_of[id(s.part)].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id') for s in drop]
for el in list(lst):
    lst.remove(el)
for el in keep_els:
    lst.append(el)
for rid in drop_rids:
    prs.part.drop_rel(rid)

prs.save(OUT)
print(OUT, os.path.getsize(OUT) // 1024, 'KB')
