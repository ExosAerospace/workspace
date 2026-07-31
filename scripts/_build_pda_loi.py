#!/usr/bin/env python3
"""Build the PD Aerospace Type-V tank LOI (v2) on PRIMARY letterhead.

v2 per Neil Milburn (2026-07-24): mutual buy/provide framing (PD Aerospace to buy, EXOS to
provide TRL-9, space-rated Type-V composite propellant tanks); full spec + price table
populated from the revised SPACEPLANE table; both parties' addresses; dual execution blocks
(EXOS + PD Aerospace) with editable signature / name-title / date-executed fields. Non-binding.
No CUI/ITAR-sensitive content (confirmed by Neil).
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx"
DST = "drafts/PD_Aerospace_Type5_Tank_LOI.docx"
NAVY = RGBColor(0x0B, 0x22, 0x40)

doc = Document(SRC)


def para(text="", bold=False, size=11, align=None, color=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Pt(18)
    if bold_lead:
        r = p.add_run("•  ")
        r.font.size = Pt(11)
        rb = p.add_run(bold_lead)
        rb.bold = True
        rb.font.size = Pt(11)
        r2 = p.add_run(text)
        r2.font.size = Pt(11)
    else:
        r = p.add_run("•  " + text)
        r.font.size = Pt(11)
    return p


# --- Header block ---
para("LETTER OF INTENT", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, space_after=2)
para("TRL-9, Space-Rated Type-V Composite Propellant Tanks — PD Aerospace Spaceplane Program",
     bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

para("Date:  July 24, 2026", size=11, space_after=2)
para("Buyer:  PD Aerospace Ltd. — 3519 Arimatsu, Midori-ku, Nagoya, Aichi 458-0924, Japan",
     size=11, space_after=2)
para("Attn:  [PD Aerospace Contact / Title]", size=11, space_after=2)
para("Via:  Toshihisa Kumazawa, EXOS Aerospace (Japan Representative)", size=11, space_after=2)
para("Provider:  EXOS Aerospace Systems & Technologies, Inc. — 299 FM 1903, Greenville, TX 75402, USA "
     "(EXOS–SSLC production program)", size=11, space_after=2)
para("Reference:  Preliminary Quote PQ260717", size=11, space_after=10)

# --- Intro ---
para("This Letter of Intent (“LOI”) sets out the mutual intent of PD Aerospace Ltd. (“PDA”) to buy, "
     "and of EXOS Aerospace Systems & Technologies, Inc. (“EXOS”), through its EXOS–SSLC production "
     "program, to provide TRL-9, space-rated Type-V composite propellant tanks for the PDA spaceplane, "
     "in accordance with the terms below, and for EXOS to reserve a production slot for PDA within the "
     "EXOS–SSLC program. This LOI is non-binding and is provided to confirm mutual intent and to enable "
     "both parties to plan capacity and program schedule ahead of a definitive purchase order.",
     size=11, space_after=8)

# --- 1. Scope of Supply ---
heading("1.  Scope of Supply")
para("EXOS will provide the following TRL-9, space-rated Type-V composite propellant tanks for the PDA "
     "spaceplane. All tank prices include NASA flight cleaning for particulates and non-volatile residue "
     "(NVR). Prices are ex-works and exclusive of packing and shipping.", size=11, space_after=6)

rows = [
    ("Tank", "Capacity", "Diam.", "Length", "MAWP (2.0 SF)", "Mass", "Unit Price (USD)"),
    ("LOX Tank\n25-AM-124-725 LOX", "950 L / 34 SCF", "630 mm / 25\"", "3150 mm / 124\"",
     "5.0 MPa / 725 psi", "55 kg / 120 lbm", "$400,488"),
    ("GHe Tank\n25-AM-76-4000 GHe", "570 L / 20 SCF", "630 mm / 25\"", "1930 mm / 76\"",
     "27.6 MPa / 4000 psi", "180 kg / 400 lbm", "$395,160"),
]
table = doc.add_table(rows=len(rows), cols=7)
try:
    table.style = "Table Grid"
except KeyError:
    pass
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = table.cell(ri, ci)
        cell.text = ""
        pr = cell.paragraphs[0]
        run = pr.add_run(val)
        run.font.size = Pt(8)
        if ri == 0:
            run.bold = True
            run.font.color.rgb = NAVY
        elif ci == 0:
            run.bold = True
para("")
para("Total for one (1) LOX tank and one (1) GHe tank: $795,648 USD, ex-works. Pricing basis reflects "
     "EXOS–SSLC’s current knowledge from orders on hand for lunar lander propellant tanks (NASA CLPS).",
     size=9.5, space_after=8, italic=True)

# --- 2. Delivery ---
heading("2.  Delivery Schedule")
para("Delivery is ex-works, twenty-four (24) weeks after receipt of order (ARO). Expedited delivery "
     "may be possible using a multi-shift production system, at a premium cost to be advised at the "
     "time of order. Where PDA orders multiple tanks for uninterrupted production, a volume/continuity "
     "discount may be available for a schedule in which each successive tank is started as early as "
     "three (3) weeks after the start of the prior tank.", size=11, space_after=8)

# --- 3. Payment ---
heading("3.  Payment Schedule")
para("Payment terms for delivery ex-works, twenty-four (24) weeks ARO:", size=11, space_after=4)
bullet("of order total with purchase order (PO)", bold_lead="25% ")
bullet("of order total at 8 weeks", bold_lead="25% ")
bullet("of order total at 16 weeks", bold_lead="25% ")
bullet("of order total on notification ready to ship, prior to shipment", bold_lead="25% ")
para("")

# --- 4. Pricing Validity & Raw-Material Escalation ---
heading("4.  Pricing Validity and Raw-Material Escalation")
para("Pricing in Section 1 is firm for a purchase order issued within [90] days of the date a "
     "definitive quotation is provided. Because the parties anticipate a purchase order no earlier "
     "than mid-2027, and given known volatility in the aerospace-grade carbon fiber and composite raw "
     "material market, pricing for any PO issued after that validity window is subject to re-quotation "
     "to reflect then-current raw-material costs. EXOS will hold pricing to the maximum extent "
     "supported by its supply chain and will provide reasonable advance notice of any adjustment.",
     size=11, space_after=8)

# --- 5. Production Slot Reservation ---
heading("5.  Production Slot Reservation")
para("Issuance of a purchase order against this LOI secures PDA’s slot in the EXOS–SSLC "
     "production program. EXOS will confirm the scheduled delivery window upon receipt of the PO and "
     "initial payment. The parties presently anticipate a target order window of mid-2027, subject to "
     "PDA’s spaceplane design progress.", size=11, space_after=8)

# --- 6. Export Control ---
heading("6.  Export Control and Regulatory Compliance")
para("The tanks and any associated technical data may be subject to U.S. export control laws and "
     "regulations, including the International Traffic in Arms Regulations (ITAR) and/or the Export "
     "Administration Regulations (EAR). Performance of any resulting purchase order, and any transfer "
     "of hardware or technical data to PDA, is expressly conditioned on EXOS obtaining all required "
     "U.S. Government export authorizations. The parties will cooperate in good faith to secure such "
     "authorizations. This LOI does not authorize the export of any controlled item or technical data.",
     size=11, space_after=8)

# --- 7. Basis, Confidentiality, Non-Binding ---
heading("7.  Basis and Nature of this LOI")
para("This LOI reflects EXOS–SSLC’s current production knowledge and order book and is "
     "provided to establish intent and reserve capacity. Final pricing, quantities, and technical "
     "scope will be confirmed in the definitive purchase order and accompanying specification. Except "
     "for the provisions on Export Control (Section 6) and Confidentiality, this LOI is non-binding "
     "and creates no obligation on either party to enter into a definitive agreement. Each party will "
     "treat the terms of this LOI and any information exchanged as confidential. This LOI is valid for "
     "[90] days from the date above unless extended in writing by the parties.", size=11, space_after=14)

# --- Dual execution ---
para("IN WITNESS WHEREOF, the parties confirm their intent and execute this Letter of Intent as of the "
     "dates set forth below.", size=11, space_after=16)

# EXOS block
heading("For and on behalf of EXOS Aerospace Systems & Technologies, Inc. (Provider)")
para("Signature: ______________________________", size=11, space_after=10)
para("Name / Title: ___________________________", size=11, space_after=10)
para("Date Executed: __________________________", size=11, space_after=16)

# PD Aerospace block
heading("For and on behalf of PD Aerospace Ltd. (Buyer)")
para("Signature: ______________________________", size=11, space_after=10)
para("Name / Title: ___________________________", size=11, space_after=10)
para("Date Executed: __________________________", size=11, space_after=2)

doc.save(DST)
print("wrote", DST)
