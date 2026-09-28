#!/usr/bin/env python3
"""Build Attachment A to the EXOS Employment Agreement — Confidential Information
and Invention Assignment Agreement (CIIAA), Rev. 2026-09-28, as .docx.

Amends and restates the prior Attachment A to add the carve-outs the 9/28
whistleblower review found missing: FAR 52.203-19 (no confidentiality terms that
restrict reporting waste, fraud or abuse), 10 USC 4701 / 41 USC 4712 protected
activity, and the DTSA 18 USC 1833(b) immunity notice. Adds an annual
re-certification (signed yearly via MC Required Reading, doc id hr-ciiaa).
"""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = sys.argv[1] if len(sys.argv) > 1 else \
    "process-library/05-human-resources/EXOS_Attachment_A_CIIAA_Rev_2026-09-28.docx"
CO = "EXOS Aerospace Systems & Technologies, Inc."
REPORT_URL = "https://mc.exosaero.space/report.html"

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def h(t, lvl=1):
    doc.add_heading(t, level=lvl)


def p(t, bold=False, italic=False):
    para = doc.add_paragraph()
    r = para.add_run(t)
    r.bold, r.italic = bold, italic
    return para


def ul(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


t = doc.add_heading(f"{CO}\nAttachment A to Employment Agreement\n"
                    "Confidential Information and Invention Assignment Agreement", 0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Revision 2026-09-28 — amends and restates the Attachment A previously executed with the Employment Agreement",
  italic=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Revision summary: adds Section 6 (Protected Activity — FAR 52.203-19, 10 U.S.C. § 4701, 41 U.S.C. § 4712), "
  "Section 7 (Defend Trade Secrets Act immunity notice, 18 U.S.C. § 1833(b)) and Section 12 (Annual "
  "Certification). All confidentiality, return-of-property and invention-assignment obligations are made "
  "expressly subject to Sections 6 and 7.", italic=True)

p(f"In consideration of my employment or continued employment by {CO} (together with its successors and "
  "assigns, the \"Company\"), my access to the Company's Confidential Information, and the compensation I "
  "receive, I agree as follows:")

# 1 ------------------------------------------------------------------
h("1. Confidential Information")
p("1.1 Definition. \"Confidential Information\" means any non-public information, in any form, that I learn or "
  "develop in connection with my employment and that relates to the Company's actual or anticipated business, "
  "research or development, including: technical data, designs, drawings, specifications, test data, "
  "processes, formulas, software, source code, know-how, and Inventions (defined below); vehicle, propulsion, "
  "cryogenic-tank and composite technology; business plans, financial information, pricing, costs, "
  "capitalization and investor information; customer, supplier, partner and payload-customer information; "
  "personnel information (other than my own terms and conditions of employment); and information the Company "
  "receives from any third party under an obligation of confidentiality.")
p("1.2 Exclusions. Confidential Information does not include information that (a) is or becomes publicly "
  "available through no wrongful act of mine, (b) I can show I knew before my employment without a "
  "confidentiality obligation, or (c) I lawfully receive from a third party without a confidentiality "
  "obligation.")
p("1.3 Nondisclosure and Non-Use. Subject to Sections 6 and 7, during and after my employment I will hold all "
  "Confidential Information in strict confidence, will not disclose it to anyone outside the Company except as "
  "authorized by the Company in the course of my duties, and will not use it except for the Company's benefit.")
p("1.4 Third-Party and Government Information. I will protect information entrusted to the Company by "
  "customers, partners and the U.S. Government (including Controlled Unclassified Information (\"CUI\")) in "
  "accordance with the applicable agreement, contract clause (including DFARS 252.204-7012) and Company "
  "procedure, and will access it only on a need-to-know basis.")
p("1.5 Export-Controlled Information. I will not export, re-export, release or transfer any technical data, "
  "software or technology controlled under the International Traffic in Arms Regulations (ITAR, 22 C.F.R. "
  "Parts 120-130) or the Export Administration Regulations (EAR, 15 C.F.R. Parts 730-774), including by "
  "release to a foreign person in the United States, except as authorized by the Company's export-compliance "
  "process.")
p("1.6 Former Employer Information. I will not bring to the Company, or use in my work, any confidential or "
  "proprietary information or trade secrets of any former employer or other person without their written "
  "authorization.")

# 2 ------------------------------------------------------------------
h("2. Inventions")
p("2.1 Definition. \"Inventions\" means discoveries, inventions, designs, improvements, developments, works of "
  "authorship, software, data, processes, techniques, know-how and trade secrets, whether or not patentable or "
  "registrable, that I conceive, reduce to practice or develop, alone or with others, during my employment.")
p("2.2 Prior Inventions. I have listed on Exhibit 1 all Inventions I made before my employment that I own or "
  "have an interest in and that relate to the Company's business (\"Prior Inventions\"). If none are listed, I "
  "represent there are none. If I incorporate a Prior Invention into Company work, I grant the Company a "
  "non-exclusive, royalty-free, irrevocable, perpetual, worldwide license to make, use, sell and otherwise "
  "exploit it as part of that work.")
p("2.3 Assignment. I agree to promptly disclose to the Company, and I hereby assign to the Company, all my "
  "right, title and interest in any Invention that (a) relates to the Company's actual or demonstrably "
  "anticipated business, research or development, (b) results from work I perform for the Company, or "
  "(c) is made using the Company's equipment, supplies, facilities, Confidential Information or time "
  "(\"Company Inventions\"). Copyrightable works I create within the scope of my employment are \"works made "
  "for hire\" under the U.S. Copyright Act.")
p("2.4 Excluded Inventions. This Agreement does not apply to an Invention that I develop entirely on my own "
  "time, without using the Company's equipment, supplies, facilities, Confidential Information or trade "
  "secrets, and that neither relates to the Company's business or actual or demonstrably anticipated research "
  "or development nor results from any work I perform for the Company.")
p("2.5 Government Contract Rights. Where a Company Invention is made in the performance of a U.S. Government "
  "contract, grant or agreement, I will cooperate with the Company in meeting any disclosure, election and "
  "licensing obligations to the Government (for example, under 35 U.S.C. §§ 200-212 and FAR 52.227-11 or "
  "DFARS 252.227-7038), and the Company's assignment rights are subject to the Government's rights.")
p("2.6 Further Assurances. During and after my employment I will assist the Company, at its expense, to "
  "obtain, perfect and enforce its rights in Company Inventions, including by signing applications and "
  "assignments. If I am unavailable or unwilling to sign, I appoint the Company as my attorney-in-fact solely "
  "to sign such documents on my behalf.")
p("2.7 Records. I will keep adequate, current written records of Company Inventions, which remain Company "
  "property.")

# 3 ------------------------------------------------------------------
h("3. Return of Company Property")
p("When my employment ends, or earlier on request, I will return all Company property, documents, data and "
  "devices, and all copies of Confidential Information in any form, and will permanently delete Confidential "
  "Information from any personal device or account after returning a copy to the Company. CUI and "
  "export-controlled information will be returned or destroyed only as directed by the Company's security "
  "procedures. This Section is subject to Sections 6 and 7.")

# 4 ------------------------------------------------------------------
h("4. No Conflicting Obligations")
p("I represent that my performance under this Agreement does not and will not breach any agreement to keep in "
  "confidence information acquired before my employment, and that I am not bound by any agreement that "
  "conflicts with this Agreement.")

# 5 ------------------------------------------------------------------
h("5. Notice of Other Engagements")
p("During my employment I will notify the Company in writing before engaging in any business activity that "
  "competes with, or relates to, the Company's business, consistent with Section 1(B) of my Employment "
  "Agreement.")

# 6 ------------------------------------------------------------------
h("6. Protected Activity — Nothing in This Agreement Restricts Reporting")
p("6.1 Protected Activity. Nothing in this Agreement, the Employment Agreement, the Employee Handbook, or any "
  "other Company agreement or policy prohibits, restricts, or requires Company permission or notice before I:")
ul([
    "report waste, fraud, abuse, gross mismanagement, a substantial and specific danger to public health or "
    "safety, or a violation of law, rule or regulation related to a Federal contract or grant, to a Member of "
    "Congress or congressional committee; an Inspector General (including the DoD and NASA Offices of "
    "Inspector General and their hotlines); the Government Accountability Office; a Federal employee "
    "responsible for contract oversight or management (including a Contracting Officer); an authorized "
    "official of the Department of Justice or other law-enforcement agency; a court or grand jury; or a "
    "management official or other Company employee with responsibility to investigate, discover or address "
    "misconduct (10 U.S.C. § 4701; 41 U.S.C. § 4712);",
    "file a charge or complaint with, communicate with, provide documents or information to, or participate "
    "in an investigation or proceeding by any Federal, state or local government agency, including the "
    "Securities and Exchange Commission, Equal Employment Opportunity Commission, National Labor Relations "
    "Board, Occupational Safety and Health Administration, and Federal Aviation Administration;",
    "receive an award from a government agency for information provided;",
    "testify truthfully in any legal proceeding; or",
    "discuss or disclose wages, hours or other terms and conditions of employment, or otherwise engage in "
    "concerted activity protected by Section 7 of the National Labor Relations Act.",
])
p("6.2 FAR 52.203-19. The Company does not require, and will not enforce, any confidentiality agreement or "
  "statement that prohibits or restricts employees from lawfully reporting waste, fraud or abuse related to "
  "the performance of a Government contract to a designated investigative or law-enforcement representative "
  "of a Federal department or agency authorized to receive such information. Any prior version of this "
  "Attachment A, or any other Company confidentiality obligation, is superseded and limited to the extent "
  "it could be read otherwise.")
p("6.3 No Retaliation. The Company will not discharge, demote, or otherwise discriminate against me as a "
  "reprisal for Protected Activity. I may report reprisal to the Inspector General of the relevant agency "
  "within three years of the reprisal. Section 2.5 of the Employee Handbook describes my rights and the "
  f"internal reporting channel (\"Report a Concern,\" {REPORT_URL}), which I may use but am not required to "
  "use before reporting externally.")
p("6.4 Handling Controlled Information While Reporting. Protected Activity does not authorize disclosure of "
  "classified information, CUI or export-controlled technical data outside authorized channels. When a "
  "report involves such information I may still make it: I will describe it generally and provide the "
  "controlled information only to an authorized recipient (for example, the relevant Inspector General or "
  "Contracting Officer) through a channel approved for that information. This Section 6 does not authorize "
  "disclosure of information protected by the Company's attorney-client privilege, except as permitted by "
  "law.")

# 7 ------------------------------------------------------------------
h("7. Defend Trade Secrets Act Immunity Notice (18 U.S.C. § 1833(b))")
p("Pursuant to the Defend Trade Secrets Act of 2016, I will not be held criminally or civilly liable under any "
  "Federal or state trade secret law for the disclosure of a trade secret that (a) is made (i) in confidence to "
  "a Federal, state or local government official, either directly or indirectly, or to an attorney, and "
  "(ii) solely for the purpose of reporting or investigating a suspected violation of law; or (b) is made in a "
  "complaint or other document filed in a lawsuit or other proceeding, if such filing is made under seal. If I "
  "file a lawsuit for retaliation by the Company for reporting a suspected violation of law, I may disclose the "
  "trade secret to my attorney and use the trade secret information in the court proceeding, if I file any "
  "document containing the trade secret under seal and do not disclose the trade secret except pursuant to "
  "court order.")

# 8 ------------------------------------------------------------------
h("8. At-Will Employment")
p("This Agreement does not create any right to continued employment. My employment remains at-will as stated "
  "in Section 5(A) of my Employment Agreement.")

# 9 ------------------------------------------------------------------
h("9. Remedies")
p("I acknowledge that a breach of Sections 1, 2 or 3 may cause the Company irreparable harm, and that the "
  "Company may seek injunctive relief in addition to any other remedy, without posting a bond to the extent "
  "permitted by law. No remedy may be sought for, and no breach arises from, Protected Activity under "
  "Section 6 or a disclosure immunized under Section 7.")

# 10 -----------------------------------------------------------------
h("10. Survival")
p("My obligations under Sections 1, 2, 3, 6, 7 and 9 survive the end of my employment, regardless of the "
  "reason.")

# 11 -----------------------------------------------------------------
h("11. General Provisions")
p("11.1 Governing Law. This Agreement is governed by the laws of the State of Texas, without regard to its "
  "choice-of-law provisions, and by applicable Federal law.")
p("11.2 Severability. If any provision is held invalid or unenforceable, it will be modified to the minimum "
  "extent necessary to make it enforceable, and the remaining provisions remain in effect. Sections 6 and 7 "
  "control over any conflicting provision.")
p("11.3 Entire Agreement; Amendment and Restatement. This Agreement, together with the Employment Agreement, "
  "is the entire agreement on its subject matter and amends and restates any prior Attachment A or "
  "Confidential Information and Invention Assignment Agreement between me and the Company. Assignments of "
  "Inventions made under a prior version remain effective. It may be modified only in a writing signed by me "
  "and an authorized officer of the Company, except that the Company may unilaterally expand the Protected "
  "Activity rights in Sections 6 and 7.")
p("11.4 Successors and Assigns. This Agreement binds my heirs and legal representatives and benefits the "
  "Company and its successors and assigns.")
p("11.5 Electronic Signature. This Agreement and each annual certification may be signed electronically, "
  "including through the Company's Mission Control Required Reading system, and an electronic signature has "
  "the same effect as a handwritten signature.")

# 12 -----------------------------------------------------------------
h("12. Annual Certification")
p("I will re-read and re-certify this Agreement each year of my employment. By each annual signature I "
  "certify that, during the prior year and to the best of my knowledge:")
ul([
    "I have complied with this Agreement, including its confidentiality, CUI and export-control obligations;",
    "I have disclosed to the Company all Company Inventions I made during the period;",
    "I have not retained Company property or Confidential Information outside Company-approved systems and "
    "devices;",
    "Exhibit 1 (Prior Inventions) remains accurate or I have provided an update; and",
    "I understand my Protected Activity rights under Sections 6 and 7.",
])

# Signature ----------------------------------------------------------
h("Employee Acknowledgment and Signature")
p("I have read this Agreement carefully, understand its terms, and sign it voluntarily. I understand that "
  "nothing in it restricts my right to report possible violations of law to the government.")
for line in ["Employee Signature: ______________________________",
             "Printed Name: ______________________________",
             "Date: ______________________________",
             "",
             f"{CO}",
             "By: ______________________________",
             "Name: John Quinn",
             "Title: Cofounder and CEO",
             "Date: ______________________________"]:
    p(line)

# Exhibit 1 ----------------------------------------------------------
doc.add_page_break()
h("Exhibit 1 — List of Prior Inventions")
p("The following is a complete list of all Prior Inventions (Section 2.2) that relate to the Company's business. "
  "Do not disclose any third party's confidential information; a general description is sufficient.")
tbl = doc.add_table(rows=6, cols=3)
tbl.style = "Table Grid"
for i, hdr in enumerate(["Title", "Date", "Identifying Number or Brief Description"]):
    tbl.rows[0].cells[i].text = hdr
p("")
p("☐ No Prior Inventions.        ☐ Additional sheets attached.")
p("Employee Signature: ______________________________    Date: ____________")

doc.save(OUT)
print(OUT)
