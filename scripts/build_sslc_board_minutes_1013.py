"""Exos record of the SSLC board meeting, Oct 13, 2026 (from transcript). Internal capture, not official SSLC minutes."""
import subprocess
from docx import Document
from docx.shared import Pt, RGBColor
WS="/Users/exomac/.openclaw/workspace"
OUT=f"{WS}/drafts/SSLC_Board_Meeting_2026-10-13_Exos_Record.docx"
NAVY=RGBColor(0x0A,0x3D,0x7C); GREY=RGBColor(0x55,0x55,0x55)
d=Document(); d.styles["Normal"].font.name="Calibri"; d.styles["Normal"].font.size=Pt(10.5)
def P(t,b=False,s=10.5,c=None,i=False,sa=4):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(sa); r=p.add_run(t); r.bold=b; r.italic=i; r.font.size=Pt(s)
    if c: r.font.color.rgb=c
def H(t): P(t,True,12,NAVY,sa=3)
def B(t): P("•  "+t,sa=2)
P("SSLC Board of Directors Meeting — October 13, 2026 (Exos record)",True,15,NAVY)
P("Exos's record from the meeting transcript (data-room/SSLC/meetings/SSLC_Board_Meeting_2026-10-13_transcript.txt). Not SSLC's official minutes. Confidential.",i=True,s=9,c=GREY)
P("Present: Jim Wertz (chair), Alice Wertz, Bob Conger, Markus Rufer, John Quinn (from Japan). Advisor: Danny Piper (advisor to SSLC and Microcosm on the acquisition; not a director). Board meeting closed 12:27; discussion continued afterward.",s=10)
H("Board actions")
B("Motion 1 (approve offer + §3 exclusivity): DEFERRED to the next meeting. Jim: his vote on the overall agreement is currently no; SSLC will send a counteroffer.")
B("Motion 2 (full read-only QuickBooks, or equivalent, access for all directors): APPROVED. Exos to make a Mission Control query agent available to directors within a few days.")
B("Motion 3 (correct AR Unbilled / customer-deposit accounting): APPROVED, 5 business days, within the current framework. An existing industry method (Exos/BBSNV deposit-to-delivery bookings) is the template for A La Carte; Markus supplies shop-floor completion data (1–2 days).")
B("Motion 4 (ERP/MRP/inventory evaluation using SSLC QuickBooks data): APPROVED, at no cost to SSLC.")
B("Motion 5 (related-party balances): DEFERRED; SSLC to review the transactions and respond in the next couple of days.")
H("Danny Piper's recommendations (basis for Heads of Terms r2)")
for t in ["Exchange a Word heads-of-terms with the headline economics; lock material terms once signed, because the financing/SEC filings cannot keep changing. Details go in the definitive agreements.",
 "Asset sale (not stock): needs board approval plus a majority of shareholders, not every holder; SSLC remains responsible to its own cap table.",
 "Non-competes only from directors and Markus (parties receiving value), about 5 years; carve-outs for the Microcosm and LOXair licenses (LOXair to be re-papered with Exos). Handle side agreements carefully, with counsel, so shareholders see them.",
 "Post-closing tax: if Exos steps up asset basis and it creates SSLC tax, Exos covers it (true-up). Add purchase-price-adjustment language (about two sentences each).",
 "$60K transaction costs are part of the purchase price; each side pays its own lawyers (Exos cannot pay SSLC's counsel). Consider paying it at term-sheet signing to fund retainers. Ed Powell (prior SSLC counsel) as a cost-effective option.",
 "Employees: no 4–6 year contracts; an assurance of continued employment offers is enough. At-will with for-cause terms in the definitive agreement.",
 "Post-closing operations (organization, facilities) are the buyer's decision; SSLC does not set terms for how Exos runs the business after closing.",

 "Prioritize agreeing the term sheet over the accounting cleanup. Compare SSLC's counteroffer (payments tied to Exos funding events, no performance requirements) against Exos's (cash + earn-outs + stock)."]: B(t)
H("Other points")
B("Jim's counteroffer: larger payments, timed after Exos's financing events (he cited May and August); to be written up and sent with comments from Danny, Bob and Markus.")
B("Jim requested a black-on-white print copy; sent during the meeting.")
B("John: the 'invested 25+ years ago' note on the payout slide is unverified and should come out.")
B("John on urgency: the fairness opinion and the RTO process cannot proceed without a signed agreement; the financing team costs about $15K/week.")
H("Action items")
for t in ["John/Exos: send the Word term sheet with current numbers (Heads of Terms r2) to Danny for redline.",
 "Exos: stand up director query access to SSLC QuickBooks data (Motion 2).",
 "Exos/BBSNV: share the deposit-to-delivery booking example for A La Carte (Motion 3).",
 "SSLC (Jim): counteroffer, with Danny's comments; Danny to resend his memo to Markus.",
 "SSLC: respond on the related-party transactions (Motion 5) within a few days.",
 "Next meeting: end of this week or early next week."]: B(t)
d.save(OUT); subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",f"{WS}/drafts",OUT],capture_output=True); print(OUT)
