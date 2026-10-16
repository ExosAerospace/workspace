#!/usr/bin/env python3
"""argon-reviewer-exo3: independent pre-send review of investor/partner communications with Gemini 4 Argon.

Jon 2026-10-15: "Go". Scope he set: investor and partner (e.g. AstroX) communications ONLY; never anything
relating to sensitive government programs or the rockets. This tool only returns findings. It never sends,
edits or attaches anything.

Usage:
  argon_review.py --file A.pptx [--file B.docx ...] [--task 3818]   # manual review; findings -> stdout + log (+ MC task note)
  argon_review.py --file X --task N --auto                           # mc-attach hook: runs only if the filename is in scope
  argon_review.py --selftest                                         # exo3 eval (offline, no API call)

Guards (deterministic, run BEFORE any API call):
  1. Sensitive-content block: export-control/CUI markings, contract numbers, rocket technical parameters -> BLOCKED.
  2. Scope gate (--auto): filename must look like investor/partner comms and must not look like a gov/technical doc.
  3. Monthly spend cap (config/argon-review.json, default $25) -> REFUSED when the month's estimated spend is reached.
"""
import argparse, datetime, json, os, re, sys, urllib.request

WS = "/Users/exomac/.openclaw/workspace"
CFG = os.path.join(WS, "config", "argon-review.json")
RULES = os.path.join(WS, "config", "argon-claim-rules.md")
STATE = os.path.join(WS, "state", "argon-spend.json")
LOGDIR = os.path.join(WS, "logs", "argon-reviews")
MODEL = "gemini-4-argon"
AGENT = "argon-reviewer-exo3"
OWNER = "John Quinn"

# Intro pricing $2/$10 per 1M tokens until ~Dec 1 2026, then $4/$20 (VentureBeat / Ars Technica, Sep-Oct 2026).
def price_per_m(today=None):
    today = today or datetime.date.today()
    return (2.0, 10.0) if today < datetime.date(2026, 12, 1) else (4.0, 20.0)

# ---- guard 1: sensitive content (never leaves the machine) -------------------------------------------------
SENSITIVE = [
    (r"\bITAR\b|\bUSML\b|\bEAR99\b|\b[0-9][A-E][0-9]{3}\b(?=.*\bECCN\b)", "export-control marking"),
    (r"export[- ]controlled|technical data\b.*\b(controlled|restricted)", "export-controlled content"),
    (r"\bCUI\b|\bFOUO\b|controlled unclassified|distribution statement [b-f]\b|\bdist(?:ribution)?[- ][b-f]\b", "CUI / distribution marking"),
    (r"\b(FA|W|N|HQ|NNX|NNM|80NSSC)[0-9A-Z]{2,4}-?\d{2}-?[A-Z]-?\d{3,5}\b", "government contract number"),
    (r"\b\d[\d,.]*\s*(lbf|kN|psia?|bar)\b|\bspecific impulse\b|\bIsp\b|chamber pressure|\bturbopump|\binjector\b|mixture ratio|\bGN&C\b|guidance law", "rocket technical parameter"),
]
def sensitive_hits(text):
    hits = []
    for pat, label in SENSITIVE:
        m = re.search(pat, text, flags=re.I)
        if m:
            hits.append(f"{label}: '{m.group(0)[:40]}'")
    return hits

# ---- guard 2: auto scope gate --------------------------------------------------------------------------------
IN_SCOPE = r"investor|japan[_ -]?round|astrox|pitch|one[_ -]?pager|newsletter|update[_ -]?email|blog|partner|thank|follow[_ -]?up|meeting[_ -]?deck|_deck|briefing"
OUT_OF_SCOPE = (r"proposal|white[_ -]?paper|sbir|sttr|rfi\b|_rfi|baa|afrl|nasa|darpa|spacewerx|tacfi|dibc|aedc|vsfb|faa|ddtc|itar|export"
                r"|counsel|procedure|bom|drawing|spec\b|_spec|cui|nda|mnda|contract|invoice|budget|cap[_ -]?table|board|resolution")
def should_review(fname):
    n = os.path.basename(fname).lower()
    if not n.endswith((".docx", ".pptx", ".txt")):
        return False
    return bool(re.search(IN_SCOPE, n)) and not re.search(OUT_OF_SCOPE, n)

# ---- guard 3: spend cap --------------------------------------------------------------------------------------
def load_cfg():
    try:
        return json.load(open(CFG))
    except Exception:
        return {"monthly_cap_usd": 25.0}
def month_key(today=None):
    return (today or datetime.date.today()).strftime("%Y-%m")
def spend(state_path=STATE, today=None):
    try:
        return float(json.load(open(state_path)).get(month_key(today), {}).get("usd", 0.0))
    except Exception:
        return 0.0
def record_spend(usd, tin, tout, state_path=STATE, today=None):
    try:
        st = json.load(open(state_path))
    except Exception:
        st = {}
    m = st.setdefault(month_key(today), {"usd": 0.0, "calls": 0, "tokens_in": 0, "tokens_out": 0})
    m["usd"] = round(m["usd"] + usd, 4); m["calls"] += 1; m["tokens_in"] += tin; m["tokens_out"] += tout
    os.makedirs(os.path.dirname(state_path), exist_ok=True)
    tmp = state_path + f".tmp{os.getpid()}"
    json.dump(st, open(tmp, "w"), indent=2); os.replace(tmp, state_path)
def over_cap(cap, state_path=STATE, today=None):
    return spend(state_path, today) >= cap

# ---- extraction ----------------------------------------------------------------------------------------------
def extract(path):
    if path.endswith(".pptx"):
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(path).slides, 1):
            out.append(f"[Slide {i}] " + " | ".join(sh.text_frame.text.replace("\n", " / ")
                       for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()))
        return "\n".join(out)
    if path.endswith(".docx"):
        from docx import Document
        d = Document(path)
        out = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                out.append(" | ".join(c.text for c in r.cells))
        return "\n".join(out)
    return open(path, errors="ignore").read()

PROMPT = """You are an independent pre-send reviewer for a U.S. aerospace company's investor and partner communications.
Review the document(s) below against the FACTS OF RECORD. Output a numbered list of concrete issues only: overclaims,
statements that conflict with the facts, internal inconsistencies between or within documents (numbers, dates, names),
and unsupported claims an investor or partner would challenge. For each issue give: document, slide/paragraph, the exact
quote, why it is a problem, and a suggested replacement. Skip style nitpicks. If everything is clean, reply exactly: CLEAN.
Do not rewrite the documents.

FACTS OF RECORD:
{rules}

DOCUMENTS:
{docs}"""

def call_argon(prompt):
    key = json.load(open(os.path.expanduser("~/.openclaw/secrets.json")))["google_default_key_main"]
    body = {"contents": [{"parts": [{"text": prompt}]}]}
    req = urllib.request.Request(f"https://generativelanguage.googleapis.com/v1beta/models/{MODEL}:generateContent?key={key}",
                                 data=json.dumps(body).encode(), headers={"Content-Type": "application/json"})
    j = json.load(urllib.request.urlopen(req, timeout=300))
    txt = "".join(p.get("text", "") for p in j["candidates"][0]["content"]["parts"] if not p.get("thought"))
    u = j.get("usageMetadata", {})
    return txt, int(u.get("promptTokenCount", 0)), int(u.get("candidatesTokenCount", 0)) + int(u.get("thoughtsTokenCount", 0))

def fleet(action, detail, outcome=None, level="info"):
    try:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from fleet_log import event
        event(AGENT, "govern", action, detail=detail, owner=OWNER, outcome=outcome, level=level)
    except Exception:
        pass

def task_note(task, body):
    try:
        key = json.load(open(os.path.expanduser("~/.openclaw/secrets.json")))["supabase_service_role_key"]
        h = {"apikey": key, "Authorization": "Bearer " + key, "Content-Type": "application/json"}
        rec = {"type": "task_note", "subject": f"task:{task}", "body": body[:3500], "owner": "Exo (Argon reviewer)", "status": "completed"}
        urllib.request.urlopen(urllib.request.Request("https://pcrgpuoljylqcumijvnw.supabase.co/rest/v1/activities",
                               data=json.dumps(rec).encode(), headers=h, method="POST"), timeout=30)
    except Exception as e:
        print(f"(task note failed: {e})", file=sys.stderr)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", action="append", required=True)
    ap.add_argument("--task")
    ap.add_argument("--auto", action="store_true", help="mc-attach hook mode: skip silently if out of scope")
    a = ap.parse_args()
    files = [f for f in a.file if os.path.exists(f)]
    if a.auto:
        files = [f for f in files if should_review(f)]
        if not files:
            return 0
    if not files:
        print("No readable files."); return 2
    texts = {os.path.basename(f): extract(f) for f in files}
    joined = "\n\n".join(f"=== {n} ===\n{t}" for n, t in texts.items())
    hits = sensitive_hits(joined)
    if hits:
        msg = "BLOCKED (not sent to Argon): " + "; ".join(hits)
        print(msg); fleet("blocked_sensitive", {"files": list(texts), "hits": hits}, outcome="blocked")
        if a.task and not a.auto:
            task_note(a.task, "Argon pre-send review " + msg)
        return 4
    cap = float(load_cfg().get("monthly_cap_usd", 25.0))
    if over_cap(cap):
        msg = f"REFUSED: Argon monthly cap ${cap:.0f} reached (spent ${spend():.2f} in {month_key()})."
        print(msg); fleet("refused_cap", {"files": list(texts)}, outcome="refused", level="warn"); return 5
    rules = open(RULES).read()
    txt, tin, tout = call_argon(PROMPT.format(rules=rules, docs=joined))
    pin, pout = price_per_m()
    usd = tin / 1e6 * pin + tout / 1e6 * pout
    record_spend(usd, tin, tout)
    clean = txt.strip().upper().startswith("CLEAN")
    n_issues = 0 if clean else len(re.findall(r"^\s*(?:\*\*)?\d+[.)]", txt, flags=re.M))
    os.makedirs(LOGDIR, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y-%m-%dT%H%M%S")
    log = os.path.join(LOGDIR, f"{stamp}_{os.path.splitext(list(texts)[0])[0][:60]}.md")
    open(log, "w").write(f"# Argon pre-send review {stamp}\nFiles: {', '.join(texts)}\nTask: {a.task or '-'}\n"
                         f"Tokens in/out: {tin}/{tout}  est ${usd:.4f}\n\n{txt}\n")
    fleet("review", {"files": list(texts), "task": a.task, "issues": n_issues, "usd": round(usd, 4), "log": log},
          outcome="clean" if clean else "issues")
    head = f"Argon pre-send review ({', '.join(texts)}): " + ("CLEAN, no issues." if clean else f"{n_issues} issue(s). Findings only; nothing was changed or sent.")
    print(head); print(txt); print(f"\n[log: {log}  est ${usd:.4f}]")
    if a.task:
        task_note(a.task, head + "\n\n" + txt)
    return 0

def selftest():
    """exo3 eval, offline. Replays the 2026-10-15 pilot: investor/partner decks in scope, gov/technical docs out,
    sensitive content blocked before any API call, spend cap enforced, price step on Dec 1."""
    import tempfile
    ok = 0
    def check(c, label):
        nonlocal ok
        if not c: raise AssertionError("FAIL: " + label)
        ok += 1
    check(should_review("Exos_AstroX_Meeting_Deck_2026-10-16_r2.pptx"), "AstroX meeting deck is in scope")
    check(should_review("Exos_Japan_Round_SBI_Investment_2026-10-16.pptx"), "Japan round investor deck is in scope")
    check(should_review("Exos_Investor_Update_Email_and_Blog_2026-10-14_CLEAN.docx"), "investor update is in scope")
    check(not should_review("Exos_Aerospace_VSFBMDZCall_Whitepaper_2026-10-15_r2.docx"), "VSFB gov whitepaper is out of scope")
    check(not should_review("Exos_NASA_SBIR_Proposal_Investor_Notes.docx"), "anything SBIR/NASA is out of scope even with 'investor'")
    check(not should_review("Exos_SPACE_COTAN_Counsel_Briefing_ITZ_2026-10-14.pptx"), "export-control counsel briefing is out of scope")
    check(not should_review("Exos_Investor_Update.pdf"), "PDFs are not reviewed")
    check(sensitive_hits("This document contains ITAR technical data"), "ITAR marking blocks")
    check(sensitive_hits("CUI // Distribution Statement D"), "CUI / Dist D blocks")
    check(sensitive_hits("Hermes contract FA9300-25-C-6022 status"), "gov contract number blocks")
    check(sensitive_hits("main engine 5,000 lbf at chamber pressure 300 psi"), "rocket technical parameters block")
    check(not sensitive_hits("1 of 3 U.S. FAA-licensed reusable launch providers; payloads returned in minutes; ¥790M Japan round"),
          "normal investor language passes")
    d = tempfile.mkdtemp(); sp = os.path.join(d, "spend.json"); t = datetime.date(2026, 10, 15)
    record_spend(24.99, 1000, 1000, state_path=sp, today=t)
    check(not over_cap(25.0, state_path=sp, today=t), "under cap allows")
    record_spend(0.02, 10, 10, state_path=sp, today=t)
    check(over_cap(25.0, state_path=sp, today=t), "at/over cap refuses")
    check(not over_cap(25.0, state_path=sp, today=datetime.date(2026, 11, 1)), "cap resets next month")
    check(price_per_m(datetime.date(2026, 11, 30)) == (2.0, 10.0) and price_per_m(datetime.date(2026, 12, 1)) == (4.0, 20.0), "price doubles Dec 1")
    check(os.path.exists(RULES) and "Part 431" in open(RULES).read(), "claim rules file present")
    print(f"argon_review selftest: {ok}/{ok} pass")

if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
    else:
        sys.exit(main())
