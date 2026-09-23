#!/usr/bin/env python3
"""inventory_capture.py — turn a photo taken on a Mission Control task into a BOM-matched
inventory record. Deterministic engine + queue. The VISION step (look at the photo, name the
component) runs in the MAIN Exo session (isolated crons can't read local files); this script
gives that session everything it needs and does the writing.

Flow:
  1. A photo is uploaded to a task (MC 'Take Photo' button) OR someone posts '@exo inv' on a task.
     -> an entry is enqueued here (kind=upload / kind=command).
  2. A launchd dispatcher (com.exos.inventory-capture) sees pending entries and wakes the main
     Exo session with the INVENTORY_CAPTURE_RUNBOOK.
  3. The main session: `next` -> Read the image -> `bom-search` -> `record` (or `ask` the uploader).

Subcommands:
  reindex-bom                          rebuild mission-control/data/bom-index.json from the BLK3 BOM xlsx
  enqueue --task ID --upload UPLID --file STORED --filename NAME --by WHO --kind upload|command
  pending                              print count of pending entries
  next                                 pop-context (does NOT delete): next pending entry as JSON + abs image path
  bom-search "keywords"                search the BOM index; JSON matches (item/pn/desc/section)
  record --queue QID --task ID --by WHO [--bom-key K] [--vendor-pn P] [--mfg-pn P]
         [--item I] [--section S] [--desc D] [--qty N] [--location L] [--note ...]
                                       upsert inventory record, mark entry done, post a ✅ note to the task
  ask --queue QID --task ID --uploader WHO --question "..."
                                       post a ❓ note to the task tagging the uploader; mark entry awaiting-reply
  status                               summary of queue + inventory
"""
import json, os, sys, re, argparse, urllib.request
from datetime import datetime, timezone

WS = "/Users/exomac/.openclaw/workspace"
MC = os.path.join(WS, "mission-control")
UPLOADS = os.path.join(MC, "uploads")
BOM_XLSX = os.path.join(UPLOADS, "1787343333564_BLK3_BOM_8-21-2026.xlsx")
BOM_INDEX = os.path.join(MC, "data", "bom-index.json")
INVENTORY = os.path.join(MC, "data", "inventory.json")
QUEUE = os.path.join(WS, "logs", "inventory-capture-queue.jsonl")
LOG = os.path.join(WS, "logs", "inventory-capture.log")
SUPA = "https://pcrgpuoljylqcumijvnw.supabase.co/rest/v1"
KEY = "sb_publishable_qBnodkDgym8-xuHPy3_qcg_0h98BYWp"
IMG_EXT = (".png", ".jpg", ".jpeg", ".heic", ".webp", ".gif")


def now():
    return datetime.now(timezone.utc).isoformat()


def logline(m):
    try:
        with open(LOG, "a") as f:
            f.write(f"{now()} {m}\n")
    except Exception:
        pass


def load(p, d):
    try:
        return json.load(open(p))
    except Exception:
        return d


def supa(method, path, body=None):
    req = urllib.request.Request(
        SUPA + "/" + path, method=method,
        data=json.dumps(body).encode() if body is not None else None,
        headers={"apikey": KEY, "Authorization": "Bearer " + KEY,
                 "Content-Type": "application/json", "Prefer": "return=representation"})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            t = r.read().decode()
            return json.loads(t) if t else []
    except Exception as e:
        logline(f"WARN supa {method} {path[:40]}: {e}")
        return None


def task_note(taskid, msg):
    supa("POST", "activities", {"type": "task_note", "subject": f"task:{taskid}",
         "body": msg, "owner": "Exo", "status": "completed"})


def flag_awaiting(taskid):
    try:
        cur = supa("GET", f"tasks?id=eq.{taskid}&select=tags,priority") or []
        tags = (cur[0].get("tags") or []) if cur else []
        if "awaiting-exo-answer" not in tags:
            tags = tags + ["awaiting-exo-answer"]
        pr = max(int(cur[0].get("priority") or 0), 88) if cur else 88
        supa("PATCH", f"tasks?id=eq.{taskid}", {"tags": tags, "priority": pr})
    except Exception as e:
        logline(f"WARN flag_awaiting {taskid}: {e}")


# ── QUEUE (append-only jsonl; each line an entry with a mutable 'status') ──
def read_queue():
    out = []
    if not os.path.exists(QUEUE):
        return out
    for ln in open(QUEUE):
        ln = ln.strip()
        if not ln:
            continue
        try:
            out.append(json.loads(ln))
        except Exception:
            pass
    return out


def write_queue(entries):
    tmp = QUEUE + ".tmp"
    with open(tmp, "w") as f:
        for e in entries:
            f.write(json.dumps(e) + "\n")
    os.replace(tmp, QUEUE)


def set_status(qid, status, extra=None):
    entries = read_queue()
    for e in entries:
        if e.get("id") == qid:
            e["status"] = status
            e["updated_at"] = now()
            if extra:
                e.update(extra)
    write_queue(entries)


def cmd_enqueue(a):
    if not a.file or not a.file.lower().endswith(IMG_EXT):
        # non-image uploads are ignored (only photos become inventory captures)
        print("skip non-image")
        return
    entries = read_queue()
    # de-dupe by upload id
    if a.upload and any(e.get("upload_id") == a.upload for e in entries):
        print("dup")
        return
    qid = "iq_" + datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S%f")
    entry = {"id": qid, "task_id": str(a.task), "upload_id": a.upload,
             "stored_name": a.file, "filename": a.filename or a.file,
             "uploaded_by": a.by or "Unknown", "kind": a.kind or "upload",
             "status": "pending", "created_at": now()}
    with open(QUEUE, "a") as f:
        f.write(json.dumps(entry) + "\n")
    logline(f"enqueue {qid} task {a.task} by {a.by} ({a.kind})")
    print(qid)


def cmd_pending(a):
    print(sum(1 for e in read_queue() if e.get("status") == "pending"))


def cmd_next(a):
    for e in read_queue():
        if e.get("status") == "pending":
            img = os.path.join(UPLOADS, e.get("stored_name", ""))
            e = dict(e)
            e["image_path"] = img
            e["image_exists"] = os.path.exists(img)
            e["bom_index"] = BOM_INDEX
            e["inventory"] = INVENTORY
            print(json.dumps(e, indent=2))
            return
    print("EMPTY")


# ── BOM INDEX ──
def cmd_reindex(a):
    import openpyxl
    wb = openpyxl.load_workbook(BOM_XLSX, data_only=True)
    rows = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        # header assumed at row index 1 (0-based) per BLK3 BOM layout
        for r in ws.iter_rows(min_row=3, values_only=True):
            if not r:
                continue
            def g(i):
                return (str(r[i]).strip() if i < len(r) and r[i] is not None else "")
            item, qty, vpn, mpn, desc, vendor = g(0), g(1), g(2), g(3), g(4), g(5)
            cost = g(8)
            if not (vpn or mpn or desc):
                continue
            rows.append({"section": sheet, "item": item, "qty": qty, "vendor_pn": vpn,
                         "mfg_pn": mpn, "description": desc, "vendor": vendor, "cost": cost})
    data = {"source": os.path.basename(BOM_XLSX), "generated_at": now(), "count": len(rows), "rows": rows}
    os.makedirs(os.path.dirname(BOM_INDEX), exist_ok=True)
    json.dump(data, open(BOM_INDEX, "w"), indent=2)
    print(f"indexed {len(rows)} BOM lines across {len(wb.sheetnames)} sheets -> {BOM_INDEX}")


def _norm(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def resolve_location_xlsx(explicit=None):
    """Pick the locations workbook to index: an explicit stored filename if given, else the most
    recently uploaded .xlsx whose name looks like a locations sheet, else the known baseline file."""
    if explicit:
        p = explicit if os.path.isabs(explicit) else os.path.join(UPLOADS, explicit)
        if os.path.exists(p):
            return p
    try:
        meta = json.load(open(UPLOADS_META))
    except Exception:
        meta = []
    cands = [m for m in meta if str(m.get("extension", "")).lower() == ".xlsx"
             and re.search(r"location", str(m.get("filename", "")), re.I)]
    cands.sort(key=lambda m: m.get("uploaded_at", ""), reverse=True)
    for m in cands:
        p = os.path.join(UPLOADS, m.get("stored_name", ""))
        if os.path.exists(p):
            return p
    return LOCATION_XLSX


def build_location_index(explicit=None):
    """Parse a Current_Inventory_Locations workbook into a flat catalog (location cell forward-filled
    per block) -> location-index.json. Then best-effort match BOM lines to a stored location by SKU or
    a distinctive description token -> bom-location-map.json (both the report and /api/inventory read it).
    Guarded: a workbook without a Parts Room / Warehouse / Trailer sheet is NOT a locations sheet and is
    skipped so an unrelated .xlsx can never wipe the good index. Returns (catalog_n, matched_n) or None."""
    import openpyxl
    src = resolve_location_xlsx(explicit)
    wb = openpyxl.load_workbook(src, data_only=True)
    areas = ["Parts Room", "Warehouse", "Trailer"]
    if not any(s in wb.sheetnames for s in areas):
        logline(f"reindex-locations SKIP: {os.path.basename(src)} has no Parts Room/Warehouse/Trailer sheet")
        return None
    cat = []
    for sh in areas:
        if sh not in wb.sheetnames:
            continue
        ws = wb[sh]
        # Header-aware column resolution (row 1), so inserting a BOM P/N column can't shift the parse.
        hdr = [("" if c is None else str(c).strip().lower()) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]

        def col_of(cands, default):
            for i, h in enumerate(hdr):
                if h in cands:
                    return i
            return default
        c_src = col_of(("source/sku", "source sku", "source"), 1)
        c_con = col_of(("contents", "description", "content"), 2)
        c_exo = col_of(("exos sku", "exos", "sku"), 3)
        c_qty = col_of(("qty", "quantity"), 4)
        c_unit = col_of(("unit",), 5)
        c_val = col_of(("value", "val"), 7)
        c_bom = col_of(("bom p/n", "bom pn", "bom p-n", "bom part", "bom part number"), None)
        need = max(8, (c_bom or 0) + 1, c_val + 1)
        loc = ""
        for row in ws.iter_rows(min_row=2, values_only=True):
            r = [("" if c is None else str(c).strip()) for c in row]
            r += [""] * (need - len(r))
            if r[0]:
                loc = r[0]
            contents, sku = r[c_con], (r[c_src] or r[c_exo])
            qty, unit, val = r[c_qty], r[c_unit], r[c_val]
            bom_pn = r[c_bom] if c_bom is not None else ""
            if not contents and not sku:
                continue
            cat.append({"area": sh, "location": loc, "contents": contents, "sku": sku,
                        "qty": qty, "unit": unit, "value": val, "bom_pn": bom_pn})
    maps = extract_location_maps(src)
    json.dump({"source": os.path.basename(src), "generated_at": now(), "count": len(cat),
               "maps": maps, "entries": cat}, open(LOCATION_INDEX, "w"), indent=2)
    # ── BOM ↔ location matching ──
    by_sku = {}
    by_bompn = {}
    for e in cat:
        if e["sku"]:
            by_sku[e["sku"].strip().upper()] = e
        if e.get("bom_pn"):
            by_bompn[e["bom_pn"].strip().upper()] = e   # authoritative: Steve's BOM P/N column
    bom = load(BOM_INDEX, {"rows": []})
    bmap = {}
    matched_authoritative = 0
    for row in bom.get("rows", []):
        vpn, mpn = row.get("vendor_pn", ""), row.get("mfg_pn", "")
        key = rec_key(vpn, mpn, row.get("item", ""), row.get("description", ""))
        hit = None; authoritative = False
        # 1) the explicit BOM P/N column wins (exact, human-verified)
        for pn in [vpn, mpn]:
            if pn and pn.strip().upper() in by_bompn:
                hit = by_bompn[pn.strip().upper()]; authoritative = True; break
        # 2) else fall back to SKU exact (>=5 chars so short/generic codes don't collide)
        if not hit:
            for pn in [vpn, mpn]:
                if pn and len(pn.strip()) >= 5 and pn.strip().upper() in by_sku:
                    hit = by_sku[pn.strip().upper()]; break
        if not hit:
            # distinctive token: alphanumeric with a digit, len>=5 (e.g. a model no) shared with contents
            dtoks = [t for t in _norm(row.get("description", "")).split()
                     if len(t) >= 5 and any(ch.isdigit() for ch in t)]
            for t in dtoks:
                for e in cat:
                    if t in _norm(e["contents"]).split():   # whole-word match, not substring
                        hit = e; break
                if hit:
                    break
        if hit:
            if authoritative:
                matched_authoritative += 1
            bmap[key] = {"location": hit["location"], "area": hit["area"],
                         "contents": hit["contents"], "source": "bom-pn" if authoritative else "map"}
    json.dump({"generated_at": now(), "count": len(bmap), "authoritative": matched_authoritative,
               "map": bmap}, open(BOM_LOC_MAP, "w"), indent=2)
    return len(cat), len(bmap)


MAPS_DIR = os.path.join(UPLOADS, "location-maps")


def _location_candidates(primary):
    """primary first, then other uploaded /location/i .xlsx (newest first), then the baseline —
    used to find map images even if the current working file (e.g. the BOM-P/N copy) dropped them."""
    out = [primary]
    try:
        meta = json.load(open(UPLOADS_META))
        cands = [m for m in meta if str(m.get("extension", "")).lower() == ".xlsx"
                 and re.search(r"location", str(m.get("filename", "")), re.I)]
        cands.sort(key=lambda m: m.get("uploaded_at", ""), reverse=True)
        out += [os.path.join(UPLOADS, m.get("stored_name", "")) for m in cands]
    except Exception:
        pass
    out.append(LOCATION_XLSX)
    seen, uniq = set(), []
    for p in out:
        if p and p not in seen and os.path.exists(p):
            seen.add(p); uniq.append(p)
    return uniq


def extract_location_maps(src):
    """Pull the physical overview map images (e.g. 'Parts Room Map', 'Warehouse Map') embedded in the
    locations workbook out to uploads/location-maps/<Sheet>.png so MC and the export can show them.
    Follows the OOXML rel chain: workbook -> sheet -> drawing -> image. Falls back across locations
    files so a copy that lost its images (openpyxl strips them) still yields the maps. Returns [{name,file}]."""
    os.makedirs(MAPS_DIR, exist_ok=True)
    for cand in _location_candidates(src):
        maps = _maps_from(cand)
        if maps:
            return maps
    return []


def _maps_from(src):
    import zipfile, re as _re, shutil
    maps = []
    try:
        with zipfile.ZipFile(src) as z:
            names = set(z.namelist())
            wbxml = z.read("xl/workbook.xml").decode("utf-8", "ignore")
            wbrels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8", "ignore")
            # attribute order varies (Id before/after Target) — parse each <Relationship> independently
            rid2target = {}
            for rel in _re.findall(r"<Relationship\b[^>]*/?>", wbrels):
                _id = _re.search(r'Id="([^"]+)"', rel); _tg = _re.search(r'Target="([^"]+)"', rel)
                if _id and _tg:
                    rid2target[_id.group(1)] = _tg.group(1)
            for sheet_tag in _re.findall(r"<sheet\b[^>]*/?>", wbxml):
                nm = _re.search(r'name="([^"]+)"', sheet_tag)
                rid = _re.search(r'r:id="([^"]+)"', sheet_tag)
                if not (nm and rid) or "map" not in nm.group(1).lower():
                    continue
                tgt = rid2target.get(rid.group(1), "")
                if not tgt:
                    continue
                sheet_path = tgt.lstrip("/")
                if not sheet_path.startswith("xl/"):
                    sheet_path = "xl/" + sheet_path
                base = os.path.basename(sheet_path)
                srels = f"xl/worksheets/_rels/{base}.rels"
                if srels not in names:
                    continue
                def _rel(base_dir, target):   # OOXML targets may be absolute (/xl/..) or relative (../..)
                    if target.startswith("/"):
                        return target.lstrip("/")
                    return os.path.normpath(os.path.join(base_dir, target)).replace("\\", "/")
                draw = _re.search(r'Target="([^"]*drawings/drawing[0-9]+\.xml)"', z.read(srels).decode("utf-8", "ignore"))
                if not draw:
                    continue
                dpath = _rel("xl/worksheets", draw.group(1))
                dbase = os.path.basename(dpath)
                drels = f"xl/drawings/_rels/{dbase}.rels"
                if drels not in names:
                    continue
                img = _re.search(r'Target="([^"]*media/image[0-9]+\.[a-zA-Z]+)"', z.read(drels).decode("utf-8", "ignore"))
                if not img:
                    continue
                ipath = _rel("xl/drawings", img.group(1))
                if ipath not in names:
                    continue
                safe = _re.sub(r"[^A-Za-z0-9]+", "_", nm.group(1)).strip("_") + os.path.splitext(ipath)[1]
                with z.open(ipath) as fsrc, open(os.path.join(MAPS_DIR, safe), "wb") as fdst:
                    shutil.copyfileobj(fsrc, fdst)
                maps.append({"name": nm.group(1), "file": safe})
    except Exception as e:
        logline(f"WARN extract_maps: {e}")
    return maps


MAP_OVERLAYS = os.path.join(MC, "data", "location-map-overlays.json")


def apply_map_overlays():
    """Draw durable team annotations (e.g. the new MC section) onto the extracted map PNGs so they
    survive re-uploads. Overlay coords are fractions of image size, keyed by map sheet name."""
    overlays = load(MAP_OVERLAYS, {})
    if not overlays:
        return
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception as e:
        logline(f"WARN overlays PIL: {e}"); return
    loci = load(LOCATION_INDEX, {"maps": []})
    for mp in loci.get("maps", []):
        ovs = overlays.get(mp.get("name", ""))
        if not ovs:
            continue
        p = os.path.join(MAPS_DIR, mp.get("file", ""))
        if not os.path.exists(p):
            continue
        try:
            im = Image.open(p).convert("RGBA")
            W, H = im.size
            d = ImageDraw.Draw(im)
            try:
                font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", max(12, int(H * 0.028)))
            except Exception:
                font = ImageFont.load_default()
            for ov in ovs:
                x0, y0 = ov["x"] * W, ov["y"] * H
                x1, y1 = (ov["x"] + ov["w"]) * W, (ov["y"] + ov["h"]) * H
                fill = ov.get("fill", "orange")
                fill = (255, 107, 0, 235) if fill == "orange" else tuple(fill) + ((255,) if len(fill) == 3 else ())
                tcol = tuple(ov.get("text_color", [255, 255, 255])) + (255,)
                outline = (20, 20, 20, 255)
                d.rectangle([x0, y0, x1, y1], fill=fill, outline=outline, width=2)
                lbl = ov.get("label", "")
                fnt = font
                fsz = ov.get("font_px")
                if fsz:
                    try: fnt = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", int(fsz))
                    except Exception: fnt = font
                if ov.get("rotate") == 90:
                    tmp = Image.new("RGBA", (max(1, int(y1 - y0)), max(1, int(x1 - x0))), (0, 0, 0, 0))
                    td = ImageDraw.Draw(tmp)
                    try:
                        tb = td.textbbox((0, 0), lbl, font=fnt); tw, th = tb[2] - tb[0], tb[3] - tb[1]
                    except Exception:
                        tw, th = len(lbl) * 6, 11
                    td.text(((tmp.width - tw) / 2, (tmp.height - th) / 2 - 2), lbl, fill=tcol, font=fnt)
                    tmp = tmp.rotate(90, expand=True)
                    im.paste(tmp, (int(x0 + ((x1 - x0) - tmp.width) / 2), int(y0 + ((y1 - y0) - tmp.height) / 2)), tmp)
                else:
                    try:
                        tb = d.textbbox((0, 0), lbl, font=fnt); tw, th = tb[2] - tb[0], tb[3] - tb[1]
                    except Exception:
                        tw, th = len(lbl) * 6, 11
                    d.text((x0 + ((x1 - x0) - tw) / 2, y0 + ((y1 - y0) - th) / 2 - 2), lbl, fill=tcol, font=fnt)
            im.convert("RGB").save(p)
            logline(f"overlay applied to {mp.get('file')} ({len(ovs)})")
        except Exception as e:
            logline(f"WARN overlay {mp.get('file')}: {e}")


def cmd_reindex_locations(a):
    res = build_location_index(getattr(a, "file", None))
    if res is None:
        print("skipped: not a locations sheet (no Parts Room/Warehouse/Trailer)"); return
    apply_map_overlays()
    n, m = res
    build_report()   # refresh the canonical report + MC page so new locations show immediately
    print(f"location catalog: {n} entries -> {LOCATION_INDEX}; BOM matched: {m} -> {BOM_LOC_MAP}; report refreshed")


def cmd_bom_search(a):
    idx = load(BOM_INDEX, {"rows": []})
    q = " ".join(a.query).lower().strip()
    toks = [t for t in re.split(r"[\s,/]+", q) if t]
    scored = []
    for row in idx.get("rows", []):
        hay = " ".join([row.get("vendor_pn", ""), row.get("mfg_pn", ""),
                        row.get("description", ""), row.get("vendor", ""),
                        row.get("section", ""), row.get("item", "")]).lower()
        sc = 0
        for t in toks:
            if t in hay:
                sc += 2
                # exact P/N hit is a strong signal
                if t == row.get("vendor_pn", "").lower() or t == row.get("mfg_pn", "").lower():
                    sc += 6
        if sc:
            scored.append((sc, row))
    scored.sort(key=lambda x: -x[0])
    print(json.dumps([{"score": s, **r} for s, r in scored[:12]], indent=2))


# ── INVENTORY ──
def inv_load():
    return load(INVENTORY, {"records": [], "updated_at": None})


def inv_save(inv):
    inv["updated_at"] = now()
    os.makedirs(os.path.dirname(INVENTORY), exist_ok=True)
    tmp = INVENTORY + ".tmp"
    json.dump(inv, open(tmp, "w"), indent=2)
    os.replace(tmp, INVENTORY)


def rec_key(vendor_pn, mfg_pn, item, desc):
    return (vendor_pn or mfg_pn or (item + ":" + desc[:24]) or desc[:24]).strip().upper()


LOCATION_XLSX = os.path.join(UPLOADS, "1790085634943_Current_Inventory_Locations.xlsx")
LOCATION_INDEX = os.path.join(MC, "data", "location-index.json")
BOM_LOC_MAP = os.path.join(MC, "data", "bom-location-map.json")
REPORT_NAME = "BLK3_Inventory_Report.xlsx"
REPORT_PATH = os.path.join(UPLOADS, REPORT_NAME)
UPLOADS_META = os.path.join(UPLOADS, "_meta.json")
REPORT_TASK = "3024"          # the Computer-section build/review task the report is pinned to


# ── Shared legend / key (Steve's warehouse codes + the BLK3 inventory-view codes). Used identically
# in the Excel export and the Mission Control Inventory page so both read the same. ──
STEVE_LEGEND = [
    ("Light Red Berry", "E6B8AF", "Inventory Subset", "Items go together"),
    ("Red", "FF0000", "Inventory Discrepancy", "Count / record does not match"),
    ("Kelly Green", "00FF00", "Trash", "Scrap / disposal"),
    ("Yellow", "FFFF00", "Uncertain", "Needs verification"),
    ("Dark Green", "6AA84F", "Possible Sale", "Candidate to sell"),
    ("Blue", "0000FF", "Home / Resting Place", "Put back in location after use"),
    ("Orange", "FF9900", "To-Do", "Requires assistance"),
]
OPS_LEGEND = [
    ("Light Orange", "FCE5CD", "SHORT vs BOM", "Required qty exceeds on-hand — reorder"),
    ("Light Green", "D9EAD3", "On-Hand / Verified", "On-hand meets the BOM requirement"),
    ("Light Grey", "EFEFEF", "Consumed / Inactive", "Used or installed; no stock held"),
]
HDR_GREY = "B7B7B7"
BAND_GREY = "F3F3F3"
SHORT_FILL = "FCE5CD"
HAVE_FILL = "D9EAD3"


def build_report():
    """Regenerate the CANONICAL inventory spreadsheet: a Legend/key tab, multiple views (Inventory,
    Shorts, On-Hand), the Locations map, and a Summary — mirroring the structure/key/legend/format of
    Steve's locations workbook. Pinned on the build task so it's always openable in MC."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception as e:
        logline(f"WARN report openpyxl: {e}"); return
    bom = load(BOM_INDEX, {"rows": []})
    inv = inv_load()
    locmap = load(BOM_LOC_MAP, {"map": {}}).get("map", {})
    by = {}
    for r in inv.get("records", []):
        for k in [r.get("vendor_pn"), r.get("mfg_pn"), r.get("key")]:
            if k:
                by[str(k).strip().upper()] = r
        if r.get("section") and r.get("item"):
            by[f"{r['section'].upper()}#{r['item']}"] = r

    hdr_font = Font(bold=True, color="000000")
    hdr_fill = PatternFill("solid", fgColor=HDR_GREY)
    band_fill = PatternFill("solid", fgColor=BAND_GREY)
    short_fill = PatternFill("solid", fgColor=SHORT_FILL)
    have_fill = PatternFill("solid", fgColor=HAVE_FILL)
    thin = Side(style="thin", color="D9D9D9"); cell_border = Border(bottom=thin)

    # Build the joined data once.
    data = []
    onhand_total = 0
    for row in bom.get("rows", []):
        sec, item = row.get("section", ""), row.get("item", "")
        vpn, mpn, desc, qty = row.get("vendor_pn", ""), row.get("mfg_pn", ""), row.get("description", ""), row.get("qty", "")
        rec = None
        for k in [vpn, mpn]:
            if k and str(k).strip().upper() in by:
                rec = by[str(k).strip().upper()]; break
        if rec is None and item:
            rec = by.get(f"{sec.upper()}#{item}")
        oh = rec.get("qty_on_hand", 0) if rec else 0
        try:
            req = float(str(qty).strip())
        except Exception:
            req = None
        shortv = round(req - oh, 3) if (req is not None and oh < req) else 0
        loc = ", ".join(rec.get("locations", [])) if rec else ""
        if not loc:
            mm = locmap.get(rec_key(vpn, mpn, item, desc))
            if mm:
                loc = f"{mm['location']} ({mm['area']}, map)"
        nph = len(rec.get("photos", [])) if rec else 0
        upd = (rec.get("updated_at", "") or "")[:19].replace("T", " ") if rec else ""
        status = "short" if shortv else ("have" if oh > 0 else "")
        onhand_total += oh
        data.append({"cells": [sec, item, vpn, mpn, desc, qty, oh, shortv or "", loc, nph or "", upd],
                     "status": status})
    short_total = sum(1 for d in data if d["status"] == "short")

    INV_COLS = ["Section", "Item", "Vendor P/N", "Mfg P/N", "Description", "Req Qty",
                "On-Hand", "Short", "Location(s)", "Photos", "Last Updated"]
    INV_W = [16, 6, 18, 18, 42, 8, 8, 8, 22, 7, 18]

    def render_table(ws, cols, widths, rows, colorby=True):
        ws.append(cols)
        for c in ws[1]:
            c.font = hdr_font; c.fill = hdr_fill
            c.alignment = Alignment(vertical="center")
        for i, d in enumerate(rows):
            ws.append(d["cells"] if isinstance(d, dict) else d)
            rr = ws[ws.max_row]
            st = d.get("status") if isinstance(d, dict) else ""
            fillc = short_fill if (colorby and st == "short") else (have_fill if (colorby and st == "have") else (band_fill if i % 2 else None))
            for c in rr:
                if fillc:
                    c.fill = fillc
                c.border = cell_border
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    wb = openpyxl.Workbook()
    # 1) LEGEND / KEY (first tab, mirrors Steve's Legend)
    lg = wb.active; lg.title = "Legend"
    lg.append(["BLK3 Inventory — Legend / Key"]); lg["A1"].font = Font(bold=True, size=14)
    lg.append(["Generated", now()[:19].replace("T", " ")])
    lg.append([])
    lg.append(["COLOR", "=", "CODE", "DESCRIPTION"])
    for c in lg[lg.max_row]:
        c.font = hdr_font; c.fill = hdr_fill
    def legend_block(title, items):
        lg.append([title]); lg[lg.max_row][0].font = Font(bold=True, italic=True)
        for name, argb, code, desc in items:
            lg.append([name, "=", code, desc])
            lg[lg.max_row][0].fill = PatternFill("solid", fgColor=argb)
            if argb in ("FF0000", "0000FF", "6AA84F"):
                lg[lg.max_row][0].font = Font(color="FFFFFF")
    legend_block("Warehouse codes (Steve's key)", STEVE_LEGEND)
    lg.append([])
    legend_block("BLK3 inventory-view codes", OPS_LEGEND)
    for col, w in zip("ABCD", [20, 4, 24, 42]):
        lg.column_dimensions[col].width = w

    # 2) INVENTORY (master), 3) SHORTS (view), 4) ON-HAND (view)
    render_table(wb.create_sheet("Inventory"), INV_COLS, INV_W, data)
    render_table(wb.create_sheet("Shorts"), INV_COLS, INV_W, [d for d in data if d["status"] == "short"])
    render_table(wb.create_sheet("On-Hand"), INV_COLS, INV_W, [d for d in data if d["status"] == "have"])

    # 5) LOCATIONS map (Steve's data, by area)
    loci = load(LOCATION_INDEX, {"entries": []})
    render_table(wb.create_sheet("Locations"),
                 ["Area", "Location", "Contents", "SKU", "BOM P/N", "Qty", "Unit", "Value"],
                 [14, 12, 46, 16, 16, 7, 8, 10],
                 [[e.get("area", ""), e.get("location", ""), e.get("contents", ""), e.get("sku", ""),
                   e.get("bom_pn", ""), e.get("qty", ""), e.get("unit", ""), e.get("value", "")]
                  for e in loci.get("entries", [])], colorby=False)

    # 5b) PHYSICAL MAP sheets — the overview images embedded in Steve's workbook (re-embedded here).
    for mp in loci.get("maps", []):
        imgp = os.path.join(MAPS_DIR, mp.get("file", ""))
        if not os.path.exists(imgp):
            continue
        try:
            from openpyxl.drawing.image import Image as XLImage
            msheet = wb.create_sheet(mp.get("name", "Map")[:31])
            msheet["A1"] = mp.get("name", "Map"); msheet["A1"].font = Font(bold=True, size=13)
            img = XLImage(imgp); img.anchor = "A3"
            msheet.add_image(img)
        except Exception as e:
            logline(f"WARN embed map {mp.get('file')}: {e}")

    # 6) SUMMARY
    s = wb.create_sheet("Summary")
    s.append(["BLK3 Inventory — canonical report"]); s["A1"].font = Font(bold=True, size=14)
    for k, v in [("Generated", now()[:19].replace("T", " ")), ("BOM lines", len(data)),
                 ("Records with on-hand", sum(1 for r in inv.get("records", []) if r.get("qty_on_hand", 0) > 0)),
                 ("Lines short", short_total), ("Total units on-hand", round(onhand_total, 3)),
                 ("Inventory records", len(inv.get("records", []))),
                 ("Location-map entries", len(loci.get("entries", [])))]:
        s.append([k, v])
    s.column_dimensions["A"].width = 24; s.column_dimensions["B"].width = 28

    wb.save(REPORT_PATH)
    _pin_report(len(data), short_total)
    logline(f"report refreshed: {len(data)} lines, {short_total} short -> {REPORT_PATH}")


def _pin_report(rows_total, short_total):
    """Upsert the report as a stable document on the build task so it shows (and updates in place)
    on that task's Documents tab in Mission Control."""
    try:
        meta = json.load(open(UPLOADS_META)) if os.path.exists(UPLOADS_META) else []
    except Exception:
        meta = []
    size = os.path.getsize(REPORT_PATH) if os.path.exists(REPORT_PATH) else 0
    entry = next((m for m in meta if m.get("id") == "blk3-inv-report"), None)
    if entry is None:
        entry = {"id": "blk3-inv-report", "filename": REPORT_NAME, "stored_name": REPORT_NAME,
                 "extension": ".xlsx", "task_id": REPORT_TASK, "project_id": None, "decision_id": None,
                 "uploaded_by": "Exo (inventory)", "current": True}
        meta.append(entry)
    entry["size"] = size
    entry["uploaded_at"] = now()
    entry["task_id"] = REPORT_TASK
    tmp = UPLOADS_META + ".tmp"
    json.dump(meta, open(tmp, "w"))
    os.replace(tmp, UPLOADS_META)


def cmd_record(a):
    inv = inv_load()
    key = a.bom_key or rec_key(a.vendor_pn or "", a.mfg_pn or "", a.item or "", a.desc or "")
    if not key:
        print("ERROR: need a key (--bom-key or --vendor-pn/--mfg-pn/--desc)")
        sys.exit(2)
    rec = next((r for r in inv["records"] if r.get("key") == key), None)
    existed = rec is not None
    # Safe count semantics (Jon 2026-09-22): an explicit --qty always wins (stock receipt).
    # Otherwise a photo of a NEW part records the first unit (+1); a photo of an ALREADY-recorded
    # part just CONFIRMS it and attaches the photo (+0) — re-photographing never double-counts.
    if a.qty not in (None, ""):
        qty = float(a.qty)
    else:
        qty = 0 if existed else 1
    entry_photo = None
    e = next((x for x in read_queue() if x.get("id") == a.queue), None) if a.queue else None
    if e:
        entry_photo = {"task": e.get("task_id"), "upload": e.get("upload_id"),
                       "file": e.get("stored_name"), "by": e.get("uploaded_by"), "at": now()}
    if rec is None:
        rec = {"key": key, "vendor_pn": a.vendor_pn or "", "mfg_pn": a.mfg_pn or "",
               "item": a.item or "", "section": a.section or "", "description": a.desc or "",
               "qty_on_hand": 0, "unit": "ea", "locations": [], "photos": [], "history": []}
        inv["records"].append(rec)
    rec["qty_on_hand"] = round(rec.get("qty_on_hand", 0) + qty, 3)
    if a.location and a.location not in rec["locations"]:
        rec["locations"].append(a.location)
    if a.section and not rec.get("section"):
        rec["section"] = a.section
    if a.desc and not rec.get("description"):
        rec["description"] = a.desc
    if a.vendor_pn and not rec.get("vendor_pn"):
        rec["vendor_pn"] = a.vendor_pn
    if a.mfg_pn and not rec.get("mfg_pn"):
        rec["mfg_pn"] = a.mfg_pn
    if entry_photo:
        rec["photos"].append(entry_photo)
    rec["history"].append({"at": now(), "by": a.by or "Exo", "delta": qty,
                           "note": a.note or "photo-captured", "queue": a.queue})
    rec["updated_at"] = now()
    inv_save(inv)
    if a.queue:
        set_status(a.queue, "recorded", {"inventory_key": key})
    if a.task:
        loc = f" · loc: {a.location}" if a.location else ""
        pn = a.vendor_pn or a.mfg_pn or ""
        pn = f" · P/N {pn}" if pn else ""
        verb = "recorded"
        tail = f"→ on-hand now **{rec['qty_on_hand']} {rec['unit']}**"
        if existed and qty == 0:
            verb = "confirmed (photo added, count unchanged)"
            tail = f"→ on-hand stays **{rec['qty_on_hand']} {rec['unit']}**"
        elif qty:
            tail = f"→ +{qty:g}, on-hand now **{rec['qty_on_hand']} {rec['unit']}**"
        task_note(a.task, f"✅ Exo (inventory): {verb} **{rec.get('description') or key}**"
                          f"{pn} {tail}{loc}. "
                          f"(from photo by {a.by or 'team'}; BOM section {rec.get('section') or '—'})")
    build_report()
    logline(f"record key={key} +{qty} -> {rec['qty_on_hand']} task {a.task}")
    print(json.dumps({"key": key, "qty_on_hand": rec["qty_on_hand"]}, indent=2))


def cmd_ask(a):
    q = a.question.strip()
    if a.queue:
        set_status(a.queue, "awaiting-reply", {"question": q})
    if a.task:
        flag_awaiting(a.task)
        task_note(a.task, f"❓ Exo (inventory): @{a.uploader or 'uploader'} — I need one detail to log the part "
                          f"you photographed: {q} Reply here on this task and I'll record it.")
    build_report()
    logline(f"ask task {a.task} uploader {a.uploader}: {q[:80]}")
    print("asked")


def cmd_report(a):
    build_report()
    print(f"report -> {REPORT_PATH}")


def cmd_consume(a):
    """Record a CONSUMPTION: parts used/installed leave stock. Decrements on-hand (or sets an
    explicit --remaining), logs the transaction, refreshes the report, posts a note."""
    inv = inv_load()
    key = a.bom_key or rec_key(a.vendor_pn or "", a.mfg_pn or "", a.item or "", a.desc or "")
    if not key:
        print("ERROR: need a key (--bom-key or --vendor-pn/--mfg-pn/--desc)"); sys.exit(2)
    rec = next((r for r in inv["records"] if r.get("key") == key), None)
    if rec is None:
        rec = {"key": key, "vendor_pn": a.vendor_pn or "", "mfg_pn": a.mfg_pn or "",
               "item": a.item or "", "section": a.section or "", "description": a.desc or "",
               "qty_on_hand": 0, "unit": "ea", "locations": [], "photos": [], "history": []}
        inv["records"].append(rec)
    qty = float(a.qty or 0)
    if a.remaining not in (None, ""):
        rec["qty_on_hand"] = round(float(a.remaining), 3)
    else:
        rec["qty_on_hand"] = round(rec.get("qty_on_hand", 0) - qty, 3)
    if a.location and a.location not in rec["locations"]:
        rec["locations"].append(a.location)
    for fld, val in [("section", a.section), ("description", a.desc), ("vendor_pn", a.vendor_pn), ("mfg_pn", a.mfg_pn)]:
        if val and not rec.get(fld):
            rec[fld] = val
    rec["history"].append({"at": now(), "by": a.by or "Exo", "delta": -qty,
                           "note": a.note or "consumed", "location": a.location or "", "queue": None})
    rec["updated_at"] = now()
    inv_save(inv)
    # short vs BOM for the note
    bom = load(BOM_INDEX, {"rows": []})
    req = None
    for row in bom.get("rows", []):
        if row.get("section", "").lower() == (rec.get("section") or "").lower() and str(row.get("item")) == str(rec.get("item")):
            try:
                req = float(str(row.get("qty")).strip())
            except Exception:
                req = None
            break
    shortnote = ""
    if req is not None and rec["qty_on_hand"] < req:
        shortnote = f" · SHORT {round(req - rec['qty_on_hand'], 3)} vs BOM (req {req:g}) → reorder"
    build_report()
    if a.task:
        loc = f" at {a.location}" if a.location else ""
        task_note(a.task, f"🔻 Exo (inventory): consumed **{qty:g} × {rec.get('description') or key}**{loc} "
                          f"→ on-hand now **{rec['qty_on_hand']:g} {rec['unit']}**{shortnote}. (per {a.by or 'team'})")
    logline(f"consume key={key} -{qty} -> {rec['qty_on_hand']} loc {a.location} task {a.task}")
    print(json.dumps({"key": key, "qty_on_hand": rec["qty_on_hand"], "short_vs_bom": (round(req - rec['qty_on_hand'], 3) if req is not None else None)}, indent=2))


def cmd_receive(a):
    """Receive parts against a QuickBooks PO: increment on-hand, set location + unit cost, stamp the
    PO#. Positive stock movement with cost/provenance (used by the Receiving module)."""
    inv = inv_load()
    key = a.bom_key or rec_key(a.vendor_pn or "", a.mfg_pn or "", a.item or "", a.desc or "")
    if not key:
        print("ERROR: need a key"); sys.exit(2)
    rec = next((r for r in inv["records"] if r.get("key") == key), None)
    if rec is None:
        rec = {"key": key, "vendor_pn": a.vendor_pn or "", "mfg_pn": a.mfg_pn or "",
               "item": a.item or "", "section": a.section or "", "description": a.desc or "",
               "qty_on_hand": 0, "unit": "ea", "locations": [], "photos": [], "history": []}
        inv["records"].append(rec)
    qty = float(a.qty or 0)
    rec["qty_on_hand"] = round(rec.get("qty_on_hand", 0) + qty, 3)
    if a.unit_cost not in (None, ""):
        try:
            rec["unit_cost"] = round(float(a.unit_cost), 4)
        except Exception:
            pass
    if a.location and a.location not in rec["locations"]:
        rec["locations"].append(a.location)
    for fld, val in [("section", a.section), ("description", a.desc), ("vendor_pn", a.vendor_pn), ("mfg_pn", a.mfg_pn)]:
        if val and not rec.get(fld):
            rec[fld] = val
    rec["history"].append({"at": now(), "by": a.by or "Exo", "delta": qty,
                           "note": f"received PO#{a.po}" if a.po else "received", "po": a.po or "",
                           "location": a.location or "", "unit_cost": a.unit_cost or "", "queue": None})
    rec["updated_at"] = now()
    inv_save(inv)
    build_report()
    logline(f"receive key={key} +{qty} @ {a.unit_cost} PO#{a.po} loc {a.location} -> {rec['qty_on_hand']}")
    print(json.dumps({"key": key, "qty_on_hand": rec["qty_on_hand"], "unit_cost": rec.get("unit_cost")}))


def cmd_status(a):
    q = read_queue()
    inv = inv_load()
    byst = {}
    for e in q:
        byst[e.get("status", "?")] = byst.get(e.get("status", "?"), 0) + 1
    print(json.dumps({"queue": byst, "inventory_records": len(inv.get("records", [])),
                      "inventory_units": round(sum(r.get("qty_on_hand", 0) for r in inv.get("records", [])), 2)}, indent=2))


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("reindex-bom")
    e = sub.add_parser("enqueue")
    e.add_argument("--task"); e.add_argument("--upload"); e.add_argument("--file")
    e.add_argument("--filename"); e.add_argument("--by"); e.add_argument("--kind", default="upload")
    sub.add_parser("pending")
    sub.add_parser("next")
    b = sub.add_parser("bom-search"); b.add_argument("query", nargs="+")
    r = sub.add_parser("record")
    for f in ["queue", "task", "by", "bom-key", "vendor-pn", "mfg-pn", "item", "section", "desc", "qty", "location", "note"]:
        r.add_argument("--" + f)
    k = sub.add_parser("ask")
    k.add_argument("--queue"); k.add_argument("--task"); k.add_argument("--uploader"); k.add_argument("--question", required=True)
    cs = sub.add_parser("consume")
    for f in ["bom-key", "vendor-pn", "mfg-pn", "item", "section", "desc", "qty", "remaining", "location", "by", "note", "task"]:
        cs.add_argument("--" + f)
    rc = sub.add_parser("receive")
    for f in ["bom-key", "vendor-pn", "mfg-pn", "item", "section", "desc", "qty", "unit-cost", "location", "po", "by", "task"]:
        rc.add_argument("--" + f)
    sub.add_parser("status")
    sub.add_parser("report")
    rl = sub.add_parser("reindex-locations"); rl.add_argument("--file")
    a = p.parse_args()
    # normalize dashed dests
    for src, dst in [("bom_key", "bom_key"), ("vendor_pn", "vendor_pn"), ("mfg_pn", "mfg_pn")]:
        pass
    {"reindex-bom": cmd_reindex, "enqueue": cmd_enqueue, "pending": cmd_pending,
     "next": cmd_next, "bom-search": cmd_bom_search, "record": cmd_record,
     "ask": cmd_ask, "status": cmd_status, "report": cmd_report, "consume": cmd_consume,
     "reindex-locations": cmd_reindex_locations, "receive": cmd_receive}[a.cmd](a)


if __name__ == "__main__":
    main()
