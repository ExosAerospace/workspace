#!/usr/bin/env python3
"""Exos x Space Grove Ventures collaboration agenda (Dr. Melissa Patton), Up/Link 1:1 Wed 2026-10-21."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/Users/exomac/.openclaw/workspace/drafts/uplink/Exos_SpaceGrove_Collaboration_Agenda_2026-10-21_r1.docx"
NAVY, GREY, RED = (0, 51, 102), (80, 80, 80), (160, 0, 0)


def run(p, text, bold=False, size=11, color=None, italic=False):
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    return r


def h(doc, text, color=NAVY):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); run(p, text, True, 13, color)


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            run(p, it[0], bold=True); run(p, it[1])
        else:
            run(p, it)


def build():
    doc = Document()
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2.5)
    for text, size, bold, color in [
        ("EXOS AEROSPACE × SPACE GROVE VENTURES", 18, True, NAVY),
        ("Collaboration Plan — Meeting Agenda (rev 1: Science Fair Payload Award)", 14, True, NAVY),
        ("Wednesday, October 21, 2026  |  3:45–4:00 PM ET  |  Up/Link, AWS Building, 12 W 39th St, NYC — Room 206, Table 22", 10, False, GREY),
        ("Exos: John Quinn, Cofounder & CEO   |   Space Grove: Dr. Melissa Patton, Founder & CEO", 10, False, GREY),
        ("Room/time per the Up/Link app as of 10/20; it has been moving meetings — check the app before going.", 9, False, RED)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, text, bold, size, color)

    h(doc, "MEETING GOAL")
    p = doc.add_paragraph()
    run(p, "Agree to co-develop a national Science Fair Payload Award: one winner from each state gets a payload on an annual "
           "Exos suborbital rideshare campaign. Bring in Nova Space (training) and the Montana spaceport work. Also agree on a "
           "two-way referral partnership. Leave with owners and a follow-up date, not a signed document.")

    h(doc, "AGENDA (15 MINUTES)")
    for t, title, pts in [
        ("1. Introductions (2 min)", "", [
            "John: Exos flies reusable suborbital rockets with full vehicle recovery, and is building BLK3 to fly payloads and test hardware in early 2027.",
            "Melissa: where Space Grove is today — the facility on the Space Coast, who is in the network, what's next."]),
        ("2. What each side brings (3 min)", "", [
            ("Space Grove: ", "real-estate activation and facility operations for space and defense companies, public-private collaboration support, a venture platform and accelerator pathways, workforce, student and university links (thespacegrove.org)."),
            ("Exos: ", "a reusable suborbital flight platform where hardware flies and is recovered for flight data; payloads come back in minutes. Exos also has an Italian affiliate (Exos Aerospace Italia, Turin)."),
            "Question for Melissa: How many of your companies are building hardware that will need flight testing in 2027?"]),
        ("3. Lead item: Science Fair Payload Award (5 min)", "", [
            ("The idea: ", "the top science fair project from each state wins a student payload on an Exos suborbital flight. All winners fly together on one annual Exos rideshare campaign."),
            ("How it is paid for: ", "each state's award is funded by sponsors (companies, state Space Grant groups, economic-development agencies, foundations). Exos sells the flight capacity at its standard payload terms. No free flights."),
            ("Who does what: ", "Space Grove leads sponsor recruiting, events and the Florida showcase. Exos provides the flight, the payload rules and recovery, and returns the payloads. Nova Space / Space Workforce Institute trains winners online and issues a credential. The Montana spaceport effort adds a hands-on workforce track and a possible future campaign site."),
            "Ask Melissa: Will Space Grove co-own the rollout plan with Exos and lead sponsor recruiting? Which 2–3 states should pilot it?"]),
        ("4. Other ways to work together (2 min)", "", [
            ("Referrals: ", "Space Grove sends flight-test candidates to Exos; Exos sends customers that need a Space Coast base to Space Grove."),
            ("Event: ", "an Exos session on her Off the Record / Behind Closed Doors podcast series about testing hardware in flight before orbit."),
            ("Italy/Europe: ", "she is an Italian Trade Agency delegate at Maker Faire Rome Oct 23–25. Explore European student and company links through Exos Italia."),
            ("Investor intros: ", "ask: 'Do you know space investors you could introduce us to? We're raising a $10M Series A.' Space Grove doesn't appear to invest its own money (it preps founders and makes intros), so ask for introductions, not a check. Follow-up: 'Which investors have you brought into deals this year?'"),
            ]),
        ("5. Next steps and owners (3 min)", "", [
            "Agree on owners: Exos (John; Lizi Quinn day-to-day) and Space Grove (Melissa or her delegate).",
            "Exos drafts a 2-page Science Fair Payload Award concept within 2 weeks and sends the payload one-pager.",
            "Exos sets up an intro call between Space Grove and Nova Space / Space Workforce Institute.",
            "Melissa sends any space investor and sponsor introductions she can make; Exos logs them in the CRM.",
            "Follow-up call in the week of Nov 2 (after her Rome trip and the Oct 28–29 Space Economy Summit at UCF).",
            "If both sides want it in writing later: mutual NDA first, then a simple MOU (reviewed by Counsel)."]),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); run(p, t, True, 11, NAVY)
        bullets(doc, pts)

    h(doc, "SCIENCE FAIR PAYLOAD AWARD: DRAFT ROLLOUT PLAN (TO BUILD TOGETHER)")
    for t, pts in [
        ("Phase 1, Nov–Dec 2026: Design", [
            "Exos + Space Grove agree on the award rules, the student payload size and mass limits (Exos sets them), judging, and the sponsor package.",
            "Bring in Nova Space / Space Workforce Institute to design a short online course for winners (payload design, flight safety, data analysis) with a credential.",
            "Pick 2–3 pilot states. Suggested: Florida (Space Grove's home), Texas (Exos home) and Montana (spaceport work).",
            "Contact state science fair organizers in the pilot states to offer it as a sponsored special award."]),
        ("Phase 2, Jan–May 2027: Pilot", [
            "Sponsors fund the pilot-state awards. Awards are given at the 2027 state fairs.",
            "Winners take the Nova/SWI course and build their payloads with Exos mentoring."]),
        ("Phase 3: First annual rideshare campaign", [
            "Pilot winners fly together on one Exos suborbital campaign. The date is set once BLK3 flies (BLK3 is targeted for early 2027).",
            "Winners and their families attend the launch. Space Grove and sponsors hold a showcase event; media coverage for sponsors."]),
        ("Phase 4, 2027–28 school year onward: National", [
            "Expand state by state as sponsors sign on, aiming for every state. Repeat the rideshare campaign every year.",
            "Montana: as the spaceport develops, add a hands-on workforce track there for winners and possibly host a future campaign."]),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); run(p, t, True, 11, NAVY)
        bullets(doc, pts)
    p = doc.add_paragraph(); run(p, "Must-haves: ", True); run(p,
        "every payload is checked for safety and export control before it flies; parent consent and youth-protection rules for minors; "
        "all prices and discounts approved by John.")

    h(doc, "NOVA SPACE AND THE MONTANA SPACEPORT (JOHN ONLY: DO NOT SHARE FIGURES)")
    bullets(doc, [
        "Nova Space sells online training for space professionals. Its Space Workforce Institute issues the certificates. Sites: <https://novaspaceinc.com>, <https://spaceworkforceinstitute.org>. Contacts: Christopher Allen (co-founder), Robert Zielinski (Director of Customer Solutions). MC #3213.",
        "Exos already proposed Nova/SWI as the education partner for the Montana spaceport effort (Rocket Trades to Space-Tech Trades credentials). The science fair award gives that partnership a national student entry point.",
        "Montana spaceport: still in development, not licensed. Exos has no role in the Department of War 'super range' announced 10/14. Hold all super-range talk until John and Edgar talk. Don't name Edgar or Montana partners to Space Grove without John's OK.",
    ])

    h(doc, "WHAT SUCCESS LOOKS LIKE 90 DAYS OUT")
    bullets(doc, [
        "Science Fair Payload Award concept agreed by Exos, Space Grove and Nova Space, with 2–3 pilot states and at least one sponsor in talks.",
        "At least 3 Space Grove companies introduced to Exos for flight-test conversations.",
        "One joint event or podcast done.",
        "A decision on whether a written referral MOU is worth doing.",
    ])

    h(doc, "BACKGROUND ON SPACE GROVE (SOURCED)")
    bullets(doc, [
        "Launched publicly at SpaceCom/Space Congress, Orlando, Jan 28–30, 2026. Focus: real-estate activation, facility operations and ecosystem alignment for space, defense and advanced tech. Source: <https://commercialspaceweek.com/press-releases/space-grove-ventures-announces-public-launch-spacecom-space-congress-signaling-new-commercial-model-space-defense-ecosystem-development>",
        "Ventures arm: facility services, deal structuring, venture studio, startup/accelerator pathways, investor opportunities. Source: <https://thespacegrove.org/venture-platform>",
        "Dr. Patton: 20+ years aligning government, industry and academia; founder of Space Grove Ventures. Source: <https://thespacegrove.org/team>",
        "Her upcoming events: Up/Link (Oct 20–21), Italian Trade Agency delegate at Maker Faire Rome (Oct 23–25), Economist Space Economy Summit at UCF (Oct 28–29). Source: <https://drmelissapatton.com/events>",
        "She said Space Grove is 'VC backed, investor backed' (SpaceCom 2026 interview). I couldn't verify the investors or fund size. Source: <https://www.youtube.com/watch?v=2XPwwy5EHW8>",
    ])

    h(doc, "RULES FOR THIS MEETING", RED)
    bullets(doc, [
        "Partnership only. Our background check (10/20) rated Space Grove CAUTION: new company (2 staff on LinkedIn), no public financials or ownership. Ask her for investor INTRODUCTIONS only, not for Space Grove to invest; don't commit money or exclusivity.",
        "Science fair flights are sold to sponsors, never given away. No price per student payload in the meeting.",
        "No pricing, no valuation, no investor names. If asked about the raise: 'raising a $10M Series A.' Nothing about a reverse merger, $215M or STRATFI.",
        "Don't call BLK3 'FAA-licensed.' Say Exos has flown licensed suborbital flights with full recovery and is moving to the FAA's newer launch rules (Part 450).",
        "BLK3 timing: 'early 2027.' Never say a company is 'on our BLK3 manifest.'",
        "Scorpius: 'Exos owns 18% of Scorpius Space Launch Company.' Nothing else.",
    ])

    doc.save(OUT); print(OUT)


if __name__ == "__main__":
    build()
