#!/usr/bin/env python3
"""Client-facing deck for the AstroX meeting (Fri 10/16/2026 11:00 JST, Ueno office). CAPABILITY-LEVEL ONLY:
no technical data, drawings, performance numbers or pricing (ITAR — any technical exchange needs a TAA first;
Empowered Official John Quinn, reviewer Denise Shirey). Same Exos styling as the SBI/650 decks.
Wording rules: Part 431 license transitioning to Part 450; BLK3 first commercial flight early 2027 (not yet flown);
SARGE commercial payloads since 2018; payloads back within minutes of landing (recovered in the rocket)."""
import os
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
WS="/Users/exomac/.openclaw/workspace"
BGIMG=os.path.join(WS,"drafts","_exos_slide_bg.png"); LOGO=os.path.join(WS,"drafts","_exos_white_official.png")
OUT=os.path.join(WS,"drafts","Exos_AstroX_Meeting_Deck_2026-10-16.pptx")
BG=RGBColor(0x0d,0x0d,0x0f); PANEL=RGBColor(0x17,0x18,0x1c); ORANGE=RGBColor(0xF2,0x6A,0x21)
WHITE=RGBColor(0xFF,0xFF,0xFF); GREY=RGBColor(0xD2,0xD2,0xD7); MUTE=RGBColor(0x9a,0x9a,0xa1)
p=Presentation(); p.slide_width=I(13.333); p.slide_height=I(7.5); BL=p.slide_layouts[6]; n=[0]
def tb(s,x,y,w,h,text,size=14,color=GREY,bold=False,align=PP_ALIGN.LEFT):
    t=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=t.text_frame; tf.word_wrap=True
    for i,line in enumerate(text if isinstance(text,list) else [text]):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.alignment=align
        r=para.add_run(); r.text=line; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; r.font.name="Calibri"
        para.space_after=Pt(size*0.45)
def rect(s,x,y,w,h,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,I(x),I(y),I(w),I(h)); sh.adjustments[0]=0.05
    sh.fill.solid(); sh.fill.fore_color.rgb=PANEL; sh.shadow.inherit=False
    if line: sh.line.color.rgb=line; sh.line.width=Pt(1.25)
    else: sh.line.fill.background()
def slide(kicker,title,notes=""):
    s=p.slides.add_slide(BL); n[0]+=1
    s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    b=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,I(0.12),I(7.5)); b.fill.solid(); b.fill.fore_color.rgb=ORANGE; b.line.fill.background()
    tb(s,0.6,0.35,10,0.4,kicker,12,ORANGE,True); tb(s,0.6,0.7,10.6,0.9,title,30,WHITE,True)
    s.shapes.add_picture(LOGO,I(11.35),I(0.38),height=I(0.42))
    tb(s,0.6,7.0,10,0.3,"EXOS AEROSPACE  ×  ASTROX  ·  16 OCTOBER 2026",9,MUTE)
    tb(s,12.2,7.0,0.6,0.3,f"{n[0]:02d}",9,MUTE,align=PP_ALIGN.RIGHT)
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s
def card(s,x,y,w,h,head,body,hs=19,bs=16,accent=False):
    rect(s,x,y,w,h,ORANGE if accent else None)
    tb(s,x+0.25,y+0.18,w-0.5,0.5,head,hs,ORANGE if accent else WHITE,True); tb(s,x+0.25,y+0.78,w-0.5,h-0.9,body,bs,GREY)
def stat(s,x,y,w,big,small):
    rect(s,x,y,w,1.75); tb(s,x+0.2,y+0.12,w-0.4,0.8,big,36,ORANGE,True); tb(s,x+0.2,y+0.9,w-0.4,0.85,small,13.5,GREY)

# 1 Title
s=p.slides.add_slide(BL); n[0]+=1
s.shapes.add_picture(BGIMG,0,0,I(13.333),I(7.5)); s.shapes.add_picture(LOGO,I(0.8),I(0.7),height=I(0.7))
tb(s,0.8,2.2,11.5,0.5,"EXOS AEROSPACE  ×  ASTROX",16,ORANGE,True)
tb(s,0.8,2.75,11.5,1.6,["Bring every test flight home —","recovery and reflight for the rockoon"],38,WHITE,True)
tb(s,0.8,4.75,11.5,0.5,"Discussion with AstroX, Inc.  ·  Tokyo  ·  16 October 2026",18,GREY)
tb(s,0.8,6.35,11.5,0.5,"John Quinn, Co-founder & CEO   ·   Toshi Kumazawa, Japan Business Development",14,GREY)
s.notes_slide.notes_text_frame.text=("Open by thanking Oda-san for meeting again after Hokkaido. Purpose: understand their test program and see where "
  "Exos recovery helps. Keep everything at capability level — no drawings, numbers or pricing today (U.S. export rules). Offer an NDA for next steps.")

# 2 Who we are
s=slide("01  |  EXOS AEROSPACE","A U.S. reusable-launch company built around recovery",
  "Say 'one of only three' exactly as written. Our license: Part 431, transitioning to Part 450. BLK3 has not flown yet — first commercial flight early 2027.")
stat(s,0.6,1.85,2.9,"1 of 3","U.S. FAA-licensed reusable rocket launch providers — with SpaceX and Blue Origin")
stat(s,3.7,1.85,2.9,"2018","commercial payloads flown on SARGE, our first reusable vehicle")
stat(s,6.8,1.85,2.9,"25+ yrs","reusable-vehicle heritage (Armadillo Aerospace lineage)")
stat(s,9.9,1.85,2.85,"Minutes","from landing to payload back in the customer's hands on BLK3")
card(s,0.6,3.95,12.15,2.75,"Our mission: make every space vehicle reusable",[
  "• Drogue and parafoil recovery flown on SARGE since 2018",
  "• BLK3 — our latest vehicle, the enhanced successor to SARGE — designed to be recovered and reflown; first commercial flight early 2027",
  "• Operating under a Part 431 FAA license, transitioning to Part 450  ·  Exos Aerospace Japan K.K. forming"],accent=True)

# 3 Their challenge
s=slide("02  |  YOUR TEST PROGRAM, AS WE UNDERSTAND IT","Every test flight releases valuable hardware at altitude",
  "Based on public information only (AstroX's Dec 2025 separation test; May 2026 JAXA attitude-control demonstration). Ask them to correct us — this is a listening slide. Questions: how do you recover today, what does a lost set cost, where do you fly (land or sea)?")
card(s,0.6,1.85,3.95,3.6,"Gondola / launcher platform","Lifted to the stratosphere and left behind after the rocket is released.")
card(s,4.7,1.85,3.95,3.6,"Attitude-control device","Your JAXA-demonstrated pointing system — expensive, and ideal to fly again.")
card(s,8.8,1.85,3.95,3.6,"Test rockets","Your first 'reach space' attempts — the flight data and hardware are most valuable after landing.")
rect(s,0.6,5.65,12.15,1.05,ORANGE)
tb(s,0.85,5.85,11.7,0.7,"Our question for you: which of these do you recover today — and what would getting them back be worth?",17,WHITE,True)

# 4 How we can help
s=slide("03  |  HOW EXOS CAN HELP","Three ways to make your program more reusable",
  "Capability level only. If they ask how — say the engineering conversation happens under NDA and a U.S. export authorization (TAA), which we would file. No numbers.")
card(s,0.6,1.85,3.95,4.85,"1 · Platform recovery",["Bring the gondola, launcher platform and attitude-control device back under a guided parafoil so they can be reflown.",
  "","Lower cost per flight, faster test cadence."],accent=True)
card(s,4.7,1.85,3.95,4.85,"2 · Test-rocket recovery",["Drogue and parafoil recovery for your early test rockets.",
  "","Get the hardware and data back for post-flight analysis."])
card(s,8.8,1.85,3.95,4.85,"3 · Bridge flights on BLK3",["Fly your customers' payloads to space on Exos BLK3 while the rockoon matures.",
  "","Keeps your customers with you; payloads back within minutes of landing."])

# 5 Start small
s=slide("04  |  A SIMPLE WAY TO START","Start small, prove the value, then grow",
  "Phase 1 is deliberately small. Do NOT quote a price today; say we will send a short proposal after the NDA. Step 2 (TAA) is our job, not theirs.")
for i,(h,b) in enumerate([("STEP 1","Mutual NDA, and a working session on your next stratospheric flights: dates, location, what you want back."),
                          ("STEP 2","Exos files the U.S. export authorization needed for joint engineering work."),
                          ("STEP 3","A small paid Phase 1: recovery plan for your next flight campaign, then recovery on flight.")]):
    card(s,0.6+i*4.1,1.95,3.9,3.4,h,b,19,16.5,accent=(i==0))
card(s,0.6,5.5,12.15,1.2,"Also open to","Flying AstroX customer payloads on BLK3 · joint work with JAXA-linked programs · Japan-based operations through Exos Aerospace Japan",17,15)

# 6 Next steps
s=slide("05  |  NEXT STEPS","Let's bring your next flight home",
  "Close on one dated next step: NDA this week, working session within 2-3 weeks. Toshi follows up in Japanese (MC #3810).")
card(s,0.6,1.85,6.0,4.4,"Proposed next steps",["1. Exchange a mutual NDA","2. Working session on your next flight campaign",
  "3. Exos sends a short Phase 1 proposal","4. Toshi coordinates in Japan"],accent=True)
card(s,6.8,1.85,5.95,4.4,"Contacts",["John Quinn — Co-founder & CEO","jquinn@exosaero.com  ·  +1-972-740-8355","",
  "Toshi Kumazawa — Japan Business Development","toshihisa@exosaero.com","","www.exosaero.com"])
tb(s,0.6,6.45,12.1,0.4,"This presentation contains general capability information only; no technical data is provided.",10,MUTE)
p.save(OUT); print("WROTE",OUT)
