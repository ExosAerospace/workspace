"""Exos counter-redline of Copernic's Oct 8, 2026 proposed redline (v5) of the Wet Lease
O&M Agreement + updated appendices. MC #3055.

Base text = Copernic's Oct 8 redline with Copernic's changes ACCEPTED (scripts/
copernic_v5_reconstruct.py). Exos revisions are real Word tracked changes (author
"Exos Aerospace") with margin comments. New Appendix H = vehicle order terms.

Inputs: Jon's 10/11 email (positions of record), Brennan Plaetzer's 10/14 reply
($5.0M path selected; one-replacement point; term from successful Pathfinder;
escrow for final 15%; maintenance budget).

Run with a python that has PyMuPDF + python-docx>=1.2 (e.g. /tmp/pdfenv/bin/python).
"""
import copy
import datetime as dt
import re
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.run import Run

sys.path.insert(0, "/Users/exomac/.openclaw/workspace/scripts")
from copernic_v5_reconstruct import paragraphs  # noqa: E402

OUT = "/Users/exomac/.openclaw/workspace/drafts/Copernic_Wet_Lease_Exos_Counter_Redline_2026-10-17.docx"
AUTHOR, INITIALS = "Exos Aerospace", "EXOS"
DATE = "2026-10-17T10:00:00Z"
_id = [9000]


def nid():
    _id[0] += 1
    return str(_id[0])


# ---------------------------------------------------------------- run helpers
def rpr(font, size, bold=False, italic=False):
    r = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    name = "Times New Roman" if font == "serif" else "Arial"
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        f.set(qn(a), name)
    r.append(f)
    if bold:
        r.append(OxmlElement("w:b"))
    if italic:
        r.append(OxmlElement("w:i"))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r.append(sz)
    return r


def make_run(text, fmt, deleted=False):
    r = OxmlElement("w:r")
    r.append(rpr(**fmt))
    parts = text.split("\n")
    for i, part in enumerate(parts):
        if i:
            r.append(OxmlElement("w:br"))
        if part:
            t = OxmlElement("w:delText" if deleted else "w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = part
            r.append(t)
    return r


def wrap(kind, run_el):
    w = OxmlElement("w:" + kind)
    w.set(qn("w:id"), nid())
    w.set(qn("w:author"), AUTHOR)
    w.set(qn("w:date"), DATE)
    w.append(run_el)
    return w


def mark_para_inserted(p):
    pPr = p._p.get_or_add_pPr()
    rp = pPr.find(qn("w:rPr"))
    if rp is None:
        rp = OxmlElement("w:rPr")
        pPr.append(rp)
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), nid())
    ins.set(qn("w:author"), AUTHOR)
    ins.set(qn("w:date"), DATE)
    rp.append(ins)


def mark_para_deleted(p):
    pPr = p._p.get_or_add_pPr()
    rp = pPr.find(qn("w:rPr"))
    if rp is None:
        rp = OxmlElement("w:rPr")
        pPr.append(rp)
    d = OxmlElement("w:del")
    d.set(qn("w:id"), nid())
    d.set(qn("w:author"), AUTHOR)
    d.set(qn("w:date"), DATE)
    rp.append(d)


def spacing(p, after=6):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)


# ---------------------------------------------------------------- edit engine
def segments(base, ops):
    """ops: ("rep", old, new) | ("ins", anchor, text) | ("del", old) | ("app", text)
    Returns [(kind, text, comment)] with kind in eq/ins/del."""
    segs = [("eq", base, None)]

    def split_at(old):
        hits = [(i, s) for i, s in enumerate(segs) if s[0] == "eq" and old in s[1]]
        assert len(hits) == 1 and hits[0][1][1].count(old) == 1, f"anchor not unique/found: {old!r}"
        i, (_, txt, _) = hits[0]
        a, b = txt.split(old)
        return i, a, b

    for op in ops:
        kind, cmt = op[0], op[-1] if isinstance(op[-1], dict) else None
        cmt = cmt["c"] if cmt else None
        if kind == "rep":
            i, a, b = split_at(op[1])
            segs[i:i + 1] = [("eq", a, None), ("del", op[1], cmt), ("ins", op[2], None), ("eq", b, None)]
        elif kind == "del":
            i, a, b = split_at(op[1])
            segs[i:i + 1] = [("eq", a, None), ("del", op[1], cmt), ("eq", b, None)]
        elif kind == "ins":
            i, a, b = split_at(op[1])
            segs[i:i + 1] = [("eq", a + op[1], None), ("ins", op[2], cmt), ("eq", b, None)]
        elif kind == "app":
            segs.append(("ins", op[1], cmt))
    return [s for s in segs if s[1]]


class Builder:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12)
        sec = self.doc.sections[0]
        from docx.shared import Inches
        sec.left_margin = sec.right_margin = Inches(1)
        sec.top_margin = sec.bottom_margin = Inches(1)
        hp = sec.header.paragraphs[0]
        hr = hp.add_run("Exos counter-redline to Copernic proposed redline v5 (Oct 8, 2026)  |  Exos draft October 17, 2026  |  For discussion, not binding")
        hr.font.size = Pt(8)
        hr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        self._footer_page_numbers(sec)

    def _footer_page_numbers(self, sec):
        fp = sec.footer.paragraphs[0]
        r = fp.add_run("Page ")
        r.font.size = Pt(8)
        for code in ("PAGE",):
            run = fp.add_run()
            run.font.size = Pt(8)
            b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
            t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = code
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
            run._r.append(b); run._r.append(t); run._r.append(e)

    def comment(self, p, run_el, text):
        self.doc.add_comment(Run(run_el, p), text=text, author=AUTHOR, initials=INITIALS)

    def para(self, segs, fmt, page_break=False, inserted=False, after=6, style=None):
        p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        spacing(p, after)
        if page_break:
            br = OxmlElement("w:r")
            b = OxmlElement("w:br"); b.set(qn("w:type"), "page")
            br.append(b)
            p._p.append(br)
        for kind, text, cmt in segs:
            fmt_k = dict(fmt)
            if kind == "eq":
                r = make_run(text, fmt_k)
                p._p.append(r)
                anchor = r
            elif kind == "ins":
                r = make_run(text, fmt_k)
                p._p.append(wrap("ins", r))
                anchor = r
            else:
                r = make_run(text, fmt_k, deleted=True)
                p._p.append(wrap("del", r))
                anchor = r
            if cmt:
                self.comment(p, anchor, cmt)
        if inserted:
            mark_para_inserted(p)
        return p

    def new_para(self, text, fmt, comment=None, after=6, page_break=False):
        """Wholly inserted (tracked) paragraph."""
        return self.para([("ins", text, comment)], fmt, page_break=page_break, inserted=True, after=after)

    def plain(self, text, fmt, after=6, page_break=False, align=None):
        p = self.para([("eq", text, None)], fmt, page_break=page_break, after=after)
        if align is not None:
            p.alignment = align
        return p

    def table(self, rows, fmt, inserted=False, cell_ops=None, widths=None, header=True):
        t = self.doc.add_table(rows=0, cols=len(rows[0]))
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            if inserted:
                trPr = t.rows[ri]._tr.get_or_add_trPr()
                ins = OxmlElement("w:ins")
                ins.set(qn("w:id"), nid()); ins.set(qn("w:author"), AUTHOR); ins.set(qn("w:date"), DATE)
                trPr.append(ins)
            for ci, val in enumerate(row):
                cp = cells[ci].paragraphs[0]
                f = dict(fmt, bold=(header and ri == 0))
                ops = (cell_ops or {}).get((ri, ci))
                segs = segments(val, ops) if ops else [("ins" if inserted else "eq", val, None)]
                for kind, text, cmt in segs:
                    r = make_run(text, f, deleted=(kind == "del"))
                    cp._p.append(r if kind == "eq" else wrap(kind, r))
                    if cmt:
                        self.comment(cp, r, cmt)
                if inserted:
                    mark_para_inserted(cp)
        self.doc.add_paragraph()
        return t


# ---------------------------------------------------------------- content
C = lambda s: {"c": s}  # noqa: E731

EDITS = {
    # Section 1.1 appendix list
    "Appendix G - Authority, Representation & Pricing\nThe updated": [
        ("ins", "Appendix G - Authority, Representation & Pricing",
         "\nAppendix H - Vehicle Order, Payment Milestones & Commercial Term",
         C("New Appendix H carries the vehicle order terms both sides aligned on (path, milestones, schedule, refunds, fee, term).")),
        ("rep", "Appendix G forms part of this Agreement.", "Appendices G and H form part of this Agreement."),
    ],
    "Section 1.2 - Effective Date and Term": [
        ("rep", "through December 31, 2030, as provided in Section 8.1,",
         "until the end of the Commercial Term under Appendix H-7, as provided in Section 8.1,",
         C("Per Brennan's Oct 14 email: the commercial term runs from successful Pathfinder, not a fixed calendar date. The 12/31/2030 references are removed throughout.")),
    ],
    '"Copernic-funded Vehicle" means': [
        ("ins", "including any vehicle supplied by Owner/Operator as its replacement",
         " under Appendix G-1a or ordered by Client under Appendix H-1(c)"),
        ("app", '\n"Commencement Date", "Commercial Term", "Commercial Launch", "Material Damage", "protection period" and "Successful Pathfinder Flight" have the meanings given in Appendices G-1a and H-7.'),
    ],
    "1. Rocket Provision:": [
        ("ins", "Post-Pathfinder vehicle hull/loss insurance is optional for Client, as provided in Appendix G-1a;",
         " after the G-1a protection period, the risk of vehicle loss sits with Client, which may insure or self-retain it;",
         C("Both sides aligned (Oct 14) on Copernic retaining post-Pathfinder vehicle risk. This states the risk allocation alongside the optional-insurance point.")),
    ],
    "Section 8.1 - Term.": [
        ("rep", "will continue until 12/31/2030 inclusive", "will continue until the end of the Commercial Term under Appendix H-7"),
        ("rep", "may be extended for additional five-year periods upon mutual agreement",
         "may be renewed in one-year or multi-year increments upon mutual written agreement",
         C("Renewal in one-year or multi-year increments by mutual agreement (Exos Oct 11 email).")),
    ],
    "Section 8.6 - Termination By Owner/Operator.": [
        ("rep", "for Client's material breach only if Client has not cured the breach within 30 days after written notice specifying it, followed by 15 days' prior written notice of termination.",
         "for cause only, on 15 days' prior written notice of termination, where: (a) Client fails to pay an undisputed amount and does not cure within 10 Business Days after written notice; (b) Client commits any other material breach and does not cure it within 30 days after written notice specifying it; (c) a Force Majeure Event prevents performance for more than 120 days; (d) a required FAA license or approval is lost or denied, or an export or other regulatory prohibition prevents performance; or (e) Client suffers a Bankruptcy.",
         C("Cause-only termination after funding, as aligned Oct 14. The causes are the ones in Exos's Oct 1 letter and Oct 11 email; the 120-day Force Majeure trigger mirrors Client's right in Section 8.2.")),
        ("rep", "Section G-4 governs any legally required or safety-related suspension.", "Section G-4 governs suspension."),
    ],
    "Section 8.7 - Termination By Owner/Operator Without Cause.": [
        ("app", " Termination by Client under Section 8.5 does not entitle Client to a refund of payments for completed milestones; Appendix H-5 governs refunds."),
    ],
    "(d) Cost: This includes": [
        ("rep", "Post-Pathfinder vehicle casualty funding requires a separate written agreement; electing not to insure the vehicle does not itself obligate Client to fund another vehicle or its repair.",
         "After the protection period in Appendix G-1a ends, vehicle loss and casualty repair are at Client's risk as provided in G-1a(c); electing not to insure the vehicle is not a breach.",
         C("Without this, no one would carry a post-Pathfinder loss. Copernic keeps the choice to insure or self-retain.")),
    ],
    "(f) The limitations and exclusive-remedy": [
        ("rep", "do not limit Owner/Operator's express Pathfinder repair or replacement obligations under G-1a, including the reasonable direct cost of that repair or replacement if Owner/Operator fails to perform, or payment or return of Payload Revenue belonging to Client under B-4, G-1b(i) or G-1c.",
         "do not limit: (i) repayment of unearned vehicle funding under Appendix H-5; (ii) Owner/Operator's Pathfinder repair-or-replace obligation under G-1a (repair to the contracted specification or one replacement per ordered vehicle), including the reasonable direct cost of that repair or replacement if Owner/Operator fails to perform, with exposure no greater than the price of the ordered vehicle, no consequential damages or lost Payload Revenue, and excluding loss caused by Client payloads or Client-furnished items; or (iii) remittance of Payload Revenue belonging to Client and received by Owner/Operator under B-4, G-1b(i) or G-1c.",
         C("Carve-outs from the general cap as set out in Exos's Oct 11 email: unearned funding and a bounded Pathfinder obligation. Remittance of Client's own Payload Revenue stays outside the cap as Copernic proposed.")),
    ],
    "IN WITNESS WHEREOF": [
        ("rep", "Appendices A through G", "Appendices A through H"),
    ],
    "A-4. Not included": [
        ("rep", "post-Pathfinder vehicle hull/loss coverage is optional for Client",
         "after the G-1a protection period, vehicle loss risk sits with Client, which may insure or self-retain it"),
    ],
    "B-1. Structure.": [
        ("ins", "Substitute Flight capacity under B-4", ", except as provided in G-1c"),
    ],
    "B-2. Vehicle commercialization": [
        ("rep", "(from the June 2026 offer)", "(Protected Pathfinder path, selected by Client)",
         C("Per Brennan's Oct 14 email: the $5.0M protected Pathfinder path is written in as the selected path. Copernic confirms the selection by signing (H-1).")),
    ],
    "Option A — one (1) ordered BLK3 vehicle": [
        ("rep", "Option A — one (1) ordered BLK3 vehicle: $5.0M per ordered vehicle.",
         "Price: $5.0M per ordered BLK3 vehicle, payable on the milestones in Appendix H-2."),
        ("rep", "on the Pathfinder (first) flight (see Appendix G, G-1a)",
         "on each ordered vehicle's Pathfinder Flight and keeps a second vehicle built and on standby (see Appendix G-1a and Appendix H-1)"),
    ],
    "Option B — two (2) ordered BLK3 vehicles": "DELETE",
    "No revenue split. Client (Copernic) books": [
        ("ins", "Client's contracted Substitute Flight capacity", ", subject to G-1c"),
    ],
    "B-6. Annual Operating Fee. Covers": [
        ("del", "approved limited-life replacements, "),
        ("ins", "and readiness/program management.",
         " Consumables and wear parts are billed at actual cost against the approved maintenance and parts budget under Appendix H-6, not inside this fee.",
         C("Splits the binding base fee from the maintenance and parts budget, as aligned Oct 14. The updated first-year budget is in H-8.")),
    ],
    "Annual Operating Fee — one vehicle": [
        ("del", " (Option B)"),
        ("app", " The fee is binding and adjusted only by US CPI; any increase above CPI requires Client's written agreement. The fee for each vehicle begins on its Commencement Date (Appendix H-7)."),
    ],
    "Vehicle hull/loss insurance after Pathfinder is optional for Client under G-1a.": [
        ("rep", "Any elected coverage, premiums and uninsured post-Pathfinder vehicle casualty funding must be separately agreed in writing.",
         "After the G-1a protection period, vehicle loss risk sits with Client (insure or self-retain); premiums for any coverage Client elects are Client's cost."),
    ],
    "D-1. Term.": [
        ("rep", "through 12/31/2030, renewable for additional five-year periods by mutual written agreement",
         "until the end of the Commercial Term under Appendix H-7, renewable in one-year or multi-year increments by mutual written agreement"),
    ],
    "D-3. Termination triggers": [
        ("rep", "by Owner/Operator for uncured material breach under Section 8.6",
         "by Owner/Operator for cause under Section 8.6 (uncured non-payment or material breach, Force Majeure over 120 days, loss or denial of FAA licensing or an export/regulatory prohibition, or Client Bankruptcy)"),
    ],
    "E-3. Liability": [
        ("rep", "apply to Pathfinder repair/replacement and Client Payload Revenue",
         "apply to repayment of unearned vehicle funding, the capped Pathfinder repair/replacement obligation, and Client Payload Revenue received by Owner/Operator"),
    ],
    "E-4. Control of the flight asset.": [
        ("app", " Client's exclusive capacity rights in a Copernic-funded Vehicle run for the Commercial Term under Appendix H-7, and end as provided in G-1a(c) if that vehicle is lost after the protection period and Client does not fund its repair or replacement."),
    ],
    "“Pathfinder Flight” means the first flight": [
        ("ins", "“Pathfinder Flight” means the first flight of each ordered vehicle.",
         " It begins at lift-off and ends when the vehicle is recovered. A scrub, hold or abort before lift-off is not a Pathfinder Flight, and before lift-off the vehicle is at Exos's risk. A “Successful Pathfinder Flight” is a Pathfinder Flight from which the vehicle is recovered without loss or Material Damage. “Material Damage” means damage that cannot be repaired to the contracted specification in the ordinary course of post-flight refurbishment.",
         C("Defines when the Pathfinder Flight starts and ends, and what 'successful' and 'materially damaged' mean, so the protection and the start of the commercial term are unambiguous.")),
        ("rep", "is lost or materially damaged during its Pathfinder Flight", "is lost or suffers Material Damage during its Pathfinder Flight"),
        ("ins", "without additional vehicle funding from Copernic",
         ", either by repairing it or by assigning the standby vehicle that Exos keeps built under the Protected Pathfinder path (Appendix H-1). The standby vehicle is targeted to fly the next scheduled flight within 90 days, subject to range availability and FAA requirements"),
        ("del", "Vehicle hull/loss insurance after Pathfinder is optional for Client; any election of coverage and any funding of an uninsured post-Pathfinder vehicle loss must be separately agreed in writing. Failure to elect vehicle hull/loss insurance is not a breach and does not itself require additional vehicle funding. ",
         C("Moved to new G-1a(c) below, with the post-Pathfinder risk allocation both sides aligned on.")),
    ],
    "Client retains the same Payload Revenue rights": [
        ("ins", "as on its Copernic-funded Vehicle.",
         " Client pays the per-flight operating cost and Reimbursable Costs allocable to that capacity on the same basis as for a flight of its Copernic-funded Vehicle."),
        ("ins", "grant Client rights over unrelated payloads on the flight.",
         " Where a Substitute Flight carries Client's committed payloads while a replacement vehicle is being built under Appendix H-1(c), Owner/Operator receives 25% of the related Payload Revenue, and no less than $250,000 per flight, in place of the per-flight operating cost. Substitute carriage is subject to available capacity and to export and regulatory approvals.",
         C("Substitute carriage on another Exos vehicle is priced like a flight of Copernic's own vehicle. The 25% / $250K-minimum terms apply only to gap flights during a Client-funded rebuild after a post-Pathfinder loss, as in Exos's Oct 1 and Oct 11 terms.")),
    ],
    "Any material breach of G-1 through G-3": [
        ("rep", "Any material breach of G-1 through G-3 is subject to the notice and cure procedure in Section 8.6.",
         "Any material breach of G-1 is subject to the notice and cure procedure in Section 8.6. A breach of G-2 or G-3 (presenting Copernic's raise as an Exos offering or endorsement, or any foreign-person financing, control, transfer or access outside Exos's CFIUS/export process) entitles Owner/Operator to suspend the affected Services immediately on written notice and, if the breach is not cured within 30 days after that notice, to terminate under Section 8.6.",
         C("Commercial disputes under G-1 stay on notice-and-cure, as Copernic proposed. Immediate suspension remains for the securities and export/CFIUS guardrails (G-2/G-3), as stated in Exos's Oct 11 email; Exos cannot carry that regulatory exposure through a cure period.")),
        ("rep", "Owner/Operator may immediately suspend only the affected Services", "Otherwise, Owner/Operator may immediately suspend only the affected Services"),
    ],
}

G1A_NEW = [
    ("(a) One replacement per ordered vehicle. Exos's Pathfinder obligation for each ordered vehicle ends when Exos has provided one repaired or replacement vehicle. The replacement vehicle's first flight is not a further Pathfinder Flight and is not protected by this G-1a. The protection period for each ordered vehicle therefore runs from lift-off of its Pathfinder Flight until the earlier of (i) a Successful Pathfinder Flight or (ii) lift-off of the replacement vehicle's first flight.",
     "Brennan asked (Oct 14) that Exos state this explicitly if its obligation ends after one replacement. It does. The $1.0M buys a second vehicle built and waiting, so a Pathfinder loss costs Copernic nothing and the program continues within about 90 days. It does not make Exos the insurer of every later first flight. A loss on or after the replacement's first flight is post-Pathfinder risk under G-1a(c), which Copernic can insure."),
    ("(b) Limits. Exos's exposure under this G-1a for each ordered vehicle is no greater than the price of that ordered vehicle. It excludes consequential damages and lost Payload Revenue, and it excludes loss or damage caused by a Client payload or a Client-furnished item.", None),
    ("(c) Risk after the protection period. After the protection period ends, the risk of loss of or damage to the vehicle sits with Client. Client may insure that risk or self-retain it. Vehicle hull/loss insurance is optional, and not electing it is not a breach. Exos has no repair or replacement obligation after the protection period unless Client funds it: routine repairs through the approved maintenance and parts budget (Appendix H-6), and casualty repair or replacement as Client elects under Appendix H-1(c). If Client has not elected in writing to fund the repair or replacement of a lost vehicle within 90 days after the loss, Client's exclusive capacity rights in that vehicle and the Annual Operating Fee for that vehicle end on that date; amounts accrued before that date remain payable.",
     "Post-Pathfinder risk sits with Copernic (aligned Oct 14). The 90-day election point stops a lost vehicle from carrying exclusive rights and fees indefinitely."),
]

E1_CELL_OPS = {
    (8, 1): [("app", " (one replacement per ordered vehicle)")],
    (8, 2): [("rep", "Post-Pathfinder vehicle hull/loss coverage optional",
              "Vehicle loss risk after the G-1a protection period (insure or self-retain)")],
}

APPX_H = [
    ("h", "Appendix H — Vehicle Order, Payment Milestones & Commercial Term"),
    ("p", "To the Wet Lease Operation and Maintenance Agreement between Copernic Space Corporation (Client) and Exos Aerospace Systems & Technologies, Inc. (Owner/Operator). Applies to each vehicle ordered under this Agreement.",
     "New appendix. Writes in the terms from Exos's Oct 11 email that Copernic confirmed in principle on Oct 14, plus the two clarifications Brennan requested."),
    ("s", "H-1. Selected path and vehicle price."),
    ("p", "(a) Client selects the Protected Pathfinder path offered in Owner/Operator's letter of October 1, 2026: $5.0M per ordered vehicle. Client confirms this selection by signing this Agreement. The $4.0M one-vehicle path in that letter is not elected, and its terms do not apply.", None),
    ("p", "(b) The price covers manufacture, integration and acceptance testing (Appendix C) of the ordered vehicle, its Pathfinder Flight, and Owner/Operator's Pathfinder protection under G-1a, for which Owner/Operator keeps a second vehicle built and on standby. The price is the same at any launch location. Operations other than at Spaceport America require a separate written agreement under A-5, with site scope (mobilization, ground support equipment, transport and logistics) billed as Reimbursable Costs.", None),
    ("p", "(c) Replacement after the protection period. If a Copernic-funded Vehicle is lost after its G-1a protection period ends, Client may order a flight-ready replacement for $4.0M, payable on the milestone percentages in H-2, with a build time of up to 12 months from receipt of the replacement deposit. G-1c applies to Client's committed payloads during that build. For contracting, insurance and budgeting purposes, the replacement cost of a flight-ready vehicle is $4.0M.", None),
    ("s", "H-2. Payment milestones (per ordered vehicle)."),
    ("t", [["Milestone", "Share", "Amount ($5.0M vehicle)", "Due"],
           ["Deposit", "50%", "$2,500,000", "On order"],
           ["Half-build", "25%", "$1,250,000", "On certified completion of half-build"],
           ["Pre-flight testing complete", "10%", "$500,000", "On certified completion of pre-flight testing"],
           ["Flight readiness", "15%", "$750,000", "Paid into escrow on completion of the Flight Readiness Review; released to Owner/Operator at lift-off of the Pathfinder Flight"]]),
    ("p", "Milestone invoices are payable within 15 days (Section 5.5) and are accompanied by the evidence in H-4. The escrow agent and escrow terms are agreed in writing before the Flight Readiness Review, and escrow fees are shared equally. If the order ends before lift-off, the escrowed amount is released under H-5. The same percentages apply to a replacement ordered under H-1(c).",
     "Final 15%: escrow at the Flight Readiness Review, released at Pathfinder lift-off, as Brennan confirmed works (Oct 14)."),
    ("s", "H-3. Schedule and outside date."),
    ("p", "Each period runs from receipt of the applicable payment or completion of the preceding milestone, not from fixed calendar dates:\n• Deposit received to half-build: 6 months.\n• Half-build to pre-flight testing complete: 4 months.\n• Testing complete to Pathfinder Flight: target 2 weeks, within a 60-day window for range activation and FAA scheduling.", None),
    ("p", "Outside date: the Pathfinder Flight occurs within 12 months of receipt of the deposit. The outside date extends day-for-day for excused delay: a Force Majeure Event, FAA licensing or range availability, export approvals, or delay caused by Client, including late payment, late payloads or a late manifest. If the Pathfinder Flight has not occurred by the outside date as extended, Client may terminate the vehicle order on 30 days' written notice, unless the Pathfinder Flight occurs within that notice period, and H-5(b) then applies.",
     "The remedy for a missed outside date (Client termination with a full refund of incomplete milestones) is new text; it gives the 12-month date the effect both sides intended."),
    ("s", "H-4. Milestone evidence."),
    ("p", "For each milestone, Owner/Operator issues a written milestone completion certificate with full photo documentation and a summary of the build and test records. Client may visit and observe on reasonable notice, and receives read-only program-status access in Owner/Operator's Mission Control system. All access is subject to U.S. export-control limits on technical data (A-5).", None),
    ("s", "H-5. Refunds."),
    ("p", "(a) Owner/Operator keeps payments for completed and verified milestones.\n(b) If the vehicle is not completed because of Owner/Operator default, including termination under H-3, Client receives a full refund of payments for incomplete milestones, with no deductions.\n(c) If non-completion results from an excused event (a Force Majeure Event, denial of FAA licensing, or an export or regulatory prohibition) or from termination by Client, Client receives a refund of payments for incomplete milestones less documented, non-cancellable third-party costs for materials and components in that milestone's budget. Owner/Operator provides the vendor invoices.\n(d) Refunds are paid within 30 days after the order ends.", None),
    ("s", "H-6. Annual Operating Fee and maintenance and parts budget."),
    ("p", "(a) Base fee: $90,000 per year for one vehicle or $165,000 per year for two, adjusted by US CPI from the second Contract Year and binding. Any increase beyond CPI requires Client's written agreement. The fee for each vehicle begins on its Commencement Date (H-7) and is pro-rated for partial years.\n(b) Maintenance and parts budget: by month 9 of each Contract Year, Owner/Operator issues the next year's maintenance and parts budget (consumables and wear parts). These items are billed at actual cost against the approved budget, and unspent budget is not charged. Spending above the approved budget requires Client's prior written approval, except for safety-of-flight or regulatory-required items, which Owner/Operator documents.\n(c) The first-year maintenance and parts budget is set out in H-8.", None),
    ("s", "H-7. Commercial Term."),
    ("p", "(a) The \"Commencement Date\" for each Copernic-funded Vehicle is the date of the first flight of that vehicle (including any replacement) from which it is recovered without loss or Material Damage. Vehicle build, testing, FAA approval and range scheduling before the Commencement Date do not count against the Commercial Term.\n(b) The \"Commercial Term\" runs for five years from the first Commencement Date. Client's exclusive capacity rights continue through the later of (i) the end of that five-year period or (ii) completion of twenty (20) Commercial Launches in total under this Agreement, provided Client is not in uncured default. A \"Commercial Launch\" is a launch of a Copernic-funded Vehicle after its Commencement Date that carries Client's payload capacity; Pathfinder Flights do not count.\n(c) The parties may renew by mutual written agreement in one-year or multi-year increments.",
     "As Brennan requested (Oct 14): the five-year term starts on successful Pathfinder, and exclusivity runs to the later of five years or 20 commercial launches while Copernic is not in uncured default. The Annual Operating Fee continues through any extension."),
    ("s", "H-8. First-year maintenance and parts budget (per vehicle, planning)."),
    ("t", [["Item", "Basis", "Budget"],
           ["Flight battery set (2 × 6S 22,000 mAh)", "Annual replacement", "$540"],
           ["High-temperature seal / o-ring set", "Annual replacement", "[$______]"],
           ["Valve assembly", "On condition; one per year budgeted", "[$______]"],
           ["Pyrotechnic / ordnance items (separation, recovery deployment)", "Shelf-life replacement", "[$______]"],
           ["Pressure-vessel recertification", "Per recertification interval", "[$______]"],
           ["Pressure regulators", "On condition", "[$______]"],
           ["Total", "", "$540 + [$______]"]]),
    ("p", "Planning figures at current vendor pricing, billed at actual cost under H-6(b). Labor, storage, readiness, FAA currency and program management are inside the base fee (B-6). Propellant and pressurant consumed at flight are in the per-flight cost and Reimbursable Costs (B-3, B-5). Insurance is not included. Owner/Operator re-issues this budget before the Pathfinder Flight.",
     "Updated maintenance budget requested Oct 14."),
    ("s", "H-9. Payload deadlines."),
    ("p", "Payload sales closed and payloads delivered for integration testing by L-90 (90 days before the scheduled launch); final launch manifest by L-60; signed FAA cross-waivers by L-90. A payload that misses these dates rolls to the next available flight. That is not a breach by either party.", None),
    ("s", "H-10. Precedence."),
    ("p", "This Appendix H controls any conflicting provision of the Agreement or the other appendices on vehicle price, payment, schedule, refunds, the Annual Operating Fee and the Commercial Term.", None),
]

COVER = [
    ("title", "Copernic Space × Exos Aerospace"),
    ("sub", "Wet Lease — Exos Counter-Redline"),
    ("small", "Responding to Copernic proposed redline v5 (October 8, 2026) and Brennan Plaetzer's email of October 14, 2026  |  Exos draft, October 17, 2026"),
    ("h", "How to read this draft"),
    ("p", "Copernic's October 8 revisions are accepted and appear as clean text, except where Exos has revised them. Exos revisions are tracked (author: Exos Aerospace), and margin comments explain each one. New Appendix H holds the vehicle order terms. Appendix C is unchanged."),
    ("h", "Points from Brennan's October 14 email"),
    ("p", "1. Selected path. The $5.0M Protected Pathfinder path is written in as the selected path (B-2, H-1). Copernic confirms the selection by signing."),
    ("p", "2. Replacement vehicle's first flight. Exos's obligation ends after one replacement per ordered vehicle, as you asked us to state explicitly (G-1a(a)). The standby vehicle is built and waiting, so a Pathfinder loss costs Copernic nothing and the program continues within a target of 90 days. The replacement's first flight is not a further Pathfinder Flight; a loss on or after that flight is post-Pathfinder risk, which Copernic may insure (G-1a(c))."),
    ("p", "3. Commercial term. The five-year term begins on successful Pathfinder. Exclusivity runs through the later of five years or 20 commercial launches, provided Copernic is not in uncured default (H-7). Build, testing, FAA approval and range scheduling do not consume the term."),
    ("p", "4. Final 15%. Paid into escrow at the Flight Readiness Review and released at Pathfinder lift-off (H-2)."),
    ("p", "5. Maintenance budget. The first-year maintenance and parts budget is at H-8."),
    ("p", "6. Insurance broker. Exos is introducing Copernic to its broker separately, for post-Pathfinder hull coverage terms."),
    ("h", "Other Exos revisions"),
    ("p", "Cause-only termination by Exos, with the causes listed (8.6). Post-Pathfinder risk allocation (Wet Lease definition, 9.1(d), A-4, B-6, E-1, G-1a(c)). Bounded carve-outs from the liability cap (11.1(f), E-3). Substitute Flight pricing (G-1c). Immediate suspension kept for the securities and export/CFIUS guardrails only (G-4). Term references conformed to the Commercial Term (1.2, 8.1, D-1)."),
    ("h", "Next steps"),
    ("p", "Copernic reviews and confirms the selected path. Counsel review at ready-to-sign, and U.S. export review before execution."),
    ("small", "This draft is for discussion. It is not a binding agreement. Terms bind only in a definitive agreement executed by both parties, subject to applicable U.S. export and regulatory approvals."),
]


def main():
    P = paragraphs()
    for p in P:
        p["text"] = p["text"].replace("third- party", "third-party")
        if p["page"] >= 34:
            p["font"] = "sans"
            if re.match(r"^G-\d[a-z]?\. ", p["text"]) and len(p["text"]) < 80:
                p["bold"], p["size"] = True, 12
            elif p["size"] == 12 and not p["bold"]:
                p["size"] = 10

    B = Builder()
    sans = lambda size=10, bold=False: {"font": "sans", "size": size, "bold": bold}  # noqa: E731

    # cover (Exos transmittal sheet, untracked; replaces Copernic's "v5" cover)
    for kind, text in COVER:
        if kind == "title":
            B.plain(text, sans(16, True), after=2)
        elif kind == "sub":
            B.plain(text, sans(13, True), after=4)
        elif kind == "small":
            B.plain(text, sans(8), after=10)
        elif kind == "h":
            B.plain(text, sans(11, True), after=4)
        else:
            B.plain(text, sans(10), after=4)

    B.plain("Wet Lease and\nO&M Agreement between\nCopernic Space Corporation and EXOS\nAerospace Systems & Technologies, Inc.",
            {"font": "serif", "size": 20, "bold": True}, page_break=True, after=12)

    used = set()
    table_rows, table_fmt = [], None
    PAGE_BREAK_PAGES = {3, 24, 25, 28, 30, 31, 32, 33, 34, 38}
    seen_pages = set()

    def flush_table():
        nonlocal table_rows
        if table_rows:
            B.table(table_rows, table_fmt, cell_ops=E1_CELL_OPS)
            table_rows = []

    for p in P:
        fmt = {"font": p["font"], "size": p["size"], "bold": p["bold"]}
        if p["table_row"]:
            table_rows.append(p["table_row"]); table_fmt = fmt
            continue
        flush_table()
        pb = p["page"] in PAGE_BREAK_PAGES and p["page"] not in seen_pages
        seen_pages.add(p["page"])
        key = next((k for k in EDITS if p["text"].startswith(k) or (len(k) > 25 and k in p["text"][:len(k) + 60])), None)
        if key and key not in used:
            used.add(key)
            ops = EDITS[key]
            if ops == "DELETE":
                para = B.para([("del", p["text"], "Removed: superseded by the per-vehicle price above, which applies to each ordered vehicle.")], fmt, page_break=pb)
                mark_para_deleted(para)
                continue
            B.para(segments(p["text"], ops), fmt, page_break=pb)
            if key.startswith("“Pathfinder Flight”"):
                for txt, cmt in G1A_NEW:
                    B.new_para(txt, fmt, comment=cmt)
            continue
        B.para([("eq", p["text"], None)], fmt, page_break=pb)
    flush_table()
    missing = set(EDITS) - used
    assert not missing, f"edits not applied: {missing}"

    # Appendix H (wholly inserted)
    first = True
    for item in APPX_H:
        kind = item[0]
        if kind == "h":
            B.new_para(item[1], sans(15, True), page_break=first, after=8); first = False
        elif kind == "s":
            B.new_para(item[1], sans(12, True), after=4)
        elif kind == "t":
            B.table(item[1], sans(9), inserted=True)
        else:
            B.new_para(item[1], sans(10), comment=item[2])

    B.doc.core_properties.author = "Exos Aerospace"
    B.doc.core_properties.title = "Copernic Wet Lease — Exos Counter-Redline (draft 2026-10-17)"
    # turn on revision view
    settings = B.doc.settings.element
    tr = OxmlElement("w:trackRevisions")
    settings.insert(0, tr)
    B.doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
