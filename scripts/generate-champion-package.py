#!/usr/bin/env python3
"""Generate STRATFI Government Champion Request Letter + Template Response."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from datetime import datetime

ASSETS = '/Users/exomac/.openclaw/workspace/assets'
OUT = '/Users/exomac/.openclaw/workspace/drafts/STRATFI-Champion-Request-Package.docx'

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.01)
    section.footer_distance = Inches(0.01)

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)

NAVY = RGBColor(0x1B, 0x3A, 0x5C)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)

def para(text, bold=False, size=None, align=None, after=None, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.color.rgb = color or DARK
    if size: run.font.size = Pt(size)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

# ================================================================
# LOGO
# ================================================================
logo = os.path.join(ASSETS, 'exos-letterhead-logo.png')
if os.path.exists(logo):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(logo, width=Inches(2.0))
    para('', after=2)

# ================================================================
# PART 1: REQUEST LETTER TO CHAMPION
# ================================================================
para('EXOS AEROSPACE SYSTEMS & TECHNOLOGIES, INC.', bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
para('299 FM 1903, Greenville, TX 75402', size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
para('www.exosaero.com', size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

para('May 12, 2026', after=12)

para('[Recipient Name]\n[Title]\n[Organization]\n[Address]', after=12)

para('Subject: Request for Government Champion Support \u2014 PY26.1 STRATFI Notice of Opportunity', bold=True, after=8)

para('Dear [Recipient Name],', after=8)

para('I am writing to request your support as a Government Champion for Exos Aerospace\u2019s submission to the SpaceWERX PY26.1 Strategic Funding Increase (STRATFI) Notice of Opportunity. Your endorsement is a prerequisite for our STRATFI application, and your organization\u2019s alignment with this technology makes you the ideal champion for this effort.')

para('Background', bold=True, after=4)
para('Exos Aerospace holds an active USAF SBIR Phase II contract (FA9300-25-C-6022, $1.9M) with the Air Force Test Center for development of the Hybrid Vehicle Rocket Reusable (HVRR) platform with a Government-Furnished Rotating Detonation Rocket Engine (RDRE). The Phase II program supports hover test demonstration of RDRE-powered flight.')

para('Our STRATFI proposal extends this work to:')
p = doc.add_paragraph(style='List Bullet')
p.add_run('Flight-qualify the RDRE on our BLK3 reusable suborbital launch vehicle (5+ instrumented flights)')
p = doc.add_paragraph(style='List Bullet')
p.add_run('Establish BLK3 as an on-demand DoD flight test service with 30-day turnaround')
p = doc.add_paragraph(style='List Bullet')
p.add_run('Scale RDRE propulsion to the BLK4 hypersonic vehicle (Mach 15\u201320, SRBM\u2013IRBM profiles)')
p = doc.add_paragraph(style='List Bullet')
p.add_run('Deliver a HVRR operational prototype for orbital transfer / upper stage applications')

para('Funding Request', bold=True, after=4)
para('We are requesting a STRATFI investment of $15,000,000 over 36 months, with private investment match exceeding $15M through our Series A equity raise ($10M, closing Q2 2026) and Euronext Dublin public listing ($12.5M, Q3 2026).')

para('What We Are Asking of You', bold=True, after=4)
para('The STRATFI program requires a Government Champion who will:')
p = doc.add_paragraph(style='List Bullet')
r = p.add_run('1. ')
r.bold = True
p.add_run('Validate the operational need for RDRE flight demonstration and reusable DoD flight test services')
p = doc.add_paragraph(style='List Bullet')
r = p.add_run('2. ')
r.bold = True
p.add_run('Provide access to the PY26.1 STRATFI DAF Notice of Opportunity Teams channel and official proposal template')
p = doc.add_paragraph(style='List Bullet')
r = p.add_run('3. ')
r.bold = True
p.add_run('Endorse the technology transition plan by signing the attached Champion Endorsement Form')

para('To make this as easy as possible, I have attached a pre-populated Champion Endorsement Form that requires only your signature and any modifications you see fit. The form is structured to meet the STRATFI program requirements with minimal effort on your part.')

para('We believe this program directly supports your organization\u2019s mission in advanced propulsion development and flight test capability. Exos is one of only three FAA-licensed reusable suborbital launch companies in the U.S., and our RDRE program \u2014 backed by an active SBIR Phase II and NASA MSFC partnership \u2014 represents a near-term pathway to flight-demonstrating detonation-based propulsion.')

para('I am available to discuss this at your convenience and can provide any additional information you require.')

para('Respectfully,', after=12)
para('John Quinn\nCofounder and Chief Executive Officer\nExos Aerospace Systems & Technologies, Inc.\njquinn@exosaero.com | +1-972-740-8355', after=4)

para('Enclosures:', bold=True, after=2)
p = doc.add_paragraph(style='List Bullet')
p.add_run('Champion Endorsement Form (pre-populated, signature required)')
p = doc.add_paragraph(style='List Bullet')
p.add_run('STRATFI Proposal Executive Summary')
p = doc.add_paragraph(style='List Bullet')
p.add_run('Exos RDRE Capability Statement')

doc.add_page_break()

# ================================================================
# PART 2: CHAMPION ENDORSEMENT FORM (TEMPLATE)
# ================================================================
para('STRATFI PY26.1 \u2014 GOVERNMENT CHAMPION ENDORSEMENT FORM', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, color=NAVY)
para('(Pre-populated for ease of completion \u2014 modify as appropriate)', size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=16, color=GRAY)

# Info block
fields = [
    ('Small Business Name:', 'Exos Aerospace Systems & Technologies, Inc.'),
    ('SBIR Phase II Contract:', 'FA9300-25-C-6022'),
    ('Awarding Agency:', 'USAF \u2014 Air Force Test Center (AFTC)'),
    ('Phase II Award Amount:', '$1,897,824'),
    ('Phase II Topic:', 'Hybrid Vehicle Rocket with Government-Furnished RDRE'),
    ('STRATFI Proposal Title:', 'BLK3/BLK4 RDRE-Powered Reusable Vehicle Program'),
    ('STRATFI Amount Requested:', '$15,000,000'),
    ('Period of Performance:', '36 months'),
]
for label, val in fields:
    p = doc.add_paragraph()
    r = p.add_run(label + ' ')
    r.bold = True; r.font.color.rgb = DARK; r.font.size = Pt(10)
    r2 = p.add_run(val)
    r2.font.color.rgb = DARK; r2.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(3)

para('', after=8)

# Endorsement section
para('GOVERNMENT CHAMPION ENDORSEMENT', bold=True, size=12, color=NAVY, after=8)

para('I, the undersigned, serve as the Government Champion for the above-referenced STRATFI proposal. In this capacity, I confirm the following:', after=6)

statements = [
    'Operational Need: The technology proposed by Exos Aerospace \u2014 specifically, the flight demonstration of Rotating Detonation Rocket Engine (RDRE) propulsion and the establishment of a reusable DoD flight test service \u2014 addresses a validated operational need within my organization\u2019s mission area.',
    'Technology Transition: I commit to supporting the transition of the technology developed under this STRATFI investment into operational use within the Department of the Air Force, United States Space Force, or other DoD component, through available contracting vehicles including Phase III SBIR, OTA, or other appropriate mechanisms.',
    'Program Oversight: I will provide appropriate government oversight and technical engagement throughout the STRATFI period of performance, including participation in program reviews and milestone assessments as required.',
    'Access and Coordination: I will facilitate Exos Aerospace\u2019s access to the PY26.1 STRATFI DAF Notice of Opportunity Teams channel and assist in obtaining the official proposal template for final submission.',
]

for i, stmt in enumerate(statements, 1):
    p = doc.add_paragraph()
    r = p.add_run(f'{i}. ')
    r.bold = True; r.font.color.rgb = DARK
    # Split into bold label and rest
    parts = stmt.split(': ', 1)
    r2 = p.add_run(parts[0] + ': ')
    r2.bold = True; r2.font.color.rgb = DARK
    r3 = p.add_run(parts[1])
    r3.font.color.rgb = DARK
    p.paragraph_format.space_after = Pt(6)

para('', after=16)

# Signature block
para('CHAMPION INFORMATION', bold=True, size=11, color=NAVY, after=8)

sig_fields = [
    ('Name:', '________________________________________'),
    ('Title:', '________________________________________'),
    ('Organization:', '________________________________________'),
    ('Email:', '________________________________________'),
    ('Phone:', '________________________________________'),
    ('', ''),
    ('Signature:', '________________________________________'),
    ('Date:', '________________________________________'),
]
for label, val in sig_fields:
    if not label:
        para('', after=8)
        continue
    p = doc.add_paragraph()
    r = p.add_run(label + '    ')
    r.bold = True; r.font.color.rgb = DARK; r.font.size = Pt(11)
    r2 = p.add_run(val)
    r2.font.color.rgb = GRAY; r2.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(8)

para('', after=16)

# Note
para('Note: This endorsement form is provided for the convenience of the Government Champion. If your organization has a preferred format for STRATFI endorsements or if additional information is required, Exos Aerospace will accommodate any modifications. Please contact John Quinn (jquinn@exosaero.com, +1-972-740-8355) with any questions.', size=9, color=GRAY, after=8)

# Footer
para('', after=8)
footer_img = os.path.join(ASSETS, 'exos-letterhead-footer.png')
if os.path.exists(footer_img):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(footer_img, width=Inches(5.5))

doc.save(OUT)
print(f'Saved: {OUT}')
