#!/usr/bin/env python3
"""Combined Lady Rocket Founding Feature Portfolio as a single editable .docx
(cover page + Tier I/II/III), letterhead header+footer, tight gap fix applied."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION

OUT = 'Lady_Rocket_Founding_Feature_Portfolio.docx'
LOGO = 'assets/exos-letterhead-logo.png'      # canonical = now tight-cropped
FOOTER = 'assets/exos-letterhead-footer.png'

GOLD = RGBColor(0xB8, 0x86, 0x0B)
NAVY = RGBColor(0x1B, 0x3A, 0x5C)
GRAY = RGBColor(0x40, 0x40, 0x40)
DARK = RGBColor(0x22, 0x22, 0x22)

TIERS = [
    ('drafts/lady-rocket-reusable-signature-spec.md', 'Founding Applicant Program — Tier One Feature'),
    ('drafts/lady-rocket-flown-artifact-vault.md', 'Founding Applicant Program — Tier Two Feature'),
    ('drafts/lady-rocket-lunar-legacy-capsule.md', 'Founding Applicant Program — Tier Three Feature'),
]

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.left_margin = Inches(1.0); sec.right_margin = Inches(0.5)
sec.top_margin = Inches(1.5); sec.bottom_margin = Inches(1.1)
sec.header_distance = Inches(0.01); sec.footer_distance = Inches(0.01)

nf = doc.styles['Normal'].font
nf.name = 'Calibri'; nf.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(4)

# letterhead header + footer (tight)
hp = sec.header.paragraphs[0]
hp.paragraph_format.space_before = Pt(0); hp.paragraph_format.space_after = Pt(0)
hp.add_run().add_picture(LOGO, height=Inches(0.75))
fpar = sec.footer.paragraphs[0]; fpar.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fpar.paragraph_format.space_before = Pt(0); fpar.paragraph_format.space_after = Pt(0)
fpar.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
fpar.paragraph_format.line_spacing = Pt(54)  # 0.75in, kills baseline descent below footer image
_fr = fpar.add_run(); _fr.font.size = Pt(1)
_fr.add_picture(FOOTER, height=Inches(0.75))


def add(text='', bold=False, size=None, align=None, italic=False, color=DARK,
        before=None, after=None):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    if align is not None: p.alignment = align
    if before is not None: p.paragraph_format.space_before = Pt(before)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p


def is_section(s):
    return bool(re.match(r'^[A-Z][A-Z0-9 ™&·\-]{3,}$', s)) and len(s) <= 40


# ---- COVER ----
add('LADY ROCKET', bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, color=GOLD, before=40, after=4)
add('Founding Applicant Feature Portfolio', bold=True, size=24,
    align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, after=4)
add('The world’s most exclusive private rocket community', italic=True, size=12,
    align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, after=18)
add('Three privileges reserved for the twenty.', italic=True, size=11,
    align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, after=14)


def ladder(num, name, desc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f'{num} — {name}  '); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAVY
    r2 = p.add_run('— ' + desc); r2.font.size = Pt(11); r2.font.color.rgb = DARK
    p.paragraph_format.space_after = Pt(10)

ladder('I', 'The Reusable Signature™', 'Your mark, flown and returned, accumulating flight heritage. Ownership.')
ladder('II', 'The Flown Artifact Vault™', 'Personal objects flown and returned with certified provenance. Renewable luxury.')
ladder('III', 'The Lunar Legacy Capsule™', 'A reserved opportunity toward the Moon, built on real lunar heritage. Legacy.')
add('Powered by Exos Aerospace · In collaboration with Copernic Space', size=9,
    align=WD_ALIGN_PARAGRAPH.CENTER, color=GOLD, before=18)


def new_page():
    doc.add_page_break()


def render_tier(path, subtitle):
    raw = open(path).read().split('\n')
    n = len(raw)
    club_txt = raw[0].strip()
    i = 0
    while i < n and raw[i].strip() != '---':
        i += 1
    mid = [x.strip() for x in raw[2:i] if x.strip()]
    hero_txt = mid[0] if mid else ''
    tag_txt = mid[1] if len(mid) > 1 else ''
    pw_txt = mid[2] if len(mid) > 2 else ''
    add(club_txt, bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, color=GOLD, after=2)
    add(subtitle, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, after=4)
    add(hero_txt, bold=True, size=19, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, after=2)
    if tag_txt: add(tag_txt, italic=True, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, after=2)
    if pw_txt: add(pw_txt, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=GOLD, after=8)
    i += 1
    buf = []
    def flush():
        if buf:
            add(' '.join(buf), after=6); buf.clear()
    while i < n:
        s = raw[i].strip()
        if s == '---' or not s:
            flush(); i += 1; continue
        if is_section(s):
            flush(); add(s, bold=True, size=12, color=NAVY, before=9, after=4); i += 1; continue
        m = re.match(r'^([A-Z][A-Za-z ™]{2,40})\.\s+(.*)$', s)
        if m and len(m.group(1).split()) <= 5 and not m.group(1).isupper():
            flush()
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.2)
            r = p.add_run(m.group(1) + '. '); r.bold = True; r.font.color.rgb = GOLD
            r2 = p.add_run(m.group(2)); r2.font.color.rgb = DARK
            p.paragraph_format.space_after = Pt(5)
            i += 1; continue
        if s.startswith('Note:'):
            flush(); add(s, italic=True, size=8.5, color=GRAY, before=8); i += 1; continue
        buf.append(s)
        i += 1
    flush()


for path, subtitle in TIERS:
    new_page()
    render_tier(path, subtitle)

doc.save(OUT)
print('WROTE', OUT)
