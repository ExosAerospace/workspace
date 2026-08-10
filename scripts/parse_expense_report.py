#!/usr/bin/env python3
"""Extract amount / period / purpose from a completed Exos expense report (.xlsx).

Maps the house template (Exos_Expense_Report_Template_2026.xlsx):
  - amount  : GRAND TOTALS grand-total cell J49 (cached value); falls back to
              summing the direct-entry expense cells if no cached total exists.
  - period  : WEEK ENDING date in C8 (named range WeekEnding).
  - purpose : Event Location (C7), else the Business Purpose note (B53).
Prints a single JSON object. Never raises to the caller — on any failure it
returns nulls so the submit flow still works (manager reads the attached file).
Usage: parse_expense_report.py <path-to-xlsx>
"""
import sys, json, datetime


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main():
    out = {"amount": None, "period": None, "purpose": None}
    if len(sys.argv) < 2:
        print(json.dumps(out)); return
    path = sys.argv[1]
    if not path.lower().endswith((".xlsx", ".xlsm", ".xls")):
        print(json.dumps(out)); return  # only spreadsheet templates are parseable
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active

        # amount: cached grand total, else sum of the direct-entry expense cells
        gt = num(ws["J49"].value)
        if gt is None:
            s, found = 0.0, False
            for rng in ("C14:I20", "C24:I28", "C33:I37", "C41:I46"):
                for row in ws[rng]:
                    for c in row:
                        n = num(c.value)
                        if n is not None:
                            s += n; found = True
            gt = round(s, 2) if found else None
        out["amount"] = gt

        # period: week-ending date
        wev = ws["C8"].value
        if isinstance(wev, (datetime.datetime, datetime.date)):
            out["period"] = "Week ending " + wev.strftime("%b %d, %Y")
        elif wev not in (None, ""):
            out["period"] = str(wev).strip()

        # purpose: event location, else business-purpose note
        for coord in ("C7", "D7", "E7", "B53"):
            v = ws[coord].value
            if v not in (None, "") and str(v).strip():
                out["purpose"] = str(v).strip(); break
    except Exception as e:  # noqa: BLE001 - never break the submit flow
        out["_err"] = str(e)[:140]
    print(json.dumps(out))


if __name__ == "__main__":
    main()
