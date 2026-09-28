#!/usr/bin/env python3
"""Build the full EXOS Employee Handbook (Rev. 2026-09-28) as .docx.

Re-issue of the 3.12.26 handbook with the whistleblower revision merged in:
new Sec. 2.5 (DFARS 252.203-7002 / 10 USC 4701 notice) and conforming
carve-outs in 2.2, 2.4, 3.6, 3.8 (incl. DTSA 1833(b) notice), 3.9 and 12.
"""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = sys.argv[1] if len(sys.argv) > 1 else \
    "process-library/05-human-resources/EXOS_Employee_Handbook_Rev_2026-09-28.docx"
CO = "EXOS Aerospace Systems & Technologies, Inc."
REPORT_URL = "https://mc.exosaero.space/report.html"

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def h(t, lvl=1):
    doc.add_heading(t, level=lvl)


def p(t, bold=False, italic=False):
    para = doc.add_paragraph()
    r = para.add_run(t)
    r.bold, r.italic = bold, italic
    return para


def ul(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


t = doc.add_heading(f"{CO}\nEmployee Handbook", 0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Revision 2026-09-28 — supersedes the Employee Handbook dated 3.12.26", italic=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Revision summary: adds Section 2.5, Whistleblower Protection & Reporting of Misconduct (written notice of "
  "employee rights under 10 U.S.C. § 4701 and 41 U.S.C. § 4712, per DFARS 252.203-7002 and FAR 52.203-17), "
  "with conforming changes to Sections 2.2, 2.4, 3.6, 3.8, 3.9 and 12.", italic=True)

# 1 ------------------------------------------------------------------
h("1. Welcome")
h("1.1 History, Goals & Culture", 2)
p(f"Welcome to {CO}")
p("EXOS is committed to advancing aerospace innovation through the development of reusable launch systems, "
  "advanced propulsion technologies, and related aerospace services. Our mission is to deliver high-quality "
  "solutions to our partners while creating an environment where employees can grow professionally and "
  "contribute meaningfully to the advancement of aerospace technology.")
p("Our employees are the foundation of our success. Each team member plays an important role in maintaining "
  "the company's reputation for technical excellence, integrity, and innovation.")
p("We expect all employees to conduct themselves professionally, collaborate respectfully, and maintain the "
  "highest standards of safety, security, and ethical conduct.")
h("1.2 Purpose of This Handbook", 2)
p(f"This handbook outlines general policies, procedures, and expectations for employees of {CO}")
p("This handbook is intended to:")
ul(["provide employees with an overview of workplace expectations", "describe company policies and procedures",
    "promote a safe, respectful, and productive workplace"])
p("This handbook does not create a contract of employment and is not intended to guarantee employment for any "
  "specific period of time.")
p("The company reserves the right to modify, suspend, or revise any policy at its discretion, except that no "
  "revision will limit the rights described in Section 2.5.")
h("1.3 At-Will Employment", 2)
p(f"Employment with {CO} is at-will. This means that either the employee or the company may terminate "
  "employment at any time, with or without cause or notice, subject to applicable law.")
p("Nothing in this handbook alters the at-will employment relationship unless a written agreement is signed "
  "by authorized company leadership.")

# 2 ------------------------------------------------------------------
h("2. Workplace Commitments")
h("2.1 Equal Opportunity Employment", 2)
p(f"{CO} is an equal opportunity employer.")
p("The company does not discriminate against employees or applicants on the basis of:")
ul(["race", "color", "religion", "sex", "national origin", "age", "disability", "marital status",
    "veteran status", "or any other status protected by law"])
p("Employment decisions are based on qualifications, merit, and business needs.")
h("2.2 Non-Harassment Policy", 2)
p("EXOS is committed to providing a professional workplace free from harassment, discrimination, or intimidation.")
p("Harassment includes any unwelcome conduct that creates an intimidating, hostile, or offensive work environment.")
p("Examples include:")
ul(["offensive jokes or comments", "slurs or derogatory remarks", "inappropriate physical conduct",
    "repeated unwanted advances", "displaying offensive material"])
p("Employees who experience or witness harassment should report the issue to management immediately.")
p("Retaliation against any employee who reports a concern, participates in an investigation, or makes a "
  "protected disclosure is strictly prohibited. See Section 2.5.")
h("2.3 Drug-Free Workplace", 2)
p("Employees may not use, possess, distribute, or be under the influence of illegal drugs or alcohol while:")
ul(["on company premises", "operating company equipment", "conducting company business"])
p("Prescription medications taken as directed are permitted.")
p("The company reserves the right to conduct drug testing where permitted by law.")
h("2.4 Open Door Policy", 2)
p("EXOS encourages open communication.")
p("Employees are encouraged to bring questions, concerns, or workplace issues to management. The company values "
  "constructive feedback and strives to resolve concerns promptly and professionally.")
p("The Open Door Policy is one option for raising concerns. It is never a precondition to, and does not limit, "
  "an employee's right to report concerns to any government agency or other recipient listed in Section 2.5.")

h("2.5 Whistleblower Protection & Reporting of Misconduct", 2)
h("2.5.1 Policy", 3)
p("EXOS is committed to integrity in all of its business, including its work on U.S. Government contracts, "
  "subcontracts, and grants. EXOS encourages employees to report suspected misconduct and strictly prohibits "
  "discharge, demotion, or any other form of discrimination or retaliation against an employee for making a "
  "protected disclosure. This section is the written notice of whistleblower rights that EXOS is required to "
  "provide under DFARS 252.203-7002 and FAR 52.203-17.")
h("2.5.2 What you may report", 3)
p("Employees are protected when they disclose information they reasonably believe is evidence of:")
ul(["gross mismanagement of a federal contract or grant;",
    "a gross waste of federal funds;",
    "an abuse of authority relating to a federal contract or grant;",
    "a violation of law, rule, or regulation related to a federal contract or grant, including the competition "
    "for or negotiation of a contract (for example: fraud, false statements, mischarging of labor or costs, "
    "falsified test or inspection records, kickbacks, bribery, gratuities, or conflicts of interest);",
    "a substantial and specific danger to public health or safety;",
    "export control violations (ITAR/EAR), mishandling of Controlled Unclassified Information, or cybersecurity "
    "incidents; or",
    "any other violation of law or of EXOS policy."])
p("A disclosure is protected even if it is made in the ordinary course of your job duties, and even if the "
  "information turns out to be mistaken, so long as you reasonably believed it when you reported it.")
h("2.5.3 Who you may report to", 3)
p("You may report internally, externally, or both. You are never required to report internally first, and you "
  "do not need EXOS's permission or notice before reporting externally.")
p("Internal channels:", bold=True)
ul(["your supervisor or any member of management;",
    "Human Resources;",
    "the Chief Executive Officer; or",
    f"the Mission Control Ethics & Compliance Report tool ({REPORT_URL}). You may submit anonymously; "
    "anonymous reports are accepted and investigated to the extent possible."])
p("External recipients protected by law (10 U.S.C. § 4701 / 41 U.S.C. § 4712):", bold=True)
ul(["a Member of Congress or a representative of a committee of Congress;",
    "an Inspector General (DoD Office of Inspector General; NASA OIG; or the IG of any other agency);",
    "the Government Accountability Office;",
    "a federal employee responsible for contract oversight or management at the relevant agency "
    "(e.g., a Contracting Officer or Contracting Officer's Representative);",
    "an authorized official of the Department of Justice or other law enforcement agency;",
    "a court or grand jury; or",
    "a management official or other employee of EXOS, or of a prime or subcontractor, who has responsibility "
    "to investigate, discover, or address misconduct."])
p("Employees may also report possible violations of law to, and communicate with, the Securities and Exchange "
  "Commission, the Equal Employment Opportunity Commission, the Occupational Safety and Health Administration, "
  "the National Labor Relations Board, or any other federal, state, or local agency, without notice to EXOS.")
p("DoD Hotline: 1-800-424-9098  |  www.dodig.mil/hotline")
p("NASA OIG Hotline: 1-800-424-9183  |  oig.nasa.gov/hotline")
h("2.5.4 No retaliation", 3)
p("EXOS will not discharge, demote, suspend, threaten, harass, reduce the pay or benefits of, or otherwise "
  "discriminate against any employee because the employee made a protected disclosure, participated in an "
  "investigation, or refused to participate in conduct the employee reasonably believes is unlawful. Any "
  "employee, supervisor, or manager who retaliates will be subject to discipline up to and including "
  "termination. Reprisal may be reported through any channel in 2.5.3.")
p("Under 10 U.S.C. § 4701 and 41 U.S.C. § 4712, an employee who believes they have been subjected to reprisal "
  "may file a complaint with the Inspector General of the agency concerned within three (3) years after the "
  "reprisal occurred. Remedies may include reinstatement, back pay, compensatory damages, and attorney fees. "
  "These rights cannot be waived by any agreement, policy, form, or condition of employment, including "
  "pre-dispute arbitration agreements.")
h("2.5.5 Management's duty to escalate; mandatory disclosure", 3)
p("Any manager who receives a report under this section must promptly forward it to the CEO (or, if the report "
  "concerns the CEO, to the Board of Directors). EXOS will investigate promptly and keep the reporter's "
  "identity confidential to the extent permitted by law and consistent with a thorough investigation. Where "
  "EXOS has credible evidence that a principal, employee, agent, or subcontractor has committed a violation of "
  "federal criminal law involving fraud, conflict of interest, bribery, or gratuities, or a violation of the "
  "civil False Claims Act, in connection with a federal contract or subcontract, EXOS will timely disclose it "
  "in writing to the agency Office of Inspector General and the Contracting Officer as required by "
  "FAR 52.203-13.")
h("2.5.6 Classified, export-controlled, and CUI information", 3)
p("Nothing in this section authorizes disclosing classified information, export-controlled technical data, or "
  "Controlled Unclassified Information to a recipient not authorized to receive it. When a report involves such "
  "information, make it through an Inspector General hotline or to an authorized government official, who can "
  "receive it through proper channels. If you are unsure, you may describe the concern in general terms without "
  "including the controlled content. Do not enter CUI or export-controlled technical data into the Mission "
  "Control report tool.")
h("2.5.7 Postings", 3)
p("EXOS displays the DoD Hotline poster and applicable agency hotline posters in common work areas and on the "
  "company intranet.")

# 3 ------------------------------------------------------------------
h("3. Company Policies & Procedures")
h("3.1 Professional Conduct", 2)
p("Employees are expected to maintain a professional standard of behavior at all times.")
p("Employees should:")
ul(["treat coworkers respectfully", "act with integrity", "cooperate with supervisors and team members",
    "perform job responsibilities diligently"])
p("Unprofessional behavior may result in disciplinary action.")
h("3.2 Dress Code", 2)
p("Employees should maintain a professional and appropriate appearance consistent with their work "
  "responsibilities and safety requirements.")
p("Certain operational environments may require specific protective equipment.")
h("3.3 Payday", 2)
p("Employees are paid bi-weekly via electronic deposit.")
p("Paychecks include wages earned minus applicable taxes and authorized deductions.")
p("Employees should report payroll discrepancies promptly.")
h("3.4 Time Tracking", 2)
p("Employees must accurately record all hours worked in the designated time tracking system.")
p("Time records should:")
ul(["be completed daily", "accurately reflect work performed", "include project descriptions when required"])
p("Falsification of time records will result in disciplinary action.")
h("3.5 Company Property", 2)
p("Company property includes equipment, vehicles, computers, software, communication systems, and other "
  "company-owned assets.")
p("Company property is provided for business use only and must be used responsibly.")
p("Employees are responsible for safeguarding company equipment and must return all company property upon "
  "termination of employment.")
h("3.6 Technology, Email & Communications Policy", 2)
p("EXOS provides employees with access to computers, networks, internet access, email, and other digital "
  "communication tools for business purposes.")
p("These systems remain the property of the company. Employees must use these systems responsibly and securely.")
p("Employees may not:")
ul(["download unauthorized software", "install unapproved applications", "access inappropriate or illegal material",
    "use company systems for illegal activity", "transmit confidential company data to unauthorized recipients"])
p("The restriction on transmitting confidential company data does not apply to disclosures protected by "
  "Section 2.5 or Section 3.8 (Protected Activity).")
p("All communications conducted using company systems may be monitored or reviewed as necessary.")
p("Employees should have no expectation of privacy when using company systems.")
h("3.7 Artificial Intelligence & External Technology Use", 2)
p("Artificial intelligence tools can assist with productivity but may introduce security risks.")
p("Employees may use AI tools for general tasks such as:")
ul(["brainstorming", "drafting content", "research on public information"])
p("However, employees may not upload or share the following with public AI tools:")
ul(["proprietary company information", "engineering or design data", "confidential business information",
    "investor or financial information", "employee personal data",
    "export-controlled or restricted technical information"])
p("Examples include tools such as ChatGPT, Claude, Gemini, Copilot, Midjourney, and similar systems.")
p("Employees are responsible for reviewing AI-generated output for accuracy and compliance with company policies.")
p("Employees must assume that information entered into public AI tools may become part of external datasets and "
  "must therefore avoid entering any company-sensitive information.")
h("3.8 Data Protection & Confidential Information", 2)
p("Employees may have access to sensitive information including:")
ul(["engineering designs", "testing data", "proprietary technology", "financial information",
    "investor communications", "internal company strategies"])
p("This information must not be shared outside the company without authorization.")
p("Unauthorized disclosure of confidential information may result in disciplinary action up to termination.")
p("Protected Activity.", bold=True)
p("Nothing in this handbook, in any EXOS confidentiality or invention assignment agreement, or in any other EXOS "
  "policy prohibits or restricts an employee from: (a) reporting waste, fraud, abuse, or any possible violation "
  "of law or regulation to a designated investigative or law enforcement representative of a federal department "
  "or agency authorized to receive such information, to an Inspector General, to Congress, or to any other "
  "government agency; (b) participating in or cooperating with any government investigation or proceeding; "
  "(c) receiving an award from a government agency for information provided; or (d) discussing wages, hours, or "
  "other terms and conditions of employment as protected by Section 7 of the National Labor Relations Act. "
  "Employees do not need prior authorization from, or to notify, EXOS before doing so. In accordance with "
  "FAR 52.203-19, EXOS hereby notifies all employees that any prior confidentiality agreement or statement is "
  "limited by, and does not supersede, these rights.")
p("Defend Trade Secrets Act Notice (18 U.S.C. § 1833(b)).", bold=True)
p("An individual shall not be held criminally or civilly liable under any federal or state trade secret law for "
  "the disclosure of a trade secret that (A) is made (i) in confidence to a federal, state, or local government "
  "official, either directly or indirectly, or to an attorney, and (ii) solely for the purpose of reporting or "
  "investigating a suspected violation of law; or (B) is made in a complaint or other document filed in a "
  "lawsuit or other proceeding, if such filing is made under seal. An individual who files a lawsuit for "
  "retaliation by an employer for reporting a suspected violation of law may disclose the trade secret to the "
  "individual's attorney and use the trade secret information in the court proceeding, if the individual files "
  "any document containing the trade secret under seal and does not disclose the trade secret, except pursuant "
  "to court order.")
h("3.9 Privacy & Monitoring", 2)
p("The company reserves the right to access company property including:")
ul(["computers", "email accounts", "files", "desks", "storage systems"])
p("Employees should not expect privacy when using company systems or facilities.")
p("EXOS will not use its monitoring rights to identify, discourage, or retaliate against an employee for a "
  "protected disclosure under Section 2.5.")
h("3.10 Export Control & Technical Data Protection", 2)
p(f"{CO} operates in an industry that may involve technologies subject to United States export control laws, "
  "including but not limited to the International Traffic in Arms Regulations (ITAR) and the Export "
  "Administration Regulations (EAR).")
p("Employees must protect controlled technical information and comply with all applicable export control laws "
  "and regulations.")
p("Export-controlled or restricted information may include, but is not limited to:")
ul(["rocket propulsion systems", "aerospace vehicle design data", "technical drawings or schematics",
    "test results and engineering analysis", "flight systems and avionics data", "manufacturing processes",
    "defense-related technologies"])
p("Employees must not share controlled information with:")
ul(["foreign nationals without authorization", "external parties without approval",
    "unauthorized individuals within the company"])
p("Controlled technical data must never be uploaded to public websites, cloud platforms, file-sharing systems, "
  "or artificial intelligence tools without prior authorization.")
p("Employees who are unsure whether information may be export-controlled must consult company management before "
  "sharing or transmitting such information.")
p("Violations of export control laws may result in severe civil or criminal penalties and may also lead to "
  "disciplinary action up to and including termination.")

# 4-11 ---------------------------------------------------------------
h("4. Employment Classification")
p("Employees may be classified as:")
ul(["Exempt", "Non-Exempt", "Full-Time", "Part-Time", "Temporary"])
p("Compensation and overtime eligibility are determined in accordance with applicable laws.")
h("5. Attendance")
p("Employees are expected to maintain regular attendance and punctuality.")
p("Employees who will be absent or late should notify their supervisor as soon as possible.")
p("Excessive absenteeism or tardiness may result in disciplinary action.")
h("6. Leave Policies")
p("Employees may be eligible for leave benefits including:")
ul(["vacation", "sick leave", "family and medical leave", "military leave", "jury duty leave"])
p("Leave policies are administered in accordance with applicable laws and company policy.")
h("7. Performance Expectations")
p("Employees are expected to perform their duties responsibly and professionally.")
p("Performance evaluations may be conducted periodically and may consider:")
ul(["quality of work", "teamwork", "reliability", "job knowledge", "compliance with company policies"])
h("8. Discipline")
p("Violation of company policies may result in disciplinary action including:")
ul(["verbal warning", "written warning", "suspension", "termination"])
p("The company reserves discretion in determining appropriate disciplinary action. No employee will be "
  "disciplined for activity protected by Section 2.5.")
h("9. Workplace Safety & Security")
p("Safety is a top priority at EXOS.")
p("Employees must follow all safety procedures and immediately report unsafe conditions.")
p("EXOS operations may involve activities such as:")
ul(["rocket testing", "energetic materials", "pressurized systems", "aerospace hardware"])
p("Employees must follow all safety instructions when participating in these activities.")
h("9.1 Security Awareness", 2)
p("Employees must remain alert to potential safety or security risks.")
p("Visitors must be escorted when required, and suspicious activity should be reported immediately.")
p("Employees must also report:")
ul(["phishing emails", "suspected cybersecurity incidents", "lost company devices", "unauthorized access to systems"])
h("9.2 Cybersecurity & Information Security", 2)
p(f"{CO} relies on secure information systems to support engineering, operations, and business activities.")
p("All employees share responsibility for protecting company information systems and digital assets.")
p("Employees must:")
ul(["maintain strong and confidential passwords", "lock computers when unattended",
    "use company systems only for authorized purposes", "report suspicious emails or phishing attempts",
    "report lost or stolen company devices immediately", "avoid downloading unauthorized software or files"])
p("Employees must not attempt to bypass security controls or access systems they are not authorized to use.")
p("Any suspected cybersecurity incident, data breach, or unauthorized system access must be reported to "
  "management immediately.")
p("Failure to follow cybersecurity practices may result in disciplinary action.")
h("10. Benefits")
p("EXOS may provide benefits such as:")
ul(["workers' compensation coverage", "other benefits as determined by the company"])
p("Benefits may change at the company's discretion.")
h("11. Termination")
p("Employees may resign voluntarily or may be terminated in accordance with the company's at-will employment policy.")
p("Upon termination employees must return:")
ul(["equipment", "keys", "access cards", "documents", "company devices"])
p("The company may conduct an exit interview.")

# 12 -----------------------------------------------------------------
h("12. Employee Acknowledgment")
p(f"I acknowledge that I have received and reviewed the {CO} Employee Handbook (Revision 2026-09-28).")
p("I understand that this handbook outlines general policies, procedures, and expectations for employees of "
  "the company.")
p(f"I acknowledge that this handbook does not create a contract of employment and that my employment with {CO} "
  "is at-will.")
p("I acknowledge that I have received written notice of my whistleblower rights and protections under "
  "10 U.S.C. § 4701 and 41 U.S.C. § 4712 (Section 2.5), and that nothing in this handbook or in any agreement "
  "with EXOS restricts my right to report possible violations of law to the government.")
p("I agree to comply with the policies described in this handbook.")
p("\nEmployee Name: ______________________________")
p("Employee Signature: __________________________")
p("Date: ______________________")

doc.save(OUT)
print(OUT)
