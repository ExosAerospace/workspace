#!/usr/bin/env python3
"""Rev 2 of the Innovation Dojo investor pitch (John's house format, from Rev1).
Locks John's 4 answers (9/27): (a) US$5.0M seed, (b) exit line, (c) FY2025 revenue +
EBITDA from QBO (accrual, unaudited), (d) BLK3-for-Japan production on intl approvals.
Edits run text in place so Toshi's styling is preserved."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = 'drafts/Exos_Aerospace_Japan_Investor_Pitch_Rev1_2026-09-25.pptx'
OUT = 'drafts/Exos_Aerospace_Japan_Investor_Pitch_Rev2_2026-09-27.pptx'
prs = Presentation(SRC); S = prs.slides

def shape(slide, name):
    return next(sh for sh in slide.shapes if sh.name == name)

def settext(slide, name, text, color_from=None):
    """Replace a textbox's text, keeping the first run's font; optional color from another box."""
    tf = shape(slide, name).text_frame
    p0 = tf.paragraphs[0]
    for p in list(tf.paragraphs[1:]): p._p.getparent().remove(p._p)
    runs = p0.runs
    for r in runs[1:]: r._r.getparent().remove(r._r)
    runs[0].text = text
    if color_from:
        src = shape(slide, color_from).text_frame.paragraphs[0].runs[0].font.color.rgb
        runs[0].font.color.rgb = src

# Slide 2 — platform: add 2025 financials (replaces BLK3 tile; BLK3 stays in the tech stack)
s = S[1]
settext(s, 'TextBox 2', 'A profitable U.S. space company with real flight heritage')
settext(s, 'TextBox 12', '$1.49M')
settext(s, 'TextBox 13', '2025 revenue · EBITDA $254K (17%)')
fn = copy.deepcopy(shape(s, 'TextBox 13')._element); s.shapes._spTree.append(fn)
fn_sh = s.shapes[-1]; fn_sh.name = 'FY25 footnote'
fn_sh.left, fn_sh.top, fn_sh.width, fn_sh.height = Inches(6.51), Inches(3.32), Inches(2.75), Inches(0.2)
fn_sh.text_frame.paragraphs[0].runs[0].text = 'FY2025, accrual basis, unaudited'
fn_sh.text_frame.paragraphs[0].runs[0].font.size = Pt(8)

# Slide 5 — seed figure
s = S[4]
settext(s, 'TextBox 30', '¥790M')
settext(s, 'TextBox 31', '≈ US$5.0M target')
settext(s, 'TextBox 32', 'at ~¥157 / US$ (27 Sep 2026)', color_from='TextBox 39')

# Slide 6 — exit line
s = S[5]
settext(s, 'TextBox 15', 'Liquidity paths: acquisition of the K.K. stake by Exos U.S. or a strategic '
                         'buyer, or a K.K.-level listing. Horizon ~3+ years.', color_from='TextBox 11')

# Securities disclaimer (Jon 9/27): K.K. raise stays separate from the U.S. parent's Reg D 504 offering
DISCLAIMER = ('Not an offer of securities of Exos Aerospace Systems & Technologies, Inc. Any offering by '
              'Exos Aerospace Japan K.K. will be made only through definitive documents to eligible investors.')
for si in (4, 5):
    sl = S[si]
    el = copy.deepcopy(shape(sl, 'TextBox 3')._element); sl.shapes._spTree.append(el)
    d = sl.shapes[-1]; d.name = 'Securities disclaimer'
    d.left, d.top, d.width, d.height = Inches(0.65), Inches(6.78), Inches(11.5), Inches(0.25)
    tf = d.text_frame
    for pp in list(tf.paragraphs[1:]): pp._p.getparent().remove(pp._p)
    rr = tf.paragraphs[0].runs
    for r in rr[1:]: r._r.getparent().remove(r._r)
    rr[0].text = DISCLAIMER; rr[0].font.size = Pt(8.5); rr[0].font.italic = True

# Slide 7 — BLK3-for-Japan status
s = S[6]
settext(s, 'TextBox 13', 'Seeking Japanese technical partners for licensed production of TRL-9 Type-5 '
                         'tanks in Japan. BLK3-for-Japan production commences upon international '
                         'approvals (incl. U.S. export), with potential reallocation of U.S. BLK3 '
                         'assets to match the speed of those approvals.')
shape(s, 'TextBox 13').width = Inches(5.34)
shape(s, 'TextBox 13').text_frame.paragraphs[0].runs[0].font.size = Pt(11.5)  # 4 lines fit the card

prs.save(OUT); print('SAVED', OUT)
