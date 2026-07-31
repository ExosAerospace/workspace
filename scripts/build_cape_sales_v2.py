#!/usr/bin/env python3
"""Sales & Marketing Capability Statement (Gerald Grimes request, 2026-07-24).
Reorganized around four target verticals; covers flight services + composite tank tech."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x40, 0x40, 0x40)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
for s in doc.sections:
    s.top_margin = Inches(0.6); s.bottom_margin = Inches(0.6)
    s.left_margin = Inches(0.7); s.right_margin = Inches(0.7)


def title(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = NAVY
    return p

def subtitle(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.font.size = Pt(11); r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(2)
    return p

def tagline(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(8)
    return p

def heading(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = ACCENT
    return p

def body(text):
    p = doc.add_paragraph(text); p.paragraph_format.space_after = Pt(4)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2); p.add_run(text)
    return p

def item(label, text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + " — "); r.bold = True; r.font.color.rgb = NAVY
    p.add_run(text)
    return p


title("EXOS AEROSPACE")
subtitle("Capability Statement  |  Flight Services & Advanced Composite Tank Technology")
tagline("One of only three FAA-licensed reusable rocket launch providers in the United States. Flight-proven since 2018.")

heading("What We Offer")
item("Flight & Launch Services",
     "Reusable suborbital flight test, wet-lease vehicle operations, payload flight (SPACEavailable…™), "
     "and hardware de-risk and qualification — real flight data that lowers program risk before orbital commitment.")
item("Advanced Composite Tank & Pressure-Structure Technology",
     "Type V linerless composite pressurized tanks and structures, delivered through Exos's ownership position in "
     "Scorpius Space Launch Company — with hardware flown to the lunar surface. Lightweight, high-cycle-life, "
     "permeation-resistant construction that scales from launch vehicles to industrial and medical gas storage.")

heading("Markets We Serve")
item("Space & Aerospace",
     "Reusable suborbital flight test and payload services to qualify and de-risk hardware before orbital commitment; "
     "Type V linerless composite tanks for launch vehicles, satellites, and in-space cryogenic storage and refueling, "
     "with demonstrated lunar-surface heritage.")
item("Defense & Government",
     "Repeatable flight-test cadence for hypersonics, reentry, and sensor programs; responsive launch profiles with "
     "minimal site infrastructure; lightweight composite pressure vessels for mobile, deployable, and space systems. "
     "Under U.S. DoD contract to build the flight vehicle for a government-furnished LOX/methane Rotating Detonation "
     "Rocket Engine (RDRE) flight demonstration.")
item("Medical",
     "Type V linerless composite pressure-vessel technology applicable to lightweight, high-pressure medical and "
     "cryogenic gas storage and transport — lower weight, high cycle life, and strong permeation resistance for "
     "portable and facility applications.")
item("Industrial & Transport",
     "Composite pressure vessels for hydrogen, CNG/RNG, and bulk industrial gas storage and mobile transport — "
     "linerless Type V construction delivering weight savings, durability, and safety over legacy metal-lined tanks.")

heading("Why Exos")
bullet("Currently selling — an operator offering suborbital payload and flight-test services today, not a future promise.")
bullet("Flight heritage since 2018 — SARGE flown four times, demonstrating repeatable reusable suborbital operations.")
bullet("BLK3 next-generation vehicle — concept-to-hover in 12 months under a USAF Phase II SBIR; 25% larger and "
       "29% lighter via integrated pressurized structures; FAA-licensed (2026).")
bullet("TRL 9 composite tank heritage — through the Scorpius ownership position; Type V linerless tanks flown to "
       "the lunar surface as fuel and LOX tanks on commercial lunar lander missions.")
bullet("Minimal-infrastructure operations — pad plus power and comms; our mobile command capability brings the rest.")

heading("Proof Points")
bullet("SARGE: four reusable suborbital flights. BLK3: concept-to-hover in 12 months, FAA-licensed.")
bullet("NASA and DoD contract performance, including a DoD contract to build the flight vehicle for a "
       "government-furnished LOX/methane RDRE flight demonstration.")
bullet("Northrop Grumman Lunar Lander Challenge — 1st place (Round 1), 2nd place (Round 2).")
bullet("Team propulsion heritage spanning hundreds of engines (including LOX/methane) across more than a dozen "
       "flying vehicles, with reusable systems proven through hundreds of successful in-flight relights.")

heading("Who We Serve")
body("Orbital launch and satellite companies qualifying hardware  •  Defense programs needing test cadence "
     "(hypersonics, reentry, sensors)  •  Research and biotech needing microgravity and recovery  •  "
     "In-space infrastructure developers (cryogenics, refueling, structures)  •  Industrial, energy, medical, "
     "and transport users of high-performance composite pressure vessels.")

heading("Company Data")
item("CAGE", "7EMY6")
item("UEI", "CK24BPZGF854")
item("NAICS Codes", "336414, 336415, 541715")
item("PSC Codes", "1725, 1340, 2645, 1675, AR15")

heading("Let's Talk")
item("John M. Quinn", "Cofounder & CEO")
body("jquinn@exosaero.com   |   www.exosaero.com   |   Greenville, TX")

out = "/Users/exomac/.openclaw/workspace/drafts/EXOS_Capability_Statement_SalesMarketing_v2.docx"
doc.save(out)
print("Saved:", out, "| paragraphs:", len(doc.paragraphs))
