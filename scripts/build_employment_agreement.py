#!/usr/bin/env python3
"""Build the EXOS Employment Agreement template (Rev. 2026-09-28) and the standard
Separation Agreement and General Release it points to, as .docx.

Why: the 9/28 whistleblower review found Employment Agreement Sec. 6(A) conditioned
termination benefits on a "full non disclosure general release of all claims in a
form provided by the Company without alteration" with no carve-out for reporting to
the government (FAR 52.203-19, 10 USC 4701, 41 USC 4712, SEC Rule 21F-17, NLRA s7).
Sec. 6(A) also referred to "subsections (b) and (c)" that are definitions, not
benefits. The text otherwise follows the executed agreement (Victor, 7/30/2026); the
only editable source on file was that signed PDF.

Existing signers are covered by Attachment A Rev. 2026-09-28 Sec. 6.5 (signed via
MC Required Reading, hr-ciiaa); this template is for new hires.
"""
import os
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT_DIR = sys.argv[1] if len(sys.argv) > 1 else "process-library/05-human-resources"
EA_OUT = os.path.join(OUT_DIR, "EXOS_Employment_Agreement_Template_Rev_2026-09-28.docx")
REL_OUT = os.path.join(OUT_DIR, "EXOS_Separation_Agreement_General_Release_Template_Rev_2026-09-28.docx")
CO = "EXOS Aerospace Systems & Technologies, Inc."
CO_CAPS = "EXOS AEROSPACE SYSTEMS & TECHNOLOGIES, INC."


def new_doc():
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    return d


def helpers(doc):
    def h(t, lvl=1):
        doc.add_heading(t, level=lvl)

    def p(t, bold=False, italic=False):
        para = doc.add_paragraph()
        r = para.add_run(t)
        r.bold, r.italic = bold, italic
        return para

    def ul(items):
        for i in items:
            doc.add_paragraph(i, style="List Bullet")
    return h, p, ul


# Carve-out shared by both documents so the wording cannot drift.
PROTECTED = [
    "report possible violations of law, or waste, fraud, abuse, gross mismanagement or a substantial and "
    "specific danger to public health or safety related to a Federal contract or grant, to any Federal, state "
    "or local government agency or official, a Member of Congress, an Inspector General (including the DoD "
    "and NASA OIG hotlines), the Government Accountability Office, a Contracting Officer or other Federal "
    "employee responsible for contract oversight, the Department of Justice or other law enforcement, or a "
    "court or grand jury (10 U.S.C. § 4701; 41 U.S.C. § 4712);",
    "file a charge or complaint with, communicate with, provide documents or information to, or participate "
    "in an investigation or proceeding by any government agency, including the SEC, EEOC, NLRB, OSHA, the "
    "Department of Labor and the FAA, without notice to or permission from the Company;",
    "receive an award from a government agency (for example, an SEC whistleblower award) for information "
    "provided;",
    "testify truthfully in any legal proceeding, or respond to a lawful subpoena or court order;",
    "discuss or disclose wages, hours or other terms and conditions of employment, or otherwise engage in "
    "concerted activity protected by Section 7 of the National Labor Relations Act; or",
    "disclose or discuss information about sexual assault or sexual harassment disputes, or any other conduct "
    "you reasonably believe to be unlawful, to the extent that right cannot be waived by law (including the "
    "Speak Out Act, 42 U.S.C. § 19401 et seq.).",
]
DTSA = ("Defend Trade Secrets Act notice (18 U.S.C. § 1833(b)): an individual will not be held criminally or "
        "civilly liable under any Federal or state trade secret law for disclosing a trade secret that is made "
        "(i) in confidence to a Federal, state or local government official, directly or indirectly, or to an "
        "attorney, solely for the purpose of reporting or investigating a suspected violation of law; or (ii) in "
        "a complaint or other document filed in a lawsuit or other proceeding, if the filing is made under seal. "
        "An individual who files a lawsuit for retaliation for reporting a suspected violation of law may "
        "disclose the trade secret to his or her attorney and use it in the court proceeding if the individual "
        "files any document containing the trade secret under seal and does not disclose it except by court "
        "order.")
CONTROLLED = ("This does not authorize disclosure of classified information, CUI or export-controlled technical "
              "data outside authorized channels; a report involving such information may still be made by "
              "describing it generally and providing it only to an authorized recipient (such as the relevant "
              "Inspector General or Contracting Officer) through an approved channel.")


# ======================================================================
# Employment Agreement template
# ======================================================================
def build_employment_agreement():
    doc = new_doc()
    h, p, ul = helpers(doc)

    p("Revision 2026-09-28 — template for new hires. Changes from the prior form: Section 6(A) rewritten (the "
      "release may not restrict reporting to the government and must use the Company's current Separation "
      "Agreement and General Release form); new Section 6(D) Protected Activity; Sections 6(B) and 7(A) refer to "
      "Attachment A, Revision 2026-09-28.", italic=True)
    p("Re: EMPLOYMENT AGREEMENT", bold=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    p("Dear [Employee First Name]")
    p(f"On behalf of {CO_CAPS}, I am pleased to offer you the position of [Position Title] for the Company. Your "
      "employment by the Company shall be governed by the following terms and conditions (this \"Agreement\"):")

    h("1. Duties and Scope of Employment.")
    p("(A) Position. For the term of your employment under this Agreement (your \"Employment\"), the Company agrees "
      "to employ you in the aforementioned position. You will report to the Company's CEO and/or COO, or to such "
      "other person as the Company subsequently may determine. You will be working out of the Company's office in "
      "Greenville, Texas. You will perform the duties and have the responsibilities and authority customarily "
      "performed and held by an employee in your position or as otherwise may be assigned or delegated to you by "
      "the Company's CEO and/or COO.")
    p("(B) Obligations to the Company. During your Employment, you shall devote your full business efforts and "
      "time to the Company. During your Employment, without the prior written approval of the Company's Chief "
      "Executive Officer, [which approval shall not be unreasonably withheld], you shall not render services in "
      "any capacity to any other person or entity and shall not act as a sole proprietor or partner of any other "
      "person or entity or own more than five percent of the stock of any other corporation. [Notwithstanding the "
      "foregoing, you may serve on corporate, civic or charitable boards or committees, deliver lectures, fulfill "
      "speaking engagements, teach at educational institutions, or manage personal investments without such "
      "advance written consent, provided that such activities do not individually or in the aggregate interfere "
      "with the performance of your duties under this Agreement.] You shall comply with the Company's policies "
      "and rules, as they may be in effect from time to time during your Employment.")
    p("(C) No Conflicting Obligations. You represent and warrant to the Company that you are under no obligations "
      "or commitments, whether contractual or otherwise, that are inconsistent with your obligations under this "
      "Agreement. In connection with your Employment, you shall not use or disclose any trade secrets or other "
      "proprietary information or intellectual property in which you or any other person has any right, title or "
      "interest and your Employment will not infringe or violate the rights of any other person. You represent "
      "and warrant to the Company that you have returned all property and confidential information belonging to "
      "any prior employer.")
    p("(D) Commencement Date. You shall commence full-time Employment as soon as reasonably practicable and in no "
      "event later than the date agreed upon in your Offer Letter.")

    h("2. Cash and Incentive Compensation.")
    p("(A) Salary. The Company shall pay you as compensation for your services the salary stated in (and agreed "
      "upon by you in) the Offer Letter. The annual compensation specified in this subsection (A), together with "
      "any modifications in such compensation that the Company may make from time to time, is referred to in this "
      "Agreement as \"Base Salary.\" [The Board or any Compensation Committee of the Board shall review your Base "
      "Salary at least annually. Effective as of the date of any change to your Base Salary, the Base Salary as so "
      "changed shall be considered the new Base Salary for all purposes of this Agreement.]")
    p("(B) [Sales Compensation. In addition to your Base Salary, you are eligible to participate in the Company's "
      "sales incentive plan (the \"Sales Incentive Plan\"), subject to your execution of an acknowledgment of the "
      "terms of your participation in the Sales Incentive Plan (such form will be included in the plan documents "
      "which will be provided to you shortly).]")

    h("3. Vacation/PTO and Employee Benefits.")
    p("During your Employment, you shall be eligible to accrue paid vacation / paid time off, pro-rated for the "
      "remainder of this calendar year, in accordance with the Company's vacation/paid time off policy, as it may "
      "be amended from time to time. During your Employment, you shall be eligible to participate in the employee "
      "benefit plans maintained by the Company and generally available to similarly situated employees of the "
      "Company, subject in each case to the generally applicable terms and conditions of the plan in question and "
      "to the determinations of any person or committee administering such plan. Please refer to the Employee "
      "Handbook for more information.")

    h("4. Business Expenses.")
    p("The Company will reimburse you for your necessary and reasonable business expenses incurred in connection "
      "with your duties hereunder upon presentation of an itemized account and appropriate supporting "
      "documentation, all in accordance with the Company's generally applicable policies.")

    h("5. Termination.")
    p("(A) Employment at Will. Your employment shall be \"at will,\" meaning that either you or the Company shall "
      "be entitled to terminate your Employment at any time and for any reason, with or without Cause. Any "
      "contrary representations that may have been made to you shall be superseded by this Agreement. This "
      "Agreement shall constitute the full and complete agreement between you and the Company on the \"at-will\" "
      "nature of your Employment, which may only be changed in an express written agreement signed by you and a "
      "duly authorized officer of the Company.")
    p("(B) Rights Upon Termination. Except as expressly provided in Section 6, upon the termination of your "
      "Employment, you shall only be entitled to the compensation and benefits earned.")

    h("6. Termination Benefits.")
    p("(A) General Release. Any other provision of this Agreement notwithstanding, no severance pay or other "
      "termination benefit the Company may offer you (beyond compensation and benefits already earned) shall be "
      "paid unless and until (i) you have executed, and have not revoked within any revocation period required by "
      "law, a general release of claims in the form of the Company's then-current Separation Agreement and General "
      "Release, and (ii) you have returned all Company property. The release will not waive claims that cannot "
      "lawfully be waived, and any confidentiality or non-disclosure term in it is subject to Section 6(D) and to "
      "Section 6 of Attachment A.")
    p("(B) Definition of \"Cause.\" For all purposes under this Agreement, \"Cause\" shall mean: (i) any [material] "
      "breach by you of this Agreement, the Confidential Information and Invention Assignment Agreement between "
      "you and the Company (Attachment A), or any other written agreement between you and the Company, if such "
      "breach causes material harm to the Company; (ii) any [material] failure by you to comply with the "
      "Company's written policies or rules, as they may be in effect from time to time during your employment, "
      "if such failure causes material harm to the Company; (iii) your repeated failure to follow reasonable and "
      "lawful instructions from the Company or the Chief Executive Officer of the Company and your failure to "
      "cure such condition after receiving 20 days advance written notice; (iv) commission, conviction of, or a "
      "plea of \"guilty\" or \"no contest\" to, a felony under the laws of the United States or any State by you "
      "[if such felony is work-related, materially impairs your ability to perform services for the Company, or "
      "results in a material loss to the Company or material damage to the reputation of the Company]; (v) your "
      "misappropriation of funds or property of the Company; (vi) neglect of your duties; or (vii) any gross or "
      "willful misconduct by you [resulting in a material loss to the Company or material damage to the "
      "reputation of the Company]. Protected Activity under Section 6(D) is never Cause.")
    p("(C) Definition of \"Permanent Disability.\" For all purposes under this Agreement, \"Permanent Disability\" "
      "shall mean your inability to perform the essential functions of your position with or without reasonable "
      "accommodation for a period of 90 consecutive days because of your physical or mental impairment.")
    p("(D) Protected Activity. Nothing in this Agreement, Attachment A, the Employee Handbook, any release under "
      "Section 6(A), or any other agreement or policy of the Company prohibits, restricts, or requires notice to "
      "or permission from the Company before you:")
    ul(PROTECTED)
    p("The Company does not require, and will not enforce, any confidentiality agreement or statement that "
      "prohibits or restricts employees from lawfully reporting waste, fraud or abuse related to the performance "
      "of a Government contract to a designated investigative or law-enforcement representative of a Federal "
      "department or agency authorized to receive such information (FAR 52.203-19). " + CONTROLLED)
    p(DTSA)

    h("7. Pre-Employment Conditions.")
    p("(A) Confidentiality Agreement. Your acceptance of this offer and commencement of employment with the "
      "Company is contingent upon the execution, and delivery to an officer of the Company, of the Company's "
      "Confidential Information and Invention Assignment Agreement (Attachment A, Revision 2026-09-28, or its "
      "then-current revision), a copy of which is enclosed for your review and execution (the \"Confidentiality "
      "Agreement\"), prior to or on your Start Date. Attachment A is re-certified annually.")
    p("(B) Right to Work. For purposes of federal immigration law, you will be required to provide to the Company "
      "documentary evidence of your identity and eligibility for employment in the United States. Such "
      "documentation must be provided to us within three (3) business days of your Start Date, or our employment "
      "relationship with you may be terminated.")
    p("(C) [Verification of Information. This offer of employment is also contingent upon the successful "
      "verification of the information you provided to the Company during your application process, as well as a "
      "general background check performed by the Company to confirm your suitability for employment. By accepting "
      "this offer of employment, you warrant that all information provided by you is true and correct to the best "
      "of your knowledge, you agree to execute any and all documentation necessary for the Company to conduct a "
      "background check and you expressly release the Company from any claim or cause of action arising out of "
      "the Company's verification of such information.]")

    h("8. Successors.")
    p("(A) Company's Successors. This Agreement shall be binding upon any successor (whether direct or indirect "
      "and whether by purchase, lease, merger, consolidation, liquidation or otherwise) to all or substantially all "
      "of the Company's business and/or assets. For all purposes under this Agreement, the term \"Company\" shall "
      "include any successor to the Company's business or assets that becomes bound by this Agreement.")
    p("(B) Your Successors. This Agreement and all of your rights hereunder shall inure to the benefit of, and be "
      "enforceable by, your personal or legal representatives, executors, administrators, successors, heirs, "
      "distributees, devisees and legatees.")

    h("9. Miscellaneous Provisions.")
    p("(A) [Indemnification. The Company shall indemnify you to the maximum extent permitted by applicable law and "
      "the Company's Bylaws with respect to your service and you shall also be covered under a directors and "
      "officers liability insurance policy paid for by the Company to the extent that the Company maintains such "
      "a liability insurance policy now or in the future.]")
    p("(B) Notice. Notices and all other communications contemplated by this Agreement shall be in writing and "
      "shall be deemed to have been duly given when personally delivered or when mailed by U.S. registered or "
      "certified mail, return receipt requested and postage prepaid. In your case, mailed notices shall be "
      "addressed to you at the home address that you most recently communicated to the Company in writing. In "
      "the case of the Company, mailed notices shall be addressed to its headquarters, and all notices shall be "
      "directed to the attention of its Secretary.")
    p("(C) Modifications and Waivers. No provision of this Agreement shall be modified, waived or discharged "
      "unless the modification, waiver or discharge is agreed to in writing and signed by you and by an authorized "
      "officer of the Company (other than you). No waiver by either party of any breach of, or of compliance with, "
      "any condition or provision of this Agreement by the other party shall be considered a waiver of any other "
      "condition or provision or of the same condition or provision at another time.")
    p("(D) Whole Agreement. No other agreements, representations or understandings (whether oral or written and "
      "whether express or implied) which are not expressly set forth in this Agreement have been made or entered "
      "into by either party with respect to the subject matter hereof. This Agreement and the Confidentiality "
      "Agreement contain the entire understanding of the parties with respect to the subject matter hereof.")
    p("(E) Withholding Taxes. All payments made under this Agreement shall be subject to reduction to reflect "
      "taxes or other charges required to be withheld by law.")
    p("(F) Choice of Law and Severability. This Agreement shall be interpreted in accordance with the laws of the "
      "State of Texas without giving effect to provisions governing the choice of law. If any provision of this "
      "Agreement becomes or is deemed invalid, illegal or unenforceable in any applicable jurisdiction by reason "
      "of the scope, extent or duration of its coverage, then such provision shall be deemed amended to the "
      "minimum extent necessary to conform to applicable law so as to be valid and enforceable or, if such "
      "provision cannot be so amended without materially altering the intention of the parties, then such "
      "provision shall be stricken and the remainder of this Agreement shall continue in full force and effect. "
      "If any provision of this Agreement is rendered illegal by any present or future statute, law, ordinance or "
      "regulation (collectively, the \"Law\") then that provision shall be curtailed or limited only to the "
      "minimum extent necessary to bring the provision into compliance with the Law. All the other terms and "
      "provisions of this Agreement shall continue in full force and effect without impairment or limitation.")
    p("(G) No Assignment. This Agreement and all of your rights and obligations hereunder are personal to you and "
      "may not be transferred or assigned by you at any time. The Company may assign its rights under this "
      "Agreement to any entity that assumes the Company's obligations hereunder in connection with any sale or "
      "transfer of all or a substantial portion of the Company's assets to such entity.")
    p("(H) Counterparts. This Agreement may be executed in two or more counterparts, each of which shall be deemed "
      "an original, but all of which together shall constitute one and the same instrument.")

    p("We look forward to working with you. To indicate your acceptance of the Company's Employment Agreement, "
      "please sign and date this letter in the space provided below and return it to me, along with a signed and "
      "dated original copy of the Confidentiality Agreement (Attachment A), on or before your first day of "
      f"Employment with {CO}.")
    for line in [CO_CAPS, "By: ______________________________", "Name: John Quinn",
                 "Title: Cofounder and CEO", "", "ACCEPTED AND AGREED:",
                 "______________________________ (Signature)",
                 "[Employee Full Name] (Name of Employee)",
                 "______________________________ (Date of Acceptance)", "",
                 "Attachment A: Confidential Information and Invention Assignment Agreement (Revision 2026-09-28)"]:
        p(line)
    doc.save(EA_OUT)
    return EA_OUT


# ======================================================================
# Separation Agreement and General Release template
# ======================================================================
def build_release():
    doc = new_doc()
    h, p, ul = helpers(doc)

    t = doc.add_heading(f"{CO}\nSeparation Agreement and General Release", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p("Template, Revision 2026-09-28. This is the release form referred to in Section 6(A) of the Employment "
      "Agreement. Fill in the bracketed items for each use. Sections 5 and 6 may not be removed or narrowed. Each "
      "completed agreement goes to outside counsel (engaged per contract at signature) before it is signed.",
      italic=True)
    p(f"This Separation Agreement and General Release (this \"Release\") is between {CO} (the \"Company\") and "
      "[Employee Full Name] (\"you\").")

    h("1. Separation Date")
    p("Your employment with the Company ends [ended] on [Separation Date]. Whether or not you sign this Release, "
      "the Company will pay all wages earned through the Separation Date and any accrued, unused paid time off "
      "payable under Company policy, and you keep any vested benefits under Company benefit plans.")

    h("2. Separation Benefits")
    p("If you sign this Release, do not revoke it, and have returned all Company property, the Company will "
      "provide: [describe severance pay / benefit continuation, amount, and payment timing]. You would not "
      "otherwise be entitled to these benefits.")

    h("3. General Release of Claims")
    p("In exchange for the Separation Benefits, you release the Company and its officers, directors, employees, "
      "agents, successors and assigns from all claims, known or unknown, arising on or before the date you sign "
      "this Release, including claims relating to your employment or its termination and claims under Title VII "
      "of the Civil Rights Act of 1964, the Age Discrimination in Employment Act (\"ADEA\"), the Americans with "
      "Disabilities Act, the Family and Medical Leave Act, the Employee Retirement Income Security Act (other than "
      "claims for vested benefits), Chapter 21 of the Texas Labor Code, and other Federal, state or local laws, "
      "each to the extent such claims may lawfully be released.")

    h("4. Claims Not Released")
    p("This Release does not waive: (a) any right that cannot lawfully be waived by private agreement; (b) claims "
      "for unemployment or workers' compensation benefits; (c) vested benefits under any Company benefit plan; "
      "(d) claims for unpaid wages or overtime under the Fair Labor Standards Act; (e) any right to "
      "indemnification or D&O insurance coverage; (f) claims arising after the date you sign; (g) your right to "
      "challenge the validity of this Release under the ADEA; or (h) any right described in Section 5, including "
      "the right to receive a government whistleblower award.")

    h("5. Protected Activity")
    p("Nothing in this Release, the Employment Agreement, Attachment A, or any other agreement or policy of the "
      "Company prohibits, restricts, or requires notice to or permission from the Company before you:")
    ul(PROTECTED)
    p("You are not required to notify the Company of any such report or communication. The Company does not "
      "require, and will not enforce, any confidentiality term that prohibits or restricts lawfully reporting "
      "waste, fraud or abuse related to the performance of a Government contract to a designated investigative or "
      "law-enforcement representative of a Federal department or agency authorized to receive such information "
      "(FAR 52.203-19). You do, however, waive any right to recover money damages from the Company in a charge or "
      "lawsuit brought by you or on your behalf, except where such a waiver is prohibited by law and except for "
      "any government whistleblower award. " + CONTROLLED)
    p(DTSA)

    h("6. Confidentiality and Non-Disparagement, Limited by Section 5")
    p("Subject always to Section 5, you agree to keep the amount of the Separation Benefits confidential, except "
      "that you may disclose it to your spouse, attorney, accountant, tax or financial advisor, or a government "
      "agency, or as required by law. Subject always to Section 5, you agree not to make statements about the "
      "Company's products or services that you know to be false and that are made with intent to harm the "
      "Company. Nothing in this Section restricts truthful statements about your terms and conditions of "
      "employment, statements protected by the National Labor Relations Act, or statements in any legal "
      "proceeding or government investigation.")

    h("7. Continuing Obligations")
    p("Your obligations under Attachment A (Confidential Information and Invention Assignment Agreement) continue "
      "after your employment ends, subject to Sections 6 and 7 of Attachment A and Section 5 of this Release. You "
      "confirm that you have returned all Company property and Confidential Information, including anything on "
      "personal devices or accounts, and that you hold no CUI or export-controlled data outside Company systems.")

    h("8. Review Period, Revocation and Advice of Counsel (ADEA / OWBPA)")
    p("You are advised in writing to consult an attorney before signing this Release. You have [21] days after "
      "receiving it to consider it [45 days if your separation is part of a group exit incentive or termination "
      "program, in which case the Company will also provide the disclosures required by 29 U.S.C. § 626(f)(1)(H)]. "
      "You may sign sooner, but not before your Separation Date. After you sign, you have 7 days to revoke by "
      "written notice to the Company's CEO; this Release takes effect on the 8th day after you sign if you do not "
      "revoke (the \"Effective Date\"). Changes to this Release, material or not, do not restart the review "
      "period.")

    h("9. General")
    p("This Release is not an admission of wrongdoing by either party. It is governed by the laws of the State of "
      "Texas. If any provision is held unenforceable, it will be limited to the minimum extent necessary and the "
      "rest of this Release remains in effect; Section 5 is never severed. This Release, the Employment Agreement "
      "and Attachment A are the entire agreement on these subjects. It may be signed in counterparts and "
      "electronically.")

    h("Signatures")
    p("By signing, you confirm that you have read this Release, understand it, were advised to consult an "
      "attorney, had the review period described in Section 8, and sign it voluntarily.")
    for line in ["Employee Signature: ______________________________",
                 "Printed Name: [Employee Full Name]",
                 "Date: ______________________________", "",
                 CO, "By: ______________________________", "Name: John Quinn",
                 "Title: Cofounder and CEO", "Date: ______________________________"]:
        p(line)
    doc.save(REL_OUT)
    return REL_OUT


if __name__ == "__main__":
    print(build_employment_agreement())
    print(build_release())
