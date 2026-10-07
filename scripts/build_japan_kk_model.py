#!/usr/bin/env python3
"""Build v16 = v15 Exos US model + Exos Aerospace Japan K.K. sub-model + US/JP consolidation.

Adds 3 sheets to a copy of the canonical v15 workbook (existing sheets untouched):
  JP Assumptions   - all Japan inputs (yellow = assumption, Jon to confirm)
  JP K.K. Model    - K.K. P&L / cash (USD + JPY), US-side Japan program, intercompany
  Consolidated US+JP - US standalone (live links to Annual P&L / CFS) + Japan, eliminations,
                       cap table / EAI ownership, CFIUS/FOCI flag, illustrative value view
Usage: /usr/bin/python3 scripts/build_japan_kk_model.py
"""
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.properties import CalcProperties

SRC = "drafts/EXOS_Financial_Model_CommissionOnly_Consortium_v15_2026-10-06.xlsx"
OUT = "drafts/EXOS_Financial_Model_v16_US_plus_JapanKK_2026-10-07.xlsx"

YC = ["C", "D", "E", "F", "G"]          # model years 1..5 on the JP sheets
UC = ["B", "C", "D", "E", "F"]          # model years 1..5 on v15 Annual P&L / Annual CFS
FY = ["FY27", "FY28", "FY29", "FY30", "FY31"]

NAVY = "1F3864"
F_IN = PatternFill("solid", fgColor="FFF2CC")      # yellow = assumption
F_HDR = PatternFill("solid", fgColor=NAVY)
F_SEC = PatternFill("solid", fgColor="D9E1F2")
F_TOT = PatternFill("solid", fgColor="E2EFDA")
BOLD = Font(name="Arial", bold=True, size=10)
NORM = Font(name="Arial", size=10)
WHITE = Font(name="Arial", bold=True, size=11, color="FFFFFF")
NOTE = Font(name="Arial", italic=True, size=9, color="595959")
USD = '_($* #,##0_);_($* (#,##0);_($* "-"??_);_(@_)'
JPY = '"¥"#,##0;[Red]-"¥"#,##0'
PCT = "0.0%"
NUM = "#,##0"
DEC = "#,##0.00"
TOPLINE = Border(top=Side(style="thin"))


class Sheet:
    def __init__(self, wb, name, title, subtitle):
        self.ws = wb.create_sheet(name)
        self.name = name
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 58
        ws.column_dimensions["B"].width = 9
        for c in YC:
            ws.column_dimensions[c].width = 15
        ws.column_dimensions["H"].width = 70
        ws["A1"] = "EXOS AEROSPACE SYSTEMS & TECHNOLOGIES, INC.  ·  EXOS AEROSPACE JAPAN K.K."
        ws["A1"].font = WHITE
        for col in "ABCDEFGH":
            ws[f"{col}1"].fill = F_HDR
        ws["A2"] = title
        ws["A2"].font = Font(name="Arial", bold=True, size=13, color=NAVY)
        ws["A3"] = subtitle
        ws["A3"].font = NOTE
        ws["A4"] = "Model year"
        ws["A5"] = "Fiscal year (Jul–Jun; Month 1 = Jul 2026, per Controls)"
        for i, c in enumerate(YC):
            ws[f"{c}4"] = i + 1
            ws[f"{c}5"] = FY[i]
            ws[f"{c}4"].font = BOLD
            ws[f"{c}5"].font = BOLD
            ws[f"{c}4"].alignment = ws[f"{c}5"].alignment = Alignment(horizontal="center")
        ws["H5"] = "Basis / source"
        ws["H5"].font = BOLD
        ws["A4"].font = ws["A5"].font = BOLD
        ws.freeze_panes = "C6"
        self.r = 7
        self.refs = {}

    def section(self, text):
        self.r += 1
        ws = self.ws
        ws[f"A{self.r}"] = text
        for col in "ABCDEFGH":
            ws[f"{col}{self.r}"].fill = F_SEC
            ws[f"{col}{self.r}"].font = BOLD
        self.r += 1

    def row(self, key, label, vals=None, fmt=USD, unit="", inp=False, note="", single=False, total=False):
        """vals: list of 5 (values or formula strings with {c} = this column, {p} = prior column),
        or one value if single=True (lands in col C)."""
        ws, r = self.ws, self.r
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = BOLD if total else NORM
        ws[f"B{r}"] = unit
        ws[f"B{r}"].font = NOTE
        cols = YC[:1] if single else YC
        if vals is not None and not isinstance(vals, list):
            vals = [vals] * len(cols)
        for i, c in enumerate(cols):
            if vals is None:
                break
            v = vals[i]
            p = YC[i - 1] if i > 0 else None
            if isinstance(v, str) and v.startswith("="):
                v = v.replace("{c}", c).replace("{p}", p or "B").replace("{u}", UC[YC.index(c)])
            cell = ws[f"{c}{r}"]
            cell.value = v
            cell.number_format = fmt
            cell.font = BOLD if total else NORM
            if inp:
                cell.fill = F_IN
            elif total:
                cell.fill = F_TOT
                cell.border = TOPLINE
        if note:
            ws[f"H{r}"] = note
            ws[f"H{r}"].font = NOTE
            ws[f"H{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        if key:
            self.refs[key] = r
        self.r += 1
        return r

    def note(self, text):
        self.ws[f"A{self.r}"] = text
        self.ws[f"A{self.r}"].font = NOTE
        self.r += 1

    def blank(self):
        self.r += 1


def ref(sheet, key, col=None, absolute=False):
    r = sheet.refs[key]
    q = f"'{sheet.name}'!"
    if col is None:
        return q + (f"$C${r}" if absolute else f"C{r}")
    return f"{q}{col}{r}"


def build():
    shutil.copyfile(SRC, OUT)
    wb = openpyxl.load_workbook(OUT)

    # ------------------------------------------------------------------ JP Assumptions
    A = Sheet(wb, "JP Assumptions", "Japan K.K. — Assumptions (all yellow cells are inputs; Jon to confirm)",
              "Every yellow cell is a management estimate unless the Basis column says SOURCED. Change inputs here only; "
              "JP K.K. Model and Consolidated US+JP recalculate.")
    yr = lambda c: f"{c}$4"

    A.section("1. CAPITAL — Japan round (¥790M) and flow to the K.K.")
    A.row("fx_close", "FX rate at close", 157, DEC, "¥/US$", True, single=True,
          note="SOURCED: IDJ Rev2 deck build (scripts/build_idj_rev2.py, 2026-09-27).")
    A.row("fx", "FX rate for projections", 155, DEC, "¥/US$", True,
          note="Working estimate (IDJ_Toshi_Answers 2026-10-06). Flat; sensitize here.")
    A.row("raise_jpy", "Japan round size", 790_000_000, JPY, "¥", True, single=True,
          note="SOURCED: ¥790M ≈ US$5.0M (IDJ Rev2 deck 9/27; Jon 9/27 structure).")
    A.row("raise_usd", "Japan round size (US$)", f"=C{A.r-1}/C{A.r-3}", USD, "US$", single=True)
    A.row("close_yr", "Close in model year #", 1, NUM, "yr", True, single=True,
          note="Assumes close in FY27 (H1 CY2027). Reg S mechanics are triggered only once a serious investor "
               "appears (Jon 9/27), then K.K. formation 4–8 wks + FEFTA prior notification.")
    A.row("down_pct", "% of Japan proceeds contributed by Exos U.S. to the K.K.", 1.0, PCT, "%", True, single=True,
          note="Seed Answers 9/17: 'all Japan-sourced capital is directed to Exos Aerospace Japan K.K.' "
               "OPEN: the 'up to $5M core US investment into the K.K.' figure is unconfirmed (task #2928 blocker).")
    A.row("premoney", "Exos U.S. equity value basis (pre-money)", 169_700_000, USD, "US$", True, single=True,
          note="SOURCED: Eqvista 409A concluded equity value $169.7M (val. date 3/31/26) = Series A figure of record (D1–D4, 8/13). "
               "Actual round pricing is a board/Counsel decision.")
    A.row("shares_pre", "Exos U.S. shares outstanding (pre-round)", 12_320_256, NUM, "sh", True, single=True,
          note="SOURCED: cap-table/Exos_CapTable_MASTER_2026-09-20.xlsx (Tie-Out A7).")
    A.row("pps", "Implied price per mirrored Exos U.S. share", f"=C{A.r-2}/C{A.r-1}", DEC, "US$/sh", single=True,
          note="EAI units mirror Exos U.S. shares (Jon 9/27). Illustrative = pre-money ÷ shares.")
    A.row("foci", "Foreign-ownership reference line (SF-328)", 0.05, PCT, "%", True, single=True,
          note="Per Jon 9/27 structuring note: SF-328 foreign-ownership line 5%; Exos holds no FCL. CFIUS turns on RIGHTS, "
               "not % — EAI members stay passive (no board/observer seat, no technical access).")

    A.section("2. USE OF PROCEEDS (K.K.) — US$, one-time")
    A.row("uop_standup", "K.K. stand-up: formation, FEFTA filing, JP/US counsel, registration", 300_000, USD, "US$", True, single=True,
          note="ESTIMATE (IDJ_Toshi_Answers bucket 1). Formation itself ¥0.4–1M; balance = counsel/FEFTA/tax set-up.")
    A.row("uop_vehicle", "Dedicated BLK3 program deposit to Exos U.S. (prepaid wet-lease)", f"=C{A.r+1}+C{A.r+2}", USD, "US$", single=True,
          note="K.K. prepays Exos U.S.; Exos U.S. builds + OWNS the vehicle (title stays U.S. — bailment/operating-lease term sheet). "
               "Prepaid is drawn down against per-flight wet-lease fees.")
    A.row("veh_build", "   of which: BLK3 vehicle build (Exos U.S. capex)", 2_500_000, USD, "US$", True, single=True,
          note="Matches 'BLK3 vehicle build $2.5M' in v15 Funding Ask (internal build cost).")
    A.row("veh_export", "   of which: export (DSP-73/TAA), container-as-zone, ship & site activation", 700_000, USD, "US$", True, single=True,
          note="ESTIMATE. Capitalized with the vehicle.")
    A.row("uop_runway", "Operating runway (office, team, BD until flight cadence)", f"=C{A.refs['raise_usd']}*C{A.refs['down_pct']}-C{A.refs['uop_standup']}-C{A.refs['uop_vehicle']}",
          USD, "US$", single=True, note="Balancing item = contributed capital less stand-up and vehicle deposit.")
    A.row("build_y1", "Vehicle build spend in close year (balance next year)", 0.4, PCT, "%", True, single=True,
          note="~12-month build-to-order from NCE/anchor commitment (Attachment D Rev 2) straddles two fiscal years.")

    A.section("3. K.K. REVENUE DRIVERS (third-party)")
    A.row("flights", "BLK3 flights from Hokkaido Spaceport (HOSPO)", [0, 1, 4, 6, 6], NUM, "#", True,
          note="ESTIMATE. First flight FY28 (spring 2028) = close + ~13 mo build. Single vehicle; Toshi draft had 1→6+/yr. "
               "Pre-HOSPO 'bridge' flights for JP payloads on the existing BLK3 from the U.S. are already in the U.S. model — not double counted.")
    A.row("price", "Payload price per flight (full manifest)", 1_000_000, USD, "US$", True,
          note="SOURCED: Exos public price $1M/flight = $7,500/kg × 135 kg.")
    A.row("fill", "Manifest fill factor", [0.75, 0.75, 0.80, 0.85, 0.90], PCT, "%", True,
          note="ESTIMATE — conservatism haircut on sold payload mass.")
    A.row("svc", "Payload integration / engineering & consulting services", [150_000, 300_000, 500_000, 700_000, 900_000], USD, "US$", True,
          note="ESTIMATE. Starts on K.K. formation (payload readiness, verification, JAXA/university programs).")
    A.row("sslc_sales", "Japan-originated SSLC / Scorpius tank sales (third-party)", [0, 500_000, 1_000_000, 1_500_000, 2_000_000], USD, "US$", True,
          note="ESTIMATE. Incremental to the U.S. model's SSLC 3rd-party ramp (Assumptions r36). Japanese-customer LOI referenced in IDJ v7 deck.")
    A.row("sslc_rate", "SSLC sales commission rate", 0.15, PCT, "%", True, single=True,
          note="SOURCED: v15 Assumptions — 15% SSLC commission (commission-only post-ROFR).")
    A.row("nce", "Japan anchor / NCE program support (gov / prefecture)", [0, 0, 0, 0, 0], USD, "US$", True,
          note="UPSIDE — held at $0 until an anchor commitment is signed (funded-vs-demonstrated gate).")

    A.section("4. K.K. COSTS")
    A.row("ic_fee", "Intercompany wet-lease fee per flight paid to Exos U.S.", 450_000, USD, "US$", True,
          note="Transfer-pricing PLACEHOLDER (arm's-length study + JP/US tax counsel required). Eliminated on consolidation.")
    A.row("range", "HOSPO range / site fee per flight", 150_000, USD, "US$", True, note="ESTIMATE (third-party, Space Cotan / HOSPO).")
    A.row("ins", "Flight & third-party liability insurance per flight", 50_000, USD, "US$", True, note="ESTIMATE.")
    A.row("heads", "Japan headcount (country mgr, BD, ops liaison, admin)", [2, 3, 4, 5, 6], NUM, "FTE", True,
          note="ESTIMATE (Toshi draft: ~3–5 roles). U.S. flight crew deploys from the U.S. (cost sits in Exos U.S.).")
    A.row("cph", "Loaded cost per FTE (Y1, +3%/yr)", 130_000, USD, "US$", True, single=True,
          note="ESTIMATE incl. ~15% employer social insurance (PwC Japan tax summary, 2025).")
    A.row("office", "Office, rent, IT & communications (Tokyo + Hokkaido liaison)", [120_000, 150_000, 180_000, 200_000, 220_000], USD, "US$", True,
          note="ESTIMATE.")
    A.row("prof", "Recurring professional fees (JP tax/audit, export compliance, counsel)", [100_000, 150_000, 150_000, 175_000, 200_000], USD, "US$", True,
          note="ESTIMATE.")
    A.row("travel", "Travel & business development", [80_000, 100_000, 120_000, 140_000, 160_000], USD, "US$", True, note="ESTIMATE.")
    A.row("tax_jp", "Japan effective corporate tax rate", 0.3064, PCT, "%", True, single=True,
          note="SOURCED: JETRO, effective rate for non-SME (paid-in capital > ¥100M), FY starting ≥ 1 Apr 2026: 30.64% — "
               "https://www.jetro.go.jp/en/invest/setting_up/section3/page3.html . NOL simplified to 100% use.")

    A.section("5. EXOS U.S. — Japan program economics")
    A.row("us_flight_cost", "Exos U.S. direct cost per Japan flight (refurb, propellant, U.S. crew deployment)", 300_000, USD, "US$", True,
          note="ESTIMATE, INTERNAL — includes Japan deployment logistics. Do not quote externally.")
    A.row("life", "Japan BLK3 depreciable life (straight-line from first flight year)", 5, NUM, "yrs", True, single=True,
          note="ESTIMATE.")
    A.row("tax_us", "U.S. marginal tax on incremental profit (only in years U.S. model pays tax)", 0.21, PCT, "%", True, single=True,
          note="SOURCED: 21% federal corporate rate (26 U.S.C. §11(b)). Applied only when Annual P&L Income Tax > 0 (NOL years = 0).")

    for c in YC:  # close-year flag
        pass

    # ------------------------------------------------------------------ JP K.K. Model
    K = Sheet(wb, "JP K.K. Model", "Exos Aerospace Japan K.K. — P&L, cash, and Exos U.S. Japan-program economics",
              "Wholly-owned subsidiary of Exos U.S. All amounts US$ unless marked ¥. Formulas only — inputs live on JP Assumptions.")
    a = lambda key, col=None: ref(A, key, col) if col else f"'{A.name}'!$C${A.refs[key]}"
    ac = lambda key: f"'{A.name}'!{{c}}{A.refs[key]}"

    K.section("A. K.K. PROFIT & LOSS (US$)")
    K.row("close", "Close-year flag", f"=IF({{c}}$4={a('close_yr')},1,0)", NUM)
    K.row("rev_fl", "Payload flight revenue (HOSPO)", f"={ac('flights')}*{ac('price')}*{ac('fill')}")
    K.row("rev_svc", "Integration / engineering services", f"={ac('svc')}")
    K.row("rev_sslc", "SSLC commission on Japan-originated tank sales", f"={ac('sslc_sales')}*{a('sslc_rate')}")
    K.row("rev_nce", "Anchor / NCE program support (upside, $0 base)", f"={ac('nce')}")
    K.row("rev", "Total K.K. revenue (third-party)", "=SUM({c}%d:{c}%d)" % (K.r - 4, K.r - 1), total=True)
    K.row("c_ic", "Intercompany wet-lease fees to Exos U.S.", f"={ac('flights')}*{a('ic_fee')}")
    K.row("c_rng", "HOSPO range / site fees", f"={ac('flights')}*{a('range')}")
    K.row("c_ins", "Flight insurance", f"={ac('flights')}*{a('ins')}")
    K.row("cos", "Total cost of service", "=SUM({c}%d:{c}%d)" % (K.r - 3, K.r - 1), total=True)
    K.row("gp", "Gross profit", f"={{c}}{K.refs['rev']}-{{c}}{K.refs['cos']}", total=True)
    K.row("gm", "Gross margin %", f"=IFERROR({{c}}{K.refs['gp']}/{{c}}{K.refs['rev']},0)", PCT)
    K.row("o_pay", "Payroll (loaded)", f"={ac('heads')}*{a('cph')}*1.03^({{c}}$4-1)")
    K.row("o_off", "Office, rent, IT", f"={ac('office')}")
    K.row("o_prof", "Professional fees", f"={ac('prof')}")
    K.row("o_trv", "Travel & BD", f"={ac('travel')}")
    K.row("o_su", "One-time stand-up (formation, FEFTA, counsel)", f"={{c}}{K.refs['close']}*{a('uop_standup')}")
    K.row("opex", "Total operating expenses", "=SUM({c}%d:{c}%d)" % (K.r - 5, K.r - 1), total=True)
    K.row("ebitda", "K.K. EBITDA", f"={{c}}{K.refs['gp']}-{{c}}{K.refs['opex']}", total=True)
    K.row("ebitda_m", "EBITDA margin %", f"=IFERROR({{c}}{K.refs['ebitda']}/{{c}}{K.refs['rev']},0)", PCT)
    K.row("cum", "Cumulative pre-tax result", f"=SUM($C{K.r-2}:{{c}}{K.r-2})")
    K.row("tax", "Japan corporate tax", f"=MAX(0,MIN({{c}}{K.refs['ebitda']},{{c}}{K.refs['cum']}))*{a('tax_jp')}",
          note="Tax only once cumulative losses are recovered (simplified NOL).")
    K.row("ni", "K.K. net income", f"={{c}}{K.refs['ebitda']}-{{c}}{K.refs['tax']}", total=True)

    K.section("B. K.K. PROFIT & LOSS (¥)")
    K.row("rev_jpy", "Total K.K. revenue (¥)", f"={{c}}{K.refs['rev']}*{ac('fx')}", JPY)
    K.row("ebitda_jpy", "K.K. EBITDA (¥)", f"={{c}}{K.refs['ebitda']}*{ac('fx')}", JPY)
    K.row("ni_jpy", "K.K. net income (¥)", f"={{c}}{K.refs['ni']}*{ac('fx')}", JPY)

    K.section("C. K.K. CASH (US$)")
    K.row("k_open", "Opening cash", "=0")
    K.row("k_eq", "Capital contribution from Exos U.S. (Japan round proceeds)", f"={{c}}{K.refs['close']}*{a('raise_usd')}*{a('down_pct')}")
    K.row("k_dep", "Prepaid wet-lease deposit paid to Exos U.S.", f"=-{{c}}{K.refs['close']}*{a('uop_vehicle')}")
    K.row("pp_open", "   Prepaid balance — opening", "=0")
    K.row("pp_use", "   Prepaid drawn against wet-lease fees (non-cash)",
          f"=MIN({{c}}{K.refs['pp_open']}-{{c}}{K.refs['k_dep']},{{c}}{K.refs['c_ic']})",
          note="Fees are settled from the prepaid until it is exhausted, then in cash.")
    K.row("pp_close", "   Prepaid balance — closing", f"={{c}}{K.refs['pp_open']}-{{c}}{K.refs['k_dep']}-{{c}}{K.refs['pp_use']}")
    K.row("k_ni", "Net income", f"={{c}}{K.refs['ni']}")
    K.row("k_addback", "Add back: fees settled from prepaid", f"={{c}}{K.refs['pp_use']}")
    K.row("k_close", "Closing cash", f"={{c}}{K.refs['k_open']}+{{c}}{K.refs['k_eq']}+{{c}}{K.refs['k_dep']}+{{c}}{K.refs['k_ni']}+{{c}}{K.refs['k_addback']}", total=True)
    # roll-forwards (prior-year links)
    ws = K.ws
    for i, c in enumerate(YC[1:], start=1):
        p = YC[i - 1]
        ws[f"{c}{K.refs['k_open']}"] = f"={p}{K.refs['k_close']}"
        ws[f"{c}{K.refs['pp_open']}"] = f"={p}{K.refs['pp_close']}"
    K.row("k_min", "Lowest K.K. year-end cash (must stay > 0)", f"=MIN(C{K.refs['k_close']}:G{K.refs['k_close']})", single=True,
          note="If negative, the round is undersized or opex too high — adjust JP Assumptions.")

    K.section("D. EXOS U.S. — JAPAN PROGRAM (incremental to the v15 U.S. standalone model)")
    K.row("u_rev", "Intercompany wet-lease fee revenue from K.K.", f"={{c}}{K.refs['c_ic']}")
    K.row("u_cost", "Direct flight cost (refurb, propellant, crew deployment)", f"={ac('flights')}*{a('us_flight_cost')}")
    K.row("u_ebitda", "Japan program EBITDA at Exos U.S.", f"={{c}}{K.r-2}-{{c}}{K.r-1}", total=True)
    K.row("u_capex", "Japan BLK3 capex (build + export/activation)",
          f"=IF({{c}}$4={a('close_yr')},{a('build_y1')},IF({{c}}$4={a('close_yr')}+1,1-{a('build_y1')},0))*{a('uop_vehicle')}")
    K.row("u_dep", "Depreciation — Japan BLK3 (from first HOSPO flight year)",
          f"=IF(SUM('{A.name}'!$C${A.refs['flights']}:{{c}}{A.refs['flights']})>0,{a('uop_vehicle')}/{a('life')},0)")
    K.row("u_pbt", "Japan program pre-tax at Exos U.S.", f"={{c}}{K.refs['u_ebitda']}-{{c}}{K.refs['u_dep']}")
    K.row("u_tax", "U.S. tax on Japan program", f"=IF('Annual P&L'!{{u}}48>0,MAX(0,{{c}}{K.r-1})*{a('tax_us')},0)")
    K.row("u_ni", "Japan program net income at Exos U.S.", f"={{c}}{K.r-2}-{{c}}{K.r-1}", total=True)
    K.blank()
    K.row("u_c_eq", "Cash: EAI purchase of mirrored Exos U.S. shares (Reg S)", f"={{c}}{K.refs['close']}*{a('raise_usd')}")
    K.row("u_c_down", "Cash: capital contribution to K.K.", f"=-{{c}}{K.refs['k_eq']}")
    K.row("u_c_dep", "Cash: prepaid wet-lease deposit received from K.K.", f"=-{{c}}{K.refs['k_dep']}")
    K.row("u_c_fee", "Cash: wet-lease fees received beyond prepaid", f"={{c}}{K.refs['u_rev']}-{{c}}{K.refs['pp_use']}")
    K.row("u_c_cost", "Cash: direct flight costs", f"=-{{c}}{K.refs['u_cost']}")
    K.row("u_c_capex", "Cash: Japan BLK3 capex", f"=-{{c}}{K.refs['u_capex']}")
    K.row("u_c_tax", "Cash: U.S. tax", f"=-{{c}}{K.refs['u_tax']}")
    K.row("u_c_net", "Net cash impact on Exos U.S. (year)", "=SUM({c}%d:{c}%d)" % (K.r - 7, K.r - 1), total=True)
    K.row("u_c_cum", "Cumulative cash impact on Exos U.S.", f"=SUM($C{K.r-1}:{{c}}{K.r-1})")

    # ------------------------------------------------------------------ Consolidated
    C = Sheet(wb, "Consolidated US+JP", "Consolidated view — Exos U.S. (v15 standalone) + Exos Aerospace Japan K.K.",
              "U.S. lines link live to Annual P&L / Annual CFS. Intercompany wet-lease fees and the prepaid deposit eliminate on consolidation. "
              "Management-prepared projections; not audited; not an offer of securities.")
    k = lambda key: f"'{K.name}'!{{c}}{K.refs[key]}"

    C.section("REVENUE")
    C.row("us_rev", "Exos U.S. standalone revenue (v15)", "='Annual P&L'!{u}14")
    C.row("ic_rev", "Exos U.S. intercompany wet-lease revenue", f"={k('u_rev')}")
    C.row("jp_rev", "Japan K.K. revenue (third-party)", f"={k('rev')}")
    C.row("elim_rev", "Elimination: intercompany wet-lease fees", f"=-{k('c_ic')}")
    C.row("rev", "CONSOLIDATED REVENUE", "=SUM({c}%d:{c}%d)" % (C.r - 4, C.r - 1), total=True)
    C.row("jp_share", "Japan-driven share of consolidated revenue", f"=IFERROR({{c}}{C.refs['jp_rev']}/{{c}}{C.refs['rev']},0)", PCT)
    C.row("rev_g", "Consolidated revenue uplift vs U.S. standalone", f"=IFERROR({{c}}{C.refs['rev']}/{{c}}{C.refs['us_rev']}-1,0)", PCT)

    C.section("EBITDA → NET INCOME")
    C.row("us_e", "Exos U.S. standalone EBITDA (v15)", "='Annual P&L'!{u}39")
    C.row("pr_e", "Japan program EBITDA at Exos U.S. (incl. intercompany fees)", f"={k('u_ebitda')}")
    C.row("jp_e", "Japan K.K. EBITDA (after intercompany fees)", f"={k('ebitda')}",
          note="Intercompany fee is revenue in the row above and cost in this row — nets to zero; no further adjustment needed.")
    C.row("ebitda", "CONSOLIDATED EBITDA", "=SUM({c}%d:{c}%d)" % (C.r - 3, C.r - 1), total=True)
    C.row("ebitda_m", "Consolidated EBITDA margin %", f"=IFERROR({{c}}{C.r-1}/{{c}}{C.refs['rev']},0)", PCT)
    C.row("da", "D&A (U.S. v15 + Japan BLK3)", f"='Annual P&L'!{{u}}42+{k('u_dep')}")
    C.row("pbt", "Consolidated profit before tax", f"={{c}}{C.refs['ebitda']}-{{c}}{C.refs['da']}", total=True)
    C.row("tax", "Income tax (U.S. v15 + U.S. Japan program + Japan)", f"='Annual P&L'!{{u}}48+{k('u_tax')}+{k('tax')}")
    C.row("pat", "CONSOLIDATED PROFIT AFTER TAX", f"={{c}}{C.refs['pbt']}-{{c}}{C.refs['tax']}", total=True)
    C.row("us_pat", "memo: Exos U.S. standalone PAT (v15)", "='Annual P&L'!{u}50")
    C.row("pat_up", "PAT uplift from Japan", f"={{c}}{C.refs['pat']}-{{c}}{C.refs['us_pat']}")

    C.section("CASH")
    C.row("us_cash", "Exos U.S. standalone closing cash (v15 Annual CFS)", "='Annual CFS'!{u}34")
    C.row("pr_cash", "Cumulative Japan-program cash at Exos U.S.", f"={k('u_c_cum')}")
    C.row("jp_cash", "Japan K.K. closing cash", f"={k('k_close')}")
    C.row("cash", "CONSOLIDATED CLOSING CASH", "=SUM({c}%d:{c}%d)" % (C.r - 3, C.r - 1), total=True,
          note="Prepaid deposit is intercompany: K.K. outflow = U.S. inflow, nets to zero here.")
    C.row("capex", "Consolidated capex (U.S. v15 + Japan BLK3)", f"=-'Annual CFS'!{{u}}21+{k('u_capex')}")
    C.row("eq_in", "Equity raised (U.S. v15 Series A + Japan EAI Reg S)", f"='Annual CFS'!{{u}}24+{k('u_c_eq')}",
          note="v15 shows $10M Series A in FY28; Y1 -$150K is a v15 entry carried as-is.")

    C.section("OWNERSHIP — EAI (Japan) units mirror Exos U.S. shares")
    C.row("sh_pre", "Shares outstanding pre-round", f"={a('shares_pre')}", NUM, single=True)
    C.row("pps", "Price per share (illustrative, 409A equity basis)", f"={a('pps')}", DEC, single=True)
    C.row("sh_jp", "Mirrored shares held via EAI (Japan round)", f"={a('raise_usd')}/C{C.r-1}", NUM, single=True)
    C.row("sh_sa", "Series A shares (U.S. $10M, same price, illustrative)", f"=SUMIF('Annual CFS'!B24:F24,\">0\")/C{C.r-2}", NUM, single=True,
          note="Illustrative only — Series A pricing is set separately.")
    C.row("sh_post", "Shares post both rounds", f"=C{C.r-4}+C{C.r-2}+C{C.r-1}", NUM, single=True)
    C.row("own_jp", "EAI (Japan) ownership of Exos U.S. — post both rounds", f"=C{C.r-3}/C{C.r-1}", PCT, single=True)
    C.row("own_jp_only", "EAI (Japan) ownership — Japan round only", f"=C{C.r-4}/(C{C.r-6}+C{C.r-4})", PCT, single=True)
    C.row("foci", "Below SF-328 foreign-ownership line?", f'=IF(MAX(C{C.r-2},C{C.r-1})<{a("foci")},"YES — below line","NO — review with Counsel")',
          "@", single=True, note="Passive-rights guardrails apply regardless of % (CFIUS turns on rights).")

    C.section("ILLUSTRATIVE VALUE VIEW — internal; Counsel (Keyes) review before any investor use")
    C.note("Not a forecast of investor returns. EV/EBITDA on FY31 consolidated EBITDA; ignores net debt, future dilution, DLOM and timing.")
    for m in (8, 12, 16):
        C.row(f"ev{m}", f"Consolidated EV @ {m}x FY31 EBITDA", f"={m}*G{C.refs['ebitda']}", single=True,
              note="16x = multiple used in v15 Valuation Summary market approach." if m == 16 else "")
        C.row(f"eai{m}", f"   Value attributable to EAI (Japan) mirrored shares @ {m}x", f"=C{C.r-1}*C{C.refs['own_jp']}", single=True)
        C.row(f"moic{m}", f"   Multiple of Japan round invested @ {m}x", f"=C{C.r-1}/{a('raise_usd')}", '0.0"x"', single=True)
    C.row("jp_ev", "Japan program FY31 EBITDA (K.K. + U.S. program) × 12x", f"=12*('{K.name}'!G{K.refs['ebitda']}+'{K.name}'!G{K.refs['u_ebitda']})", single=True,
          note="Stand-alone value created by the Japan program vs the ¥790M (~$5.0M) deployed.")

    C.section("OPEN ITEMS (Jon to confirm before investor release)")
    for t in [
        "1. 'Up to $5M core U.S. investment into the K.K.' — not in any document; model uses 100% of the ¥790M round downstreamed (JP Assumptions, '% of Japan proceeds contributed').",
        "2. Use-of-proceeds split (stand-up / BLK3 deposit / runway) and Japan flight cadence (0/1/4/6/6).",
        "3. Intercompany wet-lease fee ($450K/flight) — needs a transfer-pricing study + JP/US tax counsel.",
        "4. Round price per mirrored share — illustrative at the 409A equity basis; board + Counsel set actual terms.",
        "5. v15 inherited items: Valuation Summary flagged STALE (pre-10/6 DCF); Annual Balance Sheet check = -$186K in every year; "
        "v15 Decision D-MTP-CONSORTIUM-MA-V15 still pending.",
        "6. Investor-NDA copy: strip ChangeLog, FY2025 Actuals (forward-looking only) and the internal flight-cost input before sharing.",
    ]:
        C.note(t)

    # ------------------------------------------------------------------ ChangeLog + calc
    cl = wb["ChangeLog"]
    r = cl.max_row + 2
    cl.cell(r, 1, "v16 — Japan K.K. sub-model + US/JP consolidation (2026-10-07, task #2928)")
    rows = [
        ("JP Assumptions", "new sheet", None, "All Japan inputs (¥790M EAI Reg S round, FX, use of proceeds, HOSPO flights, K.K. costs, tax)", "@Exobot request on #2928"),
        ("JP K.K. Model", "new sheet", None, "K.K. P&L (US$/¥), cash with prepaid wet-lease, Exos U.S. Japan-program economics", "integrate Japan revenue + investment into U.S. model"),
        ("Consolidated US+JP", "new sheet", None, "U.S. v15 (live links) + K.K. − intercompany eliminations; ownership/SF-328 flag; illustrative value view", "investor-NDA consolidated view"),
        ("(all v15 sheets)", "—", None, "UNCHANGED", "v15 standalone remains the U.S. base case"),
    ]
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            cl.cell(r + i, j, v)

    # sheet order: put new sheets right after Annual Balance Sheet
    order = wb._sheets
    for name in ["Consolidated US+JP", "JP K.K. Model", "JP Assumptions"]:
        s = wb[name]
        order.remove(s)
        order.insert(order.index(wb["Annual Balance Sheet"]) + 1, s)
    for s in ("JP Assumptions", "JP K.K. Model", "Consolidated US+JP"):
        wb[s].sheet_properties.tabColor = "C00000"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(OUT)
    print(OUT)


NDA_OUT = OUT.replace(".xlsx", "_NDA_DRAFT.xlsx")


def build_nda():
    """Investor-under-NDA copy: forward-looking only, no internal change log / open items / value view."""
    wb = openpyxl.load_workbook(OUT)
    for name in ("ChangeLog", "FY2025 Actuals"):
        if name in wb.sheetnames:
            del wb[name]
    ws = wb["Consolidated US+JP"]
    cut = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("ILLUSTRATIVE VALUE VIEW"))
    for r in range(cut, ws.max_row + 1):
        for c in range(1, 9):
            cell = ws.cell(r, c)
            cell.value = None
            cell.fill = PatternFill()
    ws.cell(cut, 1, "CONFIDENTIAL — provided under NDA. Management-prepared projections (not audited); estimates, not guarantees. "
                    "Not an offer to sell or a solicitation to buy any security.").font = NOTE
    a = wb["JP Assumptions"]
    for r in range(1, a.max_row + 1):
        h = a.cell(r, 8)
        if isinstance(h.value, str) and "INTERNAL" in h.value:
            h.value = "ESTIMATE."
    for name in wb.sheetnames:
        wb[name]["H2" if name in ("JP Assumptions", "JP K.K. Model", "Consolidated US+JP") else "Z1"] = None
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(NDA_OUT)
    print(NDA_OUT)


if __name__ == "__main__":
    build()
    build_nda()
