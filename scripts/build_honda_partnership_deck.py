#!/usr/bin/env python3
"""Exos Aerospace x Honda -- Partnership Proposal (external, for Honda R&D). Hokkaido visit 13-15 Oct 2026.
Derived from scripts/build_exos_sld30_briefing.py (left untouched). Builds an editable 16:9 PPTX + HTML -> PDF twin
(Chrome headless; LibreOffice hangs on this host). Brand: black/grey + single orange accent #F26A21; no Honda logo.
Public-level information only. No Exos prices. Every factual claim cites SOURCES."""
import base64, os, subprocess, html as H
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

WS = "/Users/exomac/.openclaw/workspace"; D = os.path.join(WS, "drafts"); OUTD = os.path.join(D, "Honda")
STEM = "Exos_Honda_Partnership_Briefing_2026-10-13"
LOGO = os.path.join(D, "_exos_white_official.png")
HERO = os.path.join(D, "_exos_hero_bg.png")
PLANET_A = os.path.join(D, "_exos_planet_a.png"); PLANET_B = os.path.join(D, "_exos_planet_b.png")
COBRAND = "Exos Aerospace  ·  Honda"
DRAFT = "DRAFT — for discussion with Honda  ·  Concepts proposed; no commitments or approvals implied  ·  Public-level information only"

SOURCES = [
 "[1] Honda Global, “Honda Conducts Successful Launch and Landing Test of Experimental Reusable Rocket,” 17 Jun 2025 — global.honda/en/topics/2025/c_2025-06-17ceng.html",
 "[2] Honda R&D, “SPACETIDE 2025 — Honda’s Challenge: From Earth to Outer Space,” 2025 — global.honda/en/RandD/activity/rdtopics/spacetide2025",
 "[3] Hokkaido Spaceport (HOSPO), “HOSPO Transitions to Concession Model — SPACE COTAN Begins Operation of LC1, Runway, Hangar, and SORA on October 1,” 2026 — hokkaidospaceport.com/en/news/2306",
 "[4] Exos Aerospace, “Reusable Flight Test Infrastructure / Capabilities” — exosaero.com/capabilities (accessed Oct 2026)",
 "[5] 22 CFR §124.1 and §120.32 (ITAR), eCFR current — ecfr.gov/current/title-22/chapter-I/subchapter-M",
 "[6] Range Commanders Council, RCC 319-25, Flight Termination Systems Commonality Standard, Jun 2025 — trmc.osd.mil/wiki/download/attachments/113019893/319-25_FTS_Commonality.pdf",
 "[7] FAA, Advisory Circular 450.143-1, Reduced Reliability Flight Safety System Design, Test, and Documentation, Federal Register, 14 May 2025 — federalregister.gov/d/2025-08496",
 "[8] Hogan Lovells, “Japan’s New Chip Equipment Export Rules Take Effect,” 2023 (FEFTA Arts. 48(1), 25(1): METI approval for listed goods/technology) — hlc.com/en/publications/japans-new-chip-equipment-export-rules-take-effect",
]

CONTACT = [("John (Jon) Quinn", True), ("Cofounder & CEO", False),
           ("Exos Aerospace Systems & Technologies, Inc.", False),
           ("jquinn@exosaero.com", False), ("+1 (972) 740-8355", False), ("www.exosaero.com", False)]

S = []
S.append(dict(kind="hero", title=["Exos Aerospace × Honda", "Partnership Proposal"],
  l1="Hokkaido International Trade Zone  ·  Communications Systems  ·  Flight Termination Systems",
  l2="Prepared for Honda R&D Co., Ltd.  ·  Hokkaido, 13–15 October 2026",
  l3="Discussion draft — concepts proposed for mutual consideration"))

S.append(dict(kind="split", title="About Exos Aerospace", items=[
 "Exos Aerospace Systems & Technologies, Inc. — Greenville, Texas, USA [4]",
 "FAA-licensed reusable suborbital launch operator [4] — licensed under 14 CFR Part 431 and transitioning to Part 450",
 "Reusable flight test infrastructure for research, commercial and government payloads: fly, recover, inspect, iterate [4]",
 "Payload recovery for post-flight inspection and analysis [4]",
 "BLK3, our next reusable vehicle: preparing for flight, early 2027",
 "NIST SP 800-171 implemented; CMMC 2.0 self-assessment complete (via Atomus)",
], panel_title="With respect for Honda’s progress", panel=[
 "17 June 2025: Honda R&D launch and landing test at its Taiki, Hokkaido facility, after phased combustion and hovering tests [1][2]",
 "Reached 271.4 m and landed within 37 cm of the target, 56.6 s flight [1]",
 "Vehicle 6.3 m long, 85 cm diameter [1]",
 "Goal: technological capability for a suborbital launch by 2029 [1]",
 "We come to learn from Honda as much as to offer",
]))

S.append(dict(kind="cards", title="Three focus areas for our work together", cards=[
 ("01  ·  HOKKAIDO ITZ", "International Trade Zone at Hokkaido Spaceport",
  ["Container-as-Zone concept with HOSPO at Taiki", "Honda as neighbour, partner, supplier and customer", "Begins with a non-binding MOU"]),
 ("02  ·  COMMUNICATIONS", "Honda-built communications systems",
  ["Learn Honda’s telemetry, tracking and command approach", "Evaluation units flown on Exos reusable flights", "Path to a supply agreement"]),
 ("03  ·  FLIGHT TERMINATION", "Honda-built flight termination systems",
  ["Honda’s rocket carries a safety system that keeps it within a pre-defined flight corridor [1]", "Review against U.S. range safety standards", "Path to a supply agreement"]),
], note="Linking all three: Honda hardware ↔ Exos in-kind reusable flight services, with an Exos-initiated U.S. export authorization (TAA)."))

S.append(dict(kind="split", title="1 · Hokkaido International Trade Zone (concept)", items=[
 "Taiki Town hosts both Honda’s rocket test facility and Hokkaido Spaceport (HOSPO) [1][3]",
 "SPACE COTAN began operating HOSPO under a 15-year concession on 1 October 2026; LC1 operations are scheduled to begin by the end of FY2026 [3]",
 "Container-as-Zone: Exos (or its U.S. contractor) performs all integration of U.S. hardware inside a monitored container; partners deliver payloads only",
 "Temporary admission for the reusable vehicle and payloads — nothing permanently imported or exported",
 "First step: a non-binding MOU with SPACE COTAN / HOSPO — no controlled technical data exchanged",
], panel_title="Possible roles for Honda", panel=[
 "Co-located test and operations presence at Taiki",
 "Hardware supplier: communications, flight termination, ground support",
 "Payload customer for reusable suborbital flights",
 "Shared learning on range, weather and logistics",
], note="Concept only. No approval from the Government of Japan, the U.S. Government or HOSPO exists today; all elements are subject to review by the relevant authorities."))

S.append(dict(kind="split", title="2 · Communications systems", items=[
 "Exos objective: evaluate, and potentially acquire, Honda-built communications systems for Exos reusable vehicles",
 "Areas of interest: telemetry downlink, command uplink, GNSS tracking, and antennas that survive repeated flights",
 "Proposed path: information exchange under NDA → evaluation units → ground test → hosted flight on Exos reusable vehicles → supply agreement",
 "Value to Honda: real-flight data on Honda hardware, recovered for inspection after flight [4]",
 "To be confirmed: Japanese export approval for Honda hardware or technology (METI, Foreign Exchange and Foreign Trade Act) [8]; U.S. import classification; spectrum licensing",
], panel_title="What we do not yet know", panel=[
 "Exos has no knowledge of Honda’s communications equipment today",
 "We would be grateful for Honda’s answers to the questions on slide 11",
 "Payment proposed in kind: Exos reusable flight services (slide 7)",
]))

S.append(dict(kind="split", title="3 · Flight termination systems", items=[
 "Honda’s experimental rocket is equipped with a safety system to prevent deviation from a pre-defined flight corridor, speed and attitude conditions [1]",
 "Exos objective: evaluate, and potentially acquire, Honda-built flight termination / flight safety systems for Exos reusable vehicles",
 "U.S. Government ranges apply RCC 319, the Flight Termination Systems Commonality Standard (current edition RCC 319-25, June 2025) [6]",
 "FAA guidance AC 450.143-1 gives a means of compliance for flight termination system design, test and documentation, traced to RCC 319-19 [7]",
 "Proposed path: information exchange → review against RCC 319 / Part 450 (under an approved TAA) → evaluation units → qualification → supply agreement",
 "To be confirmed: Japanese export approval (METI) [8]; U.S. import classification",
], panel_title="What we do not yet know", panel=[
 "Exos has no knowledge of Honda’s flight termination equipment today",
 "We would be grateful for Honda’s answers to the questions on slide 10",
 "Payment proposed in kind: Exos reusable flight services (slide 7)",
]))

S.append(dict(kind="model", title="Partnering model: hardware ↔ in-kind flight services"))

S.append(dict(kind="split", title="Export compliance path: an Exos-initiated TAA", items=[
 "Under ITAR 22 CFR §124.1, U.S. Department of State (DDTC) approval of an agreement is required before a U.S. person provides defense services to foreign persons — even when the information is in the public domain [5]",
 "Exos will initiate and prepare a Technical Assistance Agreement (TAA) so Exos engineers can share reusable-rocket capabilities with Honda engineers",
 "Honda’s part: the topics of interest, the Honda entities and people involved, and the intended end use",
 "Until approval, our discussions stay non-technical and public-level. Exos may receive Honda information; it is Exos’s sharing that is controlled",
 "Honda’s own sharing with Exos may need Japanese approval (METI, FEFTA Art. 25(1)) — for Honda to confirm [8]",
 "Approval and timing are decided by the U.S. Government",
], panel_title="Exos export compliance", panel=[
 "Empowered Official: John Quinn, Cofounder & CEO",
 "Technology Control Plan for the engagement",
 "NIST SP 800-171 implemented",
]))

S.append(dict(kind="path", title="Path forward and timeline for each area",
 cols=["", "Step 1  ·  Oct–Dec 2026", "Step 2  ·  H1 2027", "Step 3  ·  2027 onward"],
 rows=[
  ("Partnership & TAA", "Mutual NDA; agree topics; Exos prepares and files the TAA", "DDTC review; framework partnering agreement with valuation schedule", "Technical exchange under the approved TAA"),
  ("Hokkaido ITZ", "Non-binding MOU with SPACE COTAN (HOSPO); define Honda’s role", "Feasibility study; consultations with Japanese and U.S. authorities", "Pilot operations, subject to all approvals"),
  ("Communications", "Honda answers to slide 11; public-level datasheets", "Evaluation units; bench and ground tests", "Hosted flights on Exos reusable vehicles; supply agreement"),
  ("Flight termination", "Honda answers to slide 10; public-level datasheets", "Review against RCC 319 / Part 450 under the TAA; evaluation units", "Qualification; supply agreement"),
 ], note="Timing is notional and depends on export approvals in both countries. BLK3: preparing for flight, early 2027."))

S.append(dict(kind="questions", title="Questions for Honda: flight termination systems", qs=[
 "Architecture: command destruct, autonomous (AFTS), or a hybrid?",
 "Termination method: thrust termination, destruct, or both?",
 "Command receiver / decoder: type, frequency band and coding?",
 "Ordnance-based or non-pyrotechnic? If ordnance, which initiators and safe-and-arm device?",
 "Which range safety standard was it qualified to (Japanese guidelines, RCC 319, other)?",
 "Redundancy: single or dual string; independent power and batteries?",
 "Flight heritage: number of flights and tests on the current design?",
 "Export classification in Japan (METI list item) and any end-use conditions?",
 "Is it available for supply outside Honda? Typical lead time and quantities?",
 "Which test and qualification data could be shared (environmental, EMI, reliability)?",
]))

S.append(dict(kind="questions", title="Questions for Honda: communications systems", qs=[
 "Telemetry downlink: frequency band (S-band or other) and modulation?",
 "Data rates for telemetry and any video?",
 "GNSS / tracking: receiver type, constellations, update rate?",
 "Command uplink: is there one, and how is it authenticated?",
 "Encryption: is the link encrypted, and to what standard?",
 "Antennas: design, placement and survival over repeated flights and landings?",
 "Ground station: Honda-owned? Compatible with common range telemetry standards?",
 "Spectrum: how are frequencies licensed in Japan; usable in the U.S.?",
 "Heritage: number of flights, and reuse cycles per unit?",
 "Availability for supply outside Honda; lead time; support and refurbishment?",
]))

S.append(dict(kind="close", title="Next steps", items=[
 "Agree a mutual NDA and a single point of contact on each side",
 "Honda’s answers to the questions on slides 10–11, at the level Honda is comfortable sharing",
 "Exos prepares the TAA scope with Honda’s input",
 "Explore a non-binding framework partnering agreement: hardware ↔ in-kind reusable flight services",
 "Consider Honda’s interest in the Hokkaido ITZ concept with SPACE COTAN",
]))

# ------------------------------------------------------------------ model diagram geometry (inches)
MODEL = [
 ("box", 0.64, 1.6, 4.3, 3.3, "HONDA PROVIDES", ["Flight termination system units", "Communications system units", "Possibly test support at Taiki", "Engineering insight into Honda hardware"], "side"),
 ("box", 8.39, 1.6, 4.3, 3.3, "EXOS PROVIDES (IN KIND)", ["Reusable suborbital flight services", "Reusable-rocket consulting under an approved TAA", "ITZ integration and operations (concept)", "Real-flight data, with hardware recovered"], "side"),
 ("mid", 5.2, 1.6, 2.93, 3.3, "AGREED VALUATION SCHEDULE", ["Both sides valued at standard commercial pricing", "Balanced and transparent", "Reviewed jointly each step"]),
 ("bar", 0.64, 5.15, 12.05, 1.25, "PRINCIPLES",
  "Start small (evaluation units ↔ hosted flights) and grow with trust  ·  In-kind services are valued, never free  ·  Each step subject to export approvals in Japan and the U.S.  ·  Separate, non-binding framework first; definitive agreements later"),
]

# ---------------------------------------------------------------- HTML twin
def b64(p): return base64.b64encode(open(p, "rb").read()).decode()
logo64, hero64, pa64, pb64 = b64(LOGO), b64(HERO), b64(PLANET_A), b64(PLANET_B)
e = lambda s: H.escape(s)
def bl(items, cls=""):
    return f'<ul class="{cls}">' + "".join(f"<li>{e(t)}</li>" for t in items) + "</ul>"
def chrome_html(title, body, band=False):
    bnd = f'<div class="band" style="background-image:url(data:image/png;base64,{pb64})"></div>' if band else ""
    return (f'<div class="slide dark">{bnd}<img class="corner" src="data:image/png;base64,{logo64}">'
            f'<div class="stitle">{e(title)}</div><div class="rule"></div>{body}<div class="draft">{e(DRAFT)}</div></div>')
def panel_html(t, items):
    return f'<div class="panel"><div class="ph">{e(t)}</div>' + "".join(f"<div>{e(x)}</div>" for x in items) + "</div>"
def note_html(s): return f'<div class="bnote">{e(s["note"])}</div>' if s.get("note") else ""

parts = []
for n, s in enumerate(S, 1):
    k = s["kind"]
    if k == "hero":
        parts.append(f'''<div class="slide hero" style="background-image:linear-gradient(180deg,rgba(0,0,0,.55),rgba(0,0,0,.35) 45%,rgba(0,0,0,.92) 70%),url(data:image/png;base64,{hero64})">
          <img class="hlogo" src="data:image/png;base64,{logo64}"><div class="cob">{e(COBRAND)}</div>
          <div class="htitle">{e(s["title"][0])}<br><span class="o">{e(s["title"][1])}</span></div>
          <div class="hsub"><div class="l1">{e(s["l1"])}</div><div class="l2">{e(s["l2"])}</div><div class="l3">{e(s["l3"])}</div></div>
          <div class="hfoot">{e(DRAFT)}</div></div>''')
    elif k == "split":
        parts.append(chrome_html(s["title"], f'<div class="lcol">{bl(s["items"], "narrow")}</div>' + panel_html(s["panel_title"], s["panel"]) + note_html(s), s.get("band")))
    elif k == "cards":
        cs = "".join(f'<div class="card"><div class="ck">{e(a)}</div><div class="ct">{e(b)}</div>' + "".join(f"<p>{e(x)}</p>" for x in c) + "</div>" for a, b, c in s["cards"])
        parts.append(chrome_html(s["title"], f'<div class="cards">{cs}</div>' + note_html(s)))
    elif k == "model":
        o = []
        for g in MODEL:
            if g[0] in ("box", "mid"):
                _, x, y, w, h, t, its = g[:7]
                cls = "side" if g[0] == "box" else "mid"
                o.append(f'<div class="g {cls}" style="left:{x}in;top:{y}in;width:{w}in;height:{h}in"><b>{e(t)}</b>' + "".join(f"<p>{e(i)}</p>" for i in its) + "</div>")
            else:
                _, x, y, w, h, t, sub = g
                o.append(f'<div class="g bar" style="left:{x}in;top:{y}in;width:{w}in;height:{h}in"><b>{e(t)}</b><span>{e(sub)}</span></div>')
        o.append('<svg class="garrow" style="left:4.98in;top:2.95in;width:0.2in;height:0.5in" viewBox="0 0 100 100" preserveAspectRatio="none"><polygon points="0,30 60,30 60,0 100,50 60,100 60,70 0,70" fill="#F26A21"/></svg>')
        o.append('<svg class="garrow" style="left:8.15in;top:2.95in;width:0.2in;height:0.5in;transform:scaleX(-1)" viewBox="0 0 100 100" preserveAspectRatio="none"><polygon points="0,30 60,30 60,0 100,50 60,100 60,70 0,70" fill="#F26A21"/></svg>')
        parts.append(chrome_html(s["title"], "".join(o)))
    elif k == "path":
        th = "".join(f"<th>{e(c)}</th>" for c in s["cols"])
        tr = "".join("<tr>" + f'<td class="rh">{e(r[0])}</td>' + "".join(f"<td>{e(c)}</td>" for c in r[1:]) + "</tr>" for r in s["rows"])
        parts.append(chrome_html(s["title"], f'<table class="pt"><tr>{th}</tr>{tr}</table>' + note_html(s)))
    elif k == "questions":
        qs = s["qs"]; half = (len(qs) + 1) // 2
        col = lambda a, off: "".join(f'<div class="q"><em>{i+off}</em><span>{e(q)}</span></div>' for i, q in enumerate(a))
        parts.append(chrome_html(s["title"], f'<div class="qcols"><div>{col(qs[:half],1)}</div><div>{col(qs[half:],half+1)}</div></div>'
                                 '<div class="bnote">We would be grateful for whatever Honda is comfortable sharing at a public, non-technical level now; detail can follow under NDA and any required approvals.</div>'))
    elif k == "close":
        src = "".join(f"<div>{e(x)}</div>" for x in SOURCES)
        parts.append(f'''<div class="slide hero" style="background-image:linear-gradient(180deg,rgba(0,0,0,.88),rgba(6,6,8,.7) 46%,rgba(0,0,0,.94)),url(data:image/png;base64,{pa64})">
          <img class="hlogo" src="data:image/png;base64,{logo64}"><div class="cob">{e(COBRAND)}</div>
          <div class="ctitle">{e(s["title"])}</div><div class="crule"></div>{bl(s["items"], "hul")}
          <div class="contact"><div class="ph">Contact</div>{"".join(f'<div class="{"cb" if b else ""}">{e(t)}</div>' for t, b in CONTACT)}
          <div class="ty">Thank you — ありがとうございます</div></div>
          <div class="sources"><b>Sources</b>{src}</div><div class="hfoot2">{e(DRAFT)}</div></div>''')

CSS = '''
@page { size:13.333in 7.5in; margin:0; }
*{margin:0;padding:0;box-sizing:border-box;-webkit-print-color-adjust:exact;print-color-adjust:exact;font-family:"Helvetica Neue",Arial,"Hiragino Sans",sans-serif;}
.slide{width:13.333in;height:7.5in;position:relative;overflow:hidden;page-break-after:always;}
.dark{background:#0d0d0f;color:#fff;padding:0.5in 0.64in;}
.band{position:absolute;left:0;right:0;bottom:0;height:0.9in;background-size:cover;background-position:center 30%;opacity:.55;}
.corner{position:absolute;top:0.45in;right:0.62in;height:0.34in;}
.stitle{font-size:28px;font-weight:800;color:#fff;letter-spacing:-0.3px;max-width:10.4in;}
.rule{width:1.15in;height:4px;background:#F26A21;margin:0.14in 0 0.3in;}
.dark ul{list-style:none;}
.dark li{position:relative;padding-left:24px;margin-bottom:12px;font-size:16.5px;line-height:1.3;color:#ececef;}
.dark li:before{content:"\\25AA";position:absolute;left:0;top:1px;color:#F26A21;font-size:14px;}
.lcol{width:7.75in;}
.panel{position:absolute;left:8.75in;top:1.6in;width:3.95in;background:#1b1b1f;border-top:4px solid #F26A21;border-radius:6px;padding:16px 20px;}
.panel .ph{font-size:12px;font-weight:800;letter-spacing:2px;text-transform:uppercase;color:#F26A21;margin-bottom:8px;}
.panel div:not(.ph){font-size:14.5px;color:#dcdce0;padding:8px 0;border-bottom:1px solid #2e2e34;line-height:1.3;}
.panel div:last-child{border-bottom:none;}
.cards{display:flex;gap:0.3in;}
.card{flex:1;background:#1b1b1f;border-top:4px solid #F26A21;border-radius:6px;padding:20px 22px;min-height:4.2in;}
.card .ck{font-size:13px;font-weight:800;letter-spacing:2px;color:#F26A21;}
.card .ct{font-size:20px;font-weight:800;color:#fff;margin:8px 0 14px;line-height:1.2;}
.card p{font-size:15.5px;color:#d2d2d7;line-height:1.35;margin-bottom:11px;padding-left:18px;position:relative;}
.card p:before{content:"\\25AA";position:absolute;left:0;color:#F26A21;}
.bnote{position:absolute;bottom:0.55in;left:0.64in;right:0.64in;font-size:11.5px;font-style:italic;color:#9a9aa1;line-height:1.35;}
.draft{position:absolute;bottom:0.25in;left:0.64in;font-size:10px;color:#77777e;}
.g{position:absolute;border-radius:6px;padding:16px 20px;}
.g b{display:block;font-size:13px;letter-spacing:2px;color:#F26A21;margin-bottom:10px;}
.g p{font-size:15.5px;color:#e2e2e6;margin-bottom:10px;padding-left:18px;position:relative;line-height:1.3;}
.g p:before{content:"\\25AA";position:absolute;left:0;color:#F26A21;}
.g.side{background:#1b1b1f;border-top:4px solid #F26A21;}
.g.mid{background:#141417;border:2px solid #F26A21;} .g.mid p{font-size:14.5px;}
.g.bar{background:#F26A21;} .g.bar b{color:#141416;margin-bottom:6px;} .g.bar span{font-size:14.5px;color:#141416;line-height:1.35;}
.garrow{position:absolute;}
table.pt{width:100%;border-collapse:separate;border-spacing:5px;margin-left:-5px;}
.pt th{font-size:12.5px;letter-spacing:1.2px;color:#F26A21;text-align:left;padding:6px 10px;border-bottom:2px solid #F26A21;}
.pt td{background:#1b1b1f;font-size:14px;color:#dcdce0;padding:12px 12px;vertical-align:top;line-height:1.3;height:0.95in;width:28%;}
.pt td.rh{background:#26262c;color:#fff;font-weight:800;font-size:14.5px;width:16%;border-left:3px solid #F26A21;}
.qcols{display:flex;gap:0.4in;}
.qcols>div{flex:1;}
.q{display:flex;gap:14px;background:#1b1b1f;border-radius:6px;padding:11px 16px;margin-bottom:11px;min-height:0.78in;align-items:center;}
.q em{font-style:normal;font-size:22px;font-weight:800;color:#F26A21;width:30px;flex:none;}
.q span{font-size:15.5px;color:#ececef;line-height:1.3;}
.hero{background-color:#000;background-size:cover;background-position:center;color:#fff;padding:0.7in;}
.hlogo{position:absolute;top:0.6in;left:0.7in;height:0.6in;}
.cob{position:absolute;top:0.78in;right:0.7in;font-size:14px;letter-spacing:1.5px;color:#d2d2d7;font-weight:600;}
.htitle{position:absolute;top:2.0in;left:0.75in;font-size:46px;font-weight:800;letter-spacing:-0.5px;line-height:1.12;}
.hsub{position:absolute;top:4.3in;left:0.78in;right:0.7in;}
.hsub .l1{font-size:21px;font-weight:700;color:#F26A21;margin-bottom:14px;}
.hsub .l2{font-size:18px;color:#e2e2e6;margin-bottom:8px;}
.hsub .l3{font-size:15px;color:#a6a6ad;}
.o{color:#F26A21;}
.hfoot{position:absolute;bottom:0.5in;left:0.78in;font-size:11px;color:#F26A21;}
.ctitle{position:absolute;top:1.55in;left:0.75in;font-size:40px;font-weight:800;color:#F26A21;}
.crule{position:absolute;top:2.3in;left:0.77in;width:1.15in;height:4px;background:#F26A21;}
ul.hul{position:absolute;top:2.55in;left:0.77in;width:6.9in;list-style:none;}
.hul li{position:relative;padding-left:24px;margin-bottom:10px;font-size:16px;color:#e6e6ea;line-height:1.3;}
.hul li:before{content:"\\25AA";position:absolute;left:0;top:1px;color:#F26A21;font-size:14px;}
.contact{position:absolute;top:1.6in;left:8.3in;width:4.35in;background:rgba(20,20,23,.9);border-top:4px solid #F26A21;border-radius:6px;padding:14px 18px;}
.contact .ph{font-size:12px;font-weight:800;letter-spacing:2px;text-transform:uppercase;color:#F26A21;margin-bottom:8px;}
.contact div:not(.ph){font-size:14px;color:#d2d2d7;line-height:1.5;} .contact .cb{color:#fff;font-weight:700;font-size:16px !important;}
.contact .ty{margin-top:10px;color:#F26A21 !important;font-weight:700;}
.sources{position:absolute;bottom:0.62in;left:0.77in;right:0.6in;font-size:8.5px;color:#a6a6ad;line-height:1.4;}
.sources b{display:block;color:#d2d2d7;letter-spacing:1.5px;text-transform:uppercase;font-size:9px;margin-bottom:2px;}
.hfoot2{position:absolute;bottom:0.36in;left:0.78in;font-size:10px;color:#F26A21;}
'''
HTML = f"<!doctype html><html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{''.join(parts)}</body></html>"
oh = os.path.join(OUTD, STEM + ".html"); op = os.path.join(OUTD, STEM + ".pdf")
open(oh, "w").write(HTML)
CH = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
subprocess.run([CH, "--headless=new", "--disable-gpu", "--no-sandbox", f"--print-to-pdf={op}", "--no-pdf-header-footer", "file://" + oh],
               check=True, capture_output=True, timeout=120)
print("PDF:", op, os.path.getsize(op))

# ---------------------------------------------------------------- PPTX (editable)
BG = RGBColor(0x0d, 0x0d, 0x0f); ORANGE = RGBColor(0xF2, 0x6A, 0x21); INK = RGBColor(0x14, 0x14, 0x16)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); GREY = RGBColor(0xD2, 0xD2, 0xD7); MUTE = RGBColor(0x9a, 0x9a, 0xa1)
G1 = RGBColor(0x1f, 0x1f, 0x24); G2 = RGBColor(0x26, 0x26, 0x2c); G0 = RGBColor(0x14, 0x14, 0x17); PANEL = RGBColor(0x1b, 0x1b, 0x1f)

def cover(src, dst, scrim_fn, W=2000, Hh=1125):
    im = Image.open(src).convert("RGB"); iw, ih = im.size; sc = max(W / iw, Hh / ih)
    im = im.resize((int(iw * sc), int(ih * sc))); iw, ih = im.size
    im = im.crop(((iw - W) // 2, (ih - Hh) // 2, (iw - W) // 2 + W, (ih - Hh) // 2 + Hh))
    m = Image.new("L", (1, Hh))
    for y in range(Hh): m.putpixel((0, y), int(scrim_fn(y / (Hh - 1)) * 255))
    Image.composite(Image.new("RGB", (W, Hh), (0, 0, 0)), im, m.resize((W, Hh))).save(dst)
TBG_HERO = os.path.join(OUTD, "_honda_bg_hero.png"); TBG_CLOSE = os.path.join(OUTD, "_honda_bg_close.png"); TBAND = os.path.join(OUTD, "_honda_band.png")
cover(HERO, TBG_HERO, lambda t: 0.55 - 0.2 * t / 0.45 if t < 0.45 else (0.35 + 0.57 * min((t - 0.45) / 0.25, 1)))
cover(PLANET_A, TBG_CLOSE, lambda t: 0.88 - 0.18 * min(t / 0.46, 1) if t < 0.46 else 0.7 + 0.24 * ((t - 0.46) / 0.54))
bim = Image.open(PLANET_B).convert("RGB"); bw, bh = bim.size; bim = bim.crop((0, int(bh * 0.08), bw, int(bh * 0.08) + int(bw * 0.9 / 13.333)))
Image.blend(Image.new("RGB", bim.size, (13, 13, 15)), bim, 0.55).save(TBAND)

p = Presentation(); p.slide_width = I(13.333); p.slide_height = I(7.5); BL = p.slide_layouts[6]
def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5):
    sh = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh
def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    return tf
def para(tf, runs, size, color, bold=False, first=False, sp=1.0, before=0, align=PP_ALIGN.LEFT):
    pr = tf.paragraphs[0] if first else tf.add_paragraph()
    pr.alignment = align; pr.line_spacing = sp; pr.space_before = Pt(before); pr.space_after = Pt(0)
    if isinstance(runs, str): runs = [(runs, color, bold)]
    for txt, c, b in runs:
        r = pr.add_run(); r.text = txt; f = r.font; f.size = Pt(size); f.bold = b; f.color.rgb = c; f.name = "Arial"
def draft(s, color=MUTE):
    t = tb(s, 0.64, 7.12, 12, 0.25); para(t, DRAFT, 9, color, first=True)
def chrome(s, title, band=False):
    rect(s, 0, 0, 13.333, 7.5, BG)
    if band: s.shapes.add_picture(TBAND, I(0), I(6.6), I(13.333), I(0.9))
    s.shapes.add_picture(LOGO, I(11.7), I(0.45), height=I(0.34))
    t = tb(s, 0.64, 0.48, 10.6, 0.7); para(t, title, 25, WHITE, True, first=True)
    rect(s, 0.64, 1.18, 1.15, 0.055, ORANGE); draft(s)
def bullets(s, items, x=0.64, y=1.55, w=12.05, h=5.0, size=15, gap=8, col=WHITE):
    tf = tb(s, x, y, w, h); first = True
    for txt in items:
        para(tf, [("▪  ", ORANGE, True), (txt, col, False)], size, col, first=first, sp=1.06, before=0 if first else gap); first = False
def panel(s, title, items, x=8.75, y=1.55, w=3.95, h=4.6):
    rect(s, x, y, w, 0.05, ORANGE); rect(s, x, y + 0.05, w, h, PANEL)
    tf = tb(s, x + 0.2, y + 0.22, w - 0.4, h - 0.3); para(tf, title.upper(), 10, ORANGE, True, first=True)
    for it in items: para(tf, it, 12.5, GREY, before=10, sp=1.05)
def note(s, sd):
    if sd.get("note"): t = tb(s, 0.64, 6.45, 12.05, 0.5); para(t, sd["note"], 10, MUTE, first=True)
def herobg(s, img):
    rect(s, 0, 0, 13.333, 7.5, BG); s.shapes.add_picture(img, I(0), I(0), I(13.333), I(7.5))
    s.shapes.add_picture(LOGO, I(0.7), I(0.6), height=I(0.6))
    t = tb(s, 7.2, 0.78, 5.43, 0.3); para(t, COBRAND, 12, GREY, True, first=True, align=PP_ALIGN.RIGHT)

for sd in S:
    k = sd["kind"]; s = p.slides.add_slide(BL)
    if k == "hero":
        herobg(s, TBG_HERO)
        t = tb(s, 0.75, 1.95, 12, 2.0); para(t, sd["title"][0], 38, WHITE, True, first=True); para(t, sd["title"][1], 38, ORANGE, True)
        t = tb(s, 0.78, 4.3, 11.8, 1.6)
        para(t, sd["l1"], 17, ORANGE, True, first=True); para(t, sd["l2"], 15, GREY, before=10); para(t, sd["l3"], 12.5, MUTE, before=6)
        t = tb(s, 0.78, 6.85, 12, 0.3); para(t, DRAFT, 10, ORANGE, first=True)
    elif k == "split":
        chrome(s, sd["title"], sd.get("band")); bullets(s, sd["items"], w=7.75, size=14.5)
        panel(s, sd["panel_title"], sd["panel"]); note(s, sd)
    elif k == "cards":
        chrome(s, sd["title"])
        for i, (a, b, c) in enumerate(sd["cards"]):
            x = 0.64 + i * 4.12
            rect(s, x, 1.55, 3.82, 0.05, ORANGE); rect(s, x, 1.6, 3.82, 4.6, PANEL)
            tf = tb(s, x + 0.22, 1.8, 3.38, 4.3)
            para(tf, a, 11, ORANGE, True, first=True); para(tf, b, 17, WHITE, True, before=6, sp=1.05)
            for x2 in c: para(tf, [("▪  ", ORANGE, True), (x2, GREY, False)], 13.5, GREY, before=10, sp=1.08)
        note(s, sd)
    elif k == "model":
        chrome(s, sd["title"])
        for g in MODEL:
            if g[0] in ("box", "mid"):
                _, x, y, w, h, t_, its = g[:7]
                if g[0] == "box": rect(s, x, y, w, 0.05, ORANGE); rect(s, x, y + 0.05, w, h - 0.05, PANEL)
                else: rect(s, x, y, w, h, G0, ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE, lw=2).adjustments[0] = 0.04
                tf = tb(s, x + 0.2, y + 0.22, w - 0.4, h - 0.3); para(tf, t_, 10.5, ORANGE, True, first=True)
                for it in its: para(tf, [("▪  ", ORANGE, True), (it, GREY, False)], 13 if g[0] == "box" else 12, GREY, before=9, sp=1.05)
            else:
                _, x, y, w, h, t_, sub = g
                rect(s, x, y, w, h, ORANGE, None, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.06
                tf = tb(s, x + 0.2, y + 0.15, w - 0.4, h - 0.25, MSO_ANCHOR.MIDDLE); para(tf, t_, 10.5, INK, True, first=True)
                para(tf, sub, 12.5, INK, before=4, sp=1.05)
        rect(s, 4.97, 2.95, 0.22, 0.5, ORANGE, None, MSO_SHAPE.RIGHT_ARROW)
        rect(s, 8.14, 2.95, 0.22, 0.5, ORANGE, None, MSO_SHAPE.LEFT_ARROW)
    elif k == "path":
        chrome(s, sd["title"])
        cw = [1.95, 3.32, 3.32, 3.32]; xs = [0.64]
        for w_ in cw[:-1]: xs.append(xs[-1] + w_ + 0.05)
        for j, c in enumerate(sd["cols"]):
            if c: t = tb(s, xs[j] + 0.1, 1.5, cw[j] - 0.2, 0.3); para(t, c.upper(), 10, ORANGE, True, first=True)
        rect(s, 0.64, 1.85, sum(cw) + 0.15, 0.03, ORANGE)
        for i, r in enumerate(sd["rows"]):
            y = 1.98 + i * 1.08
            for j, c in enumerate(r):
                rect(s, xs[j], y, cw[j], 1.0, G2 if j == 0 else PANEL)
                if j == 0: rect(s, xs[j], y, 0.04, 1.0, ORANGE)
                t = tb(s, xs[j] + 0.14, y + 0.1, cw[j] - 0.26, 0.85, MSO_ANCHOR.MIDDLE if j == 0 else MSO_ANCHOR.TOP)
                para(t, c, 12.5 if j == 0 else 11.5, WHITE if j == 0 else GREY, j == 0, first=True, sp=1.05)
        note(s, sd)
    elif k == "questions":
        chrome(s, sd["title"]); qs = sd["qs"]; half = (len(qs) + 1) // 2
        for i, q in enumerate(qs):
            c, r = i // half, i % half; x = 0.64 + c * 6.2; y = 1.55 + r * 0.92
            rect(s, x, y, 5.85, 0.8, PANEL, None, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
            t = tb(s, x + 0.15, y, 0.5, 0.8, MSO_ANCHOR.MIDDLE); para(t, str(i + 1), 18, ORANGE, True, first=True)
            t = tb(s, x + 0.65, y + 0.06, 5.05, 0.68, MSO_ANCHOR.MIDDLE); para(t, q, 13, WHITE, first=True, sp=1.05)
        t = tb(s, 0.64, 6.25, 12.05, 0.5)
        para(t, "We would be grateful for whatever Honda is comfortable sharing at a public, non-technical level now; detail can follow under NDA and any required approvals.", 10, MUTE, first=True)
    elif k == "close":
        herobg(s, TBG_CLOSE)
        t = tb(s, 0.75, 1.5, 7, 0.8); para(t, sd["title"], 34, ORANGE, True, first=True)
        rect(s, 0.77, 2.3, 1.15, 0.055, ORANGE)
        bullets(s, sd["items"], x=0.77, y=2.5, w=7.1, h=2.9, size=14, gap=6, col=GREY)
        rect(s, 8.3, 1.6, 4.35, 0.05, ORANGE); rect(s, 8.3, 1.65, 4.35, 2.75, G0)
        tf = tb(s, 8.5, 1.8, 4.0, 2.55); para(tf, "CONTACT", 10, ORANGE, True, first=True)
        for t_, b_ in CONTACT: para(tf, t_, 14 if b_ else 12, WHITE if b_ else GREY, b_, before=6 if b_ else 2)
        para(tf, "Thank you — ありがとうございます", 12, ORANGE, True, before=10)
        t = tb(s, 0.77, 5.3, 11.9, 1.5); para(t, "SOURCES", 8, GREY, True, first=True)
        for x in SOURCES: para(t, x, 7, MUTE, before=1)
        t = tb(s, 0.78, 6.98, 12, 0.3); para(t, DRAFT, 9.5, ORANGE, first=True)

out = os.path.join(OUTD, STEM + ".pptx"); p.save(out); print("PPTX:", out, os.path.getsize(out), "slides:", len(p.slides))
