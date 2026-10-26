#!/usr/bin/env python3
"""Educational Payload Sponsorship flyer (Jon 2026-10-26).

One Letter page, built from one layout spec into two outputs:
  - editable .pptx (Lizi finalizes in Canva, see feedback_human_designed_visuals)
  - .html twin for a headless-Chrome PNG preview (soffice hangs on this box)
Brand: black/grey base, orange #F26A21 accent, official white logo.
"""
import html, os, subprocess, sys
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

WS = "/Users/exomac/.openclaw/workspace"
OUTDIR = os.path.join(WS, "drafts", "edu-sponsorship")
STAMP = sys.argv[1] if len(sys.argv) > 1 else "2026-10-26"
NAME = f"Exos_Educational_Payload_Sponsorship_Flyer_{STAMP}"
LOGO = os.path.join(WS, "drafts", "_exos_white_official.png")
HEROBG = os.path.join(WS, "drafts", "_exos_hero_bg.png")

BLACK = "141416"; GREY = "26262C"; ORANGE = "F26A21"; WHITE = "FFFFFF"
LIGHT = "EFEFF1"; TEXT = "1C1C22"; MUTE = "54545D"; PALE = "CFCFD4"; SOFT = "B7B7BD"

E = []  # layout spec: (kind, x, y, w, h, props)
def rect(x, y, w, h, fill, rounded=False): E.append(("rect", x, y, w, h, dict(fill=fill, rounded=rounded)))
def img(path, x, y, w=None, h=None): E.append(("img", x, y, w, h, dict(path=path)))
def txt(x, y, w, h, lines, align="l", anchor="t", sp=1.05):
    E.append(("txt", x, y, w, h, dict(lines=lines, align=align, anchor=anchor, sp=sp)))
def R(t, sz, b=False, c=TEXT, i=False): return (t, sz, b, c, i)

# HERO
img(HEROBG, 0, 0, 8.5, 2.95)
img(LOGO, 0.5, 0.45, None, 0.5)
txt(4.6, 0.55, 3.4, 0.3, [[R("EDUCATIONAL PAYLOAD SPONSORSHIP", 9, True, SOFT)]], align="r")
txt(0.5, 1.1, 7.5, 1.2, [[R("Their experiment is built.", 31, True, WHITE)],
                         [R("Help it ", 31, True, WHITE), R("fly to space.", 31, True, ORANGE)]], sp=1.02)
txt(0.5, 2.38, 7.5, 0.4, [[R("Student teams have flight-ready science payloads waiting for a ride. Sponsor one and put your name on the mission.", 12, False, PALE)]])

# DEADLINE BAND
rect(0, 2.95, 8.5, 0.52, ORANGE)
txt(0.5, 2.95, 7.5, 0.52, [[R("FUND BY NOVEMBER 30, 2026   ", 11, True, BLACK),
                            R("so these payloads can be onboarded for our early-2027 BLK3 flight.", 11.5, False, BLACK)]], anchor="m")

# WHY IT MATTERS
rect(0.5, 3.68, 0.05, 0.78, ORANGE)
txt(0.7, 3.66, 7.3, 0.85, [[R("Real flight hardware, real data, real careers. ", 11, True, TEXT),
    R("Students design, build and integrate an experiment that rides to the edge of space (~100 km), "
      "gets 2 to 3 minutes of microgravity and comes back to them with its data. A sponsorship pays the flight "
      "for a team that has the payload but not the budget.", 10, False, MUTE)]], sp=1.12)

# TIER LADDER (countdown)
txt(0.5, 4.62, 7.5, 0.3, [[R("PICK YOUR ALTITUDE", 10, True, GREY), R("   five levels, $8,000 to $50,000", 9, False, MUTE, True)]])
tiers = [
    ("T-4", "LAUNCH PAD", "$8,000", "Funds 1 student payload (1U, up to 1 kg).", "Your name on the payload · certificate · flight-video thank-you", False),
    ("T-3", "MAX-Q", "$15,000", "Funds 2 student payloads.", "Everything above + named in the mission press release and on social", False),
    ("T-2", "MECO", "$25,000", "Funds a 3-payload team package.", "Everything above + launch-day recognition + a team briefing for your staff", False),
    ("T-1", "APOGEE", "$35,000", "Funds 4 student payloads.", "Everything above + YOUR LOGO ON THE ROCKET", True),
    ("T-0", "KÁRMÁN LINE", "$50,000", "Mission Title Sponsor: funds 6 payloads.", "Premium logo placement on the rocket · \"[Your Company] Student Flight\" naming · launch viewing invitation", True),
]
y0, rh = 4.95, 0.62
for k, (t, nm, amt, funds, perks, logo) in enumerate(tiers):
    y = y0 + k * (rh + 0.06)
    rect(0.5, y, 7.5, rh, GREY if logo else LIGHT, rounded=True)
    rect(0.5, y, 0.62, rh, ORANGE if logo else "D9D9DD", rounded=True)
    txt(0.5, y, 0.62, rh, [[R(t, 10.5, True, BLACK)]], align="c", anchor="m")
    c1, c2 = (WHITE, PALE) if logo else (TEXT, MUTE)
    txt(1.25, y + 0.08, 1.55, rh - 0.1, [[R(nm, 10.5, True, ORANGE if logo else GREY)], [R(amt, 15, True, c1)]], sp=1.0)
    txt(2.9, y + 0.09, 5.0, rh - 0.12, [[R(funds, 10, True, c1)], [R(perks, 8.8, False, c2)]], sp=1.05)

# HOW IT WORKS
yh = y0 + 5 * (rh + 0.06) + 0.08
txt(0.5, yh, 7.5, 0.28, [[R("HOW IT WORKS", 10, True, GREY)]])
steps = [("1", "Commit by Nov 30", "Choose a level. We invoice your company, or your foundation grants to the student team's school."),
         ("2", "We match a team", "We pair you with an unfunded student payload, by region, school or science area if you like."),
         ("3", "Watch it fly", "You get the flight video, the team's results and a story your people will want to share.")]
for k, (n, hd, ds) in enumerate(steps):
    x = 0.5 + k * 2.55
    txt(x, yh + 0.3, 2.4, 0.9, [[R(n + "  ", 13, True, ORANGE), R(hd, 10.5, True, TEXT)], [R(ds, 8.8, False, MUTE)]], sp=1.08)

rect(0.5, yh + 1.22, 7.5, 0.4, LIGHT, rounded=True)
txt(0.7, yh + 1.22, 7.1, 0.4, [[R("Every level above $25,000 puts your logo on the rocket. ", 9.5, True, TEXT),
    R("Bay space is limited; levels fill as commitments arrive.", 9.5, False, MUTE)]], anchor="m")

# CREDIBILITY + FOOTER
rect(0, 10.12, 8.5, 0.88, BLACK)
txt(0.5, 10.2, 4.6, 0.72, [[R("Flying payloads since 2018.", 11, True, WHITE)],
    [R("Exos has flown commercial payloads on SARGE since 2018. BLK3, our next reusable vehicle, is preparing for first flight in early 2027.", 8.3, False, PALE)]],
    anchor="m", sp=1.08)
txt(5.25, 10.2, 2.75, 0.72, [[R("Lizi Quinn · lquinn@exosaero.com", 9, True, WHITE)],
    [R("John Quinn, CEO · jquinn@exosaero.com", 9, False, PALE)],
    [R("exosaero.com  ·  @exosaero", 9, False, ORANGE)]], align="r", anchor="m", sp=1.12)


def rgb(h): return RGBColor.from_string(h)

def build_pptx(path):
    p = Presentation(); p.slide_width = I(8.5); p.slide_height = I(11)
    s = p.slides.add_slide(p.slide_layouts[6])
    for kind, x, y, w, h, pr in E:
        if kind == "rect":
            shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if pr["rounded"] else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
            if pr["rounded"]: shp.adjustments[0] = 0.12
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(pr["fill"]); shp.line.fill.background(); shp.shadow.inherit = False
        elif kind == "img":
            s.shapes.add_picture(pr["path"], I(x), I(y), I(w) if w else None, I(h) if h else None)
        else:
            tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame; tf.word_wrap = True
            tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[pr["anchor"]]
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            for n, line in enumerate(pr["lines"]):
                para = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                para.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[pr["align"]]
                para.line_spacing = pr["sp"]
                for t, sz, b, c, it in line:
                    r = para.add_run(); r.text = t; f = r.font
                    f.size = Pt(sz); f.bold = b; f.italic = it; f.color.rgb = rgb(c); f.name = "Arial"
    p.save(path)

def build_html(path):
    out = ['<html><head><meta charset="utf-8"><style>@page{size:8.5in 11in;margin:0}body{margin:0}'
           '.pg{position:relative;width:8.5in;height:11in;overflow:hidden;font-family:Arial,Helvetica,sans-serif;background:#fff}'
           '.a{position:absolute;box-sizing:border-box}</style></head><body><div class="pg">']
    for kind, x, y, w, h, pr in E:
        pos = f"left:{x}in;top:{y}in;" + (f"width:{w}in;" if w else "") + (f"height:{h}in;" if h else "")
        if kind == "rect":
            out.append(f'<div class="a" style="{pos}background:#{pr["fill"]};border-radius:{"0.08in" if pr["rounded"] else "0"}"></div>')
        elif kind == "img":
            out.append(f'<img class="a" style="{pos}" src="file://{pr["path"]}">')
        else:
            jc = {"t": "flex-start", "m": "center"}[pr["anchor"]]
            ta = {"l": "left", "r": "right", "c": "center"}[pr["align"]]
            paras = "".join("<div>" + "".join(
                f'<span style="font-size:{sz}pt;font-weight:{700 if b else 400};font-style:{"italic" if it else "normal"};color:#{c}">{html.escape(t)}</span>'
                for t, sz, b, c, it in line) + "</div>" for line in pr["lines"])
            out.append(f'<div class="a" style="{pos}display:flex;flex-direction:column;justify-content:{jc};text-align:{ta};line-height:{pr["sp"] * 1.17}">{paras}</div>')
    out.append("</div></body></html>")
    open(path, "w").write("".join(out))

if __name__ == "__main__":
    os.makedirs(OUTDIR, exist_ok=True)
    pptx = os.path.join(OUTDIR, NAME + ".pptx"); htm = os.path.join(OUTDIR, "_preview.html"); png = os.path.join(OUTDIR, "_preview.png")
    build_pptx(pptx); build_html(htm)
    subprocess.run(["/Applications/Google Chrome.app/Contents/MacOS/Google Chrome", "--headless=new", "--disable-gpu",
                    "--allow-file-access-from-files", f"--screenshot={png}", "--window-size=816,1056",
                    "--hide-scrollbars", "file://" + htm], capture_output=True, timeout=90)
    print("SAVED", pptx, os.path.getsize(pptx)); print("PREVIEW", png)
