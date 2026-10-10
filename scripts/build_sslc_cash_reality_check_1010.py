#!/usr/bin/env python3
"""SSLC year-end cash reality check (John 10/10).

Claim tested: Jim Wertz says SSLC will hold $850K cash by year-end, with Markus Rufer
describing a rate of one tank every 3 weeks at $250K each.
Output:
  drafts/SSLC_YearEnd_Cash_Reality_Check_2026-10-10.xlsx  (formula-driven, with a native line chart)
  drafts/SSLC_YearEnd_Cash_Reality_Check_2026-10-10.png  (for chat / slides)
All inputs live on the Inputs sheet and are labeled. Management assumptions are marked.
"""
import os
import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

WS = "/Users/exomac/.openclaw/workspace"
XLSX = f"{WS}/drafts/SSLC_YearEnd_Cash_Reality_Check_2026-10-10.xlsx"
PNG = f"{WS}/drafts/SSLC_YearEnd_Cash_Reality_Check_2026-10-10.png"

BLUE, ORANGE, GREY, INK = "#2B6CB0", "#F26A21", "#8A8A8A", "#1F2328"

# (label, value, source/basis)
INPUTS = [
    ("Starting cash (Oct 2, 2026)", 198000, "SSLC QuickBooks pull 10/2/2026"),
    ("Year-end target cash (Jim Wertz)", 850000, "Jim Wertz statement"),
    ("Weeks Oct 2 → Dec 31", 13, "Calendar"),
    ("Weeks per tank (Markus Rufer)", 3, "Markus Rufer statement"),
    ("Price per tank ($)", 250000, "Markus Rufer statement"),
    ("Prepaid backlog already collected ($)", 984000, "SSLC QB 10/2: $864K billed ahead + $120K deposits. DISPUTED by John: tank start/billing not captured correctly"),
    ("Direct build cost (% of tank price)", 0.52, "Management assumption (actuals-anchored model COGS)"),
    ("Annual fixed overhead ($)", 1340000, "Derived: FY26 rev $1.925M × 48% GM + $416K op loss; mgmt assumption"),
]


def build_xlsx():
    wb = openpyxl.Workbook()
    inp = wb.active
    inp.title = "Inputs"
    inp["A1"] = "SSLC year-end cash reality check: inputs"
    inp["A1"].font = Font(bold=True, size=13)
    inp.append(["Input", "Value", "Source / basis"])
    for c in inp[2]:
        c.font = Font(bold=True)
    for row in INPUTS:
        inp.append(list(row))
    for col, w in zip("ABC", (40, 14, 70)):
        inp.column_dimensions[col].width = w
    # B3..B10 = inputs in order
    for r in (3, 4, 7, 8, 10):
        inp[f"B{r}"].number_format = '$#,##0'
    inp["B9"].number_format = '0%'

    ws = wb.create_sheet("Weekly Cash")
    ws.append(["Week", "Tank delivered?", "Tank revenue ($)",
               "Best case: cash balance", "Per SSLC books (disputed): cash balance", "Wertz claim path", "Target $850K"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True)
    ws.append([0, "", 0, "=Inputs!B3", "=Inputs!B3", "=Inputs!B3", "=Inputs!B4"])
    for wk in range(1, 14):
        r = wk + 2
        ws.append([
            wk,
            f'=IF(MOD(A{r},Inputs!$B$6)=0,"Yes","")',
            f'=IF(B{r}="Yes",Inputs!$B$7,0)',
            # Best case: every delivered tank is NEW cash collected on delivery
            f"=D{r-1}+C{r}-Inputs!$B$7*Inputs!$B$9/Inputs!$B$6-Inputs!$B$10/52",
            # Realistic: deliveries first burn down the prepaid backlog (cash already received)
            f"=E{r-1}+MAX(0,SUM($C$3:C{r})-Inputs!$B$8)-MAX(0,SUM($C$3:C{r-1})-Inputs!$B$8)"
            f"-Inputs!$B$7*Inputs!$B$9/Inputs!$B$6-Inputs!$B$10/52",
            f"=Inputs!$B$3+(Inputs!$B$4-Inputs!$B$3)*A{r}/Inputs!$B$5",
            "=Inputs!$B$4",
        ])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=7):
        for c in row:
            c.number_format = '$#,##0;[Red]-$#,##0'
    for col, w in zip("ABCDEFG", (7, 10, 14, 18, 18, 16, 14)):
        ws.column_dimensions[col].width = w

    s = wb.create_sheet("Summary")
    s["A1"] = "Year-end (Dec 31) result"
    s["A1"].font = Font(bold=True, size=13)
    rows = [
        ("Tanks delivered by Dec 31 at 1 per 3 weeks", "=COUNTIF('Weekly Cash'!B3:B15,\"Yes\")"),
        ("Tank revenue by Dec 31", "=SUM('Weekly Cash'!C3:C15)"),
        ("Best case year-end cash (all tanks = new cash)", "='Weekly Cash'!D15"),
        ("Per SSLC books, if billed-ahead is accurate (DISPUTED)", "='Weekly Cash'!E15"),
        ("Wertz claim", "=Inputs!B4"),
        ("Shortfall vs claim: best case", "=Inputs!B4-B5"),
        ("Shortfall vs claim: per SSLC books (disputed)", "=Inputs!B4-B6"),
    ]
    for i, (k, f) in enumerate(rows, start=3):
        s[f"A{i}"], s[f"B{i}"] = k, f
        s[f"B{i}"].number_format = '$#,##0;[Red]-$#,##0'
    s["B3"].number_format = '0'
    s.column_dimensions["A"].width = 52
    s.column_dimensions["B"].width = 16
    s["A11"] = ("Reading: even if every tank is new cash, SSLC can't reach $850K by year-end from operations. "
                "The prepaid backlog means the next ~4 tanks were already paid for. Reaching $850K requires new "
                "customer deposits or financing, and each of those adds a delivery or repayment obligation.")
    s["A11"].alignment = Alignment(wrap_text=True)
    s.row_dimensions[11].height = 60
    s["A12"] = "Management assumptions are marked on Inputs. Projection, not a statement of SSLC's actual results."
    s["A12"].font = Font(italic=True, color="666666")

    ch = LineChart()
    ch.title = "SSLC cash path to Dec 31 vs the $850K claim"
    ch.y_axis.title = "Cash ($)"
    ch.x_axis.title = "Week from Oct 2"
    data = Reference(ws, min_col=4, max_col=7, min_row=1, max_row=15)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=15))
    ch.height, ch.width = 9, 18
    ws.add_chart(ch, "I2")

    tmp = XLSX + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, XLSX)


def series():
    v = {k: val for k, val, _ in INPUTS}
    start, target = v["Starting cash (Oct 2, 2026)"], v["Year-end target cash (Jim Wertz)"]
    weeks, per, price = v["Weeks Oct 2 → Dec 31"], v["Weeks per tank (Markus Rufer)"], v["Price per tank ($)"]
    prepaid, cogs, ovh = v["Prepaid backlog already collected ($)"], v["Direct build cost (% of tank price)"], v["Annual fixed overhead ($)"]
    burn = price * cogs / per + ovh / 52
    best, real, claim = [start], [start], [start]
    cum = 0
    for wk in range(1, weeks + 1):
        rev = price if wk % per == 0 else 0
        new_cash = max(0, cum + rev - prepaid) - max(0, cum - prepaid)
        cum += rev
        best.append(best[-1] + rev - burn)
        real.append(real[-1] + new_cash - burn)
        claim.append(start + (target - start) * wk / weeks)
    return list(range(weeks + 1)), best, real, claim, target


def build_png():
    x, best, real, claim, target = series()
    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=150)
    k = lambda a: [t / 1000 for t in a]
    ax.axhline(target / 1000, color=GREY, lw=1.5, ls="--")
    ax.plot(x, k(claim), color=GREY, lw=1.2, ls=":")
    ax.plot(x, k(best), color=BLUE, lw=2, marker="o", ms=4)
    ax.plot(x, k(real), color=ORANGE, lw=2, marker="o", ms=4)
    ax.axhline(0, color="#CCCCCC", lw=1)
    lab = dict(fontsize=9, va="center", color=INK)
    ax.text(13.2, target / 1000, f"Claim: ${target/1000:,.0f}K", **lab)
    ax.text(13.2, best[-1] / 1000, f"Best case: ${best[-1]/1000:,.0f}K\n(every tank = new cash)", **lab)
    ax.text(13.2, real[-1] / 1000, f"Per SSLC books: -${-real[-1]/1000:,.0f}K\n(if $864K billed-ahead is right; disputed)", **lab)
    for wk in (3, 6, 9, 12):
        ax.axvline(wk, color="#EEEEEE", lw=1, zorder=0)
    ax.set_xlim(0, 17.5)
    ax.set_xticks(range(0, 14))
    ax.set_xlabel("Week from Oct 2, 2026 (tank every 3 weeks at $250K: weeks 3, 6, 9, 12)", color=INK)
    ax.set_ylabel("SSLC cash ($K)", color=INK)
    ax.set_title("SSLC can't reach $850K cash by Dec 31 from operations", loc="left",
                 fontsize=13, fontweight="bold", color=INK)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.grid(axis="y", color="#EEEEEE")
    fig.text(0.01, 0.01,
             "Inputs: $198K cash 10/2 (SSLC QB); $984K prepaid backlog per SSLC books (disputed); 52% build cost and $1.34M/yr overhead "
             "are management assumptions. Projection, not SSLC's reported results.",
             fontsize=7.5, color="#666666")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    tmp = PNG + ".tmp.png"
    fig.savefig(tmp, format="png")
    os.replace(tmp, PNG)
    return best[-1], real[-1]


if __name__ == "__main__":
    build_xlsx()
    b, r = build_png()
    print(f"best {b:,.0f}  realistic {r:,.0f}")
