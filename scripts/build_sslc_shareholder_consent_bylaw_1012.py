#!/usr/bin/env python3
"""DRAFT SSLC Action by Written Consent of Shareholders + bylaw amendment (John Quinn, 10/12/2026).

Share basis: SSLC register "5-SSLC_Post_Exos_Shares_acquisition_08-2024.xls" (sheet "Exos ownership 8-2024").
Proxy: Microcosm, Inc. -> John Quinn, dated Oct 9 2024, signed Oct 18 2024 (data-room/SSLC/).
Output: drafts/SSLC_Shareholder_Written_Consent_and_Bylaw_DRAFT_2026-10-12.docx (+ .pdf via soffice)
NOT SENT. NOT EXECUTED.
"""
import os
import subprocess
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

WS = "/Users/exomac/.openclaw/workspace"
OUT = f"{WS}/drafts/SSLC_Shareholder_Written_Consent_and_Bylaw_DRAFT_2026-10-12.docx"
RED = RGBColor(0xC0, 0x00, 0x00)
GREY = RGBColor(0x60, 0x60, 0x60)

TOTAL = 872525          # 714,972 pre-Exos + 157,553 Exos (08-2024 register)
MAJ = TOTAL // 2 + 1    # 436,263
SIGNERS = [
    ("Exos Aerospace Systems & Technologies, Inc.", 157553, "By: John M. Quinn, Cofounder & CEO"),
    ("Microcosm, Inc. (by proxy)", 280000,
     "By: John M. Quinn, proxyholder under the Microcosm, Inc. Proxy dated October 9, 2024 (financial matters, including accounting)"),
    ("Robert (Bob) Conger", 73000, "Certificates #3 (44,000) and #43 (29,000)"),
    ("Markus Rufer", 91000, "Certificates #31 (2,000) and #45 (89,000)"),
]

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1)
    s.top_margin = s.bottom_margin = Inches(0.9)


def P(text="", bold=False, size=None, color=None, italic=False, align=None, sa=6, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(sa)
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def H(text):
    return P(text, bold=True, size=12, sa=4)


C = WD_ALIGN_PARAGRAPH.CENTER

# ── Cover note (internal; remove before circulation) ─────────────────────────
P("DRAFT — NOT EXECUTED — NOT SENT  ·  John Quinn review copy", bold=True, size=9, color=RED, align=C)
H("Cover note for John (remove this page before circulating)")
for t in [
    "Purpose: a shareholder action outside the board that (1) adopts a bylaw putting related-party payments, "
    "netting and reclasses under disinterested-director control, (2) requires an independent reconciliation of the "
    "disputed balances and the share register, and (3) withholds ratification of the 4/30/2026 entries. "
    "It does NOT vote on the sale to Exos (Motion 1 stays with the board; John abstains).",
    f"Vote math (08-2024 register, {TOTAL:,} shares outstanding; majority = {MAJ:,}): Exos 157,553 (18.06%) + "
    "Microcosm by proxy 280,000 (32.09%) = 437,553 (50.15%), a majority by only 1,290 shares. Adding Conger (73,000) and "
    "Rufer (91,000) brings it to 601,553 (68.94%). Recommend NOT relying on the proxy alone: get Conger and Rufer to sign.",
    "Why the 18% matters here: on SSLC's 6/3/2026 ledger (Exos 132,972; 847,944 total) Exos + proxy = 48.7%, under a majority. "
    "The 08-2024 register is the basis for 18.06%; Resolution 3 makes the Corporation reconcile the ledger to it.",
    "Checked against SSLC Bylaws (adopted June 2000; data-room/SSLC/governance/SSLC_1a_Bylaws.pdf): Art. II §10 expressly "
    "permits shareholder action by written consent of a majority of outstanding shares and recognizes proxyholders; Art. IX §1 "
    "lets holders of a majority of outstanding shares adopt or amend bylaws. The new Article is numbered X (existing Articles I-IX).",
    "Checked against SSLC Articles of Incorporation (filed 1/19/1999, Entity No. 2131102) and the Certificate of Amendment "
    "(filed 3/22/1999): one class of shares, 6,000,000 authorized; no restriction on shareholder written consent; no board "
    "size stated (so Bylaws Art. III §2 governs: 5-9 directors). Statements of Information filed 1/27/2025 and 1/18/2026: "
    "5 directors, 0 vacancies (John Quinn, Markus Rufer, Robert Conger, Alice Wertz, James Wertz); Rufer CEO; Alice Wertz "
    "Secretary and CFO.",
    "Board route first: on related-party matters Jim and Alice Wertz are interested directors, so Quinn + Rufer + Conger are "
    "3 of 5 (a quorum and a majority). Any two directors can call a special board meeting (Bylaws Art. III §8), and the board "
    "can remove or replace officers, including the CFO/Secretary, with or without cause (Art. V §4). Use this consent if "
    "Rufer or Conger will not vote at the board level.",
    "Before use: (a) Bylaws Art. II §12 requires a proxy to be 'filed with the secretary': deliver the Microcosm proxy to "
    "the Secretary (Alice Wertz) with or before the consents; (b) deliver consents to the Secretary AND every director by "
    "trackable means; (c) record date for this consent = the day the first consent is signed (Bylaws Art. II §11(b)). "
    "Sign Exos's consent last, on the day all are delivered.",
    "Exos $60K resolved (Exos QuickBooks, 10/12/2026): the $790K 'Scorpius Stock' = $450K (Nov 2021) + $60K (12/21/2023) + "
    "$250K (Mar-Apr 2024) + $30K loan converted (7/24/2024). The 10/10/2024 $60K is a SEPARATE loan (SSLC repaid $28.8K; "
    "$31.2K still due). The 18.06% and the loan are both valid.",
    "Sequence: collect signatures, deliver all signed consents to the Secretary at the same time (consents can be revoked until "
    "the majority is filed, §603(c)), then send the Notice (Exhibit B) promptly to non-consenting holders (§603(b)(2)). Expect "
    "Microcosm to revoke the proxy once it is used; actions taken before revocation remain valid (Proxy §4).",
    "Not legal advice. Sources: Cal. Corp. Code §603 (written consent), §705 (proxies), §211 (shareholder bylaw amendment), "
    "§310 (interested-director transactions), §1602 (director inspection); california.public.law. SSLC Articles, "
    "amendment and Statements of Information: data-room/SSLC/governance/.",
]:
    P("•  " + t, size=10, sa=4)
doc.add_page_break()

# ── Consent ──────────────────────────────────────────────────────────────────
P("DRAFT — NOT EXECUTED", bold=True, size=9, color=RED, align=C)
P("SCORPIUS SPACE LAUNCH COMPANY", bold=True, size=14, align=C, sa=2)
P("(a California corporation)", italic=True, align=C, sa=8)
P("ACTION BY WRITTEN CONSENT OF SHAREHOLDERS WITHOUT A MEETING", bold=True, size=12, align=C, sa=2)
P("Pursuant to Section 603 of the California Corporations Code", italic=True, align=C, sa=12)

P("The undersigned shareholders of Scorpius Space Launch Company, a California corporation (the \"Corporation\"), "
  "holding shares having not less than the minimum number of votes that would be necessary to take the following "
  "actions at a meeting at which all shares entitled to vote were present and voted, hereby consent in writing to "
  "the adoption of the following resolutions, effective on the date this consent has been signed by holders of a "
  "majority of the outstanding shares and delivered to the Secretary of the Corporation (the \"Effective Date\"):")

H("Recitals")
for t in [
    "A. Exos Aerospace Systems & Technologies, Inc. (\"Exos\") invested $790,000 in the Corporation and holds 157,553 shares "
    "(18.06%) per the Corporation's register titled \"Post Exos Shares acquisition 08-2024\".",
    "B. By a Proxy dated October 9, 2024, Microcosm, Inc. (\"Microcosm\") granted John M. Quinn its voting rights on "
    "financial matters of the Corporation, including accounting, financing and loans, until the completion of the "
    "acquisition of the Corporation by Exos or written revocation. No revocation has been delivered.",
    "C. On or about April 30, 2026, entries were recorded in the Corporation's books that (i) reclassified $426,687 of "
    "accounts payable to \"Due to Microcosm\", (ii) added $97,302 of interest to a note payable to Microcosm Astronautics "
    "Books, and (iii) netted approximately $812,699 lent by the Corporation to Microcosm (together, the \"April 2026 "
    "Entries\"). The April 2026 Entries were not presented to or approved by the proxyholder or by disinterested directors.",
    "D. The Corporation's accounts receivable unbilled (account 1110, $942,177 at October 11, 2026) and customer "
    "deposits (account 2380, $120,000) are disputed, and the Corporation's share ledger dated June 3, 2026 does not match "
    "its issued certificates.",
    "E. Accurate books are a condition to any sale of the Corporation's assets and to any distribution to its shareholders.",
]:
    P(t, indent=0.25, sa=4)

H("Resolution 1: Bylaw Amendment")
P("RESOLVED, that pursuant to Article IX, Section 1 of the Bylaws and Section 211 of the California Corporations Code, the Bylaws of the Corporation are "
  "amended to add a new Article X titled \"Financial Integrity and Related-Party Transactions\" in the form attached as "
  "Exhibit A, effective on the Effective Date.")

H("Resolution 2: Independent Reconciliation")
P("RESOLVED, that the officers of the Corporation are directed to carry out the Independent Reconciliation required "
  "by Section 10.4 of Exhibit A, and to deliver its report to every director and to any shareholder who requests it.")

H("Resolution 3: Share Register")
P("RESOLVED, that the Secretary of the Corporation is directed to reconcile the share register to the certificates "
  "issued, including without limitation Certificate #3 (Robert Conger, 44,000 shares), which is omitted from the named "
  "holders on the June 3, 2026 ledger, and Exos's holding of 157,553 shares per the 08-2024 register, and to deliver the "
  "corrected register to all shareholders within 15 days of the Effective Date.")

H("Resolution 4: No Ratification; Reservation of Rights")
P("RESOLVED, that the shareholders do not ratify the April 2026 Entries or any payment, set-off, reclassification or "
  "interest accrual in favor of a Related Party (as defined in Exhibit A) made after October 9, 2024 without the "
  "approval required by Exhibit A, and each such item shall be treated as disputed until resolved by the Independent "
  "Reconciliation. All rights of the Corporation and its shareholders are reserved.")

H("Resolution 5: General")
P("RESOLVED, that the officers of the Corporation are authorized and directed to take all actions necessary to carry "
  "out these resolutions, and that the Secretary shall file this consent with the minutes of the proceedings of the "
  "shareholders and give the notice required by Section 603(b) of the California Corporations Code.")

P("This consent may be executed in counterparts and by electronic signature, each of which is an original and all of "
  "which together are one instrument.", italic=True, size=10, sa=10)

H("Shareholder Signatures")
t = doc.add_table(rows=1, cols=4)
t.style = "Table Grid"
for i, h in enumerate(["Shareholder", "Shares", "% of 872,525", "Signature / Date"]):
    t.rows[0].cells[i].text = h
    t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
for name, sh, note in SIGNERS:
    c = t.add_row().cells
    c[0].text = f"{name}\n{note}"
    c[1].text = f"{sh:,}"
    c[2].text = f"{sh / TOTAL:.2%}"
    c[3].text = "\n\n________________\nDate: ___________"
for _ in range(3):
    c = t.add_row().cells
    c[0].text = "Name:"
    c[3].text = "\n\n________________\nDate: ___________"
tot = sum(s for _, s, _ in SIGNERS)
c = t.add_row().cells
c[0].text = "Named signers, total"
c[1].text = f"{tot:,}"
c[2].text = f"{tot / TOTAL:.2%}"
c[3].text = f"Majority required: {MAJ:,}"
for row in t.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)
doc.add_page_break()

# ── Exhibit A: bylaw ─────────────────────────────────────────────────────────
P("DRAFT — NOT EXECUTED", bold=True, size=9, color=RED, align=C)
P("EXHIBIT A", bold=True, size=12, align=C, sa=2)
P("ARTICLE X: FINANCIAL INTEGRITY AND RELATED-PARTY TRANSACTIONS", bold=True, size=12, align=C, sa=10)
BYLAW = [
    ("10.1 Definitions.",
     "\"Related Party\" means (a) any director or officer of the Corporation; (b) any holder of 10% or more of the "
     "Corporation's outstanding shares; (c) any spouse, child, parent or sibling of a person in (a) or (b); and (d) any "
     "entity in which a person in (a) through (c) holds a material financial interest or serves as a director or officer, "
     "including without limitation Microcosm, Inc. and Microcosm Astronautics Books. \"Related-Party Transaction\" means "
     "any payment, loan, advance, guarantee, lease, cost allocation, reclassification, interest accrual, set-off, netting, "
     "forgiveness or other transfer of value between the Corporation and a Related Party. \"Disinterested Director\" "
     "means a director who is not, and has no material financial interest in, a party to the transaction."),
    ("10.2 Approval Required.",
     "No Related-Party Transaction shall be entered into, paid, recorded or modified unless (a) the material facts, "
     "including written invoices or agreements supporting the amount, are disclosed to all directors, and (b) it is "
     "approved by a majority of the Disinterested Directors (who may not be fewer than two) in a manner meeting Section "
     "310(a)(2) of the California Corporations Code, or by the shareholders under Section 310(a)(1) with the shares "
     "owned by the interested Related Party not being entitled to vote. Ordinary payroll to an employee who is a Related "
     "Party at rates previously so approved does not require further approval."),
    ("10.3 Freeze Pending Reconciliation.",
     "Until the Independent Reconciliation report has been delivered and its adjustments recorded, the Corporation "
     "shall not (a) repay, net, offset, increase or accrue interest on any amount recorded as owed to a Related Party, "
     "or (b) make any distribution to shareholders."),
    ("10.4 Independent Reconciliation.",
     "Within 15 days after this Article takes effect, the Disinterested Directors shall engage an independent certified "
     "public accounting firm that has not, within the prior three years, been engaged by any Related Party or worked at "
     "the direction of any Related Party (the \"Reviewer\"). The Reviewer shall, within 45 days after engagement, "
     "reconcile and report on: (a) accounts receivable unbilled (account 1110) and customer deposits (account 2380), "
     "tank by tank, against customer purchase orders, invoices, billing milestones and deliveries; (b) all balances owed "
     "to or by Microcosm, Inc. and Microcosm Astronautics Books, including the entries recorded on or about April 30, "
     "2026, with supporting invoices and agreements; (c) the allocation of facility rent, utilities, software, office "
     "and shared services between the Corporation and any Related Party from January 1, 2024; (d) all amounts paid "
     "to any Related Party from July 1, 2024; and (e) the share register against issued certificates and paid-in "
     "capital by holder. The Corporation shall record the Reviewer's adjustments within 10 days after receiving the report."),
    ("10.5 Director Access.",
     "Every director shall receive, within three business days after written request, read-only access to the "
     "Corporation's accounting system and supporting records, consistent with Article V, Section 10 and Article VII, "
     "Section 4 of these Bylaws and Section 1602 of the California Corporations Code. No officer, employee or contractor of the Corporation shall restrict or delay that access."),
    ("10.6 Share Register.",
     "The Secretary shall maintain the share register reconciled to the certificates issued and shall furnish it to "
     "any shareholder on written request within 10 days."),
    ("10.7 Amendment.",
     "This Article may be amended or repealed only by approval of the outstanding shares (Section 152 of the California "
     "Corporations Code), and the Board shall have no power to amend or repeal it."),
]
for head, body in BYLAW:
    p = P(sa=6)
    r = p.add_run(head + " ")
    r.bold = True
    p.add_run(body)
doc.add_page_break()

# ── Exhibit B: notice to non-consenting holders ─────────────────────────────
P("DRAFT — NOT EXECUTED", bold=True, size=9, color=RED, align=C)
P("EXHIBIT B", bold=True, size=12, align=C, sa=2)
P("NOTICE OF SHAREHOLDER ACTION TAKEN WITHOUT A MEETING", bold=True, size=12, align=C, sa=2)
P("(California Corporations Code §603(b)(2))", italic=True, align=C, sa=10)
P("To the shareholders of Scorpius Space Launch Company who did not sign the written consent:")
P("On [date], holders of a majority of the Corporation's outstanding shares approved, by written consent and "
  "without a meeting: (1) a bylaw amendment requiring disinterested-director approval of related-party transactions, "
  "a freeze on related-party repayments and distributions until the books are reconciled, director access to the "
  "accounting records, and an accurate share register; (2) an independent reconciliation of the disputed accounts; "
  "(3) correction of the share register; and (4) a reservation of rights regarding related-party entries not approved "
  "as required. A copy of the consent and the bylaw amendment is enclosed.")
P("These actions protect every shareholder's share of any sale proceeds by making sure the books are accurate before "
  "any payout.")
P("Secretary, Scorpius Space Launch Company", sa=12)
P("DRAFT — INTERNAL — NOT EXECUTED — NOT SENT — prepared for John Quinn, 2026-10-12", bold=True, size=8, color=RED, align=C)

tmp = OUT + ".tmp.docx"
doc.save(tmp)
os.replace(tmp, OUT)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", f"{WS}/drafts", OUT],
               check=True, capture_output=True)
print(OUT)
