#!/usr/bin/env python3
"""Series A deck rev-2026-10-16_r2 (Jon 10/16 22:04: "canonical deck should show YTD revenue and EBITDA through end of Sept").
QuickBooks P&L accrual 1/1–9/30/2026: operating revenue $1,577,061 (Total Income $1,577,984 less $923 interest income);
NOI $244,918; interest expense $34,042; no 2026 D&A booked through Sept → EBITDA $278,036 (17.6%).
YTD Sep 2025 revenue $588,245 → +168% YoY."""
from pptx import Presentation
C="/Users/exomac/.openclaw/workspace/data-room/Exos_SeriesA/_CANONICAL"
SRC=f"{C}/Exos_SeriesA_StrategicDeck_2026-09-06_CURRENT_rev-2026-10-16.pptx"
OUT=f"{C}/Exos_SeriesA_StrategicDeck_2026-09-06_CURRENT_rev-2026-10-16_r2.pptx"
# slide 12 KPI tiles, in order: value, label
TILES={("$1.49M","FY2025 revenue"):("$1.58M","2026 YTD revenue (Jan–Sep) · +168% YoY"),
       ("100%","3-yr revenue CAGR (2022–25)"):("$278K","2026 YTD EBITDA (Jan–Sep) · 18% margin"),
       ("17%","2025 operating margin (first profit)"):("$1.49M","FY2025 revenue (first profitable year)")}
HIST=("first profitable year.","first profitable year; $1.58M in 2026 through September.")
p=Presentation(SRC); s=p.slides[11]
paras=[(sh,para) for sh in s.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs]
txt=[pa.text.strip() for _,pa in paras]
def setp(pa,new):
    pa.runs[0].text=new
    for r in pa.runs[1:]: r.text=""
done=0
for i,(sh,pa) in enumerate(paras[:-1]):
    nxt=paras[i+1][1]
    key=(pa.text.strip(),nxt.text.strip())
    if key in TILES and sh is not paras[i+1][0]:
        v,l=TILES[key]; setp(pa,v); setp(nxt,l); done+=1
for sh,pa in paras:
    if pa.text.startswith("▸  History:") and HIST[0] in pa.text:
        for r in pa.runs:
            if HIST[0] in r.text: r.text=r.text.replace(*HIST); done+=1
    if "Historical figures per QuickBooks (accrual, unaudited)" in pa.text:
        for r in pa.runs:
            r.text=r.text.replace("Historical figures per QuickBooks (accrual, unaudited);","Historical and 2026 YTD (Jan–Sep) figures per QuickBooks (accrual, unaudited); EBITDA = operating income + net interest; no 2026 depreciation booked through Sept;")
p.save(OUT); print("changes",done)
