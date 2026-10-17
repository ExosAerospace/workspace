#!/usr/bin/env python3
"""Jon 10/16 22:22: the hardware-history timeline slide (his screenshot: "Our OTV-prop solutions…" + "2029-30 EXIT Potentials")
"needs to be in every deck". Source = slide 19 of Exos_PitchDeck_FULL (MC upload 8/30), kept at
data-room/Exos_SeriesA/_RELATED_NotRoomFiles/Exos_PitchDeck_FULL_2026-08-30_TIMELINE_SOURCE.pptx.
Copied verbatim (shapes + images), locked to black bg / white text, placed as slide 4 in:
 - SBI + 650 Japan decks (r5 → r6), page-number boxes renumbered
 - canonical Strategic deck (rev-2026-10-16_r2 → _r3), replacing its older slide-4 variant."""
import copy, re
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml
from pptx.opc.package import Part
WS="/Users/exomac/.openclaw/workspace"
SRC=f"{WS}/data-room/Exos_SeriesA/_RELATED_NotRoomFiles/Exos_PitchDeck_FULL_2026-08-30_TIMELINE_SOURCE.pptx"
RELNS="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
src=Presentation(SRC).slides[18]
assert "OTV-prop" in " ".join(sh.text_frame.text for sh in src.shapes if sh.has_text_frame)
def insert_timeline(prs, at):
    blank=min(prs.slide_layouts,key=lambda l:len(l.placeholders))
    dst=prs.slides.add_slide(blank)
    for ph in list(dst.shapes): ph._element.getparent().remove(ph._element)
    idmap={}
    pkg=prs.part.package
    dg=[(rId,rel) for rId,rel in src.part.rels.items() if "diagram" in rel.reltype]
    for rId,rel in src.part.rels.items():
        if "image" in rel.reltype or "media" in rel.reltype:
            idmap[rId]=dst.part.relate_to(rel.target_part,rel.reltype)
    # SmartArt (the 2029-30 EXIT bubble): clone each diagram part under a fresh partname, with its own image rels
    newparts={}
    for rId,rel in dg:
        sp=rel.target_part; base=re.sub(r"\d+\.xml$","%d.xml",str(sp.partname))
        blob=sp.blob
        np_=Part(pkg.next_partname(base),sp.content_type,pkg,blob)
        for r2,rel2 in sp.rels.items():
            if rel2.is_external: continue
            got=np_.relate_to(rel2.target_part,rel2.reltype)
            if got!=r2: blob=blob.replace(('"%s"'%r2).encode(),('"%s"'%got).encode())
        np_._blob=blob
        idmap[rId]=dst.part.relate_to(np_,rel.reltype); newparts[rId]=(np_,rel.reltype)
    for rId,(np_,rt) in newparts.items():
        if rt.endswith("/diagramData"):
            b=np_._blob
            for o,n in idmap.items(): b=b.replace(('relId="%s"'%o).encode(),('relId="%s"'%n).encode())
            np_._blob=b
    for shp in src.shapes:
        el=copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in (RELNS+"embed",RELNS+"link",RELNS+"dm",RELNS+"lo",RELNS+"qs",RELNS+"cs"):
                v=node.get(attr)
                if v in idmap: node.set(attr,idmap[v])
        dst.shapes._spTree.append(el)
    sld=dst._element; cSld=sld.find(qn("p:cSld")); old=cSld.find(qn("p:bg"))
    if old is not None: cSld.remove(old)
    cSld.insert(0,parse_xml('<p:bg %s><p:bgPr><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'%nsdecls("p","a")))
    ov=sld.find(qn("p:clrMapOvr"))
    if ov is not None: sld.remove(ov)
    sld.append(parse_xml('<p:clrMapOvr %s><a:overrideClrMapping bg1="dk1" tx1="lt1" bg2="dk2" tx2="lt2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/></p:clrMapOvr>'%nsdecls("p","a")))
    for shp in dst.shapes:
        if shp.has_text_frame:
            for para in shp.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.color is not None and run.font.color.type is None: run.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
    lst=prs.slides._sldIdLst; items=list(lst); lst.remove(items[-1]); lst.insert(at,items[-1])
for k in ["SBI_Investment","650_Fund"]:
    f=f"{WS}/drafts/Exos_Japan_Round_{k}_2026-10-16_r5.pptx"; out=f.replace("_r5","_r6")
    prs=Presentation(f)
    pages={i:[sh for sh in s.shapes if sh.has_text_frame and re.fullmatch(r"\d\d",sh.text_frame.text.strip() or "x") and int(sh.text_frame.text.strip())==i] for i,s in enumerate(prs.slides,1)}
    insert_timeline(prs,3)
    for i,shs in pages.items():
        if i>=4:
            for sh in shs: sh.text_frame.paragraphs[0].runs[0].text=f"{i+1:02d}"
    prs.save(out); print("WROTE",out,len(prs.slides))
C=f"{WS}/data-room/Exos_SeriesA/_CANONICAL"
f=f"{C}/Exos_SeriesA_StrategicDeck_2026-09-06_CURRENT_rev-2026-10-16_r2.pptx"; out=f.replace("_r2.pptx","_r3.pptx")
prs=Presentation(f)
old=prs.slides[3]; assert "SARGE vehicle" in " ".join(sh.text_frame.text for sh in old.shapes if sh.has_text_frame)
insert_timeline(prs,4)  # add first (fresh partname), then drop the old slide 4
lst=prs.slides._sldIdLst; el=list(lst)[3]; prs.part.drop_rel(el.get(qn("r:id"))); lst.remove(el)
prs.save(out); print("WROTE",out,len(prs.slides))
