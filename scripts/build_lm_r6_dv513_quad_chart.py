#!/usr/bin/env python3
"""LM SBIR 2026.R6 — DAF26BX06-DV513 (Autonomous On-Orbit Logistics & Sustainment) quad chart for the
Lockheed Martin Letter-of-Support / Phase II teaming ask (MC #3008). Non-proprietary: no LM NDA in place.

Layout helpers are reused from build_gungnir_ftv_quad_chart.py; content lives in C.
Output: drafts/LM_SBIR_2026R6/ (.pptx editable; PDF via soffice).
"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from build_gungnir_ftv_quad_chart import tb, section, rect, line, header, footer, _runs, NAVY, GRAY, DARK, LIGHT, WHITE, ACCENT, IMG
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = 'drafts/LM_SBIR_2026R6/Exos_LM_R6_DV513_QuadChart_DRAFT_2026-09-30.pptx'

C = {
 'title': 'Exos Aerospace  |  Recoverable µg Cryo-Transfer Flight Demo',
 'subtitle': 'Retiring the on-orbit refuel function: tank-to-tank cryogenic propellant transfer in microgravity on a reusable suborbital platform',
 'tag': 'SBIR 2026.R6  |  DAF26BX06-DV513 (D2P2)',
 'cta': 'Autonomous On-Orbit Logistics & Sustainment',
 'img_caption': 'BLK3 reusable VTVL vehicle, tethered hover test, Greenville TX',
 'conops': ['1  Ground 1-g transfer baseline\n    (Greenville test set)', '2  Integrate in BLK3 payload bay',
            '3  µg coast: tank-to-tank\n    cryo transfer, instrumented', '4  Recover · inspect · refly'],
 'q1_tagline': 'Refuel is the gating function of an on-orbit logistics enterprise. Exos flies the transfer hardware in microgravity, brings it back, and flies it again.',
 'q2': [
  ('Objective', ['Demonstrate <b>tank-to-tank cryogenic propellant transfer in microgravity</b> (mass gauging, thermal, pressure, flow) and compare it against a characterized 1-g ground baseline.',
                 'Deliver recovered flight hardware and a flight dataset that de-risks the <b>refuel</b> leg of DV513\'s resupply / refuel / repair / recover architecture.']),
  ('Innovation', ['<b>Recoverable, repeatable</b> space-environment test points for a fraction of the cost and schedule of an orbital campaign.',
                  'Linerless composite <b>Type-V cryo tankage</b> (the class that flew as fuel and LOX tanks on IM-1 Odysseus and IM-2 Athena) is the leading tank candidate.']),
  ('Team', ['<b>Exos</b> (prime / integrator, BLK3 flight + Greenville transfer test set) · <b>Princeton Cryogenics</b> (transfer system, cryo-thermal) · <b>Scorpius Space Launch Co.</b> (Type-V tankage, subject to trade).',
            '<b>Lockheed Martin (proposed):</b> active Phase II teammate. Receiver-side / architecture requirements, data review, and the Phase III transition path.']),
 ],
 'ask': '<b>Ask of LM:</b> Letter of Support + Phase II teaming on DV513.',
 'cost_rows': [
  ['Exos labor: PM, systems eng., integration, test', '355'],
  ['Princeton Cryogenics: transfer system & cryo-thermal', '185'],
  ['Source/receiver tank set (Type-V candidate)', '130'],
  ['Transfer hardware, QD, instrumentation', '95'],
  ['Ground transfer campaign (Greenville)', '80'],
  ['Flight integration + suborbital demo ops', '95'],
  ['Travel, reserve, G&A', '45'],
  ['Total (12 months)', '985'],
 ],
 'milestones': 'M1 SRR + cryogen/surrogate decision (Mo 2) · M2 PDR + tank down-select (Mo 4) · M3 CDR (Mo 7) · M4 ground 1-g transfer campaign (Mo 9) · M5 µg flight transfer + recovery (Mo 12) · final report + transition package',
 'cost_note': 'Proposal-stage estimate. TABA (if allowed) is additive. LM Phase II scope/share to be set jointly.',
 'q4': ['<b>Success:</b> an instrumented µg cryogenic transfer, recovered and inspected, with a flight-vs-1-g dataset and a transition package delivered.'],
 'transition_ladder': [
  ('D2P2 (DV513)', 'recoverable µg cryo-transfer flight demo, LM on the team', '$985K / 12 mo'),
  ('Phase III with LM', 'transfer + tankage scaled into LM on-orbit logistics / depot architectures', 'LM-integrated'),
  ('USSF sustainment', 'refuel capability for LEO/MEO/GEO/cislunar assets; repeat flights qualify upgrades', 'SSC / SpaceWERX'),
  ('Commercial OTV', 'Type-V tank sets already baseline tankage for a commercial OTV program', 'dual-use pull'),
 ],
 'footer': 'Distribution: non-proprietary summary; no export-controlled data (no LM NDA in place). Nothing is awarded. BLK3 received its FAA launch license Mar 2026 and is preparing for flight; '
           'Exos has flown commercial payloads since 2018 (SARGE). Type-V tanks are flight-proven as tankage; in-space cryo transfer on this architecture has not been demonstrated and is the objective of this effort.',
 'slide2': {
  'subtitle': 'Backup: why Exos, and what we are asking of Lockheed Martin',
  'left': [
   ('Why Exos', ['One of only three FAA-licensed vertical-takeoff reusable rocket operators with demonstrated recovery/reflight (with SpaceX and Blue Origin).',
                 'Commercial payloads flown since 2018 (SARGE). BLK3 is our latest FAA-licensed vehicle (Mar 2026).',
                 'Minimal infrastructure: needs a pad plus power/comms. Fast, recoverable, repeatable test cadence.',
                 'Greenville propellant-transfer ground test set for the 1-g baseline before any flight.',
                 'Current USAF SBIR Phase II performer (FA9300-25-C-6022).']),
   ('Tankage heritage', ['Scorpius linerless composite Type-V cryo tanks flew as fuel and LOX tanks on IM-1 Odysseus and IM-2 Athena (TRL 9 as tankage).',
                         'Exos holds an 18% stake in Scorpius; Type-V tanks are a leading candidate, subject to a fit/cost/schedule trade.']),
  ],
  'right': [
   ('Proposed LM role (Phase II)', ['Define receiver-side and interface requirements from LM on-orbit logistics / servicing architectures.',
                                    'Participate in design reviews (SRR / PDR / CDR) and flight data review.',
                                    'Co-own the Phase III transition plan into LM programs and USSF sustainment.']),
   ('Ask', ['Letter of Support for the DV513 D2P2 submission.',
            'Confirmation of LM Phase II teaming intent; Exos will report the government award outcome to LM.',
            'An NDA if LM wants technical detail beyond this non-proprietary summary.']),
   ('Point of contact', ['John Quinn, Cofounder & CEO · jquinn@exosaero.com · +1 (972) 740-8355 · www.exosaero.com']),
  ],
 },
}


def build():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    header(s, C['title'], C['subtitle'], C['tag'])
    tb(s, 9.6, 0.5, 3.55, 0.4, C['cta'], 8.5, WHITE, align=PP_ALIGN.RIGHT)
    MX, MY = 6.2, 3.95
    line(s, MX, 1.0, MX, 7.08); line(s, 0.1, MY, 13.23, MY)

    # Q1 — picture + ConOps
    s.shapes.add_picture(IMG, Inches(0.2), Inches(1.05), height=Inches(1.95))
    tb(s, 0.2, 3.0, 3.5, 0.22, C['img_caption'], 7, GRAY)
    cx = 3.75
    tb(s, cx, 1.05, 2.4, 0.25, 'Concept of Operations', 10, NAVY, True)
    for i, step in enumerate(C['conops']):
        rect(s, cx, 1.35 + i * 0.44, 2.35, 0.38, LIGHT if i != 2 else ACCENT)
        tb(s, cx + 0.02, 1.36 + i * 0.44, 2.33, 0.36, step, 8, DARK if i != 2 else WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 0.2, 3.22, 5.95, 0.7, C['q1_tagline'], 9.5, NAVY, True)

    # Q2 — description
    section(s, MX + 0.08, 1.0, 13.23 - MX - 0.1, 2.62, C['q2'], None, 8.4)
    tb(s, MX + 0.08, 3.64, 13.23 - MX - 0.1, 0.28, C['ask'], 8.8, ACCENT)

    # Q3 — cost & schedule
    tb(s, 0.2, 4.0, 5.9, 0.25, 'Cost and Schedule ($K, 12-month D2P2)', 10.5, NAVY, True)
    rows = [['Element', '$K']] + C['cost_rows']
    t = s.shapes.add_table(len(rows), 2, Inches(0.2), Inches(4.28), Inches(5.9), Inches(1.9)).table
    t.columns[0].width = Inches(4.9); t.columns[1].width = Inches(1.0)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell = t.cell(i, j); cell.text = ''
            p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            _runs(p, v, 7.6, WHITE if i == 0 else DARK, i == 0 or i == len(rows) - 1)
            cell.margin_top = cell.margin_bottom = Inches(0.0)
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i == len(rows) - 1 else WHITE)
    tb(s, 0.2, 6.3, 5.9, 0.5, '<b>Milestones:</b> ' + C['milestones'], 7.4)
    tb(s, 0.2, 6.83, 5.9, 0.25, C['cost_note'], 6.8, GRAY)

    # Q4 — success & transition
    tb(s, MX + 0.08, 4.0, 6.9, 0.25, 'Success & Transition', 10.5, NAVY, True)
    tb(s, MX + 0.08, 4.27, 6.9, 0.5, C['q4'], 8.6)
    y = 4.8
    for k, (step, what, money) in enumerate(C['transition_ladder']):
        rect(s, MX + 0.1, y, 0.32, 0.4, NAVY)
        tb(s, MX + 0.1, y, 0.32, 0.4, str(k + 1), 11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        tb(s, MX + 0.47, y - 0.02, 4.9, 0.46, f'<b>{step}</b> — {what}', 8.1)
        tb(s, MX + 5.4, y - 0.02, 1.65, 0.46, money, 8, NAVY, True, PP_ALIGN.RIGHT)
        y += 0.52
    footer(s, C['footer'], 1)

    s2 = prs.slides.add_slide(blank)
    header(s2, 'Exos × Lockheed Martin  |  DAF26BX06-DV513', C['slide2']['subtitle'], C['tag'])
    section(s2, 0.3, 1.1, 6.3, 5.9, C['slide2']['left'], None, 11)
    section(s2, 6.8, 1.1, 6.3, 5.9, C['slide2']['right'], None, 11)
    footer(s2, C['footer'], 2)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT); print('saved', OUT)


if __name__ == '__main__':
    build()
