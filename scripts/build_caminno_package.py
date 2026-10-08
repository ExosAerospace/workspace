#!/usr/bin/env python3
"""Build the Exos <-> CAMINNO document set on PRIMARY letterhead (+ PDFs).

Sources (drafts/caminno/*.md) -> drafts/caminno/*_2026-10-08.docx/.pdf:
  - Email_Intro_Neil_Anika_TrajectoryCrossVerification  (intro email draft)
  - MNDA_Exos_CAMINNO                                    (mutual NDA)
  - ITAR_Package_Exos_CAMINNO                            (gate checklist, DD, cert, TDTCA)
Markdown subset: '# ' title, '## ' heading, '- ' bullet, '| a | b |' table rows
(first row = header), '<<PAGEBREAK>>', **bold**. Same renderer as build_jivajet_mnda.py.
"""
import os, re, subprocess
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

WS = "/Users/exomac/.openclaw/workspace"
SRC = os.path.join(WS, "assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx")
DIR = os.path.join(WS, "drafts/caminno")
DATE = "2026-10-08"
NAVY = RGBColor(0x0B, 0x22, 0x40)
DOCS = ["Email_Intro_Neil_Anika_TrajectoryCrossVerification",
        "MNDA_Exos_CAMINNO", "ITAR_Package_Exos_CAMINNO"]


def add_runs(p, text, size=None):
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        r = p.add_run(part[2:-2] if bold else part)
        r.bold = bold or None
        if size:
            r.font.size = Pt(size)


def _borders(t):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}')
        for k, v in (('val', 'single'), ('sz', '4'), ('space', '0'), ('color', '808080')):
            x.set(qn(f'w:{k}'), v)
        b.append(x)
    tblPr.append(b)


def flush_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    n = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=n)
    _borders(t)
    for i, row in enumerate(cells):
        for j in range(n):
            c = t.cell(i, j); c.text = ''
            p = c.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(1); pf.space_after = Pt(1); pf.line_spacing = 1.0
            add_runs(p, row[j] if j < len(row) else '', size=8.5)
            if i == 0:
                for r in p.runs:
                    r.bold = True
    doc.add_paragraph()


def build(name):
    doc = Document(SRC)
    body = doc.element.body
    for el in list(body):
        if el.tag.endswith("}p") or el.tag.endswith("}tbl"):
            body.remove(el)
    table = []
    for raw in open(os.path.join(DIR, name + ".md")).read().split('\n'):
        s = raw.rstrip()
        if s.startswith('|'):
            table.append(s); continue
        if table:
            flush_table(doc, table); table = []
        if s.strip() == '<<PAGEBREAK>>':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); continue
        if s.strip() == '---':
            continue
        if s.strip() == '':
            continue
        if s.startswith('# '):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(s[2:]); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY
            continue
        if s.startswith('## ') and '↔' in s:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(s[3:]); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
            continue
        if s.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
            r = p.add_run(s[3:]); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
            continue
        if s.startswith('- '):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, '•  ' + s[2:]); continue
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        add_runs(p, s)
    if table:
        flush_table(doc, table)
    out = os.path.join(DIR, f"{name}_{DATE}.docx")
    doc.save(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", DIR, out],
                   check=True, capture_output=True)
    print("WROTE", out, "+ .pdf")


if __name__ == "__main__":
    for d in DOCS:
        build(d)
