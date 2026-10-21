#!/usr/bin/env python3
"""Canonical internal documents for the orbital first-flight payload-bay prospect (Jon 2026-10-21, Discord DM 06:16–07:06).
Builds three editable .docx files in drafts/:
  First_Flight_Payload_Concepts_<date>.docx      — what to fly on an orbital maiden launch (payloads that pay on any outcome)
  AEDC_Velocity_OnRamp_Brief_<date>.docx         — partner on-ramp construct, OFAC-gated compliance support, 30-second script
  Loft_Orbital_Fit_Assessment_<date>.docx        — where Loft fits (phase 2), and where it doesn't
Sources are cited inline (Hard Rule: sourced facts only)."""
import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches

WS = "/Users/exomac/.openclaw/workspace"
D = datetime.date.today().strftime("%Y-%m-%d")
NAVY = RGBColor(0x0A, 0x16, 0x28)


def new(title, subtitle):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.9)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    p = doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
    p = doc.add_paragraph(); r = p.add_run(subtitle); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return doc


def h(doc, t):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAVY
    p.paragraph_format.space_before = Pt(10)


def para(doc, t, italic=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.italic = italic; return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(" " + it[1])
        else:
            p.add_run(it)


def sources(doc, items):
    h(doc, "Sources")
    for t, u in items:
        p = doc.add_paragraph(style="List Number"); p.add_run(t + " — "); r = p.add_run(u); r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x99)


SUB = f"Exos Aerospace · INTERNAL · Prepared by Exo for John Quinn · {datetime.date.today():%B %d, %Y} · Prospect: orbital first-flight payload-bay buyer (name TBD)"

# ---------------------------------------------------------------- 1. Payload concepts
doc = new("First-Flight Payload Concepts", SUB)
h(doc, "The core idea")
para(doc, "The prospect plans to buy the payload bay on an orbital launcher's first flight at a discount, in exchange for "
     "follow-on bookings once the vehicle has proven itself. Maiden flights fail often: nearly one in two first flights of a "
     "new rocket goes wrong [1]. So the payloads that make this profitable are ones that earn money whether the launch "
     "succeeds or fails.")
h(doc, "Concepts, ranked")
bullets(doc, [
    ("1. Reentry body that pays either way (lead concept).", "If the rocket reaches orbit, the upper stage deorbits the body and it returns reentry data from about Mach 25. If the rocket fails, the body separates on the anomaly and flies its own reentry from whatever speed it reached — high-Mach test data a hypersonics program would normally pay a full launch for. Data is telemetered during descent, so no ocean recovery is needed. Proven buyer: AFRL gave Varda a 4-year, $48M Prometheus contract to fly reentry capsules as hypersonic testbeds [2][3]."),
    ("2. Ejectable crash-survivable black box.", "Records the ascent environment (vibration, shock, acoustics, thermal). On success the launcher buys the data for its payload user's guide; on failure the launcher, its insurers and the FAA investigation need it. Builds on Exos recovery work and Exos rocket-launched radiosondes."),
    ("3. Munitions and energetics qualification.", "Fuzes, igniters, propellant grains and guidance electronics qualified against real launch loads (plus radiation and vacuum if orbit is reached); inert hypersonic test rods fit concept 1. EXPORT/ITAR: Empowered Official sign-off before pitching."),
    ("4. Missile-defense and space-tracking targets.", "Radar calibration spheres or reentry-vehicle-shaped objects released for MDA, Space Force or commercial trackers to calibrate and exercise sensors. Fits Exos's test-side prime missile-defense lane."),
    ("5. Memorial and collectible payloads.", "Precedent: Vulcan's first flight carried more than 150 Celestis memorial capsules as a secondary payload [4][5]. High margin, almost no engineering."),
    ("6. Paid ballast.", "First flights often carry dummy mass (e.g. Zhuque-3's first flight carried a 'dummy payload' [6]); replace it with paying mass such as shielding samples, materials panels or a water/propellant depot demo."),
    ("7. Payloads that stay on the upper stage.", "New Glenn's first flight kept Blue Ring Pathfinder attached to the second stage for the whole mission [7][8]. No separation system to qualify: hosted electronics, flight heritage for commercial parts, and a drag-sail deorbit demo the launcher wants anyway."),
])
h(doc, "Exos role and structure")
bullets(doc, [
    ("Prime integrator and broker:", "Exos (or the prospect) buys the bay at a discount, fills it with outcome-proof payloads (1–4), and resells slices."),
    ("Option on follow-on flights:", "the bay purchase carries a fixed price or first refusal on flights 2–5 — the prospect's upside."),
    ("BLK3 suborbital pre-qualification:", "each payload is qualified on BLK3 first (first flights planned early 2027), cutting payload-side risk before the orbital attempt."),
    ("No free flights:", "every slot is priced, including ballast and memorial slots."),
])
h(doc, "Recommendation")
para(doc, "Lead with concepts 1 and 2 packaged as a 'Maiden Flight Data Package': a real DoD buyer, value on any outcome, and "
     "squarely in Exos's recovery and reentry strength. Add 3 and 4 for defense customers and 5 as low-cost margin.")
sources(doc, [
    ("BBC Future, 'What are the odds of a successful space launch?' (2023)", "https://www.bbc.com/future/article/20230518-what-are-the-odds-of-a-successful-space-launch"),
    ("C4ISRNET, Varda lands third capsule with AFRL hypersonic data (May 2025)", "https://www.c4isrnet.com/space/2025/05/14/varda-lands-third-space-capsule-carrying-key-hypersonic-flight-data"),
    ("Varda W-3 mission page (Prometheus / AFRL)", "https://www.varda.com/mission/w-3"),
    ("Celestis, Enterprise Flight on Vulcan's first launch", "https://www.celestis.com/blog/star-trek-first-family-embarks-on-final-infinite-journey-on-a-rocket-named-vulcan"),
    ("Wikipedia, Vulcan Centaur (Cert-1, Jan 8 2024)", "https://en.wikipedia.org/wiki/Vulcan_Centaur"),
    ("Space Stats, orbital launches 2025 (Zhuque-3 first flight)", "https://spacestatsonline.com/launches/year/2025"),
    ("Blue Origin, Blue Ring Pathfinder payload for NG-1", "https://www.blueorigin.com/news/blue-ring-pathfinder-payload"),
    ("Wikipedia, New Glenn (NG-1, Jan 16 2025)", "https://en.wikipedia.org/wiki/New_Glenn"),
])
f1 = f"{WS}/drafts/First_Flight_Payload_Concepts_{D}.docx"; doc.save(f1)

# ---------------------------------------------------------------- 2. AEDC on-ramp brief
doc = new("AEDC Velocity Alliance — Partner On-Ramp Brief", SUB)
h(doc, "Construct")
bullets(doc, [
    ("Vehicle:", "AEDC Velocity Alliance, an Other Transaction consortium run by AEDC with the Air Force Test Center under 10 U.S.C. §4022 and §2808a, supporting a multibillion-dollar portfolio of test-infrastructure projects [1][2]. Exos is an active member (Opportunity 2 membership effective 30 Sep 2026). Work arrives as individual project announcements (POAs) on AcqBot."),
    ("Fit:", "Exos's Velocity positioning is flight test as a low-cost complement to ground test. The reentry body gives AEDC hypersonic flight data to check against its tunnels; the black box gives ascent-environment data."),
    ("Structure:", "Exos prime and integrator (reentry body, black box, recovery, payload integration, DoD interface); prospect = firm-fixed-price sub supplying launch access (first-flight bay + follow-on options). Exos keeps a substitution right so one maiden-flight failure doesn't sink a Velocity deliverable."),
    ("Paths in:", "(1) Named partner under Exos on a POA Exos bids as prime — fastest, recommended. (2) Own membership later — the Request for Whitepapers stays open to new applicants through 30 Sep 2032 (FA9101-26-9-Velocity_Alliance)."),
])
h(doc, "Gate and compliance support (Jon 10/21 06:39)")
para(doc, "Assume the prospect lacks what a Velocity submission requires. Exos supports him in getting there ONLY if OFAC / "
     "restricted-party screening comes back clean.")
bullets(doc, [
    ("Step 1:", "He signs the ExoSign Partner Due Diligence Intake & Background Check Authorization."),
    ("Step 2 (same day):", "OFAC lists, U.S. Consolidated Screening List and SAM.gov exclusions on him, his company, and listed owners/officers. Any match stops everything and goes to John only."),
    ("Step 3 if clean:", "background report and company diligence (registry, litigation, media, beneficial ownership)."),
    ("Step 4 if all clean, Exos assists:", "SAM.gov registration (UEI + CAGE); JCP certification (DD Form 2345); CMMC Level 2 self-assessment in SPRS (offer our Atomus setup); mutual NDA then teaming agreement; Empowered Official export review before any reentry or vehicle data is shared."),
    ("Interim:", "Exos bids as prime and names him as partner. Whether a partner with compliance still pending can be named on a POA is not yet confirmed — check the RFWP before relying on it."),
    ("Cost logic:", "supporting his compliance is real Exos work — the strongest argument for the $5M strategic investment."),
])
h(doc, "30-second introduction script (about 80 words)")
para(doc, "\"Here's the bigger opportunity. Exos is a member of the Air Force's AEDC Velocity Alliance, a contract vehicle for "
     "a multibillion-dollar test modernization program that's open through 2032. We bring flight testing as the low-cost "
     "complement to ground testing. Your first-flight payload bay fits that exactly. We'd bring your company in as our "
     "launch-access partner, and Exos would handle payload integration and the DoD interface. The cleanest way to align us is "
     "a $5 million strategic investment in Exos. Want me to walk you through it?\"", italic=True)
h(doc, "If he leans in")
bullets(doc, [
    ("What the $5M pays for:", "the integration and DoD work he'd otherwise build himself — payload integration plans, the reentry body and black-box payloads, and AEDC project responses with Exos as prime."),
    ("Why Exos is prime:", "already a member with registrations and compliance in place; he gets on through us as a named partner without waiting to join."),
    ("Next step:", "he signs the due-diligence form, then NDA and teaming terms, with the investment on a separate track."),
])
h(doc, "Keep out of the conversation")
bullets(doc, [
    "\"$2.6B\" — briefed at the members-only Industry Day; not found in any public source. Say \"multibillion-dollar.\"",
    "Member count (82) and POA mechanics — from member-only notices.",
    "\"Awarded\" or \"contract won\" — Exos is a member; projects are bid individually.",
    "Making the investment a condition of the partnership — offer them side by side; a tie reads as pay-to-play to DoD.",
    "Investment terms — Rule 504 offering documents only, private conversation, no promised returns; David Mitchell (securities approver) reviews before anything goes in writing.",
    "BLK3 — \"first flights planned for early 2027\"; never bare \"licensed.\"",
])
sources(doc, [
    ("Arnold AFB, 'AEDC Velocity Alliance advances modernization for future test needs'", "https://www.arnold.af.mil/News/Article-Display/Article/4574989/aedc-velocity-alliance-advances-modernization-for-future-test-needs"),
    ("Aerospace Testing International, AEDC establishes test modernization consortium", "https://www.aerospacetestinginternational.com/news/arnold-engineering-development-complex-establishes-test-modernization-consortium.html"),
])
f2 = f"{WS}/drafts/AEDC_Velocity_OnRamp_Brief_{D}.docx"; doc.save(f2)

# ---------------------------------------------------------------- 3. Loft fit
doc = new("Loft Orbital — Fit Assessment for the First-Flight Construct", SUB)
h(doc, "Verdict")
para(doc, "Yes, as phase 2 — not on the first flight. Bring Loft in only after the prospect's screening is clean and his "
     "follow-on flight rights are confirmed.")
h(doc, "Why not flight 1")
para(doc, "Loft sells reliability: it hosts several customers' payloads on one satellite and flies proven rideshares. YAM-9 was "
     "its fourth consecutive SpaceX Transporter launch, and Loft says it is manifested through 2027 [1]. A maiden launch with "
     "roughly coin-flip odds doesn't fit that model. Flight 1 stays Exos's: reentry body and black box.")
h(doc, "Where Loft fits")
bullets(doc, [
    ("1. Buyer for the follow-on slots.", "The prospect's bet is discounted capacity on flights 2–5. Loft fills satellites with customer payloads and needs launches, so it is the natural anchor tenant for that capacity — turning his option into revenue. (Inference: that Loft wants launch supply beyond SpaceX is not verified.)"),
    ("2. Orbital government channel.", "Loft Orbital Federal is one of 12 vendors on the Space Force STEP 2.0 IDIQ (10 years, $237M ceiling), whose scope includes payload integration, launch and on-orbit support [2][3]. Pairs with AEDC Velocity: AEDC = test, STEP 2.0 = on-orbit demonstration."),
    ("3. One payload interface from suborbital to orbital.", "Loft uses a standard payload adapter (the Hub) [1]. If BLK3's bay mirrors that interface, a payload qualified on BLK3 drops straight into a Loft satellite on the orbital flights. Requires Loft to share the Hub interface spec."),
])
h(doc, "How to engage")
bullets(doc, [
    ("Partner / consortium, not acquisition:", "Loft is well funded — $170M Series C in January 2025 [4] — and is already in the Exos capability-gap tracker as the hosted-payload partner."),
    ("Export:", "Loft formed a UAE joint venture in 2024 [5]; Empowered Official review before sharing BLK3 interface or controlled technical data."),
])
sources(doc, [
    ("Loft Orbital, 'YAM-9: Advancing Missions, Compute, and Innovation in Orbit'", "https://loftorbital.com/yam-9-benchmarking-the-future-of-ai-enabled-space-infrastructure"),
    ("Space Systems Command, STEP 2.0 award to 12 vendors", "https://www.ssc.spaceforce.mil/Newsroom/Article-Display/Article/4171553/space-test-experiments-platform-20-to-bolster-ussf-spacecraft-procurement"),
    ("Washington Technology, STEP 2.0 scope (May 2025)", "https://www.washingtontechnology.com/contracts/2025/05/space-force-picks-dozen-237m-tech-development-demonstration-contract/405029"),
    ("TechCrunch, Loft Orbital $170M after $500M of bookings (Jan 2025)", "https://techcrunch.com/2025/01/15/loft-orbital-lands-a-fresh-170-million-after-logging-over-500-million-of-bookings/"),
    ("TechCrunch, Loft Orbital UAE joint venture (Aug 2024)", "https://techcrunch.com/2024/08/25/loft-orbital-forms-joint-venture-with-uae-based-firm-to-scale-satellite-production-in-the-middle-east/"),
])
f3 = f"{WS}/drafts/Loft_Orbital_Fit_Assessment_{D}.docx"; doc.save(f3)
print(f1); print(f2); print(f3)
