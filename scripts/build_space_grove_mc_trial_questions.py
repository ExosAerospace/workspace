#!/usr/bin/env python3
"""Mission Control platform trial: questions for Space Grove Ventures (external, Exos letterhead).
Questions-only derivative of drafts/mc-industry/build_sheet.py: drops the internal 'why' column, the JP Morgan
and Space Grove illustrative answers, and internal Exos references."""
import subprocess, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

WS = "/Users/exomac/.openclaw/workspace"
BODY = f"{WS}/drafts/uplink/_body_SpaceGrove_MC_Trial_Questions.docx"
OUT = f"{WS}/drafts/uplink/Exos_SpaceGrove_MissionControl_Trial_Questions_2026-10-21.docx"
NAVY, GREY = (0, 51, 102), (80, 80, 80)

SECTIONS = [
    ("1. Your organization and the trial scope", [
        "How many people, and in which locations, would use Mission Control during the trial?",
        "Who would sponsor the trial, and who would run it day to day?",
        "What would a successful 90 days look like? Please name three measurable outcomes.",
        "What should be explicitly out of scope for the trial?"]),
    ("2. Mission and decision rights", [
        "What is Space Grove's North Star in one sentence?",
        "Which decisions only you make? Which can staff make? What, if anything, could an AI agent decide?",
        "How quickly do decisions need an answer, and what should happen if one waits too long?",
        "Who approves in your place when you are unavailable?"]),
    ("3. What you track", [
        "What are the 3–5 main things your team tracks (for example companies, partners, sponsors, events, programs)?",
        "What stages does each move through, and who moves it forward?",
        "Could you share one real example from start to finish (redacted is fine)?",
        "Which numbers matter on each record (value, dates, commitments)?"]),
    ("4. Which modules to turn on", [
        "Which do you need: tasks, decisions, contacts/CRM, pipeline, documents/data room, e-signature, calendar and follow-ups, events, finance, people, procedures, research, AI agents?",
        "Which current tools would Mission Control replace, and which would it sit beside?",
        "What would make your team stop using a new tool in week two?"]),
    ("5. People and access", [
        "What roles do you have, and what should each see (and not see)?",
        "Will outside users need access (portfolio founders, sponsors, partners, board)?",
        "How does your team sign in today (Google, Microsoft, other; multi-factor)?"]),
    ("6. Systems and data", [
        "Which email and calendar platform do you use?",
        "Which CRM, accounting, file storage, e-signature and chat tools do you use?",
        "Which system is the official record for each thing you track?",
        "Can you export the last 12 months of core data (contacts, companies, deals) to start the trial with?"]),
    ("7. Security and data handling", [
        "Which regulations, contracts or internal policies govern your data (for example NDAs with founders, investor privacy)?",
        "Where may your data be hosted? Is a hosted Mission Control instance acceptable?",
        "Are there AI providers you require or do not allow?",
        "What must be logged, and for how long?",
        "Who needs to approve going live?"]),
    ("8. AI agents and automation", [
        "What are the 10 most repetitive tasks your team does each week, and who does them?",
        "What, if anything, may an AI agent send outside the organization without a person approving it?",
        "What overnight or scheduled work would be most valuable (briefs, deadline watchers, digests)?",
        "What does a wrong answer cost you? Please give an example."]),
    ("9. Communication", [
        "Where does your team communicate today, and where should Mission Control send notifications?",
        "Who wants a daily brief, and in what format?"]),
    ("10. Documents and brand", [
        "Please share your logo, colors, fonts, templates and any required disclaimers.",
        "Who approves final versions of documents that go outside the organization?"]),
    ("11. Reporting", [
        "Which recurring reports do you produce today (weekly, monthly, sponsor or board updates)?",
        "Which recurring meetings should Mission Control prepare agendas for?"]),
    ("12. Trial logistics", [
        "Preferred trial length and start date?",
        "Who on your side would handle the few setup steps (sign-in, account access, approvals)?",
        "When would you decide whether to continue after the trial?"]),
]


def run(p, text, bold=False, size=11, color=None):
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    return r


def build():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for text, size, bold, color in [
        ("Mission Control Platform Trial", 17, True, NAVY),
        ("Discovery Questions for Space Grove Ventures", 13, True, NAVY),
        ("Prepared by Exos Aerospace  |  October 21, 2026", 10, False, GREY)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, text, bold, size, color)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    run(p, "Mission Control is the AI-assisted operations platform Exos runs on, covering tasks, decisions, contacts, "
           "pipeline, documents and AI agents. Your answers let us set up a trial instance configured for Space Grove "
           "rather than a blank tool. Short answers are fine, and anything you would rather cover on a call can be "
           "left blank.")
    p = doc.add_paragraph()
    run(p, "Helpful to share if available: ", True)
    run(p, "a simple org chart, one example record from start to finish, your list of software tools, and your brand kit.")

    n = 0
    for title, qs in SECTIONS:
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
        run(p, title, True, 12, NAVY)
        for q in qs:
            n += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3); p.paragraph_format.keep_with_next = True
            run(p, f"{n}.  ", True); run(p, q)
            a = doc.add_paragraph(); a.paragraph_format.left_indent = Inches(0.3)
            run(a, "Answer: ", False, 10, GREY)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    run(p, "Contact: ", True); run(p, "John Quinn, Cofounder & CEO, Exos Aerospace — jquinn@exosaero.com")
    doc.save(BODY)
    subprocess.run([sys.executable, f"{WS}/scripts/apply_exos_letterhead.py", BODY, OUT], check=True)
    print(OUT, n)


if __name__ == "__main__":
    build()
