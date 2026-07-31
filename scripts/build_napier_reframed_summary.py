#!/usr/bin/env python3
"""Reframed BLK3 payload-revenue financing summary for Napier Park Special Opportunities (Alex Maldutis).
Menu-of-structures framing (creative), not a rigid $750K factor. Clean investor memo docx."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
NAVY=RGBColor(0x0F,0x2A,0x4A); GOLDT=RGBColor(0x9A,0x6B,0x10); GREYT=RGBColor(0x44,0x44,0x44)
TEMPLATE="assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx"
doc=Document(TEMPLATE) if os.path.exists(TEMPLATE) else Document()
doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(10.5)
# clear the template's single empty body paragraph so content starts cleanly under the letterhead
if doc.paragraphs and not doc.paragraphs[0].text.strip():
    el=doc.paragraphs[0]._element; el.getparent().remove(el)
def H(t,l):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(t); r.font.bold=True; r.font.color.rgb=NAVY; r.font.size=Pt(14 if l==1 else 12)
    return p
def P(t,bold=False,italic=False,size=10.5):
    p=doc.add_paragraph(); r=p.add_run(t); r.font.bold=bold; r.italic=italic; r.font.size=Pt(size); return p
def B(lead,rest=""):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Pt(18); p.paragraph_format.space_after=Pt(3)
    b=p.add_run("•  "); b.font.size=Pt(10.5)
    if rest:
        r=p.add_run(lead+" "); r.font.bold=True; r.font.size=Pt(10.5); r2=p.add_run(rest); r2.font.size=Pt(10.5)
    else:
        r=p.add_run(lead); r.font.size=Pt(10.5)
    return p

tp=doc.add_paragraph(); r=tp.add_run("BLK3 Payload-Revenue Financing"); r.font.size=Pt(20); r.font.bold=True; r.font.color.rgb=NAVY
sp=doc.add_paragraph(); r=sp.add_run("Structured Facility Concept  ·  A Menu for Discussion"); r.font.size=Pt(12.5); r.font.bold=True; r.font.color.rgb=GOLDT
mp=doc.add_paragraph(); r=mp.add_run("Prepared for Napier Park Global, Special Opportunities (Alex Maldutis)  ·  Referral: Mike Chacos, Trade Vector  ·  Exos Aerospace  ·  2026-07-22  ·  Proprietary"); r.font.size=Pt(9); r.font.color.rgb=GREYT

H("1. The Opportunity in Brief",1)
P("Exos Aerospace is one of only three FAA-licensed reusable rocket launch providers in the United States, and the only one "
  "currently selling suborbital payload services. We are financing BLK3 flight cadence against priced payload receivables and "
  "the reusable flight asset itself, bridging today's committed demand to larger contracted government awards. We are open to "
  "getting creative on structure and want your read on what fits your book.")

H("2. The Revenue We Finance Against",1)
B("Pricing:","$7,500/kg. A full 135 kg BLK3 bay = $1,012,500 gross revenue per flight.")
B("Payment waterfall:","50% deposit at contract execution, 40% at T-90 to T-120 days, 10% post-flight. A single full-bay deposit = $506,250 at signing.")
B("Anchor already performing:","the entX Launch Services Agreement (Flight-2 radioisotope-heater-unit payload) is signed and paying, not projected: over $400K received to date, including a $328K payment that landed today.")
B("Reusable upside:","bookings beyond 135 kg trigger a follow-on flight on the same airframe three days later at under $50K marginal cost, creating a second full-margin bay.")
B("Rideshare diversification:","the 135 kg bay is most commonly sold as a rideshare of up to ~25 customers (averaging ~5.4 kg each) sharing one flight, marketed as SPACEavailable...™. For a factor this is an ideal profile: up to 25 independent receivables arriving on a rolling basis, low counterparty concentration, and continued paydown even if any single customer slips. The bay can also fill as a few larger anchor deals.")

H("3. The Ask - a ~$2M Facility on the Format We Already Run",1)
P("We are seeking a committed facility of approximately $2M that does two jobs: (1) purchases our payload and contract "
  "receivables on the same factoring format Exos already uses, and (2) stands as the finance backing letter we need to bid "
  "DoD matching-fund opportunities (Section 4). Proposed starting terms, extended from our current government-receivable factoring:",bold=True)
B("Instrument:","true-sale purchase of receivables (factoring), not a loan. No equity, no board rights, no personal guarantee.")
B("Advance:","near-par, roughly 97.6 to 99.1 cents on the dollar of face value.")
B("Pricing:","a discount priced by days-to-payment at a low-teens APR. Our current government-receivable factoring runs 13.5 to 14.0% APR.")
B("Settlement:","funded the next business day after each trade.")
B("Mechanics:","each signed Launch Services Agreement, starting with the contracted entX Flight-2 payload, is a factorable receivable; as the manifest books, each new LSA adds to the facility.")
P("Where we would value your creativity:",bold=True)
B("A committed forward facility","to purchase future signed LSAs as the manifest fills, so financing keeps pace with bookings.")
B("Best-rate treatment","for the government-backed receivables (entX milestones, and the NASA IDIQ and DoD awards as they land) relative to commercial payload receivables.")
B("Oversell upside","- bookings beyond 135 kg trigger a follow-on flight, whose receivables add collateral and recovery capacity.")
P("The structure is self-liquidating and short-duration, with a government-contract backstop (Section 4). We are anchoring to a "
  "format that is proven and already funds our business, and we are glad to build around it with you.",italic=True)

H("4. The Backstops - and Why the Facility Unlocks Award Flow",1)
B("DoD matching contracts (the key point):","many DoD matching-fund opportunities require proof of financing in place before we can even apply. A ~$2M in-place facility covers our currently forecasted matching opportunities. We win roughly one in five we pursue, but the finance backing letter is a prerequisite to bid at all, so the facility directly unlocks award flow that then repays it.")
B("DoD / TACFI path","- a ~$4M award path on the HVRR-RDRE Phase II basis (USAF FA9300-25-C-6022).")
B("NASA IDIQ launch-services contract","(anticipated 2H 2026) - an indefinite-delivery vehicle for suborbital launch services.")
P("Any one of these awards materially exceeds the facility and would repay it in full. The facility bridges the gap between committed "
  "payload demand today and the timing of these larger contracts, and it is the very thing that lets us pursue them, which is the "
  "short-duration, contract-anchored profile Special Opportunities underwrites.",italic=True)

H("5. Why Exos, Why Now",1)
B("Scarcity:","one of three FAA-licensed reusable providers; the only one currently selling payload services.")
B("Under-supplied demand:","the largest incumbent (Blue Origin New Shepard) is paused and the parabolic alternative (ZERO-G) is grounded, into a 60+ named-target pipeline.")
B("Active inbound:","demand is live, not theoretical, with Lockheed Martin reaching out to us yesterday about BLK3 capacity.")
B("Credibility:","veteran-led; Scorpius lunar-heritage tanks; active USAF SBIR Phase II; Trade Vector (Mike Chacos) financing our USAF contract.")

H("6. Proposed Next Step",1)
B("A short working session","to align on structure and your underwriting requirements.")
B("We provide:","the entX LSA, payload pricing and waterfall, pipeline and contract book, per-flight economics, and government-contract status.")

doc.add_paragraph("")
f=doc.add_paragraph(); r=f.add_run("John Quinn, Cofounder & CEO  ·  Exos Aerospace Systems & Technologies, Inc.  ·  jquinn@exosaero.com  ·  972-740-8355"); r.font.size=Pt(9); r.font.color.rgb=GREYT
n=doc.add_paragraph(); r=n.add_run("Figures reflect Exos internal working models and projected payload revenue; not a guarantee of future performance. Past flight performance is documented in the Exos data room."); r.font.size=Pt(8); r.italic=True; r.font.color.rgb=GREYT

os.makedirs("drafts",exist_ok=True)
out="drafts/Napier_Park_BLK3_Structured_Facility_Concept.docx"; doc.save(out); print("BUILT",out)
