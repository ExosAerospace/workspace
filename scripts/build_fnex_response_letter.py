"""FNEX 9/25 Advisory Agreement — plain-English response letter (Word on PRIMARY letterhead).
Run: python3 scripts/build_fnex_response_letter.py
v5 (Jon 9/26 #2596): written AS a response to Ainigma; board-approved-terms framing; breakup fee ACCEPTED narrowly.
Source of positions: drafts/Exos_FNEX_Fee_Dollars_and_Pushback_v4_2026-09-26.docx (task #2596).
"""
import sys, os, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from _letterhead_fix import apply_letterhead_fixes
LH="assets/EXOS_Letterhead_Greenville_TX_address_PRIMARY.docx"
OUT="drafts/Exos_FNEX_Agreement_Response_Letter_DRAFT_2026-09-26.docx"
ORANGE=RGBColor(0xF2,0x6A,0x21); DARK=RGBColor(0x1A,0x1A,0x1A); GREY=RGBColor(0x55,0x55,0x55)

L=[
("draft","DRAFT v5 — for Counsel review before sending. Bracketed [ ] items are John's to set."),
("p","September [__], 2026"),
("gap",),
("addr","Michael Sheldon and Graham Bruce"),
("addr","Ainigma Capital / FNEX"),
("addr","michael@ainigmacapital.com  ·  graham@ainigmacapital.com"),
("gap",),
("p","**Re: Exos response to the September 25 draft Advisory Agreement**"),
("p","Michael and Graham,"),
("p","Thank you for the draft, and for the time you have both put into Exos over the past two months. We want to get this signed and get to work, so I'll be direct about what we can and can't sign."),
("p","**The terms our board approved are the terms in your August 20 proposal and Michael's September 2 letter.** That is the deal I took to the board, and it is the deal I have authority to sign. I can't change those economics on my own, and if I take a materially different agreement back to the board I expect it to be turned down. Several points in the September 25 draft change those economics. Below is what we need in the agreement so it matches what the board approved, plus the one point where we have moved toward you."),

("h","1.  Stock fee issued at the deal price, with no discount"),
("b","Your proposal was 2% of enterprise value in stock. The draft issues that stock at a 30% discount, which turns 2% into roughly 2.86%. At a $250 million value, a stock fee shown as $5 million becomes about $7.14 million of stock at the deal price."),
("b","Your September 2 model issued your fee shares at the merger price, and that is what the board approved. **We can't accept a discount of any size.** Please issue the fee shares at the same price the PIPE investors pay."),

("h","2.  Fees earned at closing, not at signing"),
("b","The draft makes the fee payable when a merger agreement is signed. Your September 2 letter had it paid \"at the closing of the merger and receipt of PIPE funds,\" and that is the structure the board approved."),
("b","**Please restore it:** 0.5% cash and the 2% stock at closing, paid from PIPE funds actually received, and the remaining 0.5% cash after the follow-on raise. If the transaction doesn't close, no fee, stock or expenses are owed, even if a pre-merger PIPE or note has been funded. Any pre-merger money must either sit in escrow until closing or convert into Exos shares on terms we agree now. It can't become debt the company has to repay."),

("h","3.  Minimum value and PIPE"),
("b","**We accept the $250 million minimum enterprise value.**"),
("b","The PIPE needs to be at least [$15 million] funded at closing, and the company has to keep at least $10 million after every cost paid at closing: your cash fee, the placement fee, legal, audit, listing and expenses. A $10 million PIPE doesn't meet that test. If closing costs would take us below $10 million, the unpaid part of your cash fee moves to the follow-on raise."),

("h","4.  PIPE placement fee of 7%"),
("b","We agree to the 7% placement fee we discussed, on PIPE capital FNEX actually places, paid in cash at closing from those proceeds."),
("b","Please write it into the agreement as your complete compensation for the placement: no warrants, no other placement or structuring fees, and nothing on capital Exos raises itself (our Rule 504 offering, Series A, strategic or government funding). The fees on \"Additional Transactions\" during the tail should come out."),

("h","5.  Breakup fee, as we discussed"),
("b","**We agree to a breakup fee to protect your work.** If you deliver a transaction on the agreed terms and Exos walks away from it to take a better go-public or merger deal elsewhere within six months, Exos pays you [$200,000]."),
("b","It should be written narrowly so it covers exactly that case. It doesn't apply if a transaction fails for reasons outside Exos's control or doesn't meet the agreed terms, and it doesn't apply to Exos's own Rule 504 offering, Series A, strategic or government funding. It is a single fixed amount with no expenses added. Proposed language is attached."),

("h","6.  Expenses: definition, budget and cap"),
("b","The September 25 draft is the first version with an expense clause. Before we agree to one, please send (a) what counts as a reimbursable expense, (b) an itemized budget for the engagement, and (c) a hard cap of no more than [$____]. Anything over $150 is approved in advance, as your draft says. Expenses are paid only at closing from PIPE proceeds and are credited against the fee."),

("h","7.  Tie the stock to results"),
("b","On August 20 you said \"the higher it goes, the better for you, the better for us.\" We agree, and we'd like the agreement to say so. The fee shares vest in thirds: at closing with at least [$15 million] raised; when the stock is listed on NASDAQ and trades at or above the deal price within [12] months; and when it trades at [150%] of the deal price within [24] months. Unvested shares are cancelled at their deadline, and all fee shares carry a [12]-month lock-up on the same terms as the founders. Proposed language is attached."),

("h","8.  Remaining drafting points"),
("b","**Fee base:** enterprise value means the shares issued to Exos holders times the price at closing. Please remove the language reaching consideration \"of every kind\" received by Exos or our shareholders after the term, and exclude the PIPE and later capital."),
("b","**Term and tail:** a flat 6-month term with no automatic renewal and no added notice period, and a 12-month tail limited to investors you actually introduce, binding the company only and not our shareholders."),
("b","**Legal:** Texas law, mutual and capped indemnity, remove \"or its assigns,\" a key-person clause naming Michael and Graham, and confirmation that everyone receiving a fee is FINRA/SEC registered."),

("h","Next step"),
("p","Nothing here asks for more than you proposed. It is your August 20 and September 2 terms, which our board has approved, plus the 7% placement fee and the breakup fee we worked out together. If you send a revised draft on this basis, I can take it back to the board as the deal they already approved, and we'll turn it quickly."),
("gap",),
("p","Best Regards,"),
("gap",),
("addr","John Quinn"),
("addr","Cofounder & CEO"),
("addr","Exos Aerospace Systems & Technologies, Inc."),
("addr","jquinn@exosaero.com  ·  +1-972-740-8355"),
("gap",),
("h","Attachment — proposed clause language"),
("q","\"Equity Fee. Two percent (2%) of MEV, payable in Target common stock issued at the per-share price paid by investors in the PIPE (without discount) (the 'Equity Fee Shares'). The Equity Fee Shares shall be issued at Closing and shall vest as follows: (i) one-third upon Closing of the Transaction with a PIPE of not less than [$15,000,000] funded at Closing; (ii) one-third upon the Target's common stock being listed on the NASDAQ Capital Market or higher and its 30-trading-day volume-weighted average price equalling or exceeding [100%] of the PIPE price at any time within [12] months after Closing; and (iii) one-third upon the 30-trading-day VWAP equalling or exceeding [150%] of the PIPE price at any time within [24] months after Closing. Equity Fee Shares not vested by the applicable deadline are forfeited and cancelled without consideration. All Equity Fee Shares, vested or not, are subject to a lock-up of not less than [12] months after Closing and on terms no more favourable than those applying to the Company's founders. No Fees of any kind are earned or payable unless and until the Transaction closes.\""),
("q","\"Cash Fee. One percent (1%) of MEV, earned only upon Closing, of which one-half (0.5%) is payable at Closing solely from PIPE proceeds actually received by the Company, and one-half (0.5%) is payable upon receipt of the first post-Closing financing; provided that the Company retains not less than $10,000,000 of PIPE proceeds after all Transaction costs payable at Closing, with any shortfall deferred to the first post-Closing financing. If the Transaction does not close, no Fees or expenses of any kind are owed, whether or not any PIPE or pre-merger financing has been funded.\""),
("q","\"Placement Fee. For a PIPE closed concurrently with the Transaction, FNEX shall receive a cash placement fee of seven percent (7%) of gross PIPE proceeds actually received by the Company from investors introduced by FNEX, payable at Closing from those proceeds. The Placement Fee is FNEX's sole compensation for the PIPE; no warrants, options, or other fees shall be payable. No Placement Fee is payable on capital raised by the Company from any other source, including its Rule 504 offering, Series A, strategic, or government funding.\""),
("q","\"Breakup Fee. If, during the Term, FNEX has presented the Company with a Transaction meeting the terms of this Agreement (including a MEV of not less than $250,000,000 and a PIPE of not less than [$15,000,000] committed at Closing), and the Company elects not to proceed with that Transaction and, within six (6) months after that election, enters into a definitive agreement for a competing merger, reverse merger, SPAC or public-listing transaction with a party not introduced by FNEX, the Company shall pay FNEX a one-time fee of [$200,000] (the 'Breakup Fee'). The Breakup Fee is FNEX's sole remedy in that case and is not increased by expenses. No Breakup Fee is payable if the Transaction fails for reasons outside the Company's control, or on any capital the Company raises through its Rule 504 offering, Series A, or strategic or government funding.\""),
]

doc=Document(LH); apply_letterhead_fixes(doc)
def runs(p,text,size,color,italic=False):
    for i,part in enumerate(re.split(r"\*\*",text)):
        if not part: continue
        r=p.add_run(part); r.font.name="Calibri"; r.font.size=Pt(size); r.font.color.rgb=color
        r.font.bold=(i%2==1); r.font.italic=italic
for it in L:
    k=it[0]; p=doc.add_paragraph(); f=p.paragraph_format
    if k=="draft": runs(p,it[1],9,ORANGE,True); f.space_after=Pt(10)
    elif k=="p": runs(p,it[1],11,DARK); f.space_after=Pt(8)
    elif k=="addr": runs(p,it[1],11,DARK); f.space_after=Pt(0)
    elif k=="gap": f.space_after=Pt(2)
    elif k=="h": runs(p,it[1],12,ORANGE); p.runs[0].font.bold=True; f.space_before=Pt(10); f.space_after=Pt(4); f.keep_with_next=True
    elif k=="b":
        f.left_indent=Inches(0.3); f.first_line_indent=Inches(-0.2); f.space_after=Pt(4)
        r=p.add_run("•  "); r.font.size=Pt(11); r.font.color.rgb=ORANGE; runs(p,it[1],11,DARK)
    elif k=="q": f.left_indent=Inches(0.3); f.space_after=Pt(6); runs(p,it[1],10,GREY,True)
doc.save(OUT); print("wrote",OUT)
