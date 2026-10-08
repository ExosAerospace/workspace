#!/usr/bin/env python3
"""dibc_leverage_watch.py — track how Exos leverages its DIBC membership (Jon 2026-10-08, hub MC #870).

DIBC solicitations are members-only and land in John's M365 inbox (he is the Primary POC), where the
Proton-based pollers can't see them. This job:
  1. Reads M365 message HEADERS only (Mail.ReadBasic: from, subject, receivedDateTime; never bodies) for
     mail from ATI/DIBC or with DIBC in the subject. Dedupes by message id in state/dibc-leverage-state.json.
  2. Classifies each new hit: solicitation (RPP/RFP/amendment/Q&A/due) vs other, plus Exos fit terms.
  3. New SOLICITATION -> appends a dated note to #870 and DMs Jon the single hands-on step (forward the
     RPP to @Exobot). Non-solicitation mail is logged only (R-NONEWS: no news = silence).
  4. Rebuilds the tracker workbook (Lanes / Solicitation log / Award intel, from config/dibc-leverage.json)
     and attaches it to #870 via mc-attach.py whenever its content changes.

  dibc_leverage_watch.py              # run (launchd daily)
  dibc_leverage_watch.py --dry-run    # print hits, write nothing
  dibc_leverage_watch.py --build      # rebuild + attach the workbook only
"""
import hashlib, json, os, subprocess, sys, urllib.parse, urllib.request
from datetime import datetime, timedelta, timezone

WS = "/Users/exomac/.openclaw/workspace"
CFG = os.path.join(WS, "config", "dibc-leverage.json")
STATE = os.path.join(WS, "state", "dibc-leverage-state.json")
SECRETS = os.path.expanduser("~/.openclaw/secrets.json")
SB = "https://pcrgpuoljylqcumijvnw.supabase.co/rest/v1"
MAILBOX = "jquinn@exosaero.com"
OUT = os.path.join(WS, "drafts", "Exos_DIBC_Leverage_Tracker.xlsx")
SEED_SINCE = "2026-08-15T00:00:00Z"   # first run backfills from the membership application
JON_DM = "user:1501632195163590739"


def load(path, default):
    try:
        return json.load(open(path))
    except Exception:
        return default


def save(path, obj):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path + ".tmp", "w") as f:
        json.dump(obj, f, indent=1)
    os.replace(path + ".tmp", path)


def graph_token(s):
    login = s.get("ms_graph_login_host", "login.microsoftonline.us")
    graph = s.get("ms_graph_api_host", "graph.microsoft.us")
    data = urllib.parse.urlencode({"client_id": s["ms_graph_client_id"], "client_secret": s["ms_graph_client_secret"],
                                   "scope": f"https://{graph}/.default", "grant_type": "client_credentials"}).encode()
    req = urllib.request.Request(f"https://{login}/{s['ms_graph_tenant_id']}/oauth2/v2.0/token", data=data, method="POST")
    return json.loads(urllib.request.urlopen(req, timeout=30).read())["access_token"]


def scan_headers(s, since):
    graph = s.get("ms_graph_api_host", "graph.microsoft.us")
    tok = graph_token(s)
    qs = urllib.parse.urlencode({"$select": "id,from,subject,receivedDateTime", "$top": "200",
                                 "$filter": f"receivedDateTime ge {since}"}, quote_via=urllib.parse.quote, safe="")
    url = f"https://{graph}/v1.0/users/{MAILBOX}/messages?{qs}"   # all folders: rules may file ATI mail
    pages = 0
    while url and pages < 40:
        d = json.loads(urllib.request.urlopen(urllib.request.Request(url, headers={"Authorization": "Bearer " + tok}), timeout=45).read())
        for m in d.get("value", []):
            yield m
        url = d.get("@odata.nextLink"); pages += 1


def classify(cfg, sender, subject):
    sl, sub = sender.lower(), subject.lower()
    if not (any(sl.endswith("@" + d) or sl.endswith("." + d) for d in cfg["watch_sender_domains"])
            or any(t in sub for t in cfg["watch_subject_terms"])):
        return None
    if sl.endswith("pm.me") or sl.endswith("exosaero.com"):
        return None   # our own Exo/Jon traffic about DIBC, not consortium mail
    kind = "solicitation" if any(t in sub for t in cfg["solicitation_terms"]) else "other"
    fit = sorted({t for t in cfg["fit_terms"] if t in sub})
    return kind, fit


def sb(method, path, body=None):
    k = load(SECRETS, {})["supabase_service_role_key"]
    h = {"apikey": k, "Authorization": "Bearer " + k, "Content-Type": "application/json"}
    r = urllib.request.Request(SB + "/" + path, method=method, headers=h,
                               data=json.dumps(body).encode() if body is not None else None)
    raw = urllib.request.urlopen(r, timeout=30).read().decode()
    return json.loads(raw) if raw.strip() else None


def note_hub(cfg, hit):
    tid = cfg["hub_task"]
    cur = sb("GET", f"tasks?id=eq.{tid}&select=description")[0]["description"] or ""
    fit = ", ".join(hit["fit"]) or "none flagged from subject"
    add = (f"\n\n[{hit['received'][:10]} DIBC SOLICITATION, auto-logged by dibc_leverage_watch] "
           f"\"{hit['subject']}\" from {hit['from']}. Exos fit terms: {fit}. "
           f"NEXT: John forwards the RPP + attachments to @Exobot; Exo screens fit and drafts against it.")
    sb("PATCH", f"tasks?id=eq.{tid}", {"description": cur + add})


def dm_jon(hit):
    msg = (f"📥 New DIBC solicitation in your Outlook: \"{hit['subject']}\" ({hit['received'][:10]}). "
           f"Forward it with attachments to @Exobot and I'll screen fit and draft against it. Logged on MC #870.")
    subprocess.run(["/usr/bin/python3", os.path.join(WS, "scripts", "deliver.py"), "-m", msg,
                    "--target", JON_DM, "--source", "dibc-leverage-watch"], timeout=120)


def build(cfg, state, attach=True):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    hdr_fill, hdr_font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")

    def sheet(ws, cols, rows, widths):
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font = hdr_fill, hdr_font
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    ws = wb.active; ws.title = "Lanes"
    sheet(ws, ["ID", "Leverage lane", "Status", "Next action", "Owner", "MC task", "Gate"],
          [[l["id"], l["lane"], l["status"], l["next"], l["owner"], l["mc"], l["gate"]] for l in cfg["lanes"]],
          [5, 42, 34, 60, 20, 14, 18])
    hits = sorted(state.get("hits", {}).values(), key=lambda h: h["received"], reverse=True)
    sheet(wb.create_sheet("Solicitation log"),
          ["Received", "From", "Subject", "Type", "Exos fit terms", "Action"],
          [[h["received"][:16].replace("T", " "), h["from"], h["subject"], h["kind"], ", ".join(h["fit"]),
            "Forward RPP to @Exobot" if h["kind"] == "solicitation" else "logged"] for h in hits],
          [17, 28, 70, 13, 26, 24])
    sheet(wb.create_sheet("Award intel"), ["Date", "What", "Source"],
          [[a["date"], a["what"], a["url"]] for a in cfg["award_intel"]], [12, 80, 70])
    meta = wb.create_sheet("About")
    for r in [["Hub task", f"MC #{cfg['hub_task']}"], ["DIBC member since", cfg["member_since"]], ["OTA", cfg["ota"]],
              ["Renewal", cfg["renewal_note"]],
              ["How it updates", "Daily: Exo reads Outlook headers (sender/subject/date only) for ATI/DIBC mail, logs it here, "
                                 "and flags new solicitations on #870 + Discord. Lanes are edited in config/dibc-leverage.json."],
              ["Built", datetime.now().strftime("%Y-%m-%d %H:%M CT")]]:
        meta.append(r)
    meta.column_dimensions["A"].width, meta.column_dimensions["B"].width = 20, 110

    # Fingerprint excludes the "Built" stamp so unchanged content doesn't re-version the attachment.
    fp = hashlib.sha256(json.dumps([cfg["lanes"], cfg["award_intel"], hits], sort_keys=True).encode()).hexdigest()
    if fp == state.get("built_fp") and os.path.exists(OUT):
        return False
    wb.save(OUT)
    if attach:
        subprocess.run(["/usr/bin/python3", os.path.join(WS, "scripts", "mc-attach.py"), "--task", str(cfg["hub_task"]),
                        "--file", OUT, "--doc-key", "exos-dibc-leverage-tracker", "--by", "Exo",
                        "--title", "DIBC Leverage Tracker"], check=True, timeout=120)
    state["built_fp"] = fp
    return True


def selftest():
    """Output eval: classifier routes correctly, own traffic is excluded, and every alert is unique per message id."""
    cfg, state = load(CFG, None), load(STATE, {"hits": {}})
    cases = [("dibc@ati.org", "DIBC RPP: Solid Rocket Motor Composite Case Capacity", "solicitation"),
             ("dibconsortium@ati.org", "DIBC Welcomes - Exos", "other"),
             ("jquinnexo@pm.me", "Re: DIBC RPP draft", None),
             ("someone@vendor.com", "Quarterly newsletter", None)]
    bad = [(s, sub) for s, sub, want in cases if (classify(cfg, s, sub) or (None,))[0] != want]
    ids = list(state.get("hits", {}))
    ok = not bad and len(ids) == len(set(ids)) and all(l.get("next") for l in cfg["lanes"])
    print("PASS" if ok else f"FAIL {bad}")
    sys.exit(0 if ok else 1)


def main():
    if "--selftest" in sys.argv:
        return selftest()
    dry = "--dry-run" in sys.argv
    cfg, state = load(CFG, None), load(STATE, {"hits": {}})
    if "--build" in sys.argv:
        changed = build(cfg, state); save(STATE, state); print("built" if changed else "unchanged"); return
    s = load(SECRETS, {})
    if not s.get("ms_graph_client_id"):
        print("ms_graph creds not configured"); return
    since = state.get("last_scan") or SEED_SINCE
    since_dt = datetime.strptime(since, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc) - timedelta(days=2)
    first_run = "last_scan" not in state
    new = []
    for m in scan_headers(s, since_dt.strftime("%Y-%m-%dT%H:%M:%SZ")):
        if m["id"] in state["hits"]:
            continue
        sender = (((m.get("from") or {}).get("emailAddress") or {}).get("address") or "").lower()
        c = classify(cfg, sender, m.get("subject") or "")
        if not c:
            continue
        hit = {"received": m["receivedDateTime"], "from": sender, "subject": m.get("subject") or "", "kind": c[0], "fit": c[1]}
        new.append(hit)
        if not dry:
            state["hits"][m["id"]] = hit
    for h in new:
        print(f"{'[dry] ' if dry else ''}{h['kind']:12} {h['received'][:10]} {h['from']}  {h['subject']}")
    if dry:
        return
    # Backfill run seeds the log silently; afterwards only NEW solicitations alert.
    if not first_run:
        for h in (x for x in new if x["kind"] == "solicitation"):
            note_hub(cfg, h); dm_jon(h)
    state["last_scan"] = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    build(cfg, state)
    save(STATE, state)


if __name__ == "__main__":
    main()
