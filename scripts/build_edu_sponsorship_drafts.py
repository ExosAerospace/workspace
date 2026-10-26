#!/usr/bin/env python3
"""Educational Payload Sponsorship outreach drafts (Jon 2026-10-26, MC #4050).

Reads drafts/edu-sponsorship/targets.json and writes one editable .docx with a
To/Cc/Subject + body per target. Drafts only: Jon reviews and sends from his
own Outlook. Nothing is sent from here.
"""
import json, os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor

WS = "/Users/exomac/.openclaw/workspace"
D = os.path.join(WS, "drafts", "edu-sponsorship")
STAMP = sys.argv[1] if len(sys.argv) > 1 else "2026-10-26"
OUT = os.path.join(D, f"Exos_Educational_Payload_Sponsorship_Outreach_Drafts_{STAMP}.docx")
LIZI = "lquinn@exosaero.com"
SUBJECT = "Sponsor a student payload to space (commitments by Nov 30)"


def first_name(t):
    n = (t.get("contact_name") or "").strip()
    return n.split()[0] if n and n.lower() not in ("not found", "n/a") else ""


def short_org(o):
    o = o.split(" / ")[0].split(" (")[0].strip()
    return "Rocket" if o.lower().startswith("rocket") else o


def addr(t):
    note = (t.get("contact_email") or "").lower()
    if "media" in note or "first robotics" in note or t["org"].startswith("BAE"):
        return ""  # not a grants inbox (Moody media contact, BAE FIRST-only), so use the portal/form
    m = re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", t.get("contact_email") or "")
    return m.group(0) if m else ""


def body(t):
    org = short_org(t["org"])
    hi = f"{first_name(t)}," if first_name(t) else f"Hello {org} team,"
    if t.get("category") == "foundation":
        why = (f"This sits squarely in what {org} already funds: hands-on STEM where students build something real. "
               "Grants can go straight to the student team's school, and we handle flight integration from there.")
    else:
        why = ("These are the engineers your teams will be recruiting in a few years, and they will have designed, "
               "built and flown real flight hardware before they graduate.")
    if org == "Rocket":
        why = ("Detroit students would design, build and fly real hardware to the edge of space and get it back with "
               "their data. That is the kind of career-pathway story Rocket already invests in.")
        ask = ("With a name like Rocket, this one felt obvious to us. I'd like to offer Rocket the Mission Title Sponsor "
               "level ($50,000): six student payloads, premium logo placement on the rocket and the flight named "
               "\"Rocket Student Flight.\" We would fill it with Detroit student teams.")
    else:
        ask = ("Levels run from $8,000, which funds one student payload, to $50,000 for Mission Title Sponsor with six "
               f"payloads. Anything above $25,000 puts {org}'s logo on the rocket. The one-page flyer is attached.")
    return "\n\n".join([
        hi,
        "I'm John Quinn, CEO of Exos Aerospace. We have student teams with science payloads built and ready to fly on "
        "BLK3, our next reusable rocket, in early 2027. What they don't have is the funding for the ride, and I'd "
        f"like to ask {org} to sponsor some of them.",
        ask,
        why,
        "We need commitments by November 30 to onboard these payloads for the flight. Lizi Quinn on our team (cc'd) "
        "is matching sponsors to student teams and can pair you with a school, region or science area that matters to you.",
        "Would you have 15 minutes in the next week or two?",
        "John",
    ])


def main():
    targets = json.load(open(os.path.join(D, "targets.json")))
    if isinstance(targets, dict):
        targets = targets.get("targets") or next(v for v in targets.values() if isinstance(v, list))
    targets.sort(key=lambda t: (t.get("priority") or 9, t.get("category", ""), t["org"]))
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    doc.add_heading("Educational Payload Sponsorship — outreach drafts", 1)
    doc.add_paragraph(f"DRAFTS for John's review and send (MC #4050). Attach the flyer to each. "
                      f"{len(targets)} targets, priority order. Where no verified email exists, the program's "
                      "portal or contact form is listed instead.")
    for t in targets:
        doc.add_page_break()
        doc.add_heading(f"{t['org']}  (P{t.get('priority', '?')}, {t.get('category', '')})", 2)
        to = addr(t)
        if not to:
            to = f"[address not found — use {t.get('contact_form_url') or t.get('program_url') or 'program site'}]"
        who = " / ".join(x for x in [t.get("contact_name"), t.get("contact_title")] if x and x.lower() != "not found")
        for k, v in [("To", f"{who} <{to}>" if who else to), ("Cc", LIZI), ("Subject", SUBJECT),
                     ("Attach", "Exos_Educational_Payload_Sponsorship_Flyer_2026-10-26.pptx (export to PDF from Canva once Lizi finalizes)")]:
            p = doc.add_paragraph(); r = p.add_run(f"{k}: "); r.bold = True; p.add_run(v)
        if t.get("how_to_apply"):
            p = doc.add_paragraph(); r = p.add_run(f"How they fund: {t['how_to_apply']}"); r.italic = True
            r.font.color.rgb = RGBColor(0x54, 0x54, 0x5D)
        for para in body(t).split("\n\n"):
            doc.add_paragraph(para)
        src = t.get("relevant_fact_with_source_url")
        if src:
            p = doc.add_paragraph(); r = p.add_run(f"Research note (not for the email): {src}"); r.italic = True
            r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x54, 0x54, 0x5D)
    doc.save(OUT)
    print("SAVED", OUT, len(targets))


if __name__ == "__main__":
    main()
