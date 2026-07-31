#!/usr/bin/env python3
"""Generate the sales & marketing-focused Exos Capability Statement (Grimes channel).
Draft for Jon's review in Series A Discord before it goes to GG Ventures.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GRAY = RGBColor(0x44, 0x44, 0x44)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
OUT = "/Users/exomac/.openclaw/workspace/drafts/EXOS_Capability_Statement_SalesMarketing_DRAFT.docx"

d = Document()
for s in d.sections:
    s.top_margin = Inches(0.6); s.bottom_margin = Inches(0.6)
    s.left_margin = Inches(0.7); s.right_margin = Inches(0.7)
st = d.styles["Normal"].font; st.name = "Calibri"; st.size = Pt(10.5)

def para(text, size=10.5, bold=False, color=None, after=4, before=0, align=None, italic=False):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align: p.alignment = align
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p

def head(text):
    para(text, size=11.5, bold=True, color=ACCENT, before=8, after=3)

def bullet(label, text):
    p = d.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + " — "); r.bold = True; r.font.size = Pt(10.5)
    r2 = p.add_run(text); r2.font.size = Pt(10.5)

# ---- Masthead ----
para("EXOS AEROSPACE", size=22, bold=True, color=NAVY, after=0)
para("Capability Statement  |  Flight Services & Payload Solutions",
     size=11, bold=True, color=GRAY, after=2)
para("One of only three FAA-licensed reusable rocket launch providers in the United States.",
     size=10.5, italic=True, color=ACCENT, after=8)

# ---- Positioning line ----
para("Flight-proven, reusable suborbital access that de-risks your hardware before it "
     "reaches orbit — faster iteration, lower program risk, real flight data you can act on.",
     size=10.5, after=8)

# ---- What we offer ----
head("SERVICES WE PROVIDE")
bullet("Payload Flight Services — SPACEavailable…™",
       "Reserve payload space on upcoming BLK3 suborbital flights. Microgravity exposure, "
       "high-altitude environments, and payload recovery — for research, biotech, and technology validation.")
bullet("Wet-Lease Flight Test",
       "We operate and maintain the vehicle; you bring the payload or system. Turnkey execution, "
       "minimal integration burden, and rapid campaign timelines for commercial and defense customers.")
bullet("Hardware De-Risk & Qualification",
       "Validate components and subsystems in real flight (boost loads under 4g, ~7g transient) before "
       "committing to costly orbital or full-scale programs.")
bullet("Reusable Platform Access & Recovery",
       "Repeatable, cadence-focused operations with payload recovery and operational learning — built "
       "for iteration, not one-off launches.")
bullet("Custom Test Campaigns",
       "Modular payload bay and tailored flight profiles, with minimal launch-site infrastructure "
       "requirements — pad plus power/comms; our mobile command capability brings the rest.")

# ---- Why Exos ----
head("WHY EXOS")
bullet("Currently selling",
       "Exos is an operator actively offering suborbital payload and flight-test services today — "
       "not a future promise.")
bullet("Flight heritage since 2018",
       "SARGE flown four times, demonstrating repeatable reusable suborbital operations.")
bullet("BLK3 next-generation vehicle",
       "Developed concept-to-hover in 12 months under a USAF Phase II SBIR — 25% larger and 29% lighter "
       "via integrated pressurized structures — and FAA-licensed (2026).")
bullet("De-risk mission",
       "We make orbital, defense, and in-space programs cheaper and faster by testing first.")
bullet("Advanced structures depth",
       "Through its ownership position in Scorpius Space Launch Company, Exos brings Type V linerless "
       "composite pressurized-structure heritage, including hardware flown to the lunar surface.")

# ---- Proof ----
head("PROOF POINTS")
bullet("Flight record", "SARGE: four reusable suborbital flights. BLK3: concept-to-hover in 12 months, FAA-licensed.")
bullet("Government performance", "NASA and DoD contract execution; under DoD contract to build the flight "
       "vehicle for a government-furnished LOX/methane Rotating Detonation Rocket Engine (RDRE) flight demonstration.")
bullet("Competition", "Northrop Grumman Lunar Lander Challenge — 1st place (Round 1), 2nd place (Round 2).")
bullet("Team heritage", "Hundreds of engines built (including LOX/methane) across more than a dozen flying "
       "vehicles, with reusable systems proven through hundreds of successful in-flight relights.")

# ---- Who it's for ----
head("WHO WE SERVE")
para("Orbital launch & satellite companies qualifying hardware  •  Defense programs needing test cadence "
     "(hypersonics, reentry, sensors)  •  Research & biotech needing microgravity and recovery  •  "
     "In-space infrastructure developers (cryogenics, refueling, structures).", size=10.5, after=8)

# ---- Contact ----
head("LET'S TALK")
para("John M. Quinn  —  Cofounder & CEO", size=10.5, bold=True, after=1)
para("jquinn@exosaero.com   |   www.exosaero.com   |   Greenville, TX", size=10.5, after=3)
para("CAGE 7EMY6  •  UEI CK24BPZGF854  •  NAICS 336414 / 336415 / 541715  •  "
     "PSC 1725 / 1340 / 2645 / 1675 / AR15", size=8.5, color=GRAY, after=0)

d.save(OUT)
print("wrote", OUT)
