#!/usr/bin/env python3
"""Convert the monthly investor report markdown to a formatted .docx.

Builds on the company letterhead (assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx)
so the banner header/footer are inherited. That template ships with only the Normal
style, so headings/lists/tables are rendered manually rather than via named styles.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
OUT = sys.argv[2]

NAVY = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x0B, 0x5C, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LETTERHEAD = "assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx"

if os.path.exists(LETTERHEAD):
    doc = Document(LETTERHEAD)
    # Strip the template's empty body (header/footer + section props stay intact).
    body = doc.element.body
    for el in list(body):
        if el.tag.endswith("}p") or el.tag.endswith("}tbl"):
            body.remove(el)
    sec = doc.sections[0]
    sec.top_margin = Inches(1.5)      # clear the banner
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
else:
    doc = Document()
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], m, Inches(0.9))

base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(10.5)

def add_runs(par, text):
    """Parse **bold** segments into runs."""
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = par.add_run(seg[2:-2]); r.bold = True
        else:
            par.add_run(seg)

def add_heading(level, txt):
    sizes = {1: 15, 2: 12.5, 3: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level <= 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(txt)
    r.bold = True
    r.font.size = Pt(sizes.get(level, 11))
    r.font.color.rgb = NAVY if level <= 2 else ACCENT
    return p

def add_bullet(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.17)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("•  ")
    add_runs(p, text)

def add_number(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(f"{num}.  ")
    add_runs(p, text)

def _shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)

def _borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "BFBFBF")
        b.append(e)
    tblPr.append(b)

lines = open(SRC, encoding="utf-8").read().splitlines()
n = len(lines)
i = 0
in_title_block = True

while i < n:
    line = lines[i].rstrip()

    if line.strip() == "[[PAGEBREAK]]":
        doc.add_page_break()
        i += 1
        continue

    if re.match(r"^[━─—_]{5,}$", line) or line.strip() == "---":
        in_title_block = False
        i += 1
        continue

    if in_title_block and line and not line.startswith("#"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.bold = True
        if i == 0:
            r.font.size = Pt(16); r.font.color.rgb = NAVY
        else:
            r.font.size = Pt(11); r.font.color.rgb = ACCENT
        i += 1
        continue

    if not line.strip():
        i += 1
        continue

    m = re.match(r"^(#{1,6})\s+(.*)$", line)
    if m:
        in_title_block = False
        add_heading(len(m.group(1)), m.group(2).strip())
        i += 1
        continue

    if line.lstrip().startswith("|"):
        tbl_lines = []
        while i < n and lines[i].lstrip().startswith("|"):
            tbl_lines.append(lines[i].strip())
            i += 1
        rows = [r for r in tbl_lines if not re.match(r"^\|[\s:|-]+\|$", r)]
        cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
        if cells:
            ncol = max(len(r) for r in cells)
            t = doc.add_table(rows=0, cols=ncol)
            t.autofit = True
            _borders(t)
            for ri, row in enumerate(cells):
                crow = t.add_row().cells
                for ci in range(ncol):
                    cp = crow[ci].paragraphs[0]
                    add_runs(cp, row[ci] if ci < len(row) else "")
                    if ri == 0:
                        _shade(crow[ci], "1F2A44")
                        for rr in cp.runs:
                            rr.bold = True; rr.font.color.rgb = WHITE
        continue

    mb = re.match(r"^\s*[-*]\s+(.*)$", line)
    if mb:
        add_bullet(mb.group(1))
        i += 1
        continue

    mn = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
    if mn:
        add_number(mn.group(1), mn.group(2))
        i += 1
        continue

    p = doc.add_paragraph()
    add_runs(p, line)
    i += 1

doc.save(OUT)
print("WROTE", OUT)
