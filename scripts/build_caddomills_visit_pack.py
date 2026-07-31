#!/usr/bin/env python3
"""Build the Grimes/Segreto/Caddo Mills visit pack:
  1. DD Dossier (internal)
  2. Visit Playbook / talking points (internal)
  3. Exos Target Facility Profile (leave-behind, on PRIMARY letterhead)
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

WS = "/Users/exomac/.openclaw/workspace"
DRAFTS = os.path.join(WS, "drafts")
LETTERHEAD = os.path.join(WS, "assets", "EXOS_Letterhead_Greenville_TX_PRIMARY.docx")
NAVY = RGBColor(0x0B, 0x22, 0x40)
GREY = RGBColor(0x44, 0x44, 0x44)

def H(doc, text, size=15, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY
    return p

def SUB(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = GREY
    return p

def body(doc, text, size=10.5):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    _emit(p, text, size); return p

def bullet(doc, text, size=10.5, level=0):
    try:
        p = doc.add_paragraph(style="List Bullet")
    except KeyError:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
        text = "•  " + text
    p.paragraph_format.space_after = Pt(2)
    if level: p.paragraph_format.left_indent = Inches(0.5 + 0.25*level)
    _emit(p, text, size); return p

def _emit(p, text, size=10.5):
    import re
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        r = p.add_run(seg); r.font.size = Pt(size)
        if i % 2: r.bold = True

def title_block(doc, title, subtitle):
    p = doc.add_paragraph(); r = p.add_run(title)
    r.bold = True; r.font.size = Pt(19); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph(); r2 = p2.add_run(subtitle)
    r2.italic = True; r2.font.size = Pt(10.5); r2.font.color.rgb = GREY
    doc.add_paragraph()

# ============================================================= DOC 1: DOSSIER
def dossier():
    doc = Document()
    title_block(doc, "Due Diligence Dossier — Grimes / Segreto / Caddo Mills",
        "INTERNAL. Prepared for John Quinn & Neil Milburn. Series A / facility track. "
        "Sources: public web, corporate registries, GG Ventures LLC and Capstone Commercial websites, "
        "City of Caddo Mills / EDC materials, July 2026. Unverified items are flagged.")

    H(doc, "1. Bottom Line Up Front")
    body(doc, "Gerald Grimes (**GG Ventures LLC**) and Daniel Segreto (**Capstone Commercial**) are an "
        "industrial-real-estate + private-capital pairing out of Dallas. This is, at its core, a **build-to-suit / "
        "site-and-capital assembly play** anchored on **Caddo Mills** (Hunt County, ~10 minutes west of our Greenville "
        "HQ on I-30). Grimes brings industrial development and private capital; Segreto brings DFW site selection, land, "
        "and municipal entitlement relationships; the undisclosed **'principal'** is the anchor money/tenant/developer "
        "they intend to introduce to the City of Caddo Mills.")
    body(doc, "**Why it matters to Exos:** it maps directly onto the footprint Phil described, a facility to roll a "
        "vehicle out, test, and roll it back in, with expanded composites production and larger winding equipment for "
        "tanks and rocket structures. A build-to-suit lease keeps that capex off the Series A and preserves cash for "
        "flight hardware. Type B EDC incentives and airport-adjacent industrial land are real levers here.")
    body(doc, "**What to keep separate:** the real-estate opportunity (near-term, credible) from any suggestion that "
        "they raise Series A capital for us (compliance-gated, unproven fit). Treat as two tracks. Details in §6.")

    H(doc, "2. Gerald Grimes — GG Ventures LLC")
    bullet(doc, "**Firm:** GG Ventures LLC, a self-described *private investment, warehouse management, and logistics "
        "consulting firm* focused on industrial real estate and supply-chain / 3PL operations.")
    bullet(doc, "**HQ:** 10500 Metric Dr., Ste 119, Dallas, TX 75243  |  +1 972-345-3952  |  gerald@ggventuresllc.com")
    bullet(doc, "**Capabilities claimed:** industrial real estate acquisition and asset management; Class A warehouse / "
        "distribution; 3PL operations; facility sourcing across a network of 1,500+ warehouse locations in North "
        "America; **private capital and fund management**; 20+ years' experience.")
    bullet(doc, "**Concrete, verifiable project:** developing **Building 1 at the Garland Innovation Center — 143,537 SF**, "
        "targeting e-commerce, manufacturing, and medical distribution. Confirms they actually develop / place capital "
        "into Class A industrial, not just broker it.")
    bullet(doc, "**Person:** Gerald Grimes appears tied to multiple Dallas-area entities over ~20 years (e.g., GSS "
        "Solutions LLC, Unlimited Solutions LLC). [DD GAP] No public aerospace or venture-equity track record found.")

    H(doc, "3. Daniel Segreto — Capstone Commercial Real Estate Group")
    bullet(doc, "**Role:** Industrial & Flex Specialist / Vice President at Capstone Commercial (Dallas). Investment "
        "sales plus landlord-representation leasing. Direct: 214.566.9052  |  dsegreto@capstonecommercial.com")
    bullet(doc, "**Background:** in industrial/flex CRE since 2009; prior NAI Robert Lynn; ~12 years at Ford Motor "
        "Company parts distribution (logistics/distribution fluency). Bilingual EN/ES.")
    bullet(doc, "**Municipal savvy:** served on the City of Little Elm Planning & Zoning / Board of Adjustments "
        "(2015–2019). Useful signal, he understands entitlements, zoning, and how to work a city on a site.")
    bullet(doc, "**Firm:** Capstone Commercial, Dallas CRE firm since 2000; industrial/office/land across DFW; millions "
        "of SF leased/sold. Co-founders **Steve Burris, CCIM** (Managing Broker/President) and **Sean Porter** "
        "(Principal — Strategy, Acquisitions & Investment Sales).")

    H(doc, "4. The 'Principal' — the key unknown  [DD GAP]")
    body(doc, "Grimes/Segreto reference bringing 'our principal' to meet the City of Caddo Mills. Identity, role, and "
        "balance sheet are undisclosed. This is the single most important thing to qualify before we invest time or "
        "expose anything. Possible profiles: (a) a build-to-suit developer/landlord who would fund and own the "
        "facility and lease it to Exos; (b) an anchor capital source / family office; (c) a distribution/industrial "
        "tenant unrelated to us who anchors the park. Our posture and disclosures differ materially depending on which. "
        "**Ask directly and early (see Playbook §Questions).**")

    H(doc, "5. Caddo Mills — the site opportunity")
    bullet(doc, "**Location:** Hunt County, on the I-30 corridor ~10 min west of Exos' Greenville HQ, keeps our "
        "workforce and preserves the Greenville test/hover assets while adding room to scale.")
    bullet(doc, "**Growth:** one of Texas' fastest-growing cities per capita — roughly 1,504 (2020) to ~5,969 (2025), "
        "+~297%. Avg household income >$90K. Aggressive **Type B EDC** (jobs-focused; can fund land/infrastructure).")
    bullet(doc, "**EDC board (2026):** Derrek Ruff (Pres), Rosalyn Kelly (VP), Lane Parrish (Treas), Jeb Payne, "
        "Nick Meno, Justin Rudd; Cody Hawkins (Ex-Officio). A November 2025 council item considered appointing a new "
        "CMEDC Director. [VERIFY current director before the city meeting.]")
    bullet(doc, "**Caddo Mills Municipal Airport (7F3):** city-owned, WWII-origin field, **4,000 x 150 ft runway**, "
        "slated for expansion with light-industrial development, hangars, corporate-aircraft and distribution use "
        "('Airport FBO Initiative'). Airfield adjacency is a genuine differentiator for our logistics and future ops.")
    bullet(doc, "**Momentum:** master-planned communities (Trailstone, Caddo Downs, The Mills), the 60+ acre Pit Stop "
        "Entertainment District, Western Regional Wastewater plant, FM 36 realignment, and I-30/FM 1565 upgrades, "
        "meaning utilities and road capacity are being built out now.")

    H(doc, "6. Strategic fit & the two-track framing")
    SUB(doc, "Track A — Facility / real estate (credible, near-term)")
    bullet(doc, "GG Ventures + Capstone can plausibly assemble land, entitle it, and deliver a **build-to-suit** "
        "high-bay facility, then lease to Exos. That is exactly their demonstrated competence (Garland Innovation "
        "Center). Pair with Type B EDC incentives and this is a real path to the footprint we need without a large "
        "Exos capex outlay.")
    bullet(doc, "Our leverage: Exos is a marquee, job-creating, one-of-three FAA-licensed reusable-rocket anchor "
        "tenant, exactly what a growth city's EDC and a spec developer want. **We are the demand they are selling.** "
        "That is worth incentives and favorable lease economics.")
    SUB(doc, "Track B — Capital / investors (compliance-gated, unproven)")
    bullet(doc, "They also pitch 'investment, strategic partnerships, and capital formation.' GG Ventures does private "
        "capital / fund management for **real estate**, not aerospace growth equity. Do **not** assume a Series A lead. "
        "If they propose to introduce or raise investor capital for Exos, that implicates broker-dealer / finder rules.")
    bullet(doc, "**Standing Exos rules apply:** no retainers and no success fees on capital (SOX / broker-dealer / "
        "Reg D). Any capital-intro arrangement routes through securities counsel before anything is agreed. Warm "
        "investor intros are welcome; a paid finder arrangement is not.")

    H(doc, "7. Risks & red flags")
    bullet(doc, "**Undisclosed principal (§4)** — do not expose facility internals, roadmap, or controlled data until "
        "identity and intent are known and US-person status is confirmed.")
    bullet(doc, "**Export control / ITAR + CUI.** Winding equipment, rocket structures, and test operations touch "
        "ITAR-controlled technical data and our CMMC/CUI boundary. The mutual NDA is in place, but an NDA is not a "
        "TAA, keep site discussions to non-controlled facility parameters (square footage, power, doors, crane, "
        "utilities). No technical data, CAD/STEP, or process detail. Confirm all their attendees are US persons.")
    bullet(doc, "**Broker-as-finder risk (§6B).** Keep a CRE broker from acting as an unregistered securities finder.")
    bullet(doc, "**Conflated agendas.** They blend real estate and capital in one motion; we separate them so the "
        "facility deal is not held hostage to an investor conversation (or vice-versa).")
    bullet(doc, "**Entity/standing unverified.** [DD GAP] Confirm GG Ventures LLC good standing (TX SOS), Grimes' "
        "principal role, and any litigation/liens before term-sheet stage.")

    H(doc, "8. Open DD checklist (before / at the meeting)")
    for item in [
        "Who is the principal? Name, entity, role (developer / capital / tenant), and balance-sheet capacity.",
        "Are all their principals and the anchor US persons? (export-control gate).",
        "GG Ventures LLC — TX SOS standing, ownership, and Grimes' formal title; litigation/lien check.",
        "Track record: completed build-to-suit / industrial deliveries beyond Garland Innovation Center.",
        "Caddo Mills site specifics: parcel(s), acreage, zoning, utilities/power capacity, airport-parcel control.",
        "Current CMEDC Director and the incentive tools available (Type B) for a manufacturing anchor.",
        "Their economic model: do they want Exos as a build-to-suit tenant, an equity story, or both?",
        "Timeline: when is the principal's City of Caddo Mills visit, and what do they want from Exos that day?",
    ]:
        bullet(doc, item)

    H(doc, "9. Recommended posture")
    body(doc, "Take the meeting, lead with the **facility/real-estate track**, and treat the capital conversation as "
        "exploratory and compliance-gated. Qualify the principal hard. Offer Exos as the anchor demand and let the "
        "EDC + developer compete to land us. Keep all technical/controlled detail off the table until we know exactly "
        "who is in the room. Concrete asks and talking points are in the companion **Visit Playbook**.")

    out = os.path.join(DRAFTS, "Grimes_Segreto_CaddoMills_DD_Dossier.docx")
    doc.save(out); return out

# ============================================================ DOC 2: PLAYBOOK
def playbook():
    doc = Document()
    title_block(doc, "Visit Playbook & Talking Points — Grimes / Segreto Meeting",
        "INTERNAL. Next-week meeting ahead of the principal's City of Caddo Mills visit. "
        "Attendees (Exos): John Quinn (lead), Neil Milburn (ops/facility), Phil Eaton (host/site). Read with the DD Dossier.")

    H(doc, "Meeting objective")
    body(doc, "Qualify the principal and the real opportunity, and position Exos as the anchor tenant a growth city and "
        "a spec developer want to land, without exposing controlled technical detail or making any capital-side "
        "commitment. Leave with (1) the principal's identity and intent, (2) a date for the joint City of Caddo Mills "
        "visit, and (3) agreement on a facility-first, two-track process.")

    H(doc, "Roles")
    bullet(doc, "**John Quinn** — leads; capital strategy, vision, the ask. Controls the capital conversation and keeps it exploratory.")
    bullet(doc, "**Neil Milburn** — facility/operational requirements; sanity-checks any site against what a rocket program actually needs.")
    bullet(doc, "**Phil Eaton** — continuity as host; relationship warmth already established.")

    H(doc, "Agenda (60–75 min)")
    bullet(doc, "Welcome + recap since the visit and the executed mutual NDA (5 min).")
    bullet(doc, "Their side: GG Ventures / Capstone model, the Caddo Mills thesis, and **who the principal is** (15).")
    bullet(doc, "Exos in brief: who we are, heritage, trajectory, and why the footprint matters now (10).")
    bullet(doc, "Facility track: what Exos needs, build-to-suit vs. lease, EDC incentives, airport adjacency (20).")
    bullet(doc, "Capital track: exploratory only, how they work, what's compliant on our side (10).")
    bullet(doc, "Next steps: the city-visit date and a short written follow-up (5).")

    H(doc, "Exos talking points (approved framing)")
    bullet(doc, "**One of only three FAA-licensed reusable rocket launch providers in the U.S.** — and the one currently "
        "selling suborbital payload services.")
    bullet(doc, "**Flight heritage since 2018** — Exos is the commercial successor to Armadillo Aerospace; a proven, "
        "reusable platform, not a pre-revenue concept. (Never frame as 'first flight in 2027.')")
    bullet(doc, "**Physical-AI / re-industrialization thesis** — Exos builds the reusable platforms and mass that the "
        "$20T physical-world economy needs in space; a de-risking test platform for orbital programs.")
    bullet(doc, "**Scorpius heritage** — Type-5 cryogenic tank technology with lunar flight heritage (flew on Intuitive "
        "Machines Odysseus and Athena). Speaks to advanced composites and the winding capability we're scaling.")
    bullet(doc, "**Next phase** — Series A capital formation and a facility built to scale production and test cadence; "
        "Caddo Mills keeps our Greenville workforce while giving us room to grow.")

    H(doc, "Facility requirements to convey (non-sensitive only)")
    bullet(doc, "High-bay integration hall sized to roll a vehicle **out to test and back in** under cover, single-flow.")
    bullet(doc, "Expanded **composites production** with room for **larger filament-winding equipment** for tanks and "
        "rocket structures.")
    bullet(doc, "Test infrastructure / stand area with appropriate standoff, power, and utilities.")
    bullet(doc, "Overhead crane / high door clearance; heavy power service; water and wastewater capacity.")
    bullet(doc, "Airfield adjacency (Caddo Mills 7F3) as a logistics and operations advantage.")
    bullet(doc, "**Preferred structure: build-to-suit / long-term lease** — keeps capex off the Series A. Frame this as our default.")

    H(doc, "Questions to ask THEM")
    bullet(doc, "**Who is the principal** — name, entity, and are they the developer, the capital, or the anchor tenant?")
    bullet(doc, "What's the principal's balance sheet / how do they fund a build-to-suit of this size?")
    bullet(doc, "Are all your principals and the anchor **US persons**? (we'll explain the export-control reason plainly.)")
    bullet(doc, "What Caddo Mills parcels do you control or have optioned, and what's the zoning/utilities status?")
    bullet(doc, "What EDC incentives are realistically on the table for a manufacturing anchor, and who's the current CMEDC Director?")
    bullet(doc, "On capital: how do you work, and what exactly are you proposing on the investment side?")
    bullet(doc, "What's the date for the principal's City of Caddo Mills visit, and what do you want from Exos that day?")

    H(doc, "Landmines — do NOT")
    bullet(doc, "Do **not** share technical data, CAD/STEP, process detail, or roadmap specifics. Facility parameters only.")
    bullet(doc, "Do **not** agree to any retainer, success fee, or finder arrangement on capital. Route to securities counsel.")
    bullet(doc, "Do **not** let the facility deal get tied to an investor commitment (or vice-versa). Two separate tracks.")
    bullet(doc, "Do **not** commit to a site or timeline in the room; everything subject to Exos DD and board process.")
    bullet(doc, "Do **not** disclose facility internals to the principal until identity + US-person status are confirmed.")

    H(doc, "The ask / desired outcomes")
    bullet(doc, "Principal identified and qualified.")
    bullet(doc, "A firm date for the joint City of Caddo Mills visit, with Exos hosting the same day.")
    bullet(doc, "Agreement to a facility-first, two-track process with capital kept exploratory and compliant.")
    bullet(doc, "A short written recap from them (parcels, incentives concept, principal) so we can advance DD.")

    out = os.path.join(DRAFTS, "Grimes_Segreto_Visit_Playbook.docx")
    doc.save(out); return out

# =================================================== DOC 3: FACILITY (letterhead)
def facility():
    doc = Document(LETTERHEAD) if os.path.exists(LETTERHEAD) else Document()
    p = doc.add_paragraph(); r = p.add_run("Target Facility Profile — Hunt County / Caddo Mills Corridor")
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph(); r2 = p2.add_run(
        "Prepared for discussion with GG Ventures LLC and Capstone Commercial  |  July 2026  |  Non-proprietary")
    r2.italic = True; r2.font.size = Pt(10); r2.font.color.rgb = GREY
    doc.add_paragraph()

    body(doc, "Exos Aerospace is one of only three FAA-licensed reusable rocket launch providers in the United States, "
        "with flight heritage dating to 2018. As we enter our next phase of production scale-up and flight cadence, we "
        "are evaluating a facility in the Hunt County / Caddo Mills corridor that keeps our Greenville workforce while "
        "providing room to grow. Our strong preference is a build-to-suit or long-term lease structure.")

    H(doc, "Core requirements")
    bullet(doc, "**High-bay integration hall** sized to move a launch vehicle from the build floor out to test and back "
        "under cover in a single flow.")
    bullet(doc, "**Expanded composites production** with floor space and clear height for larger filament-winding "
        "equipment used to produce propellant tanks and rocket structures.")
    bullet(doc, "**Test area** with appropriate standoff, utilities, and power for integrated testing.")
    bullet(doc, "**Heavy utilities:** robust electrical service, water and wastewater capacity, and compressed-air / gas infrastructure.")
    bullet(doc, "**Material handling:** overhead crane capacity and tall door clearance for large structures and vehicles.")

    H(doc, "Site attributes we value")
    bullet(doc, "Proximity to our Greenville operations (workforce retention and shared assets).")
    bullet(doc, "I-30 corridor access for logistics and oversized transport.")
    bullet(doc, "Airfield adjacency (Caddo Mills Municipal Airport, 7F3) for operational and logistics flexibility.")
    bullet(doc, "A supportive municipal and economic-development environment able to move quickly.")

    H(doc, "What Exos brings")
    bullet(doc, "A marquee advanced-manufacturing anchor: skilled aerospace jobs, capital investment, and national visibility.")
    bullet(doc, "A proven, revenue-generating reusable platform, not a speculative venture.")
    bullet(doc, "A partner that, in the founders' words, moves quickly.")

    doc.add_paragraph()
    body(doc, "Facility parameters in this document are non-proprietary and provided for site-planning discussion only. "
        "Technical specifications are available under appropriate agreements.", size=9)

    out = os.path.join(DRAFTS, "Exos_Target_Facility_Profile_CaddoMills.docx")
    doc.save(out); return out

if __name__ == "__main__":
    os.makedirs(DRAFTS, exist_ok=True)
    for f in (dossier(), playbook(), facility()):
        print("wrote", f)
