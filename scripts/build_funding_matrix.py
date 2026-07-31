#!/usr/bin/env python3
"""Exos suborbital-science cross-agency FUNDING MATRIX.
Landscape one-pager (branded PNG+PDF) + editable Word table.
Program x Agency/Vehicle x $ ceiling x deadline x status. Figures verified 2026-07-22."""
import os
from PIL import Image, ImageDraw, ImageFont
W,H=2560,1980
NAVY=(9,20,42); GOLD=(232,170,60); ORANGE=(233,122,50); WHITE=(255,255,255)
LIGHT=(238,243,248); GREY=(206,217,231); PANEL=(19,40,70); PANEL_BD=(46,82,130)
ROWA=(16,32,58); ROWB=(21,40,70); BAND=(30,54,92)

def F(kind,size):
    paths={'reg':"/System/Library/Fonts/Supplemental/Arial.ttf",
           'bold':"/System/Library/Fonts/Supplemental/Arial Bold.ttf",
           'black':"/System/Library/Fonts/Supplemental/Arial Black.ttf"}
    p=paths[kind]
    if os.path.exists(p):
        try: return ImageFont.truetype(p,size)
        except: pass
    return ImageFont.truetype({'reg':"DejaVuSans.ttf",'bold':"DejaVuSans-Bold.ttf",'black':"DejaVuSans-Bold.ttf"}[kind],size)

def bg():
    img=Image.new("RGB",(W,H),NAVY); d=ImageDraw.Draw(img)
    for y in range(H):
        t=y/H; d.line([(0,y),(W,y)],fill=(int(9+8*t),int(20+14*t),int(42+22*t)))
    return img

def wrap(d,s,f,mw):
    out=[]
    for para in s.split("\n"):
        ln=""
        for w in para.split():
            t=(ln+" "+w).strip()
            if d.textlength(t,font=f)<=mw: ln=t
            else: out.append(ln); ln=w
        out.append(ln)
    return out

def txt_wrap(d,x,y,s,f,fill,mw,lh=32):
    yy=y
    for ln in wrap(d,s,f,mw): d.text((x,yy),ln,font=f,fill=fill); yy+=lh
    return yy

M=80
# columns: (label, width)
COLS=[("AGENCY / PROGRAM",560),("VEHICLE",300),("CEILING (Phase I / II)",420),("NEXT DEADLINE",380),("STATUS / NOTES",740)]
CX=[M]
for _,w in COLS: CX.append(CX[-1]+w)

SECTIONS=[
 ("MICROGRAVITY / BIOMEDICAL  —  BLK3 as the screen/de-risk platform (Purdue PI)", [
   ("NIH  —  STTR (PA-27-102)","STTR grant","~$314K / ~$2.1M","Sep 8, 2026","PRIMARY. Requires research-institution PI (>=30%). Fast-Track + Phase IIB open."),
   ("NASA ROSBio (Appendix D)","Research grant","~$200-400K/yr","Next appendix - confirm","Directly funds suborbital/parabolic flight biology. Cleanest bio money."),
   ("NASA Flight Opportunities / TechFlights","Flight-test grant","up to $750K","Annual - confirm next release","Buys the suborbital flight test itself (tech maturation)."),
   ("NSF  —  Eng. / Biomanufacturing","Research grant","varies","Program-specific","Cellular biomanufacturing + fluid physics in microgravity."),
   ("DoD  —  CDMRP / USAMRDC / DHA","Medical research","varies","Annual cycles","Aerospace medicine, medical countermeasures."),
   ("ISS National Lab","CSP pathway only","slot sales","No solicitation (dead)","Sell BLK3 screening slots to already-funded NIH/NCATS PIs."),
 ]),
 ("STRUCTURAL HEALTH MONITORING (FBG)  —  Purdue CMSC / Dr. Dianyun Zhang", [
   ("NASA SBIR  —  AERO.4.S26B","SBIR","$225K / ~$750-850K","FY26-27 BAA - confirm window","LOS IN HAND (Zhang). Flight-measurement technologies."),
   ("Navy  —  ONR / NAVAIR","BAA / white paper","varies","Rolling BAA","Fiber-optic SHM flight test - relationship already warm."),
   ("AFRL  —  RXC","Open BAA (FA8650-18-S-5010)","varies","Open thru Sep 20, 2027","Composite materials SHM; long open window."),
   ("FAA","Research","varies","Program-specific","Reusable composite certification data need."),
   ("Army DEVCOM / NIST / DOE","Various","varies","Various","Secondary pull-through on the same dataset."),
 ]),
 ("ATMOSPHERIC SAMPLING (SACSS 40-100 km)  —  Purdue / Dr. Steven Collicott", [
   ("NOAA CPO  —  ERB / AC4","Research grant","~$150-500K/yr","LOI ~Sep 2026 (confirm)","PRIMARY, science-led. Exos + Purdue Co-PI; leads with the 40-100 km gap."),
   ("NASA SBIR  —  INSTALG.2.S26A","SBIR","$225K / ~$750-850K","FY26-27 BAA - confirm window","LOS IN HAND (Collicott). Suborbital instruments & sensors."),
   ("NOAA SBIR","SBIR","~$190K Phase I","NOFO ~Nov 2026","~80% reuse of the NASA INSTALG package."),
   ("NSF  —  Geospace / Aeronomy","Research grant","varies","Program-specific","Mesosphere / lower-thermosphere science."),
   ("EPA / USSF / AFRL","Various","varies","Various","Same dataset; reentry-debris / space-domain angle."),
 ]),
]

img=bg(); d=ImageDraw.Draw(img)
# header
d.text((M,54),"EXOS",font=F('black',50),fill=WHITE)
d.text((M+225,64),"AEROSPACE",font=F('reg',42),fill=GOLD)
d.text((W-M,80),"SUBORBITAL-SCIENCE FUNDING MATRIX  |  Prepared 2026-07-22  |  Proprietary",font=F('reg',26),fill=GREY,anchor="ra")
d.line([(M,150),(W-M,150)],fill=PANEL_BD,width=3)
d.text((M,178),"Cross-Agency Funding Matrix",font=F('black',62),fill=WHITE)
d.text((M,262),"Three BLK3-platform programs, mapped to every agency, vehicle, ceiling, and deadline - sequenced for parallel submission.",font=F('reg',30),fill=GREY)
y=330
# column header
hh=56
d.rounded_rectangle([M,y,W-M,y+hh],radius=10,fill=(12,26,50),outline=PANEL_BD,width=2)
for i,(lab,w) in enumerate(COLS):
    d.text((CX[i]+18,y+15),lab,font=F('bold',25),fill=GOLD)
    if i>0: d.line([(CX[i],y),(CX[i],y+hh)],fill=PANEL_BD,width=1)
y+=hh+8

fcell=F('reg',25); fkey=F('bold',26); fnote=F('reg',24)
def rowheight(r):
    a=len(wrap(d,r[0],fkey,COLS[0][1]-30))
    n=len(wrap(d,r[4],fnote,COLS[4][1]-30))
    return max(a,n,1)*32+22

for stitle,rows in SECTIONS:
    # section band
    bh=46
    d.rounded_rectangle([M,y,W-M,y+bh],radius=8,fill=BAND,outline=PANEL_BD,width=2)
    d.text((M+18,y+9),stitle,font=F('bold',28),fill=WHITE)
    y+=bh+4
    for ri,r in enumerate(rows):
        rh=rowheight(r)
        d.rectangle([M,y,W-M,y+rh],fill=(ROWA if ri%2==0 else ROWB))
        # cells
        yy=txt_wrap(d,CX[0]+18,y+11,r[0],fkey,WHITE,COLS[0][1]-30)
        txt_wrap(d,CX[1]+18,y+11,r[1],fcell,GREY,COLS[1][1]-30)
        # ceiling emphasize
        ceil_fill=GOLD if ("$" in r[2] or "K" in r[2]) else GREY
        txt_wrap(d,CX[2]+18,y+11,r[2],F('bold',25),ceil_fill,COLS[2][1]-30)
        dl_fill=ORANGE if any(m in r[3] for m in ["Sep","Nov","2026","2027"]) else GREY
        txt_wrap(d,CX[3]+18,y+11,r[3],F('bold',24),dl_fill,COLS[3][1]-30)
        txt_wrap(d,CX[4]+18,y+11,r[4],fnote,GREY,COLS[4][1]-30)
        for i in range(1,len(COLS)):
            d.line([(CX[i],y),(CX[i],y+rh)],fill=(30,50,84),width=1)
        y+=rh
    y+=12

# sequencing box
by=y+4
d.rounded_rectangle([M,by,W-M,by+250],radius=14,fill=(15,28,52),outline=GOLD,width=4)
d.text((M+28,by+22),"RECOMMENDED SEQUENCING",font=F('bold',30),fill=GOLD)
seq=[("NOW  (Jul-Sep)","NIH STTR (Sep 8) - lock the Purdue PI first  ·  NASA SBIR AERO + INSTALG (LOS in hand)  ·  NOAA ERB/AC4 LOI (~Sep)"),
     ("NEXT  (Oct-Dec)","NOAA SBIR (~Nov)  ·  NOAA ERB/AC4 full proposal (~Dec)  ·  NASA ROSBio / TechFlights (confirm windows)"),
     ("PARALLEL / NON-DILUTIVE","Navy ONR-NAVAIR + AFRL RXC (open BAAs)  ·  sell ISS-NL screening slots to funded PIs  ·  NASA Flight Opportunities")]
yy=by+72
for lab,txt in seq:
    d.text((M+28,yy),lab,font=F('bold',26),fill=WHITE)
    d.text((M+430,yy),txt,font=F('reg',25),fill=GREY)
    yy+=56
# footer
d.rectangle([0,H-64,W,H],fill=(7,14,30))
d.text((M,H-48),"John Quinn, Cofounder & CEO   |   jquinn@exosaero.com   |   Ceilings are total-cost caps; verify each NOFO/BAA at release. STTR/SBIR splits: STTR SB>=40% / RI>=30%.",font=F('reg',22),fill=GREY)

os.makedirs("drafts",exist_ok=True)
png="drafts/Exos_Suborbital_Funding_Matrix.png"; pdf="drafts/Exos_Suborbital_Funding_Matrix.pdf"
img.save(png); img.convert("RGB").save(pdf,resolution=200.0)
print("BUILT",png)

# ---------- editable Word table ----------
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc=Document()
sec=doc.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=Inches(14),Inches(8.5)
for m in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(sec,m,Inches(0.4))
h=doc.add_heading("Exos Suborbital-Science Cross-Agency Funding Matrix",level=0)
doc.add_paragraph("Prepared 2026-07-22 · Proprietary. Ceilings are total-cost caps; verify each NOFO/BAA at release. STTR/SBIR cost split: small business >=40% (STTR) / >=67% (SBIR Phase I); research institution >=30% (STTR).")
def shade(cell,hexc):
    t=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:fill'),hexc); t.append(s)
tbl=doc.add_table(rows=1,cols=5); tbl.style='Light Grid Accent 1'
hdr=tbl.rows[0].cells
for i,(lab,_) in enumerate(COLS):
    hdr[i].text=lab;
    for p in hdr[i].paragraphs:
        for r_ in p.runs: r_.font.bold=True; r_.font.size=Pt(9)
for stitle,rows in SECTIONS:
    sc=tbl.add_row().cells; a=sc[0].merge(sc[1]).merge(sc[2]).merge(sc[3]).merge(sc[4])
    a.text=stitle
    for p in a.paragraphs:
        for r_ in p.runs: r_.font.bold=True; r_.font.size=Pt(10); r_.font.color.rgb=RGBColor(0x1F,0x3B,0x66)
    shade(a,"DCE6F2")
    for r in rows:
        c=tbl.add_row().cells
        for i,val in enumerate(r):
            c[i].text=val
            for p in c[i].paragraphs:
                for r_ in p.runs: r_.font.size=Pt(8.5)
doc.add_paragraph("")
p=doc.add_paragraph(); run=p.add_run("Recommended sequencing"); run.bold=True; run.font.size=Pt(11)
for lab,txt in seq:
    pp=doc.add_paragraph(style='List Bullet'); r1=pp.add_run(lab+":  "); r1.bold=True; r1.font.size=Pt(9.5); r2=pp.add_run(txt); r2.font.size=Pt(9.5)
docx="drafts/Exos_Suborbital_Funding_Matrix.docx"; doc.save(docx)
print("BUILT",docx)
