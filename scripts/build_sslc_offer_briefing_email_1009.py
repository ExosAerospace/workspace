#!/usr/bin/env python3
"""Build SSLC_Offer_Briefing_Email_DRAFT_2026-10-09.docx
John Quinn → Jim Wertz and the SSLC Board. DRAFT — not sent.
Subject: Exos offer for SSLC — summary
10/9 update numbers: ① $4,360K | ② $5,360K | ③ $6.3M–$9.5M | ④ $12M–$19M (2031 projection)
"""
import os, shutil, subprocess
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

WS   = "/Users/exomac/.openclaw/workspace"
TMPL = os.path.join(WS, "assets", "EXOS_Letterhead_Greenville_TX_PRIMARY.docx")
RAW  = os.path.join(WS, "drafts", "_sslc_briefing_email_raw.docx")
OUT  = os.path.join(WS, "drafts", "SSLC_Offer_Briefing_Email_DRAFT_2026-10-09.docx")

NAVY  = RGBColor(0x1F, 0x3A, 0x68)
BLACK = RGBColor(0x00, 0x00, 0x00)
RED   = RGBColor(0xC0, 0x00, 0x00)
GREY  = RGBColor(0x60, 0x60, 0x60)

shutil.copy(TMPL, RAW)
doc = Document(RAW)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)

def P(text="", bold=False, size=None, color=None, italic=False,
      sa=3, sb=0, indent=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(sa)
    p.paragraph_format.space_before = Pt(sb)
    if indent: p.paragraph_format.left_indent = Inches(indent)
    if align:  p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold = bold; r.italic = italic
        if size:  r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run("•  " + text)
    r.font.size = Pt(10.5)
    return p

# ── Draft / not-sent banner ───────────────────────────────────────────────────
P("DRAFT — NOT SENT  ·  Internal use only  ·  John Quinn / Exos Aerospace",
  bold=True, size=9, color=RED, align=WD_ALIGN_PARAGRAPH.CENTER, sa=4)

# ── Memo header ───────────────────────────────────────────────────────────────
P("To:", bold=True, size=10.5, sa=1)
P("Jim Wertz and members of the SSLC Board of Directors", size=10.5, sa=3, indent=0.5)

P("From:", bold=True, size=10.5, sa=1)
P("John M. Quinn, Cofounder & CEO, Exos Aerospace Systems & Technologies, Inc.", size=10.5, sa=3, indent=0.5)

P("Date:", bold=True, size=10.5, sa=1)
P("October 9, 2026", size=10.5, sa=3, indent=0.5)

P("Subject:", bold=True, size=10.5, sa=1)
P("Exos offer for SSLC — summary", size=10.5, sa=6, indent=0.5)

# ── Horizontal rule (via border paragraph) ────────────────────────────────────
# Use a light horizontal separator
hr = doc.add_paragraph()
hr.paragraph_format.space_after = Pt(6)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
pPr = hr._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single')
bottom.set(qn('w:sz'), '6')
bottom.set(qn('w:space'), '1')
bottom.set(qn('w:color'), '1F3A68')
pBdr.append(bottom)
pPr.append(pBdr)

# ── Body ─────────────────────────────────────────────────────────────────────
P("Jim and members of the Board,", size=10.5, sa=4)

P("Attached is Exos's offer to acquire SSLC's assets under Section 3 of our May 2025 Asset "
  "Acquisition Agreement. In short:", size=10.5, sa=4)

bullet("At close: $4.36M in cash, transaction costs, credit for Exos's existing stake, "
       "payment of SSLC's debts at face, and a working-capital bridge — plus $1.0M of Exos "
       "stock (EAI units at the current 409A value), for $5.36M at financing close. "
       "The EAI units are restricted at issuance (Rule 504 Reg D) and are structured to provide "
       "part of the public float, becoming tradeable at the RTO listing subject to the same "
       "lockup as Exos principals.")

bullet("Over five years: earn-outs on IM tank revenue, five milestones (the first at $5.1M of "
       "FY2026-27 revenue), and additional Exos stock, bringing total value to $6.3M–$9.5M.")

bullet("Upside: if holders keep their Exos stock through 2031, the projected total is $12M–$19M. "
       "This is a projection, not a guarantee.")

P("", sa=4)  # spacer

P("The up-front price reflects what the numbers support today; the earn-outs are how that value "
  "grows as goals are met. With $198K in cash and a $5.1M year requiring double last quarter's "
  "pace, this offer removes the funding risk and gives SSLC holders a direct share of the upside.",
  size=10.5, sa=6)

P("SSLC employees. Jim, to your question: Exos shall offer employment to SSLC employees, including "
  "management, support, engineering, and other technical and manufacturing employees, on terms no "
  "less favorable than their current employment terms with SSLC (including compensation and "
  "benefits), subject to each employee satisfying Exos's standard hiring policies and Exos's "
  "completion of confirmatory due diligence. In addition, Exos will develop an option pool for key "
  "SSLC employees, to be mutually agreed in the Definitive Agreement.",
  size=10.5, sa=6)

P("We ask the Board to confirm §3 exclusivity so we can proceed to close.", size=10.5, sa=8)

# ── Signature ─────────────────────────────────────────────────────────────────
P("John M. Quinn", bold=True, size=10.5, sa=1)
P("Cofounder & CEO, Exos Aerospace Systems & Technologies, Inc.", size=10.5, sa=1)

# ── Footer note ───────────────────────────────────────────────────────────────
P("", sa=8)
P("Detail in model: EXOS_SSLC_OptionB_Acquisition_Model_2026-10-07.xlsx  ·  "
  "Offer summary: Exos_SSLC_OptionB_Offer_Summary_2026-10-07.pdf  ·  "
  "MC task: #2414",
  size=8, italic=True, color=GREY, sa=2)

P("DRAFT — INTERNAL — NEGOTIATING TERMS — NOT SENT — John Quinn, CEO, Exos Aerospace — 2026-10-09",
  bold=True, size=8, color=RED, align=WD_ALIGN_PARAGRAPH.CENTER, sa=0)

# ── save raw, apply letterhead ────────────────────────────────────────────────
doc.save(RAW)
subprocess.run(
    ["/usr/bin/python3",
     os.path.join(WS, "scripts/apply_exos_letterhead.py"), RAW, OUT],
    check=True
)
os.remove(RAW)
print(f"DOCX: {OUT}")
print("DRAFT — NOT SENT — briefing email complete.")
