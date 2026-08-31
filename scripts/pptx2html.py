#!/usr/bin/env python3
"""pptx2html.py <src.pptx> <out.html>

Renders a PowerPoint deck to a clean, print-friendly HTML *content* view — one
block per slide (slide number, title, body text, tables, and speaker notes).

This is deliberately a CONTENT view, not a pixel-perfect render: LibreOffice/soffice
hangs in this environment, so Mission Control's /viewpdf pipeline converts HTML→PDF
via headless Chrome (see server.js ensurePdfRendition). This script is the pptx
equivalent of scripts/docx2html.py, letting a deck attached to a Decision be reviewed
inline (approve ordering / figures / narrative) without downloading PowerPoint.

Runs on the system python (/usr/bin/python3) where python-pptx is installed.
"""
import sys
import html


def esc(s):
    return html.escape(str(s if s is not None else ""))


def shape_text(shape):
    """Best-effort text extraction from a single shape (incl. grouped shapes)."""
    out = []
    try:
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # GROUP
            for sub in shape.shapes:
                out.extend(shape_text(sub))
            return out
    except Exception:
        pass
    if getattr(shape, "has_table", False):
        try:
            rows = []
            for r in shape.table.rows:
                cells = [esc(c.text.strip()) for c in r.cells]
                rows.append("<tr>" + "".join("<td>" + c + "</td>" for c in cells) + "</tr>")
            out.append("<table class='t'>" + "".join(rows) + "</table>")
        except Exception:
            pass
        return out
    if getattr(shape, "has_text_frame", False):
        for para in shape.text_frame.paragraphs:
            txt = "".join(run.text for run in para.runs).strip()
            if not txt:
                txt = (para.text or "").strip()
            if txt:
                lvl = getattr(para, "level", 0) or 0
                out.append(("bullet", lvl, txt))
    return out


def render(src, out_path):
    from pptx import Presentation
    prs = Presentation(src)
    slides = list(prs.slides)
    parts = []
    parts.append("""<!DOCTYPE html><html><head><meta charset="utf-8">
<style>
  @page { margin: 0.5in; }
  body { font-family: -apple-system, 'Segoe UI', Roboto, Arial, sans-serif; color:#111; line-height:1.4; margin:0; }
  .deckmeta { color:#666; font-size:12px; margin:0 0 16px; }
  h1.docttl { font-size:20px; margin:0 0 2px; }
  .slide { border:1px solid #d0d0d0; border-radius:8px; padding:14px 18px; margin:0 0 16px; page-break-inside:avoid; }
  .snum { font-size:11px; letter-spacing:.06em; text-transform:uppercase; color:#8a8a8a; font-weight:700; margin-bottom:6px; }
  .stitle { font-size:17px; font-weight:700; margin:0 0 8px; color:#0a2a5c; }
  ul { margin:6px 0 0; padding-left:22px; }
  li { margin:3px 0; }
  li.l1 { list-style:circle; margin-left:14px; }
  li.l2 { list-style:square; margin-left:28px; }
  table.t { border-collapse:collapse; margin:8px 0; font-size:12.5px; }
  table.t td { border:1px solid #bbb; padding:4px 8px; }
  .notes { margin-top:10px; padding-top:8px; border-top:1px dashed #ccc; color:#555; font-size:12px; }
  .notes b { color:#333; }
</style></head><body>""")
    parts.append('<h1 class="docttl">Deck content review</h1>')
    parts.append('<div class="deckmeta">%d slides · content-only view for approval (ordering, figures, narrative). Download the original for exact layout/graphics.</div>' % len(slides))

    for i, slide in enumerate(slides, 1):
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except Exception:
            pass
        title_shape_id = None
        try:
            title_shape_id = id(slide.shapes.title) if slide.shapes.title is not None else None
        except Exception:
            pass

        items = []
        for shape in slide.shapes:
            if title_shape_id is not None and id(shape) == title_shape_id:
                continue
            items.extend(shape_text(shape))

        body_html = []
        bullets = []
        for it in items:
            if isinstance(it, tuple) and it and it[0] == "bullet":
                _, lvl, txt = it
                bullets.append((lvl, txt))
            else:  # already-rendered html (e.g., table)
                if bullets:
                    body_html.append(_bullets_html(bullets)); bullets = []
                body_html.append(it)
        if bullets:
            body_html.append(_bullets_html(bullets))

        notes = ""
        try:
            if slide.has_notes_slide:
                nt = (slide.notes_slide.notes_text_frame.text or "").strip()
                if nt:
                    notes = nt
        except Exception:
            pass

        parts.append('<div class="slide">')
        parts.append('<div class="snum">Slide %d</div>' % i)
        if title:
            parts.append('<div class="stitle">%s</div>' % esc(title))
        parts.append("".join(body_html) if body_html else '<div style="color:#999;font-size:12px">(no extractable text — graphic/image slide; see original)</div>')
        if notes:
            parts.append('<div class="notes"><b>Notes:</b> %s</div>' % esc(notes))
        parts.append('</div>')

    parts.append("</body></html>")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("".join(parts))


def _bullets_html(bullets):
    lis = []
    for lvl, txt in bullets:
        cls = "l1" if lvl == 1 else ("l2" if lvl and lvl >= 2 else "")
        lis.append('<li class="%s">%s</li>' % (cls, esc(txt)))
    return "<ul>" + "".join(lis) + "</ul>"


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.stderr.write("usage: pptx2html.py <src.pptx> <out.html>\n")
        sys.exit(2)
    render(sys.argv[1], sys.argv[2])
