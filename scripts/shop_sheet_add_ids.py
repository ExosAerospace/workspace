#!/usr/bin/env python3
"""Write the permanent Exos ID column into Steve's shop workbook (exo3 shop-stable-ids, Jon 2026-09-30:
"add the ID column to Steve's sheet ... They should not have the rights to delete the column without my
decision").

From a source Current_Inventory_Locations*.xlsx:
  - ids come from data/shop-id-registry.json via the reindex matcher (aborts if any line would get a NEW id
    or any registry id would retire — the sheet must describe exactly what the registry holds);
  - an "Exos ID" column is APPENDED after the last used column on each item sheet (never inserted:
    openpyxl insert_cols does not rewrite formulas), header styled like its left neighbour;
  - each item sheet is protected so columns cannot be deleted/inserted, while everything else Steve needs
    stays allowed (format, insert/delete rows, sort, filter, hyperlinks) and every cell is unlocked;
  - the protection password is generated once into ~/.config/exos/shop-sheet-protect.pw (0600) and is
    never printed.
Then verifies the saved file by reopening it (ids, protection flags, formulas unchanged vs source).

  shop_sheet_add_ids.py --src X.xlsx [--out Y.xlsx] [--register-meta] [--set-required]
  --data-dir/--pw-file (selftest isolation)
"""
import argparse, copy, json, os, secrets, sys, time, datetime
import openpyxl
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import reindex_shop_inventory as R

UPLOADS = os.path.join(R.WS, "mission-control/uploads")
PW_FILE = os.path.expanduser("~/.config/exos/shop-sheet-protect.pw")
SPARE_COLS = 20      # unlocked spare columns right of the id column
UNLOCKED = Protection(locked=False)


def get_password(pw_file):
    if os.path.exists(pw_file):
        return open(pw_file).read().strip()
    os.makedirs(os.path.dirname(pw_file), exist_ok=True)
    pw = secrets.token_urlsafe(18)
    fd = os.open(pw_file, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "w") as f:
        f.write(pw + "\n")
    os.chmod(pw_file, 0o600)
    return pw


def last_used_col(ws):
    last = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, "") and c.column > last:
                last = c.column
    return last


def formulas(ws):
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")}


def protect(ws, pw):
    p = ws.protection
    p.sheet = True
    p.password = pw
    p.deleteColumns = True      # blocked
    p.insertColumns = True      # blocked
    for k in ("formatCells", "formatColumns", "formatRows", "insertRows", "deleteRows", "sort",
              "autoFilter", "insertHyperlinks", "selectLockedCells", "selectUnlockedCells"):
        setattr(p, k, False)    # allowed


def build(src, out, reg, pw):
    items, per_cat, sheets = R.parse_workbook(src)
    summary, retired, _ = R.assign_ids(items, reg)
    if summary["new"] or retired:
        raise SystemExit(f"source does not match the registry (new={summary['new']}, retired={len(retired)}); "
                         "run reindex_shop_inventory.py on it first")
    wb = openpyxl.load_workbook(src)
    src_formulas = {}
    id_cols = {}
    for title in R.ITEM_SHEETS:
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        src_formulas[title] = formulas(ws)
        hrow = sheets.get(title, {}).get("header_row", 1)
        col = sheets.get(title, {}).get("id_col") or (last_used_col(ws) + 1)
        id_cols[title] = col
        h = ws.cell(hrow, col)
        left = ws.cell(hrow, col - 1) if col > 1 else None
        if left is not None and left.has_style:
            h.font = copy.copy(left.font); h.fill = copy.copy(left.fill); h.border = copy.copy(left.border)
            h.alignment = copy.copy(left.alignment); h.number_format = left.number_format
        h.value = R.ID_HEADER
        ws.column_dimensions[get_column_letter(col)].width = 13
        # unlock every existing cell + column defaults (covers blank/spare rows) so rows stay editable/deletable
        max_c = max(ws.max_column, col) + SPARE_COLS
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max_c):
            for c in row:
                c.protection = UNLOCKED
        for ci in range(1, max_c + 1):
            ws.column_dimensions[get_column_letter(ci)].protection = UNLOCKED
        protect(ws, pw)
    for it in items:
        c = wb[it["_sheet"]].cell(it["_row"], id_cols[it["_sheet"]])
        c.value = it["id"]
        c.protection = UNLOCKED
    wb.save(out)
    return items, id_cols, src_formulas


def verify(out, items, id_cols, src_formulas):
    """Reopen and check. -> list of failures ([] = ok)."""
    fails = []
    wb = openpyxl.load_workbook(out)
    items2, _, sheets2 = R.parse_workbook(out)
    if [i["_sheet_id"] for i in items2] != [i["id"] for i in items]:
        fails.append("ids read back from the saved file differ from the assigned ids")
    for title, col in id_cols.items():
        ws = wb[title]
        if sheets2.get(title, {}).get("id_col") != col:
            fails.append(f"{title}: Exos ID header not found at column {col}")
        p = ws.protection
        if not p.sheet:
            fails.append(f"{title}: sheet protection off")
        if not (p.deleteColumns and p.insertColumns):
            fails.append(f"{title}: column delete/insert not blocked")
        if p.deleteRows or p.insertRows or p.sort or p.autoFilter or p.formatCells:
            fails.append(f"{title}: a row/format/sort/filter permission is blocked")
        if formulas(ws) != src_formulas[title]:
            fails.append(f"{title}: formulas changed vs source")
        locked = sum(1 for row in ws.iter_rows() for c in row if c.protection.locked)
        if locked:
            fails.append(f"{title}: {locked} locked cells")
    return fails


def register_meta(out):
    meta_p = os.path.join(UPLOADS, "_meta.json")
    meta = json.load(open(meta_p))
    entry = {
        "id": secrets.token_hex(8),
        "filename": os.path.basename(out).split("_", 1)[1],
        "stored_name": os.path.basename(out),
        "size": os.path.getsize(out),
        "extension": ".xlsx",
        "task_id": None,
        "project_id": None,
        "decision_id": None,
        "uploaded_by": "Exo (shop-stable-ids)",
        "uploaded_at": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z",
    }
    meta.append(entry)
    with open(meta_p + ".tmp", "w") as f:
        f.write(json.dumps(meta, indent=2))
    os.replace(meta_p + ".tmp", meta_p)
    return entry


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--src")
    ap.add_argument("--out")
    ap.add_argument("--data-dir", default=R.DATA)
    ap.add_argument("--pw-file", default=PW_FILE)
    ap.add_argument("--register-meta", action="store_true")
    ap.add_argument("--set-required", action="store_true")
    a = ap.parse_args(argv)
    src = a.src or R.newest_locations_xlsx()
    out = a.out or os.path.join(UPLOADS, f"{int(time.time() * 1000)}_Current_Inventory_Locations_EXOS-ID_"
                                         f"{datetime.date.today().isoformat()}.xlsx")
    reg_p = os.path.join(a.data_dir, "shop-id-registry.json")
    reg = R.load_registry(reg_p)
    if not reg:
        raise SystemExit(f"no registry at {reg_p}")
    pw = get_password(a.pw_file)
    items, id_cols, src_formulas = build(src, out, reg, pw)
    fails = verify(out, items, id_cols, src_formulas)
    res = {"src": os.path.basename(src), "out": out, "ids": len(items), "id_cols": id_cols,
           "formulas_compared": sum(len(v) for v in src_formulas.values()), "verify_failures": fails}
    if fails:
        print(json.dumps(res, indent=1)); return 2
    if a.register_meta:
        res["meta_entry"] = register_meta(out)
    if a.set_required:
        reg = R.load_registry(reg_p)
        reg.update(id_column_required=True, id_column_set_at=R.now_iso(), id_column_file=os.path.basename(out))
        json.dump(reg, open(reg_p + ".tmp", "w"), indent=1, ensure_ascii=False)
        os.replace(reg_p + ".tmp", reg_p)
        res["id_column_required"] = True
    print(json.dumps(res, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
