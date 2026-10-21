#!/usr/bin/env python3
"""Orbital On-Ramp Questionnaire (Jon 2026-10-21 06:53) — short partner questionnaire that decides the construct:
Exos as operator (wet lease) vs. partner as vehicle operator/licensee (Exos = manufacturer/integrator) vs. hybrid.
Editable .docx on PRIMARY letterhead; answer cells left blank for the partner."""
import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WS = "/Users/exomac/.openclaw/workspace"
LH = WS + "/assets/EXOS_Letterhead_Greenville_TX_address_PRIMARY.docx"
OUT = WS + f"/drafts/Orbital_OnRamp_Questionnaire_{datetime.date.today():%Y-%m-%d}_r2.docx"
NAVY = RGBColor(0x0A, 0x16, 0x28)

doc = Document(LH)
body = doc.element.body
for el in list(body):
    if el.tag != qn('w:sectPr'): body.remove(el)
for s in doc.sections:
    s.top_margin = Inches(1.7); s.bottom_margin = Inches(1.5); s.left_margin = s.right_margin = Inches(1.0)
    s.different_first_page_header_footer = False
    sp = s._sectPr
    for ref in list(sp.findall(qn('w:headerReference'))) + list(sp.findall(qn('w:footerReference'))):
        if ref.get(qn('w:type')) == 'first': sp.remove(ref)
    for tp in list(sp.findall(qn('w:titlePg'))): sp.remove(tp)

def para(text="", bold=False, size=10.5, color=None, align=None, space_after=6, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after); return p

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)

def borders(t):
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + e); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), '9BB3D6'); b.append(el)
    tblPr.append(b)

n = [0]
def section(title, qs):
    para(title, bold=True, size=12, color=NAVY, space_after=4)
    t = doc.add_table(rows=0, cols=3); t.autofit = False; borders(t)
    for q, hint in qs:
        n[0] += 1
        c = t.add_row().cells
        c[0].width, c[1].width, c[2].width = Inches(0.35), Inches(3.0), Inches(3.15)
        c[0].text = str(n[0]); shade(c[0], 'E8EEF6'); shade(c[1], 'F5F8FC')
        c[1].text = ""; r = c[1].paragraphs[0].add_run(q); r.bold = True; r.font.size = Pt(9.5)
        if hint:
            h = c[1].add_paragraph().add_run(hint); h.italic = True; h.font.size = Pt(8.5); h.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        for p in c[0].paragraphs:
            for rr in p.runs: rr.font.size = Pt(9.5)
    para("", space_after=4)

from docx.enum.text import WD_ALIGN_PARAGRAPH
para("ORBITAL ON-RAMP QUESTIONNAIRE", bold=True, size=15, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Exos Aerospace Systems & Technologies, Inc. — Partner Structuring", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10, italic=True)
para("Thank you for your interest in pairing Exos BLK3 suborbital flights with your payload bookings on orbital first flights. "
     "Your answers help us propose the right structure for you. Depending on your registrations and capabilities, that may be: "
     "(A) Exos operates the vehicle for you; (B) your company operates as the vehicle operator and licensee, with Exos as "
     "manufacturer and integrator; or (C) a combination. The form takes about 10 minutes. Write \"No\" or \"N/A\" where something "
     "doesn't apply. Information you provide is treated as confidential.", size=10)
para("Company: ____________________________   Completed by: ____________________   Title: ______________   Date: __________", size=9.5, space_after=10)

section("1. Company and U.S. Government registrations", [
    ("Company legal name, state/country of formation, year formed, and headquarters address.", ""),
    ("SAM.gov registration: active? UEI and CAGE code.", "Enter \"Not registered\" if none."),
    ("Registered with the State Department's Directorate of Defense Trade Controls (DDTC) under ITAR (22 CFR Part 122)? Registration code and expiration date.", "Also name your Empowered Official, if any."),
    ("Export authorizations held now or in the past 5 years (ITAR licenses or agreements such as DSP-5 or TAA; EAR licenses)?", "List type and number; no technical detail needed."),
    ("Joint Certification Program (DD Form 2345) certification number, and CMMC level or NIST SP 800-171 SPRS score.", ""),
    ("Security clearances — facilities: does the company (or a parent/affiliate) hold a facility security clearance (FCL)? Level, sponsoring agency, and whether you hold approved safeguarding capability or accredited secure space.", "Yes/No and level only. Do not include any classified information."),
    ("Security clearances — personnel: how many of your employees or key contractors hold active personnel security clearances, at what levels, and in which roles? Any pending investigations or sponsorships?", "Counts, levels and roles only. No names or classified information needed at this stage."),
    ("U.S. Government contracts, grants or Other Transaction agreements in the past 5 years.", "Agency, contract number, value, period, and your role (prime or sub)."),
    ("Government consortium memberships, and your small-business status (SB, SDVOSB, 8(a), HUBZone, WOSB) and NAICS codes.", ""),
])
section("2. Ownership and foreign involvement", [
    ("Is the company U.S.-owned and controlled? List each owner of 10% or more, with percentage and citizenship.", ""),
    ("Any foreign ownership, control, financing or board seats? Any non-U.S. persons who would need access to technical data?", "If yes, give countries and roles."),
])
section("3. Operator capability", [
    ("Does the company hold, or has it applied for, an FAA launch or reentry license (14 CFR Part 450) or an experimental permit? License or application number and status.", ""),
    ("Launch operations experience: key people, their roles (flight safety, range, mission assurance), and years of experience.", ""),
    ("Can the company obtain FAA-required third-party liability insurance? Current aerospace coverage and broker.", ""),
    ("Agreements with spaceports, ranges or launch sites (existing or in negotiation).", ""),
    ("Which structure do you prefer: (A) Exos operates, (B) you operate as licensee, or (C) combination / undecided? Why?", ""),
])
section("4. Orbital first-flight plan", [
    ("Target orbital launch vehicle(s) and provider(s), planned first-flight date, and capacity booked.", "Status: signed contract, LOI, or in discussion."),
    ("Follow-on flight options or rights secured with the orbital provider.", ""),
    ("Payloads you plan to fly (for example reentry bodies, flight data recorders, qualification articles), and any customers already identified.", ""),
    ("How the first-flight payload bay purchase is funded, and your timeline.", ""),
    ("Primary point of contact: name, title, email and phone.", ""),
])
para("Please return the completed questionnaire to John Quinn, Cofounder & CEO, at jquinn@exosaero.com. Exos may ask for supporting "
     "documents (for example SAM.gov, DDTC or FAA records). Formal due diligence and a separate background check authorization follow "
     "before any teaming arrangement.", size=9.5, italic=True)
doc.core_properties.title = "Orbital On-Ramp Questionnaire"; doc.core_properties.author = "Exos Aerospace"
doc.save(OUT); print(OUT, n[0], "questions")
