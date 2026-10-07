#!/usr/bin/env python3
"""
refresh_send_decisions.py — bring every open/delegated send-Decision up to the newest, copy-clean draft.

Jon 10/7: "look at all our current proposals due to be sent and revise them accordingly and resubmit
for send. How do we not send an outdated one?"

For each open/delegated Decision that references a stale draft (proposal_freshness_check):
  1. take the newest same-doc copy in drafts/
  2. apply COPY_RULES; if anything changed, save a new revision (today's date, or _vN if today exists)
  3. re-point the Decision at that file, set proposedAction = the draft's text, fix the title
  4. collapse duplicate Decisions that resolve to the same draft (older one -> answered "Superseded")
Investor (OPP_FUND-*) Decisions are re-pointed only; they stay held for the securities approver.

  refresh_send_decisions.py            dry run (report only)
  refresh_send_decisions.py --apply
"""
import datetime, importlib.util, json, os, re, sys

WS = "/Users/exomac/.openclaw/workspace"
DEC = os.path.join(WS, "mission-control", "data", "decisions.json")
TODAY = datetime.date.today().isoformat()
sp = importlib.util.spec_from_file_location("pfc", os.path.join(WS, "scripts", "proposal_freshness_check.py"))
pfc = importlib.util.module_from_spec(sp); sp.loader.exec_module(pfc)
from docx import Document


def task_closed(tid):
    """True if the Decision's spawned MC task is Done/Cancelled — never re-stage a send for a closed task."""
    if not tid:
        return False
    try:
        sys.path.insert(0, os.path.join(WS, "scripts")); import capstone_capture as cc
        r = cc.supa("GET", f"tasks?select=state&id=eq.{int(tid)}")
        return bool(r) and r[0].get("state") in ("Done", "Cancelled")
    except Exception:
        return False


def revise(path):
    """Apply copy rules to a docx; return (path_to_use, hits). Writes a new revision only if changed."""
    doc = Document(path); hits = []
    for p in doc.paragraphs:
        new, h = pfc.copy_fix(p.text)
        if h:
            hits += h
            for r in p.runs[1:]:
                r.text = ""
            if p.runs:
                p.runs[0].text = new
    if not hits:
        return path, []
    base = re.sub(r"_\d{4}-\d{2}-\d{2}(_v\d+)?\.docx$", "", os.path.basename(path))
    out = os.path.join(pfc.DRAFTS, f"{base}_{TODAY}.docx"); n = 2
    while os.path.exists(out):
        out = os.path.join(pfc.DRAFTS, f"{base}_{TODAY}_v{n}.docx"); n += 1
    doc.save(out)
    return out, hits


def main(apply):
    raw = json.load(open(DEC)); arr = raw if isinstance(raw, list) else raw["decisions"]
    byid = {d.get("id"): d for d in arr}
    now = datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%S.000Z")
    report, by_file, cache = [], {}, {}
    for did, title, p in pfc.open_decision_files():
        if not pfc.check(p):
            continue
        it = byid[did]
        if task_closed(it.get("spawned_task")):
            report.append((did, os.path.relpath(p, WS), "(skipped: task closed)", [], False))
            if apply:
                it.update(status="answered", answer="Withdrawn — source task closed", answeredAt=now, updated=now,
                          note=(it.get("note") or "") + f" [withdrawn {TODAY}: task #{it['spawned_task']} closed; no send]")
            continue
        latest = pfc.latest_copy(p)
        if latest not in cache:
            cache[latest] = revise(latest) if apply else (latest, pfc.copy_fix("\n".join(x.text for x in Document(latest).paragraphs))[1])
        newp, hits = cache[latest]
        rel_old = os.path.relpath(p, WS); rel_new = os.path.relpath(newp, WS)
        fund = "OPP_FUND-" in rel_new
        newtitle, th = pfc.copy_fix(it.get("title", ""))
        report.append((did, rel_old, rel_new, hits + th, fund))
        by_file.setdefault(rel_new, []).append(did)
        if apply:
            for k in ("context", "note", "question"):
                if it.get(k):
                    it[k] = pfc.copy_fix(str(it[k]).replace(rel_old, rel_new))[0]
            if rel_new not in json.dumps(it):
                it["context"] = (it.get("context", "") + f"\n\nDraft staged: {rel_new}").strip()
            it["title"] = newtitle
            if not fund:
                it["proposedAction"] = "\n".join(x.text for x in Document(newp).paragraphs)
            it["updated"] = now
            it["note"] = (it.get("note") or "") + f" [refreshed {TODAY}: current draft {os.path.basename(rel_new)}; freshness + copy rules checked]" + \
                (" [investor item: held for securities approver, Series A outreach paused]" if fund else "")
    dups = {f: ids for f, ids in by_file.items() if len(ids) > 1}
    for f, ids in dups.items():
        keep = max(ids, key=lambda i: byid[i].get("created", ""))
        for i in ids:
            if i != keep and apply:
                d = byid[i]
                d.update(status="answered", answer="Superseded", answeredAt=now, updated=now,
                         note=(d.get("note") or "") + f" [superseded by {keep}: same draft {os.path.basename(f)}]")
    if apply:
        tmp = DEC + ".tmp"
        json.dump(raw, open(tmp, "w"), indent=2); os.replace(tmp, DEC)
    for did, o, n, h, fund in report:
        print(f"{did} | {os.path.basename(o)} -> {os.path.basename(n)}{' | FIXES: ' + '; '.join(h) if h else ''}{' | INVESTOR-HOLD' if fund else ''}")
    for f, ids in dups.items():
        print(f"DUP {os.path.basename(f)}: {ids} -> keep newest")
    print(f"{len(report)} Decision(s) {'refreshed' if apply else 'would refresh'}")


if __name__ == "__main__":
    main("--apply" in sys.argv)
