#!/usr/bin/python3
"""Validation memo: FNEX Advisory Agreement (Board Copy, 10/5) vs the 10/5 'fix it' call transcript."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors

D = "/Users/exomac/.openclaw/workspace/drafts"
BASE = "Exos_FNEX_Agreement_FixIt_Validation_2026-10-05"
TITLE = "FNEX Advisory Agreement — Board Copy vs. 'Fix It' Call (10/5, 5:37 PM CDT)"
INTRO = ("Source documents on MC task #2596: 'FNEX Advisory Agreement latest Exos markup Board Copy.docx' and 'fix it.docx' "
         "(30-min transcript, John Quinn / Michael Sheldon). Every point agreed on the call was checked against the clause text, "
         "tracked insertions included. Result: 12 points are already in the agreement, and 9 are missing or wrong and need a fix "
         "before John sends it to Raquel, Graham and Michael. The proposed replacement wording is below. "
         "INTERNAL — not reviewed by Counsel.")

OK = [
    "Fees: MEV defined at the Transaction price per share at Closing ('Transaction', not 'merger', in §3/§3a).",
    "No FNEX placement fee on the PIPE or later placements; 'in no event shall exceed' removed; 3% 'targeted'.",
    "Fees payable to FNEX solely for the Registered Representative (Raquel); Michael's name and 'or its assigns' removed.",
    "Share sale: tranche 1 leak-out ≤10% of daily volume, no sales below the price; tranche 2 ratable over 18 months, "
    "2x early release, NO 'continued advisory engagement' condition.",
    "Breakup: $60K/month accrual tied to Schedule A milestones; 'after that election' = after Exos's decision; greater of "
    "accrued fee or 0.5%, capped at the Fees.",
    "Breakup: Exos's own raises excluded; 'itself' removed; '$10M Series A' wording used.",
    "Strategic: no fee on a JV, license or commercial partnership; existing Exos relationships excluded.",
    "'Not payable twice on the same transaction' language removed (Michael: it confuses things).",
    "Term: 6 months, no auto-renewal; 3-month milestone review.",
    "Tail period: 12 months, Introduced Party list. Indemnity cap = total cash Fees (previously agreed).",
    "Governing law: Texas. §19 PIPE/Exos protections: left as-is for now (Michael: 'address it later').",
    "Schedule A: milestones (a) through (d) listed, with dates [TBC].",
]
GAPS = [
    ("1. §3 No Placement Fee — PIPE agent cost",
     "Says PIPE third-party costs are 'approximately 3%' (stated twice).",
     "Agreed: the PIPE placement agent will be closer to 7% ('seven is about midpoint'), and ATM/later placements are targeted at 3%. John: 'Agreed.'",
     "\"…Third-party placement-agent costs of the PIPE are targeted at approximately seven percent (7%), and at-the-market (ATM) or later-placement costs at approximately three percent (3%), in each case approved in writing by EXOS in advance…\" Delete the duplicate 3% sentence."),
    ("2. §3 No Placement Fee vs. the new 5% strategic fee",
     "Bars FNEX compensation on 'any later placement or financing of EXOS'.",
     "Agreed: 5% cash on a strategic investment that FNEX introduces outside the PIPE.",
     "Add: \"…except as provided under Strategic Transactions for an FNEX-Introduced Investment.\""),
    ("3. Strategic Transactions — garbled",
     "Reads '…Fees calculated under this Section 3 on the value of that Change of Control (five percent (5%) in cashNo fee…'. This mixes the two cases, and the parenthesis is never closed.",
     "Agreed: (a) Change of Control with an FNEX-introduced party = 1% cash + 2% shares, IN PLACE OF the Transaction Fee ('you can't sell control twice'); (b) an FNEX-introduced equity investment outside the PIPE = 5% of the amount invested, CASH only, IN ADDITION to the Transaction Fee. Delete the 'never in addition / in place of' and capping language.",
     "(a) \"…FNEX shall receive the Fees under Section 3 (1% cash, 2% shares) on the value of that Change of Control, in place of, and not in addition to, any Fee on the Transaction.\" (b) \"If FNEX first introduces in writing during the Term an investor that makes an equity investment in EXOS outside the PIPE, and that investment closes, EXOS shall pay FNEX five percent (5%) of the amount invested, in cash only, in addition to any Fee on the Transaction.\""),
    ("4. Fees Earned Only at Closing + Expenses",
     "'No Fee, Awarded Shares or expenses are earned or payable unless Closing occurs'. The Expenses clause says '…payable at Closing'.",
     "Agreed: tying expenses to Closing 'doesn't make sense'. Pre-approved expenses are paid as they are incurred, so FNEX is not pre-funding Exos.",
     "Delete 'or expenses' from Fees Earned Only at Closing. Expenses: \"…reimbursed within thirty (30) days of FNEX's invoice with receipts, whether or not Closing occurs, up to $75,000 in aggregate.\""),
    ("5. Breakup — funding carve-out",
     "Breakup fee excused for 'failure to fund at an agreed amount'.",
     "Agreed: drop 'agreed size'. FNEX is not paid if Exos raises no money from any source. If Exos IS funded from any source, the breakup fee applies.",
     "\"…(including failure to obtain funding from any source, …)\". Also fix the stray ')' after '…government funding' (Michael: drop the parentheses)."),
    ("6. Key persons",
     "Breakup cross-references 'Termination by EXOS any Key Persons clauses', but §9 has no key-person clause.",
     "Agreed: 'any key persons' in lowercase as boilerplate. The key persons (Michael, Graham, Raquel) go in a SEPARATE side agreement, outside the FINRA document, that lets Exos terminate if one of them is not active.",
     "Change to \"…or under any key persons agreement\". Counsel to draft the side letter (Exos may terminate this Agreement without fee or breakup liability if any key person ceases to be actively involved)."),
    ("7. 'Cause' is undefined",
     "Tail says '[No Tail Period fee is payable if EXOS terminates this Agreement for Cause.]', but 'Cause' is never defined.",
     "Michael asked for a capitalized, defined, FAIR Cause ('there's cause and there's cause').",
     "Define it in §9: \"'Cause' means FNEX's (i) material breach not cured within 15 days of notice, (ii) loss of FINRA/SEC registration or a bar, (iii) fraud, gross negligence or willful misconduct, or (iv) violation of securities law in connection with this Agreement.\" Remove the brackets."),
    ("8. Leftover 'merger price per share'",
     "The Share Sale clause ('no sales below the merger price per share', '2x the merger price per share') and §19 still say 'merger price'.",
     "Agreed at 0:20: use 'Transaction price per share', a defined term.",
     "Replace with 'Transaction price per share' throughout (§19 can follow when it is reopened)."),
    ("9. Term — post-closing role",
     "'…negotiate in good faith… FNEX's continued engagement after Closing'.",
     "Agreed: in spirit only. Any post-closing advisory role is a separate future deal with its own economics, and must not read as a 2-year engagement.",
     "\"…and any post-Closing advisory engagement, which would be under a separate written agreement.\""),
]
MINOR = ("Housekeeping: Late Fees says 'from Seller' and should say 'from EXOS'. The FEES clause still has bracketed [TBC] items: "
         "the 0.5%/0.5% cash split tied to $40M of FNEX-placed financings, and 'one-half cash, one-half stock' on the breakup. "
         "Neither was discussed on the fix-it call.")
NEXT = ("Next step: John approves these 9 fixes. Exo then applies them as tracked changes on the Board Copy, and John sends it to Raquel, Graham and "
        "Michael ('this is what I'm taking to my board'). Model updated: Exos_FNEX_Model_EXOS_Redline_v7_2026-10-05.xlsx.")

def build_docx():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    doc.add_heading(TITLE, 1)
    doc.add_paragraph(INTRO)
    doc.add_heading("Missing or wrong — fix before sending (9)", 2)
    t = doc.add_table(rows=1, cols=4); t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Clause", "Board Copy now", "Agreed on call", "Proposed fix"]):
        t.rows[0].cells[i].text = h
    for g in GAPS:
        cells = t.add_row().cells
        for i, v in enumerate(g): cells[i].text = v
    doc.add_paragraph(MINOR)
    doc.add_heading("Already in the agreement (12)", 2)
    for o in OK: doc.add_paragraph(o, style="List Bullet")
    doc.add_heading("Model impact (v6 → v7)", 2)
    doc.add_paragraph("PIPE placement-agent cost: 3% → 7% (third party, not FNEX). ATM/later placements: 7% → 3%. FNEX travel is now paid as incurred, "
                      "so it is a sunk cost even with no Closing. New 5% cash strategic-investment fee input (default $0). "
                      "Net to Exos at a $15M PIPE / $350M MEV goes from ≈$11.7M to ≈$11.1M, which still clears the $10M Net Closing Cash condition.")
    doc.add_paragraph(NEXT)
    p = os.path.join(D, BASE + ".docx"); doc.save(p); return p

def build_pdf():
    ss = getSampleStyleSheet(); sm = ParagraphStyle("sm", parent=ss["Normal"], fontSize=7.5, leading=9)
    p = os.path.join(D, BASE + ".pdf")
    el = [Paragraph(TITLE, ss["Heading2"]), Paragraph(INTRO, ss["Normal"]), Spacer(1, 6),
          Paragraph("Missing or wrong — fix before sending (9)", ss["Heading3"])]
    data = [[Paragraph(f"<b>{h}</b>", sm) for h in ["Clause", "Board Copy now", "Agreed on call", "Proposed fix"]]]
    data += [[Paragraph(v.replace("&", "&amp;"), sm) for v in g] for g in GAPS]
    tb = Table(data, colWidths=[85, 120, 150, 185], repeatRows=1)
    tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F4A261"))]))
    el += [tb, Spacer(1, 4), Paragraph(MINOR, sm), Paragraph("Already in the agreement (12)", ss["Heading3"])]
    el += [Paragraph("• " + o, sm) for o in OK]
    el += [Paragraph("Model impact (v6 → v7)", ss["Heading3"]),
           Paragraph("PIPE agent 3%→7% (third party); ATM/later 7%→3%; travel paid as incurred; new 5% cash strategic fee input. "
                     "Net to Exos at $15M PIPE / $350M MEV ≈$11.7M → ≈$11.1M (still clears the $10M condition).", sm),
           Spacer(1, 4), Paragraph(NEXT, sm)]
    SimpleDocTemplate(p, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36).build(el)
    return p

print(build_docx()); print(build_pdf())
