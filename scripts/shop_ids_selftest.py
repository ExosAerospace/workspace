#!/usr/bin/env python3
"""Selftest for permanent shop ids + the Exos ID column guard (exo3 eval for shop-id-guard-exo3).
Runs entirely in a temp dir (own registry/inventory/logs, throwaway password file, --decision-dry-run,
--no-fleet-log): never touches live data, never files a real Decision.

  python3 scripts/shop_ids_selftest.py      -> prints PASS/FAIL per check, exit 0 only if all pass
"""
import json, os, shutil, subprocess, sys, tempfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import reindex_shop_inventory as R

REINDEX = os.path.join(HERE, "reindex_shop_inventory.py")
ADDIDS = os.path.join(HERE, "shop_sheet_add_ids.py")
HDR = ["Location", "Source/SKU", "Contents", "Exos SKU", "Qty", "Value", "Total"]
results = []


def check(name, ok, why=""):
    results.append((name, bool(ok)))
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  -- {why}"))


def base_rows():
    wh, pr = [], []
    for i in range(40):
        wh.append([f"A{i // 4}.{i % 4}" if i % 3 else "", f"SRC-{i:03d}", f"Warehouse part {i}", f"EX{i:04d}", i + 1, 10])
    for k in range(3):                       # duplicates in one location (occurrence tiebreak)
        wh.append(["Cab2", "", "Bolt kit", "", 1, 5])
    for i in range(20):
        pr.append([f"1A-{i}", "", f"Parts room item {i}", "", 2, 3])
    return {"Warehouse": wh, "Parts Room": pr}


def make_wb(path, sheets, ids=None):
    """sheets: {title: [row lists]}; ids: {title: [id or ''] aligned with rows} -> adds an Exos ID column."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(HDR + ([R.ID_HEADER] if ids is not None else []))
        for n, r in enumerate(rows):
            rr = n + 2
            row = list(r) + [f"=E{rr}*F{rr}"]
            if ids is not None:
                row.append(ids[title][n])
            ws.append(row)
    wb.create_sheet("Legend").append(["COLOR", "=", "CODE"])
    wb.save(path)


def run(args, data, logs):
    p = subprocess.run([sys.executable, REINDEX, "--data-dir", data, "--log-dir", logs, "--no-fleet-log",
                        "--decision-dry-run", "--json"] + args, capture_output=True, text=True)
    try:
        return p.returncode, json.loads(p.stdout.strip().splitlines()[-1])
    except Exception:
        return p.returncode, {"stdout": p.stdout, "stderr": p.stderr}


def snap(src, dst):
    shutil.rmtree(dst, ignore_errors=True)
    shutil.copytree(src, dst)


def read(p):
    return open(p).read() if os.path.exists(p) else None


def main():
    t = tempfile.mkdtemp(prefix="shopids-")
    base_data, logs = os.path.join(t, "data0"), os.path.join(t, "logs")
    os.makedirs(base_data); os.makedirs(logs)
    rows = base_rows()
    v1 = os.path.join(t, "v1_Current_Inventory_Locations.xlsx")
    make_wb(v1, rows)

    # seed: first load on an empty registry allocates SHOP-00001..
    rc, r0 = run(["--file", v1], base_data, logs)
    ids0 = r0.get("ids", [])
    check("seed load allocates one id per line", rc == 0 and len(ids0) == 63 and len(set(ids0)) == 63, r0)

    # 1. reload with same sheet -> 0 id changes
    d = os.path.join(t, "d1"); snap(base_data, d)
    rc, r = run(["--file", v1], d, logs)
    check("same sheet reload -> 0 id changes", rc == 0 and r["ids"] == ids0 and r["summary"]["new"] == 0
          and r["summary"]["retired"] == 0, r.get("summary"))

    # 2. row inserted in the middle -> only the new row gets a new id
    d = os.path.join(t, "d2"); snap(base_data, d)
    rows2 = json.loads(json.dumps(rows)); rows2["Warehouse"].insert(20, ["A9.9", "", "Brand new valve", "", 1, 1])
    v2 = os.path.join(t, "v2_Current_Inventory_Locations.xlsx"); make_wb(v2, rows2)
    rc, r = run(["--file", v2], d, logs)
    exp = ids0[:20] + ["SHOP-00064"] + ids0[20:]
    check("row inserted mid-sheet -> only the new row gets a new id", rc == 0 and r["ids"] == exp
          and r["summary"]["new"] == 1, r.get("summary"))

    # 3. row moved to a new location -> id kept
    d = os.path.join(t, "d3"); snap(base_data, d)
    rows3 = json.loads(json.dumps(rows)); moved = rows3["Warehouse"].pop(5); moved[0] = "Z9.9"
    rows3["Warehouse"].append(moved)
    v3 = os.path.join(t, "v3_Current_Inventory_Locations.xlsx"); make_wb(v3, rows3)
    rc, r = run(["--file", v3], d, logs)
    ok = rc == 0 and r["summary"]["new"] == 0 and r["summary"]["retired"] == 0 and sorted(r["ids"]) == sorted(ids0)
    check("row moved to new location -> id kept", ok and r["summary"]["moved"] >= 1, r.get("summary"))
    reg3 = json.load(open(os.path.join(d, "shop-id-registry.json")))
    check("moved row's registry fingerprint updated to new location",
          reg3["ids"][ids0[5]]["location"] == "z9.9", reg3["ids"][ids0[5]])

    # step 2 tool: Exos ID column + protection; id_column_required=true
    dg = os.path.join(t, "dg"); snap(base_data, dg)
    pw = os.path.join(t, "pw", "shop.pw")
    v1id = os.path.join(t, "v1id_Current_Inventory_Locations_EXOS-ID.xlsx")
    p = subprocess.run([sys.executable, ADDIDS, "--src", v1, "--out", v1id, "--data-dir", dg, "--pw-file", pw,
                        "--set-required"], capture_output=True, text=True)
    pwval = read(pw).strip() if read(pw) else "<none>"
    check("add-ids tool verifies clean (ids, protection, formulas)", p.returncode == 0
          and '"verify_failures": []' in p.stdout, p.stdout[-400:] + p.stderr[-400:])
    check("password file 0600 and never printed", oct(os.stat(pw).st_mode & 0o777) == "0o600"
          and pwval not in p.stdout and pwval not in p.stderr)
    wb = openpyxl.load_workbook(v1id)
    ws = wb["Warehouse"]
    check("id column appended (not inserted): header at col H, formulas intact",
          ws.cell(1, 8).value == R.ID_HEADER and ws["G2"].value == "=E2*F2" and ws.protection.deleteColumns
          and not ws.protection.deleteRows)
    items_id, _, _ = R.parse_workbook(v1id)
    sheet_ids = {"Warehouse": [], "Parts Room": []}
    for it in items_id:
        sheet_ids[it["_sheet"]].append(it["_sheet_id"])
    check("every line carries its registry id in the sheet", [i["_sheet_id"] for i in items_id] == ids0)

    # 4. with the column: same file -> all kept, 0 changes
    d = os.path.join(t, "d4"); snap(dg, d)
    rc, r = run(["--file", v1id], d, logs)
    check("id-column reload -> kept all, 0 id changes", rc == 0 and r["ids"] == ids0
          and r["summary"]["kept"] == 63, r.get("summary"))

    # 5. with the column: new row (blank id) inserted mid-sheet + a row moved -> accepted
    d = os.path.join(t, "d5"); snap(dg, d)
    r5, i5 = json.loads(json.dumps(rows)), json.loads(json.dumps(sheet_ids))
    r5["Warehouse"].insert(10, ["B1.1", "", "Steve's new line", "", 1, 1]); i5["Warehouse"].insert(10, "")
    r5["Parts Room"][3][0] = "9Z-1"
    v5 = os.path.join(t, "v5.xlsx"); make_wb(v5, r5, i5)
    rc, r = run(["--file", v5], d, logs)
    check("id column + blank new row + moved row -> accepted, 1 new id, rest kept",
          rc == 0 and r["summary"]["new"] == 1 and r["summary"]["kept"] == 63
          and r["ids"][:10] == ids0[:10] and r["ids"][11:] == ids0[10:], r.get("summary"))

    # 6. column removed -> rejected + decision payload; registry + inventory untouched
    d = os.path.join(t, "d6"); snap(dg, d)
    before_reg, before_inv = read(os.path.join(d, "shop-id-registry.json")), read(os.path.join(d, "shop-inventory.json"))
    rc, r = run(["--file", v1, "--uploader", "Steve Norsworthy"], d, logs)
    dp = r.get("decision_payload") or {}
    check("column removed -> REJECTED (exit 3)", rc == 3 and r.get("rejected") and
          any("column missing" in x for x in r.get("reasons", [])), r)
    check("decision payload built per spec", dp.get("title") == "Shop inventory upload rejected: Exos ID column missing/changed"
          and dp.get("priority") == "high" and dp.get("assignedTo") == "jquinn@exosaero.com"
          and dp.get("approver") == "John Quinn" and dp.get("dedupeKey") == f"shop-id-guard:{os.path.basename(v1)}"
          and dp.get("options", [""])[0].startswith("Keep rejecting") and len(dp.get("options", [])) == 2
          and "Steve Norsworthy" in dp.get("context", "") and os.path.basename(v1) in dp.get("context", ""), dp)
    check("rejected upload: decision NOT filed in test (dry-run)", r.get("decision") == "dry-run (not filed)", r.get("decision"))
    check("rejected upload: registry + shop-inventory.json unchanged",
          read(os.path.join(d, "shop-id-registry.json")) == before_reg
          and read(os.path.join(d, "shop-inventory.json")) == before_inv)
    lg = read(os.path.join(logs, "shop-id-guard.log")) or ""
    check("rejection reason logged", "REJECTED" in lg and os.path.basename(v1) in lg)
    gs = json.load(open(os.path.join(d, "shop-id-guard.json")))
    check("rejected file recorded (auto-pick skips it)", os.path.basename(v1) in gs.get("rejected", {}))

    # 7. >2% blank/duplicate/unknown ids -> rejected
    d = os.path.join(t, "d7"); snap(dg, d)
    i7 = json.loads(json.dumps(sheet_ids)); i7["Warehouse"][0] = ""; i7["Warehouse"][1] = i7["Warehouse"][2]
    i7["Parts Room"][0] = "SHOP-99999"
    v7 = os.path.join(t, "v7.xlsx"); make_wb(v7, rows, i7)
    rc, r = run(["--file", v7], d, logs)
    check(">2% blank/duplicate/unknown ids -> REJECTED", rc == 3 and any("Exos ID (blank 1, duplicate 1, unknown 1)"
          in x for x in r.get("reasons", [])), r.get("reasons"))

    # 8. mass removal (>25% of lines) with the column intact -> rejected (no bulk deletion without Jon)
    d = os.path.join(t, "d8"); snap(dg, d)
    r8 = {"Warehouse": rows["Warehouse"][:10], "Parts Room": rows["Parts Room"]}
    i8 = {"Warehouse": sheet_ids["Warehouse"][:10], "Parts Room": sheet_ids["Parts Room"]}
    v8 = os.path.join(t, "v8.xlsx"); make_wb(v8, r8, i8)
    rc, r = run(["--file", v8], d, logs)
    check("bulk removal >25% -> REJECTED", rc == 3 and any("retire" in x for x in r.get("reasons", [])), r.get("reasons"))

    # 9. a few rows deleted (normal use) -> accepted, those ids retired (not deleted)
    d = os.path.join(t, "d9"); snap(dg, d)
    r9, i9 = json.loads(json.dumps(rows)), json.loads(json.dumps(sheet_ids))
    del r9["Parts Room"][0]; del i9["Parts Room"][0]
    v9 = os.path.join(t, "v9.xlsx"); make_wb(v9, r9, i9)
    rc, r = run(["--file", v9], d, logs)
    reg9 = json.load(open(os.path.join(d, "shop-id-registry.json")))
    gone = sheet_ids["Parts Room"][0]
    check("row deleted -> accepted, id retired and kept in registry", rc == 0 and r["summary"]["retired"] == 1
          and reg9["ids"][gone]["status"] == "retired")
    check("retired id never reused (next_id unchanged)", reg9["next_id"] == 64, reg9["next_id"])

    # 10. Jon's accept path: --override-guard re-matches the rejected file by location+contents
    d = os.path.join(t, "d10"); snap(os.path.join(t, "d6"), d)
    rc, r = run(["--file", v1, "--override-guard"], d, logs)
    gs = json.load(open(os.path.join(d, "shop-id-guard.json")))
    check("override accepts + re-matches (0 id changes) + clears rejection", rc == 0 and r["ids"] == ids0
          and os.path.basename(v1) not in gs.get("rejected", {}), r.get("summary"))

    shutil.rmtree(t, ignore_errors=True)
    n_ok = sum(1 for _, ok in results if ok)
    print(f"\n{n_ok}/{len(results)} passed")
    return 0 if n_ok == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
