#!/usr/bin/env python3
"""build_montana_partner_guide.py — "Montana Partner Guide — Working with Exos (exo3)".

Partner-facing onboarding guide for Montana Spaceport & Test Range (Edgar Johansson,
Sean McClung, Isaiah Centifanto): how Exos works (exo3 / Mission Control), their scoped
MC access, the shared tasks (#3303-#3310), Program Operating Agreement Q&A, and how to
work with Exos most effectively. Jon email-command 2026-10-01. Parent MC task #3287.

GUARDRAILS (partner-facing): no ITAR/technical data beyond what Montana already holds
(Neil's 9/18 Launch Site Basics); no prices/phase dollar figures (they live in the POA);
no investor/finance content; Part 431 held / Part 450 in progress; BLK3 licensed, not
yet flown; Montana site = planned/in development; no free flights.

  Word (editable, on PRIMARY letterhead):  PYTHONNOUSERSITE=1 python3.12 scripts/build_montana_partner_guide.py docx
  PDF twin (reportlab + locked stationery): python3 scripts/build_montana_partner_guide.py pdf
"""
import os, sys

WS = "/Users/exomac/.openclaw/workspace"
OUT = os.path.join(WS, "drafts", "Montana_Partner_Guide_Working_with_Exos_exo3_2026-10-01")
LH = os.path.join(WS, "assets", "EXOS_Letterhead_Greenville_TX_PRIMARY.docx")

TITLE = "Montana Partner Guide — Working with Exos"
SUB = ("For Montana Spaceport & Test Range: Edgar Johansson, Sean McClung, Isaiah Centifanto  |  "
       "October 1, 2026")

C = [
 ("h1", "1. Why this guide"),
 ("p", "Exos runs its programs inside Mission Control, our system of record, using an AI-native operating "
       "model we call exo3. As of today you each have your own sign-in to a Montana partner view. This guide "
       "covers how that works, answers the questions we expect on the Program Operating Agreement, and "
       "sets out how we think Montana and Exos can move fastest together."),

 ("h1", "2. How Exos works: the exo3 model in one page"),
 ("b", "People set direction and make every commitment. Exo, our AI operations layer, does the tracking, "
       "drafting, routing and follow-up. Nothing contractual, financial, governmental or export-related is "
       "decided by AI. A named Exos officer decides and signs."),
 ("b", "Every workstream is a task with a named human owner. Every decision is recorded with who decided and "
       "when, and every action is logged. That audit trail is the point: it lets a small team run a "
       "licensed launch operation with large-company discipline."),
 ("b", "What it means for you: one place instead of scattered email threads. Post a note on the relevant "
       "task; Exo routes it to the Exos owner, and the answer comes back on the same task where all three of "
       "you can see it. Email still works. Mission Control is simply faster and keeps the record."),
 ("b", "Exos people (John, Neil, Phil, Lizi) send anything formal to you themselves. Our AI does not send "
       "commitments on its own."),

 ("h1", "3. Your Mission Control access"),
 ("p", "You will see: program status, partner documents (including this guide and the 'Fly with Exos → Fly "
       "from Montana' one-pager), a read-only list of Montana payload prospects, and the shared tasks below, "
       "each with notes you can read and add to. You will not see Exos-internal records. Access is logged and "
       "each account is individual, so please don't share sign-ins."),
 ("n", "Sign-in: (1) open mc.exosaero.space and enter your montanaspaceport.com email. A one-time code "
       "arrives from noreply@notify.cloudflare.com (check spam). (2) Sign in with the temporary password John "
       "sends you separately. (3) Choose your own password (8+ characters). You land on the Montana partner view."),

 ("h1", "4. The shared task list"),
 ("t", [["#", "Workstream", "Montana lead", "Exos owner"],
        ["3303", "Welcome & Q&A (start here)", "All", "John Quinn"],
        ["3304", "Program Operating Agreement: path to signature", "Edgar", "John Quinn"],
        ["3305", "Task Order No. 1: site evaluation, candidate sites + site data", "Sean / Isaiah", "Neil Milburn"],
        ["3306", "Governor's signing agreement + delegation letters", "Edgar", "John Quinn"],
        ["3307", "FAA/DOT RFI FAA-2026-9736: aligned comments (due Oct 26)", "Edgar / Isaiah", "John Quinn"],
        ["3308", "Payload customer introductions", "Edgar", "Phil Eaton"],
        ["3309", "University & workforce program", "Edgar", "John Quinn"],
        ["3310", "Funding stack: federal grants + facility bonds", "Edgar", "John Quinn"]]),
 ("p", "Montana leads are our suggestion. Please reassign by posting a note on #3303."),

 ("h1", "5. Your Program Operating Agreement: questions & answers"),
 ("q", "What is Montana acquiring?"),
 ("a", "Two BLK3 reusable suborbital vehicles under the Exos National Charter Enterprise (NCE) program, on a "
       "build-operate-train-transfer model. Phase 1: BLK3.1 is built and flight-tested at Exos's current "
       "FAA-licensed site, with Montana's team trained in integration, ground operations and recovery as the "
       "build completes. Phase 2: BLK3.2, Montana-site licensing support, and a demonstration launch from "
       "Montana. Operations then transition to Montana, with Exos providing ongoing technical support under "
       "a wet-lease arrangement."),
 ("q", "What does it cost, and when do we pay?"),
 ("a", "Phase amounts, payment timing and Task Order No. 1 terms are exactly as written in the agreement. Any "
       "change is agreed in writing with John Quinn. Ask clause-specific questions on task #3304."),
 ("q", "Does the State of Montana have to appropriate money?"),
 ("a", "No. The agreement requires no state or municipal funding. The program is designed to be funded by "
       "federal grants and contracts, education partnerships, commercial launch revenue and private "
       "investment. Montana's role is to champion the program, help with federal applications, and provide "
       "the political framework for FAA site licensing. Task #3310 tracks the funding stack."),
 ("q", "What starts the clock?"),
 ("a", "Signature plus the Phase 1 funding commitment is T0 and starts the vehicle build. Task Order No. 1, the "
       "Montana launch-site evaluation, starts on execution and runs 60 days: kickoff briefing (day 14), draft "
       "report (day 45), final ranked-site report plus a one-page congressional summary (day 60)."),
 ("q", "What's still open before signature?"),
 ("a", "From our side: your legal contracting entity and address (the draft shows 'TBD'), your authorized "
       "signatory and authority, the effective date, the Phase 1 funding path, and whether the Governor's "
       "signing agreement comes before signature or runs alongside it. Please post these on #3304."),
 ("q", "Who holds which licenses?"),
 ("a", "BLK3 holds an active FAA reusable launch vehicle license under 14 CFR Part 431 (March 2026), and Exos is "
       "converting to Part 450, the current, site-flexible framework, ahead of 2027 flights. Flying from a new "
       "site in Montana is done under Part 450, which is the path we are already on. If Montana operates the "
       "spaceport site itself, that site needs an FAA launch site operator license (Part 420), and Exos "
       "provides the operator-side coordination package. The FAA gives real priority to sites with an active, "
       "licensed operator engaged. That is the leverage this partnership gives Montana."),
 ("q", "What does a launch site need?"),
 ("a", "Very little fixed infrastructure. BLK3 takes off and lands vertically at the same pad, and Exos's mobile "
       "command module brings the rest. Neil Milburn's 9/18 'BLK3 Launch Site Basics' brief to Sean has the "
       "pad and hazard-area basics. Site-specific flight-safety analysis, including wind-driven impact "
       "dispersion and downrange population, is an Exos deliverable under Task Order No. 1. Site choice is "
       "the biggest single driver of the flight-safety systems a Montana launch requires, so good candidate "
       "data early matters most."),
 ("q", "What technical information will we receive, and when?"),
 ("a", "BLK3 is U.S. export-controlled launch-vehicle technology. Detailed technical data and procedures are "
       "never shared through Mission Control or email. They are released only for program work (training, "
       "integration, operations), to U.S. persons, under an export-compliance plan Exos sets up with you "
       "before training begins. Please tell us early if any team member, university participant or partner "
       "is not a U.S. person, so we can plan for it rather than discover it later."),
 ("q", "What can we say publicly, to customers or to legislators?"),
 ("a", "Accurate and simple: Exos is one of only three FAA-licensed reusable rocket launch providers in the "
       "United States and has flown commercial payloads since 2018 (SARGE). BLK3, its latest FAA-licensed "
       "vehicle, is preparing for flight. The Montana spaceport is planned and in development. Montana "
       "customers fly on BLK3 from Exos's current site in 2027, then from Montana once the site is licensed. "
       "Joint announcements and any statement naming Exos are cleared by John Quinn first."),
 ("q", "How do payload customers work?"),
 ("a", "You make the introduction (task #3308). The Exos payload team scopes the payload and quotes the slot. "
       "Slots are sold, not donated. That revenue is what flies the vehicle. Universities typically fund "
       "through Space Grant, EPSCoR, or industry sponsors, and we help them find the path."),

 ("h1", "6. How to work with us most effectively"),
 ("b", "Put it on the task. One question or update per note, on the task it belongs to. That is the fastest "
       "route to an answer and a decision, and all three of you see the same record."),
 ("b", "Name one Montana owner per workstream. Our suggestion: Edgar for agreement, state and federal "
       "relationships; Sean for site surveys and operations; Isaiah for site data, technology, and keeping "
       "Montana's records in step with ours."),
 ("b", "Make the agreement signature-ready now. Close the entity, signatory and funding-path items on #3304, "
       "and run the Governor's letter in parallel rather than in sequence. An executed agreement with a "
       "licensed operator is what federal grant reviewers and the FAA look for first."),
 ("b", "Feed site data early. Candidate sites with coordinates, land control, access, power and comms, nearby "
       "population and airspace constraints (#3305). Task Order No. 1 can only be as good as that data."),
 ("b", "Use the federal window. File Montana's site/range-viability comment on Docket FAA-2026-9736 before "
       "Oct 26. Exos files its operator comment separately, aligned but not joint (#3307)."),
 ("b", "Stack the funding. Grants follow executed agreements and political support. Work the Governor's "
       "letter and delegation letters (#3306) alongside federal applications and facility bonds (#3310)."),
 ("b", "Build the payload pipeline now. Every warm introduction from Montana that books a 2027 slot is a "
       "customer who later flies from Montana (#3308)."),
 ("b", "Keep a short cadence. We suggest a 30-minute tag-up every two weeks, with the agenda pulled straight "
       "from the open notes on these tasks."),

 ("h1", "7. Exos contacts"),
 ("b", "John Quinn, Cofounder & CEO: agreement, commercial terms, state/federal relationships. jquinn@exosaero.com, +1-972-740-8355"),
 ("b", "Neil Milburn, Co-Founder & COO: site evaluation, FAA licensing, range and operations. nmilburn@exosaero.com"),
 ("b", "Phil Eaton, Operations Manager: payload scoping and integration. peaton@exosaero.com"),
 ("b", "Lizi Quinn: program coordination. lquinn@exosaero.com"),
 ("n", "Exos Aerospace Systems & Technologies, Inc. · 299 FM 1903, Greenville, TX 75402 · www.exosaero.com. "
       "This guide is a summary for convenience. The executed Program Operating Agreement governs."),
]


def build_docx():
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    NAVY, DARK, GREY = RGBColor(0x0A, 0x3D, 0x7C), RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x55, 0x55, 0x55)
    doc = Document(LH)
    body = doc.element.body
    for el in list(body):                      # keep only sectPr (header/footer live there)
        if el.tag != qn("w:sectPr"):
            body.remove(el)

    def run(p, text, size=10.5, bold=False, color=DARK, italic=False):
        r = p.add_run(text); r.font.name = "Calibri"; r.font.size = Pt(size)
        r.bold = bold; r.italic = italic; r.font.color.rgb = color
        return r

    def par(after=4, before=0, left=None):
        p = doc.add_paragraph(); f = p.paragraph_format
        f.space_after = Pt(after); f.space_before = Pt(before)
        if left is not None: f.left_indent = Inches(left)
        return p

    run(par(after=2), TITLE, 18, True, NAVY)
    run(par(after=10), SUB, 9.5, False, GREY, italic=True)
    for kind, val in C:
        if kind == "h1":
            run(par(after=4, before=10), val, 13, True, NAVY)
        elif kind == "p":
            run(par(), val)
        elif kind == "n":
            run(par(after=6), val, 9.5, False, GREY, italic=True)
        elif kind == "b":
            p = par(after=3, left=0.25); p.paragraph_format.first_line_indent = Inches(-0.18)
            run(p, "▪  ", 10.5, True, NAVY); run(p, val)
        elif kind == "q":
            run(par(after=1, before=5), val, 10.5, True, DARK)
        elif kind == "a":
            run(par(after=4, left=0.15), val)
        elif kind == "t":
            t = doc.add_table(rows=len(val), cols=len(val[0]))
            tblPr = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
            for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
                x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), "BFBFBF"); b.append(x)
            tblPr.append(b)
            for i, row in enumerate(val):
                for j, txt in enumerate(row):
                    c = t.cell(i, j); c.text = ""
                    run(c.paragraphs[0], txt, 9.5, i == 0, NAVY if i == 0 else DARK)
                    if i == 0:
                        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "E8EEF6")
                        c._tc.get_or_add_tcPr().append(sh)
            for row in t.rows:
                for j, w in enumerate((0.55, 3.4, 1.2, 1.35)):
                    row.cells[j].width = Inches(w)
            par(after=2)
    doc.save(OUT + ".docx"); print("wrote", OUT + ".docx")


def build_pdf():
    sys.path.insert(0, os.path.join(WS, "scripts"))
    from exos_letterhead import stationery, MARGINS
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.colors import HexColor
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from xml.sax.saxutils import escape as X
    F, FB, FI = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"
    for nm, path in (("Arial", "/System/Library/Fonts/Supplemental/Arial.ttf"),
                     ("Arial-Bold", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
                     ("Arial-Italic", "/System/Library/Fonts/Supplemental/Arial Italic.ttf")):
        if os.path.exists(path): pdfmetrics.registerFont(TTFont(nm, path))
    if "Arial" in pdfmetrics.getRegisteredFontNames(): F, FB, FI = "Arial", "Arial-Bold", "Arial-Italic"
    NAVY, DARK, GREY = HexColor("#0A3D7C"), HexColor("#1A1A1A"), HexColor("#555555")
    S = lambda **k: ParagraphStyle("x", fontName=k.pop("f", F), fontSize=k.pop("s", 10), leading=k.pop("l", 13.2), textColor=k.pop("c", DARK), **k)
    st = {"title": S(f=FB, s=17, l=21, c=NAVY, spaceAfter=2), "sub": S(f=FI, s=9, c=GREY, spaceAfter=10),
          "h1": S(f=FB, s=12.5, l=16, c=NAVY, spaceBefore=9, spaceAfter=4, keepWithNext=1), "p": S(spaceAfter=4),
          "n": S(f=FI, s=9, l=12, c=GREY, spaceAfter=6), "b": S(leftIndent=14, bulletIndent=2, spaceAfter=3),
          "q": S(f=FB, spaceBefore=5, spaceAfter=1, keepWithNext=1), "a": S(leftIndent=10, spaceAfter=4)}
    E = [Paragraph(X(TITLE), st["title"]), Paragraph(X(SUB), st["sub"])]
    for kind, val in C:
        if kind == "t":
            cell = S(s=9, l=11.5)
            data = [[Paragraph(("<b>%s</b>" % X(x)) if i == 0 else X(x), cell) for x in row] for i, row in enumerate(val)]
            t = Table(data, colWidths=[0.5*inch, 3.35*inch, 1.2*inch, 1.45*inch])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, HexColor("#BFBFBF")),
                                   ("BACKGROUND", (0, 0), (-1, 0), HexColor("#E8EEF6")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            head = E.pop()                      # bind the section heading to its table
            E += [KeepTogether([head, t]), Spacer(1, 4)]
        elif kind == "b":
            E.append(Paragraph(X(val), st["b"], bulletText="▪"))
        else:
            E.append(Paragraph(X(val), st[kind]))
    doc = SimpleDocTemplate(OUT + ".pdf", pagesize=letter, title=TITLE, author="Exos Aerospace", **MARGINS)
    doc.build(E, onFirstPage=stationery, onLaterPages=stationery)
    print("wrote", OUT + ".pdf")


if __name__ == "__main__":
    {"docx": build_docx, "pdf": build_pdf}[sys.argv[1] if len(sys.argv) > 1 else "docx"]()
