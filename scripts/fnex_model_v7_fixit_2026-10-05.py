#!/usr/bin/python3
"""FNEX model v7: apply the 10/5 'fix it' call (John Quinn / Michael Sheldon, 5:37 PM CDT) to model v6.
Input  drafts/Exos_FNEX_Model_EXOS_Redline_v6_2026-10-05.xlsx (not modified)
Output drafts/Exos_FNEX_Model_EXOS_Redline_v7_2026-10-05.xlsx"""
import os, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

WS = "/Users/exomac/.openclaw/workspace/drafts"
SRC = os.path.join(WS, "Exos_FNEX_Model_EXOS_Redline_v6_2026-10-05.xlsx")
OUT = os.path.join(WS, "Exos_FNEX_Model_EXOS_Redline_v7_2026-10-05.xlsx")
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")

wb = openpyxl.load_workbook(SRC)
aa, summ, sc, eng = wb["Additional Assumptions"], wb["Summary"], wb["Individual Scenarios"], wb["Engine"]

def mark(c, fill=GREEN):
    c.fill = fill

# 1. PIPE third-party placement agent ~7% (Michael: "closer to seven... seven is about midpoint"; John: "Agreed")
aa["A36"] = ("AGREED 10/5 fix-it: third-party placement-agent cost on the PIPE — ~7% target (midpoint; "
             "no 'in no event exceed' cap; Exos selects and pre-approves; NOT payable to FNEX)")
aa["B36"] = 0.07
mark(aa["A36"]); mark(aa["B36"])
summ["A62"] = "Less: Third-party PIPE placement-agent cost (~7% target; NO FNEX placement fee — agreed 10/5 fix-it)"
mark(summ["A62"])

# 2. ATM / later placements ~3% target, third-party only
for col in "BCD":
    sc[f"{col}48"] = 0.03; mark(sc[f"{col}48"])
sc["A48"] = "Third-party ATM / later-placement cost (% of new money post-merger) — 3% target, agreed 10/5; NOT FNEX"
mark(sc["A48"])
aa["A35"] = "Third-party ATM / later-placement cost on post-merger rounds (3% target, agreed 10/5 — NOT payable to FNEX)"
mark(aa["A35"])
eng["A47"] = "Third-party ATM / Placement Cost (3% target of new money post-merger; not FNEX)"
summ["A64"] = "Less: Third-party ATM / later-placement cost, post-merger rounds (3% target, NOT FNEX)"
mark(summ["A64"])

# 3. Expenses: pre-approved travel, reimbursed as incurred — NOT contingent on Closing
aa["A109"] = ("EXOS: FNEX Reimbursable Expenses — travel in conduct of deal business only, pre-approved in writing, "
              "$75K cap; AGREED 10/5 fix-it: reimbursed AS INCURRED (not at Closing) and owed even if Closing does not occur")
mark(aa["A109"])
summ["A65"] = "Less: Other Transaction Costs + FNEX pre-approved travel ($75K cap; paid as incurred, owed even with no Closing)"
mark(summ["A65"])

# 4. Breakup: exclusion is now 'no funding from any source'
aa["A112"] = ("EXOS: FNEX Breakup Fee — MAXIMUM accrued ($) — payable ONLY if Exos walks for a competing deal signed "
              "within 6 months; NOT payable if Exos obtains no funding from any source (fix-it 10/5), on FNEX "
              "non-performance, key-person termination, or capital from Exos's 504 / $10M Series A / strategic / government funding")
mark(aa["A112"])

# 5. NEW: 5% cash fee on FNEX-introduced strategic investment outside the PIPE (additive, cash only)
rows = [
    (114, "FIX-IT 10/5 — FNEX-INTRODUCED STRATEGIC INVESTMENT (outside the PIPE)", None),
    (115, "Strategic equity investment first introduced in writing by FNEX, closed outside the PIPE ($) — input", 0),
    (116, "FNEX fee on that investment — CASH ONLY, % of amount invested (agreed 5%; replaces the 2% stock; no stock in private co)", 0.05),
    (117, "FNEX strategic-investment fee ($) — ADDITIVE to Transaction Fees (two separate transactions)", "=B115*B116"),
    (118, "Change of Control with an FNEX-introduced party: Fees = 1% cash + 2% shares of value, IN PLACE OF the Transaction Fee (control sold once)", None),
]
for r, label, val in rows:
    aa[f"A{r}"] = label; mark(aa[f"A{r}"])
    if val is not None:
        aa[f"B{r}"] = val; mark(aa[f"B{r}"])
aa["A114"].font = Font(bold=True)
aa["B116"].number_format = "0.0%"; aa["B115"].number_format = aa["B117"].number_format = "$#,##0"

summ["A67"] = "Memo — FNEX 5% cash fee on FNEX-introduced strategic investment outside the PIPE (additive; not deducted above)"
summ["B67"] = "=-'Additional Assumptions'!B117"
summ["B67"].number_format = summ["B66"].number_format
mark(summ["A67"]); mark(summ["B67"])

# 6. Change-log tab
cl = wb.create_sheet("EXOS Changes 10-05 fix-it", 0)
lines = [
    ("EXOS REDLINE OF FNEX MODEL — v7, 2026-10-05 PM (after the 'fix it' call, John Quinn / Michael Sheldon, 5:37 PM CDT, 30 min)", None),
    ("LEGEND: GREEN = agreed on the fix-it call · YELLOW = still open. All v6 (10/5 AM) changes remain — see 'EXOS Changes 10-05'.", None),
    ("GREEN 1. PIPE third-party placement agent ≈7% target (Additional Assumptions B36: 3% → 7%). Michael: 'the pipe will probably be closer to seven… seven is about midpoint'; John: 'Agreed.' No 'in no event exceed' cap. Still NOT an FNEX fee.", GREEN),
    ("GREEN 2. ATM / later placements ≈3% target, third-party only (Individual Scenarios B48:D48: 7% → 3%; feeds Additional Assumptions B35 / Engine row 47).", GREEN),
    ("GREEN 3. FNEX travel expenses: pre-approved, $75K cap, reimbursed AS INCURRED — not 'payable at Closing', and owed even if Closing never occurs (John: 'payable at closing doesn't really make sense… you should be paid for your expenses as we go'). Amount unchanged; now a sunk cost in the downside case.", GREEN),
    ("GREEN 4. NEW: 5% CASH fee on a strategic equity investment FNEX first introduces in writing that closes outside the PIPE (Additional Assumptions B115:B117; Summary memo B67). Additive to the Transaction Fees, cash only, no stock. Default input $0.", GREEN),
    ("GREEN 5. Change of Control with an FNEX-introduced party = 1% cash + 2% shares IN PLACE OF the Transaction Fee (not additive — 'you can't sell control twice'). No model change; contract term.", GREEN),
    ("GREEN 6. Breakup fee exclusion re-cut: no breakup fee if Exos obtains no funding from ANY source (replaces 'failure to fund the PIPE at the agreed size'). Max accrued unchanged ($60K × 6 = $360K, B112). Exos's own 504 / $10M Series A / strategic / government raises never trigger it.", GREEN),
    ("GREEN 7. Fee shares: tranche 2 vests ratably over 18 months with NO 'continued advisory engagement' condition; early release at 2x merger price. Post-closing advisory = a separate future deal. No model change.", GREEN),
    ("Net to Exos at a $15M PIPE, $350M MEV (v6 → v7): $15.0M − $1.75M cash fee at Closing (0.5%) − $1.05M placement agent (7%, was $0.45M at 3%) − $0.075M travel − ~$1.0M other costs ≈ $11.1M (was ≈ $11.7M). Still clears the $10M Net Closing Cash condition by ≈ $1.1M.", YELLOW),
    ("Unchanged / still open: merger price shows $2.00 (John approved a $2.50 floor, D-7281923, pending Michael); 0.5% deferred cash tied to $40M FNEX-placed financings [TBC]; §19 PIPE definitions deferred by Michael; Schedule A dates [TBC].", YELLOW),
    ("Open in Excel to recalculate (file saved by script; cached values cleared).", None),
]
cl.column_dimensions["A"].width = 160
for i, (t, fill) in enumerate(lines, 1):
    c = cl.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top")
    if fill: c.fill = fill
    if i == 1: c.font = Font(bold=True, size=12)

wb.save(OUT)
print(OUT)
