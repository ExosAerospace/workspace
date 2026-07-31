#!/usr/bin/env python3
"""md_to_docx.py — convert a Markdown deliverable to a clean Word .docx.
Usage: python3 scripts/md_to_docx.py <input.md> [output.docx]
Jon works from Word, never .md. Lightweight MD handling (headings/bold/bullets)."""
import sys, os, re
from docx import Document
from docx.shared import Pt, RGBColor

def convert(src, dst=None):
    dst = dst or os.path.splitext(src)[0] + ".docx"
    doc = Document()
    for raw in open(src, encoding="utf-8").read().split("\n"):
        line = raw.rstrip("\n")
        if not line.strip():
            doc.add_paragraph(); continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            p = doc.add_paragraph(); r = p.add_run(m.group(2))
            r.bold = True; r.font.size = Pt(16 - 1.5*len(m.group(1)))
            r.font.color.rgb = RGBColor(0x0B,0x22,0x40); continue
        if re.match(r"^\s*[-*•]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            _emit(p, re.sub(r"^\s*[-*•]\s+","",line)); continue
        p = doc.add_paragraph(); _emit(p, line)
    doc.save(dst); return dst

def _emit(p, text):
    # bold **x**; strip inline md markers
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        r = p.add_run(seg.replace("`",""))
        if i % 2: r.bold = True

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: md_to_docx.py <input.md> [output.docx]"); sys.exit(1)
    print("wrote", convert(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None))
