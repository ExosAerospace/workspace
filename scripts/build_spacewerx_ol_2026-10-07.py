#!/usr/bin/env python3
"""SpaceWERX On-Orbit Space Logistics (OOSL) Challenge — 10/7 reviewer-lessons revision (MC task #3044).

Rebuilds the 9/23 internal drafts against the LIVE DSIP topic DAF26BX06-DV513 ("Autonomous On Orbit Logistics and
Sustainment Architecture", DoW SBIR 2026 CSO 26.BX R6, Direct-to-Phase-II only, SBIR only, ITAR-flagged, CMMC L2 Self,
closes 21 Oct 2026 12:00 ET) and the NASA SBIR AERO.4.S26B reviewer-lessons gate
(drafts/NASA_SBIR_AERO4S26B_Reviewer_Lessons_2026-10-07.docx).
Outputs (same stems as the 9/23 MC uploads, so versions chain):
  drafts/Exos_SpaceWERX_OrbitalLogistics_Draft_Submittals_2026-10-07.docx
  drafts/Exos_SpaceWERX_OrbitalLogistics_Strategy_2026-10-07.docx
Format = the 9/23 docs' internal-draft format (Letter, 0.9/0.8 in margins, banner), now repeated in a header on
every page with a page-numbered footer. PDFs via soffice --headless --convert-to pdf.
Run: /usr/bin/python3 scripts/build_spacewerx_ol_2026-10-07.py
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DARK = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x55, 0x55, 0x55)


class D:
    def __init__(self, banner):
        self.doc = Document()
        for s in self.doc.sections:
            s.left_margin = s.right_margin = Inches(0.9)
            s.top_margin = s.bottom_margin = Inches(0.8)
        st = self.doc.styles["Normal"]
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(4)
        h = self.doc.sections[0].header.paragraphs[0]
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run(banner); r.font.size = Pt(8); r.bold = True; r.font.color.rgb = GREY
        f = self.doc.sections[0].footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = f.add_run("Exos Aerospace Systems & Technologies, Inc. · Internal review draft · Page "); r.font.size = Pt(8)
        self._field(f, "PAGE")

    def _field(self, p, code):
        r = p.add_run(); r.font.size = Pt(8)
        for tag, txt in (("begin", None), (None, code), ("end", None)):
            if tag:
                r._r.append(r._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): tag}))
            else:
                e = r._r.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"}); e.text = txt; r._r.append(e)

    def p(self, *parts, size=10.5, bold=False, italic=False, align=None, after=4, color=None):
        p = self.doc.add_paragraph()
        if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if align == "j": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(after)
        for x in parts:
            b = bold
            if isinstance(x, tuple): x, b = x[0], True
            r = p.add_run(x); r.font.size = Pt(size); r.bold = b; r.italic = italic
            if color: r.font.color.rgb = color
        return p

    def title(self, t, sub=None):
        self.p(t, size=15, bold=True, color=DARK, after=2)
        if sub: self.p(sub, size=11.5, after=6)

    def h1(self, t):
        p = self.p(t, size=12.5, bold=True, color=DARK, after=3)
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True

    def h2(self, t):
        p = self.p(t, size=11, bold=True, after=2)
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True

    def body(self, *parts):
        return self.p(*parts, align="j")

    def quote(self, t):
        p = self.p(t, size=9.5, italic=True, align="j")
        p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.right_indent = Inches(0.35)

    def bullet(self, t, lead=None):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        if lead:
            r = p.add_run(lead + " "); r.bold = True; r.font.size = Pt(10.5)
        r = p.add_run(t); r.font.size = Pt(10.5)

    def table(self, rows, widths, size=9):
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        for j, w in enumerate(widths): t.columns[j].width = Inches(w)
        for i, row in enumerate(rows):
            tr = t.rows[i]._tr.get_or_add_trPr(); tr.append(tr.makeelement(qn("w:cantSplit"), {}))
            for j, v in enumerate(row):
                c = t.cell(i, j); c.width = Inches(widths[j])
                cp = c.paragraphs[0]; cp.paragraph_format.space_after = Pt(0)
                r = cp.add_run(str(v)); r.font.size = Pt(size); r.bold = (i == 0)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def rule(self):
        self.p("_" * 92, size=8, after=2)

    def save(self, path):
        self.doc.save(path); print("WROTE", path)


TOPIC = "DAF26BX06-DV513"
CLOSE = "21 October 2026, 12:00 p.m. ET"

# =====================================================================================================
# 1) DRAFT SUBMITTALS -> D2P2 TECHNICAL VOLUME DRAFT (primary) + reserve outline
# =====================================================================================================
d = D("EXOS AEROSPACE — INTERNAL REVIEW DRAFT · NOT FOR EXTERNAL RELEASE · REVISED BY EXO 2026-10-07 "
      "(supersedes 2026-09-23)")
d.title("SpaceWERX On-Orbit Space Logistics Challenge — Draft Submittal (revised)",
        f"DSIP topic {TOPIC} · Direct-to-Phase-II (D2P2) · SBIR · window closes {CLOSE}")
d.p("Review draft for John. Capability-level only; no controlled specifications. [____] = input only John (or the "
    "named owner) can supply. Part 0 is internal and is removed before upload; Part 1 is the technical-volume draft.",
    size=10)
d.rule()

# ---------------- PART 0 (internal) ----------------
d.h1("Part 0 — Pre-submission gate (internal; delete before upload)")
d.body("Applied the NASA SBIR AERO.4.S26B reviewer-lessons gate (10/7) to the 9/23 drafts after pulling the live "
       f"topic text for {TOPIC} from DSIP. Three findings changed the plan:")
d.bullet("The topic is SBIR Direct-to-Phase-II only (no Phase I, no STTR). Draft C as an STTR is not possible; it is "
         "re-scoped below as an optional SBIR entry.", "D2P2 + SBIR only.")
d.bullet("Phase I-equivalent work must be documented: stakeholder engagement with a USSF customer, a feasibility study, "
         "an integration pathway, regulatory considerations, commercial dual use, and a Phase II prototype plan.",
         "Phase I-equivalent proof is mandatory.")
d.bullet("The challenge page states that identical submissions may not be applied to multiple focus areas. The 9/23 "
         "Submittal B (OTV focus area) is dropped: that focus area requires an OTV prototype (modular chassis, "
         "autonomous navigation, multi-stop routing), and Exos does not build OTVs. Its tank-set content is folded "
         "into the primary entry as OTV compatibility.", "One distinct entry per focus area.")
d.table([
    ["#", "Gate item", "Status", "What changed / what is open"],
    ["1", "Topic fit", "FIXED", "Topic and focus-area text quoted (Sec. 1); every objective mapped. Lead use = propellant "
     "storage, condition monitoring, metering and custody (Focus Area 1), not launch."],
    ["2", "Structure", "OPEN", "Ordered to the topic's required Phase I-equivalent elements and the five focus-area items. "
     "The DAF D2P2 instructions (DAF_SBIR_26BX_D2P2_R6_v2.pdf) could not be downloaded outside DSIP: confirm the "
     "volume list, page limit, and TOC there and re-order Part 1 verbatim. [____ owner: Exo with John's DSIP login]"],
    ["3", "Narrative", "FIXED", "Technical approach and work plan written as prose (Sec. 5), not only tables."],
    ["4", "People", "OPEN", "Key personnel written from verified bios. PI must be an Exos employee: recommended Neil "
     "Milburn [____ John confirm PI]. Scorpius lead engineer bio [____]."],
    ["5", "Work plan", "FIXED", "Eight numbered tasks with owner, months, deliverable, success criterion; schedule table. "
     "Numeric targets are proposed [____ Neil confirm]."],
    ["6", "Budget", "OPEN", "No dollars in the technical volume. One source of truth = DSIP cost volume: $[____]. Travel "
     "limited to two trips. Scorpius share must keep Exos at ≥50% of the work [____ confirm split]. Commission "
     "credit on Scorpius purchases must be shown as a credit in the cost volume."],
    ["7", "Dependencies", "FIXED / OPEN", "No Government facility on the critical path (Exos Texas test site + Scorpius). "
     "No IP license needed: Scorpius supplies tanks as a subcontractor and keeps its background IP [____ teaming "
     "agreement + letter]. BLK3 flight demo is optional, behind a go/no-go."],
    ["8", "Maturity", "FIXED", "Risk table with cited evidence; early go/no-go at month 4 (sensor survival) and month 12 "
     "(flight-demo decision)."],
    ["9", "Signatures / endorsements", "OPEN (Jon)", "DSIP cover-sheet certifications by the authorized official, "
     "foreign-affiliation disclosure (incl. Exos Italia Srl), fraud/waste/abuse certification, Scorpius letter, "
     "and Customer Memorandum if R6 requires one [____ all John]."],
], [0.3, 1.2, 0.9, 4.3], size=8.5)
d.h2("Overclaim and fact corrections made to the 9/23 text")
for t in [
    "\"Exos holds 18% + ROFR\" removed: the ROFR has expired. Scorpius = affiliate in which Exos holds an 18% interest.",
    "\"Exos flies and reflies BLK3 with fast (2-day) turnaround\" removed: BLK3 has not flown; a two-day turnaround is a "
    "design goal, not demonstrated.",
    "\"NASA-SBIR FBG structural-health-monitoring line\" removed: that NASA SBIR proposal was not selected. FBG sensing "
    "is now a Phase II task behind a month-4 go/no-go, with a fallback.",
    "\"Microcracking defeated\" changed to \"testing to date has shown no microcracking\"; the RT-cure resin system and "
    "composite fabrication are credited to Scorpius, not Exos.",
    "\"Zero-boil-off\" removed; boil-off is measured, not promised.",
    "Nebula/Intuitive Machines OTV contract named only with consent [____]; the IM-1/IM-2 tank heritage is public and "
    "cited.",
    "Exos is positioned as an enabling-technology supplier, not an OTV builder or propulsion provider.",
]:
    d.bullet(t)
d.h2("Blanks only John (or the named owner) can fill")
for t in [
    "[____] PI (recommended: Neil Milburn, COO).",
    "[____] Proposed amount and period of performance (draft assumes 15 months).",
    "[____] USSF stakeholder engagement record (names, offices, dates) beyond the RSLP RIDE Industry Day.",
    "[____] Customer Memorandum: confirm whether the R6 D2P2 instructions require one and, if so, the signer.",
    "[____] Scorpius teaming letter, quote, and lead engineer; Exos/Scorpius work split (Exos ≥50%).",
    "[____] Princeton CryoTech: include as subcontractor (letter + quote) or drop.",
    "[____] Name Intuitive Machines / Nebula as a commercial reference: yes or no.",
    "[____] SPRS score / CMMC Level 2 (Self) assessment on file (topic requirement).",
    "[____] Neil: export (ITAR/EAR) review of the final volume; confirm numeric test targets in Sec. 5.",
    "[____] John: DSIP certifications and submit. John is in Japan from 8 Oct; set who files and when.",
]:
    d.bullet(t)
d.rule()

# ---------------- PART 1 technical volume ----------------
d.h1("Part 1 — Technical Volume (D2P2) draft")
d.table([
    ["Topic", f"{TOPIC} — Autonomous On Orbit Logistics and Sustainment Architecture (DoW SBIR 2026 CSO 26.BX)"],
    ["Proposal title", "Instrumented Linerless Composite Cryogenic Tank Sets for Depot and OTV Propellant Storage, "
                       "Condition Monitoring, and Custody Tracking"],
    ["Focus area / framework", "Focus Area 1, Propellant Management · Logistics Function approach"],
    ["Offeror", "Exos Aerospace Systems & Technologies, Inc., 299 FM 1903, Greenville, TX 75402 · CAGE 7EMY6 · "
                "UEI CK24BPZGF854 · small business"],
    ["Principal Investigator", "[____ recommended: Neil Milburn, COO]"],
    ["Period of performance", "15 months (proposed) [____]"],
    ["Cost", "Stated in the cost volume only: $[____]"],
], [1.6, 5.1], size=9.5)

d.h2("1. Topic fit")
d.body("The topic asks for prototype-ready technology that supports routine refueling and resupply of U.S. space "
       "assets. Focus Area 1 states:")
d.quote("\"A resilient on-orbit logistics enterprise requires the ability to store, transfer, meter, monitor and "
        "manage propellant across multiple orbital regimes. Relevant propellant classes may include cryogenic "
        "propellants ... D2P2 prototypes should demonstrate one or more of the following: autonomous propellant "
        "transfer methods, standardized or interoperable refueling interfaces, depot integration, propellant "
        "distribution operations, propellant quality monitoring ...\" Enabling technologies listed include "
        "\"Propellant storage and sustainment technologies, including boil-off mitigation\" and \"Sensors and "
        "diagnostics for propellant condition monitoring.\" (DSIP topic " + TOPIC + ")")
d.body("Exos uses the Logistics Function framework. The function is cryogenic propellant storage at a depot or on an "
       "OTV, with continuous condition monitoring, metering, and custody tracking. Each Phase II objective maps to "
       "the topic text:")
d.table([
    ["Phase II objective", "Topic language it answers"],
    ["O1. Instrumented flight-heritage cryogenic tank set (storage node)", "Propellant storage and sustainment; boil-off mitigation"],
    ["O2. Propellant condition and structural-health sensing on the tank", "Propellant quality monitoring; sensors and diagnostics for propellant condition"],
    ["O3. Mass gauging, metering, and custody record for fill, hold, and transfer", "Precision metering and custody tracking during storage and transfer"],
    ["O4. Autonomous fill/hold/transfer sequencing with anomaly response", "Autonomous control systems for fueling safety, sequencing, and anomaly response"],
    ["O5. Depot- and OTV-compatible interfaces and data", "Compatibility with depot operations, OTVs, and client spacecraft"],
], [3.3, 3.4], size=9)

d.h2("2. Identification and significance of the problem")
d.body("Spacecraft launch with a fixed propellant load. A depot or OTV that stores cryogens for resupply must hold them "
       "for long periods, know how much is on board, show that it is in good condition, and pass it on with a "
       "custody record. Tanks drive dry mass and leak paths. Composite cryogenic tanks save mass but have historically "
       "carried a risk of microcracking under thermal cycling, which is hard to detect in service. A depot operator "
       "today has no low-mass tank set that reports its own structural state and propellant condition. That gap "
       "limits how much a logistics planner can trust a stored propellant inventory, which is the planning input the "
       "In-Domain Logistics enterprise needs.")

d.h2("3. Focus Area 1 required content")
d.bullet("Exos's Phase II work uses liquid nitrogen as a safe cryogenic simulant for test. The tank architecture flew "
         "with liquid oxygen and liquid methane on the Intuitive Machines IM-1 and IM-2 landers (Sec. 4.2), both "
         "commercially available terrestrial propellants. In-space resource use is outside this Phase II scope.",
         "Supply chain and sourcing.")
d.bullet("Linerless Type-V composite tanks store the cryogen. Phase II measures boil-off over defined hold periods "
         "against a thermal model, rather than claiming zero boil-off. Interfaces are designed to accept an existing "
         "or emerging standard refueling interface [____ select interface with the USSF stakeholder in Task 1].",
         "Transfer and storage.")
d.bullet("Fiber-optic temperature and strain sensing plus mass gauging give a running inventory and condition record. "
         "Each fill, hold, and transfer event is logged with mass, temperature history, and structural-state data, "
         "so a receiving node gets the propellant's custody history. Contamination control follows standard cryogenic "
         "cleanliness practice for the test set.", "Metering, quality control, and custody.")
d.bullet("Exos addresses the storage, quality-management, and metering portion of the value chain. The tank set "
         "serves depots and reusable OTVs; it does not replace either.", "Value chain role and integration.")
d.bullet("Lower tank dry mass leaves more mass for propellant delivered. Self-reporting tanks let a planner trust the "
         "inventory without extra inspection sorties. Both raise the maneuver endurance that the topic seeks.",
         "Economics and military utility.")

d.h2("4. Phase I-equivalent work (required by the topic)")
d.h2("4.1 Stakeholder engagement")
d.body("Exos attended the Space Systems Command Space Access RIDE Industry Day (Kirtland AFB / virtual, 22 September "
       "2026), where the Space Access portfolio presented Servicing, Mobility & Logistics as one of its three mission "
       "pillars. [____ John: list direct engagement with SSC SYD 80 Space Mobility & Logistics (AL) or other USSF "
       "logistics stakeholders: names, offices, dates, and what they said they need. The topic requires a clear USSF "
       "customer; this section fails the gate until it is filled.]")
d.h2("4.2 Feasibility study")
d.body("The feasibility study (attached as [____ Exo to draft; Neil to approve]) rests on flown hardware. Scorpius "
       "Space Launch Company's Pressurmaxx all-composite Type-V cryogenic propellant tanks flew on Intuitive Machines' "
       "IM-1 lander, which landed on the Moon in February 2024, and on IM-2 (Athena), launched 26 February 2025 "
       "(Scorpius press releases, 14 March 2024 and 3 March 2025; CompositesWorld, \"Scorpius Type 5 tanks underpin "
       "second Intuitive Machines mission to the moon,\" 2025). Testing to date of the room-temperature-cure resin "
       "system used in these tanks has shown no microcracking. The study will cover thermal and structural analysis "
       "of the Phase II tank set, sensor survivability, gauging accuracy, and integration with depot and OTV "
       "interfaces, with the technical risks listed in Section 6.")
d.h2("4.3 Integration pathway")
d.body("The tank set integrates as the storage element of a depot or OTV. Its data products (inventory, condition, "
       "custody log) are delivered in an open format so a depot or fleet health-monitoring system can ingest them. "
       "Task 1 sets the interface and data standard with the USSF stakeholder.")
d.h2("4.4 Regulatory considerations")
d.body("The topic is ITAR-designated. Exos is a U.S.-owned small business and will handle technical data under ITAR "
       "and EAR, with an export review before release. Ground cryogenic testing follows Exos's existing test-site "
       "safety procedures. Any flight test on BLK3 goes through Exos's FAA launch license and payload review. "
       "Exos will meet the topic's CMMC Level 2 (Self) requirement [____ SPRS score on file].")
d.h2("4.5 Commercial dual-use potential")
d.body("The same instrumented tank set serves commercial depot operators, OTV builders, lunar lander builders, and "
       "satellite manufacturers. Scorpius tanks are already sold commercially to lunar lander and OTV customers "
       "[____ name Intuitive Machines/Nebula only with consent]. Self-reporting tanks also give insurers and "
       "operators condition data for life-extension decisions.")
d.h2("4.6 Phase II prototype plan")
d.body("Section 5 gives the plan: an instrumented two-tank storage and transfer module, ground-tested with liquid "
       "nitrogen for hold, boil-off, gauging, transfer, and thermal cycling, with an optional suborbital "
       "low-gravity transfer test on BLK3 after a go/no-go.")

d.h2("5. Technical approach and work plan")
d.body("Exos will build and test a prototype storage-and-transfer module made of two Scorpius Type-V tanks with "
       "fiber-optic sensing, mass gauging, and Exos-developed control and data software. Work starts with a "
       "low-cost coupon and sub-scale test to prove the sensors survive cryogenic cycling before any full-size "
       "hardware is bought. If the sensors fail that test, Exos falls back to conventional temperature and strain "
       "gauges and capacitance gauging, and the program continues. The tank set is then designed, built, and "
       "ground-tested at Exos's Texas test site with liquid nitrogen. Tests measure boil-off over set hold "
       "periods, gauging accuracy against weighed reference loads, tank-to-tank transfer, and structural state "
       "across repeated thermal cycles. Software runs fill, hold, and transfer automatically and flags anomalies. "
       "A month-12 review decides whether to fly a low-gravity transfer test on BLK3. That flight is optional: the "
       "ground data alone completes the prototype, so the program does not depend on a flight date.")
d.table([
    ["Task", "Objective / method", "Owner", "Months", "Deliverable", "Success criterion (proposed)"],
    ["1", "Kickoff; requirements, interface and data standard with USSF stakeholder", "PI", "1–2",
     "Requirements and interface document", "Stakeholder concurs with baseline"],
    ["2", "Sensor survival: fiber-optic sensors on Type-V coupons and a sub-scale bottle, LN2 cycling", "PI + Scorpius",
     "2–4", "Coupon test report; GO/NO-GO 1", "Sensors read through 50 cycles within tolerance of reference gauges"],
    ["3", "Tank-set and test-module design (two tanks, transfer line, interfaces)", "CTO", "3–6",
     "Design review packages", "Design closes with margins on proof pressure and thermal loads"],
    ["4", "Fabricate and proof-test tanks; install sensors", "Scorpius + PI", "6–9",
     "Two instrumented tanks; proof-test records", "Both tanks pass proof test with no leak"],
    ["5", "Control and data software: autonomous sequencing, anomaly flags, custody log", "CTO", "6–13",
     "Software build and demo", "Runs a full fill-hold-transfer cycle without manual steps"],
    ["6", "Ground cryogenic test campaign (hold, boil-off, gauging, transfer, cycling)", "COO", "9–12",
     "Test report; GO/NO-GO 2", "Gauging error ≤2% of full load; boil-off within model band; no leak after cycling"],
    ["7", "Optional: suborbital low-gravity transfer test on BLK3", "COO", "13–15",
     "Flight data package (if flown)", "Transfer data recorded; hardware recovered"],
    ["8", "Integration pathway, transition plan, final report", "PI", "13–15",
     "Final report; Phase III plan", "Delivered and accepted"],
], [0.4, 1.95, 0.75, 0.55, 1.35, 1.7], size=8)
d.h2("Schedule")
qs = ["M1–3", "M4–6", "M7–9", "M10–12", "M13–15"]
span = {1: [0], 2: [0, 1], 3: [0, 1], 4: [1, 2], 5: [1, 2, 3, 4], 6: [2, 3], 7: [4], 8: [4]}
rows = [["Task"] + qs]
for k in range(1, 9):
    rows.append([str(k)] + ["■" if i in span[k] else "" for i in range(5)])
rows.append(["Gates", "", "GO/NO-GO 1 (M4)", "", "GO/NO-GO 2 (M12)", "Final"])
d.table(rows, [0.6, 1.2, 1.2, 1.2, 1.2, 1.2], size=8.5)
d.body("Travel is limited to two trips: kickoff with the USSF stakeholder and the final demonstration review. Other "
       "reviews are virtual.")

d.h2("6. Risks, evidence, and early decisions")
d.table([
    ["Risk", "Evidence today", "Mitigation", "Decision point"],
    ["Sensors do not survive cryogenic cycling or bonding", "Fiber-optic sensing of composite structures is not yet "
     "proven on these tanks", "Coupon and sub-scale test first; fallback to conventional gauges", "GO/NO-GO 1, month 4"],
    ["Composite tank leak or microcracking under cycling", "Tank type flew on IM-1 and IM-2 (cited Sec. 4.2); no "
     "microcracking seen in resin testing to date", "Proof test every tank; cycle test with sensing", "Task 4 proof test"],
    ["Single-source tank supplier schedule", "Scorpius is the tank source", "Order long-lead items in month 3; "
     "sub-scale article before full size; Exos witnesses proof tests", "Month 6 design review"],
    ["Flight-test availability", "BLK3 first flight campaign planned for early 2027", "Flight test optional; ground "
     "data completes the prototype", "GO/NO-GO 2, month 12"],
], [1.6, 1.9, 1.95, 1.25], size=8.5)

d.h2("7. Related work")
d.body("Verified Exos Air Force contracts (USASpending.gov): FA9300-25-C-6022, SBIR Phase II, Suborbital Hover "
       "Vehicle–Reusable Rocket (2025); FA8649-21-P-0838, AFWERX SBIR Phase II, reusable hypersonic launch-vehicle "
       "prototype (2021); FA9300-24-P-6006, SBIR Phase I (2024); FA8649-20-P-0799, AFWERX SBIR Phase I (2020); and "
       "FA9300-26-C-6011, AFRL Project Hermes (executed September 2026). Exos holds an FAA launch license for its "
       "SARGE/BLK3 vehicle family; SARGE lifted off four times from Spaceport America in 2018–2019.")

d.h2("8. Key personnel")
d.bullet("[____ confirm as PI] Cofounder and COO. Longtime Armadillo Aerospace veteran; Exos's FAA liaison; leads "
         "operations, program execution, and Exos's composite-tank technical work with Scorpius. Proposed role: PI, "
         "Tasks 1, 2, 4, 6–8.", "Neil Milburn.")
d.bullet("Cofounder and CTO. Has built hundreds of rocket engines and dozens of flying vehicles; owns Exos vehicle "
         "and fluid-system design. Role: Tasks 3 and 5.", "Russell Blink.")
d.bullet("Cofounder and CEO. Eleven years building Exos; 14.5 years in the U.S. Navy Submarine Service. Role: "
         "program executive, USSF stakeholder lead, transition.", "John Quinn.")
d.bullet("[____ name, title, years, directly related tank work; letter of commitment].", "Scorpius lead engineer.")

d.h2("9. Facilities, subcontractors, and work share")
d.body("Exos: engineering, integration, and software in Greenville, Texas; cryogenic and engine test site at Caddo "
       "Mills, Texas. Scorpius Space Launch Company (Torrance, California; Exos affiliate, 18% Exos interest): tank "
       "fabrication and proof test under subcontract. [____ Princeton CryoTech: cryogenic instrumentation, if "
       "teamed.] No Government facility is on the critical path. Exos performs at least 50% of the work [____ "
       "confirm against the cost volume].")

d.h2("10. Transition and Phase III")
d.body("Phase III moves the tank set into USSF depot and OTV programs and commercial sales. Exos sells Scorpius tank "
       "hardware and integrates the sensing, gauging, and custody system. Targets are USSF Space Access / Space "
       "Mobility & Logistics programs, depot and OTV builders, and satellite manufacturers.")

d.h2("11. Prior, current, or pending support")
d.body("[____ John: disclose any pending proposal with overlapping scope, e.g., cryogenic-transfer work pitched to "
       "NASA Flight Opportunities or TACFI. State none if none.]")
d.rule()

# ---------------- reserve entry ----------------
d.h1("Reserve entry (optional, separate focus area) — Logistics Health Monitoring")
d.body("Title: Self-Monitoring Composite Pressure Structures with Serialized Health Records. Focus Area 6 asks for "
       "\"novel materials for monitorable components\" and \"data fusion across sensor data, maintenance logs, and "
       "digital twins, with serialized pedigree and historical tracking.\" This entry would center on the health "
       "record and digital twin of a fleet of composite pressure vessels, not on propellant storage, so it is not "
       "identical to the primary entry. SBIR D2P2 only (no STTR on this topic). File only if the primary entry is "
       "complete and a second Phase I-equivalent package and stakeholder can be documented.")
d.h2("Dropped: 9/23 Submittal B (OTV focus area)")
d.body("Focus Area 2 requires an OTV prototype with modular chassis, autonomous navigation, and multi-stop routing. "
       "Exos does not build OTVs, so the entry would fail topic fit. OTV compatibility is objective O5 of the "
       "primary entry instead.")
d.save("drafts/Exos_SpaceWERX_OrbitalLogistics_Draft_Submittals_2026-10-07.docx")

# =====================================================================================================
# 2) STRATEGY BRIEF (revised)
# =====================================================================================================
s = D("EXOS AEROSPACE — INTERNAL REVIEW DRAFT · NOT FOR EXTERNAL RELEASE · REVISED BY EXO 2026-10-07 "
      "(supersedes 2026-09-23)")
s.title("SpaceWERX On-Orbit Space Logistics Challenge",
        "Opportunity evaluation and recommended slate — revised against the live DSIP topic and the 10/7 lessons gate")
s.rule()
s.h1("Bottom line")
s.body("Still a strong fit, but narrower than the 9/23 brief said. The live topic, ", (TOPIC,), " (\"Autonomous On "
       "Orbit Logistics and Sustainment Architecture\"), is SBIR Direct-to-Phase-II only, ITAR-designated, and "
       f"requires CMMC Level 2 (Self). It closes {CLOSE}. Recommendation: file one entry in Focus Area 1, Propellant "
       "Management, built on flight-heritage Scorpius Type-V tanks with Exos sensing, gauging, custody tracking, and "
       "autonomous sequencing. Hold an optional second entry in Focus Area 6, Logistics Health Monitoring. Drop the "
       "OTV entry. The biggest gap is the topic's mandatory Phase I-equivalent proof, above all a documented USSF "
       "stakeholder.")
s.h1("What changed since 9/23")
s.table([
    ["9/23 brief said", "Corrected (source)"],
    ["Submit A and B under SBIR, C under STTR", "Topic is SBIR D2P2 only; no STTR, no Phase I (DSIP topic text)"],
    ["Five named focus areas", "Six focus areas; identical submissions may not span focus areas (afwerxchallenge.com/on-orbit)"],
    ["B: OTV tank sets", "OTV focus area requires an OTV prototype; Exos does not build OTVs; dropped"],
    ["Exos holds 18% + ROFR", "ROFR expired; Exos holds an 18% interest; commission-only sales role (Jon, 10/6)"],
    ["BLK3 flies and reflies, 2-day turnaround", "BLK3 not yet flown; first flight campaign early 2027; 2-day turnaround is a design goal"],
    ["NASA SBIR FBG line", "That NASA SBIR proposal was not selected (10/5); FBG now behind a go/no-go"],
    ["Microcracking defeated; zero boil-off", "\"No microcracking in testing to date\"; boil-off measured, not promised"],
    ["~$20M program", "Program press figures vary ($20M SSC, May 2026; $24M National Defense, Sept 2026); amount per award [____]"],
], [2.6, 4.1], size=9)
s.h1("Focus-area fit (six areas)")
s.table([
    ["Focus area", "Exos fit", "Decision"],
    ["1 Propellant Management", "Strong: storage, condition monitoring, metering, custody on flown tank type", "PRIMARY entry"],
    ["2 Orbital Transfer Vehicles", "Weak: needs an OTV prototype", "Drop; OTV compatibility folded into entry 1"],
    ["3 Network Mechanics", "Weak: not an Exos lane", "No entry"],
    ["4 Orbital Warehousing", "Weak: needs depot construction and autonomous material handling", "No entry"],
    ["5 Reuse, Repair, Refuel, Recovery", "Partial: tanks are servicing components, but focus is robotic repair", "No entry"],
    ["6 Logistics Health Monitoring", "Moderate: self-monitoring composite components, serialized health records", "Optional second entry"],
], [2.0, 3.2, 1.5], size=9)
s.h1("Team")
s.bullet("Exos (prime, PI, ≥50% of work): sensing, gauging, custody software, integration, ground test, optional "
         "BLK3 flight test.")
s.bullet("Scorpius Space Launch Company (subcontractor; Exos holds an 18% interest): Type-V tank fabrication and proof "
         "test. Scorpius keeps its background IP, so no license is on the critical path. Needs a teaming letter "
         "and quote. Exos's 15% sales commission on Scorpius hardware must be credited in the cost volume.")
s.bullet("Princeton CryoTech (optional subcontractor): cryogenic instrumentation. Include only with a letter and "
         "quote.")
s.bullet("Intuitive Machines: commercial reference only, named only with consent.")
s.h1("Compliance and registration")
for t in [
    "DSIP account, SAM.gov registration (CAGE 7EMY6 / UEI CK24BPZGF854), SBA SBC Control ID: confirm current.",
    "CMMC Level 2 (Self) required by the topic: confirm the SPRS score is posted.",
    "ITAR-designated topic: Neil's export review of the final volume before upload.",
    "Foreign-affiliation disclosure: include Exos Italia Srl (hibernating).",
    "Instructions: read DAF_SBIR_26BX_D2P2_R6_v2.pdf in DSIP for the volume list, page limit, and whether a Customer "
    "Memorandum is required; re-order the technical volume to it verbatim.",
    "Final check: download the compiled DSIP package and read it page by page; have someone who did not write it "
    "score it against the topic.",
]:
    s.bullet(t)
s.h1("Timeline (John in Japan from 8 Oct)")
s.table([
    ["Date", "Action", "Owner"],
    ["By 8 Oct", "Go/no-go on filing; pick PI; authorize filer; list USSF engagements", "John"],
    ["9–12 Oct", "Feasibility study draft; R6 instructions reconciled; Scorpius letter and quote requested", "Exo / Neil"],
    ["13–16 Oct", "Cost volume (single source of truth); work split ≥50% Exos; export review", "Exo / Kyle / Neil"],
    ["17–19 Oct", "Package compiled; independent read-through; certifications", "Exo / John (remote)"],
    ["20 Oct", "File with one day of buffer before the 21 Oct 12:00 ET close", "Authorized filer"],
], [1.1, 4.3, 1.3], size=9)
s.h1("Decisions needed from John")
for t in [
    "File the Focus Area 1 entry: yes or no. Optional Focus Area 6 entry: yes or no.",
    "PI (recommended Neil Milburn), proposed amount $[____], and period of performance.",
    "USSF stakeholder engagement to cite (the topic requires a USSF customer).",
    "Name Intuitive Machines / Nebula: yes or no. Include Princeton CryoTech: yes or no.",
    "Who files in DSIP while John is in Japan.",
]:
    s.bullet(t)
s.p("Sources: DSIP topic " + TOPIC + " (pulled 7 Oct 2026); afwerxchallenge.com/on-orbit; SSC newsroom (May 2026); "
    "Scorpius press releases (14 Mar 2024, 3 Mar 2025); CompositesWorld (2025); USASpending.gov (Exos awards).",
    size=8.5, italic=True)
s.save("drafts/Exos_SpaceWERX_OrbitalLogistics_Strategy_2026-10-07.docx")
