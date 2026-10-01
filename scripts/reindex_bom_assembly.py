#!/usr/bin/env python3
"""reindex_bom_assembly.py — index a vehicle BOM workbook (Exos BOM layout: one sheet per section,
header on row 2: Item | Quantity | Vendor P/N | Manufacturer P/N | Description | Vendor | 2nd | 3rd | Cost ea)
into mission-control/data/bom-index-<assembly>.json for the ERP MRP loader (erp-foundation/mrp.js).
The BLK3 index (bom-index.json) is still built by inventory_capture.py reindex-bom; this handles the other
vehicles (HVRR, ...) listed in mission-control/data/mrp-assemblies.json.
Usage: python3 scripts/reindex_bom_assembly.py --assembly HVRR [--xlsx path]
"""
import argparse, json, os
from datetime import datetime, timezone
WS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(WS, "mission-control", "data")

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--assembly", required=True); ap.add_argument("--xlsx")
    a = ap.parse_args()
    cfg = json.load(open(os.path.join(DATA, "mrp-assemblies.json")))
    asm = next((x for x in cfg["assemblies"] if x["id"] == a.assembly), None)
    if not asm: raise SystemExit(f"{a.assembly} not in mrp-assemblies.json")
    xlsx = a.xlsx or os.path.join(WS, asm["source_xlsx"])
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    rows = []
    for sheet in wb.sheetnames:
        if "top level" in sheet.lower():
            continue   # top-level sheet lists the section sub-assemblies themselves, not parts
        for r in wb[sheet].iter_rows(min_row=3, values_only=True):
            if not r: continue
            g = lambda i: (str(r[i]).strip() if i < len(r) and r[i] is not None else "")
            item, qty, vpn, mpn, desc, vendor, cost = g(0), g(1), g(2), g(3), g(4), g(5), g(8)
            if not (vpn or mpn or desc): continue
            rows.append({"section": sheet, "item": item, "qty": qty, "vendor_pn": vpn, "mfg_pn": mpn,
                         "description": desc, "vendor": vendor, "cost": cost})
    out = os.path.join(DATA, asm["bom_index"])
    json.dump({"assembly": a.assembly, "source": os.path.basename(xlsx), "generated_at": datetime.now(timezone.utc).isoformat(),
               "count": len(rows), "rows": rows}, open(out, "w"), indent=1)
    print(f"{a.assembly}: indexed {len(rows)} BOM lines -> {out}")

if __name__ == "__main__":
    main()
