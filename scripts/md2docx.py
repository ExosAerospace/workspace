import re, sys
from docx import Document
from docx.shared import Pt

SRC=sys.argv[1] if len(sys.argv)>1 else "/Users/exomac/.openclaw/workspace/DARPA_DSO_HR001125S0013_ExecSummary_FINAL.md"
OUT=sys.argv[2] if len(sys.argv)>2 else SRC.rsplit('.',1)[0]+".docx"

doc=Document(); st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(11)
def add_runs(p,text):
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r=p.add_run(part[2:-2]); r.bold=True
        else: p.add_run(part)
lines=open(SRC).read().split('\n'); i=0
isT=lambda l:l.strip().startswith('|')
while i<len(lines):
    line=lines[i].rstrip(); s=line.strip()
    if s=='' or s=='---': i+=1; continue
    if isT(line):
        rows=[]
        while i<len(lines) and isT(lines[i]):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]); i+=1
        rows=[r for r in rows if not all(set(c)<=set('-: ') for c in r)]
        if rows:
            n=len(rows[0]); t=doc.add_table(rows=0,cols=n); t.style='Light Grid Accent 1'
            for ri,r in enumerate(rows):
                cells=t.add_row().cells
                for ci in range(n):
                    cells[ci].text=''; p=cells[ci].paragraphs[0]; add_runs(p, r[ci] if ci<len(r) else '')
                    if ri==0:
                        for run in p.runs: run.bold=True
        continue
    if s.startswith('#'):
        m=re.match(r'(#+)\s+(.*)',s); lvl=len(m.group(1)); txt=re.sub(r'\*\*(.+?)\*\*',r'\1',m.group(2))
        doc.add_heading(txt, level=0 if lvl==1 else min(lvl-1,4)); i+=1; continue
    if s.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); add_runs(p,s[2:]); i+=1; continue
    m=re.match(r'(\d+)\.\s+(.*)',s)
    if m: p=doc.add_paragraph(style='List Number'); add_runs(p,m.group(2)); i+=1; continue
    if s.startswith('_') and s.endswith('_'):
        p=doc.add_paragraph(); r=p.add_run(s.strip('_')); r.italic=True; i+=1; continue
    p=doc.add_paragraph(); add_runs(p,s); i+=1
doc.save(OUT); print("WROTE",OUT)
