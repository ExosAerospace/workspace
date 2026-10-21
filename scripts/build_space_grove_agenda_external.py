#!/usr/bin/env python3
"""External (send-to-Melissa) version of the Exos x Space Grove collaboration agenda, on Exos letterhead.
Internal notes (background check, guardrails, John-only) are deliberately omitted. Body built here, then
scripts/apply_exos_letterhead.py puts the PRIMARY letterhead header/footer on every page."""
import subprocess, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

WS = "/Users/exomac/.openclaw/workspace"
BODY = f"{WS}/drafts/uplink/_body_SpaceGrove_Agenda_external.docx"
OUT = f"{WS}/drafts/uplink/Exos_SpaceGrove_Collaboration_Agenda_External_2026-10-21.docx"
NAVY, GREY = (0, 51, 102), (80, 80, 80)


def run(p, text, bold=False, size=11, color=None, italic=False):
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    return r


def h(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
    run(p, text, True, 12.5, NAVY)


def sub(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True
    run(p, text, True, 11, NAVY)


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            run(p, it[0], bold=True); run(p, it[1])
        else:
            run(p, it)


def build():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

    for text, size, bold, color in [
        ("Exos Aerospace × Space Grove Ventures", 17, True, NAVY),
        ("Collaboration Plan — Meeting Agenda", 13, True, NAVY),
        ("Wednesday, October 21, 2026  |  3:45 PM ET  |  Up/Link: The SpaceTech Business Summit, New York City", 10, False, GREY),
        ("John Quinn, Cofounder & CEO, Exos Aerospace   |   Dr. Melissa Patton, Founder & CEO, Space Grove Ventures", 10, False, GREY)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, text, bold, size, color)

    h(doc, "Purpose")
    doc.add_paragraph(
        "Explore how Exos Aerospace and Space Grove Ventures can work together. Our lead proposal is a national "
        "Science Fair Payload Award: the winning project from each state earns a student payload on an Exos suborbital "
        "flight, with all winners flying together on an annual Exos rideshare campaign. We would build it with Nova Space "
        "and our Montana spaceport workforce effort. We would also like to set up a simple two-way referral partnership.")

    h(doc, "Agenda")
    for t, pts in [
        ("1. Introductions", [
            "Exos Aerospace: reusable suborbital rockets with full vehicle recovery. Our BLK3 vehicle, targeted to fly in early 2027, carries payloads and test hardware, and recovered payloads are back in hand within minutes of landing.",
            "Space Grove Ventures: the Space Coast platform, network and plans for 2027."]),
        ("2. What each organization brings", [
            ("Space Grove: ", "facility activation and operations, public-private collaboration, a venture platform and accelerator pathways, and links to workforce programs, students and universities."),
            ("Exos: ", "a reusable suborbital flight platform for flying and recovering hardware and experiments, payload integration support, and an Italian affiliate, Exos Aerospace Italia (Turin)."),
            ]),
        ("3. Science Fair Payload Award (lead item)", [
            "Concept, sponsor model and roles (see the draft rollout plan below).",
            "Would Space Grove co-own the rollout with Exos and lead sponsor recruitment?",
            "Which 2–3 states should pilot it?"]),
        ("4. Other ways to work together", [
            ("Referrals: ", "Space Grove introduces companies that need suborbital flight testing to Exos; Exos introduces companies that need a Space Coast base to Space Grove."),
            ("Events: ", "an Exos session at a Space Grove event or podcast on testing hardware in flight before orbit."),
            ("Italy and Europe: ", "student and company links through Exos Aerospace Italia, building on Space Grove's Italian Trade Agency connection."),
            ("Investor introductions: ", "Exos is raising a $10M Series A. Are there space investors in your network you would be willing to introduce us to? We would welcome the same for Science Fair Payload Award sponsors."),
            ]),
        ("5. Next steps and owners", [
            "Name a lead on each side.",
            "Exos drafts a two-page Science Fair Payload Award concept within two weeks for Space Grove's review.",
            "Exos arranges an introduction call between Space Grove and Nova Space.",
            "Space Grove shares any space investor and sponsor introductions it can make.",
            "Follow-up call during the week of November 2.",
            "Mutual NDA, then a short memorandum of understanding if both sides want to formalize."]),
    ]:
        sub(doc, t); bullets(doc, pts)

    h(doc, "Science Fair Payload Award — Draft Rollout Plan")
    doc.add_paragraph(
        "The top science fair project in each state wins a student payload on an Exos suborbital flight. Each year, all "
        "state winners fly together on one Exos rideshare campaign. Awards are funded by sponsors such as companies, "
        "state Space Grant consortia, economic-development agencies and foundations, who receive recognition at the "
        "state fairs and at the launch.")
    sub(doc, "Roles")
    bullets(doc, [
        ("Space Grove Ventures: ", "sponsor recruitment, events, the Florida showcase and ecosystem partners."),
        ("Exos Aerospace: ", "the flight, payload rules and size limits, integration, mentoring, recovery and return of payloads and data."),
        ("Nova Space / Space Workforce Institute: ", "an online course for winners covering payload design, flight safety and data analysis, with a credential (novaspaceinc.com)."),
        ("Montana spaceport workforce effort: ", "a hands-on workforce track for winners and, as the spaceport develops, a possible future campaign site."),
    ])
    for t, pts in [
        ("Phase 1, November–December 2026: Design", [
            "Agree on award rules, judging, payload limits and the sponsor package.",
            "Design the Nova Space / Space Workforce Institute winner course.",
            "Select 2–3 pilot states (proposed: Florida, Texas and Montana) and contact their state science fair organizers about a sponsored special award."]),
        ("Phase 2, January–May 2027: Pilot", [
            "Sponsors fund the pilot-state awards, which are presented at the 2027 state fairs.",
            "Winners complete the course and build their payloads with Exos mentoring."]),
        ("Phase 3: First annual rideshare campaign", [
            "Pilot winners fly together on one Exos suborbital campaign, scheduled once BLK3 enters service.",
            "Winners and families attend the launch, with a showcase event for sponsors and media."]),
        ("Phase 4, 2027–28 school year onward: National", [
            "Add states as sponsors join, working toward every state, with one rideshare campaign each year.",
            "Add the Montana hands-on workforce track as the spaceport develops."]),
    ]:
        sub(doc, t); bullets(doc, pts)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    run(p, "Program requirements: ", True)
    run(p, "every payload passes Exos safety and export-control review before flight; parental consent and "
           "youth-protection policies apply to all student participants.")

    h(doc, "Contacts")
    bullets(doc, [
        "John Quinn, Cofounder & CEO, Exos Aerospace — jquinn@exosaero.com — www.exosaero.com",
        "Dr. Melissa Patton, Founder & CEO, Space Grove Ventures — thespacegrove.org",
    ])

    doc.save(BODY)
    subprocess.run([sys.executable, f"{WS}/scripts/apply_exos_letterhead.py", BODY, OUT], check=True)
    subprocess.run([sys.executable, f"{WS}/scripts/apply_exos_letterhead.py", OUT, OUT, "--check"], check=False)
    print(OUT)


if __name__ == "__main__":
    build()
