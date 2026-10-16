#!/usr/bin/env python3
"""Japan round (SBI Investment + 650 Fund) data-room upload checklist + cover-email drafts (John 10/16/2026)."""
import os, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Pt, RGBColor
WS="/Users/exomac/.openclaw/workspace"; DR="data-room/Exos_SeriesA"
X=f"{WS}/drafts/Exos_Japan_Round_DataRoom_Upload_Checklist_2026-10-16.xlsx"
D=f"{WS}/drafts/Exos_Japan_Round_SBI_650_Cover_Emails_DRAFT_2026-10-16.docx"
H=["#","Action","Peony folder","Exact file to upload (local path)","Who sees it","Why / note"]
UP=[
 ("0","DO FIRST","—","MNDA via MC DocuSign panel (/docusign), MNDA template. SBI: include the MNPI clause (parent SBI Holdings is listed on TSE, 8473). 650 Fund: no MNPI clause. Counterparty signs first, Exos countersigns.","SBI Investment; 650 Fund","No NDA with either firm on file. The decks promise the room \"under NDA\". Grant room access only after the MNDA is executed."),
 ("1","UPLOAD","1 Exec Summary","drafts/Exos_Japan_Round_SBI_Investment_2026-10-16_r4.pptx","SBI only (also email to Yang Yang)","The deck you presented, corrected: Nebula/IM line → \"first tank sale for an orbital transfer vehicle\" (figures of record); EBITDA $254K → net income $188K from the filed 2025 return."),
 ("2","UPLOAD","1 Exec Summary","drafts/Exos_Japan_Round_650_Fund_2026-10-16_r4.pptx","650 only (also email to Dominik)","Same two corrections. Each deck is personalized, so give each firm its own deck. If both share one room, put each deck in a per-investor subfolder or send it by email only."),
 ("3","UPLOAD","3.5 409a",f"{DR}/_CANONICAL/EXOS_409A_Eqvista_2026-03-31_CURRENT.pdf","Both","Pricing basis: EAI units are priced off the Exos U.S. valuation ($169.7M, $8.66/sh). Keep the 2024/2025 vintages if they're already there."),
 ("4","UPLOAD","3.2 Projections",f"{DR}/_CANONICAL/EXOS_Financial_Model_Full_Revenue_FIXED_v12_FY2025actuals_rev-2026-10-10.xlsx","Both","Model of record. Includes the FY2025 Actuals tab (Form 1120). Do NOT use the v16 US+Japan K.K. model: it still carries ROFR wording."),
 ("5","UPLOAD","5 Market",f"{DR}/05_Market_Information/5.1_Market_Research_Exos_2026-08_rev-2026-10-16.docx","Both","NEW revision today: license wording → Part 431 transitioning to Part 450; dropped an unsourced \"only\" claim. Replaces rev-2026-10-10 in Peony."),
 ("6","UPLOAD","5 Market",f"{DR}/05_Market_Information/5.2_Competitive_Analysis_Exos_2026-08_rev-2026-10-10.docx","Both","Current. Its FY2025 numbers match the filed return."),
 ("7","UPLOAD","1 Exec Summary",f"{DR}/_CANONICAL/Exos_Policy_Tailwind_NSTP_2026-08-21.pdf","Both","U.S. policy tailwind brief. No valuation in it; room-safe."),
 ("8","UPLOAD","6 Business Dev",f"{DR}/06_Business_Development/Exos_BD_Engine_Pipeline_Stats_CURRENT.docx","Both","Demonstrates the AI operating platform (core asset #4); covers U.S. and Japan/APAC."),
 ("9","HOLD — your call","1 Exec Summary",f"{DR}/_CANONICAL/Exos_SeriesA_StrategicDeck_2026-09-06_CURRENT_rev-2026-10-10.pptx","—","This is the U.S. Series A (Reg D) deck. Showing it to Reg S buyers of EAI units blurs the two offerings (Counsel's call). It also says \"licensed and hover-demonstrated\" BLK3 and $1.55M revenue, so it needs the 10/15 wording fix plus the FY2025 decision before it goes to anyone."),
 ("10","HOLD — your call","3.3 Cap Table","data-room/Leonid_TACFI_Match/Exos_Capitalization_Table.xlsx","—","Last dated 6/24; the manifest says to confirm the current version. Names individual holders. Share at the term-discussion stage, after Kyle confirms it."),
 ("11","DO NOT UPLOAD","—",f"{DR}/Scorpius_Acquisition/* (SSLC BS/PL, overview, sales/expense PDFs)","—","SSLC's confidential financials under the 7/27/2021 NDA, with the deal still in negotiation."),
 ("12","DO NOT UPLOAD","—",f"{DR}/06_Business_Development/Exos_Contracts_Applied_For_2026-09-07.docx","—","Lists U.S. defense solicitations (incl. USAF FMS). Not for foreign investors; the Japan materials stay non-technical and defense-light."),
 ("13","DO NOT UPLOAD","—","drafts/EXOS_Financial_Model_v16_US_plus_JapanKK_2026-10-07_NDA_DRAFT.xlsx; drafts/Exos_Aerospace_Japan_Investor_Pitch_Rev2_2026-09-27.pptx","—","v16 has ROFR language (banned). Rev2 pitch is superseded by the 10/16 meeting decks."),
 ("14","LATER (Counsel)","—","EAI Operating Agreement + Reg S subscription / offering docs","—","Per your 9/27 rule, Reg S execution starts only once a serious investor exists. SBI asking for the room may be that trigger: your call on when to start Counsel."),
 ("15","REMOVE if present","—","Any pre-10/10 deck, model v9 and earlier, 5.1 rev-2026-10-10, anything with ROFR / Nebula / RTO / \"$25M\"","—","One file per artifact (manifest rule)."),
]
wb=Workbook(); ws=wb.active; ws.title="Upload checklist"
ws.append(["Exos — Japan Round data room (SBI Investment · 650 Fund) — upload checklist, 2026-10-16"]); ws["A1"].font=Font(bold=True,size=13)
ws.append([]); ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F3A68")
col={"UPLOAD":"E2EFDA","DO FIRST":"FFF2CC","HOLD — your call":"FCE4D6","DO NOT UPLOAD":"F8CBAD","LATER (Counsel)":"DDEBF7","REMOVE if present":"EDEDED"}
for r in UP:
    ws.append(list(r))
    for c in ws[ws.max_row]: c.alignment=Alignment(wrap_text=True,vertical="top")
    ws.cell(ws.max_row,2).fill=PatternFill("solid",fgColor=col[r[1]])
for k,w in zip("ABCDEF",[4,16,15,62,20,70]): ws.column_dimensions[k].width=w
ws.freeze_panes="A4"; wb.save(X)

d=Document(); st=d.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10.5)
def P(t,b=False,c=None,sa=6):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(sa); r=p.add_run(t); r.bold=b
    if c: r.font.color.rgb=c
    return p
RED=RGBColor(0xC0,0,0)
P("DRAFT: for John to send from jquinn@exosaero.com · not sent",True,RED)
P("EMAIL 1: SBI Investment",True)
P("To: Yang Yang <yayang@sbigroup.co.jp>\nCc: Kosuke Nakajima <kousunak@sbigroup.co.jp>; Toshi Kumazawa <toshihisa@exosaero.com>\nAttach: Exos_Japan_Round_SBI_Investment_2026-10-16_r4.pptx\nSubject: Exos Aerospace: deck from today's meeting, and next steps")
P("Dear Ms. Yang and Mr. Nakajima,")
P("Thank you for meeting with us today in Tokyo. As you requested, the presentation we walked through is attached.")
P("We would be glad to open our data room to the SBI team. The first step is a mutual NDA, which I will send through our e-signature system shortly. Once it is signed, I will send data-room access. The room includes our 409A valuation, financial model, market research and competitive analysis.")
P("As we discussed, the Japan round is offered as units of Exos Aerospace Investments, LLC under Regulation S, separate from any U.S. offering. Definitive terms will come through offering documents prepared by our counsel.")
P("We would also welcome a follow-up call with your investment team, and any introductions within SBI's corporate venture partners where there is a real business fit.")
P("Best regards,\nJohn Quinn\nCofounder & CEO, Exos Aerospace Systems & Technologies, Inc.\n+1-972-740-8355 · www.exosaero.com",sa=14)
P("EMAIL 2: 650 Fund",True)
P("To: Dominik Steiner <dominik@650.fund>\nCc: Toshi Kumazawa <toshihisa@exosaero.com>\nAttach: Exos_Japan_Round_650_Fund_2026-10-16_r4.pptx\nSubject: Exos Aerospace: deck from today's meeting at JETRO")
P("Dear Dominik,")
P("Thank you for the time today at JETRO. The deck we discussed is attached.")
P("I'm happy to give 650 access to our data room. I'll send a mutual NDA by e-signature first, and room access will follow as soon as it is signed.")
P("The Japan round is offered as units of Exos Aerospace Investments, LLC under Regulation S, separate from any U.S. offering, with definitive terms through counsel-prepared offering documents.")
P("I'd welcome a follow-up call at your convenience.")
P("Best regards,\nJohn Quinn\nCofounder & CEO, Exos Aerospace Systems & Technologies, Inc.\n+1-972-740-8355 · www.exosaero.com")
raw=D.replace(".docx","_raw.docx"); d.save(raw)
subprocess.run(["/usr/bin/python3",f"{WS}/scripts/apply_exos_letterhead.py",raw,D],check=False)
if not os.path.exists(D): os.rename(raw,D)
elif os.path.exists(raw): os.remove(raw)
print(X); print(D)
