#!/usr/bin/env python3
"""Build Exos Manufacturer & Operator Statement (Lady Rocket / Italy) as letterhead .docx + PDF twin.
Word:  PYTHONNOUSERSITE=1 python3.12 scripts/build_ladyrocket_exor_statement.py docx
PDF:   python3 scripts/build_ladyrocket_exor_statement.py pdf
"""
import os, sys, zipfile
sys.path.insert(0, os.path.dirname(__file__))
from _ladyrocket_exor_statement_content import TITLE, SUBTITLE, INTRO, SECTIONS, SIGNER

WS = os.path.expanduser("~/.openclaw/workspace")
LH = os.path.join(WS, "assets", "EXOS_Letterhead_Greenville_TX_PRIMARY.docx")
BASE = os.path.join(WS, "drafts", "Exos_LadyRocket_Manufacturer_Statement_Italy_2026-10-11_v2")


def build_docx():
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    NAVY, BODY = RGBColor(0x0A, 0x3D, 0x7C), RGBColor(0x1A, 0x1A, 0x1A)
    d = Document(LH)

    def para(text, size=10.5, bold=False, color=BODY, italic=False, space=6, style=None):
        p = d.add_paragraph(style=style) if style else d.add_paragraph()
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.color.rgb = color; r.font.name = "Calibri"
        p.paragraph_format.space_after = Pt(space)
        return p

    def borders(t):
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{e}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "8A9BB0")):
                el.set(qn(f"w:{k}"), v)
            b.append(el)
        tblPr.append(b)

    para(TITLE, 15, True, NAVY, space=2)
    para(SUBTITLE, 9, False, NAVY, italic=True, space=10)
    para(INTRO)
    for head, items in SECTIONS:
        para(head, 11.5, True, NAVY, space=4)
        for it in items:
            if isinstance(it, str):
                para(it)
            elif it[0] == "bullets":
                for b in it[1]:
                    try:
                        para(b, style="List Bullet", space=3)
                    except KeyError:
                        para("•  " + b, space=3)
            elif it[0] == "table":
                _, hdr, rows = it
                t = d.add_table(rows=1 + len(rows), cols=len(hdr)); borders(t)
                for i, h in enumerate(hdr):
                    c = t.rows[0].cells[i]; c.text = ""
                    r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NAVY
                for ri, row in enumerate(rows, 1):
                    for ci, val in enumerate(row):
                        c = t.rows[ri].cells[ci]; c.text = ""
                        r = c.paragraphs[0].add_run(val); r.font.size = Pt(9.5); r.bold = (ci == 0)
                para("", space=2)
    para("", space=8)
    for line in SIGNER:
        para(line, 10.5, line == SIGNER[0], space=0)
    d.save(BASE + ".docx")
    print("wrote", BASE + ".docx")


def build_pdf():
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, ListFlowable, ListItem
    import tempfile
    tmp = tempfile.mkdtemp()
    with zipfile.ZipFile(LH) as z:
        media = sorted(n for n in z.namelist() if n.startswith("word/media/"))
        z.extractall(tmp, media)
    hdr_img, ftr_img = os.path.join(tmp, "word/media/image1.png"), os.path.join(tmp, "word/media/image2.png")
    W, H = letter

    def on_page(c, doc):
        hw = W; hh = hw * 542 / 2625; fh = hw * 431 / 2625
        c.drawImage(hdr_img, 0, H - hh, hw, hh, mask="auto")
        c.drawImage(ftr_img, 0, 0, hw, fh, mask="auto")

    NAVY = colors.HexColor("#0A3D7C"); BODY = colors.HexColor("#1A1A1A")
    S = lambda **k: ParagraphStyle("x", **{**dict(fontName="Helvetica", fontSize=10, leading=13.5,
                                                  textColor=BODY, spaceAfter=6), **k})
    st_title = S(fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=NAVY)
    st_sub = S(fontName="Helvetica-Oblique", fontSize=8.5, textColor=NAVY, spaceAfter=10)
    st_h = S(fontName="Helvetica-Bold", fontSize=11.5, textColor=NAVY, spaceBefore=6, spaceAfter=4)
    st_b, st_cell = S(), S(fontSize=9, leading=11.5, spaceAfter=0)
    st_cellb = S(fontName="Helvetica-Bold", fontSize=9, leading=11.5, spaceAfter=0)
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;")

    flow = [Paragraph(esc(TITLE), st_title), Paragraph(esc(SUBTITLE), st_sub), Paragraph(esc(INTRO), st_b)]
    for head, items in SECTIONS:
        flow.append(Paragraph(esc(head), st_h))
        for it in items:
            if isinstance(it, str):
                flow.append(Paragraph(esc(it), st_b))
            elif it[0] == "bullets":
                flow.append(ListFlowable([ListItem(Paragraph(esc(b), st_b), leftIndent=12) for b in it[1]],
                                         bulletType="bullet", start="•", leftIndent=12))
            elif it[0] == "table":
                _, hdr, rows = it
                data = [[Paragraph(esc(h), st_cellb) for h in hdr]] + \
                       [[Paragraph(esc(v), st_cellb if i == 0 else st_cell) for i, v in enumerate(r)] for r in rows]
                t = Table(data, colWidths=[1.9 * inch, 4.6 * inch])
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8A9BB0")),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF6")),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
                flow += [t, Spacer(1, 8)]
    flow.append(Spacer(1, 14))
    for line in SIGNER:
        flow.append(Paragraph(esc(line) or "&nbsp;", st_cellb if line == SIGNER[0] else S(spaceAfter=0)))
    doc = SimpleDocTemplate(BASE + ".pdf", pagesize=letter, topMargin=1.7 * inch, bottomMargin=1.35 * inch,
                            leftMargin=1.0 * inch, rightMargin=1.0 * inch, title=TITLE, author="Exos Aerospace")
    doc.build(flow, onFirstPage=on_page, onLaterPages=on_page)
    print("wrote", BASE + ".pdf")


if __name__ == "__main__":
    {"docx": build_docx, "pdf": build_pdf}[sys.argv[1]]()
