#!/usr/bin/env python3
"""Build the Acrisure insurer-facing 'Vehicle Lease & Custody Structure' summary on
PRIMARY letterhead. Pricing redacted per broker's request (structure only)."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx"
DST = "drafts/Acrisure_Vehicle_Lease_Structure_Summary.docx"
NAVY = RGBColor(0x0B, 0x22, 0x40)

doc = Document(SRC)

def para(text="", bold=False, size=11, align=None, color=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return p

def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAVY
    return p

def bullet(lead, text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.left_indent = Pt(18)
    p.add_run("•  ").font.size = Pt(11)
    if lead:
        rb = p.add_run(lead); rb.bold = True; rb.font.size = Pt(11)
    r2 = p.add_run(text); r2.font.size = Pt(11)
    return p

# Header
para("VEHICLE LEASE & CUSTODY STRUCTURE", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, space_after=2)
para("Summary for Insurer — Prepared for Acrisure Aerospace", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

para("Date:  July 28, 2026", space_after=2)
para("Prepared by:  Philip Eaton, Exos Aerospace Systems & Technologies, Inc.", space_after=2)
para("For:  Steven Aceves / Adaly Ramon, Acrisure Aerospace", space_after=2)
para("Re:  Structure supporting $10M physical-loss coverage on the dedicated vehicle component set", space_after=2)
para("Note:  Structure only. Commercial rates/lease pricing are omitted, as requested.", italic=True, size=10, space_after=8)

heading("1.  What the Program Is")
para("Exos builds and leases a suborbital research launch vehicle to a specific customer to fly missions "
     "for that customer's research needs. Exos remains the owner of the vehicle at all times; title never "
     "passes to the customer. As part of the arrangement, the customer is responsible for procuring and "
     "funding a $10M policy covering loss of the vehicle while it is held in Exos's possession.")

heading("2.  Asset Status While Insured — Not Operational (Key Point)")
para("In the state in which Exos holds it, the vehicle is NOT an operational, flight-ready rocket and "
     "cannot fly. It is a dedicated set of long-lead components held for one customer — functionally a "
     "bonded set of spare parts, not an assembled vehicle. Specifically:")
bullet("Present: ", "a spare set of tanks plus a computer/avionics section with the engine section installed.")
bullet("Not installed: ", "the pressurization section, the payload section, and the recovery section — all added later.")
bullet("Consequence: ", "in this stored configuration the vehicle is incapable of flight. Final assembly and "
       "testing occur only after the customer authorizes proceed.")
para("These are the long-lead-time items the customer leaves in Exos's possession — dedicated to that "
     "customer like bonded inventory — so this is properly a 'spare parts' / stored-inventory exposure, "
     "not an operational-aircraft/hull exposure.", space_after=8)

heading("3.  Coverage Scope Requested")
bullet("Covered: ", "physical loss of or damage to the dedicated component set (the tanks, computer/avionics "
       "and engine section) while stored at Exos's facility, 299 FM 1903, Greenville, TX 75402 — a fixed, "
       "monitored premises. Agreed value: $10,000,000.")
bullet("Coverage ends at ship-out: ", "once the customer gives authorization to proceed and the material "
       "ships from Exos's building to the launch site for final assembly, coverage is no longer required.")
bullet("Explicitly NOT requested: ", "transit coverage after the vehicle leaves Exos's building, launch-site "
       "assembly, pre-flight operations, or flight/launch coverage. Those are outside this policy.")

heading("4.  Lease Structure (Pricing Omitted)")
bullet("Owner / Lessor: ", "Exos Aerospace Systems & Technologies, Inc. — retains title and ownership throughout.")
bullet("Lessee: ", "the research customer, who directs the flight program and funds the loss-of-vehicle policy.")
bullet("Type: ", "true operating lease — not a finance/capital lease; Exos retains title and residual value. "
       "No purchase or title-transfer to the customer.")
bullet("Insurance obligation: ", "customer procures and maintains the $10M loss-of-vehicle policy; Exos named "
       "as owner/loss payee for its insurable interest in the stored component set.")
bullet("Reversion: ", "on lease expiry, termination, or breach, the vehicle and all Exos-owned equipment "
       "revert to Exos.")

heading("5.  Bottom Line for Underwriting")
para("Exos is seeking coverage on a dedicated, bonded set of spare rocket components stored at a fixed Texas "
     "premises and incapable of flight in that state — a stored-inventory / spare-parts exposure with an "
     "agreed value of $10M and a hard cutoff at ship-out. A template lease reflecting the structure above "
     "can be provided on request to confirm the ownership, insurable-interest, and insurance-obligation terms.")

para("", space_after=2)
para("Prepared for insurer underwriting use. Non-binding summary; definitive terms are set by the executed "
     "lease and policy.", italic=True, size=9, color=NAVY)

doc.save(DST)
print("wrote", DST)
