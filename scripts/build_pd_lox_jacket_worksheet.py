#!/usr/bin/env python3
"""Build the PD Aerospace LOX tank insulating-jacket fabrication estimate worksheet (.docx).
Budgetary Materials + Labor-Hours framework, modeled on the Project Khonsu jackets.
Numeric cells are intentionally left [TBD] — this is an estimating framework for Phil to
populate with Khonsu actuals + the PD tank geometry, not an invented price."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/Users/exomac/.openclaw/workspace/drafts/PD_Aerospace_LOX_Jacket_Estimate_Worksheet.docx"

doc = Document()
for s in doc.styles:
    pass
normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(10.5)

def h(text, size=13, space_before=10):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(4)
    return p

def body(text, bullet=False):
    p = doc.add_paragraph(text, style="List Bullet" if bullet else None)
    p.paragraph_format.space_after = Pt(2)
    return p

# Title
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("PD Aerospace — LOX Tank Insulating Jacket"); r.bold = True; r.font.size = Pt(16)
st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = st.add_run("Fabricate & Install — Budgetary Materials + Labor-Hours Worksheet"); r.font.size = Pt(11); r.italic = True
meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Requested by Neil Milburn (COO)  •  Estimator: Phil  •  Reference build: Project Khonsu  •  Status: BUDGETARY / prospective order (2027, TBD)")
r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = RGBColor(0x66,0x66,0x66)

h("1. Purpose & status")
body("Budgetary Materials and Labor-Hours estimate to fabricate and install an insulating jacket "
     "for the PD Aerospace LOX tank whose budget design they have reviewed favorably. Approach mirrors "
     "the insulating jackets Exos built for Project Khonsu. PD is hinting at an order next year (TBD) — "
     "this is a rough-order-of-magnitude (ROM) figure to support that conversation, not a firm quote.")

h("2. Inputs needed before numbers are real (open items)")
body("Populate these before the estimate can be trusted. Two are hard blockers:", )
body("PD tank geometry — from the budget design Neil forwarded: outer diameter, overall length, dome "
     "geometry, TOTAL insulated surface area (sq ft), and number/size/location of ports, penetrations, "
     "and mounting interfaces. (Blocker — drives every area-scaled line below.)", bullet=True)
body("Insulation spec — type and thickness (e.g. SOFI spray foam, bonded rigid closed-cell foam, or "
     "MLI + purge) and service condition (LOX, approx. -297 F / -183 C). (Blocker.)", bullet=True)
body("Outer jacket / facing — material and finish (composite wrap, GRP shell, aluminized restraint, etc.).", bullet=True)
body("Fabrication vs. field install — shop-built jacket vs. install at PD's site; travel/fit-up scope.", bullet=True)
body("Quantity & schedule — one unit vs. more; need date. Affects tooling amortization.", bullet=True)

h("3. Reference basis — Project Khonsu (enter actuals)")
body("Pull the Khonsu jacket build records; the PD estimate scales from these. Fill the right column:")
tbl = doc.add_table(rows=1, cols=2); tbl.style = "Light Grid Accent 1"
tbl.rows[0].cells[0].paragraphs[0].add_run("Khonsu reference metric").bold = True
tbl.rows[0].cells[1].paragraphs[0].add_run("Actual (fill in)").bold = True
for metric in ["Insulated surface area (sq ft)", "Insulation type & thickness used",
               "Total materials cost ($)", "Total fabrication labor (hrs)",
               "Total installation labor (hrs)", "Tooling / pattern hrs (one-time)",
               "Scrap / rework allowance applied (%)"]:
    row = tbl.add_row().cells; row[0].text = metric; row[1].text = "[  ]"

h("4. Materials BOM template")
body("Quantities are area-driven off Section 2 geometry. Enter unit costs; extend for total.")
m = doc.add_table(rows=1, cols=5); m.style = "Light Grid Accent 1"
for i, col in enumerate(["Material line item", "Spec / basis", "Qty basis", "Est. unit $", "Ext $"]):
    m.rows[0].cells[i].paragraphs[0].add_run(col).bold = True
for item in [
    ("Insulation media (foam / MLI / blanket)", "per Sec.2 spec", "sq ft x thickness"),
    ("Adhesive / bonding & primer system", "per media", "sq ft"),
    ("Vapor barrier / sealant", "cryo-rated", "sq ft + seams"),
    ("Outer jacket / facing material", "per Sec.2", "sq ft"),
    ("Closeout strips, standoffs, bands / fasteners", "restraint", "lin ft / ea"),
    ("Penetration boots & port closeouts", "per penetration", "ea"),
    ("Consumables (mixing, masking, PPE, cure media)", "shop", "lot"),
    ("Tooling / patterns (if new for PD geometry)", "one-time", "lot"),
]:
    row = m.add_row().cells
    row[0].text = item[0]; row[1].text = item[1]; row[2].text = item[2]
    row[3].text = "[TBD]"; row[4].text = "[TBD]"
tot = m.add_row().cells; tot[0].paragraphs[0].add_run("MATERIALS SUBTOTAL").bold = True
tot[4].text = "[TBD]"

h("5. Labor-hours template")
body("Enter hours per operation. Where possible drive from Khonsu hrs/sq ft x PD area (Sec.6).")
l = doc.add_table(rows=1, cols=3); l.style = "Light Grid Accent 1"
for i, col in enumerate(["Operation", "Basis (hrs/sq ft or fixed)", "Est. hrs"]):
    l.rows[0].cells[i].paragraphs[0].add_run(col).bold = True
for op in [
    ("Engineering / pattern & template development", "fixed (one-time)"),
    ("Surface prep & masking", "hrs/sq ft"),
    ("Cut & fit insulation", "hrs/sq ft"),
    ("Bond / apply & cure", "hrs/sq ft"),
    ("Outer jacket layup / wrap & closeouts", "hrs/sq ft"),
    ("Penetration & port closeouts", "hrs/penetration"),
    ("Install on tank / fit-up", "hrs/sq ft or fixed"),
    ("Inspection & rework allowance", "% of above"),
]:
    row = l.add_row().cells; row[0].text = op[0]; row[1].text = op[1]; row[2].text = "[TBD]"
tot = l.add_row().cells; tot[0].paragraphs[0].add_run("TOTAL LABOR HOURS").bold = True; tot[2].text = "[TBD]"

h("6. Scaling method (Khonsu → PD)")
body("Area-scaled ops:  PD hrs (or $) = Khonsu value x (PD insulated area / Khonsu insulated area).", bullet=True)
body("Fixed / one-time items (engineering, patterns, new tooling) do NOT scale — carry as flat adders.", bullet=True)
body("If PD insulation spec or jacket material differs from Khonsu, adjust the affected line before scaling.", bullet=True)

h("7. Assumptions & exclusions")
body("Excludes the tank itself, proof/leak/cryo qualification testing, freight, and on-site travel unless "
     "line-itemed. Assumes shop fabrication and single unit. Budgetary accuracy target ± [ ]% — refine "
     "once PD geometry and Khonsu actuals are entered.")

h("8. Owner actions")
body("Phil: pull Khonsu jacket build records into Section 3; confirm insulation/jacket spec for LOX service.", bullet=True)
body("Neil: forward the PD tank budget design / drawing so surface area and penetrations can be quantified.", bullet=True)
body("Then this worksheet resolves to a single Materials $ + Labor-hrs ROM for the PD conversation.", bullet=True)

doc.save(OUT)
print("WROTE", OUT)
