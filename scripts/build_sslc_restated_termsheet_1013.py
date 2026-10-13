"""SSLC Asset Acquisition — Restated Offer Term Sheet (task #2414, 2026-10-13).
Restates the 5/13/2025 signed Asset Acquisition Agreement (Quinn/Wertz) with the Option B structure revisions (r4/r5, 10/12)."""
import os, subprocess
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
WS="/Users/exomac/.openclaw/workspace"
RAW=f"{WS}/drafts/_sslc_restated_ts_raw.docx"
OUT=f"{WS}/drafts/Exos_SSLC_Asset_Acquisition_Offer_Term_Sheet_RESTATED_2026-10-13_r1.docx"
NAVY=RGBColor(0x0A,0x3D,0x7C); INK=RGBColor(0x1A,0x1A,0x1A); RED=RGBColor(0xB0,0x1E,0x1E); GREY=RGBColor(0x55,0x55,0x55)
d=Document()
st=d.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10)
def P(t,size=10,bold=False,color=INK,align=None,sa=4,sb=0,italic=False):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(sa); p.paragraph_format.space_before=Pt(sb)
    if align is not None: p.alignment=align
    r=p.add_run(t); r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color; r.font.name="Calibri"; return p
def H(t): return P(t,11.5,True,NAVY,sb=8,sa=3)
def table(rows,widths,header=True):
    t=d.add_table(rows=0,cols=len(widths)); t.style="Table Grid"; t.autofit=False
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,txt in enumerate(row):
            cells[j].width=Inches(widths[j]); cells[j].text=""
            r=cells[j].paragraphs[0].add_run(txt); r.font.size=Pt(9); r.font.name="Calibri"
            if header and i==0: r.font.bold=True; r.font.color.rgb=NAVY
            elif j==0: r.font.bold=True
    for j,w in enumerate(widths): t.columns[j].width=Inches(w)
    d.add_paragraph().paragraph_format.space_after=Pt(2)

P("CONFIDENTIAL  ·  DRAFT  ·  NON-BINDING  ·  NEGOTIATING TERMS",9,True,RED,WD_ALIGN_PARAGRAPH.CENTER)
P("ASSET ACQUISITION — RESTATED OFFER TERM SHEET",15,True,NAVY,WD_ALIGN_PARAGRAPH.CENTER,sa=2)
P("Scorpius Space Launch Company (SSLC) assets → Exos Aerospace Systems & Technologies, Inc.",10.5,True,INK,WD_ALIGN_PARAGRAPH.CENTER,sa=2)
P("October 13, 2026  ·  r1  ·  Restates the Asset Acquisition Agreement signed by John M. Quinn (Exos) and James R. Wertz (SSLC) on May 13, 2025, with the structure revisions negotiated since (Option B, 10/7–10/12/2026)",9,False,GREY,WD_ALIGN_PARAGRAPH.CENTER,sa=8,italic=True)

H("1. Purpose")
P("This term sheet keeps the intent of the executed agreements (the October 9, 2024 Amendment, executed October 18, 2024, and the May 2, 2025 Asset Acquisition Agreement, SSLC board-approved May 2, 2025 and signed May 13, 2025): a cash-free / debt-free purchase of all SSLC assets, property, trademarks, trade secrets and IP, with credit for Exos's $790,000 investment (18.06%). It restates the price structure to add equity and SSLC-growth earn-outs. It is the basis for M&A counsel to draft the Definitive Agreement. Where it differs from the 2025 Agreement, the parties will document the change in a signed amendment.")

H("2. Signed 2025 terms → restated terms")
table([
 ["Term","May 13, 2025 Agreement (signed)","Restated offer (this term sheet)"],
 ["Parties","Buyer: Exos (Texas). Seller: SSLC (California).","Unchanged."],
 ["Structure","Cash-free / debt-free purchase of all assets, property, trademarks, trade secrets and IP (§1).","Unchanged. Asset purchase, not stock. Exos assumes no SSLC debts or liabilities. They stay with SSLC and are paid from SSLC proceeds (debt payoff removed from the offer 10/10/2026)."],
 ["Excluded assets","Microcosm's right to use PRESSURMAXX tank technology and resin for its launch vehicles and spacecraft. LOXair LLC's non-aerospace use under the October 2023 Technology License (§1).","Unchanged."],
 ["Cash at close","$2.51M (§2)","$2,510,000. Unchanged."],
 ["Seller transaction costs","$60,000 (§2)","$60,000. Unchanged."],
 ["Exos's existing stake","Exos relinquishes or retires its 18% (credit for $790K prior investment) (§2)","Unchanged. $790,000 non-cash credit. Exos's 157,553 shares are retired at close."],
 ["Total commitment at close","$3.36M ($2.57M closing + $790K prior investment)","$3,360,000 + $1,000,000 in Exos EAI units = $4,360,000 at financing close (NEW: equity)."],
 ["Equity at close","None","NEW: $1.0M of Exos Aerospace Investments (EAI) units for key SSLC personnel (not pro-rata to the cap table). Issued under Exos's Rule 504 Reg D offering at the 409A value (~0.59% fully diluted at ~$170M). Restricted and legended at issuance. Counts toward the $10M rolling 12-month Rule 504 cap. U.S. persons only (kept separate from the Reg S Japan round)."],
 ["Earn-outs (FY27–FY31)","None","NEW, see §3. Low $950K / Base $2.525M / High $4.1M."],
 ["Exclusivity","Non-exclusive until Exos has a financing agreement. Then 90 days of exclusivity. Extensions not unreasonably withheld while Exos works in good faith (§3).","Exos invokes §3 based on its exclusive financing agreement with Ainigma/FNEX. 90 days from SSLC's confirmation, plus a 60-day extension if significant progress is made. Extensions not unreasonably withheld."],
 ["Closing","Within 90 days of exclusivity, subject to §5 conditions (§6).","Target January 15, 2027 (Q1 FY27), subject to the same conditions and to Exos financing close."],
 ["Holdback","5% escrow on the closing payment, released 12 months after closing absent claims (§7).","Unchanged."],
 ["Reps & warranties","Standard reps, warranties and indemnities, subject to due diligence (§4).","Unchanged. Asset-focused: sole IP ownership free of liens, trade secrets protected, ITAR/DDTC compliance, contracts assignable, title to tooling, no undisclosed liabilities on the assets."],
 ["Conditions","Definitive Agreement. Due diligence to Buyer's satisfaction. Board and shareholder approvals. Regulatory and third-party consents (§5).","Unchanged. Diligence includes read-only QuickBooks access and review of the April 30, 2026 entries. Exos's approval is by disinterested directors (Quinn abstains), supported by the Eqvista fairness opinion."],
 ["Employees","Not addressed","NEW: Exos offers employment to SSLC employees on terms no less favorable than today (pay and benefits), subject to Exos hiring policies and diligence. Key-employee option pool to be set in the Definitive Agreement."],
 ["Fees","Each party pays its own costs beyond $60K of seller expenses (§9).","Unchanged. 5% outside service fee on cash proceeds."],
 ["Confidentiality / law","NDA dated July 27, 2021 (§8). Delaware law; California law for employment matters (§9).","Unchanged."],
],[1.25,2.45,2.8])

H("3. Earn-outs (FY27–FY31) — none pay unless the trigger is met")
table([
 ["Earn-out","Trigger","Low","Base","High"],
 ["Tank revenue","5% of SSLC's Intuitive Machines tank revenue each year, capped at $300K/yr","$750K","$1,250K","$1,500K"],
 ["M1","SSLC FY2026-27 revenue ≥ $5.1M → $200K (paid FY28 Q1 once verified)","$200K","$200K","$200K"],
 ["M2","Nebula OTV flight tank set delivered to and accepted by Intuitive Machines → $200K","—","$200K","$200K"],
 ["M3","SSLC annual revenue ≥ $10M → $250K","—","$250K","$250K"],
 ["M4","SSLC annual revenue ≥ $20M → $250K","—","—","$250K"],
 ["M5","50% cut in tank build cost (dome and process automation proven) → $200K","—","—","$200K"],
 ["Exos stock at RTO listing","0.0% / 0.25% / 0.50% fully diluted, issued at listing; lockup same as Exos principals","$0","$625K","$1,500K"],
 ["Total earn-outs","","$950K","$2,525K","$4,100K"],
 ["TOTAL 5-YEAR OFFER VALUE","$4.36M at close + earn-outs","$5.31M","$6.885M","$8.46M"],
],[1.35,3.05,0.7,0.7,0.7])
P("The low/base/high amounts are management scenarios. Exos stock values depend on an RTO listing that may not happen and are not guaranteed. Model: EXOS_SSLC_OptionB_Acquisition_Model_2026-10-07.xlsx.",8.5,italic=True,color=GREY)

H("4. Open items for counsel and the parties")
for t in [
 "Signed amendment to the 2025 Agreement adopting §§2–3 above (equity and earn-outs), signed by SSLC (Wertz) and Exos.",
 "SSLC's written confirmation of §3 exclusivity, and the date it starts.",
 "Exos's $31,200 loan to SSLC (11% interest, past due). SSLC's board denied cancellation, so the loan stands. Its payoff mechanics at close go in the Definitive Agreement.",
 "Rule 504: confirm how much of the $10M 12-month capacity is left before any EAI issuance.",
 "Eqvista fairness opinion: agreement engaged 8/14/2026, still unsigned. Books corrections (April 30, 2026 entries) are pending.",
 "Status of the high-pressure helium tank owed under fully paid Exos PO 395 (due 9/24/2025 under 2025 §1).",
]: P("•  "+t,9.5,sa=2)

P("Non-binding summary of negotiating terms. It creates no obligation until a Definitive Agreement is signed, and it is not legal advice. The May 13, 2025 Agreement remains in effect except as amended in writing.",8.5,italic=True,color=GREY,sb=6)
P("DRAFT — FOR REVIEW — John M. Quinn, Cofounder & CEO, Exos Aerospace — 2026-10-13",8.5,True,RED,WD_ALIGN_PARAGRAPH.CENTER,sb=4)
d.save(RAW)
subprocess.run(["/usr/bin/python3",f"{WS}/scripts/apply_exos_letterhead.py",RAW,OUT],check=True)
os.remove(RAW)
subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",f"{WS}/drafts",OUT],capture_output=True)
print(OUT)
