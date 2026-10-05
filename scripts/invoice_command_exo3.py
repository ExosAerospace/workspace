#!/usr/bin/env python3
"""
invoice_command_exo3.py — "@Exobot invoice <product>" email command (EXOS-FIN-003-0 Tasks 2->3).

Jon approved 2026-10-03. Lizi (or Jon) forwards a buyer's email reply to Exo with a line like

    @Exobot invoice watch
    Name: Jane Doe, Email: jane@example.com, Qty: 2

(or inline: "@Exobot invoice watch Jane Doe jane@example.com 2"). This module is called by
scripts/mc-email-command.py AFTER its sender authentication + export-control screen and BEFORE
any LLM hand-off. It is deterministic (no LLM):

  1. authorized senders only (config/invoice-offerings.json authorized_senders) — others: ignore + log
  2. Message-ID dedupe (offering ledger `processed`) — the same message can never create two invoices
  3. parse product + buyer full name, billing email, qty (email falls back to the forwarded
     buyer message's Reply-To/From). Missing/ambiguous -> reply asking for exactly what's missing
  4. catalog lookup (unknown product -> reply), cap check over Reserved+Invoiced+Paid in the ledger
     plus any hand-entered rows in the product workbook (over cap -> "sold out / only N left")
  5. CSL denied-party screen on the buyer (scripts/csl_screen.py). Hit or screen error -> HOLD,
     no invoice, Jon alerted (fail-closed)
  6. qb_invoice.prepare(): QB invoice created UNSENT + PDF + "EXO3: Approve & send invoice" Decision
     (approver Jon). Nothing reaches the buyer until Jon answers Approve & send (qb_invoice watch).
  7. ledger entry Reserved (+ workbook mirror), confirmation emailed back to the (internal) sender

Status after that is driven by qb_invoice.py watch: Approve & send -> Invoiced, Cancel -> Cancelled
(unit released), QB balance 0 -> Paid  (see ledger_mark()).

Owner: John Quinn. Eval: `python3 scripts/invoice_command_exo3.py --selftest`.
Telemetry: fleet_log (logs/fleet-events.jsonl) + decision_trace + logs/invoice-command.log.

CLI:
  invoice_command_exo3.py --selftest
  invoice_command_exo3.py --simulate <file.eml> [--dry-run] [--sandbox]
      --sandbox  runs against temp copies of the ledger + workbook (never the real ones)
  invoice_command_exo3.py --status [product]
"""
import argparse, copy, datetime, fcntl, json, os, re, shutil, sys, tempfile, uuid
from email.utils import getaddresses

WS = "/Users/exomac/.openclaw/workspace"
sys.path.insert(0, os.path.join(WS, "scripts"))

AGENT = "com.exos.mc-email-command"   # host agent (registry); sub-capability below
CAPABILITY = "invoice-command-exo3"
OFFERINGS = os.path.join(WS, "config/invoice-offerings.json")
LEDGER = os.path.join(WS, "mission-control/data/offering-ledger.json")
USERS_FILE = os.path.join(WS, "mission-control/users.json")
DISPOSITIONS = os.path.join(WS, "mission-control/data/csl-dispositions.json")
LOG = os.path.join(WS, "logs/invoice-command.log")
JON = "jquinn@exosaero.com"
ACTIVE = ("preparing", "Reserved", "Invoiced", "Paid")   # count against the cap
INTERNAL_DOMAINS = ("exosaero.com", "pm.me", "proton.me", "protonmail.com")
TELEMETRY = True   # selftest turns this off so evals never write fleet/trace/log records

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9\-]+(?:\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?1[\s.\-]?)?\(?\d{3}\)?[\s.\-]?\d{3}[\s.\-]\d{4}(?!\d)")
CMD_RE = re.compile(r"@\s*ex(?:o(?:bot|bto|boto|bott)?|bot)\b[\s:,\-]*invoice\b[ \t:,\-]*([^\n]*)", re.I)
# Start of quoted / forwarded material: everything above is the sender's own text.
FWD_RE = re.compile(r"(?im)^\s*(?:_{8,}|-{2,}\s*(?:forwarded message|original message)|begin forwarded message"
                    r"|from:\s*\S.*|on\s.{4,200}?wrote:\s*$|>)"
                    r"|\bfrom:\s*[^\n]{0,200}?\b(?:sent|date):"                # Outlook/Gmail header run inline
                    r"|-{2,}\s*(?:forwarded message|original message)|begin forwarded message")
NAME_L = r"(?:buyer'?s?\s+(?:full\s+)?name|customer\s+name|full\s+name|bill\s+to|buyer|customer|name)"
EMAIL_L = r"(?:billing\s+e-?mail(?:\s+address)?|buyer'?s?\s+e-?mail|e-?mail(?:\s+address)?)"
QTY_L = r"(?:qty|quantity|units|count|how\s+many)"
PHONE_L = r"(?:(?:buyer'?s?\s+)?(?:phone|tel|telephone|mobile|cell)(?:\s+(?:number|no\.?|#))?)"
ADDR_L = r"(?:(?:billing|mailing|shipping|street|home|buyer'?s?)\s+address|address)"
ALL_L = rf"(?:{NAME_L}|{EMAIL_L}|{QTY_L}|{PHONE_L}|{ADDR_L})"
EXOS_PHONES = {"9727408355", "4309664005"}   # Jon mobile, Exos main: never a buyer's number
NUM_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10}
NAME_STOP = {"for", "to", "qty", "quantity", "units", "unit", "x", "please", "pls", "the", "buyer", "name", "email",
             "e-mail", "invoice", "of", "at", "is", "a", "an", "pcs", "pc", "billing", "customer", "thanks", "thank", "you"}
TRAIL_STOP = {"please", "pls", "thanks", "thank", "you", "x", "qty", "quantity", "units", "unit", "pcs", "pc", "for", "to"}
SIGNOFF = re.compile(r"(?i)\b(exos|aerospace|cofounder|co-founder|ceo|coo|cfo|thanks|thank you|regards|best|cheers|"
                     r"sent from|mobile|address|website|linkedin|youtube|inc\b|llc\b|srl\b)")


def now():
    return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def log(m):
    if not TELEMETRY:
        return
    os.makedirs(os.path.dirname(LOG), exist_ok=True)
    with open(LOG, "a") as f:
        f.write(f"{now()} {m}\n")


def _fleet(action, outcome, detail, level="info", corr=None):
    if not TELEMETRY:
        return
    try:
        import fleet_log
        fleet_log.event(AGENT, "orchestrate", action, corr=corr, owner="John Quinn", outcome=outcome,
                        level=level, detail=dict(detail, capability=CAPABILITY))
    except Exception:
        pass


def _trace(corr, inputs, policy, chosen, decision_id=None, write_back=""):
    if not TELEMETRY:
        return
    try:
        import decision_trace
        decision_trace.emit(decision_id=decision_id, agent=AGENT + ":" + CAPABILITY, corr=corr, inputs=inputs,
                            policy_applied=policy, chosen=chosen, write_back=write_back,
                            final_call="agent", approver="John Quinn")
    except Exception:
        pass


def jload(p, default):
    try:
        with open(p) as f:
            return json.load(f)
    except Exception:
        return default


def load_offerings(path=None):
    return jload(path or OFFERINGS, {"authorized_senders": [], "offerings": {}})


# ── parsing ───────────────────────────────────────────────────────────────────
def split_own_text(body):
    """(sender's own text, forwarded/quoted remainder)."""
    body = body or ""
    m = FWD_RE.search(body)
    return (body[:m.start()], body[m.start():]) if m else (body, "")


_BLOCK_TAGS = re.compile(r"(?i)<\s*(?:br|hr)\b[^>]*>|</\s*(?:p|div|tr|li|h[1-6]|table|blockquote)\s*>")


def message_text(msg, cap=20000):
    """Body text for the invoice parser: text/plain if present, else HTML with block tags turned into
    newlines (Outlook sends HTML-only; a tag-stripped blob loses the line structure that separates the
    command, the signature, and the forwarded header). Handles single-part text/html too."""
    plain, htmls = [], []
    for part in (msg.walk() if msg.is_multipart() else [msg]):
        if part.get_content_disposition() == "attachment":
            continue
        ct = part.get_content_type()
        try:
            txt = part.get_payload(decode=True).decode(part.get_content_charset() or "utf-8", errors="replace")
        except Exception:
            continue
        (plain if ct == "text/plain" else htmls if ct == "text/html" else []).append(txt)
    if plain:
        text = "\n".join(plain)
    else:
        import html as _h
        h = "\n".join(htmls)
        h = re.sub(r"(?is)<(script|style|head).*?</\1>", " ", h)
        h = _BLOCK_TAGS.sub("\n", h)
        h = re.sub(r"(?s)<[^>]+>", " ", h)
        text = _h.unescape(h)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u200b", "").replace("\xa0", " ")
    text = "\n".join(re.sub(r"[ \t]+", " ", ln).strip() for ln in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()[:cap]


def find_command(body):
    """The invoice command line in the sender's OWN text (never in quoted/forwarded material — a quoted
    earlier '@Exobot invoice' in a reply chain must not re-fire). Returns the regex match or None."""
    own, _ = split_own_text(body)
    return CMD_RE.search(own)


def is_invoice_command(body):
    return find_command(body) is not None


def _internal(addr):
    return (addr or "").lower().rsplit("@", 1)[-1] in INTERNAL_DOMAINS


def _qty_val(s):
    s = (s or "").strip().lower()
    m = re.search(r"\d+", s)
    if m:
        return int(m.group())
    for w, n in NUM_WORDS.items():
        if re.search(rf"\b{w}\b", s):
            return n
    return None


def _clean_name(s):
    s = EMAIL_RE.sub(" ", s or "")
    s = re.sub(r"[<>()\[\]\"]", " ", s)
    s = re.sub(r"\s+", " ", s).strip(" ,;:|-.\t")
    return s


def _name_like(s, exclude):
    """A 2-6 word personal/company name — not a signature line, not an Exos roster name."""
    s = _clean_name(s)
    if not s or re.search(r"\d", s) or SIGNOFF.search(s):
        return None
    words = [w for w in s.split() if w]
    if not (2 <= len(words) <= 6) or not all(re.fullmatch(r"[A-Za-z][A-Za-z.'\-]*|&|and", w) for w in words):
        return None
    if s.lower() in exclude:
        return None
    return s


def roster_names():
    u = jload(USERS_FILE, [])
    u = u if isinstance(u, list) else u.get("users", [])
    return {(x.get("name") or "").strip().lower() for x in u if x.get("name")}


def resolve_product(token_text, offerings):
    """Longest alias that the text after 'invoice' starts with -> (key, remainder) or (None, first_word)."""
    t = (token_text or "").strip()
    best = None
    for key, o in offerings.items():
        for al in sorted(set([key] + o.get("aliases", [])), key=len, reverse=True):
            if re.match(rf"(?i){re.escape(al)}(?![A-Za-z0-9])", t) and (best is None or len(al) > len(best[1])):
                best = (key, al)
    if best:
        return best[0], t[len(best[1]):]
    first = (t.split() or [""])[0].strip(" ,;:.")
    return None, first


def parse_command(body, offerings, exclude_names=()):
    """Parse an invoice command. Returns dict:
       product (key|None), product_word, name, email, qty, email_source, missing[list], ambiguous[list]."""
    exclude = {n.lower() for n in exclude_names if n}
    own, fwd = split_own_text(body)
    m = CMD_RE.search(own)
    res = {"product": None, "product_word": "", "name": None, "email": None, "qty": None, "phone": None, "address": None,
           "email_source": None, "missing": [], "ambiguous": []}
    if not m:
        res["missing"].append("the command line '@Exobot invoice <product>'")
        return res
    key, rest = resolve_product(m.group(1), offerings)
    res["product"], res["product_word"] = key, (key or rest)
    if key is None:
        rest = m.group(1)[len(rest):] if rest else m.group(1)

    names, emails, qtys, phones, addrs = [], [], [], [], []
    # (a) labeled fields anywhere in the sender's own text (explicit labels are safe to read anywhere)
    lab_re = re.compile(rf"(?i)(?<![\w])({ALL_L})\s*[:=]\s*(.+?)(?=[\s,;|]+{ALL_L}\s*[:=]|\n|$)")
    labeled_spans = []
    for lm in lab_re.finditer(own):
        label, val = lm.group(1).lower(), lm.group(2).strip().strip(",;|")
        labeled_spans.append((lm.start(), lm.end()))
        if re.fullmatch(EMAIL_L, label, re.I):
            emails += EMAIL_RE.findall(val)
        elif re.fullmatch(PHONE_L, label, re.I):
            phones += [ph.strip() for ph in PHONE_RE.findall(val)]
        elif re.fullmatch(ADDR_L, label, re.I):
            if re.search(r"\d", val) and len(val) >= 8:
                addrs.append(val.strip())
        elif re.fullmatch(QTY_L, label, re.I):
            q = _qty_val(val)
            if q is not None:
                qtys.append(q)
        else:
            emails += EMAIL_RE.findall(val)
            ws_ = []
            for w in _clean_name(val).split():
                if not re.fullmatch(r"[A-Za-z][A-Za-z.'\-]*|&|and", w) or SIGNOFF.search(w) or len(ws_) >= 6:
                    break
                ws_.append(w)
            if ws_:
                names.append(" ".join(ws_))

    # (b) inline on the command line, after the product word (only the part not covered by labels)
    inline = rest if not re.search(rf"(?i){ALL_L}\s*[:=]", rest) else re.split(rf"(?i){ALL_L}\s*[:=]", rest)[0]
    if inline.strip():
        emails += EMAIL_RE.findall(inline)
        t = EMAIL_RE.sub(" ", inline)
        for qm in re.finditer(r"(?i)\b(?:qty|quantity|x)\s*[:=]?\s*(\d{1,3})\b|\b(\d{1,3})\s*(?:units?|watch(?:es)?|pcs?)\b"
                              r"|(?<![\w@.])(\d{1,3})(?![\w@.])", t):
            qtys.append(int(next(g for g in qm.groups() if g)))
        t = re.sub(r"(?i)\b(?:qty|quantity|x)\s*[:=]?\s*\d{1,3}\b|\b\d{1,3}\s*(?:units?|watch(?:es)?|pcs?)\b|\d", " ", t)
        for rn in sorted(exclude, key=len, reverse=True):   # an Exos signature name run onto the line
            if rn:
                t = re.sub(rf"(?i)\b{re.escape(rn)}\b", " ", t)
        words = [w for w in re.split(r"[\s,;|:]+", t) if w]
        while words and words[0].lower().strip(".") in NAME_STOP:    # "for Jane Doe", "buyer Jane Doe"
            words.pop(0)
        while words and words[-1].lower().strip(".") in TRAIL_STOP:  # "Jane Doe please"
            words.pop()
        nm = _name_like(" ".join(words), exclude) if words else None
        if nm:
            names.append(nm)
        elif words and len(words) == 1 and words[0][:1].isalpha():
            names.append(words[0])   # a single word: kept so we can ask for the FULL name

    # (c) unlabeled field lines directly under the command line (stop at the first non-field line —
    #     that is where the signature starts). Only used for fields still missing.
    after = own[m.end():].splitlines()
    taken = 0
    for ln in after:
        s = ln.strip()
        if not s:
            continue
        if re.search(rf"(?i){ALL_L}\s*[:=]", s) or taken >= 3:
            break
        if EMAIL_RE.fullmatch(s.strip("<>")):
            emails.append(s.strip("<>")); taken += 1; continue
        qm = re.fullmatch(r"(?i)(?:qty|quantity|x)?\s*[:=]?\s*(\d{1,3})\s*(?:units?|watch(?:es)?|pcs?)?", s)
        if qm:
            qtys.append(int(qm.group(1))); taken += 1; continue
        nm = _name_like(s, exclude)
        if nm and not names:
            names.append(nm); taken += 1; continue
        break

    # (d) billing email fallback: Reply-To / From of the forwarded buyer message
    if not [e for e in emails if not _internal(e)] and not emails:
        hdr = re.search(r"(?im)^\s*from:\s*(.+)$", fwd or "")
        if hdr:
            block = (fwd or "")[hdr.start():hdr.start() + 800]
            rt = re.search(r"(?im)^\s*reply-to:\s*(.+)$", block)
            for src, line in (("forwarded Reply-To", rt.group(1) if rt else ""), ("forwarded From", hdr.group(1))):
                cands = [e for e in EMAIL_RE.findall(line) if not _internal(e)]
                if cands:
                    emails.append(cands[0]); res["email_source"] = src
                    break

    # (e) phone fallback (task #3418): exactly one non-Exos phone number in the forwarded buyer email
    if not phones and fwd:
        cands = {re.sub(r"\D", "", ph)[-10:]: ph.strip() for ph in PHONE_RE.findall(fwd)}
        cands = {k: v for k, v in cands.items() if k not in EXOS_PHONES}
        if len(cands) == 1:
            phones.append(next(iter(cands.values()))); res["phone_source"] = "forwarded email"

    def uniq(xs, key=lambda x: x):
        out, seen = [], set()
        for x in xs:
            k = key(x)
            if k not in seen:
                seen.add(k); out.append(x)
        return out

    emails = uniq(emails, key=lambda e: e.lower())
    ext = [e for e in emails if not _internal(e)]
    emails = ext or emails   # an Exos address only counts if it is the only one given (e.g. a Jon test)
    qtys = uniq(qtys)
    names = uniq([n for n in names if n.lower() not in exclude], key=lambda n: n.lower())

    if len(names) > 1:
        res["ambiguous"].append("buyer name (I see " + " / ".join(names) + ")")
    elif names:
        if len(names[0].split()) < 2:
            res["missing"].append(f"the buyer's FULL name (I only have \"{names[0]}\")")
        else:
            res["name"] = names[0]
    else:
        res["missing"].append("the buyer's full name")
    if len(emails) > 1:
        res["ambiguous"].append("billing email (I see " + " / ".join(emails) + ")")
    elif emails:
        res["email"] = emails[0]
        res["email_source"] = res["email_source"] or "command"
    else:
        res["missing"].append("the buyer's billing email")
    if len(qtys) > 1:
        res["ambiguous"].append("quantity (I see " + " / ".join(str(q) for q in qtys) + ")")
    elif qtys:
        if qtys[0] < 1:
            res["ambiguous"].append("quantity (must be 1 or more)")
        else:
            res["qty"] = qtys[0]
    else:
        res["missing"].append("the quantity")
    # Task #3418 (10/5): an invoice must carry the customer's address, phone and email.
    phones = uniq(phones, key=lambda x: re.sub(r"\D", "", x)[-10:])
    if len(phones) > 1:
        res["ambiguous"].append("phone number (I see " + " / ".join(phones) + ")")
    elif phones:
        res["phone"] = phones[0]
    else:
        res["missing"].append("the buyer's phone number")
    addrs = uniq(addrs, key=lambda x: x.lower())
    if len(addrs) > 1:
        res["ambiguous"].append("mailing address (I see " + " / ".join(addrs) + ")")
    elif addrs:
        res["address"] = addrs[0]
    else:
        res["missing"].append("the buyer's mailing address (street, city, state, ZIP)")
    return res


# ── ledger (Reserved -> Invoiced -> Paid / Cancelled) ─────────────────────────
class _Locked:
    def __init__(self, path):
        self.path = path

    def __enter__(self):
        os.makedirs(os.path.dirname(self.path), exist_ok=True)
        real = os.path.abspath(self.path) == os.path.abspath(LEDGER)
        self.lf = open(os.path.join(WS, "logs", ".offering-ledger.lock") if real else self.path + ".lock", "w")
        fcntl.flock(self.lf, fcntl.LOCK_EX)
        self.data = jload(self.path, None) or {"_doc": "Offering reservations ledger (EXOS-FIN-003-0). Written by "
                                                        "invoice_command_exo3.py + qb_invoice.py watch.",
                                                "offerings": {}, "processed": {}}
        self.data.setdefault("offerings", {}); self.data.setdefault("processed", {})
        return self.data

    def save(self):
        tmp = self.path + ".tmp"
        with open(tmp, "w") as f:
            json.dump(self.data, f, indent=1)
        os.replace(tmp, self.path)

    def __exit__(self, *a):
        try:
            fcntl.flock(self.lf, fcntl.LOCK_UN); self.lf.close()
        except Exception:
            pass


def _wb_path(off):
    wb = (off or {}).get("workbook") or {}
    p = wb.get("path")
    return os.path.join(WS, p) if p and not os.path.isabs(p) else p


def workbook_manual_units(off, wb_path=None):
    """Rows hand-filled in the product workbook (Buyer set, no [exo:...] tag) — they count against the cap."""
    wb_cfg = (off or {}).get("workbook") or {}
    p = wb_path or _wb_path(off)
    if not p or not os.path.exists(p):
        return 0
    try:
        import openpyxl
        ws = openpyxl.load_workbook(p, read_only=True)[wb_cfg["sheet"]]
        c = wb_cfg["cols"]; n = 0
        for r in range(wb_cfg["first_row"], wb_cfg["last_row"] + 1):
            buyer = ws.cell(r, c["buyer"]).value
            notes = str(ws.cell(r, c["notes"]).value or "")
            if buyer and str(buyer).strip() and "[exo:" not in notes:
                n += 1
        return n
    except Exception as e:
        log(f"WARN workbook read failed {p}: {e}")
        return 0


def units_used(ledger, product):
    return sum(int(e.get("qty", 0)) for e in ledger["offerings"].get(product, {}).get("entries", [])
               if e.get("status") in ACTIVE)


def sync_workbook(off, ledger, product, wb_path=None):
    """Mirror ledger entries into the workbook's per-unit slot rows (one row per unit). Rows we own carry an
    [exo:<entry>] tag in Notes; hand-entered rows are never touched. Best-effort: ledger is the source of truth."""
    wb_cfg = (off or {}).get("workbook") or {}
    p = wb_path or _wb_path(off)
    if not p or not os.path.exists(p):
        return False
    try:
        import openpyxl
        wb = openpyxl.load_workbook(p)
        ws = wb[wb_cfg["sheet"]]; c = wb_cfg["cols"]
        rows = list(range(wb_cfg["first_row"], wb_cfg["last_row"] + 1))
        free = []
        for r in rows:
            notes = str(ws.cell(r, c["notes"]).value or "")
            if "[exo:" in notes:
                for k in ("buyer", "email", "reserved_date", "paid_date", "notes"):
                    ws.cell(r, c[k]).value = None
            if not (ws.cell(r, c["buyer"]).value or "").__str__().strip():
                free.append(r)
        for e in ledger["offerings"].get(product, {}).get("entries", []):
            if e.get("status") not in ("Reserved", "Invoiced", "Paid"):
                continue
            for k in range(int(e["qty"])):
                if not free:
                    log(f"WARN workbook full — {e['id']} unit {k + 1} not mirrored"); break
                r = free.pop(0)
                ws.cell(r, c["buyer"]).value = e["buyer"]
                ws.cell(r, c["email"]).value = e["email"]
                ws.cell(r, c["reserved_date"]).value = (e.get("reserved_at") or "")[:10]
                if e.get("paid_at"):
                    ws.cell(r, c["paid_date"]).value = e["paid_at"][:10]
                ws.cell(r, c["notes"]).value = (f"[exo:{e['id']}] {e['status']} · QB invoice #{e.get('invoice_doc') or '?'}"
                                               f" · {e.get('decision_id') or ''} · unit {k + 1}/{e['qty']}")
        tmp = p + ".tmp.xlsx"
        wb.save(tmp); os.replace(tmp, p)
        return True
    except Exception as e:
        log(f"WARN workbook sync failed {p}: {e}")
        return False


def ledger_mark(rec, status, ledger_path=None, offerings_path=None, when=None):
    """Hook for qb_invoice.py watch: status in Invoiced / Paid / Cancelled for an offering invoice."""
    off_meta = (rec or {}).get("offering") or {}
    if not off_meta.get("entry_id"):
        return False
    product = off_meta.get("product")
    lock = _Locked(ledger_path or LEDGER)
    data = lock.__enter__()
    try:
        ent = next((e for e in data["offerings"].get(product, {}).get("entries", []) if e["id"] == off_meta["entry_id"]), None)
        if not ent:
            return False
        ent["status"] = status
        ent[{"Invoiced": "invoiced_at", "Paid": "paid_at", "Cancelled": "cancelled_at"}.get(status, "updated_at")] = when or now()
        ent.setdefault("history", []).append({"at": now(), "status": status, "by": "qb_invoice watch"})
        lock.save()
        off = load_offerings(offerings_path)["offerings"].get(product)
        if not ledger_path or ledger_path == LEDGER:   # only the real ledger mirrors to the real workbook
            sync_workbook(off, data, product)
        _fleet("offering_status", status, {"entry": ent["id"], "invoice": ent.get("invoice_doc"), "product": product})
        log(f"ledger {ent['id']} -> {status} (invoice #{ent.get('invoice_doc')})")
        return True
    finally:
        lock.__exit__()


# ── CSL ───────────────────────────────────────────────────────────────────────
def csl_check(name, screen_fn=None):
    """('clear'|'hit'|'error', detail). Dispositioned-cleared parties pass. Errors fail CLOSED."""
    try:
        import csl_screen
        screen_fn = screen_fn or csl_screen.screen
        disp = jload(DISPOSITIONS, {})
        if (disp.get(csl_screen.norm(name)) or {}).get("status") == "cleared":
            return "clear", "previously dispositioned cleared"
        hits = screen_fn(name)
        if hits:
            h = hits[0]
            return "hit", f"{h.get('listed_name')} [{h.get('source')}] score {h.get('score')} {h.get('countries') or ''}".strip()
        return "clear", "no CSL match (11 U.S. lists)"
    except Exception as e:
        return "error", f"screen failed: {e}"


# ── handler ───────────────────────────────────────────────────────────────────
def _fmt_missing_reply(p, product_word, offerings):
    have = []
    if p.get("name"): have.append(f"Name: {p['name']}")
    if p.get("email"): have.append(f"Email: {p['email']}")
    if p.get("qty"): have.append(f"Qty: {p['qty']}")
    if p.get("phone"): have.append(f"Phone: {p['phone']}")
    if p.get("address"): have.append(f"Address: {p['address']}")
    asks = p["missing"] + [f"a single, unambiguous {a}" for a in p["ambiguous"]]
    ex = (f"@Exobot invoice {product_word or 'watch'}\nName: {p.get('name') or 'Jane Doe'}, "
          f"Email: {p.get('email') or 'jane@example.com'}, Qty: {p.get('qty') or 1}\n"
          f"Phone: {p.get('phone') or '214-555-0100'}\nAddress: {p.get('address') or '123 Main St, Dallas, TX 75201'}")
    return ("No invoice was created. I need:\n" + "\n".join(f"- {a}" for a in asks)
            + ("\n\nWhat I did get: " + "; ".join(have) if have else "")
            + "\n\nPlease forward the buyer's email again with the command filled in, for example:\n\n" + ex)


def handle(mid, sender, subject, body, *, dry_run=False, ledger_path=None, offerings_path=None, wb_path=None,
           prepare_fn=None, screen_fn=None, names_exclude=None):
    """Process one invoice command. Returns a result dict:
       outcome, reply (text for the sender or None), jon_alert (text or None), discord (text or None), detail."""
    cfg = load_offerings(offerings_path)
    offerings = cfg.get("offerings", {})
    sender = (sender or "").lower().strip()
    corr = uuid.uuid4().hex[:12]
    R = {"outcome": None, "reply": None, "jon_alert": None, "discord": None, "detail": {}, "corr": corr}
    ledger_path = ledger_path or LEDGER
    # dry-run never writes the REAL ledger/workbook; a sandbox (temp-copy) ledger still records so the
    # duplicate / cap paths can be exercised end-to-end without touching production state.
    write = (not dry_run) or os.path.abspath(ledger_path) != os.path.abspath(LEDGER)
    wb_ok = (not dry_run) or bool(wb_path)

    def done(outcome, level="info", **kw):
        R.update(outcome=outcome, **kw)
        log(f"[{outcome}] mid={mid} sender={sender} subj={subject!r} {json.dumps(R['detail'])[:400]}"
            + (" (dry-run)" if dry_run else ""))
        _fleet("invoice_command", outcome, dict(R["detail"], mid=mid, sender=sender, dry_run=dry_run), level=level, corr=corr)
        return R

    # 1. authorized senders only
    if sender not in {a.lower() for a in cfg.get("authorized_senders", [])}:
        return done("unauthorized", level="warn")

    # 2. duplicate protection (Message-ID)
    with _Locked(ledger_path) as L:
        prev = L["processed"].get(mid)
    if prev:
        R["detail"] = {"previous": prev}
        return done("duplicate")

    def mark_processed(outcome, **extra):
        lk = _Locked(ledger_path); d = lk.__enter__()
        try:
            d["processed"][mid] = dict({"at": now(), "outcome": outcome, "sender": sender, "subject": subject[:160]}, **extra)
            lk.save()
        finally:
            lk.__exit__()

    # 3. parse
    excl = set(names_exclude if names_exclude is not None else roster_names())
    p = parse_command(body, offerings, excl)
    R["detail"] = {k: p.get(k) for k in ("product", "name", "email", "qty", "email_source", "phone", "address")}
    if p["product"] is None:
        known = ", ".join(sorted(offerings)) or "(none)"
        word = p.get("product_word") or ""
        if not word:
            reply = f"No invoice was created: I couldn't tell which product to invoice. Use '@Exobot invoice <product>'; available: {known}."
        else:
            reply = (f"No invoice was created: \"{word}\" isn't an approved product in the invoice catalog. "
                     f"Available: {known}. New offerings need Jon's approval first (EXOS-FIN-003-0 Task 1).")
        if write: mark_processed("unknown_product")
        return done("unknown_product", reply=reply)
    off = offerings[p["product"]]
    if p["missing"] or p["ambiguous"]:
        R["detail"].update(missing=p["missing"], ambiguous=p["ambiguous"])
        if write: mark_processed("missing_info")
        return done("missing_info", reply=_fmt_missing_reply(p, p["product"], offerings))

    qty, name, email = p["qty"], p["name"], p["email"]
    cap = int(off["cap"])
    maxreq = int(off.get("max_per_request", cap))

    # 4. cap check + reservation placeholder (under the ledger lock -> no double-allocation)
    lk = _Locked(ledger_path); data = lk.__enter__()
    try:
        manual = workbook_manual_units(off, wb_path)
        used = units_used(data, p["product"]) + manual
        left = max(cap - used, 0)
        R["detail"].update(used=used, manual_workbook_units=manual, cap=cap, left_before=left)
        if qty > maxreq or qty > left:
            if write:
                data["processed"][mid] = {"at": now(), "outcome": "sold_out", "sender": sender, "subject": subject[:160]}
                lk.save()
            msg = (f"No invoice was created: the {off['name']} offering is SOLD OUT ({cap} of {cap} units reserved/invoiced/paid)."
                   if left == 0 else
                   f"No invoice was created: {name} asked for {qty}, but only {left} of {cap} {p['product']} unit(s) are left. "
                   f"Forward again with Qty: {left} (or fewer) if they want what's left.")
            if qty > maxreq and left >= qty:
                msg = f"No invoice was created: {qty} exceeds the per-request limit of {maxreq}."
            lk.__exit__(); lk = None
            return done("sold_out", reply=msg)
        same_buyer = [e for e in data["offerings"].get(p["product"], {}).get("entries", [])
                      if e.get("status") in ACTIVE and e.get("email", "").lower() == email.lower()]
        entry = {"id": f"{p['product'][:3].upper()}-{datetime.datetime.now().strftime('%y%m%d%H%M%S')}-{uuid.uuid4().hex[:4]}",
                 "product": p["product"], "buyer": name, "email": email, "qty": qty, "unit_price": float(off["unit_price"]),
                 "amount": round(qty * float(off["unit_price"]), 2), "status": "preparing", "mid": mid,
                 "requested_by": sender, "requested_at": now(), "email_source": p["email_source"]}
        if write:
            data["offerings"].setdefault(p["product"], {"entries": []})["entries"].append(entry)
            data["processed"][mid] = {"at": now(), "outcome": "preparing", "entry": entry["id"], "sender": sender,
                                      "subject": subject[:160]}
            lk.save()
    finally:
        if lk:
            lk.__exit__()

    def set_entry(status, **kw):
        if not write:
            entry.update(status=status, **kw); return
        lk2 = _Locked(ledger_path); d = lk2.__enter__()
        try:
            for e in d["offerings"][p["product"]]["entries"]:
                if e["id"] == entry["id"]:
                    e.update(status=status, **kw); entry.update(e)
            d["processed"][mid].update(outcome=status, **{k: v for k, v in kw.items() if k in ("invoice_doc", "decision_id")})
            lk2.save()
            snapshot = copy.deepcopy(d)
        finally:
            lk2.__exit__()
        if wb_ok and status in ("Reserved", "Cancelled", "failed", "csl_hold"):
            sync_workbook(off, snapshot, p["product"], wb_path)

    # 5. CSL denied-party screen (fail closed)
    cs, cdetail = csl_check(name, screen_fn)
    R["detail"]["csl"] = f"{cs}: {cdetail}"
    if cs != "clear":
        set_entry("csl_hold", csl=f"{cs}: {cdetail}")
        (lambda *a, **k: None if dry_run else _trace(*a, **k))(corr, [name, email, p["product"], str(qty)], "EXOS-FIN-003-0 Task 3 step 1: CSL screen before invoicing; hit/error stops",
               f"HOLD — no invoice ({cs})", write_back="mission-control/data/offering-ledger.json")
        what = "matched the U.S. Consolidated Screening List" if cs == "hit" else "could not be screened (screen error)"
        return done("csl_hold", level="warn",
                    reply=(f"No invoice was created for {name}: the buyer {what}. It's on hold for Jon's review; "
                           f"please don't reply to the buyer about payment until Jon clears it."),
                    jon_alert=(f"CSL HOLD on an @Exobot invoice request (no invoice created).\n"
                               f"Buyer: {name} <{email}> · {qty} x {p['product']} · requested by {sender}\n"
                               f"Screen: {cs} — {cdetail}\nSubject: {subject}\n"
                               f"If this is a name collision, disposition it in mission-control/data/csl-dispositions.json "
                               f"(or tell Exo) and have the request forwarded again."))

    # 6. prepare the QB invoice (UNSENT) + Approve & send Decision (approver Jon)
    unit = float(off["unit_price"]); total = round(qty * unit, 2)
    desc = off["desc_template"].format(qty=qty)
    note = (f"Offering: {off['name']} ({p['product']}) — {qty} x ${unit:,.2f} = ${total:,.2f}, {off['terms']}. "
            f"Requested by {sender} via '@Exobot invoice' email command (EXOS-FIN-003-0); CSL: clear. "
            f"Cap {cap}: {cap - left + qty} of {cap} units committed if approved; {left - qty} left. Ledger entry {entry['id']}."
            + (f" NOTE: this buyer already has {sum(int(e['qty']) for e in same_buyer)} unit(s) active "
               f"({', '.join('#' + str(e.get('invoice_doc')) for e in same_buyer)}) — check this is not a duplicate request."
               if same_buyer else ""))
    try:
        if prepare_fn is None:
            import qb_invoice
            prepare_fn = qb_invoice.prepare
        ns = argparse.Namespace(order=None, milestone=None, customer=name, email=email, amount=f"{total:.2f}", desc=desc,
                                phone=p.get("phone"), address=p.get("address"),
                                item=off["qb_item"], terms=off["terms"], note=note, doc_number=None, dry_run=dry_run,
                                qty=qty, unit_price=f"{unit:.2f}", offering={"product": p["product"], "entry_id": entry["id"]})
        rec = prepare_fn(ns) or {}
    except BaseException as e:   # qb_invoice raises SystemExit on validation errors
        err = str(e) or e.__class__.__name__
        set_entry("failed", error=err[:300])
        return done("error", level="error",
                    reply=(f"No invoice could be prepared for {name} ({qty} x {p['product']}): {err[:300]}\n"
                           f"Jon has been notified. Please don't forward it again until he confirms."),
                    jon_alert=(f"@Exobot invoice FAILED for {name} <{email}> ({qty} x {p['product']}), requested by {sender}.\n"
                               f"Error: {err[:400]}\nLedger entry {entry['id']} marked failed (not counted against the cap). "
                               f"Check QuickBooks for a half-created unsent invoice before retrying."))
    doc, dec = rec.get("doc_number"), rec.get("decision_id") or ("(dry-run: not filed)" if dry_run else None)
    set_entry("Reserved", invoice_doc=doc, qb_id=rec.get("qb_id"), decision_id=dec, reserved_at=now(), item=off["qb_item"])
    left_after = left - qty
    (lambda *a, **k: None if dry_run else _trace(*a, **k))(corr, [name, email, p["product"], str(qty)],
           "EXOS-FIN-003-0 Task 2->3: authorized sender + cap + CSL clear -> prepare UNSENT invoice; Jon approves send",
           f"prepared QB invoice #{doc} ${total:,.2f}, Decision {dec}", decision_id=dec if not dry_run else None,
           write_back="QB invoice (unsent) + mission-control/data/offering-ledger.json")
    R["detail"].update(invoice=doc, decision=dec, amount=total, left_after=left_after, entry=entry["id"])
    reply = (f"Invoice prepared (NOT sent) for {name}.\n\n"
             f"- Invoice #: {doc}\n- Amount: ${total:,.2f} ({qty} x ${unit:,.2f}, {off['terms']})\n"
             f"- Billing email: {email}" + (f" (taken from the {p['email_source']})" if p["email_source"] not in (None, "command") else "") + "\n"
             f"- Decision: {dec} (Jon: Approve & send / Hold / Cancel invoice)\n"
             f"- Logged: Reserved in the {p['product']} ledger (entry {entry['id']})\n"
             f"- Remaining: {left_after} of {cap} {p['product']} unit(s) available\n\n"
             f"QuickBooks emails the invoice with its pay link only after Jon approves the Decision. "
             f"Nothing has gone to the buyer yet."
             + (f"\n\nHeads-up: this buyer already had {sum(int(e['qty']) for e in same_buyer)} unit(s) reserved; "
                f"Jon will see that on the Decision." if same_buyer else ""))
    return done("success", reply=reply,
                discord=(f"🧾 @Exobot invoice: {name} · {qty} x {p['product']} · ${total:,.2f} · QB #{doc} prepared UNSENT → "
                         f"{dec} for your Approve & send. {left_after}/{cap} left. (requested by {sender})"))


# ── self-test (ExO3 Trusted-Evals) ────────────────────────────────────────────
def _selftest():
    global TELEMETRY
    TELEMETRY = False
    fails = []
    offs = load_offerings()["offerings"]
    ex = {"lizi quinn", "john quinn"}

    def chk(label, body, **want):
        got = parse_command(body, offs, ex)
        bad = {k: (got.get(k), v) for k, v in want.items() if (bool(got.get(k)) if v is True else got.get(k)) != v}
        print(("PASS " if not bad else "FAIL ") + label + ("" if not bad else f"  {bad}  missing={got['missing']} amb={got['ambiguous']}"))
        if bad:
            fails.append(label)

    chk("labeled one-line", "@Exobot invoice watch\nName: Jane Doe, Email: jane@example.com, Qty: 2\n\nLizi Quinn\nExos Aerospace",
        product="watch", name="Jane Doe", email="jane@example.com", qty=2)
    chk("inline", "@Exobot invoice watch Jane Doe jane@example.com 2", product="watch", name="Jane Doe", email="jane@example.com", qty=2)
    chk("inline qty word", "@exobot invoice watches for Mary Ann Smith, mas@x.org, qty 3", product="watch",
        name="Mary Ann Smith", email="mas@x.org", qty=3)
    chk("multi-line unlabeled + signature", "@Exobot invoice watch\nRobert Walker\nrw@postal.example\n1\n \n Lizi Quinn \n Exos Aerospace Systems",
        name="Robert Walker", email="rw@postal.example", qty=1)
    chk("labeled multi-line w/ billing email", "Hi Exo\n@Exobot invoice flight watch\nBuyer: Carol Adams\nBilling email: carol@x.com\nQuantity: one\nThanks!",
        product="watch", name="Carol Adams", email="carol@x.com", qty=1)
    fwd = ("@Exobot invoice watch Name: Ben Russell Qty: 1\n \n Lizi Quinn \n Exos Aerospace\n"
           "________________________________\n From: Ben Russell <brussell@tradeway.example> \n Sent: Friday, October 3, 2026 9:00 AM \n"
           " To: Lizi Quinn <lquinn@exosaero.com> \n Subject: RE: BLK3 flight watch \n I'd like one!\n")
    chk("email fallback to forwarded From", fwd, name="Ben Russell", email="brussell@tradeway.example", qty=1, email_source="forwarded From")
    fwd_rt = fwd.replace(" Sent:", " Reply-To: ben.personal@home.example\n Sent:")
    chk("email fallback prefers forwarded Reply-To", fwd_rt, email="ben.personal@home.example", email_source="forwarded Reply-To")
    chk("missing qty", "@Exobot invoice watch Name: Jane Doe Email: jane@example.com", qty=None)
    p = parse_command("@Exobot invoice watch Name: Jane Doe Email: jane@example.com" + "\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", offs, ex)
    ok = p["missing"] == ["the quantity"]
    print(("PASS " if ok else "FAIL ") + "missing qty -> asks only for quantity", p["missing"]); ok or fails.append("missing-list")
    # task #3418: phone + address are required; labeled fields parse, phone falls back to the forwarded buyer email
    chk("phone+address labeled", "@Exobot invoice watch\nName: Jane Doe, Email: jane@example.com, Qty: 1\nPhone: (214) 555-0100\nAddress: 12 Elm St, Dallas, TX 75201",
        phone="(214) 555-0100", address="12 Elm St, Dallas, TX 75201")
    p = parse_command("@Exobot invoice watch Name: Jane Doe Email: jane@example.com Qty: 1", offs, ex)
    ok = any("phone" in m for m in p["missing"]) and any("mailing address" in m for m in p["missing"])
    print(("PASS " if ok else "FAIL ") + "no phone/address -> asks for both", p["missing"]); ok or fails.append("contact-required")
    chk("phone from forwarded buyer signature (not Exos numbers)",
        "@Exobot invoice watch Name: Timm Anderson Email: mrrocket@mac.com Qty: 1 Address: 9 Oak Ln, Houston, TX 77002\n"
        "Lizi Quinn\nMob. +1-972-740-8355\n\nFrom: Timm Anderson <mrrocket@mac.com>\nSent: Friday\nLizi, I want a watch!\nTimm Anderson\n832-551-7138",
        phone="832-551-7138")
    p = parse_command("@Exobot invoice watch Jane jane@example.com 1", offs, ex)
    ok = p["name"] is None and any("FULL name" in m for m in p["missing"])
    print(("PASS " if ok else "FAIL ") + "single-word name -> asks for full name"); ok or fails.append("full-name")
    p = parse_command("@Exobot invoice watch Jane Doe a@x.com b@y.com 1", offs, ex)
    ok = any("billing email" in a for a in p["ambiguous"])
    print(("PASS " if ok else "FAIL ") + "two emails -> ambiguous"); ok or fails.append("amb-email")
    p = parse_command("@Exobot invoice watch Jane Doe jane@x.com qty 2 3 units", offs, ex)
    ok = any("quantity" in a for a in p["ambiguous"])
    print(("PASS " if ok else "FAIL ") + "two qtys -> ambiguous"); ok or fails.append("amb-qty")
    p = parse_command("@Exobot invoice hoodie Jane Doe jane@x.com 1", offs, ex)
    ok = p["product"] is None and p["product_word"] == "hoodie"
    print(("PASS " if ok else "FAIL ") + "unknown product detected"); ok or fails.append("unknown")
    q = ("@exo thanks, looks good\n\nOn Fri, Oct 3, 2026 at 9:00 AM Exo wrote:\n> You directed Exo by email:\n"
         "> \"@Exobot invoice watch Jane Doe jane@x.com 2\"\n")
    ok = not is_invoice_command(q)
    print(("PASS " if ok else "FAIL ") + "quoted earlier command in a reply chain does NOT fire"); ok or fails.append("quoted")
    ok = not is_invoice_command("@exo please prepare a response to this invoice question") and is_invoice_command("@exobot invoice watch")
    print(("PASS " if ok else "FAIL ") + "only '@Exobot invoice' fires (not 'invoice' elsewhere)"); ok or fails.append("trigger")

    # handler paths against a temp ledger, stubbed QB + CSL (no network, no QB writes)
    tmp = tempfile.mkdtemp(prefix="invcmd-selftest-")
    try:
        led = os.path.join(tmp, "ledger.json")
        calls = []

        def fake_prepare(ns):
            calls.append(ns)
            return {"doc_number": f"9{len(calls):03d}", "qb_id": "T", "decision_id": f"D-TEST{len(calls)}"}
        clear = lambda n: []
        hit = lambda n: [{"listed_name": "BAD ACTOR", "source": "SDN", "score": 1.0, "countries": "IR"}]
        kw = dict(ledger_path=led, wb_path=os.path.join(tmp, "none.xlsx"), prepare_fn=fake_prepare, screen_fn=clear, names_exclude=ex)
        L = "lquinn@exosaero.com"
        r = handle("<m1>", L, "Fw: watch", "@Exobot invoice watch Jane Doe jane@x.com 2\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        ok = r["outcome"] == "success" and "Remaining: 8 of 10" in r["reply"] and calls[-1].qty == 2 and calls[-1].amount == "1350.00" \
            and calls[-1].desc == "Exos flight watch pre-order (qty 2). Pre-order — ships after the BLK3 flight, targeted 2027." \
            and calls[-1].terms == "Due on receipt"
        print(("PASS " if ok else "FAIL ") + "success path", r["outcome"]); ok or fails.append("success")
        r = handle("<m1>", L, "Fw: watch", "@Exobot invoice watch Jane Doe jane@x.com 2\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        ok = r["outcome"] == "duplicate" and len(calls) == 1
        print(("PASS " if ok else "FAIL ") + "duplicate Message-ID -> no second invoice"); ok or fails.append("dup")
        r = handle("<m2>", "phil@exosaero.com", "Fw", "@Exobot invoice watch Jane Doe jane@x.com 1\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        ok = r["outcome"] == "unauthorized" and r["reply"] is None and len(calls) == 1
        print(("PASS " if ok else "FAIL ") + "unauthorized sender ignored"); ok or fails.append("unauth")
        r = handle("<m3>", L, "Fw", "@Exobot invoice watch Jane Doe 1\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        ok = r["outcome"] == "missing_info" and "billing email" in r["reply"] and len(calls) == 1
        print(("PASS " if ok else "FAIL ") + "missing info -> reply, no invoice"); ok or fails.append("missing")
        r = handle("<m4>", L, "Fw", "@Exobot invoice watch Bad Actor bad@x.com 1\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **dict(kw, screen_fn=hit))
        ok = r["outcome"] == "csl_hold" and r["jon_alert"] and len(calls) == 1
        print(("PASS " if ok else "FAIL ") + "CSL hit -> hold + Jon alert, no invoice"); ok or fails.append("csl")
        r = handle("<m5>", L, "Fw", "@Exobot invoice watch Err Person err@x.com 1\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201",
                   **dict(kw, screen_fn=lambda n: (_ for _ in ()).throw(RuntimeError("net down"))))
        ok = r["outcome"] == "csl_hold" and len(calls) == 1
        print(("PASS " if ok else "FAIL ") + "CSL screen error -> fail closed"); ok or fails.append("csl-err")
        r = handle("<m6>", L, "Fw", "@Exobot invoice watch Big Buyer big@x.com 9\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        ok = r["outcome"] == "sold_out" and "only 8 of 10" in r["reply"] and len(calls) == 1
        print(("PASS " if ok else "FAIL ") + "over cap -> 'only N left', no invoice"); ok or fails.append("cap")
        r = handle("<m7>", L, "Fw", "@Exobot invoice watch Last Eight le@x.com 8\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        r2 = handle("<m8>", L, "Fw", "@Exobot invoice watch One More om@x.com 1\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        ok = r["outcome"] == "success" and r2["outcome"] == "sold_out" and "SOLD OUT" in r2["reply"]
        print(("PASS " if ok else "FAIL ") + "cap reached -> sold out"); ok or fails.append("soldout")
        rec = {"offering": {"product": "watch", "entry_id": json.load(open(led))["offerings"]["watch"]["entries"][0]["id"]}}
        ledger_mark(rec, "Cancelled", ledger_path=led)
        r3 = handle("<m9>", L, "Fw", "@Exobot invoice watch One More om@x.com 1\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        ok = r3["outcome"] == "success" and "Remaining: 1 of 10" in r3["reply"]
        print(("PASS " if ok else "FAIL ") + "Cancel releases units back to the cap"); ok or fails.append("cancel")
        r = handle("<m10>", L, "Fw", "@Exobot invoice watch Q Err qe@x.com 1\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201",
                   **dict(kw, prepare_fn=lambda ns: (_ for _ in ()).throw(SystemExit("QB Item 'Sales' not found"))))
        st = json.load(open(led))
        ent = [e for e in st["offerings"]["watch"]["entries"] if e["mid"] == "<m10>"][0]
        ok = r["outcome"] == "error" and ent["status"] == "failed" and r["jon_alert"]
        print(("PASS " if ok else "FAIL ") + "QB prepare error -> failed (not counted), Jon alerted"); ok or fails.append("qb-err")
        r = handle("<m10>", L, "Fw", "@Exobot invoice watch Q Err qe@x.com 1\nPhone: 214-555-0100\nAddress: 1 Main St, Dallas, TX 75201", **kw)
        ok = r["outcome"] == "duplicate"
        print(("PASS " if ok else "FAIL ") + "failed message is NOT auto-retried (no half-created duplicate)"); ok or fails.append("qb-err-dup")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if fails:
        print("SELFTEST FAILED:", fails)
        return 1
    print("SELFTEST OK")
    return 0


def _simulate(path, dry_run, sandbox):
    """Run a saved .eml through the SAME path mc-email-command uses (minus IMAP/auth), printing the
    reply/alerts instead of sending them."""
    import email as _email
    import importlib.util
    spec = importlib.util.spec_from_file_location("mcc", os.path.join(WS, "scripts/mc-email-command.py"))
    mcc = importlib.util.module_from_spec(spec); spec.loader.exec_module(mcc)
    msg = _email.message_from_bytes(open(path, "rb").read())
    body = mcc.body_text(msg)
    frm = [a.lower() for _n, a in getaddresses([msg.get("From", "")])]
    kw = {"dry_run": dry_run}
    tmp = None
    if sandbox:
        tmp = tempfile.mkdtemp(prefix="invcmd-sim-")
        led = os.path.join(tmp, "offering-ledger.json")
        if os.path.exists(LEDGER):
            shutil.copy2(LEDGER, led)
        kw["ledger_path"] = led
        off = load_offerings()["offerings"].get("watch")
        src = _wb_path(off)
        if src and os.path.exists(src):
            kw["wb_path"] = os.path.join(tmp, os.path.basename(src)); shutil.copy2(src, kw["wb_path"])
    print("is_invoice_command:", is_invoice_command(body))
    r = handle((msg.get("Message-ID") or "").strip(), frm[0] if frm else "", str(msg.get("Subject", "")), body, **kw)
    print(json.dumps({k: v for k, v in r.items()}, indent=1))
    if tmp:
        print("sandbox dir:", tmp)
    return r


def main():
    a = sys.argv[1:]
    if "--selftest" in a:
        raise SystemExit(_selftest())
    if "--simulate" in a:
        _simulate(a[a.index("--simulate") + 1], "--dry-run" in a, "--sandbox" in a)
        return
    if "--status" in a:
        prod = a[a.index("--status") + 1] if len(a) > a.index("--status") + 1 else "watch"
        off = load_offerings()["offerings"][prod]
        d = jload(LEDGER, {"offerings": {}})
        d.setdefault("offerings", {})
        used = units_used(d, prod); manual = workbook_manual_units(off)
        print(f"{prod}: {used} ledger + {manual} hand-entered = {used + manual} of {off['cap']} committed")
        for e in d["offerings"].get(prod, {}).get("entries", []):
            print(f"  {e['id']} {e['status']:<10} {e['qty']} x {e['buyer']} <{e['email']}> #{e.get('invoice_doc')} {e.get('decision_id')}")
        return
    print(__doc__)


if __name__ == "__main__":
    main()
