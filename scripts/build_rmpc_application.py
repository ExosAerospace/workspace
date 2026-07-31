#!/usr/bin/env python3
"""RMPC Phase-1 application (v3 — FFR guidelines + iBOSS iSSI/iFEX integration) into one document."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
NAVY=RGBColor(0x0b,0x1a,0x3a); GOLD=RGBColor(0xb8,0x86,0x1f); DARK=RGBColor(0x1a,0x1a,0x1a); GREY=RGBColor(0x55,0x55,0x55); RED=RGBColor(0x8a,0x1a,0x1a)
OUT="drafts/Exos_RMPC_Phase1_Application.docx"

S1="""The payload is a self-contained robotic propellant-transfer ORU: a single CrossLink-compatible unit containing a source tank, a receiver tank, and an androgynous fluid coupling. The FFR arm grapples and repositions the ORU via its passive CrossLink cleat; the source-to-receiver fluid mate is made by the iBOSS iSSI/iFEX interface, and the demonstration is the active robotic coupling mate/demate with a metered propellant transfer — the gating step for in-space refueling and depots — not passive carriage.

Development is well past concept. Tank-to-tank transfer has already been demonstrated at the component/breadboard level on Exos's propellant-transfer test set, retiring the fluid path, metering, and coupling-mate physics on the ground. The Scorpius Type-V linerless composite tankage is TRL-9, with flight heritage as fuel and LOX containment on the Intuitive Machines Odysseus and Athena lunar missions. The iBOSS iSSI interface was demonstrated in space in 2022; its iFEX fuel-exchange variant is an androgynous, multi-propellant coupling. What remains is packaging these flight-relevant elements into a sub-12.5 kg CrossLink ORU with self-managed thermal control within the interface power limit — achievable in eight months, with the ~2028 FFR flight providing qualification margin.

An on-orbit robotic-arm demonstration is essential because only sustained microgravity reveals autonomous coupling mate/demate and metered transfer with true zero-g fluid behavior across the interface — leak-tight sealing, contact and alignment under compliance — which ground fixtures and seconds-long suborbital arcs cannot reproduce.

For SWaP and thermal reasons we use a benign storable working fluid on orbit; iFEX is multi-propellant, so the same interface carries forward to cryogenic depots, and cryogenic fluid physics is matured on Exos's parallel BLK3 suborbital transfer flight.

Evidence: ground transfer data on the Exos test set (diligence); Scorpius Type-V flight heritage in public Intuitive Machines IM-1/IM-2 records; iSSI 2022 in-space demonstration (iBOSS)."""

S2="""The demonstration advances NASA's 2026 Civil Space Shortfalls SF32 (affordable, resilient supply chain), SF28 (robotics, automation, autonomy), and SF31 (technology demonstration environments): autonomous robotic propellant-coupling mate-and-transfer, via the flight-demonstrated iBOSS iSSI/iFEX interface, is the gating fluid step for refueling and depots.

A flight evaluates what ground and suborbital testing cannot: sustained-microgravity fluid behavior across the coupling — ullage position, phase distribution, and transfer dynamics — together with robotic mate dynamics, contact and alignment loads, and leak-tight sealing achieved in true zero-g with no ground fixturing. Suborbital arcs give only seconds of low-g; a complete metered transfer needs minutes of continuous microgravity.

Our payload needs a flight now because the fluid path is ground-demonstrated, the Type-V tankage is flight-proven, and the iSSI interface has space heritage — the sole remaining risk is the robotic mate-and-transfer on orbit, and closing it now positions U.S. depot and refueling architectures ahead of accelerating demand."""

S3="""We build on flight-proven and flight-demonstrated hardware and reduce new development to ORU integration. The fluid architecture derives from Exos's existing propellant-transfer test set; the tankage is the TRL-9 Scorpius Type-V; the robotic fluid coupling is the iBOSS iSSI/iFEX interface, provided via iBOSS.space US LLC. Exos primes integration; Scorpius supplies Type-V tankage; Princeton CryoTech supports fluid management; iBOSS supplies the iSSI/iFEX coupling.

We design explicitly to the FFR guidelines: mass under 12.5 kg; within the CrossLink cross-section; a passive CrossLink Cleat interface (28 V/3 A, RS422) for arm handling, since the arm mates via CrossLink, with iSSI/iFEX as the internal source-to-receiver fluid coupling; thermally isolated with self-managed thermal (MLI blanket plus heaters) inside the interface power limit; all moving and deployable elements tethered against debris; and design to launch loads and the LEO thermal and radiation environment. The receiver is integral to the single ORU, since only one deck CrossLink position is available.

ConOps: the arm installs the ORU, the iSSI/iFEX coupling mates the source and receiver, a metered transfer follows, then demate — an active manipulation.

Risks and mitigations: iFEX is an emerging interface — a ground mate/demate and transfer campaign plus a benign storable fluid; SWaP and thermal within 12.5 kg and the power limit — a miniaturized dual-tank and storable fluid; foreign ownership and export control — contract through iBOSS.space US LLC under our ITAR process; arm reach and workspace — design to the published FFR CAD and confirm in Phase 2."""

S4="""The payload combines background IP owned separately by each team member, integrated by Exos as prime.

Exos Aerospace owns the system-integration and flight-test IP, the propellant-transfer test set, and the ORU architecture, including the foreground integration IP developed for this demonstration; Exos is also an 18% owner of Scorpius. Scorpius Space Launch Company owns the Type-V linerless composite tank IP, used under our teaming arrangement. Princeton CryoTech owns the cryogenic fluid-management methods it contributes. iBOSS owns the patented iSSI and iFEX interface IP, licensed to the team through iBOSS.space US LLC for the demonstration.

The FFR platform, the Motiv robotic arm, and the NASA-furnished separable deck interface are NASA/provider IP; we design to the published interface specifications, as permitted by the challenge.

Permissions: all technical exchange is governed by an executed mutual NDA; a teaming agreement allocates background versus foreground IP and data rights, with each party retaining its background IP. Controlled technical data is exchanged under our ITAR/export-control process, contracting through iBOSS.space US LLC. Any SBIR-derived data carries applicable SBIR data rights.

Post-challenge: Exos offers the demonstrated capability as an integrated flight-test and refueling service; Scorpius Type-V tankage and the iSSI/iFEX interface remain commercially available and licensable, enabling depot, satellite-servicing, and refueling primes to adopt the capability through teaming or license."""

S5="""Two deliberate design choices and one integration note warrant explanation.

First, working fluid. For the orbital demonstration we baseline a benign, storable working fluid rather than a cryogen. Within the mandatory 12.5 kg mass, an interface-limited power budget (about 28 V/3 A), and full thermal isolation from the bus — and given the payload persists on orbit from an early-2028 launch until the FFR rendezvous — sustaining a cryogen is not credible. A storable fluid demonstrates the ISAM-critical step within the box; because the iBOSS iFEX coupling is multi-propellant (liquid and cryogenic), the same interface carries forward to cryogenic depots, and cryogen-specific fluid physics is retired on Exos's parallel BLK3 suborbital transfer flight. This is a scoping choice, not a capability gap.

Second, manipulation. Because the FFR arm interfaces only through the CrossLink cleat and only one deck CrossLink position is available, the receiver is integral to our ORU, and the coupling mate/demate is performed by the iSSI/iFEX interface driven by the arm's install motion — an active manipulation, not passive carriage.

Integration and maturity: we adopt the iBOSS iSSI interface, demonstrated in space in 2022, with its iFEX fuel-exchange variant as the fluid coupling; iFEX is an emerging interface that this flight would help mature. iBOSS participates in the FFR mission, aligning our payload with the platform ecosystem. Controlled data is exchanged through iBOSS.space US LLC under our export-control process.

Assumptions (per the guidelines): arm reach and workspace from the provided CAD; a LEO total-ionizing-dose assumption with margin; reasonable EMI/EMC compliance; self-managed hot- and cold-case thermal within the interface power limit; and robustness to changes in inclination, duration, and downlink.

Our approach targets the hardest, highest-value ISAM step and builds on flight-proven TRL-9 tankage, a ground-demonstrated transfer, and a space-demonstrated robotic interface — retiring real refueling and depot risk rather than a laboratory analog."""

FIELDS=[
 ("1","Payload Development Status",300,"Explain how this payload has been successfully tested (component or breadboard level) to demonstrate current development status and readiness for a flight test in eight months. Describe how the payload would benefit from an on-orbit demonstration with a robotic arm. Provide evidence.",S1),
 ("2","Benefit of Hosted Orbital Flight",150,"How could an on-orbit flight demonstration with the FFR robotic arm advance ISAM capabilities? What would a flight test enable you to evaluate that ground testing cannot? Why does your payload need a test flight now?",S2),
 ("3","Payload Development Plan",250,"How you plan to build your payload, how you meet the technical-guideline requirements, and the risks and mitigations.",S3),
 ("4","Intellectual Property",None,"Who owns the IP; permissions for existing/off-the-shelf technology; team-member ownership; and how others may utilize the technology after the challenge.",S4),
 ("5","Additional Information",None,"Outstanding issues; where the payload does not meet technical-guideline requirements and why; and how the approach better meets the objectives or exceeds the guidelines.",S5),
]

doc=Document()
for s in doc.sections:
    s.left_margin=s.right_margin=Inches(0.85); s.top_margin=s.bottom_margin=Inches(0.7)
def p(text="",size=10.5,bold=False,color=DARK,after=6,before=0,italic=False):
    pa=doc.add_paragraph(); pa.paragraph_format.space_after=Pt(after); pa.paragraph_format.space_before=Pt(before)
    r=pa.add_run(text); r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color; r.font.name="Calibri"
    return pa

p("EXOS AEROSPACE  —  NASA TechLeap RMPC Phase-1 Application (Draft, rev. FFR guidelines + iBOSS iSSI/iFEX)", size=14.5, bold=True, color=NAVY, after=1)
p("Robotic Propellant-Transfer Coupling Demonstration (self-contained CrossLink ORU)  ·  Team: Exos Aerospace (prime) · Scorpius Space Launch Company · Princeton CryoTech · iBOSS (via iBOSS.space US LLC — iSSI/iFEX)", size=10, italic=True, color=GREY, after=2)
p("PROPRIETARY — Contains Exos/partner proprietary technical information.", size=8.5, bold=True, color=RED, after=4)

for num,title,limit,prompt,body in FIELDS:
    wc=len(body.split())
    lbl=f"{wc} words" if limit is None else f"{limit}-word field · {wc} words"
    over=" OVER" if (limit and wc>limit) else ""
    pa=doc.add_paragraph(); pa.paragraph_format.space_before=Pt(12); pa.paragraph_format.space_after=Pt(1)
    r=pa.add_run(f"{num}.  {title}"); r.font.size=Pt(13); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name="Calibri"
    rl=pa.add_run(f"    ({lbl}{over})"); rl.font.size=Pt(9); rl.font.bold=True; rl.font.color.rgb=(RED if over else GOLD); rl.font.name="Calibri"
    p(prompt, size=8.5, italic=True, color=GREY, after=5)
    for para in body.strip().split("\n\n"):
        p(para.strip(), size=10.5, color=DARK, after=6)
    print(f"  S{num}: {wc} words (limit {limit}){over}")

def head6(num,title,lbl):
    pa=doc.add_paragraph(); pa.paragraph_format.space_before=Pt(12); pa.paragraph_format.space_after=Pt(1)
    r=pa.add_run(f"{num}.  {title}"); r.font.size=Pt(13); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name="Calibri"
    rl=pa.add_run(f"    ({lbl})"); rl.font.size=Pt(9); rl.font.bold=True; rl.font.color.rgb=GOLD; rl.font.name="Calibri"
def bullet(text,size=10.5,indent=0.28,lead="•  "):
    pa=doc.add_paragraph(); pa.paragraph_format.left_indent=Inches(indent); pa.paragraph_format.space_after=Pt(3)
    r=pa.add_run(lead); r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=GOLD; r.font.name="Calibri"
    r2=pa.add_run(text); r2.font.size=Pt(size); r2.font.color.rgb=DARK; r2.font.name="Calibri"

PP=[
 "Month 1: Execute team agreements and the iBOSS license via iBOSS.space US LLC; complete the ITAR/export screen on iSSI/iFEX data; pull the FFR CAD; freeze requirements to the CrossLink interface, 12.5 kg, power, and thermal limits (SRR).",
 "Months 1-2: Lock interface control documents with iBOSS and the FFR team (arm reach/workspace, deck interface, iSSI/iFEX fluid coupling).",
 "Months 2-3: PDR - ORU architecture: dual source/receiver tanks, iSSI/iFEX coupling, passive CrossLink cleat, metering, MLI/heater thermal control.",
 "Month 3: CDR; order long-lead items (iFEX unit, Type-V cartridge, valves).",
 "Months 4-5: Build and integrate; run the ground robotic mate/demate and transfer campaign on Exos's test set with an arm emulator - the primary risk-retirement gate.",
 "Months 5-6: Environmental qualification - vibration, thermal-vacuum, EMI/EMC, radiation assessment.",
 "Months 6-7: End-to-end functional demonstration (mate, metered transfer, demate); anomaly resolution.",
 "Months 7-8: Flight-readiness review; deliver flight-ready payload and data package; support FFR integration.",
 "Risk mitigation: baseline a benign storable fluid (SWaP/thermal); retire iFEX-mate risk early via the Month 4-5 ground campaign; hold schedule margin through parallel workstreams and reuse of flight-proven Type-V tankage.",
]
head6("6","Project Plan","200-word field · %d words" % sum(len(b.split()) for b in PP))
p("In bulleted chronological order, the project plan and timeline to a flight-ready payload in eight months, including risk mitigation.", size=8.5, italic=True, color=GREY, after=5)
for b in PP: bullet(b)
print("  S6 (Project Plan):", sum(len(b.split()) for b in PP), "words")

BUDG=[
 ("Phase 1 - Design & Planning (Months 1-3)","$130K",[
   ("Systems engineering & design (SRR/PDR/CDR, ICDs)","$80K"),
   ("iBOSS iSSI/iFEX license + early integration support","$35K"),
   ("ITAR/export screening, teaming & legal","$15K")]),
 ("Phase 2 - Procurement & Build (Months 3-5)","$355K",[
   ("iBOSS iFEX/iSSI flight coupling hardware","$80K"),
   ("Scorpius Type-V source + receiver cartridge (flight)","$115K"),
   ("ORU structure, valves, metering, sensors","$60K"),
   ("CrossLink cleat, avionics, harness","$30K"),
   ("Thermal control (MLI, heaters, controllers)","$20K"),
   ("Integration labor","$50K")]),
 ("Phase 3 - Test & Qualification (Months 5-7)","$120K",[
   ("Ground robotic mate/demate + transfer campaign (arm emulator, fixtures)","$40K"),
   ("Environmental qualification (vibration, thermal-vacuum, EMI/EMC, radiation)","$50K"),
   ("Test labor + anomaly resolution","$30K")]),
 ("Phase 4 - Flight Readiness & Delivery (Months 7-8)","$45K",[
   ("Flight-readiness review + data package","$25K"),
   ("FFR integration support + travel","$20K")]),
]
head6("7","Proposed Budget","line items by phase · total $725K")
p("Specific line items by phase, including contingency for risk mitigation.", size=8.5, italic=True, color=GREY, after=5)
for phase,sub,items in BUDG:
    pa=doc.add_paragraph(); pa.paragraph_format.space_before=Pt(5); pa.paragraph_format.space_after=Pt(2)
    r=pa.add_run(phase+"  -  "+sub); r.font.size=Pt(10.5); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name="Calibri"
    for item,amt in items: bullet(f"{item}:  {amt}", size=10, indent=0.45, lead="- ")
def totline(label,amt,color=DARK,bold=True):
    pa=doc.add_paragraph(); pa.paragraph_format.space_after=Pt(1); pa.paragraph_format.space_before=Pt(4)
    r=pa.add_run(f"{label}:  {amt}"); r.font.size=Pt(10.5); r.font.bold=bold; r.font.color.rgb=color; r.font.name="Calibri"
totline("Base subtotal","$650K")
totline("Contingency / risk reserve (~11.5%)","$75K")
totline("TOTAL","$725K",color=GOLD)
p("Funding: ~$500K prize-funded + ~$225K team cost-share/in-kind (reused flight-proven Type-V heritage tooling, Exos's existing propellant-transfer test set, iBOSS iSSI background). ROM figures - firm iFEX and Type-V with iBOSS/Scorpius quotes before submission.", size=8.5, italic=True, color=GREY, before=3)

BUDGET_DETAIL="""Our rough-order total to develop the payload and demonstration is approximately $725,000, which exceeds the $500,000 award. We fund the ~$225,000 balance through team cost-share and in-kind contributions, not additional cash requirements: reused flight-proven hardware and facilities (Scorpius Type-V heritage tooling and Exos's existing propellant-transfer test set, already built and demonstrated) contributed in-kind; iBOSS iSSI background hardware and IP contributed through the teaming arrangement; and Exos internal R&D and partner cost-share covering any residual cash. The $500,000 award funds the direct cash build - the iFEX/iSSI flight coupling, the Scorpius Type-V flight cartridge, ORU fabrication, qualification testing, and integration - while the team's reuse of proven assets covers the remainder. The ~$75,000 contingency (11.5%) is held against the primary risk, iFEX-coupling maturity."""
head6("8","Budget Detail - Funding Above $500K","~%d words" % len(BUDGET_DETAIL.split()))
p("If total cost exceeds $500,000, explain how you plan to fund the remaining costs.", size=8.5, italic=True, color=GREY, after=5)
p(BUDGET_DETAIL, size=10.5, color=DARK)

p("Open items before Aug 12 submission: formalize the iBOSS teaming/license (via iBOSS.space US LLC) and the foreground-IP clause across Exos/Scorpius/Princeton/iBOSS; run the ITAR/export screen on iSSI/iFEX technical-data exchange; pull FFR CAD to fix arm reach/workspace assumptions; firm iFEX and Type-V costs with vendor quotes.", size=8.5, italic=True, color=RED, before=10)
doc.save(OUT); print("BUILT", OUT)
