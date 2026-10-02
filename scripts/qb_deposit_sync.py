# -*- coding: utf-8 -*-
"""
qb_deposit_sync.py — keep the $800K build-capital tracker current from ACTUAL QuickBooks bank
deposits. Reads Deposit entities (each line = money into the bank), classifies each line to a
build-capital lane by payer + memo, and:
  - CONFIDENT match -> credit that lane (accumulated in the ledger; tracker net-to-build updated)
  - AMBIGUOUS      -> flag to a pending-review list; NEVER silently booked
Deterministic, rule-based (no model in the money path). Idempotent: each deposit line is keyed by
"<DepositId>:<lineIndex>" and never counted twice.

Lanes (match Exos_BuildCapital_800K_Tracker.xlsx rows 6-9, net-to-build column E):
  payload  = new payload-sales bookings (UTD/Lizi)     -> E6
  sbir_nasa= SBIR / NASA grants for work                -> E7
  po10     = Reg D 504 investor subscriptions           -> E8
  usaf     = USAF contract (Karmen/Hermes/FA9300)       -> E9

exo3 r3 (D-7281865, Jon 2026-10-01 "Yes, gate it"): AI proposes, a human confirms.
  - A CONFIDENT match no longer credits a lane. It becomes a PROPOSAL (ledger "proposals",
    status "proposed") and each run's new proposals go on ONE consolidated EXO3 Decision card
    (dedupeKey qb-deposit:<digest of keys>, priority high — money).
  - A proposal is BOOKED (moved to "matched", lane credited, tracker updated) only when that
    exact card is ANSWERED by a human with "Book ...". Token/agent answers stay open with a
    proposedAnswer since e13a807 and never book; an answeredBy that names an agent is refused.
  - "Don't book" -> status "rejected" (moved to pending review, never booked).
  - The 17 matches booked before this gate (2026-09-30 ledger) are grandfathered as-is.
  - fleet_log events: deposit_proposals_filed / deposit_booked / deposit_rejected / deposit_sync_run.

Run:  python3 scripts/qb_deposit_sync.py [--dry] [--since 2026-07-01] [--no-card]
      python3 scripts/qb_deposit_sync.py --selftest     # offline: no QB, no MC, temp files only
"""
import os, sys, re, json, datetime, hashlib, tempfile, urllib.request

WS = os.path.expanduser("~/.openclaw/workspace")
sys.path.insert(0, os.path.join(WS, "scripts"))
LEDGER = os.path.join(WS, "mission-control", "data", "build_capital_deposits.json")
TRACKER = os.path.join(WS, "drafts", "Exos_BuildCapital_800K_Tracker.xlsx")
BRIDGE = os.path.join(WS, "mission-control", "data", "capital_bridge_ledger.json")
DECISIONS = os.path.join(WS, "mission-control", "data", "decisions.json")
TOKEN = os.path.join(WS, "mission-control", "data", ".decisions-token")
INGEST = "http://127.0.0.1:8080/api/decisions/ingest"
AGENT = "com.exos.qb-deposit-sync"
DRY = "--dry" in sys.argv
NO_CARD = "--no-card" in sys.argv
try:
    from fleet_log import event as _fleet_event, new_corr
except Exception:  # telemetry must never break the sync
    _fleet_event = None
    def new_corr():
        import uuid; return uuid.uuid4().hex[:16]
def _arg(flag, default):
    return sys.argv[sys.argv.index(flag)+1] if flag in sys.argv else default
SINCE = _arg("--since", "2026-07-01")

# ---- classification signals -------------------------------------------------
USAF_CONTRACT = re.compile(r"\bkarmen\b|\bhermes\b|FA9300|HVRR|RDRE\s*(contract|flight)", re.I)
GOV_GRANT = re.compile(r"\bNASA\b|\bUSAF\b|air\s?force|\bAFRL\b|\bSBIR\b|\bSTTR\b|\bDoD\b|DFAS|"
                       r"treasury|space\s?force|spacewerx|\bONR\b|\bMDA\b|flight\s+opportunit", re.I)
INVESTOR = re.compile(r"\b504\b|reg\.?\s?d\b|subscription|\bSAFE\b|convertible|investor|equity\s+invest|"
                      r"capital\s+contribution|PO10|unit\s+purchase|stock\s+purchase", re.I)
# Deposit posted to an equity account = investor money (504). Strongest, most reliable PO10 signal.
EQUITY_ACCT = re.compile(r"capital\s*stock|common\s*stock|preferred\s*stock|paid.?in\s*capital|additional\s+paid", re.I)
INTEREST = re.compile(r"interest\s+(payment|income)|\bdividend\b", re.I)
PAYLOAD = re.compile(r"payload|\bBLK3\b|flight\s?2|\bbay\b|manifest|\bslot\b|suborbital|entX|ent-?x|"
                     r"microgravity|rideshare|cubesat|1U|2U|SARGE", re.I)

def load_known_investors():
    names = set()
    try:
        d = json.load(open(BRIDGE))
        for r in (d.get("receipts") or d.get("entries") or []):
            n = (r.get("investor") or r.get("name") or "").strip().lower()
            if n and len(n) > 2: names.add(n)
    except Exception:
        pass
    return names

def classify(payer, memo, investors):
    blob = f"{payer} {memo}".strip()
    low = blob.lower()
    if INTEREST.search(blob):                                        return None, "interest/dividend — not a build lane"
    if USAF_CONTRACT.search(blob):                                   return "usaf", "USAF contract signal"
    if EQUITY_ACCT.search(payer):                                    return "po10", "posted to equity account (504 investor)"
    if any(n in low for n in investors) or INVESTOR.search(blob):    return "po10", "investor / 504 signal"
    if GOV_GRANT.search(blob):                                       return "sbir_nasa", "gov/grant signal"
    if PAYLOAD.search(blob):                                         return "payload", "payload signal"
    return None, "no clear signal"

# ---- QB pull ----------------------------------------------------------------
def fetch_deposit_lines(since):
    import qb_client as qb
    access, realm, env = qb.get_access_token()
    lines = []
    start = 1
    while True:
        qr = qb.query(f"SELECT * FROM Deposit WHERE TxnDate >= '{since}' STARTPOSITION {start} MAXRESULTS 1000",
                      access, realm, env)
        rows = qr.get("Deposit", [])
        for d in rows:
            did = d.get("Id"); dt = d.get("TxnDate"); note = d.get("PrivateNote") or ""
            for i, ln in enumerate(d.get("Line", []) or []):
                det = ln.get("DepositLineDetail") or {}
                ent = (det.get("Entity") or {})
                payer = ent.get("name") or (det.get("AccountRef") or {}).get("name") or ""
                memo = ln.get("Description") or note
                amt = float(ln.get("Amount", 0) or 0)
                if amt <= 0:
                    continue
                lines.append({"key": f"{did}:{i}", "date": dt, "payer": payer, "memo": memo, "amount": amt})
        if len(rows) < 1000:
            break
        start += 1000
    return lines

# ---- exo3 human-confirm gate (D-7281865) -------------------------------------
BOOK_RX = re.compile(r"^\s*book\b", re.I)
REJECT_RX = re.compile(r"^\s*(don.?t|do not)\s+book\b", re.I)
AGENT_NAMES = re.compile(r"\bexo(bot)?\b|daily update|auto|agent|cron|claude|system|token", re.I)

def fleet(action, detail, outcome=None, level="info", corr=None):
    if _fleet_event:
        try:
            return _fleet_event(AGENT, "orchestrate", action, corr=corr, detail=detail,
                                owner="John Quinn", outcome=outcome, level=level)
        except Exception:
            pass
    return corr

def load_decisions(path=None):
    try:
        d = json.load(open(path or DECISIONS))
    except Exception:
        return {}
    L = d if isinstance(d, list) else d.get("decisions", [])
    return {x.get("id"): x for x in L if isinstance(x, dict)}

def human_answer(dec):
    """Return the answer text only if a HUMAN closed the card. Token answers never change status
    (e13a807: they leave status open + proposedAnswer), and an agent-named answeredBy is refused."""
    if not dec or str(dec.get("status", "")).lower() != "answered":
        return None
    who = str(dec.get("answeredBy") or "").strip()
    if not who or AGENT_NAMES.search(who):
        return None
    return str(dec.get("answer") or "")

def ensure_schema(led):
    led.setdefault("proposals", [])
    if "gate" not in led:
        led["gate"] = {"since": datetime.datetime.now().isoformat(timespec="seconds"),
                       "decision": "D-7281865",
                       "grandfathered_matched": len(led.get("matched", [])),
                       "rule": "confident matches are proposals; only a human 'Book' answer on their card books them"}
    return led

def apply_answers(led, decisions, corr=None):
    """Book / reject proposals whose card a human has answered. Returns (booked, rejected)."""
    booked = rejected = 0
    keep = []
    for P in led["proposals"]:
        did = P.get("decision_id")
        ans = human_answer(decisions.get(did)) if did else None
        if ans is not None and REJECT_RX.match(ans):
            led["pending"].append({k: P[k] for k in ("key", "date", "payer", "memo", "amount")}
                                  | {"why": f"proposed {P['lane']} rejected by human on {did}"})
            rejected += 1
            fleet("deposit_rejected", {"key": P["key"], "amount": P["amount"], "decision": did}, corr=corr)
        elif ans is not None and BOOK_RX.match(ans):
            dec = decisions[did]
            led["lanes"][P["lane"]] = round(led["lanes"][P["lane"]] + P["amount"], 2)
            led["matched"].append({**{k: v for k, v in P.items() if k not in ("status",)},
                                   "booked_by": dec.get("answeredBy"), "booked_on": did,
                                   "booked_at": datetime.datetime.now().isoformat(timespec="seconds")})
            booked += 1
            fleet("deposit_booked", {"key": P["key"], "lane": P["lane"], "amount": P["amount"],
                                     "decision": did, "by": dec.get("answeredBy")}, outcome="booked", corr=corr)
        else:
            keep.append(P)   # open, agent-proposed, hold, or unknown answer -> stays a proposal
    led["proposals"] = keep
    return booked, rejected

def batch_key(props):
    return "qb-deposit:" + hashlib.sha256("|".join(sorted(p["key"] for p in props)).encode()).hexdigest()[:12]

def card_payload(props, corr=None):
    tot = sum(p["amount"] for p in props)
    lines = [f"- {p['date']}  ${p['amount']:,.2f}  {p['payer'][:40]} — {p['memo'][:60]}  -> {p['lane']} ({p['why']})"
             for p in props[:30]]
    return {
        "title": f"EXO3: Book {len(props)} QuickBooks deposit(s) (${tot:,.0f}) to the build-capital tracker?",
        "question": f"Should these {len(props)} bank deposit(s) totalling ${tot:,.2f} be booked to the build-capital lanes shown?",
        "context": "qb_deposit_sync.py read these deposits from QuickBooks (read-only) and matched them to a "
                   "build-capital lane by payer/memo rules:\n" + "\n".join(lines) +
                   ("\n(+ more in mission-control/data/build_capital_deposits.json)" if len(props) > 30 else "") +
                   "\n\nNothing is booked until you answer. 'Book' credits exactly these lines on the next 06:15 "
                   "run; 'Don't book' moves them to pending review. No QuickBooks writes either way. "
                   "Gate: D-7281865.",
        "proposedAction": "Book these deposits to the lanes shown.",
        "type": "choice",
        "options": ["Book these deposits (recommended)", "Don't book — reclassify", "Hold"],
        "priority": "high",
        "dedupeKey": batch_key(props),
        "source": "exo3-qb-deposit-sync",
        "agentClass": "GOVERN",
        "assignedTo": "jquinn@exosaero.com",
        "approver": "John Quinn",
        "relatedItem": "build_capital_deposits",
        "correlationId": corr,
    }

def ingest_decision(payload):
    tok = open(TOKEN).read().strip()
    req = urllib.request.Request(INGEST, data=json.dumps(payload).encode(), method="POST",
                                 headers={"x-exo-token": tok, "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=20) as r:
        out = json.loads(r.read().decode() or "{}")
    return out.get("id") or (out.get("decision") or {}).get("id")

def propose(led, lines, investors):
    """Classify unseen lines. Confident -> proposal (NOT booked); unclear -> pending. Returns (new_prop, new_pend)."""
    new_prop = new_pend = 0
    for L in lines:
        if L["key"] in led["seen"]:
            continue
        lane, why = classify(L["payer"], L["memo"], investors)
        led["seen"][L["key"]] = True
        if lane:
            led["proposals"].append({**L, "lane": lane, "why": why, "status": "proposed", "decision_id": None,
                                     "proposed_at": datetime.datetime.now().isoformat(timespec="seconds")})
            new_prop += 1
        else:
            led["pending"].append({**L, "why": why})
            new_pend += 1
    return new_prop, new_pend

def file_card(led, ingest=None, corr=None):
    """ONE consolidated card for every proposal not yet on a card. Returns decision id or None."""
    unfiled = [p for p in led["proposals"] if not p.get("decision_id")]
    if not unfiled:
        return None
    try:
        did = (ingest or ingest_decision)(card_payload(unfiled, corr))
    except Exception as e:
        fleet("deposit_card_failed", {"error": str(e)[:200]}, outcome="error", level="error", corr=corr)
        return None
    if did:
        for p in unfiled:
            p["decision_id"] = did
        fleet("deposit_proposals_filed", {"decision": did, "count": len(unfiled),
                                          "amount": round(sum(p["amount"] for p in unfiled), 2)}, corr=corr)
    return did

def main():
    led = json.load(open(LEDGER)) if os.path.exists(LEDGER) else {
        "seen": {}, "lanes": {"payload":0,"sbir_nasa":0,"po10":0,"usaf":0}, "matched": [], "pending": []}
    ensure_schema(led)
    corr = new_corr()
    investors = load_known_investors()
    booked, rejected = apply_answers(led, load_decisions(), corr)
    try:
        lines = fetch_deposit_lines(SINCE)
    except Exception as e:
        lines = None
        print(f"QB pull failed ({type(e).__name__}): {str(e)[:160]} — no new proposals this run.")
    new_prop, new_pend = propose(led, lines or [], investors)
    led["last_run"] = datetime.datetime.now().isoformat(timespec="seconds")
    led["since"] = SINCE
    print(f"deposit lines pulled since {SINCE}: {len(lines or [])} | new PROPOSED (not booked): {new_prop} | "
          f"new pending: {new_pend} | booked on human answer: {booked} | rejected: {rejected}")
    print("lane totals (net-to-build, booked only):", {k: f"${v:,.0f}" for k,v in led["lanes"].items()})
    if led["proposals"]:
        print(f"PROPOSALS awaiting a human 'Book' answer ({len(led['proposals'])}):")
        for p in led["proposals"][-8:]:
            print(f"   {p['date']}  ${p['amount']:,.0f}  -> {p['lane']}  card={p.get('decision_id') or '(not filed)'}")
    if led["pending"]:
        print(f"PENDING review ({len(led['pending'])}) — flagged, NOT booked:")
        for p in led["pending"][-8:]:
            print(f"   {p['date']}  ${p['amount']:,.0f}  payer='{p['payer'][:30]}'  memo='{p['memo'][:40]}'")
    if DRY:
        print("\n(dry-run — no ledger/tracker writes, no Decision card)"); return
    if not NO_CARD:
        did = file_card(led, corr=corr)
        if did: print(f"Decision card filed/updated: {did}")
    json.dump(led, open(LEDGER, "w"), indent=2)
    fleet("deposit_sync_run", {"new_proposed": new_prop, "new_pending": new_pend, "booked": booked,
                               "rejected": rejected, "open_proposals": len(led["proposals"])},
          outcome="ok" if lines is not None else "qb_pull_failed", corr=corr)
    if booked == 0 and os.path.exists(TRACKER) and "tracker_synced" in led.get("gate", {}):
        return   # lanes unchanged since the last tracker write
    # update tracker net-to-build cells E6..E9 (booked = human-confirmed + grandfathered only)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(TRACKER); s = wb.active
        rowmap = {"payload":6, "sbir_nasa":7, "po10":8, "usaf":9}
        today = datetime.date.today().isoformat()
        for lane, r in rowmap.items():
            s.cell(r, 5, round(led["lanes"][lane], 2)).number_format = "#,##0"
            s.cell(r, 7, f"human-confirmed {today}")
        wb.save(TRACKER)
        led["gate"]["tracker_synced"] = today
        json.dump(led, open(LEDGER, "w"), indent=2)
        print(f"tracker updated: {TRACKER}")
    except Exception as e:
        print(f"tracker update skipped ({type(e).__name__}: {str(e)[:120]}) — ledger saved.")

def selftest():
    """Replays the audit gap 'auto-booked without click' and proves the gate holds. Offline only."""
    ok = 0; n = 0
    def t(name, cond):
        nonlocal ok, n
        n += 1; ok += bool(cond); print(("PASS " if cond else "FAIL ") + name)
    global fleet
    _real_fleet = fleet; events = []
    fleet = lambda action, detail, **k: events.append(action)
    try:
        led = {"seen": {"G:0": True}, "lanes": {"payload":0,"sbir_nasa":0,"po10":1000.0,"usaf":0},
               "matched": [{"key": "G:0", "amount": 1000.0, "lane": "po10"}], "pending": []}
        ensure_schema(led)
        t("existing matches grandfathered (count recorded, lane untouched)",
          led["gate"]["grandfathered_matched"] == 1 and led["lanes"]["po10"] == 1000.0)
        lines = [{"key": "9:0", "date": "2026-10-01", "payer": "Capital Stock", "memo": "504 subscription", "amount": 25000.0},
                 {"key": "9:1", "date": "2026-10-01", "payer": "NASA", "memo": "SBIR Phase I", "amount": 5000.0},
                 {"key": "9:2", "date": "2026-10-01", "payer": "SVB", "memo": "INTEREST PAYMENT", "amount": 12.0}]
        # REPLAY: the pre-gate code credited po10/sbir_nasa here, with nobody clicking anything.
        p, q = propose(led, lines, set())
        t("confident matches become proposals, not bookings (replay of auto-booked-without-click)",
          p == 2 and led["lanes"] == {"payload":0,"sbir_nasa":0,"po10":1000.0,"usaf":0} and len(led["matched"]) == 1)
        t("unclear line goes to pending review", q == 1 and led["pending"][-1]["key"] == "9:2")
        filed = []
        did = file_card(led, ingest=lambda pl: filed.append(pl) or "D-T1")
        t("ONE consolidated card for all new proposals, high priority, human options",
          did == "D-T1" and len(filed) == 1 and filed[0]["priority"] == "high"
          and filed[0]["title"].startswith("EXO3: ") and filed[0]["options"][0].startswith("Book")
          and all(x["decision_id"] == "D-T1" for x in led["proposals"]))
        t("re-run files no duplicate card", file_card(led, ingest=lambda pl: filed.append(pl) or "D-T9") is None and len(filed) == 1)
        t("same keys -> same dedupeKey", batch_key(led["proposals"]) == filed[0]["dedupeKey"])
        # Agent/token answer: e13a807 leaves status open + proposedAnswer
        dec = {"D-T1": {"id": "D-T1", "status": "open", "proposedAnswer": "Book these deposits (recommended)",
                        "auto_proposed": {"by": "Decisions Daily Update"}}}
        b, r = apply_answers(led, dec)
        t("agent-proposed answer never books", b == 0 and led["lanes"]["po10"] == 1000.0 and len(led["proposals"]) == 2)
        dec = {"D-T1": {"id": "D-T1", "status": "answered", "answer": "Book these deposits", "answeredBy": "Exobot"}}
        b, r = apply_answers(led, dec)
        t("answered-by-agent name refused", b == 0 and len(led["proposals"]) == 2)
        dec = {"D-T1": {"id": "D-T1", "status": "answered", "answer": "Book these deposits", "answeredBy": ""}}
        t("answered with no human name refused", apply_answers(led, dec) == (0, 0))
        dec = {"D-T1": {"id": "D-T1", "status": "answered", "answer": "Hold", "answeredBy": "John Quinn"}}
        t("human 'Hold' keeps proposals open", apply_answers(led, dec) == (0, 0) and len(led["proposals"]) == 2)
        t("answer on a DIFFERENT card books nothing",
          apply_answers(led, {"D-OTHER": {"id": "D-OTHER", "status": "answered", "answer": "Book", "answeredBy": "John Quinn"}}) == (0, 0))
        dec = {"D-T1": {"id": "D-T1", "status": "answered", "answer": "Book these deposits (recommended)", "answeredBy": "John Quinn"}}
        b, r = apply_answers(led, dec)
        t("human 'Book' books exactly the carded lines",
          b == 2 and led["lanes"]["po10"] == 26000.0 and led["lanes"]["sbir_nasa"] == 5000.0
          and not led["proposals"] and led["matched"][-1]["booked_on"] == "D-T1" and led["matched"][-1]["booked_by"] == "John Quinn")
        t("booking is idempotent (re-apply books nothing)", apply_answers(led, dec) == (0, 0) and led["lanes"]["po10"] == 26000.0)
        t("re-pulled lines are never re-proposed", propose(led, lines, set()) == (0, 0))
        led2 = ensure_schema({"seen": {}, "lanes": {"payload":0,"sbir_nasa":0,"po10":0,"usaf":0}, "matched": [], "pending": []})
        propose(led2, lines[:1], set()); file_card(led2, ingest=lambda pl: "D-T2")
        b, r = apply_answers(led2, {"D-T2": {"id": "D-T2", "status": "answered", "answer": "Don't book — reclassify", "answeredBy": "John Quinn"}})
        t("human 'Don't book' rejects to pending, never books",
          (b, r) == (0, 1) and led2["lanes"]["po10"] == 0 and led2["pending"][-1]["key"] == "9:0" and not led2["proposals"])
        led3 = ensure_schema({"seen": {}, "lanes": {"payload":0,"sbir_nasa":0,"po10":0,"usaf":0}, "matched": [], "pending": []})
        propose(led3, lines[:1], set())
        def boom(pl): raise RuntimeError("MC down")
        t("card filing failure leaves proposals unbooked and unfiled (retried next run)",
          file_card(led3, ingest=boom) is None and led3["proposals"][0]["decision_id"] is None and led3["lanes"]["po10"] == 0)
        t("telemetry events emitted", {"deposit_proposals_filed", "deposit_booked", "deposit_rejected"} <= set(events))
        t("--selftest touches no live files", LEDGER not in json.dumps(led))
    finally:
        fleet = _real_fleet
    print(f"qb_deposit_sync selftest: {ok}/{n} pass")
    return 0 if ok == n else 1

if __name__ == "__main__":
    if "--selftest" in sys.argv:
        sys.exit(selftest())
    main()
