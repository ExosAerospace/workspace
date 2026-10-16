#!/usr/bin/env python3
"""FAA Docket FAA-2026-9736 — Exos comment, r4 (2026-10-16).
r4: per John 10/16, Japan acceptance of FAA Part 450 is NOT stated as fact (heard secondhand, unverified);
adds the verified Aug 31 2026 FAA-NSPS Memorandum of Cooperation [8] and softens the recognition wording.
Restructured to the 'Public Comment on Federal Regulations' guide (Intro / Background / Analysis /
Recommendations / Conclusion), keyed to the RFI's numbered questions (91 FR 54909), adds the
Hokkaido Spaceport (HOSPO) visit of 13-15 Oct 2026 and an FAA-license international-recognition
recommendation. Also builds the regulations.gov submittal field sheet (comment-box text + form fields).
Run: PYTHONNOUSERSITE=1 python3.12 scripts/build_faa_rfi_9736_comment_r4.py"""
import os
from docx import Document
from docx.shared import Pt, RGBColor

LH = "assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx"
DATE = "2026-10-16"
OUT = f"drafts/Exos_FAA_RFI_9736_Spaceport_Corridor_Comment_{DATE}_r4.docx"
OUT_FIELDS = f"drafts/Exos_FAA_RFI_9736_Regulations_gov_Submittal_Sheet_{DATE}_r2.docx"
NAVY = RGBColor(0x0A, 0x3D, 0x7C)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)

SOURCES = [
    "FAA/DOT, Siting New Spaceports and Use of Priority Airspace for Critical Space Launch Corridors, RFI, 91 FR 54909 "
    "(Aug. 25, 2026): https://www.federalregister.gov/documents/2026/08/25/2026-17290",
    "FAA Office of Commercial Space Transportation, International Affairs: "
    "https://www.faa.gov/space/additional_information/international_affairs",
    "International Bar Association, Overview of the proposed amendments to Japan's Space Activities Act (2026): "
    "https://www.ibanet.org/overview-of-the-proposed-amendments-to-Japans-Space-Activities-Act-and-beyond",
    "Chambers and Partners, Space Law 2026 — Japan, Trends and Developments: "
    "https://practiceguides.chambers.com/practice-guides/space-law-2026/japan/trends-and-developments",
    "Hokkaido Spaceport (HOSPO), official site: https://hokkaidospaceport.com",
    "Nippon.com, Hokkaido Town Aims to Become East Asian Space Launch Hub: https://www.nippon.com/en/japan-topics/g02556",
    "Cabinet Office of Japan, Seventh Space Development and Utilization Award — Taiki Town: "
    "https://www8.cao.go.jp/space/english/award/prize/e02.pdf",
    "FAA and National Space Policy Secretariat, Cabinet Office of Japan, Memorandum of Cooperation on International "
    "Commercial Space Transportation (signed Aug. 31, 2026; non-binding): "
    "https://www.faa.gov/space/additional_information/international_affairs/FAA-NSPS-Memorandum-of-Cooperation.pdf",
]


def mk(doc):
    def para(text, size=10, color=DARK, bold=False, italic=False, after=5, before=0):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(after); pf.space_before = Pt(before)
        r = p.add_run(text)
        r.font.name = "Calibri"; r.font.size = Pt(size)
        r.font.color.rgb = color; r.font.bold = bold; r.font.italic = italic
        return p

    def h(text, size=12):
        para(text, size=size, color=NAVY, bold=True, after=2, before=7)

    def lead(text, body):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(text + " ")
        r.font.name = "Calibri"; r.font.size = Pt(10); r.font.color.rgb = DARK; r.font.bold = True
        r2 = p.add_run(body)
        r2.font.name = "Calibri"; r2.font.size = Pt(10); r2.font.color.rgb = DARK
        return p
    return para, h, lead


# ======================= COMMENT LETTER =======================
doc = Document(LH)
para, h, lead = mk(doc)

para("DRAFT r4 FOR JOHN QUINN REVIEW — NOT YET FILED. Delete this line before upload to regulations.gov.",
     size=8.5, color=GREY, italic=True, after=7)
para("Submitted electronically via www.regulations.gov", bold=True, after=1)
para("Docket No. FAA-2026-9736  |  Federal Aviation Administration, U.S. Department of Transportation", after=1)
para("Date: [insert date on filing — comments due October 26, 2026]", color=GREY, after=6)
para("Re: Request for Information — Siting New Spaceports and Use of Priority Airspace for "
     "Critical Space Launch Corridors (91 FR 54909, August 25, 2026)", bold=True, color=NAVY, size=11, after=6)

h("Responder Information (RFI Section 4)", size=11)
para("Name / affiliation: Exos Aerospace Systems & Technologies, Inc. (“Exos”), a U.S. small business commercial "
     "launch operator", after=1)
para("Address: 299 FM 1903, Greenville, TX 75402", after=1)
para("Website: https://www.exosaero.com", after=1)
para("Point of contact: John Quinn, Cofounder & Chief Executive Officer — jquinn@exosaero.com — (430) 966-4005",
     after=6)

h("1. Introduction — who we are and why we are commenting")
para("Exos is a U.S. commercial launch operator built around reusable, vertical-takeoff suborbital rockets. We hold an "
     "FAA Part 431 launch license and are transitioning to Part 450. We have flown reusable rocket systems in "
     "the Armadillo/SARGE lineage since 2018, and our next vehicle, BLK3, is preparing for first flight in early 2027. "
     "Our vehicle needs only a launch pad, power, and communications, so our view of “a spaceport that actually gets "
     "used” comes from the operator side of the fence. We are working on spaceport siting today in West Texas (the "
     "Concho Valley and the Midland region). From October 13 to 15, 2026, our CEO visited Hokkaido Spaceport (HOSPO) "
     "in Taiki, Japan, to evaluate it as an allied host site. That visit gave us a direct comparison of how a "
     "foreign host nation is building, funding, and regulating a commercial spaceport. We offer these comments as "
     "policy input. We understand this RFI is not a Screening Information Request or a Request for Offers, and "
     "nothing in this submission is an offer.")

h("2. Background — RFI questions addressed")
para("We respond to Question 1 (siting criteria, including island sites and international equities), Question 2 "
     "(why existing commercial spaceports are underutilized), Question 4 (new locations, costs, and public-private "
     "funding), Question 5 (critical space launch corridors), and Question 6 (questions not asked). Our principal "
     "Question 6 item is international recognition of FAA licensing, which affects how U.S. operators can use "
     "allied host spaceports.")

h("3. Analysis and evidence")
lead("Q1 — Operator siting criteria.",
     "From an operator’s standpoint, a site is genuinely usable when it has low-population overflight and open, "
     "uncongested airspace, distributed-range geometry that contains hazard areas, and a low cost of activation. A "
     "reusable suborbital operator needs a pad, power, and communications, not an orbital-scale complex, to begin "
     "flying. Siting requirements built only around orbital-class users screen out the operators who can generate "
     "flight activity at a new site first.")
lead("Q1 — Island and coastal sites; the HOSPO comparison.",
     "HOSPO shows that a coastal site with open water on the launch azimuths can attract foreign operators quickly. "
     "It faces the open Pacific to the east and south, supporting eastward LEO and southward polar launches [5][6]. "
     "It hosted a launch by a Taiwanese company, jtSPACE, in July 2025. The rocket was shipped by sea and trucked to "
     "the site, and the operator chose HOSPO over an Australian alternative on transport distance, time, and cost "
     "[6]. The lesson for U.S. siting is that port access and logistics cost can decide a site’s competitiveness as "
     "much as launch geometry does.")
lead("Q1 — International equities.",
     "U.S. launch operators who fly from a foreign host site must hold an FAA license and may also need the host "
     "country’s license. The FAA itself identifies this as a possible dual-licensing burden, and notes that U.S. law "
     "and FAA regulations can serve as a framework for other countries [2]. Japan amended its Space Activities Act in "
     "June 2026 to cover rocket-only launches, test flights, and reuse. Each launch from Japan still requires a "
     "permit from the Prime Minister, and a framework for suborbital flight is reserved for a future amendment "
     "[3][4]. On August 31, 2026, the FAA and the National Space Policy Secretariat of Japan’s Cabinet Office signed a "
     "non-binding Memorandum of Cooperation. It covers information exchange and exploring commonality in safety "
     "standards and compliance evaluation for launch, reentry, and suborbital flight [8]. As published, neither the "
     "amended Act nor the Memorandum provides for Japan to accept an FAA license in place of its own permit. Allied "
     "regimes are still being written, which makes this a good time for the United States to offer Part 450 as a "
     "reference those regimes can build on.")
lead("Q2 — Why licensed spaceports are underutilized.",
     "In our experience, site requirements and cost structures at many spaceports were shaped around heavyweight, "
     "expendable orbital operations. Those requirements over-burden the reusable and suborbital operators best "
     "positioned to generate regular cadence. Reusable suborbital flights are lower-energy, contain their hazard "
     "areas, and repeat. A new or dormant site can use them to build an operating record and community confidence "
     "before graduating to orbital-class activity. The West Texas locations we are evaluating are concrete candidates "
     "for this anchor-tenant approach.")
lead("Q2 — A campaign-based cadence model.",
     "The most direct way to turn a licensed site into a used site is to fly it often. Exos has developed, and "
     "discussed with NASA, a flight-campaign concept in which one mobilization to a site supports multiple flights "
     "spaced days apart within a quarter. As the model matures, we expect campaign cadence to approach monthly and "
     "to extend to qualified allied host locations, including Japan and Italy. Policy built around recurring "
     "campaigns, not one-off launches, is what makes new and dormant spaceports economically viable.")
lead("Q4 — Public-private funding and workforce.",
     "HOSPO is a public-private model. Taiki Town and its operator, SPACE COTAN, have developed the site with "
     "national support, and SPACE COTAN was selected for Japan’s Space Strategy Fund (¥10.5 billion) to raise launch "
     "frequency through multi-vehicle handling, weather forecasting, communications, and cryogenic capability [6][7]. "
     "That is the kind of targeted, cadence-focused public investment the Department’s National Infrastructure "
     "Development Office (NIDO) could structure for U.S. sites. We encourage weighting partnership models that pair a "
     "spaceport with state, university, and institutional participation, including workforce training, flight-test "
     "services, and vehicle wet-lease capability, so that a site builds a durable local base instead of depending on "
     "a single anchor launch company.")
lead("Q5 — Requisites of a critical launch corridor.",
     "As the RFI notes, the most valuable attribute of a corridor for aviation is predictability: advance notice of "
     "the time, size, and duration of the Aircraft Hazard Area [1]. Reusable suborbital trajectories are short and "
     "geographically contained. Their airspace footprint is substantially smaller than a long-duration orbital "
     "ascent, which makes them a model for corridor efficiency rather than a burden on the system.")
lead("Q5 — Priority, activation, fees, and emergencies.",
     "Any “priority” treatment should follow a flight’s actual hazard window and cadence, not the operator’s size. A "
     "high-tempo reusable operator needs repeatable corridor activation, not a weeks-ahead coordination cycle for "
     "every flight of a previously flown profile. Fee or restriction structures that penalize high-cadence small "
     "operators would price out the reusable model national policy seeks to grow. Contingency procedures should take "
     "advantage of the fast, contained abort characteristics of suborbital reusable vehicles.")
lead("Costs, benefits, and the small-business perspective.",
     "The costs of an orbital-class requirement set fall hardest on small reusable operators, while the benefits "
     "of their cadence (site revenue, workforce, community acceptance) go to the spaceport and the region. A "
     "graduated framework aligns who pays with who benefits.")

h("4. Recommendations")
recs = [
    "Adopt a tiered or graduated spaceport framework under 14 CFR part 420 so that lightweight reusable suborbital "
    "operations can activate a site without bearing orbital-class infrastructure obligations (Q1, Q2, Q4).",
    "Treat reusable suborbital campaigns as a recognized anchor-tenant path for new and dormant spaceports, including "
    "the West Texas candidates (Q2, Q4).",
    "Weight port and wide-load logistics access in siting evaluations alongside launch geometry, as the HOSPO "
    "experience shows (Q1, Q4).",
    "Use NIDO public-private partnership structures for targeted cadence investments such as weather, "
    "communications, and multi-vehicle pad capability, modeled on allied programs such as Japan’s Space Strategy Fund "
    "(Q4).",
    "Provide cadence-based, pre-coordinated corridor activation for licensed reusable operations, with compressed "
    "timelines for previously demonstrated trajectories, and base priority on hazard window and cadence, not "
    "operator size (Q5).",
    "Avoid corridor fee and restriction structures that disproportionately burden high-cadence small operators "
    "(Q5).",
    "(Q6) Pursue bilateral regulatory-recognition arrangements with allied launch states. Japan is a natural "
    "starting point: the August 31, 2026 Memorandum of Cooperation already provides a channel [8], and Japan’s "
    "suborbital framework is still being drafted. The goal would be for a host nation to be able to rely on FAA "
    "Part 450 safety determinations as a basis for its own permit. This reduces dual licensing for U.S. operators, extends "
    "U.S. regulatory leadership, and turns allied spaceports into an extension of U.S. launch capacity.",
    "(Q6) Add this question to the Department’s analysis: how can FAA licensing and the Office of Spaceports’ "
    "international outreach support U.S. operators at allied host spaceports without duplicative review?",
]
for i, r in enumerate(recs, 1):
    para(f"R{i}. {r}", after=3)

h("5. Conclusion")
para("Spaceports become operational when operators can fly from them often and affordably. Reusable suborbital "
     "operators are the fastest way to give a new or dormant site an operating record, and the United States can "
     "extend that advantage abroad by offering Part 450 as a reference allied host nations can build on. In summary, we "
     "recommend a graduated siting framework (R1–R3), cadence-focused public-private investment (R4), "
     "cadence-based corridor activation and fair fees (R5–R6), and international recognition of FAA licensing "
     "(R7–R8). Exos would welcome the opportunity to provide further detail on any of these points.")
para("Respectfully submitted,", after=10, before=5)
para("John Quinn — Cofounder & Chief Executive Officer", bold=True, after=1)
para("Exos Aerospace Systems & Technologies, Inc.  |  299 FM 1903, Greenville, TX 75402  |  jquinn@exosaero.com",
     color=GREY, after=6)

h("Sources", size=10)
for i, s in enumerate(SOURCES, 1):
    para(f"[{i}] {s}", size=8.5, color=GREY, after=1)

os.makedirs("drafts", exist_ok=True)
doc.save(OUT)
print("WROTE", OUT)

# ======================= REGULATIONS.GOV SUBMITTAL SHEET =======================
COMMENT_BOX = (
    "Exos Aerospace Systems & Technologies, Inc. (Greenville, TX; www.exosaero.com) submits the attached comments on "
    "Docket FAA-2026-9736, Siting New Spaceports and Use of Priority Airspace for Critical Space Launch Corridors "
    "(91 FR 54909). Exos is a U.S. small-business operator of reusable, vertical-takeoff suborbital rockets. We hold "
    "an FAA Part 431 launch license and are transitioning to Part 450. We are working on spaceport siting in West "
    "Texas, and we visited Hokkaido Spaceport (HOSPO), Japan, as an allied host site in October 2026.\n\n"
    "Summary of recommendations (full analysis and sources in the attachment):\n"
    "R1. Adopt a graduated spaceport framework under 14 CFR part 420 so lightweight reusable suborbital operations can "
    "activate a site without orbital-class infrastructure obligations (Q1, Q2, Q4).\n"
    "R2. Recognize reusable suborbital campaigns as an anchor-tenant path for new and dormant spaceports (Q2, Q4).\n"
    "R3. Weight port and wide-load logistics access in siting alongside launch geometry (Q1, Q4).\n"
    "R4. Use NIDO public-private partnerships for targeted cadence investments such as weather, communications, and "
    "multi-vehicle pads, modeled on allied programs such as Japan's Space Strategy Fund (Q4).\n"
    "R5. Provide cadence-based, pre-coordinated corridor activation for licensed reusable operations, with compressed "
    "timelines for previously demonstrated trajectories; base priority on hazard window and cadence, not operator "
    "size (Q5).\n"
    "R6. Avoid corridor fees and restrictions that disproportionately burden high-cadence small operators (Q5).\n"
    "R7. (Q6) Pursue bilateral regulatory-recognition arrangements with allied launch states, starting with Japan "
    "under the August 31, 2026 FAA-Cabinet Office Memorandum of Cooperation while Japan's suborbital framework is "
    "being drafted, so host nations could rely on FAA Part 450 safety determinations as a basis for their own permits. This reduces dual licensing and extends U.S. regulatory leadership.\n"
    "R8. (Q6) Analyze how FAA licensing and Office of Spaceports outreach can support U.S. operators at allied host "
    "spaceports without duplicative review.\n\n"
    "Point of contact: John Quinn, Cofounder & CEO, jquinn@exosaero.com."
)

fd = Document()
p2, h2, _ = mk(fd)
p2("Exos — regulations.gov submittal sheet, Docket FAA-2026-9736 (INTERNAL, do not upload)", size=13, bold=True,
   color=NAVY, after=4)
p2("Comment page: https://www.regulations.gov/docket/FAA-2026-9736 → open the RFI notice → “Comment”. "
   "Comments due October 26, 2026 (91 FR 54909).", after=8)
rows = [
    ("Comment* (text box)", "Paste the text in Section B below. "
     f"Length: {len(COMMENT_BOX):,} characters. The full comment is the attachment."),
    ("What is your comment about? (if shown)", "Select the closest category offered, e.g., Private Industry / "
     "Business. Some dockets do not show this field."),
    ("Attach Files", f"{os.path.basename(OUT).replace('.docx', '.pdf')} (export the approved Word file to PDF, "
     "remove the DRAFT line first). Limit: up to 20 files, 10 MB each; PDF and DOCX are accepted."),
    ("Email Address", "jquinn@exosaero.com — tick “receive email confirmation and tracking number”. "
     "The email is not posted publicly."),
    ("Tell us about yourself", "An Organization"),
    ("Organization Name", "Exos Aerospace Systems & Technologies, Inc."),
    ("Public-disclosure check", "Everything in the box and attachment is posted publicly. Use the company line "
     "(430) 966-4005, not a personal mobile. No proprietary, export-controlled, or ITAR technical data is included."),
    ("After submitting", "Save the tracking number and attach the confirmation to MC #2452."),
]
t = fd.add_table(rows=1, cols=2)
t.style = "Table Grid"
t.rows[0].cells[0].text = "regulations.gov field"
t.rows[0].cells[1].text = "What to enter"
for a, b in rows:
    c = t.add_row().cells
    c[0].text = a; c[1].text = b
h2("Section B — text for the Comment box")
for chunk in COMMENT_BOX.split("\n"):
    if chunk.strip():
        p2(chunk, after=3)
h2("Section C — how the comment meets the Public-Comment guide")
for line in [
    "Introduction with distinguishing experience → Section 1 (operator, Part 431→450, West Texas siting, HOSPO visit).",
    "Background identifying what we comment on → Section 2 maps to RFI Questions 1, 2, 4, 5, 6.",
    "Analysis with evidence and citations → Section 3, with numbered public sources [1]–[8].",
    "Recommendations with specific changes → Section 4, R1–R8, each tagged to an RFI question.",
    "Conclusion recapping recommendations → Section 5.",
    "Missing perspective (small business) and costs vs. benefits → Section 3, last paragraph.",
]:
    p2("• " + line, after=2)
fd.save(OUT_FIELDS)
print("WROTE", OUT_FIELDS, "comment-box chars:", len(COMMENT_BOX))
