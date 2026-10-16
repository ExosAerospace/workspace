#!/usr/bin/env python3
"""Convert an Exo Word briefing (python-docx built: bold 16pt title, bold 12pt section heads, List Bullet, tables)
into an editable 16:9 Exos-branded PowerPoint. Sections become slides, auto-paginated; tables become slide tables.
  docx_briefing_to_pptx.py in.docx out.pptx"""
import sys, math, os
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

WS="/Users/exomac/.openclaw/workspace"
BGIMG=os.path.join(WS,"drafts","_exos_slide_bg.png"); LOGO=os.path.join(WS,"drafts","_exos_white_official.png")
BG=RGBColor(0x0d,0x0d,0x0f); PANEL=RGBColor(0x17,0x18,0x1c); PANEL2=RGBColor(0x1f,0x20,0x25); ORANGE=RGBColor(0xF2,0x6A,0x21)
WHITE=RGBColor(0xFF,0xFF,0xFF); GREY=RGBColor(0xD2,0xD2,0xD7); MUTE=RGBColor(0x9a,0x9a,0xa1)

def items(doc):
    out=[]
    for el in doc.element.body.iterchildren():
        tag=el.tag.split('}')[1]
        if tag=="tbl":
            t=Table(el,doc); out.append(("table",[[c.text.strip() for c in r.cells] for r in t.rows])); continue
        if tag!="p": continue
        p=Paragraph(el,doc); txt=p.text.strip()
        if not txt: continue
        r=p.runs[0] if p.runs else None; sz=r.font.size.pt if r is not None and r.font.size else None
        if sz==16 and r.bold: out.append(("title",txt))
        elif sz==12 and r.bold: out.append(("head",txt))
        elif p.style.name.startswith("List"):
            lead=r.text.strip() if (r is not None and r.bold and len(p.runs)>1) else ""
            out.append(("bullet",(lead,txt[len(lead):].strip() if lead else txt)))
        else: out.append(("para",(txt, bool(r is not None and r.italic))))
    return out

def convert(src,dst):
    doc=Document(src); its=items(doc)
    p=Presentation(); p.slide_width=I(13.333); p.slide_height=I(7.5); BL=p.slide_layouts[6]; n=[0]
    title=next((t for k,t in its if k=="title"),os.path.basename(src))
    def tb(s,x,y,w,h,size,color,bold=False):
        t=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=t.text_frame; tf.word_wrap=True; return tf
    def run(par,text,size,color,bold=False,italic=False):
        r=par.add_run(); r.text=text; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; r.font.italic=italic; r.font.name="Calibri"
    def base(head,cont=False):
        s=p.slides.add_slide(BL); n[0]+=1
        s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
        b=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,I(0.12),I(7.5)); b.fill.solid(); b.fill.fore_color.rgb=ORANGE; b.line.fill.background()
        tf=tb(s,0.6,0.45,10.5,1.0,26,WHITE); tf.paragraphs[0].text=""; run(tf.paragraphs[0],head+("  (cont.)" if cont else ""),26,WHITE,True)
        s.shapes.add_picture(LOGO,I(11.35),I(0.48),height=I(0.42))
        ft=tb(s,0.6,7.02,11,0.3,9,MUTE); run(ft.paragraphs[0],title.upper(),9,MUTE)
        nt=tb(s,12.2,7.02,0.6,0.3,9,MUTE); nt.paragraphs[0].alignment=PP_ALIGN.RIGHT; run(nt.paragraphs[0],f"{n[0]:02d}",9,MUTE)
        return s
    # title slide
    s=p.slides.add_slide(BL); n[0]+=1
    s.shapes.add_picture(BGIMG,0,0,I(13.333),I(7.5)); s.shapes.add_picture(LOGO,I(0.8),I(0.7),height=I(0.7))
    tf=tb(s,0.8,2.5,11.6,1.8,36,WHITE); run(tf.paragraphs[0],title,36,WHITE,True)
    sub=next((t[0] for k,t in its if k=="para"),"")
    tf=tb(s,0.8,4.5,11.6,1.5,16,GREY); run(tf.paragraphs[0],sub,16,GREY)
    # sections
    secs=[]; cur=None
    for k,v in its:
        if k=="title": continue
        if k=="head": cur=[v,[]]; secs.append(cur); continue
        if cur is None: continue   # subtitle para already on title slide
        cur[1].append((k,v))
    BODY_H=5.35; CPL=135  # chars/line at 15pt over 12.1in
    def est(k,v,size=17.5):
        txt=(v[0]+" "+v[1]) if k=="bullet" else v[0]
        lines=math.ceil(len(txt)/(CPL*15/size)); return lines*size*1.32/72+size*0.5/72
    for head,body in secs:
        text_blocks=[]; 
        for k,v in body:
            if k=="table": text_blocks.append(("table",v))
            else: text_blocks.append((k,v))
        page=[]; h=0; first=True
        def flush_text(page,first):
            s=base(head,not first)
            box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,I(0.6),I(1.45),I(12.15),I(BODY_H+0.1)); box.adjustments[0]=0.03
            box.fill.solid(); box.fill.fore_color.rgb=PANEL; box.line.fill.background(); box.shadow.inherit=False
            tf=tb(s,0.85,1.65,11.7,BODY_H-0.2,17.5,GREY)
            for i,(k,v) in enumerate(page):
                par=tf.paragraphs[0] if i==0 else tf.add_paragraph(); par.space_after=Pt(10)
                if k=="bullet":
                    run(par,"▸  ",17.5,ORANGE,True)
                    if v[0]: run(par,v[0]+" ",17.5,WHITE,True)
                    run(par,v[1],17.5,GREY)
                else: run(par,v[0],17.5,MUTE if v[1] else GREY,False,v[1])
        for k,v in text_blocks:
            if k=="table":
                if page: flush_text(page,first); first=False; page=[]; h=0
                hdr,rows=v[0],v[1:]; ncol=len(hdr)
                lens=[max(len(r[c]) for r in v) for c in range(ncol)]
                W=12.15; fixed={0:0.45}
                if hdr[-1].lower().startswith('owner'): fixed[ncol-1]=1.45
                rest=W-sum(fixed.values()); fl=sum(lens[c] for c in range(ncol) if c not in fixed)
                ws=[fixed.get(c, rest*lens[c]/fl) for c in range(ncol)]
                def rh(r): return max(math.ceil(len(r[c])/max(1,(ws[c]*72/12.5*2.0)))*0.19 for c in range(ncol))+0.15
                chunks=[]; ch=[]; hh=0.45
                for r in rows:
                    if ch and hh+rh(r)>BODY_H: chunks.append(ch); ch=[]; hh=0.45
                    ch.append(r); hh+=rh(r)
                if ch: chunks.append(ch)
                for ch in chunks:
                    s=base(head,not first); first=False
                    gt=s.shapes.add_table(len(ch)+1,ncol,I(0.6),I(1.45),I(W),I(0.4)).table
                    for c in range(ncol): gt.columns[c].width=I(ws[c])
                    for ri,r in enumerate([hdr]+ch):
                        for c in range(ncol):
                            cell=gt.cell(ri,c); cell.fill.solid(); cell.fill.fore_color.rgb=ORANGE if ri==0 else (PANEL if ri%2 else PANEL2)
                            cell.margin_left=cell.margin_right=I(0.08); cell.margin_top=cell.margin_bottom=I(0.05)
                            tf=cell.text_frame; tf.word_wrap=True; tf.paragraphs[0].text=""
                            run(tf.paragraphs[0],r[c],13 if ri==0 else 12.5,WHITE if (ri==0 or c in(0,1)) else GREY,ri==0 or c==1)
                continue
            e=est(k,v)
            if page and h+e>BODY_H: flush_text(page,first); first=False; page=[]; h=0
            page.append((k,v)); h+=e
        if page: flush_text(page,first)
    p.save(dst); print("WROTE",dst,n[0],"slides")

if __name__=="__main__": convert(sys.argv[1],sys.argv[2])
