#!/usr/bin/env python3
"""Build the UTD x Exos Capstone MOU on the PRIMARY letterhead.
Content rev. 2026-08-21 — amended per UTD kickoff call (no in-course monetary
mechanics; research track added; grading de-quota'd; 16 weeks; Aug 28 start).
Run: PYTHONNOUSERSITE=1 python3.12 scripts/build_utd_mou_docx.py"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

LH = "assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx"
OUT = "drafts/UTD_Exos_Capstone_MOU_2026-08-21.docx"
NAVY = RGBColor(0x0A, 0x3D, 0x7C); DARK = RGBColor(0x1A, 0x1A, 0x1A)
doc = Document(LH)

def run(p, t, size=11, bold=False, color=DARK, italic=False):
    r = p.add_run(t); r.font.name = "Calibri"; r.font.size = Pt(size)
    r.bold = bold; r.italic = italic; r.font.color.rgb = color; return r

def title(t, size=15):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1); run(p, t, size, True, NAVY); return p
def sub(t):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8); run(p, t, 12, True, DARK); return p
def h(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(2); run(p, t, 12, True, NAVY); return p
def body(segs, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    for s, b in segs: run(p, s, 11, b, DARK)
    return p
def bullet(segs):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.28); p.paragraph_format.space_after = Pt(2)
    run(p, "•  ", 11, False, DARK)
    for s, b in (segs if isinstance(segs, list) else [(segs, False)]): run(p, s, 11, b, DARK)
    return p

title("Memorandum of Understanding")
sub("Exos × UT Dallas — “EXO3 Payload Capstone” (Fall 2026)")
body([("Parties: ", True), ("Exos Aerospace Systems & Technologies, Inc. (“Exos”) and The University of Texas at Dallas (“UTD”).", False)], after=1)
body([("Exos lead: ", True), ("Lizi Quinn (relationship owner + public-content approver); John Quinn, Cofounder & CEO (sponsor, supportive).", False)], after=1)
body([("UTD lead: ", True), ("Edwin Van Der Vlist (capstone program coordinator), with the assigned section instructor of record.", False)], after=1)
body([("Status: ", True), ("DRAFT for review — route to counsel before signature.", False)], after=8)

h("1. Purpose")
body([("UTD provides a Fall 2026 marketing capstone team to act as Exos’s agency of record, marketing and filling the open ~135 kg on the BLK3 Flight 2 payload bay (early 2027, Spaceport America, NM; launch est. April 2027) under the SPACEavailable…™ campaign, directing Exos’s EXO3 AI agent fleet. The engagement is a graded academic project combining applied marketing execution with primary and secondary research; students are graded on the quality of their strategy, research, and process — not on sales closed.", False)])

h("2. Roles")
bullet([("Lizi Quinn (Exos) ", True), ("— owns the UTD relationship and is the approval gate for anything that goes public (media, website, outreach naming Exos).", False)])
bullet([("John Quinn (Exos) ", True), ("— sponsor; supportive; engaged only for critical-impact decisions.", False)])
bullet([("Exo (EXO3 agent) ", True), ("— the operating engine; gates student design/build requests, generates code/creative in an isolated sandbox, and runs the progress tracker.", False)])
bullet([("Edwin Van Der Vlist / section instructor (UTD) ", True), ("— coordinate the program and assign the section instructor of record, who forms teams, uploads the roster, sits in on client meetings, and sets grades; parallel rights with Jon + Lizi to see/manage student accounts; joint determiner (with Exos) of material contribution for non-cash recognition eligibility only.", False)])

h("3. Student access & security")
bullet([("Each student receives a scoped Mission Control login limited to the “Payload Sales” board + the progress tracker. No access to Exos CRM, financials, decisions, documents, or internal systems.", False)])
bullet([("Access is monitored: any attempt to exceed scope is logged and alerts Exos + Lizi.", False)])
bullet([("Restricted from students: Narravity and Napier Park (walled off).", False)])
bullet([("Design-request loop: students request builds; Exo generates in an isolated capstone sandbox (never Exos production); changes are reviewed and reversible; nothing executes on Exos systems.", False)])

h("4. Timeline (16 weeks; first class Aug 28, 2026; instructor sets team formation)")
bullet([("Weeks 1–3 (research ramp): ", True), ("onboard without overload; students research the ideal customer profile, Exos’s value proposition and market positioning, and the payload opportunity before any outreach — a deliberate phase-in.", False)])
bullet([("Weeks 1–2: ", True), ("team sets its own space-themed name, defines ICP/segments; instructor-required client kickoff meeting.", False)])
bullet([("Weeks 2–3: ", True), ("Exos site visit (Greenville, TX, 3–4 hrs — Jon/Lizi + BD team address the team + facility tour). Subject to UTD’s >30-mile off-campus authorization (Greenville ≈ 50 mi); the instructor secures approval and Exos coordinates dates.", False)])
bullet([("Week 3: ", True), ("Challenge kickoff — AI-only vs Hybrid both live (framed as a learning ramp, not a sales quota).", False)])
bullet([("Weeks 4–16: ", True), ("campaign build, human-in-the-loop outreach validation, weekly research + pipeline gates, handoffs, final presentation + reusable GTM playbook. A live weekly-gate tracker shows progress to the team and the instructor.", False)])

h("5. The Challenge")
body([("A controlled A/B on matched prospect cohorts: AI-only (fleet autonomous) vs Hybrid (students directing the fleet), same tools and approval gate, metrics pre-registered. It tests whether human-directed AI beats autonomous AI on pipeline quality and process outcomes. This is a research exercise; its results are not a grading quota.", False)])

h("6. Deliverables, research & grading")
body([("Required research (every project carries both):", True)], after=2)
bullet([("Primary research: ", True), ("exercise Exos’s live outreach pipeline (~86–89 prepared items), validate the AI fleet’s fit scoring (~85% confidence), and deliver independent feedback on the process — what works, what to improve.", False)])
bullet([("Secondary research: ", True), ("market landscape for suborbital payload demand, plus an independent value-proposition / positioning analysis for Exos.", False)])
bullet([("Team indicators (informational, not the grade): ", True), ("pipeline participation, research quality, campaign assets shipped, and AI-effectiveness (leverage, reply rate, attribution). Bookings are tracked as program interest but are not a graded target.", False)])
bullet([("Individual score (role-fair, /100): ", True), ("primary-track impact 35, contribution-to-team-outcome 20, AI-leverage 15, consistency 15, instructor+peer assessment 15. Five equal tracks — Growth/Outreach, Web & Pipeline Engineering, Creative/Content, Data/Analytics & Ops, and Research & Positioning — each scored in its own currency; specializing to any one track is a full, unpenalized path. Grades are set solely by the UTD instructor.", False)])

h("7. Individual accountability")
body([("It is a team effort, and each person is expected to materially contribute. Where Exos exercises its discretion to offer any non-cash token of appreciation, material contribution — individual score ≥ 60 in the student’s track, affirmed by both the instructor and Exos — informs eligibility.", False)])

h("8. Recognition")
body([("No funds, cards, per-diems, or monetary instruments are issued to students at any point during the course, and no cash is paid for sales or project accomplishments. Exos may, at its sole discretion and consistent with UTD policy, provide small non-cash tokens of appreciation. This MOU promises no reward, prize, trip, internship, or employment.", False)])

h("9. Exos contribution")
body([("Substantial in-kind value: the EXO3 platform + agent-fleet access, Exo’s build time, brand assets and media, and CEO + BD team time (including an on-site facility visit and weekly check-ins).", False)])

h("10. Post-course opportunities")
body([("After the graded project concludes, Exos may — at its sole discretion and entirely independent of the course — consider a student for an internship or role. Nothing in this MOU is a promise or commitment of internship or employment, and no such opportunity is compensation for coursework.", False)])

h("11. Intellectual property")
body([("Exos owns all capstone deliverables (site, copy, campaign assets, pipeline, research reports). UTD confirms its capstone terms do not conflict.", False)])

h("12. Media")
body([("UTD is responsible for local and national media and audience-building around the “AI challenged by UTD” story. Any media naming Exos or its capabilities passes Exos (Lizi) review first — positioning rules apply (flying since 2018; one of only three FAA-licensed reusable rocket launch providers in the U.S.; never “first commercial flight”; no overstatement; export-safe).", False)])

h("13. Sources & attribution")
body([("Primary source: the Exos BLK3 Payload User’s Guide. Context: historical Exos + Armadillo Aerospace video. Any material used from Armadillo Aerospace must credit “Armadillo Aerospace” (Exos is Armadillo’s commercial successor).", False)])

h("14. Compliance")
bullet([("Student data / FERPA: ", True), ("minimal collection; the UTD instructor uploads from the UTD side.", False)])
bullet([("Export diligence: ", True), ("the roster captures US-citizen / non-US status as a documentation step, not a disqualifier — the student scope is public/commercial marketing data only (no ITAR/CUI); non-US students participate fully.", False)])
bullet([("No monetary payments to students: ", True), ("the program issues no cards, per-diems, wages, commissions, or cash to students during the course; recognition is non-cash and hosted by Exos, consistent with UTD’s state-institution obligations.", False)])
bullet([("Securities: ", True), ("any public investor/offering language is counsel-gated, must fit the applicable Reg D exemption’s solicitation rules, and is authored by Exos + counsel only — never by students.", False)])
bullet([("Hosting: ", True), ("the SPACEavailable microsite is served at spaceavailable.exosaero.com (static, reviewed-before-deploy).", False)])

h("15. Term & signatures")
body([("Effective on signature; runs the Fall 2026 semester. Either party may terminate for cause with written notice; student data handled per FERPA on termination.", False)], after=14)

sg = doc.add_paragraph(); run(sg, "Exos Aerospace Systems & Technologies, Inc.", 11, True, DARK)
p = doc.add_paragraph(); run(p, "By: ____________________________   Name: John Quinn   Title: Cofounder & CEO   Date: __________", 11)
doc.add_paragraph()
p = doc.add_paragraph(); run(p, "The University of Texas at Dallas", 11, True, DARK)
p = doc.add_paragraph(); run(p, "By: ____________________________   Name: ____________________   Title: __________   Date: __________", 11)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
run(p, "DRAFT — route to counsel (IP, student data, media, launch-site liability, non-cash-recognition treatment) before signature.", 9.5, False, DARK, italic=True)

doc.save(OUT); print("saved", OUT)
