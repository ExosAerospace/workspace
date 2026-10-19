#!/usr/bin/env python3
"""10/19 SSLC email drafts: (1) Paula Lessnick (BBSNV) — action Motions 2–4; (2) SSLC Board — information + motions for 10/20."""
from docx import Document
from docx.shared import Pt, Inches

OUT_PAULA = "drafts/SSLC_Email_to_PLessnick_Motions2-4_2026-10-19.docx"
OUT_BOARD = "drafts/SSLC_Email_to_Board_Meeting_Motions_2026-10-19.docx"


def new_doc():
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1)
    d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)
    return d


def para(d, t, *, bold=False, italic=False, size=11, after=8):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)


def header(d, rows):
    para(d, "DRAFT — NOT SENT · for John Quinn’s review", italic=True, size=9)
    for k, v in rows:
        p = d.add_paragraph(); a = p.add_run(k + "  "); a.bold = True; p.add_run(v); p.paragraph_format.space_after = Pt(2)
    para(d, "—" * 40, after=8)


# ---------------------------------------------------------------- 1. Paula Lessnick
d = new_doc()
header(d, [
    ("From:", "John Quinn <jquinn@exosaero.com>"),
    ("To:", "Paula Lessnick <paula.lessnick@bbsnv.com>"),
    ("Cc:", "Markus Rufer <mrufer@smad.com>"),
    ("Subject:", "SSLC Board Motions 2–4 (Oct 13): deposit and progress-billing methodology for SSLC"),
])
for t in [
    "Paula,",
    "Thank you for offering to help. At its October 13 meeting the Scorpius Space Launch Company (SSLC) Board approved three motions that touch "
    "your work with us. Here is what I am asking of you on each, and what is not needed.",
    "Motion 3: correct the AR Unbilled and customer-deposit accounting (due Tuesday, October 20). This is the main ask. SSLC bills about 40% "
    "at the start of each tank, and those billings are not being closed against the final invoice when the tank ships, so its “billings in "
    "excess of revenue” balance keeps growing. The Board agreed that SSLC’s bookkeeper, A La Carte Solutions, should keep the books, but that it "
    "would help to give them a worked example of how a company in the same business books this correctly. Please send me, by end of day "
    "Monday, October 19 if you can:",
    "1. A one-to-two page write-up of how BBSNV books Exos’s deposits and progress billings: the account a deposit goes to when received, how it "
    "moves as work is performed, and the entry when the item is delivered or the flight is flown.\n"
    "2. Sample journal entries for one tank-style order: deposit at start, progress billing, final invoice and delivery.\n"
    "3. A short note on when revenue is recognized: at delivery, or over time as work is done if the contract gives a right to payment for work "
    "performed. Deposits received before the work is done are a liability until then, under ASC 606 [1][2].",
    "Please keep it generic and based on Exos’s own practice. I will pass it to Markus Rufer, SSLC’s CEO, for A La Carte to apply. If a "
    "30-minute call with Markus and A La Carte would help, I would appreciate it; please suggest two times on Tuesday or Wednesday.",
    "Motion 2: read-only access to SSLC’s books for its directors. Done, no action for you. Access was established through Exos Mission "
    "Control on October 19, and Markus and Bob have their login instructions. It is for SSLC directors only, and I will not share SSLC’s "
    "books with BBSNV unless the SSLC Board authorizes it.",
    "Motion 4: ERP/MRP evaluation. Exos is evaluating an ERP/MRP system for SSLC at no cost to SSLC, on the same Mission Control platform. When we have a shortlist, I would like "
    "15 minutes of your view on whether its work-in-progress and percent-complete tracking will feed the revenue entries in item 3. Nothing is "
    "needed this week.",
    "One more request: two of SSLC’s directors, Markus Rufer and Robert (Bob) Conger, may ask for a short call before Tuesday to "
    "understand the booking issues, including the April 30, 2026 related-party entries. Could you make 30 to 45 minutes available on "
    "Monday or Tuesday morning (Pacific)? I will send you the relevant ledger extracts first, and I will tell them you are Exos’s "
    "accountant, so your view is a second opinion rather than an independent review.",
    "BBSNV’s time on this is for Exos’s account unless the SSLC Board agrees otherwise. SSLC information you receive is confidential.",
    "Thank you,\nJohn",
    "John Quinn\nCEO, Exos Aerospace Systems & Technologies, Inc.\nDirector, Scorpius Space Launch Company",
]:
    para(d, t)
para(d, "Sources for John’s reference (not part of the email):", bold=True, size=9, after=2)
for s in ["[1] Deloitte DART, Revenue Roadmap 8.4, Revenue recognized over time (ASC 606-10-25-27): "
          "https://dart.deloitte.com/USDART/home/codification/revenue/asc606-10/roadmap-revenue-recognition/chapter-8-step-5-determine-when/8-4-revenue-recognized-over-time",
          "[2] Deloitte DART, Revenue Roadmap 1.8, Step 5: recognize revenue when control transfers: "
          "https://dart.deloitte.com/USDART/home/codification/revenue/asc606-10/roadmap-revenue-recognition/chapter-1-overview/1-8-step-5-recognize-revenue",
          "Motions 2–4 as approved: SSLC Board meeting transcript, October 13, 2026 (data-room/SSLC/meetings)."]:
    para(d, s, size=9, after=2)
d.save(OUT_PAULA)

# ---------------------------------------------------------------- 2. SSLC Board
d = new_doc()
header(d, [
    ("From:", "John Quinn <jquinn@exosaero.com>"),
    ("To:", "Jim Wertz <jim@smad.com>; Jim Wertz <jwertz@smad.com>; Alice Wertz <awertz@smad.com>; Markus Rufer <mrufer@smad.com>; "
            "Robert Conger <rconger@smad.com>"),
    ("Subject:", "SSLC Board: notice, agenda and motions for the special meeting on Tuesday, October 20"),
    ("Attachments:", "SSLC_Director_Violation_Notice_to_Board_2026-10-19.docx; SSLC_RelatedParty_Reclass_Findings_2026-10-19.docx; "
                     "SSLC_Board_Motions_Voting_Eligibility_2026-10-19.docx"),
])
for t in [
    "Directors,",
    "Please find attached the papers for Tuesday’s meeting. Alice, as Secretary, please send notice under Article III, Section 8 by [____ day/time] "
    "for a special meeting on Tuesday, October 20 at [____] PT by Teams, and enter this email and its attachments in the minute book. If notice "
    "has not gone out by then, I will call the meeting with a second director, as Section 8 allows.",
    "Attachments:\n"
    "1. Notice to the Board (October 19), setting out each item under the Bylaws.\n"
    "2. Findings on the related-party reclassifications: the April 30, 2026 entries, what they did to SSLC’s balance sheet, and why they do not "
    "meet basic accounting standards.\n"
    "3. Motions and voting eligibility: a plain-English description of each motion, which directors cannot vote, and the rule that says so.",
    "Agenda and motions",
    "1. Motion 1: the Exos offer, §3 exclusivity and right of first refusal, or SSLC’s written counteroffer (vote). I am interested and will abstain.\n"
    "2. Motions 2–4, approved October 13: status. Motion 2 is done: access was established through Mission Control on October 19, and Markus and Bob "
    "were each emailed their login instructions that day for read-only access to SSLC’s books and records. Motion 3, the corrected AR "
    "Unbilled and customer-deposit accounting, due Tuesday. Paula Lessnick of BBSNV is preparing a worked example of the deposit and "
    "progress-billing entries for A La Carte. Motion 4, the ERP/MRP evaluation, is in progress at Exos’s cost on the same Mission Control platform.\n"
    "3. Motion 5: the related-party balances owed to Microcosm, MDAB and the Wertz family ($857,419 of $1,020,690) (vote).\n"
    "4. New resolutions RP-1 to RP-6 on the self-dealing entries: a formal finding; reversal of the April 30, 2026 entries; a receivable for "
    "Microcosm’s share of the building; a freeze on insider payments; an independent CPA review; and posting controls (roll-call votes).\n"
    "5. The Bookstore interest waiver: signed MDAB release, reversing entry and resolution.\n"
    "6. The October 16 letter of intent for SSLC and Microcosm: receive it; resolve that the October 27, 2015 60/40 split does not bind SSLC; "
    "written confirmation of which tank-making assets SSLC owns; "
    "independent financial adviser.\n"
    "6a. Motion NC-1, NewCap Partners: Danny has given sound advice to Jim, SSLC and Exos. So that it is handled correctly, record that "
    "NewCap is engaged by Jim and earns 5% on a sale of SSLC and/or Microcosm, put the engagement letter in the minute book, and have any "
    "SSLC-paid part of the fee approved by disinterested directors or by shareholders with the sale.\n"
    "7. The Exos commission: confirm the Board’s April 10, 2025 approval (15% on new business, 5% on repeat customers, paid on receipt of funds) "
    "stands, and vote on the written agreement. I am interested and will abstain.\n"
    "8. Payment schedule for the Exos loan and the Space Tech Expo invoice. I am interested and will abstain.\n"
    "9. Information items: monthly statements June–September 2026 with an officer’s certificate; terms of the Intuitive Machines order; an "
    "accounting of hardware orders.",
    "Who cannot vote. Under Corporations Code §310, a director with a material financial interest in a transaction, or who sits on the other "
    "company’s board, cannot have their vote counted toward approving it. On Microcosm, MDAB and Wertz items, that is Jim, Alice and Bob: all "
    "three are Microcosm directors, and Microcosm’s stock distribution shows Alice with 50.52%, Jim 20.31% and Bob 7.81%. Markus holds 0.78% of "
    "Microcosm, which is not material; he should disclose it, and his vote counts. On Exos items, that is me. The attached motions paper sets "
    "this out item by item. I ask that each director state their interests at the start of the meeting and that the minutes record each "
    "director’s vote.",
    "If the Board does not take up these items on Tuesday, Exos, which holds 18.06% of SSLC’s shares, intends to call a special meeting of "
    "shareholders under Article II, Section 3.",
    "Regards,\nJohn",
    "John Quinn\nDirector, Scorpius Space Launch Company",
]:
    para(d, t)
d.save(OUT_BOARD)
print(OUT_PAULA); print(OUT_BOARD)
