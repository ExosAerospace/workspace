#!/usr/bin/env python3.12
"""AstroX briefing for John (10/15). Editable .docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
OUT="drafts/AstroX_Briefing_2026-10-15_r2.docx"
NAVY=RGBColor(0x0A,0x3D,0x7C)
d=Document(); s=d.styles["Normal"]; s.font.name="Calibri"; s.font.size=Pt(10.5)
for sec in d.sections: sec.left_margin=sec.right_margin=Inches(0.9)
def h(t):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=NAVY; p.paragraph_format.space_before=Pt(8)
def para(t,i=False):
    p=d.add_paragraph(); r=p.add_run(t); r.italic=i; p.paragraph_format.space_after=Pt(4)
def bl(t,lead=None):
    p=d.add_paragraph(style="List Bullet")
    if lead: p.add_run(lead).bold=True
    p.add_run(t); p.paragraph_format.space_after=Pt(2)
t=d.add_paragraph(); r=t.add_run("AstroX, Inc. — Briefing"); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=NAVY
para("Prepared for John Quinn  |  October 15, 2026  |  INTERNAL  |  Met Shobu Oda (CEO) in person, Hokkaido Spaceport event, Oct 13–14", i=True)

h("Bottom line")
para("AstroX is a small, Series-A-funded Japanese launch startup building a 'rockoon' — a rocket lifted by balloon to the stratosphere and launched from a gondola. "
     "Every test drops expensive balloon-side hardware from 20–40 km. Exos's recovery and reflight know-how fits that pain directly, and BLK3 can carry their customers' "
     "payloads to space before their own system is operational. Good fit (4/5 in our vehicle-sales program), modest budget — start with a small, paid recovery engagement.")

h("Company snapshot")
bl(" AstroX, Inc. — HQ 1-87 Honmachi, Odaka-ku, Minamisoma-shi, Fukushima; Tokyo office (Taito-ku); R&D centers in Chiba and Fukushima. ~11 employees (PitchBook).","Who:")
bl(" balloon-launched rocket for small satellites. Dec 2025: aerial separation test of a dummy rocket from the balloon launcher; goal of reaching space by end-2026.","What:")
bl(" May 28, 2026 joint business demonstration with JAXA of a suspended attitude-control device for stratospheric balloons (https://global.jaxa.jp/press/2026/05/20260528-1_e.html).","JAXA tie:")
bl(" Series A closed Feb 2026 at ¥2.32B (~US$15.6M); backers include ANRI, DEEPCORE, Daiwa Corporate Investment, SMBC, Mizuho, One Capital, Inclusion Japan (https://japanstartupobserver.substack.com/p/astrox-closes-series-a-at-232-billion). ~US$18M total raised per CB Insights. Target launch cost under ¥500M.","Funding:")
bl(" enough for 2026 stratospheric tests, thin for orbital development — price any offer accordingly.","Read:")

h("People")
bl(" Founder & Representative Director, CEO — met you in person. shobu_oda@astrox.jp. Decision-maker.","Shobu Oda —")
bl(" Director & CTO; professor at Chiba Institute of Technology.","Yutaka Wada —")
bl(" Chief Strategy Officer (since Jan 2025).","Sumito Togo —")
bl(" Executive Officer; ran the first H-IIA launch at Mitsubishi Heavy Industries — senior credibility on their side.","Takashi Maemura —")
para("Source: AstroX team page (https://astrox.jp/team?lang=en). Emails other than Oda's are unverified.", i=True)

h("Their pain → our offer")
bl(" the gondola/launcher platform, the JAXA attitude-control device and test rockets fall from altitude each flight; no public recovery method. Losing reusable balloon-side hardware works against their low-cost target.","Pain:")
bl(" recovery of the gondola/launcher platform and attitude-control device (drogue + GPS-guided parafoil) so they can be reflown.","Offer 1 — Platform recovery:")
bl(" recovery of their first 'reach space' test rockets for post-flight analysis.","Offer 2 — Test-rocket recovery:")
bl(" fly their customers' payloads on BLK3 before the rockoon is operational — builds their customer pipeline and our manifest interest.","Offer 3 — Bridge flights:")
bl(" SARGE parafoil/drogue recovery flights (2018–19); Exos operates under a Part 431 license that is transitioning to Part 450. BLK3 is our next vehicle, and its first flight is targeted for early 2027 — it has not flown yet; hover test is planned.","Credentials to use (exact wording):")

h("Connections worth using")
bl(" Innovative Space Carrier (our #1 target, also met in Hokkaido) opened a factory in Minamisoma — same city.","Minamisoma:")
bl(" shared investor with ISC (Animal Spirits, per our target research) — a possible warm path.","Investors:")
bl(" both are JAXA-linked; a joint 'recovery for Japanese launch startups' offer could cover both.","Package:")

h("Due diligence and rules")
bl(" denied-party (CSL) clean for AstroX and Oda; CRM DD 'caution' only because early-stage and small — nothing adverse found.","Screening:")
bl(" Japan is a cleared ally, but recovery-system design help to a foreign launch developer is likely an ITAR defense service → TAA before any technical exchange (Empowered Official John Quinn, reviewer Denise Shirey). Capability-level talk is fine now.","Export:")
bl(" no pricing until you approve; BLK3 timing 'early 2027'; no Hokkaido launch date.","Commercial:")

h("Talking points for the next call")
for x in ["'Every flight you recover is hardware you don't rebuild. Recovery is what we do — it's how we make every space vehicle reusable.'",
          "'We can start small: recover your gondola and attitude-control device on your next stratospheric tests.'",
          "'Until your rocket is flying, your customers can fly on ours — that keeps them with you.'"]:
    bl(" "+x)
h("Questions to ask Oda-san")
for x in ["How do you recover the gondola and attitude-control device today, and what does a lost set cost you?",
          "When are your next stratospheric flights, and where (over land or sea)?",
          "Do you have customers waiting for a ride before the rockoon is ready?",
          "Who decides on outside services — you, Wada-san or Togo-san?"]:
    bl(" "+x)

h("Next steps")
bl(" Toshi: Japanese thank-you + propose a recovery-capability call (task #4 on Toshi's trip list; MC tasks wait on your approval of D-7282121).","1.")
bl(" Exo: one-page recovery capability sheet (no technical data) for that call.","2.")
bl(" If interest: small paid Phase 1 on the Honda/NYK model, TAA filed for technical scope.","3.")
para("Records: CRM (Shobu Oda, met in person 10/14); vehicle-sales target #17 (drafts/vehiclesales/targets_master_50.json); card data/business-cards/2026-10-14_Shobu_Oda_AstroX.jpg.", i=True)
d.save(OUT); print("WROTE",OUT)
