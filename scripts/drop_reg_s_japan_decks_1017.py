#!/usr/bin/env python3
"""Jon 10/17: drop Reg S from the Japan decks (no committed Reg S investors; add back only if required). r7 -> r8."""
import re
from pptx import Presentation
DISC="Not an offer to sell or a solicitation of an offer to buy securities. Any offering will be made only through definitive offering documents. Financials FY2025, accrual basis, unaudited."
def set_para(p,t):
    rs=p.runs
    rs[0].text=t
    for r in rs[1:]: r._r.getparent().remove(r._r)
for n in ("SBI_Investment","650_Fund"):
    src=f"drafts/Exos_Japan_Round_{n}_2026-10-16_r7.pptx"; dst=f"drafts/Exos_Japan_Round_{n}_2026-10-17_r8.pptx"
    prs=Presentation(src)
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            for p in list(sh.text_frame.paragraphs):
                t="".join(r.text for r in p.runs)
                if not p.runs: continue
                if t.startswith("Not an offer to sell") and "Regulation S" in t: set_para(p,DISC)
                elif t.strip()=="JAPAN ROUND — REG S": set_para(p,"JAPAN ROUND")
                elif "Offshore Regulation S" in t: p._p.getparent().remove(p._p)
                elif "Reg S round" in t: set_para(p,t.replace("Term discussion on the Reg S round","Term discussion"))
    prs.save(dst)
    left=[(i,sh.text_frame.text[:80]) for i,s in enumerate(prs.slides,1) for sh in s.shapes if sh.has_text_frame and re.search(r"Reg(ulation)?\.? ?S\b",sh.text_frame.text)]
    print(dst,"remaining Reg S:",left)
