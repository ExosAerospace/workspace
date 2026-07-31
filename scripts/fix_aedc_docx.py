#!/usr/bin/env python3
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn

PATH = 'Exos_AEDC_Velocity_Alliance_Whitepaper.docx'
d = Document(PATH)

for s in d.sections:
    sectPr = s._sectPr
    pgSz = sectPr.find(qn('w:pgSz'))
    print("BEFORE pgSz:", {k.split('}')[-1]: v for k, v in (pgSz.attrib.items() if pgSz is not None else [])})
    # Force US Letter portrait
    if pgSz is None:
        pgSz = sectPr.makeelement(qn('w:pgSz'), {})
        sectPr.append(pgSz)
    pgSz.set(qn('w:w'), '12240')   # 8.5"
    pgSz.set(qn('w:h'), '15840')   # 11"
    if qn('w:orient') in pgSz.attrib:
        del pgSz.attrib[qn('w:orient')]
    # Margins
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(0.5)
    s.top_margin = Inches(2.0)
    s.bottom_margin = Inches(1.4)
    s.header_distance = Inches(0.01)
    s.footer_distance = Inches(0.01)
    # Remove any text-column gutter weirdness: force single column, sane spacing
    cols = sectPr.find(qn('w:cols'))
    if cols is not None:
        for a in list(cols.attrib):
            del cols.attrib[a]
        cols.set(qn('w:num'), '1')
        cols.set(qn('w:space'), '720')

# Clear stray indents on body paragraphs (keep our intentional 0.3 bullets)
for p in d.paragraphs:
    pf = p.paragraph_format
    if pf.right_indent is not None:
        pf.right_indent = None

d.save(PATH)

# verify
d2 = Document(PATH)
for s in d2.sections:
    print("AFTER  page_w:", round(s.page_width.inches, 2),
          "L:", round(s.left_margin.inches, 2),
          "R:", round(s.right_margin.inches, 2),
          "usable:", round((s.page_width - s.left_margin - s.right_margin).inches, 2))
print("SAVED", PATH)
