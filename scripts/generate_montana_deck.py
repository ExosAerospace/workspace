#!/usr/bin/env python3
"""Generate the Exos Aerospace x Montana partnership PPTX deck."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# Colors
NAVY = RGBColor(0x0B, 0x1D, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_BLUE = RGBColor(0x1E, 0x90, 0xFF)
LIGHT_GRAY = RGBColor(0xF0, 0xF2, 0xF5)
MED_GRAY = RGBColor(0x8A, 0x8A, 0x8A)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
TABLE_HEADER_BG = RGBColor(0x0B, 0x1D, 0x3A)
TABLE_ROW_ALT = RGBColor(0xE8, 0xEE, 0xF7)
TABLE_ROW_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height


def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=18, color=DARK_TEXT,
                bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_multiline(slide, left, top, width, height, lines, font_size=16, color=DARK_TEXT,
                  bold=False, line_spacing=1.3, font_name="Calibri", alignment=PP_ALIGN.LEFT):
    """lines is a list of (text, bold, font_size_override, color_override) tuples or just strings."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if isinstance(line, str):
            txt, b, fs, c = line, bold, font_size, color
        else:
            txt = line[0]
            b = line[1] if len(line) > 1 else bold
            fs = line[2] if len(line) > 2 else font_size
            c = line[3] if len(line) > 3 else color
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(fs)
        p.font.color.rgb = c
        p.font.bold = b
        p.font.name = font_name
        p.alignment = alignment
        p.space_after = Pt(fs * 0.4)
    return txBox


def add_accent_bar(slide, left, top, width=Inches(0.08), height=Inches(0.8)):
    return add_shape(slide, left, top, width, height, ACCENT_BLUE)


def slide_number(slide, num):
    add_textbox(slide, Inches(12.5), Inches(7.0), Inches(0.7), Inches(0.4),
                str(num), font_size=11, color=MED_GRAY, alignment=PP_ALIGN.RIGHT)


def add_table(slide, left, top, width, rows_data, col_widths):
    """rows_data: list of lists. First row is header."""
    num_rows = len(rows_data)
    num_cols = len(rows_data[0])
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, Inches(0.45 * num_rows))
    table = table_shape.table

    for ci, cw in enumerate(col_widths):
        table.columns[ci].width = cw

    for ri, row in enumerate(rows_data):
        for ci, cell_text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = cell_text
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.name = "Calibri"
                if ri == 0:
                    p.font.bold = True
                    p.font.color.rgb = WHITE
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = TABLE_HEADER_BG
                else:
                    p.font.color.rgb = DARK_TEXT
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = TABLE_ROW_ALT if ri % 2 == 1 else TABLE_ROW_WHITE
                p.alignment = PP_ALIGN.LEFT
            cell.text_frame.paragraphs[0].space_before = Pt(4)
            cell.text_frame.paragraphs[0].space_after = Pt(4)
            cell.margin_left = Inches(0.1)
    return table_shape


# ============================================================
# SLIDE 1 — TITLE
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide, NAVY)

# Accent stripe
add_shape(slide, Inches(0), Inches(0), Inches(0.15), H, ACCENT_BLUE)

add_textbox(slide, Inches(1.0), Inches(1.5), Inches(11), Inches(1.2),
            "EXOS AEROSPACE", font_size=52, color=WHITE, bold=True)
add_textbox(slide, Inches(1.0), Inches(2.7), Inches(11), Inches(0.8),
            "Reusable Flight Test Infrastructure for the Space & Defense Industrial Base",
            font_size=24, color=ACCENT_BLUE, bold=False)

# Divider line
add_shape(slide, Inches(1.0), Inches(3.8), Inches(4), Inches(0.04), ACCENT_BLUE)

add_multiline(slide, Inches(1.0), Inches(4.2), Inches(10), Inches(2.0), [
    ("Partnership Opportunity: Montana Spaceport & Test Range", True, 28, WHITE),
    ("", False, 12),
    ("Presented to: Montana Spaceport & Test Range Leadership", False, 16, MED_GRAY),
    ("April 2026", False, 16, MED_GRAY),
])
slide_number(slide, 1)


# ============================================================
# SLIDE 2 — WHO IS EXOS AEROSPACE?
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT_BLUE)

add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
            "WHO IS EXOS AEROSPACE?", font_size=32, color=NAVY, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.5),
            "One of only 3 U.S.-licensed reusable suborbital launch companies",
            font_size=18, color=ACCENT_BLUE, bold=True)

# Left column — key facts
add_accent_bar(slide, Inches(0.8), Inches(1.8), height=Inches(3.5))
add_multiline(slide, Inches(1.1), Inches(1.8), Inches(5.5), Inches(3.5), [
    ("Founded in Greenville, TX", True, 16),
    ("FAA reusable launch license renewed March 2026", False, 14),
    ("", False, 8),
    ("The only operator currently flying wet-lease flight test services", True, 16),
    ("", False, 8),
    ("Profitable since December 2025", True, 16),
    ("Revenue doubled 4 consecutive years", False, 14),
    ("", False, 8),
    ("Active government contracts:", True, 16),
    ("USAF SBIR Phase II", False, 14),
    ("DoD RDRE propulsion flight demo", False, 14),
    ("NASA REDDI program", False, 14),
])

# Right column — value prop box
box = add_shape(slide, Inches(7.2), Inches(1.8), Inches(5.5), Inches(3.2), LIGHT_GRAY)
add_multiline(slide, Inches(7.5), Inches(1.9), Inches(5.0), Inches(3.0), [
    ("WHAT WE DO", True, 18, NAVY),
    ("", False, 8),
    ("Customers fly subsystems, not missions.", True, 15, DARK_TEXT),
    ("", False, 6),
    ("We provide the vehicle. They test avionics, propulsion,", False, 14),
    ("GNC, reentry systems, and hypersonic technologies", False, 14),
    ("at a fraction of the cost and timeline.", False, 14),
    ("", False, 10),
    ("CORE VALUE:", True, 15, ACCENT_BLUE),
    ("Reduce program risk", False, 14),
    ("Compress timelines", False, 14),
    ("Cut cost of failure", False, 14),
])

slide_number(slide, 2)


# ============================================================
# SLIDE 3 — THE MARKET GAP
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT_BLUE)

add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
            "THE MARKET GAP WE FILL", font_size=32, color=NAVY, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.5),
            "The flight test bottleneck is real", font_size=18, color=ACCENT_BLUE, bold=True)

# Problem bullets
add_multiline(slide, Inches(0.8), Inches(1.7), Inches(5.5), Inches(1.5), [
    ("Wind tunnels are saturated nationwide", False, 15),
    ("Computational modeling has known limits", False, 15),
    ("Flight truth is the constraint", True, 16, ACCENT_BLUE),
    ("No substitute for actual flight data", False, 14, MED_GRAY),
])

# Table
table_data = [
    ["Challenge", "Exos Solution"],
    ["Multi-year test vehicle development", "Ready-to-fly reusable platform"],
    ["$50M+ custom flight test programs", "Wet-lease access from ~$1M per flight"],
    ["18-24 month lead times", "Weeks to integration, months to flight"],
    ["Single-use test articles", "Reusable — fly, recover, fly again"],
]
add_table(slide, Inches(0.8), Inches(3.5), Inches(11.5), table_data,
          [Inches(5.75), Inches(5.75)])

# Customers line
add_textbox(slide, Inches(0.8), Inches(6.2), Inches(11), Inches(0.5),
            "Customers: Defense primes | Emerging launch companies | University research | Government agencies",
            font_size=14, color=MED_GRAY)

slide_number(slide, 3)


# ============================================================
# SLIDE 4 — BLK3 PSRLV PLATFORM
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT_BLUE)

add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
            "OUR PLATFORM: BLK3 PSRLV", font_size=32, color=NAVY, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.5),
            "Partially Reusable Suborbital Launch Vehicle", font_size=18, color=ACCENT_BLUE, bold=True)

# Left — specs
add_accent_bar(slide, Inches(0.8), Inches(1.8), height=Inches(4.0))
add_multiline(slide, Inches(1.1), Inches(1.8), Inches(5.5), Inches(4.0), [
    ("CAPABILITIES", True, 16, NAVY),
    ("", False, 6),
    ("Suborbital flight profile reaching 100+ km altitude", False, 14),
    ("Payload capacity for flight test articles", False, 14),
    ("Vertical launch, parachute recovery, reusable", False, 14),
    ("Integrated telemetry, tracking, and data return", False, 14),
    ("FAA-licensed operations", False, 14),
    ("", False, 10),
    ("FLIGHT HERITAGE", True, 16, NAVY),
    ("", False, 6),
    ("Multiple successful flights and recoveries", False, 14),
    ("Proven integration with customer payloads", False, 14),
    ("Flight-qualified avionics, GNC, and recovery systems", False, 14),
    ("", False, 10),
    ("ROADMAP", True, 16, NAVY),
    ("", False, 6),
    ("Next BLK3 build in progress", False, 14),
    ("Expanding fleet capacity to meet growing demand", False, 14),
])

# Right — placeholder for vehicle image
img_box = add_shape(slide, Inches(7.2), Inches(1.8), Inches(5.3), Inches(4.5), LIGHT_GRAY)
add_textbox(slide, Inches(7.5), Inches(3.5), Inches(4.7), Inches(1.0),
            "[BLK3 PSRLV Vehicle Photo]", font_size=16, color=MED_GRAY,
            alignment=PP_ALIGN.CENTER)

slide_number(slide, 4)


# ============================================================
# SLIDE 5 — SCORPIUS STRATEGIC ASSET
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT_BLUE)

add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
            "STRATEGIC ASSET: SCORPIUS", font_size=32, color=WHITE, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.8),
            "World's only TRL 9 Type 5 cryogenic tank technology — demonstrated on the Moon",
            font_size=18, color=ACCENT_BLUE, bold=True)

# Key points
add_multiline(slide, Inches(0.8), Inches(2.0), Inches(5.8), Inches(3.5), [
    ("LUNAR HERITAGE", True, 16, ACCENT_BLUE),
    ("Type 5 cryo tanks flew as fuel and LOX tanks on", False, 14, WHITE),
    ("Intuitive Machines Odysseus and Athena lunar missions", True, 14, WHITE),
    ("", False, 8),
    ("APPLICATIONS", True, 16, ACCENT_BLUE),
    ("Reusability", False, 14, WHITE),
    ("In-space refueling", False, 14, WHITE),
    ("Long-duration cryogenic storage", False, 14, WHITE),
    ("In-space architecture", False, 14, WHITE),
    ("", False, 8),
    ("No known competitor at this TRL with lunar heritage", True, 15, ACCENT_BLUE),
])

# Montana angle box
box = add_shape(slide, Inches(7.2), Inches(2.0), Inches(5.5), Inches(2.5), RGBColor(0x12, 0x2A, 0x4F))
add_multiline(slide, Inches(7.5), Inches(2.1), Inches(5.0), Inches(2.3), [
    ("WHY THIS MATTERS FOR MONTANA", True, 16, ACCENT_BLUE),
    ("", False, 6),
    ("Scorpius cryogenic tank technology could be", False, 14, WHITE),
    ("manufactured and tested at Montana facilities —", False, 14, WHITE),
    ("creating high-value aerospace manufacturing", False, 14, WHITE),
    ("jobs anchored by proven, flight-qualified hardware.", False, 14, WHITE),
])

# Placeholder for IM lunar lander image
img_box = add_shape(slide, Inches(7.2), Inches(4.8), Inches(5.5), Inches(2.0), RGBColor(0x12, 0x2A, 0x4F))
add_textbox(slide, Inches(7.5), Inches(5.3), Inches(5.0), Inches(1.0),
            "[Intuitive Machines Lunar Lander Image]", font_size=14, color=MED_GRAY,
            alignment=PP_ALIGN.CENTER)

slide_number(slide, 5)


# ============================================================
# SLIDE 6 — WHY MONTANA?
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT_BLUE)

add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
            "WHY MONTANA?", font_size=32, color=NAVY, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.5),
            "Natural advantages that align perfectly with Exos operations",
            font_size=18, color=ACCENT_BLUE, bold=True)

# 3 columns
col_w = Inches(3.6)
col_gap = Inches(0.4)
col_starts = [Inches(0.8), Inches(0.8) + col_w + col_gap, Inches(0.8) + 2 * (col_w + col_gap)]

# Column 1 — Geographic
add_shape(slide, col_starts[0], Inches(1.7), col_w, Inches(0.5), NAVY)
add_textbox(slide, col_starts[0], Inches(1.7), col_w, Inches(0.5),
            "  GEOGRAPHIC & REGULATORY", font_size=14, color=WHITE, bold=True)
add_multiline(slide, col_starts[0] + Inches(0.1), Inches(2.3), col_w - Inches(0.2), Inches(3.5), [
    ("High elevation (~5,400 ft at Butte)", False, 13),
    ("Favorable for launch operations", False, 12, MED_GRAY),
    ("", False, 6),
    ("Low population density", False, 13),
    ("Ideal for range safety & FAA licensing", False, 12, MED_GRAY),
    ("", False, 6),
    ("Remote terrain, clear airspace", False, 13),
    ("Minimal commercial aviation conflict", False, 12, MED_GRAY),
    ("", False, 6),
    ("No state sales tax", True, 13, ACCENT_BLUE),
    ("Major cost advantage for equipment", False, 12, MED_GRAY),
])

# Column 2 — Infrastructure
add_shape(slide, col_starts[1], Inches(1.7), col_w, Inches(0.5), NAVY)
add_textbox(slide, col_starts[1], Inches(1.7), col_w, Inches(0.5),
            "  EXISTING INFRASTRUCTURE", font_size=14, color=WHITE, bold=True)
add_multiline(slide, col_starts[1] + Inches(0.1), Inches(2.3), col_w - Inches(0.2), Inches(3.5), [
    ("AeroTec Test Facility (MADA)", True, 13),
    ("Established ground test capability", False, 12, MED_GRAY),
    ("", False, 6),
    ("Montana Connections TIFID", False, 13),
    ("Tax increment financing for development", False, 12, MED_GRAY),
    ("", False, 6),
    ("Growing aerospace cluster", False, 13),
    ("GE Aerospace JV, Lattice Materials", False, 12, MED_GRAY),
    ("", False, 6),
    ("Montana State University", True, 13, ACCENT_BLUE),
    ("Top-3 academic small sat program", False, 12, MED_GRAY),
])

# Column 3 — State support
add_shape(slide, col_starts[2], Inches(1.7), col_w, Inches(0.5), NAVY)
add_textbox(slide, col_starts[2], Inches(1.7), col_w, Inches(0.5),
            "  STATE SUPPORT", font_size=14, color=WHITE, bold=True)
add_multiline(slide, col_starts[2] + Inches(0.1), Inches(2.3), col_w - Inches(0.2), Inches(3.5), [
    ("406 JOBS Initiative", True, 13),
    ("Advanced mfg as priority sector", False, 12, MED_GRAY),
    ("", False, 6),
    ("Job Growth Incentive Credit", False, 13),
    ("50% FICA on qualifying new hires", False, 12, MED_GRAY),
    ("", False, 6),
    ("CDBG-ED Grants", False, 13),
    ("For aerospace infrastructure", False, 12, MED_GRAY),
    ("", False, 6),
    ("NASA MTSGC at MSU", True, 13, ACCENT_BLUE),
    ("Space Grant Consortium", False, 12, MED_GRAY),
])

slide_number(slide, 6)


# ============================================================
# SLIDE 7 — WHAT EXOS WANTS TO DO IN MONTANA
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT_BLUE)

add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
            "WHAT EXOS WANTS TO DO IN MONTANA", font_size=32, color=NAVY, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.5),
            "Establish Montana as a premier flight test and propulsion development hub",
            font_size=18, color=ACCENT_BLUE, bold=True)

# 3 phase boxes
phases = [
    ("PHASE 1 — TEST OPERATIONS", "Year 1-2", [
        "Conduct static fire testing and propulsion",
        "development at AeroTec/MADA",
        "",
        "RDRE ground test campaigns",
        "",
        "Bring existing DoD/NASA contract",
        "work to Montana facilities",
        "",
        "Establish local engineering and",
        "operations team",
    ]),
    ("PHASE 2 — SUBORBITAL FLIGHT OPS", "Year 2-3", [
        "Work with FAA toward launch site",
        "licensing for suborbital flights",
        "",
        "Leverage Montana geography for",
        "expanded flight corridors",
        "",
        "Wet-lease flight test services from",
        "Montana for defense & commercial",
        "",
        "TX + MT = schedule flexibility",
        "and surge capacity",
    ]),
    ("PHASE 3 — ECOSYSTEM ANCHOR", "Year 3-5", [
        "Scorpius cryogenic tank manufacturing",
        "and test operations",
        "",
        "University partnerships",
        "(MSU, Montana Tech)",
        "",
        "Workforce training pipeline for",
        "aerospace technicians & engineers",
        "",
        "Attract supply chain and adjacent",
        "aerospace companies to the region",
    ]),
]

for i, (title, timeline, bullets) in enumerate(phases):
    x = Inches(0.8) + i * (Inches(3.8) + Inches(0.3))
    w = Inches(3.8)

    # Header
    add_shape(slide, x, Inches(1.7), w, Inches(0.5), NAVY)
    add_textbox(slide, x, Inches(1.7), w, Inches(0.5),
                f"  {title}", font_size=13, color=WHITE, bold=True)

    # Timeline badge
    add_shape(slide, x, Inches(2.2), Inches(1.2), Inches(0.35), ACCENT_BLUE)
    add_textbox(slide, x, Inches(2.2), Inches(1.2), Inches(0.35),
                f"  {timeline}", font_size=11, color=WHITE, bold=True)

    # Bullets
    lines = []
    for b in bullets:
        if b == "":
            lines.append(("", False, 6))
        else:
            lines.append((b, False, 13))
    add_multiline(slide, x + Inches(0.15), Inches(2.7), w - Inches(0.3), Inches(4.0), lines)

slide_number(slide, 7)


# ============================================================
# SLIDE 8 — ECONOMIC IMPACT
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT_BLUE)

add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
            "ECONOMIC IMPACT", font_size=32, color=NAVY, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.5),
            "What Exos brings to Montana", font_size=18, color=ACCENT_BLUE, bold=True)

# 4 quadrants
quad_data = [
    (Inches(0.8), Inches(1.7), "DIRECT JOBS", [
        ("Phase 1: 10-15 positions", False, 14),
        ("Engineering and operations", False, 12, MED_GRAY),
        ("", False, 6),
        ("Phase 2: 25-40 positions", False, 14),
        ("Launch ops, payload integration, range safety", False, 12, MED_GRAY),
        ("", False, 6),
        ("Phase 3: 50-75+ positions", True, 14, ACCENT_BLUE),
        ("Manufacturing, test, launch, administration", False, 12, MED_GRAY),
    ]),
    (Inches(7.0), Inches(1.7), "FEDERAL DOLLARS", [
        ("Active DoD contracts", True, 14),
        ("RDRE flight demo, SBIR Phase II", False, 12, MED_GRAY),
        ("", False, 6),
        ("NASA program participation", False, 14),
        ("Research and test campaigns", False, 12, MED_GRAY),
        ("", False, 6),
        ("Up to $50M follow-on potential", True, 14, ACCENT_BLUE),
        ("STRATFI / TACFI / OTA pathways", False, 12, MED_GRAY),
    ]),
    (Inches(0.8), Inches(4.5), "ECOSYSTEM DEVELOPMENT", [
        ("Supply chain attracts supporting businesses", False, 14),
        ("University research creates talent pipeline", False, 14),
        ("", False, 6),
        ("Positions Montana alongside TX, FL, VA", True, 14, ACCENT_BLUE),
        ("as a recognized aerospace operations state", False, 12, MED_GRAY),
    ]),
    (Inches(7.0), Inches(4.5), "WORKFORCE TRAINING", [
        ("Aerospace technician certification", False, 14),
        ("University & community college partnerships", False, 14),
        ("Veteran transition programs", False, 14),
        ("", False, 6),
        ("Aerospace / defense crossover skills", True, 13, ACCENT_BLUE),
    ]),
]

for x, y, title, lines in quad_data:
    add_shape(slide, x, y, Inches(5.5), Inches(0.4), NAVY)
    add_textbox(slide, x, y, Inches(5.5), Inches(0.4),
                f"  {title}", font_size=13, color=WHITE, bold=True)
    add_multiline(slide, x + Inches(0.15), y + Inches(0.5), Inches(5.2), Inches(2.2), lines)

slide_number(slide, 8)


# ============================================================
# SLIDE 9 — WHY EXOS + MONTANA WINS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT_BLUE)

add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
            "WHY EXOS + MONTANA WINS", font_size=32, color=NAVY, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.5),
            "The partnership logic is clear", font_size=18, color=ACCENT_BLUE, bold=True)

# Partnership table
table_data = [
    ["Montana Provides", "Exos Provides"],
    ["Test infrastructure (AeroTec/MADA)", "FAA-licensed reusable launch capability"],
    ["State incentives & tax advantages", "Proven flight heritage & government contracts"],
    ["Geographic advantages for range safety", "Immediate customer demand (DoD, NASA, commercial)"],
    ["University talent pipeline", "Aerospace workforce training programs"],
    ["Political will & state support", "Anchor tenant credibility for ecosystem growth"],
]
add_table(slide, Inches(0.8), Inches(1.7), Inches(11.5), table_data,
          [Inches(5.75), Inches(5.75)])

# Together box
box = add_shape(slide, Inches(0.8), Inches(4.8), Inches(11.5), Inches(2.2), LIGHT_GRAY)
add_multiline(slide, Inches(1.2), Inches(4.9), Inches(10.5), Inches(2.0), [
    ("TOGETHER:", True, 18, NAVY),
    ("", False, 6),
    ("Montana gets a real, revenue-generating aerospace anchor — not a speculative bet", False, 15),
    ("Exos gets geographic diversification, cost advantages, and expanded capacity", False, 15),
    ("DoD/NASA get additional test range options in a strategically valuable location", False, 15),
    ("The region builds a durable aerospace economy with compounding returns", False, 15),
])

slide_number(slide, 9)


# ============================================================
# SLIDE 10 — NEXT STEPS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.15), H, ACCENT_BLUE)

add_textbox(slide, Inches(1.0), Inches(0.5), Inches(11), Inches(0.6),
            "NEXT STEPS", font_size=32, color=WHITE, bold=True)
add_textbox(slide, Inches(1.0), Inches(1.1), Inches(11), Inches(0.5),
            "Let's move this forward", font_size=18, color=ACCENT_BLUE, bold=True)

steps = [
    ("1", "SITE VISIT", "Exos leadership tours AeroTec facility and meets MADA team"),
    ("2", "TECHNICAL ASSESSMENT", "Evaluate facility capabilities against Exos test requirements"),
    ("3", "MOU / LETTER OF INTENT", "Formalize partnership intent and define terms"),
    ("4", "FAA ENGAGEMENT", "Begin pre-application discussions for launch site licensing"),
    ("5", "FUNDING COORDINATION", "Align state incentives, federal grants, and Exos investment"),
]

for i, (num, title, desc) in enumerate(steps):
    y = Inches(2.0) + i * Inches(0.9)
    # Number circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), y, Inches(0.5), Inches(0.5))
    circle.fill.solid()
    circle.fill.fore_color.rgb = ACCENT_BLUE
    circle.line.fill.background()
    tf = circle.text_frame
    tf.paragraphs[0].text = num
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    add_textbox(slide, Inches(1.8), y - Inches(0.05), Inches(3), Inches(0.35),
                title, font_size=16, color=WHITE, bold=True)
    add_textbox(slide, Inches(1.8), y + Inches(0.3), Inches(10), Inches(0.35),
                desc, font_size=14, color=MED_GRAY)

# Divider
add_shape(slide, Inches(1.0), Inches(6.3), Inches(4), Inches(0.04), ACCENT_BLUE)

# Contact
add_multiline(slide, Inches(1.0), Inches(6.5), Inches(5), Inches(0.9), [
    ("John Quinn", True, 16, WHITE),
    ("CEO & Cofounder, Exos Aerospace Systems & Technologies", False, 12, MED_GRAY),
    ("+1-972-740-8355  |  jon@exosaero.com  |  exosaero.com", False, 12, MED_GRAY),
])

slide_number(slide, 10)


# ============================================================
# SAVE
# ============================================================
out_path = os.path.expanduser("~/.openclaw/workspace/Exos_Montana_Partnership_Deck.pptx")
prs.save(out_path)
print(f"Saved: {out_path}")
