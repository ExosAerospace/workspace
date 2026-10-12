#!/usr/bin/env python3
"""DRAFT letter, John Quinn -> Jim Wertz: the April 30, 2026 related-party entries, as found and their impact
(John 10/12/2026). Shows the math, the appearance of self-dealing, and gives Jim an opportunity to correct.
The proxy is deliberately NOT mentioned (held in reserve per John).
Figures: SSLC QuickBooks (Balance Sheet 10/11/2026; Transaction Detail by Account 1/1/2024-10/12/2026), SSLC
08-2024 share register, Bylaws (2000), Statements of Information (2025, 2026).
Output: drafts/SSLC_Letter_to_JWertz_April2026_Entries_DRAFT_2026-10-12.docx (+ .pdf). NOT SENT.
"""
import os
import subprocess
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WS = "/Users/exomac/.openclaw/workspace"
RAW = f"{WS}/drafts/_sslc_wertz_letter_raw.docx"
OUT = f"{WS}/drafts/SSLC_Letter_to_JWertz_April2026_Entries_DRAFT_2026-10-12.docx"
RED = RGBColor(0xC0, 0x00, 0x00)
NAVY = RGBColor(0x1F, 0x3A, 0x68)

# Distribution base: holders other than Exos (Exos takes its 18% as the $790K stake credit, Option A)
NON_EXOS = 714972
WERTZ_GROUP = 383000          # Microcosm 280,000 + Wertz family 103,000 (08-2024 register)
OTHERS = NON_EXOS - WERTZ_GROUP
HOLDERS = [("Robert Conger", 73000), ("Markus Rufer", 91000), ("John Kulpa", 30332), ("Ed Keith", 22000),
           ("All other non-Wertz holders", OTHERS - 73000 - 91000 - 30332 - 22000)]
STEP1 = 351000 + 97302        # uncredited facility share + capitalized interest
STEP2 = STEP1 + 518201        # + netting offsets not yet supported by invoices

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)


def P(text="", bold=False, size=None, color=None, italic=False, align=None, sa=5, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(sa)
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def H(text):
    return P(text, bold=True, size=11.5, color=NAVY, sa=3)


def B(text):
    return P("•  " + text, indent=0.2, sa=3)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hex_)
    tcPr.append(s)


def table(header, rows, widths, bold_last=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        shade(c, "DCE6F1")
    for ri, r in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
            if bold_last and ri == len(rows) - 1:
                cells[i].paragraphs[0].runs[0].bold = True
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(9.5)
                if i > 0 and p.text[:1] in "$−—":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    P(sa=4)
    return t


k = lambda v: f"${v / 1000:,.0f}K"

P("DRAFT — NOT SENT  ·  Prepared for John Quinn review  ·  Confidential", bold=True, size=9, color=RED,
  align=WD_ALIGN_PARAGRAPH.CENTER)
P("October [__], 2026", sa=8)
P("James R. Wertz\nDirector, Scorpius Space Launch Company\nPresident, Microcosm, Inc.\n3111 Lomita Blvd, Torrance, CA 90505", sa=8)
P("Re: The April 30, 2026 Microcosm and Bookstore entries, and an opportunity to correct them before the Board acts",
  bold=True, sa=8)

P("Jim,")
P("I am writing to you directly, before raising this with the full Board or the shareholders, because I would rather "
  "we fix this together. In reviewing SSLC's QuickBooks for the acquisition, I found a set of entries recorded on "
  "April 30, 2026 that move value from SSLC's shareholders to Microcosm, Inc. and Microcosm Astronautics Books. I was "
  "not told about them and they were not brought to the Board. Below is what the books show, what it means for the "
  "shareholders, and what I am asking you to do.")

H("1. What the books show")
table(["April 30, 2026 entry (per SSLC's QuickBooks)", "Amount", "Effect"], [
    ["Accounts payable to Microcosm reclassified to 2385 Due to Microcosm", "$426,687",
     "Vendor invoices become a fixed insider balance to be paid at close"],
    ["Interest capitalized into 2700 Notes Payable, Bookstore", "$97,302", "New liability to Microcosm Astronautics Books"],
    ["SSLC's loans to Microcosm (2019 to Oct 2025) booked, then fully netted", "$812,699",
     "SSLC's receivable from Microcosm eliminated"],
    ["   offset against: Loan Payable, Microcosm", "$294,499", "Needs the loan documents"],
    ["   offset against: pre-6/2025 'reimbursables'", "$423,965", "No invoices provided"],
    ["   offset against: 'Judy's time'", "$61,290", "No invoices provided"],
    ["   offset against: reimbursable", "$32,946", "No invoices provided"],
], [3.3, 0.9, 2.3])
P("In addition, since July 2024 SSLC has paid 100% of the Lomita Blvd rent and utilities (it paid about 50% before), "
  "while Microcosm continues to occupy the building and keep its name on it. Microcosm's former share, about $316K of "
  "rent and $35K of utilities through August 2026, has never been credited against the $451,429 shown as Due to "
  "Microcosm. Over the same period, SSLC paid Microcosm entities about $1.31M (2024 to 2025), while Exos's $31,200 loan "
  "remains past due.", sa=6)

H("2. Why this looks like self-dealing")
P("I am not assuming bad intent. But anyone reviewing this, including a buyer, a fairness-opinion provider or a "
  "court, would see the following:")
B("Every entry favors entities you and Alice control, and every one increases what SSLC pays them at close, ahead of "
  "the shareholders.")
B("The entries were made under the direction of interested officers (you as Microcosm's President; Alice as SSLC's "
  "CFO and Secretary) and were not disclosed to or approved by the disinterested directors (Cal. Corp. Code §310).")
B("The accounting firm was retained over my objection and worked directly with you and Alice, while my director access "
  "to the books was delayed for months. The bylaws give every director an \"absolute right\" to inspect the books "
  "(Art. VII §4; Art. V §10; Cal. Corp. Code §1602).")
B("The timing: the entries were booked while a sale was pending, when insider balances paid at face come straight out "
  "of the proceeds going to shareholders.")
B("Insider interest was capitalized while the one outside lender, Exos, went unpaid.")

H("3. What it costs the other shareholders")
P(f"Every dollar paid to Microcosm or the Bookstore at close comes out of the proceeds shared by all {NON_EXOS:,} "
  f"non-Exos shares. The Microcosm/Wertz group holds {WERTZ_GROUP:,} of those shares, so about "
  f"{OTHERS / NON_EXOS:.0%} of each such dollar comes from the other holders.")
table(["", "Insider claims", "Change"], [
    ["As booked: Due to Microcosm $451,429 + Bookstore $320,582 + Wertz loan $92,500", "$864,511", "—"],
    ["Corrected for the facility share (~$351K) and the capitalized interest ($97K)", k(864511 - STEP1), f"−{k(STEP1)}"],
    ["...and if the $518K of netting offsets cannot be supported by invoices", "Microcosm owes SSLC", f"−{k(STEP2)}"],
], [4.0, 1.25, 1.25])
rows = []
for name, sh in HOLDERS:
    rows.append([f"{name} ({sh:,} shares)", k(STEP1 * sh / NON_EXOS), k(STEP2 * sh / NON_EXOS)])
rows.append(["Total, holders outside the Microcosm/Wertz group", k(STEP1 * OTHERS / NON_EXOS), k(STEP2 * OTHERS / NON_EXOS)])
table(["Lost by each holder if not corrected", "Step 1", "Step 1 + 2"], rows, [4.0, 1.25, 1.25], bold_last=True)
P("Base case at a January 15, 2027 close: about $1.51M of net cash to holders as booked, about $1.86M with the facility "
  "credit alone. Shares per SSLC's 08-2024 register; amounts are estimates until SSLC provides the support requested.",
  italic=True, size=9, sa=6)

H("4. The opportunity to make it right")
P("I would like to bring this to the Board as a correction you initiated. By Friday, October 16, 2026, please confirm "
  "in writing that you and Microcosm will:")
for i, t in enumerate([
    "Treat the April 30, 2026 entries as disputed (reversed or flagged) until they are reviewed, and provide the "
    "invoices and agreements behind each offset and the Bookstore and Wertz notes.",
    "Credit Microcosm's share of rent and utilities from July 2024 against the Due to Microcosm balance, or show the "
    "written agreement that says otherwise.",
    "Agree that no Microcosm, Bookstore or Wertz balance will be paid, netted, secured or accrue further interest "
    "before close without approval by the disinterested directors.",
    "Support an independent accountant, chosen by the disinterested directors, to reconcile these balances and the "
    "billed-ahead account.",
    "Give me full read-only QuickBooks access by October 15, as already requested.",
], 1):
    P(f"{i}.  {t}", indent=0.2, sa=3)
P("If we do this together, the shareholders, including the minority holders who have been with SSLC for 25 years, see "
  "accurate books and a clean path to close, and the question of how these entries were made does not have to go "
  "any further.", sa=5)
P("If we cannot resolve it by then, I will put the full detail before the Board and the shareholders and act on it as "
  "a director and on behalf of Exos as an 18.06% shareholder. Exos reserves all of its rights.", sa=8)
P("I would rather make the phone call than send the next letter. Call me any time.", sa=10)
P("Sincerely,", sa=18)
P("John M. Quinn", bold=True, sa=1)
P("Director, Scorpius Space Launch Company\nCofounder & CEO, Exos Aerospace Systems & Technologies, Inc.", sa=8)
P("Figures are from SSLC's QuickBooks (Balance Sheet 10/11/2026; Transaction Detail by Account 1/1/2024–10/12/2026) "
  "and SSLC's share register. Confidential under the July 27, 2021 NDA.", italic=True, size=8.5)

doc.save(RAW)
subprocess.run(["/usr/bin/python3", f"{WS}/scripts/apply_exos_letterhead.py", RAW, OUT], check=True)
os.remove(RAW)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", f"{WS}/drafts", OUT],
               check=True, capture_output=True)
print(OUT, f"step1={STEP1} step2={STEP2} others={OTHERS} ({OTHERS / NON_EXOS:.3f})")
