#!/usr/bin/env python3
"""Synthesize logs/mc-exo3-audit/slice-{A,B,C,D}.json into a Word report."""
import json, collections
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE = "logs/mc-exo3-audit"
OUT = "drafts/MC_exo3_Compliance_Audit_2026-09-27.docx"
S = {k: json.load(open(f"{BASE}/slice-{k}.json")) for k in "ABCD"}
TITLES = {"A": "Home, governance & admin pages", "B": "Business development pages",
          "C": "Operations, finance & ERP pages", "D": "Background agent fleet"}
COLOR = {"COMPLIANT": RGBColor(0x1B, 0x7F, 0x3B), "PARTIAL": RGBColor(0xB2, 0x6B, 0x00),
         "NON-COMPLIANT": RGBColor(0xB0, 0x1E, 0x1E)}

def verdict_of(sec):
    v = sec.get("verdict") or ""
    return v.upper().replace("NON_COMPLIANT", "NON-COMPLIANT")

def listify(x):
    if x is None: return []
    return x if isinstance(x, list) else [x]

def fixtext(f):
    if isinstance(f, dict):
        return f"[{f.get('priority','?')}/{f.get('effort','?')}] {f.get('fix') or f.get('text') or json.dumps(f)}"
    return str(f)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
doc.add_heading("Mission Control — exo3 Compliance Audit", 0)
doc.add_paragraph("Audit date: 2026-09-27 · Mode: read-only · Standard: governance/AGENT_SPECS.md, "
                  "governance/registry.json (8-property envelopes), agent-registry.json, govlog start/end/error, "
                  "human gate on external sends. INTERNAL — contains security findings; do not forward.")

# --- summary table
doc.add_heading("1. Scorecard", 1)
t = doc.add_table(rows=1, cols=5); t.style = "Light Grid Accent 1"
for i, h in enumerate(["Slice", "Sections", "Compliant", "Partial", "Non-compliant"]):
    t.rows[0].cells[i].text = h
tot = collections.Counter()
for k in "ABCD":
    c = collections.Counter(verdict_of(s) for s in S[k]["sections"])
    tot += c
    r = t.add_row().cells
    r[0].text = f"{k} — {TITLES[k]}"; r[1].text = str(len(S[k]["sections"]))
    r[2].text = str(c["COMPLIANT"]); r[3].text = str(c["PARTIAL"]); r[4].text = str(c["NON-COMPLIANT"])
r = t.add_row().cells
r[0].text = "TOTAL"; r[1].text = str(sum(tot.values()))
r[2].text = str(tot["COMPLIANT"]); r[3].text = str(tot["PARTIAL"]); r[4].text = str(tot["NON-COMPLIANT"])

doc.add_heading("Bottom line", 2)
for b in [
    "Registration is solid: all 148 launchd agents are in agent-registry.json (exo3 gate R1 holds).",
    "Telemetry is not: ~100 of 148 launchd agents emit nothing to governance/events.jsonl, including the "
    "DECIDE core (overnight resolver, decision triage, note router, legacy triage, CRM DD engine) and financial writers.",
    "All 33 openclaw crons are unregistered and untelemetered; 28 put a model in the critical path.",
    "Mission Control's own server writes are outside govlog (glog() only used for PDF renders) — no correlation IDs on human edits.",
    "Access control is the biggest real risk: several write endpoints accept any logged-in user, and the Supabase "
    "publishable key can read the full CRM. This ties directly to the MC-protection work.",
    "Good news: no autonomous external send to a non-Exos recipient was found. External sends are gated via Decisions or go to Jon.",
]:
    doc.add_paragraph(b, style="List Bullet")

# --- ranked fix list
doc.add_heading("2. Ranked fix list", 1)
RANKED = [
 ("P1", "S", "Security", "Lock Supabase: enable RLS / revoke anon read on CRM tables; stop hard-coding the publishable key (backup.html:235, mc-asset-backup.py:30). Nightly export read 7,191 contacts with it."),
 ("P1", "S", "Security", "Decisions gate: /api/decisions/answer must check the assignee/approver; stop the shared X-Exo-Token letting an LLM cron answer decisions; remove the 'Jon' fallback attribution (server.js:4974-4990, 1519-1524)."),
 ("P1", "S", "Security", "Expenses: /api/expense/decide lets any logged-in user approve/deny and triggers email (server.js:2081). Restrict to the approver."),
 ("P1", "S", "Security", "UTD students (early_adopter) can PATCH any task and read any timeline; scope writes to their project (server.js:1647, 2840, 2946)."),
 ("P1", "S", "Security", "Change-password has no current-password check (server.js:1228)."),
 ("P1", "S", "Security", "/api/search bypasses member scoping (deals, all tasks, all decisions) (server.js:2564-2596)."),
 ("P1", "M", "Security", "Exobot chat runs the full main agent for every member with no per-role tool envelope, registry entry or govlog (server.js:5358)."),
 ("P1", "M", "Resilience", "Backups: every DR copy lives on the same Mac; plaintext secrets copied into backups unencrypted. Move encrypted copies off-machine (the Studio/Mini split is the natural fix)."),
 ("P1", "S", "Governance", "build_procedure.py publishes new procedures as 'Approved' and intl_release 'Yes' by default — bypasses your approval and defaults export to open (:214-223)."),
 ("P1", "S", "Governance", "CSL denied-party auto-clear is decided by a local LLM with no human confirm (csl_adjudicate.py:96-104)."),
 ("P1", "S", "Governance", "crm_dd_engine: a model 'fail' blocks the CRM before a human confirms; route to a Decision first."),
 ("P1", "S", "Governance", "Inventory capture: LLM writes the inventory of record without human confirm (inventory_capture.py:720)."),
 ("P1", "M", "Governance", "Self-Improvement Implementer (weekly openclaw cron) edits prompts/workspace autonomously; make it propose-only via Decisions with diffs."),
 ("P1", "S", "Telemetry", "Instrument the DECIDE core via governance/instrument_fleet.py (overnight_resolver, decision_triage, decision-note-router, legacy_triage_decisions, crm_dd_engine)."),
 ("P1", "M", "Telemetry", "Register the 33 openclaw crons in agent-registry.json and mirror run status into events.jsonl."),
 ("P1", "M", "Telemetry", "Bulk-instrument the remaining ~95 silent launchd scripts, financial writers first (qb-deposit-sync, capital-bridge-funds, quote-close-stamp, payload-commission-trigger, expense-*)."),
 ("P1", "S-M", "Telemetry", "One glog() wrapper on all MC mutating routes so every human edit carries a correlation ID; add an audit trail for admin identity/role changes."),
 ("P1", "S-M", "Data integrity", "Replace UI hard-deletes (CRM, companies, deals, templates, APAC) with soft-delete + role check + audit."),
 ("P1", "S", "Data integrity", "build-dossiers.py wipes UI edits to curated dossiers; preserve UI-owned fields."),
 ("P1", "S", "Claims", "RDRE overclaim on exos-prototype.html:487 (public-site prototype) — conflicts with the not-an-RDRE-provider guardrail."),
 ("P1", "S", "Quiet drift", "Failed deliveries logged as success: research_agent, payload_sales_engine post_discord, followup_nudge_digest, payload_commission_trigger; safety disposition-task creation failures silently dropped (server.js:2160-2180)."),
 ("P2", "S", "Security", "E-sign: sender attribution + atomic envelopes.json (server.js:3951, 248); receiving 'received_by' taken from request, not session."),
 ("P2", "S", "Data integrity", "Atomic/locked writes for shared JSON stores: decisions.json (9 direct writers), expense-reports, inventory, bom-index, required-reading, envelopes."),
 ("P2", "S", "Telemetry", "task_ledger.py: emit start before exec and end/error after, so the 17+ wrapped jobs get hang detection; un-chain multi-script plists."),
 ("P2", "S", "Registry", "Fix the agent-registry generator (resolve script after ' -- ', chained plists) and fill reusability_scope (132 TBD); cross-link agents to governance classes."),
 ("P2", "S", "Ops", "HR mirror (bamboo_sync.js) isn't scheduled — directory stale since 2026-08-10. Assets/Production pages show seed data with no owner."),
 ("P3", "S", "Hygiene", "Retire crm-reminder-lizi leftovers and expired one-off reminder plists; fill the Fleet page with the launchd agents."),
]
t = doc.add_table(rows=1, cols=5); t.style = "Light Grid Accent 1"
for i, h in enumerate(["#", "Pri", "Effort", "Area", "Fix"]): t.rows[0].cells[i].text = h
for n, (p, e, a, f) in enumerate(RANKED, 1):
    c = t.add_row().cells
    c[0].text, c[1].text, c[2].text, c[3].text, c[4].text = str(n), p, e, a, f
doc.add_paragraph("Note: the audit flagged weekly-employee-report sending HR data to megawattcontrols@gmail.com. "
                  "The script lists that address as one of Jon's own aliases, so this is likely intended — confirm or drop it.")

# --- section-by-section
doc.add_heading("3. Section-by-section verdicts", 1)
for k in "ABC":
    doc.add_heading(f"Slice {k} — {TITLES[k]}", 2)
    for sec in S[k]["sections"]:
        v = verdict_of(sec)
        p = doc.add_paragraph(); run = p.add_run(f"{sec['name']} — {v}"); run.bold = True
        if v in COLOR: run.font.color.rgb = COLOR[v]
        if sec.get("page"): doc.add_paragraph(f"Page: {sec['page']}")
        hd = listify(sec.get("headers"))
        if hd: doc.add_paragraph("Headers: " + " | ".join(map(str, hd))[:1200])
        ag = listify(sec.get("agents"))
        if ag:
            names = [a.get("name", str(a)) if isinstance(a, dict) else str(a) for a in ag]
            doc.add_paragraph("Runs underneath: " + "; ".join(names)[:1200])
        for g in listify(sec.get("gaps")):
            doc.add_paragraph("Gap: " + str(g), style="List Bullet")
        for f in listify(sec.get("fixes")):
            doc.add_paragraph("Fix: " + fixtext(f), style="List Bullet")

doc.add_heading("Slice D — Background agent fleet (181 agents)", 2)
D = S["D"]
doc.add_paragraph(D["summary"]["headline"])
t = doc.add_table(rows=1, cols=4); t.style = "Light Grid Accent 1"
for i, h in enumerate(["Agent", "Kind / schedule", "Verdict", "Gaps"]): t.rows[0].cells[i].text = h
for sec in sorted(D["sections"], key=lambda s: (verdict_of(s) != "NON-COMPLIANT", s["name"])):
    c = t.add_row().cells
    c[0].text = sec["name"]; c[1].text = f"{sec.get('kind','')} · {sec.get('schedule','')}"
    c[2].text = verdict_of(sec); c[3].text = "; ".join(map(str, listify(sec.get("gaps"))))[:400]
doc.add_heading("Orphans / cleanup", 3)
for n in D["orphans"].get("notes", []): doc.add_paragraph(n, style="List Bullet")

doc.add_paragraph("\nSource data: logs/mc-exo3-audit/slice-{A,B,C,D}.json (each gap carries file:line references). "
                  "No code, data, jobs or git state were changed by the audit.")
doc.save(OUT)
print(OUT, dict(tot), len(RANKED))
