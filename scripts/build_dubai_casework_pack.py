#!/usr/bin/env python3
"""Dubai fraud — congressional casework + parallel federal referral pack (Jon @Exobot 2026-10-11).
Builds letterhead PDFs (ExoSign-ready signature lines) + one editable .docx of everything, and pre-fills
Sen. Cruz's official privacy form (ink signature required by that office — NOT ExoSign-able).
Run: python3 scripts/build_dubai_casework_pack.py"""
import os, sys
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
import docx
from docx.shared import Pt

WS = "/Users/exomac/.openclaw/workspace"
sys.path.insert(0, WS + "/scripts")
from exos_letterhead import stationery, MARGINS

OUT = WS + "/drafts/dubai_casework_2026-10-11"
DATE = "October 12, 2026"
NAVY = HexColor("#0A3D7C"); DARK = HexColor("#1A1A1A"); GREY = HexColor("#555555")
ss = getSampleStyleSheet()
ST = {
    "body": ParagraphStyle("b", parent=ss["Normal"], fontSize=10.5, textColor=DARK, alignment=TA_JUSTIFY, leading=14, spaceAfter=8),
}
ST["l"] = ParagraphStyle("l", parent=ST["body"], alignment=TA_LEFT, spaceAfter=1)
ST["g"] = ParagraphStyle("g", parent=ST["l"], fontSize=9.5, textColor=GREY)
ST["h"] = ParagraphStyle("h", parent=ss["Heading2"], fontSize=11.5, textColor=NAVY, spaceBefore=6, spaceAfter=3, leading=13)
ST["bl"] = ParagraphStyle("bl", parent=ST["body"], leftIndent=14, bulletIndent=2, spaceAfter=3)
ST["t"] = ParagraphStyle("t", parent=ss["Heading1"], fontSize=14, textColor=NAVY, spaceAfter=6)

IC3 = "7c585d207f3547893dc5fe1d0919d8c"
SIG = ["Cofounder &amp; Chief Executive Officer", "Exos Aerospace Systems &amp; Technologies, Inc.",
       "299 FM 1903, Greenville, TX 75402",
       "+1 (972) 740-8355&nbsp;&nbsp;·&nbsp;&nbsp;jquinn@exosaero.com&nbsp;&nbsp;·&nbsp;&nbsp;www.exosaero.com"]

# ── shared facts (verified against case file / sent-mail log) ────────────────
TIMELINE = [
    "<b>December 2025</b>: first call to the FBI Dallas Field Office.",
    "<b>April 6, 2026</b>: U.S. Embassy Abu Dhabi (American Citizen Services) notified.",
    f"<b>April 23, 2026</b>: IC3 complaint filed (#{IC3}); FBI Dallas called again; written congressional inquiry "
    "requests emailed to the casework teams of Senator Cornyn and Senator Cruz.",
    "<b>May 18, 2026</b>: IC3 complaint re-submitted.",
    "<b>June 23, 2026</b>: voluntary interview with two Supervisory Special Agents of the FBI Legal Attaché office in "
    "Dubai on a U.S. Department of State referral. The agents opened by acknowledging that the State-to-FBI referral "
    "had taken too long.",
    "<b>June 25, 2026</b>: written follow-up to the Legal Attaché agents.",
    "<b>As of October 1, 2026</b>: no case agent assigned. The Dallas Field Office number reaches only the general "
    "tip line, and Exos has no direct point of contact at the FBI.",
]
DURATION = "nearly ten months"
CASE = (f"Between December 2021 and April 2023, Exos was defrauded of <b>$613,097.63</b> in wires sent from U.S. "
        "JPMorgan Chase accounts, through BNY Mellon as correspondent bank, to Emirates Islamic Bank in the UAE "
        "(18 U.S.C. §§ 1343, 1956, 1962). The principal subject is Ralf Jung (alias &ldquo;Frank Kahn&rdquo;, German "
        "national, prior German financing-fraud conviction). Exos and its investors are the victims. Exos won twice in "
        "Dubai court (High Court August 2023, Case #207/2023/5263; affirmed by the Supreme Court August 2025) and filed "
        "Dubai Police criminal complaint #225004184901 (November 23, 2025). No funds have been returned.")
FOREIGN = ("Every other affected country is acting: Germany secured a connected conviction on July 27, 2026; Ireland's "
           "An Garda Síochána Economic Crime Unit confirmed on August 30, 2026 that an updated Mutual Legal Assistance "
           "request is going to the UAE; and the UAE courts have ruled for Exos twice. The United States is the only "
           "affected jurisdiction where the referral has not become an open case.")
URGENT = ("Since about August 11, 2026 the subjects have sent WhatsApp threats from UAE numbers to the lead cooperating "
          "victim. Victims report assets being moved to Germany, the United Kingdom, Turkey, Brazil and Indonesia, "
          "and into a residential property in the United States.")


def letter_doc(addr, re_lines, salutation, paras, cc, enclosures=None):
    """paras: list of (kind, text) where kind in body/h/bl."""
    return dict(addr=addr, re=re_lines, sal=salutation, paras=paras, cc=cc, encl=enclosures or [])


def timeline_block():
    return [("h", f"Exos has been asking the FBI for help for {DURATION}")] + [("bl", t) for t in TIMELINE]


# ── 1. Congressional casework requests ───────────────────────────────────────
ASK_CONG = ("I respectfully ask your office to make a congressional inquiry to the Federal Bureau of Investigation "
            "(Office of Congressional Affairs and the Dallas Field Office) asking: (1) the status of the referral "
            f"(IC3 #{IC3}, and the State Department referral handled by FBI Legal Attaché Dubai in June 2026); (2) "
            "whether it has been assigned to a case agent; and (3) for a direct point of contact so Exos can deliver "
            "its evidence. Exos has preserved the complete wire records and bank statements, the signed contracts, the "
            "forged proof-of-funds, both court judgments, the police complaint and the threat messages, and will "
            "produce them immediately.")
PARALLEL = ("In parallel, Exos is referring the matter to the U.S. Secret Service Dallas Field Office and to Homeland "
            "Security Investigations Dallas, and is filing a supplemental IC3 report with the new facts.")

CRUZ = letter_doc(
    ["The Honorable Ted Cruz", "United States Senator for Texas", "Attn: Casework",
     "961 J.J. Pickle Federal Building, 300 E. 8th Street", "Austin, TX 78701", "casework@cruz.senate.gov"],
    ["Casework Request: Status Inquiry to the FBI on an International Wire-Fraud Referral",
     "Constituent: John Quinn / Exos Aerospace Systems &amp; Technologies, Inc. (Greenville, Hunt County)",
     f"IC3 #{IC3} · Dubai Police complaint #225004184901 · Prior request to your office: April 23, 2026"],
    "Dear Senator Cruz:",
    [("body", "I am John Quinn, Cofounder and CEO of Exos Aerospace Systems &amp; Technologies, Inc. in Greenville, "
              "Texas. On April 23, 2026 I wrote to your casework team about an international fraud against Exos and its "
              "investors. I have not received a reply, and the FBI has still not assigned the matter to an agent. I am "
              "asking again, with a signed privacy release enclosed."),
     ("h", "The case"), ("body", CASE), ("body", FOREIGN), ("body", URGENT)]
    + timeline_block()
    + [("h", "Request"), ("body", ASK_CONG), ("body", PARALLEL),
       ("body", "Thank you for your help. You can reach me directly at +1 (972) 740-8355 or jquinn@exosaero.com.")],
    "cc: Robert Cass (co-victim; rob@irishsustainablehomes.com)",
    ["Signed Privacy Act release (Senator Cruz Standard Information and Privacy Act Form)"])

SELF = letter_doc(
    ["The Honorable Keith Self", "U.S. House of Representatives, Texas 3rd District", "Attn: Casework, Hunt County Office",
     "2410 Lee Street, Suite 202", "Greenville, TX 75401"],
    ["Casework Request: Status Inquiry to the FBI on an International Wire-Fraud Referral",
     "Constituent: John Quinn / Exos Aerospace Systems &amp; Technologies, Inc. (Greenville, Hunt County)",
     f"IC3 #{IC3} · Dubai Police complaint #225004184901"],
    "Dear Congressman Self:",
    [("body", "I am John Quinn, Cofounder and CEO of Exos Aerospace Systems &amp; Technologies, Inc., based in "
              "Greenville in your district. I am asking for your office's help with an international fraud against "
              "Exos and its investors. For almost ten months the FBI has not assigned the matter to an agent. A signed "
              "privacy release is enclosed."),
     ("h", "The case"), ("body", CASE), ("body", FOREIGN), ("body", URGENT)]
    + timeline_block()
    + [("h", "Request"), ("body", ASK_CONG), ("body", PARALLEL),
       ("body", "I would be glad to meet your staff at your Greenville office, which is a few minutes from our "
                "facility. You can reach me directly at +1 (972) 740-8355 or jquinn@exosaero.com.")],
    "cc: Robert Cass (co-victim; rob@irishsustainablehomes.com)",
    ["Signed Privacy Act release"])

CORNYN = letter_doc(
    ["The Honorable John Cornyn", "United States Senator for Texas", "Attn: Casework, Dallas Regional Office",
     "5001 Spring Valley Road, Suite 1125 E", "Dallas, TX 75244"],
    ["Follow-Up to My April 23, 2026 Casework Request: FBI Status Inquiry, International Wire-Fraud Referral",
     "Constituent: John Quinn / Exos Aerospace Systems &amp; Technologies, Inc. (Greenville, Hunt County)",
     f"IC3 #{IC3} · Dubai Police complaint #225004184901"],
    "Dear Senator Cornyn:",
    [("body", "On April 23, 2026 I wrote to your casework team asking for help getting the FBI to act on an "
              "international fraud against Exos Aerospace and its investors. I understand your office is not opening "
              "new cases as your term ends. This letter follows up on that April request, which your office already "
              "has, and adds a signed privacy release so your staff can make the inquiry before the end of the term."),
     ("h", "The case"), ("body", CASE), ("body", FOREIGN), ("body", URGENT)]
    + timeline_block()
    + [("h", "Request"), ("body", ASK_CONG), ("body", PARALLEL),
       ("body", "Thank you for your years of support for Texas's space and defense industry. You can reach me directly "
                "at +1 (972) 740-8355 or jquinn@exosaero.com.")],
    "cc: Robert Cass (co-victim; rob@irishsustainablehomes.com)",
    ["Signed Privacy Act release"])

# ── 2. Federal referrals ─────────────────────────────────────────────────────
ASK_REF = ("I respectfully ask that your office review this referral for investigation of the wire-fraud and "
           "money-laundering conduct and of the U.S. assets reported above, coordinate with the FBI (IC3 #{ic3}) as "
           "appropriate, and give Exos a point of contact. Exos has preserved and will produce immediately the complete "
           "wire records and bank statements, the signed contracts (English and Arabic), the forged Emirates Islamic "
           "Bank proof-of-funds, both Dubai judgments, the police complaint and the threat messages. I will come to "
           "your office or host agents at our Greenville, Texas facility.").format(ic3=IC3)
WIRES = [
    "Dec 6, 2021: $449,279.00, Chase → Infinity Group FZCO, Emirates Islamic Bank IBAN AE040340003708412206302",
    "Dec 13, 2021: $8,218.63, Chase → same account",
    "May 20, 2022: $50,000.00 (+$100 fees), Chase → Infinity Group FZCO via BNY Mellon",
    "Sep 23, 2022: $70,000.00, Chase → Infinity Group FZCO via BNY Mellon",
    "Apr 21, 2023: $35,000.00, same-day ACH → Beyond Infinity",
    "Jun 16, 2022: $500.00 (related routing). Total: $613,097.63",
]
SUBJECTS = ("Subjects: Ralf Jung (alias &ldquo;Frank Kahn&rdquo;; German passport C4KN7VH02; prior German "
            "financing-fraud conviction); Ferhat Dogan (Turkish passport U235923290); Tulgar (UAE); and the "
            "entities Aras Consultancy FZCO / The Infinity Group FZCO / Beyond Infinity / Aras Group (Dubai Marina).")


def referral(addr, agency_line, sal, why):
    return letter_doc(
        addr,
        ["Referral: International Wire Fraud and Money Laundering Through U.S. Banks (18 U.S.C. §§ 1343, 1956, 1962)",
         "Victim: Exos Aerospace Systems &amp; Technologies, Inc. (Greenville, TX) and its investors",
         f"Related: FBI IC3 #{IC3} · Dubai Police complaint #225004184901 · Dubai Case #207/2023/5263"],
        sal,
        [("body", f"I am John Quinn, Cofounder and CEO of Exos Aerospace Systems &amp; Technologies, Inc. I am referring "
                  f"to {agency_line} an international wire-fraud and money-laundering scheme that took $613,097.63 from "
                  f"Exos through U.S. banks. {why}"),
         ("h", "The scheme"), ("body", CASE)]
        + [("bl", w) for w in WIRES]
        + [("body", SUBJECTS),
           ("h", "Why now"), ("body", FOREIGN), ("body", URGENT + " A U.S. property tied to the proceeds gives a direct "
                                                  "U.S. hook for seizure or forfeiture; Exos is obtaining its details "
                                                  "and will provide them as soon as they are received."),
           ("body", f"Exos reported this to the FBI {DURATION} ago (first call December 2025; IC3 filed April 23, 2026; "
                    "interview with FBI Legal Attaché Dubai June 23, 2026). As of October 1, 2026 no case agent has "
                    "been assigned. This referral is not meant to compete with the FBI. It is meant to make sure a U.S. "
                    "agency with jurisdiction is acting while the assets can still be traced."),
           ("h", "Request"), ("body", ASK_REF),
           ("body", "Thank you. You can reach me directly at +1 (972) 740-8355 or jquinn@exosaero.com.")],
        "cc: Robert Cass (co-victim; rob@irishsustainablehomes.com)")


USSS = referral(["Special Agent in Charge", "U.S. Secret Service, Dallas Field Office", "Attn: Financial Crimes / Cyber Fraud Task Force",
                 "125 E. John Carpenter Freeway, Suite 300", "Irving, TX 75062"],
                "the U.S. Secret Service Dallas Field Office", "Dear Special Agent in Charge:",
                "The Secret Service investigates wire fraud and financial crimes against U.S. businesses and payment systems.")
HSI = referral(["Travis Pickard, Special Agent in Charge", "Homeland Security Investigations, Dallas",
                "Attn: Financial Crimes / Illicit Finance Group", "125 E. John Carpenter Freeway, Suite 800", "Irving, TX 75062"],
               "Homeland Security Investigations Dallas", "Dear SAC Pickard:",
               "HSI investigates transnational money laundering and cross-border movement of criminal proceeds.")

# ── 3. Privacy Act release (Exos-prepared; for Self + Cornyn) ────────────────
def release(member, office):
    return dict(member=member, office=office)


RELEASES = [release("Representative Keith Self (TX-03)", "Hunt County Office, 2410 Lee Street, Suite 202, Greenville, TX 75401"),
            release("Senator John Cornyn", "Dallas Regional Office, 5001 Spring Valley Road, Suite 1125 E, Dallas, TX 75244")]


# ── renderers ────────────────────────────────────────────────────────────────
def pdf_letter(path, L):
    E = []
    P = lambda t, s="body": E.append(Paragraph(t, ST[s]))
    S = lambda x=0.10: E.append(Spacer(1, x * inch))
    P(DATE, "l"); S()
    for t in L["addr"]:
        P(t, "l")
    S()
    P(f"<b><font color='#0A3D7C'>RE:&nbsp;&nbsp;</font>{L['re'][0]}</b>", "l")
    for t in L["re"][1:]:
        P("&nbsp;&nbsp;&nbsp;&nbsp;" + t, "g")
    S(); P(L["sal"], "l"); S(0.06)
    for k, t in L["paras"]:
        E.append(Paragraph(t, ST[k], bulletText="•" if k == "bl" else None))
    S(0.12); P("Respectfully,", "l"); S(0.42)
    P("______________________________", "l"); P("John Quinn", "l")
    for t in SIG:
        P(t, "l")
    S()
    if L["encl"]:
        P("Enclosure: " + "; ".join(L["encl"]), "g")
    P(L["cc"], "g")
    SimpleDocTemplate(path, pagesize=letter, **MARGINS).build(E, onFirstPage=stationery, onLaterPages=stationery)


def pdf_release(path, R):
    E = []
    P = lambda t, s="body": E.append(Paragraph(t, ST[s]))
    S = lambda x=0.10: E.append(Spacer(1, x * inch))
    P("Privacy Act Release and Authorization", "t")
    P(f"To: {R['member']}, {R['office']}", "l"); S(0.15)
    for lab, val in [("Name", "John Quinn"), ("Organization", "Exos Aerospace Systems &amp; Technologies, Inc. (Cofounder &amp; CEO)"),
                     ("Address", "299 FM 1903, Greenville, TX 75402"), ("Phone", "+1 (972) 740-8355"),
                     ("Email", "jquinn@exosaero.com"),
                     ("Federal agencies involved", "Federal Bureau of Investigation (Dallas Field Office; Legal Attaché Dubai; IC3); "
                                                   "U.S. Secret Service (Dallas); Homeland Security Investigations (Dallas); U.S. Department of State"),
                     ("Reference numbers", f"IC3 #{IC3}; Dubai Police complaint #225004184901; Dubai Case #207/2023/5263"),
                     ("Nature of inquiry", "Status of an international wire-fraud / money-laundering referral against Exos "
                                           "Aerospace and its investors ($613,097.63); whether a case agent has been assigned; "
                                           "a direct point of contact.")]:
        P(f"<b>{lab}:</b> {val}", "l"); S(0.04)
    S(0.15)
    P(f"Under the Privacy Act of 1974 (5 U.S.C. § 552a), I authorize {R['member']} and their staff to request and receive "
      "information about me and about my complaint and referral described above from the federal agencies listed, and I "
      "authorize those agencies to release that information to the Member's office for this inquiry. I authorize the office "
      "to discuss this matter with Robert Cass, a co-victim. I certify that the information I have provided is true and "
      "correct to the best of my knowledge.")
    S(0.45)
    P("______________________________&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;Date: ______________", "l")
    P("John Quinn", "l")
    SimpleDocTemplate(path, pagesize=letter, **MARGINS).build(E, onFirstPage=stationery, onLaterPages=stationery)


IC3_TEXT = [
    ("h", "How to file"),
    ("body", "IC3 does not let you edit a complaint or attach files. Its FAQ says to file a new complaint and state that it adds "
             "information to the earlier one (ic3.gov/Home/FAQ). Paste each section below into the matching part of the form at "
             "complaint.ic3.gov. Keep the originals; agents will ask for them directly. Do not enter SSN or date of birth (IC3 "
             "instructs against it)."),
    ("h", "1. Who is filing"),
    ("body", "Victim filing on behalf of a business: Exos Aerospace Systems &amp; Technologies, Inc., 299 FM 1903, Greenville, TX "
             "75402. Filer: John Quinn, Cofounder &amp; CEO, +1 (972) 740-8355, jquinn@exosaero.com."),
    ("h", "3. Financial transactions"), *[("bl", w) for w in WIRES],
    ("body", "Originating bank: JPMorgan Chase (U.S.); correspondent: BNY Mellon; beneficiary bank: Emirates Islamic Bank, UAE."),
    ("h", "4. Subjects"), ("body", SUBJECTS),
    ("h", "5. Description of incident (paste in full)"),
    ("body", f"THIS IS A NEW COMPLAINT FILED TO ADD INFORMATION TO MY PREVIOUS COMPLAINT #{IC3} (filed April 23, 2026; "
             "re-submitted May 18, 2026). The original complaint describes an advance-fee financing fraud in which Exos "
             "Aerospace wired $613,097.63 from U.S. Chase accounts through BNY Mellon to Emirates Islamic Bank (UAE) "
             "between December 2021 and April 2023. New information since the original filing:"),
    ("bl", "June 23, 2026: I gave a voluntary interview to two Supervisory Special Agents of FBI Legal Attaché Dubai on a "
           "U.S. Department of State referral. No case agent has been assigned since. As of October 1, 2026 the Dallas Field "
           "Office number reaches only the tip line."),
    ("bl", "July 27, 2026: German authorities obtained a conviction connected to the same fraud ring. Ralf Jung already has a "
           "prior German financing-fraud conviction."),
    ("bl", "August 30, 2026: Garda Billy Leydon of An Garda Síochána's Economic Crime Unit (Ireland) confirmed an updated "
           "Mutual Legal Assistance request to UAE authorities (billy.d.leydon@garda.ie)."),
    ("bl", "Since about August 11, 2026: the subjects have sent WhatsApp threats from UAE mobile numbers to the lead "
           "cooperating victim, Robert Cass (Ireland)."),
    ("bl", "A recovery-fraud scheme is targeting the same victims through mustafa66dogan11@gmail.com, asking a 30% fee. The "
           "surname matches subject Ferhat Dogan."),
    ("bl", "Co-victims report that the subjects are moving assets to Germany, the United Kingdom, Turkey, Brazil and Indonesia "
           "through front companies, and into a residential property in the United States. The details are being obtained "
           "and will be provided."),
    ("bl", "Dubai: High Court judgment for Exos (August 2023, Case #207/2023/5263), affirmed by the Supreme Court (August 2025); "
           "Dubai Police criminal complaint #225004184901 (November 23, 2025). Funds were moved out of the judgment-debtor "
           "accounts while the cases were pending, and nothing has been recovered. Victims are seeking a UAE travel ban "
           "on Ralf Jung and linking their cases to a criminal case."),
    ("bl", "Parallel referrals are being made to the U.S. Secret Service (Dallas) and Homeland Security Investigations "
           "(Dallas), and congressional casework inquiries are being made with Senator Cruz, Senator Cornyn and "
           "Representative Keith Self."),
    ("body", "Exos asks that this complaint be forwarded to the FBI Dallas Field Office and linked to the earlier complaint, "
             "and that a case agent contact Exos. Evidence available on request: wire records, bank statements, signed "
             "contracts, forged proof-of-funds, court judgments, the police complaint and the threat messages."),
    ("h", "6. Other information"),
    ("body", "Prior law enforcement contact: FBI Dallas (December 2025; April 23, 2026), FBI Legal Attaché Dubai (June 23, 2026), "
             "U.S. Embassy Abu Dhabi (April 6, 2026), Dubai Police, An Garda Síochána (Ireland), and German prosecutors in Koblenz."),
]


def pdf_ic3(path):
    E = [Paragraph("IC3 Supplemental Complaint: Text to Paste", ST["t"]),
         Paragraph(f"Adds to IC3 #{IC3} · Prepared {DATE} · John files online at complaint.ic3.gov", ST["g"]), Spacer(1, 8)]
    for k, t in IC3_TEXT:
        E.append(Paragraph(t, ST[k], bulletText="•" if k == "bl" else None))
    SimpleDocTemplate(path, pagesize=letter, **MARGINS).build(E, onFirstPage=stationery, onLaterPages=stationery)


def fill_cruz(src, dst):
    import pypdf
    r = pypdf.PdfReader(src); w = pypdf.PdfWriter(); w.append(r)
    w.update_page_form_field_values(w.pages[0], {
        "Name Mr  Mrs  Ms  Other please specify": "Mr. John Quinn (Cofounder & CEO, Exos Aerospace)",
        "Address": "299 FM 1903",
        "CityStateZip": "Greenville, TX 75402",
        "Preferred Phone": "(972) 740-8355",
        "Email": "jquinn@exosaero.com",
        "Social Security Number": "N/A (not an SSA/IRS matter)",
        "Federal Agency Involved": "FBI (Dallas Field Office / Legal Attache Dubai / IC3); USSS; HSI",
        "Requested Benefits": "Status inquiry on fraud referral, IC3 #" + IC3,
        "If yes which offices": "Sen. Cornyn; Sen. Cruz (casework email)",
        "Date4_af_date": "04/23/2026",
        "Check Box1": "/Yes", "Check Box4": "/Yes",
    }, auto_regenerate=False)
    w.set_need_appearances_writer(True)
    w.write(dst)


def docx_all(path, items):
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(10.5)
    strip = lambda t: (t.replace("<b>", "").replace("</b>", "").replace("&amp;", "&").replace("&nbsp;", " ")
                       .replace("&ldquo;", "“").replace("&rdquo;", "”").replace("<font color='#0A3D7C'>", "").replace("</font>", ""))
    for i, (title, L) in enumerate(items):
        if i:
            d.add_page_break()
        d.add_heading(title, level=1)
        if L == "IC3":
            for k, t in IC3_TEXT:
                d.add_heading(strip(t), level=2) if k == "h" else d.add_paragraph(strip(t), style="List Bullet" if k == "bl" else None)
            continue
        d.add_paragraph(DATE)
        d.add_paragraph("\n".join(strip(a) for a in L["addr"]))
        d.add_paragraph("RE: " + "\n".join(strip(x) for x in L["re"]))
        d.add_paragraph(L["sal"])
        for k, t in L["paras"]:
            d.add_heading(strip(t), level=2) if k == "h" else d.add_paragraph(strip(t), style="List Bullet" if k == "bl" else None)
        d.add_paragraph("Respectfully,\n\n\nJohn Quinn\n" + "\n".join(strip(s) for s in SIG))
        if L["encl"]:
            d.add_paragraph("Enclosure: " + "; ".join(L["encl"]))
        d.add_paragraph(L["cc"])
    d.save(path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    files = {
        "Exos_Casework_Request_Sen_Cruz_2026-10-12.pdf": CRUZ,
        "Exos_Casework_Request_Rep_Self_2026-10-12.pdf": SELF,
        "Exos_Casework_FollowUp_Sen_Cornyn_2026-10-12.pdf": CORNYN,
        "Exos_Referral_USSS_Dallas_2026-10-12.pdf": USSS,
        "Exos_Referral_HSI_Dallas_2026-10-12.pdf": HSI,
    }
    for f, L in files.items():
        pdf_letter(f"{OUT}/{f}", L); print("WROTE", f)
    pdf_release(f"{OUT}/Exos_Privacy_Release_Rep_Self_2026-10-12.pdf", RELEASES[0])
    pdf_release(f"{OUT}/Exos_Privacy_Release_Sen_Cornyn_2026-10-12.pdf", RELEASES[1])
    pdf_ic3(f"{OUT}/Exos_IC3_Supplement_Text_2026-10-12.pdf")
    fill_cruz(f"{OUT}/cruz_privacy_form.pdf", f"{OUT}/Cruz_Privacy_Form_PREFILLED_InkSign_2026-10-12.pdf")
    docx_all(f"{OUT}/Exos_Dubai_Fraud_Casework_and_Referrals_EDITABLE_2026-10-12.docx",
             [("Sen. Cruz: casework request", CRUZ), ("Rep. Keith Self: casework request", SELF),
              ("Sen. Cornyn: follow-up", CORNYN), ("U.S. Secret Service Dallas: referral", USSS),
              ("HSI Dallas: referral", HSI), ("IC3 supplemental complaint text", "IC3")])
    print("done")
