#!/usr/bin/env python3
"""Editable Word (PRIMARY letterhead) for the Exos VSFB MDZ whitepaper.
Run: PYTHONNOUSERSITE=1 python3.12 scripts/build_vsfb_mdz_docx.py
"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import vsfb_mdz_content as K

LH = "assets/EXOS_Letterhead_Greenville_TX_PRIMARY.docx"
OUT = "drafts/VSFB_MDZ/Exos_VSFB_MDZ_Whitepaper_%s%s.docx" % (K.DATE, K.VER)
NAVY = RGBColor(0x0A, 0x3D, 0x7C); DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55); RED = RGBColor(0x9A, 0x1B, 0x1B)

doc = Document(LH)
# strip any template body paragraphs (keep sectPr)
body = doc.element.body
for el in list(body):
    if el.tag.endswith('}p') or el.tag.endswith('}tbl'):
        body.remove(el)

def runs(p, text, size=10.5, color=DARK, italic=False, bold_all=False):
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part: continue
        b = part.startswith("**") and part.endswith("**")
        r = p.add_run(part[2:-2] if b else part)
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Calibri"
        r.font.bold = b or bold_all; r.font.italic = italic
    return p

def para(text, size=10.5, color=DARK, align=None, before=0, after=5, italic=False, bold=False, indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent is not None: pf.left_index = None; pf.left_indent = Inches(indent)
    if align: p.alignment = align
    return runs(p, text, size, color, italic, bold)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)

def borders(tbl, color="9AA5B1"):
    tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), color); b.append(x)
    tblPr.append(b)

def table(rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); borders(t)
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Inches(widths[j])
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            runs(c.paragraphs[0], val, 9, RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else DARK, bold_all=(i == 0))
            if i == 0: shade(c, "0A3D7C")
            elif "[JON" in val or "[ENG" in val or "[PLANNING" in val: shade(c, "FFF4CC")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def decision(text):
    t = doc.add_table(rows=1, cols=1); borders(t, "C9A227"); c = t.cell(0, 0); c.width = Inches(6.5); shade(c, "FFF4CC")
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    runs(c.paragraphs[0], text, 9.5, RED, bold_all=False)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# DRAFT marking in header (keeps letterhead image)
for s in doc.sections:
    hp = s.header.add_paragraph(); hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(K.MARK + "  ·  VSFBMDZCall"); r.font.size = Pt(8); r.font.color.rgb = RED; r.font.bold = True

for blk in K.C:
    k = blk[0]
    if k == "title": para(blk[1], 20, NAVY, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4, bold=True)
    elif k == "title2": para(blk[1], 15, NAVY, WD_ALIGN_PARAGRAPH.CENTER, before=10, after=4, bold=True)
    elif k == "sub": para(blk[1], 11, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=3, italic=True)
    elif k == "cover": para(blk[1], 10.5, DARK, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    elif k == "spacer": para("", after=10)
    elif k == "note": para(blk[1], 8.5, GREY, after=6, italic=True)
    elif k == "h1": para(blk[1], 13, NAVY, before=12, after=4, bold=True).paragraph_format.keep_with_next = True
    elif k == "h2": para(blk[1], 11, NAVY, before=7, after=3, bold=True).paragraph_format.keep_with_next = True
    elif k == "p": para(blk[1], 10.5, DARK, WD_ALIGN_PARAGRAPH.JUSTIFY, after=5)
    elif k == "bl":
        for b in blk[1]: para("•  " + b, 10.5, DARK, after=3, indent=0.25)
    elif k == "src": para(blk[1], 8.5, DARK, after=2)
    elif k == "table": table(blk[1], blk[2])
    elif k == "decision": decision(blk[1])
    elif k == "pb": doc.add_page_break()
    elif k == "img":
        from docx.shared import Inches as _In
        doc.add_picture(blk[1], width=_In(6.5))
        if len(blk)>2: para(blk[2], 8.5, GREY, after=6, italic=True)

doc.core_properties.title = "Exos Aerospace — VSFBMDZCall Whitepaper"
doc.core_properties.author = "Exos Aerospace Systems & Technologies, Inc."
doc.save(OUT); print("wrote", OUT)
