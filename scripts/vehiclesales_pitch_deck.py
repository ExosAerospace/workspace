#!/usr/bin/python3
"""Per-company PowerPoint pitch generator for the BLK3 Vehicle Sales-exo3 program.

Usage:
  /usr/bin/python3 scripts/vehiclesales_pitch_deck.py <targets.json> [--company NAME | --all]
        [--out drafts/vehiclesales/decks] [--no-pdf] [--force]

targets.json: a list of target records (or {"targets": [...]}) with fields:
  company, country, website, scope ("vehicle"|"recovery"|"both"), vehicle_and_status,
  recovery_status (optional), funding_total, latest_round, funds_availability, gov_awards,
  key_contacts, decision_maker, pain_point, pitch_angle, fit_score, priority, sources.

Output: Exos_Pitch_<Company>_<YYYY-MM-DD>.pptx (+ .pdf if soffice completes in time).

Guardrails (baked into the fixed copy and linted over every generated string; a lint hit
blocks the deck unless --force):
  - BLK3 has NOT flown. Never say it has.
  - No "only / first / fastest" superlatives; no "one of only three ..." phrase.
  - No ROFR. Exos is not a propulsion/engine provider. No free flights.
  - No internal cost figures; no export-controlled technical detail (recovery = capability level).
  - Foreign targets get "Subject to U.S. export control review." footer on every slide.
Nothing is sent anywhere; this only writes local files.
"""
import argparse, datetime, json, os, re, shutil, subprocess, sys, tempfile

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

WS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO = os.path.join(WS, "drafts", "_exos_white_official.png")
HERO = os.path.join(WS, "drafts", "_exos_hero_bg.png")
FALLBACK_LOGO = os.path.join(WS, "mission-control", "brand-assets", "exos-logo.png")

# Exos brand: black/grey base, single orange accent (reference_exos_brand_palette)
BG = RGBColor(0x0D, 0x0D, 0x0F)
PANEL = RGBColor(0x1C, 0x1C, 0x20)
ORANGE = RGBColor(0xF2, 0x6A, 0x21)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0xD2, 0xD2, 0xD7)
MUTE = RGBColor(0x9A, 0x9A, 0xA1)
FONT = "Arial"

SW, SH = 13.333, 7.5  # 16:9 inches

US_NAMES = {"us", "usa", "u.s.", "u.s.a.", "united states", "united states of america", "america"}

HEADLINES = {
    "vehicle": "Fly sooner: a licensed reusable rocket + a market",
    "recovery": "Don't lose your first vehicles: recover every test",
    "both": "Fly sooner and recover every test: a licensed reusable rocket, a market, and recovery",
}

CONTACT = {
    "name": "John Quinn",
    "title": "Cofounder & CEO, Exos Aerospace Systems & Technologies",
    "email": "jquinn@exosaero.com",
    "phone": "+1-972-740-8355",
    "web": "www.exosaero.com",
}

LINT = [
    (r"\bBLK\s*3\b[^.]{0,40}\b(has|have|had)\s+(flown|flew|launched)", "claims BLK3 has flown"),
    (r"\bBLK\s*3\b[^.]{0,20}\b(flew|flight[- ]proven|flight[- ]heritage)", "claims BLK3 has flown"),
    (r"\bone of only\b", "'one of only' superlative"),
    (r"\b(the|world'?s|america'?s|nation'?s|only)\s+(only|first|fastest)\b", "only/first/fastest superlative"),
    (r"\bfirst[- ]ever\b|\bfastest\b", "first/fastest superlative"),
    (r"\bonly\s+(company|operator|provider|licensed|reusable|vehicle|one)\b", "'only' superlative"),
    (r"\bROFR\b|right of first (refusal|offer)", "ROFR"),
    (r"\bfree\s+(flight|flights|launch|launches|ride|rides)\b|\bat no (cost|charge)\b", "free flight"),
    (r"\b(engine|propulsion|RDRE)\s+(provider|supplier|sale|sales)\b", "propulsion/engine provider claim"),
]
# Phrases that are explicitly negated in our own copy and must not trigger the lint.
LINT_ALLOW = [r"not (an? )?(engine|propulsion)", r"no (engine|propulsion)"]


def lint(text):
    hits = []
    for pat, why in LINT:
        for m in re.finditer(pat, text, re.I):
            ctx = text[max(0, m.start() - 25): m.end() + 5]
            if any(re.search(a, ctx, re.I) for a in LINT_ALLOW):
                continue
            hits.append(f"{why}: ...{text[max(0, m.start()-30):m.end()+30]}...")
    return hits


def s(v, default=""):
    """Normalize a field (str/list/dict/number) to a display string."""
    if v is None:
        return default
    if isinstance(v, (list, tuple)):
        return "; ".join(s(x) for x in v if s(x))
    if isinstance(v, dict):
        return "; ".join(f"{k}: {s(x)}" for k, x in v.items())
    out = str(v).strip()
    # Customer-facing text: drop citation markers like [3] / [S12] / [3][4] (sources go in the footnote).
    out = re.sub(r"\s*\[(?:S?\d+(?:,\s*S?\d+)*)\]", "", out).strip()
    return out or default


def is_foreign(rec):
    c = s(rec.get("country")).lower()
    return bool(c) and c not in US_NAMES


def safe_name(name):
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_") or "Company"


# ---------------------------------------------------------------- layout helpers
def fit_pt(text, w_in, h_in, start, minimum=10, line_h=1.22):
    """Shrink font until the estimated wrapped height fits the box."""
    size = start
    paras = text.split("\n")
    while size > minimum:
        cpl = max(8, int(w_in * 72 / (size * 0.53)))  # chars per line, Arial avg width
        lines = sum(max(1, -(-len(p) // cpl)) for p in paras)
        if lines * size * line_h / 72 + 0.1 * len(paras) * size / 18 <= h_in:
            return size
        size -= 1
    return minimum


def bg(slide, color=BG):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def box(slide, x, y, w, h, text="", size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, fit=True, minimum=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.04))
    if fit and text:
        size = fit_pt(text, w - 0.1, h - 0.08, size, minimum)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        f = r.font
        f.size, f.bold, f.name = Pt(size), bold, FONT
        f.color.rgb = color
    return tb


def rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def bullets(slide, x, y, w, h, items, size=18, minimum=11, color=GREY, gap=6):
    """Orange square bullets; one paragraph per item, auto-shrunk to fit."""
    items = [i for i in items if i]
    joined = "\n".join("   " + i for i in items)
    size = fit_pt(joined, w - 0.15, h - 0.1, size, minimum, line_h=1.3)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r0 = p.add_run()
        r0.text = "■  "
        r0.font.size, r0.font.name = Pt(size * 0.7), FONT
        r0.font.color.rgb = ORANGE
        # allow "Lead: rest" with bold white lead-in
        if ": " in item and len(item.split(": ", 1)[0]) <= 40:
            lead, rest = item.split(": ", 1)
            r1 = p.add_run()
            r1.text = lead + ": "
            r1.font.size, r1.font.bold, r1.font.name = Pt(size), True, FONT
            r1.font.color.rgb = WHITE
            item = rest
        r = p.add_run()
        r.text = item
        r.font.size, r.font.name = Pt(size), FONT
        r.font.color.rgb = color
    return tb


def logo(slide, x=SW - 2.05, y=0.32, w=1.6):
    path = LOGO if os.path.exists(LOGO) else (FALLBACK_LOGO if os.path.exists(FALLBACK_LOGO) else None)
    if path:
        slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    else:
        box(slide, x, y, w, 0.5, "EXOS AEROSPACE", 16, WHITE, True, fit=False)


def chrome(slide, rec, n, title, kicker=None):
    """Standard content-slide frame: title, orange rule, logo, footer."""
    bg(slide)
    if kicker:
        box(slide, 0.6, 0.35, 9.5, 0.35, kicker.upper(), 12, ORANGE, True, fit=False)
    box(slide, 0.6, 0.68, 10.3, 0.85, title, 32, WHITE, True, minimum=20)
    rect(slide, 0.6, 1.58, 1.4, 0.06, ORANGE)
    logo(slide)
    footer(slide, rec, n)


def footer(slide, rec, n, footnote=None):
    y = SH - 0.45
    rect(slide, 0.6, y - 0.08, SW - 1.2, 0.01, RGBColor(0x3A, 0x3A, 0x40))
    left = "Exos Aerospace  |  Proprietary & confidential  |  Discussion draft"
    if is_foreign(rec):
        left += "  |  Subject to U.S. export control review."
    box(slide, 0.6, y, 10.8, 0.3, left, 9, MUTE, fit=False)
    box(slide, SW - 1.6, y, 1.0, 0.3, str(n), 9, MUTE, align=PP_ALIGN.RIGHT, fit=False)
    if footnote:
        box(slide, 0.6, y - 0.62, SW - 1.2, 0.5, footnote, 9, MUTE, minimum=7)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- slides
def build(rec, out_dir, date):
    company = s(rec.get("company"), "Company")
    scope = s(rec.get("scope"), "vehicle").lower()
    if scope not in HEADLINES:
        raise ValueError(f"{company}: scope must be vehicle|recovery|both, got {scope!r}")
    want_vehicle = scope in ("vehicle", "both")
    want_recovery = scope in ("recovery", "both")
    foreign = is_foreign(rec)
    src = rec.get("sources") or []
    src_txt = s(src)
    texts = []  # every visible/notes string, for lint

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]

    def T(t):
        texts.append(t)
        return t

    # 1. Title -----------------------------------------------------------
    sl = prs.slides.add_slide(blank)
    bg(sl)
    if os.path.exists(HERO):
        pic = sl.shapes.add_picture(HERO, 0, 0, width=prs.slide_width, height=prs.slide_height)
        # dark scrim for legibility
        scrim = rect(sl, 0, 0, SW, SH, BG)
        scrim.fill.fore_color.rgb = BG
        _set_alpha(scrim, 45)
    logo(sl, x=0.7, y=0.6, w=2.3)
    box(sl, 0.7, 2.35, 11.9, 1.3, T(f"{company} × Exos Aerospace"), 44, WHITE, True, minimum=28)
    rect(sl, 0.7, 3.75, 1.6, 0.07, ORANGE)
    box(sl, 0.7, 3.95, 11.5, 1.3, T(HEADLINES[scope]), 26, GREY, minimum=18)
    box(sl, 0.7, 5.6, 11.0, 0.4, T(f"BLK3 Vehicle Sales program  |  Discussion draft  |  {date}"), 14, MUTE, fit=False)
    footer(sl, rec, 1)
    notes(sl, T(
        f"Open with the ask: we want to help {company} get flying sooner"
        + (" and protect early hardware through recovery" if want_recovery else "")
        + ". Exos is an FAA-licensed reusable launch operator. BLK3 is our next vehicle; its first flight "
        "campaign is planned for early 2027 (it has not flown yet; say so plainly if asked). "
        "This is a discussion draft, not an offer."))

    # 2. Your challenge --------------------------------------------------
    sl = prs.slides.add_slide(blank)
    chrome(sl, rec, 2, T("Your challenge"), kicker=company)
    pain = T(s(rec.get("pain_point"), "Getting to flight on schedule while protecting capital."))
    status = T(s(rec.get("vehicle_and_status"), ""))
    rect(sl, 0.6, 2.0, 7.3, 3.9, PANEL)
    box(sl, 0.85, 2.15, 6.8, 0.4, "WHAT WE HEARD", 12, ORANGE, True, fit=False)
    box(sl, 0.85, 2.6, 6.8, 3.15, pain, 22, WHITE, minimum=13)
    rect(sl, 8.2, 2.0, 4.53, 3.9, PANEL)
    box(sl, 8.45, 2.15, 4.0, 0.4, "VEHICLE & STATUS", 12, ORANGE, True, fit=False)
    stat = status
    if want_recovery and rec.get("recovery_status"):
        stat += "\nRecovery: " + T(s(rec.get("recovery_status")))
    box(sl, 8.45, 2.6, 4.05, 3.15, stat, 17, GREY, minimum=11)
    footer(sl, rec, 2, footnote=("Sources: " + src_txt) if src_txt else None)
    notes(sl, T(
        f"Reflect their situation back in their words before pitching. Confirm the pain point: {pain} "
        "Ask what a slip of another 6-12 months costs them (customers, investors, milestone payments). "
        "Use only the public sources in the footnote; do not speculate about their finances."))

    # 3. Our purpose -----------------------------------------------------
    sl = prs.slides.add_slide(blank)
    chrome(sl, rec, 3, T("Our purpose: Make Every Space Vehicle Reusable"), kicker="Exos Aerospace")
    cols = [
        ("MASSIVE TRANSFORMATIVE PURPOSE", "Make Every Space Vehicle Reusable. Reuse turns a launch from a one-shot event into a repeatable, affordable test and flight cadence."),
        ("INTEGRATED PRESSURIZED STRUCTURES", "Our core competency: tanks and primary structure designed and built as one integrated, pressurized airframe."),
        ("DESIGNED FOR SIMPLICITY", "BLK3 is built with fewer than 5,000 components, which keeps build, inspection and turnaround manageable."),
    ]
    _cards(sl, cols, T)
    notes(sl, T(
        "Exos exists to make every space vehicle reusable. Our core competency is integrated pressurized "
        "structures. BLK3 is built with fewer than 5,000 components. Keep this at capability level; "
        "do not discuss design detail, materials processes or specifications (export control)."))

    # 4. The offer -------------------------------------------------------
    sl = prs.slides.add_slide(blank)
    chrome(sl, rec, 4, T("The offer"), kicker={"vehicle": "Vehicle", "recovery": "Recovery", "both": "Vehicle + recovery"}[scope])
    if want_vehicle and want_recovery:
        _offer_two_col(sl, T)
    elif want_vehicle:
        _cards(sl, VEHICLE_OPTIONS, T, top=2.0, h=3.6, body=15)
        box(sl, 0.6, 5.85, 12.1, 0.45, T("Every option is structured so Exos remains the FAA licensee and operator unless the FAA approves otherwise."), 13, MUTE, minimum=10)
    else:
        _cards(sl, RECOVERY_STEPS_CARDS, T, top=2.0, h=3.6, body=15)
        box(sl, 0.6, 5.85, 12.1, 0.45, T(RECOVERY_SYSTEM_LINE), 13, MUTE, minimum=10)
    footer(sl, rec, 4)
    n = []
    if want_vehicle:
        n.append("Vehicle: walk the four structures (outright sale of BLK3 structures; white-label under the Exos "
                 "license and operations; dedicated wet lease; licensed build). Let them pick the one that fits their "
                 "balance sheet. Exos does not sell engines or propulsion. Do not quote internal costs.")
    if want_recovery:
        n.append("Recovery: Exos recovery system (drogue/ballute, guided parafoil, GPS tracking) at capability level "
                 "only, no specs. Path is integration study, then ground test, then flight.")
    n.append("Pricing is quoted per program; no flights are offered free. Keep it to options, not commitments.")
    notes(sl, T(" ".join(n)))

    # 5. Why Exos --------------------------------------------------------
    sl = prs.slides.add_slide(blank)
    chrome(sl, rec, 5, T("Why Exos"), kicker="Credentials")
    why = [
        "FAA-licensed reusable launch operator: 14 CFR §431 license renewed March 13, 2026.",
        "SARGE flight record: our SARGE reusable suborbital rocket flew 4 times (2018–2019).",
        "BLK3: ~135 kg payload class; first flight campaign planned for early 2027.",
        "Get it back: we recover the vehicle and the payload, so hardware and data come home.",
        "A market on day one: access to Exos payload customers; BLK3 public price is $1M per flight.",
    ]
    bullets(sl, 0.6, 2.0, 12.1, 4.0, [T(x) for x in why], size=20, minimum=13, gap=10)
    notes(sl, T(
        "Credentials, stated carefully: FAA license under 14 CFR 431, renewed March 13, 2026. SARGE flew four times "
        "in 2018-2019. BLK3 has not flown; its first flight campaign is planned for early 2027. The $1M per flight "
        "BLK3 figure is our public price. Avoid superlatives and uniqueness claims entirely."))

    # 6. What it means for <Company> ------------------------------------
    sl = prs.slides.add_slide(blank)
    chrome(sl, rec, 6, T(f"What it means for {company}"), kicker="Fit")
    items = []
    if rec.get("funds_availability") or rec.get("funding_total") or rec.get("latest_round"):
        # Never show our internal funds-availability rating to the prospect; public round facts only.
        fund = s(rec.get("latest_round")) or s(rec.get("funding_total"))
        if fund:
            items.append(T("Your funding: " + fund + ". An Exos program can be phased so spend tracks your milestones."))
    if rec.get("gov_awards"):
        items.append(T("Milestones: " + s(rec.get("gov_awards")) + ". Flight data from a licensed operator can support milestone reviews (subject to sponsor approval)."))
    if rec.get("pitch_angle"):
        items.append(T("Our proposal: " + s(rec.get("pitch_angle"))))
    if want_vehicle:
        items.append(T("Revenue sooner: fly customers under an established FAA license while your own vehicle matures."))
    if want_recovery:
        items.append(T("Protect early hardware: recover test articles and data instead of losing them."))
    bullets(sl, 0.6, 2.0, 12.1, 3.6, items, size=19, minimum=12, gap=9)
    box(sl, 0.6, 5.75, 12.1, 0.45, T("BLK3 has not yet flown; its first flight campaign is planned for early 2027. BLK3-dependent items are contingent on that campaign."), 12, MUTE, minimum=9)
    footer(sl, rec, 6)
    notes(sl, T(
        f"Tie the offer to {company}'s own milestones and capital position. Decision maker: "
        f"{s(rec.get('decision_maker'), 'unverified')}. Key contacts: {s(rec.get('key_contacts'), 'unverified')}. "
        "Do not promise that Exos flights satisfy any government milestone; that is the sponsor's call. "
        "Any BLK3-dependent commitment is contingent on BLK3's first flight campaign (planned early 2027)."))

    # 7. Path forward ----------------------------------------------------
    sl = prs.slides.add_slide(blank)
    chrome(sl, rec, 7, T("Path forward"), kicker="Next steps")
    steps = [
        ("1  DISCOVERY CALL", "Confirm goals, timeline, budget window and the right structure. Mutual NDA if needed."),
        ("2  SCOPED STUDY / LOI", "Paid scoped study or LOI, contingent on BLK3 flight milestones. Defines scope, schedule and price."),
        ("3  PROGRAM", "Definitive agreement and execution: " + ("vehicle program" if scope == "vehicle" else "recovery program" if scope == "recovery" else "vehicle and recovery program") + " with milestone payments."),
    ]
    _cards(sl, steps, T, top=2.0, h=3.3, body=16, arrows=True)
    if foreign:
        box(sl, 0.6, 5.6, 12.1, 0.5, T("For non-U.S. parties, every step is subject to U.S. export control review and any required authorizations before technical data is shared."), 13, MUTE, minimum=10)
    notes(sl, T(
        "Close for a discovery call date. Step 2 is an LOI or paid study contingent on BLK3 flight; flights are always paid and "
        "no exclusive or preferential rights are granted. Step 3 is the definitive program."
        + (" Foreign party: no technical data until export review (Empowered Official) clears it." if foreign else "")))

    # 8. Contact ---------------------------------------------------------
    sl = prs.slides.add_slide(blank)
    bg(sl)
    if os.path.exists(HERO):
        sl.shapes.add_picture(HERO, 0, 0, width=prs.slide_width, height=prs.slide_height)
        sc = rect(sl, 0, 0, SW, SH, BG)
        _set_alpha(sc, 45)
    logo(sl, x=0.7, y=0.6, w=2.3)
    box(sl, 0.7, 2.2, 11, 0.9, T("Let's talk"), 40, WHITE, True, fit=False)
    rect(sl, 0.7, 3.15, 1.6, 0.07, ORANGE)
    box(sl, 0.7, 3.4, 11, 2.2, T(f"{CONTACT['name']}\n{CONTACT['title']}\n{CONTACT['email']}  |  {CONTACT['phone']}\n{CONTACT['web']}"), 22, GREY, minimum=16)
    footer(sl, rec, 8)
    notes(sl, T("Leave with a concrete next step and date. Follow-up comes from John Quinn."))

    # lint ---------------------------------------------------------------
    problems = [h for t in texts for h in lint(t)]
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"Exos_Pitch_{safe_name(company)}_{date}.pptx")
    return prs, path, problems


VEHICLE_OPTIONS = [
    ("SALE", "Purchase BLK3 vehicle structures for your program."),
    ("WHITE-LABEL", "Your brand and customers; flown under the Exos FAA license and operations."),
    ("DEDICATED WET LEASE", "A BLK3 dedicated to your missions, crewed and operated by Exos."),
    ("LICENSED BUILD", "Build BLK3-derived vehicles under an Exos license and support agreement."),
]
RECOVERY_STEPS_CARDS = [
    ("RECOVERY SYSTEM", "Drogue/ballute deceleration, guided parafoil and GPS tracking, adapted to your vehicle."),
    ("1  INTEGRATION STUDY", "Fit, interfaces and concept of operations for your vehicle."),
    ("2  GROUND TEST", "Prove deployment and integration on the ground."),
    ("3  FLIGHT", "Recover your test vehicles and their data."),
]
RECOVERY_SYSTEM_LINE = "Capability-level summary only; technical details follow export review and an NDA."


def _offer_two_col(sl, T):
    rect(sl, 0.6, 2.0, 5.95, 3.95, PANEL)
    rect(sl, 6.78, 2.0, 5.95, 3.95, PANEL)
    box(sl, 0.85, 2.12, 5.5, 0.4, T("VEHICLE: FOUR WAYS TO FLY"), 13, ORANGE, True, fit=False)
    bullets(sl, 0.85, 2.55, 5.5, 3.3, [T(f"{a.title()}: {b}") for a, b in VEHICLE_OPTIONS], size=15, minimum=11, gap=6)
    box(sl, 7.03, 2.12, 5.5, 0.4, T("RECOVERY: GET YOUR TEST VEHICLES BACK"), 13, ORANGE, True, fit=False)
    rec_items = [
        "Exos recovery system: drogue/ballute, guided parafoil, GPS tracking.",
        "Integration study: fit, interfaces, concept of operations.",
        "Ground test: prove deployment and integration.",
        "Flight: recover test vehicles and data.",
    ]
    bullets(sl, 7.03, 2.55, 5.5, 3.3, [T(x) for x in rec_items], size=15, minimum=11, gap=6)
    box(sl, 0.6, 6.05, 12.1, 0.4, T("Exos remains FAA licensee/operator. " + RECOVERY_SYSTEM_LINE), 12, MUTE, minimum=9)


def _cards(sl, cards, T, top=2.0, h=3.6, body=17, arrows=False):
    n = len(cards)
    gap = 0.3
    w = (SW - 1.2 - gap * (n - 1)) / n
    for i, (head, txt) in enumerate(cards):
        x = 0.6 + i * (w + gap)
        rect(sl, x, top, w, h, PANEL)
        rect(sl, x, top, w, 0.06, ORANGE)
        box(sl, x + 0.22, top + 0.25, w - 0.44, 0.75, T(head), 14 if n > 3 else 15, ORANGE, True, minimum=10)
        box(sl, x + 0.22, top + 1.05, w - 0.44, h - 1.25, T(txt), body, GREY, minimum=11)
        if arrows and i < n - 1:
            box(sl, x + w - 0.05, top + h / 2 - 0.3, gap + 0.1, 0.6, "›", 28, ORANGE, True, align=PP_ALIGN.CENTER, fit=False)


def _set_alpha(shape, pct_opaque):
    """Make a solid-filled autoshape semi-transparent (pct_opaque 0-100)."""
    from pptx.oxml.ns import qn
    from lxml import etree
    srgb = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(pct_opaque * 1000)))


def to_pdf(pptx_path, timeout=120):
    soffice = shutil.which("soffice") or "/opt/homebrew/bin/soffice"
    if not os.path.exists(soffice):
        return None
    prof = tempfile.mkdtemp(prefix="lo_prof_")
    cmd = [soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf",
           "--outdir", os.path.dirname(pptx_path), pptx_path]
    try:
        p = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
        p.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(p.pid, 9)
        except Exception:
            pass
        return None
    finally:
        shutil.rmtree(prof, ignore_errors=True)
    pdf = os.path.splitext(pptx_path)[0] + ".pdf"
    return pdf if os.path.exists(pdf) else None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("targets")
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--company")
    g.add_argument("--all", action="store_true")
    ap.add_argument("--out", default=os.path.join(WS, "drafts", "vehiclesales", "decks"))
    ap.add_argument("--date", default=datetime.date.today().isoformat())
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--force", action="store_true", help="write deck even if guardrail lint finds issues")
    a = ap.parse_args()

    data = json.load(open(a.targets))
    recs = data.get("targets", data) if isinstance(data, dict) else data
    if a.company:
        recs = [r for r in recs if s(r.get("company")).lower() == a.company.lower()]
        if not recs:
            sys.exit(f"company {a.company!r} not found in {a.targets}")
    elif not a.all:
        sys.exit("specify --company NAME or --all")

    rc = 0
    for r in recs:
        prs, path, problems = build(r, a.out, a.date)
        if problems:
            print(f"[GUARDRAIL] {s(r.get('company'))}:", *problems, sep="\n  - ", file=sys.stderr)
            if not a.force:
                print("  -> deck NOT written (fix the input record or use --force)", file=sys.stderr)
                rc = 2
                continue
        prs.save(path)
        print("PPTX", path)
        if not a.no_pdf:
            pdf = to_pdf(path)
            print("PDF ", pdf if pdf else "(skipped: soffice failed or timed out)")
    sys.exit(rc)


if __name__ == "__main__":
    main()
