#!/usr/bin/env python3
"""HVRR BOM supplier DD coverage (Jon 2026-10-22: "one DD cleared supplier for every part").

HVRR = contract FA9300-25-C-6022 (USAF/AFRL) -> defense -> supplier DD applies.
  names            write the normalized supplier list that still needs a defense screen
                   (drafts/hvrr-supplier-dd/screen_names.txt) and mark every HVRR supplier defense_use
  report           build drafts/hvrr-supplier-dd/HVRR_Supplier_DD_Coverage_<date>.xlsx:
                   every BOM line -> supplier -> DD status -> covered? + the action list
Screening itself: python3 scripts/supplier_dd_exo3.py --names-file <names> --defense --program HVRR --no-cards
"""
import json, os, re, sys, datetime
from collections import defaultdict
WS = "/Users/exomac/.openclaw/workspace"
BOM = os.path.join(WS, "mission-control/data/bom-index-hvrr.json")
REG = os.path.join(WS, "mission-control/data/supplier-dd-registry.json")
OUT = os.path.join(WS, "drafts/hvrr-supplier-dd")
CANON = [("mcmaster", "McMaster-Carr"), ("digikey", "Digi-Key"), ("pasterknack", "Pasternack"), ("pasternak", "Pasternack"),
         ("pasternack", "Pasternack"), ("mouser", "Mouser Electronics"), ("radwell", "Radwell International"),
         ("swag", "Swagelok"), ("amazon", "Amazon"), ("summitracing", "Summit Racing"), ("aircraftspruce", "Aircraft Spruce"),
         ("morganewton", "Morgan Newton"), ("morgannewton", "Morgan Newton"), ("peigen", "PEI-Genesis"), ("hobbytown", "HobbyTown"),
         ("httpsparkplugsmorrisonandmarvin", "Morrison & Marvin"), ("wwwshireeninc", "Shireen Inc"), ("rcbattert", "RC Battery"),
         ("scorpius", "Scorpius Space Launch Company"), ("sslc", "Scorpius Space Launch Company"),
         ("exosstrongenterprises", "Strong Enterprises"), ("militaryfastners", "Military Fasteners"),
         ("desgnatronics", "Designatronics Inc."), ("arrow", "Arrow Electronics"), ("msc", "MSC Industrial Supply"),
         ("midwesthose", "Midwest Hose & Specialty"), ("chemiquip", "Chemiquip Products"), ("ratermann", "Ratermann Manufacturing"),
         ("ebay", "eBay")]
IN_HOUSE = {"exos"}
MARKETPLACE = {"Amazon", "eBay"}   # no single seller to diligence -> name a distributor for defense buys
UNKNOWN = {"", "???"}

def slug(s): return "".join(c.lower() if c.isalnum() else "-" for c in (s or "")).strip("-")[:60]
def key(s): return re.sub(r"(inc|llc|corp|corporation|co|ltd|company)$", "", re.sub(r"[^a-z0-9]", "", (s or "").lower()))
def reg_index(reg):
    """registry lookup by letters-only name, preferring a GREEN/cleared record ("DigiKey" == "Digi-Key")."""
    idx = {}
    for x in reg["suppliers"]:
        k = key(x.get("name")); cur = idx.get(k)
        if not cur or (x.get("tier") == "GREEN" and cur.get("tier") != "GREEN"): idx[k] = x
    return idx

def canon(v):
    v = (v or "").strip().strip("`").strip()
    k = re.sub(r"[^a-z]", "", v.lower())
    if k in IN_HOUSE: return "EXOS (in-house)"
    for pre, c in CANON:
        if k.startswith(pre): return c
    return v

def alternatives(v):
    c = canon(v)
    if " or " in c: return [canon(x) for x in c.split(" or ")]
    return [c]

def rows():
    return json.load(open(BOM))["rows"]

def suppliers():
    out = set()
    for r in rows():
        for a in alternatives(r.get("vendor")):
            if a not in UNKNOWN and a != "EXOS (in-house)" and a not in MARKETPLACE: out.add(a)
    return sorted(out)

def cmd_names():
    reg = json.load(open(REG)); idx = reg_index(reg)
    need = []
    for n in suppliers():
        s = idx.get(key(n))
        if s: s["defense_use"] = True; s["programs"] = sorted(set(s.get("programs", [])) | {"HVRR"})
        if not s or s.get("tier") != "GREEN": need.append(n)
    need += [x for x in ("Omega Engineering",) if x not in need]   # proposed source for an unnamed BOM line
    tmp = REG + ".tmp"; json.dump(reg, open(tmp, "w"), indent=2); os.replace(tmp, REG)
    os.makedirs(OUT, exist_ok=True)
    open(os.path.join(OUT, "screen_names.txt"), "w").write("\n".join(need) + "\n")
    print(f"{len(suppliers())} HVRR suppliers; {len(need)} need a defense screen -> {OUT}/screen_names.txt")

def status_of(n, idx):
    if n == "EXOS (in-house)": return "In-house (make)", True
    if n in UNKNOWN: return "No supplier named on BOM", False
    if n in MARKETPLACE: return "Marketplace — name a distributor", False
    s = idx.get(key(n))
    if not s: return "Not screened", False
    if s.get("tier") == "GREEN" or s.get("status") == "cleared": return "DD cleared", True
    if s.get("status") in ("held", "held_at_order"): return f"Decision at order ({s.get('tier')}: {'; '.join((s.get('hard_reasons') or []) + (s.get('flags') or []))[:120]})", False
    return f"{s.get('tier')} — {s.get('status')}", False

def cmd_report():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    idx = reg_index(json.load(open(REG)))
    wb = Workbook(); ws = wb.active; ws.title = "BOM coverage"
    hdr = ["Section", "Item", "Description", "Vendor P/N", "Qty", "BOM vendor", "Supplier (normalized)", "DD status", "Covered", "Action"]
    ws.append(hdr)
    gaps = defaultdict(list); n_cov = 0; total = 0
    for r in rows():
        total += 1
        alts = alternatives(r.get("vendor"))
        st = [(a,) + status_of(a, idx) for a in alts]
        ok = next((x for x in st if x[2]), None)
        a, s, cov = ok or st[0]
        act = "" if cov else ("Name a supplier" if a in UNKNOWN else "Name a distributor (authorized)" if a in MARKETPLACE
                              else "Exo: find an alternate cleared supplier / decision at order")
        if cov: n_cov += 1
        else: gaps[a].append(r)
        ws.append([r.get("section"), r.get("item"), r.get("description"), r.get("vendor_pn"), r.get("qty"), r.get("vendor"), a, s, "Yes" if cov else "No", act])
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0A1628")
    for col, w in zip("ABCDEFGHIJ", (16, 6, 48, 16, 6, 22, 26, 40, 8, 40)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    try: prop = json.load(open(os.path.join(OUT, "proposed_sources.json")))["lines"]
    except Exception: prop = []
    if prop:
        ps = wb.create_sheet("Proposed sources")
        ps.append(["BOM line", "Vendor P/N", "BOM vendor", "Exo-proposed supplier", "Proposed supplier DD", "Why / verify"])
        for r in rows():
            v = (r.get("vendor") or "").strip()
            if canon(v) not in UNKNOWN and canon(v) not in MARKETPLACE: continue
            p = next((x for x in prop if (x.get("match") and x["match"].lower() in (r.get("description") or "").lower())
                      or (x.get("vendor_pn") and x["vendor_pn"] == r.get("vendor_pn"))), None)
            if p: ps.append([r.get("description"), r.get("vendor_pn"), v or "(blank)", p["proposed"], status_of(canon(p["proposed"]) if p["proposed"] != "EXOS (in-house)" else "EXOS (in-house)", idx)[0], p["why"]])
        for c in ps[1]: c.font = Font(bold=True)
        for col, w in zip("ABCDEF", (50, 18, 12, 26, 30, 70)): ps.column_dimensions[col].width = w
    g = wb.create_sheet("Gaps by supplier")
    g.append(["Supplier", "Parts affected", "DD status", "Next step"])
    for a, rs in sorted(gaps.items(), key=lambda x: -len(x[1])):
        g.append([a or "(blank)", len(rs), status_of(a, idx)[0],
                  "Name a supplier on the BOM" if a in UNKNOWN else "Pick an authorized distributor" if a in MARKETPLACE
                  else "Decision card raised when a PO is placed; Exo sources a cleared alternate"])
    for c in g[1]: c.font = Font(bold=True)
    g.column_dimensions["A"].width = 34; g.column_dimensions["C"].width = 60; g.column_dimensions["D"].width = 60
    s = wb.create_sheet("Summary", 0)
    for row in (["HVRR BOM supplier DD coverage", datetime.date.today().isoformat()],
                ["Contract", "FA9300-25-C-6022 (USAF/AFRL) — defense, so supplier DD applies"],
                ["BOM lines", total], ["Lines with a DD-cleared (or in-house) supplier", n_cov], ["Lines not yet covered", total - n_cov],
                ["Suppliers with gaps", len(gaps)],
                ["Rule", "Every part needs at least one DD-cleared supplier before it is ordered on HVRR. Uncleared suppliers get a Decision card when the PO is placed."]):
        s.append(row)
    s["A1"].font = Font(bold=True, size=14); s.column_dimensions["A"].width = 44; s.column_dimensions["B"].width = 90
    fn = os.path.join(OUT, f"HVRR_Supplier_DD_Coverage_{datetime.date.today().isoformat()}.xlsx")
    wb.save(fn); print(fn, f"covered {n_cov}/{total}, {len(gaps)} supplier gaps")

if __name__ == "__main__":
    {"names": cmd_names, "report": cmd_report}[sys.argv[1]]()
