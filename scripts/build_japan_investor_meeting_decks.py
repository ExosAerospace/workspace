#!/usr/bin/env python3
"""Japan investor meeting decks, Fri 10/16/2026 (Tokyo): SBI Investment (CVC Dept, Ms. Yang Yang) 14:00 JST
and 650 Fund (Dominik Steiner, @JETRO) 16:00 JST. Editable 16:9 pptx, Exos brand (black/grey + orange).
Claim rules: Part 431 license transitioning to Part 450 (never bare 'FAA-licensed BLK3'); BLK3 first commercial
flight early 2027; SARGE commercial payloads since 2018; round = EAI LLC units, offshore Reg S, mirror Exos U.S. shares."""
import os, sys
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

WS="/Users/exomac/.openclaw/workspace"
BGIMG=os.path.join(WS,"drafts","_exos_slide_bg.png"); LOGO=os.path.join(WS,"drafts","_exos_white_official.png")
BG=RGBColor(0x0d,0x0d,0x0f); PANEL=RGBColor(0x17,0x18,0x1c); ORANGE=RGBColor(0xF2,0x6A,0x21)
WHITE=RGBColor(0xFF,0xFF,0xFF); GREY=RGBColor(0xD2,0xD2,0xD7); MUTE=RGBColor(0x9a,0x9a,0xa1)
DISCL=("Not an offer to sell or a solicitation of an offer to buy securities. Units of Exos Aerospace Investments, LLC are "
       "offered only to non-U.S. persons in offshore transactions under Regulation S, separate from any U.S. offering, and only "
       "through definitive offering documents. Financials FY2025, accrual basis, unaudited.")

AUD={
 "sbi":{"out":"drafts/Exos_Japan_Round_SBI_Investment_2026-10-16.pptx","for":"SBI Investment Co., Ltd.  ·  CVC Department",
        "when":"Tokyo  ·  16 October 2026"},
 "650":{"out":"drafts/Exos_Japan_Round_650_Fund_2026-10-16.pptx","for":"650 Fund  ·  Dominik Steiner",
        "when":"JETRO Tokyo  ·  16 October 2026"},
}

def build(key):
    a=AUD[key]; p=Presentation(); p.slide_width=I(13.333); p.slide_height=I(7.5); BL=p.slide_layouts[6]; n=[0]
    def tb(s,x,y,w,h,text,size=14,color=GREY,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
        t=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=t.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
        for i,line in enumerate(text if isinstance(text,list) else [text]):
            para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.alignment=align
            r=para.add_run(); r.text=line; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; r.font.name="Calibri"
            para.space_after=Pt(size*0.45)
        return t
    def rect(s,x,y,w,h,fill=PANEL,line=None):
        sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,I(x),I(y),I(w),I(h)); sh.adjustments[0]=0.05
        sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.shadow.inherit=False
        if line: sh.line.color.rgb=line; sh.line.width=Pt(1.25)
        else: sh.line.fill.background()
        return sh
    def slide(kicker,title,notes=""):
        s=p.slides.add_slide(BL); n[0]+=1
        s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
        s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,I(0.12),I(7.5)).fill.solid()
        bar=s.shapes[-1]; bar.fill.fore_color.rgb=ORANGE; bar.line.fill.background()
        tb(s,0.6,0.35,10,0.4,kicker,12,ORANGE,True); tb(s,0.6,0.7,12,0.9,title,30,WHITE,True)
        s.shapes.add_picture(LOGO,I(11.35),I(0.38),height=I(0.42))
        tb(s,0.6,7.0,10,0.3,f"EXOS AEROSPACE  ·  JAPAN ROUND  ·  {a['for']}",9,MUTE)
        tb(s,12.2,7.0,0.6,0.3,f"{n[0]:02d}",9,MUTE,align=PP_ALIGN.RIGHT)
        if notes: s.notes_slide.notes_text_frame.text=notes
        return s
    def card(s,x,y,w,h,head,body,hs=18,bs=15.5,accent=False):
        rect(s,x,y,w,h,line=ORANGE if accent else None)
        tb(s,x+0.25,y+0.18,w-0.5,0.5,head,hs,ORANGE if accent else WHITE,True)
        tb(s,x+0.25,y+0.72,w-0.5,h-0.85,body,bs,GREY)
    def stat(s,x,y,w,big,small):
        rect(s,x,y,w,1.75); tb(s,x+0.2,y+0.12,w-0.4,0.8,big,36,ORANGE,True); tb(s,x+0.2,y+0.9,w-0.4,0.85,small,13.5,GREY)

    # 1 Title
    s=p.slides.add_slide(BL); n[0]+=1
    s.shapes.add_picture(BGIMG,0,0,I(13.333),I(7.5)); s.shapes.add_picture(LOGO,I(0.8),I(0.7),height=I(0.7))
    tb(s,0.8,2.2,11.5,0.5,"JAPAN ROUND  ·  ¥790M (≈ US$5.0M)",16,ORANGE,True)
    tb(s,0.8,2.75,11.5,1.6,["Making every space vehicle reusable —","built in Japan, together with Japan"],38,WHITE,True)
    tb(s,0.8,4.75,11.5,0.5,f"Prepared for {a['for']}",18,GREY)
    tb(s,0.8,5.25,11.5,0.5,a["when"],14,MUTE)
    tb(s,0.8,6.35,11.5,0.5,"John Quinn, Co-founder & CEO   ·   Toshi Kumazawa, Japan Business Development",14,GREY)
    s.notes_slide.notes_text_frame.text="Thank them for the time. One line: we are a profitable U.S. reusable-launch company building a Japan operating company, and we came to Japan this week to put it together."

    # 2 Who we are
    s=slide("01  |  WHO WE ARE","A profitable U.S. reusable-launch company with real flight heritage",
      "Say 'one of only three' exactly as written. Do not say BLK3 is licensed — say our license is Part 431, transitioning to Part 450.")
    stat(s,0.6,1.85,2.9,"1 of 3","U.S. FAA-licensed reusable rocket launch providers — with SpaceX and Blue Origin")
    stat(s,3.7,1.85,2.9,"2018","commercial payloads flown on SARGE, our first reusable vehicle")
    stat(s,6.8,1.85,2.9,"25+ yrs","reusable-vehicle heritage (Armadillo Aerospace lineage)")
    stat(s,9.9,1.85,2.85,"$1.49M","2025 revenue · net income $188K — profitable (filed 2025 tax return; unaudited)")
    card(s,0.6,3.95,6.0,2.75,"What we do",["• Reusable suborbital launch and flight test for payloads, universities and launch companies",
        "• Recovery systems that bring vehicles and hardware back to fly again",
        "• Operating under a Part 431 FAA license, transitioning to Part 450"])
    card(s,6.8,3.95,5.95,2.75,"Group footprint",["• Exos Aerospace — Greenville, Texas (HQ, manufacturing)",
        "• Exos Aerospace Italia, Srl — Torino",
        "• Exos Aerospace Japan K.K. — Japan operating company (forming)",
        "• 18% stockholder in Scorpius Space Launch Co. (cryogenic tanks)"])

    # 3 Technology
    s=slide("02  |  TECHNOLOGY","Proven hardware — not a promise",
      "BLK3 has not flown yet: 'first commercial flight of BLK3, our latest vehicle, early 2027'. Scorpius tanks are a Scorpius (SSLC) product, Exos is an 18% stockholder.")
    card(s,0.6,1.85,4.0,4.85,"BLK3 reusable rocket",["The enhanced successor to SARGE.","• Out-of-autoclave carbon composite airframe",
        "• Under 5,000 components — built for low cost and fast turnaround","• Designed to be recovered by drogue + parafoil and reflown",
        "• First commercial flight of BLK3: early 2027"],accent=True)
    card(s,4.75,1.85,4.0,4.85,"Type-5 cryo tanks (Scorpius)",["Linerless composite tanks — lighter than metal.",
        "• Flew as fuel and LOX tanks on Intuitive Machines' Odysseus (2024) and Athena (2025) lunar landers",
        "• First tank sale for an orbital transfer vehicle"])
    card(s,8.9,1.85,3.85,4.85,"Recovery know-how",["Our core skill: getting expensive hardware back.",
        "• Drogue and parafoil recovery flown on SARGE since 2018",
        "• Directly useful to Japan's new launch companies, which lose test hardware on every flight"])

    # 4 Why Japan
    s=slide("03  |  WHY JAPAN — AND WHAT HAPPENED THIS WEEK","The Japan platform is already moving",
      "This week: Hokkaido Spaceport event 10/13-14. NYK = 'in discussion', no agreement yet. Do not name AstroX or other companies' plans in detail; keep it at 'Japanese launch startups'.")
    card(s,0.6,1.85,6.0,2.35,"Hokkaido Spaceport (HOSPO)",["Collaboration with Space Cotan to operate BLK3 from HOSPO on a wet-lease basis",
        "for Japanese government, university and commercial customers."],accent=True)
    card(s,6.8,1.85,5.95,2.35,"Logistics",["In discussion with NYK Line to ship BLK3 and ground equipment from Texas to Japan",
        "— a Japanese partner on the critical path."])
    card(s,0.6,4.35,6.0,2.35,"Customer signal",["• LOI for a fuel-tank order from a Japanese space-development company",
        "• Recovery and flight-test discussions with Japanese launch startups met this week"])
    card(s,6.8,4.35,5.95,2.35,"Bridge service",["Japanese payloads can fly on Exos BLK3 from the U.S. starting early 2027,",
        "while the Japan vehicle and approvals are completed."])

    # 5 The round
    s=slide("04  |  THE INVESTMENT","A structured round to build Exos Aerospace Japan",
      "Investors are PASSIVE: no board seat, no access to technical data. That is what keeps U.S. export and CFIUS reviews simple — present it as a feature. Pricing: off the Exos U.S. valuation ($169.7M equity value, independent 409A, 31 Mar 2026). Do not discuss the RTO.")
    rect(s,0.6,1.85,4.3,4.85,line=ORANGE)
    tb(s,0.85,2.05,3.8,0.4,"JAPAN ROUND — REG S",12,ORANGE,True); tb(s,0.85,2.5,3.8,0.9,"¥790M",54,WHITE,True)
    tb(s,0.85,3.55,3.8,0.5,"≈ US$5.0M target (at ~¥157/US$)",14,GREY)
    tb(s,0.85,4.25,3.8,2.3,["• Units in Exos Aerospace Investments, LLC","• EAI units mirror Exos U.S. shares",
        "• Priced off the Exos U.S. valuation","• Offshore Regulation S — separate from the U.S. offering"],15,GREY)
    card(s,5.1,1.85,7.65,2.55,"Use of funds",["1. Stand up Exos Aerospace Japan K.K. — office and local team",
        "2. Secure a BLK3 wet-lease for HOSPO   3. Bridge flights for Japanese customers from the U.S.",
        "4. Licensed Type-5 tank production with Japanese partners"])
    card(s,5.1,4.55,7.65,2.15,"Structure",["Exos Aerospace Japan K.K. is wholly owned by Exos U.S. (same model as Exos Aerospace Italia).",
        "Investors hold a passive financial stake: no board seat or technical-data access, which keeps U.S. export review straightforward."])
    tb(s,0.6,6.72,12.1,0.3,DISCL,8,MUTE)

    # 6 Returns
    s=slide("05  |  RETURNS","Grow recurring revenue, then a defined exit",
      "Horizon ~3+ years. Exit through an Exos U.S. event (listing or acquisition), realized via EAI. Keep it there; do not give a listing date.")
    for i,(h,b) in enumerate([("NOW — THE ROUND","Stand up the Japan company; first wet-lease and bridge flights for Japanese customers."),
                              ("SCALE","Recurring revenue: flights · Type-5 tank and engine sales · recovery and flight-test services."),
                              ("EXIT","Liquidity through an Exos U.S. event — public listing or acquisition — realized via EAI LLC. Horizon ~3+ years.")]):
        card(s,0.6+i*4.1,1.95,3.9,3.0,h,b,18,16,accent=(i==2))
    card(s,0.6,5.15,12.15,1.5,"Non-dilutive layers alongside the equity",
        ["Payload pre-sales (reserve BLK3 capacity) · Japanese government programs · strategic partner contracts"],18,15)
    tb(s,0.6,6.72,12.1,0.3,DISCL,8,MUTE)

    # 7 Audience-specific
    if key=="sbi":
        s=slide("06  |  WHY SBI","Capital plus the right Japanese corporate partners",
          "Source for the CVC list: SBI Investment Open Innovation page (sbinvestment.co.jp/en/features/open-innovation) — 24 CVC funds, ¥182.5B as of 17 Apr 2026. These are POTENTIAL fits only; we have no contact with these companies. Ask Ms. Yang which CVC partners look at space.")
        card(s,0.6,1.85,5.2,4.85,"What we are asking",["• An investment from SBI in the ¥790M Japan round",
            "• Introductions to SBI's corporate venture partners where there is a real business fit",
            "• Advice on Japanese co-investors and the right round structure"],accent=True)
        card(s,6.0,1.85,6.75,4.85,"Possible fits among SBI's CVC partners",[
            "• SUBARU — aerospace division; flight-test and composite structures",
            "• Nikon — optics and space sensors; payload flights",
            "• Nippon Express (NX Global Innovation Fund) and SANKYU — heavy logistics for vehicles and ground equipment",
            "• Mitsui Kinzoku — advanced materials for composite and cryogenic tanks",
            "Exos brings a Japan operating company, HOSPO operations and proven U.S. hardware."],16,14.5)
    else:
        s=slide("06  |  WHY 650","A Japan-relevant deep-tech company in the 'missing middle'",
          "650 Fund II leads pre-seed/seed only; Series A+ is follow-on. So the ask is an allocation sized for them ($250K-$2.5M range Steiner calls the 'missing middle') plus network, not a lead. Fit: Deep Tech pillar (materials) and Japan thesis. Sources: 650.fund/fundii; SCCIJ interview with Steiner.")
        card(s,0.6,1.85,5.2,4.85,"What we are asking",["• An allocation in the ¥790M Japan round sized for 650",
            "• Introductions to Japanese co-investors and corporate partners",
            "• 650's cross-border know-how: U.S., Switzerland/Europe and Japan"],accent=True)
        card(s,6.0,1.85,6.75,4.85,"Why it fits 650",[
            "• Deep tech: linerless composite cryogenic tanks and reusable flight systems",
            "• A clear Japan thesis: Exos Aerospace Japan K.K., HOSPO operations, Japanese customers and partners",
            "• Cross-border: U.S. headquarters, Exos Italia in Torino, Japan operating company",
            "• Profitable company, so the capital goes into growth, not survival"],16,14.5)

    # 8 Next steps
    s=slide("07  |  NEXT STEPS","Build Exos Aerospace Japan with us",
      "Close by agreeing one next step with a date. Offer the data room after an NDA. Toshi follows up in Japanese.")
    card(s,0.6,1.85,6.0,4.0,"Proposed next steps",["1. Share the data room under NDA","2. Follow-up call with the investment team",
        "3. Visit Exos in Greenville, Texas, or see BLK3 testing","4. Term discussion on the Reg S round"],accent=True)
    card(s,6.8,1.85,5.95,4.0,"Contacts",["John Quinn — Co-founder & CEO","jquinn@exosaero.com  ·  +1-972-740-8355","",
        "Toshi Kumazawa — Japan Business Development","toshihisa@exosaero.com","","www.exosaero.com"])
    tb(s,0.6,6.72,12.1,0.3,DISCL,8,MUTE)
    out=os.path.join(WS,a["out"]); p.save(out); print("WROTE",out)

for k in (sys.argv[1:] or AUD): build(k)
