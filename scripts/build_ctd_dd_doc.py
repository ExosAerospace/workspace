#!/usr/bin/env python3
"""Capture the MC CTD Data-Room DD tracker into one Word document (confidential, internal)."""
import json
from collections import Counter
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC="mission-control/data/ctd-dd.json"
OUT="drafts/CTD_DueDiligence_Tracker.docx"
NAVY=RGBColor(0x0b,0x1a,0x3a); GOLD=RGBColor(0xb8,0x86,0x1f); DARK=RGBColor(0x1a,0x1a,0x1a); GREY=RGBColor(0x55,0x55,0x55); RED=RGBColor(0x8a,0x1a,0x1a); GREEN=RGBColor(0x1c,0x6b,0x2f)
data=json.load(open(SRC)); meta=data.get("meta",{}); sections=data.get("sections",[])
doc=Document()
for s in doc.sections:
    s.left_margin=s.right_margin=Inches(0.7); s.top_margin=s.bottom_margin=Inches(0.6)

def p(text="",size=10.5,bold=False,color=DARK,after=5,before=0,italic=False,align=None):
    pa=doc.add_paragraph(); pa.paragraph_format.space_after=Pt(after); pa.paragraph_format.space_before=Pt(before)
    if align: pa.alignment=align
    r=pa.add_run(text); r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color; r.font.name="Calibri"
    return pa

STATUS_COLOR={"received":GREEN,"complete":GREEN,"done":GREEN,"requested":GOLD,"pending":GOLD,"outstanding":RED,"missing":RED,"n/a":GREY,"na":GREY}
def scolor(st): return STATUS_COLOR.get((st or "").strip().lower(), DARK)

# Header / marking
p("CONFIDENTIAL  ·  PROPRIETARY  ·  INTERNAL USE ONLY", size=8.5, bold=True, color=RED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
p(meta.get("title","CTD — Data-Room Due Diligence Tracker"), size=16, bold=True, color=NAVY, after=2)
# meta block
mb=[("Company", meta.get("company","")),("Counterparty", meta.get("counterparty","")),
    ("Review window", meta.get("window","")),("Sequence", meta.get("sequence","")),
    ("Started", meta.get("started","")),("Updated", meta.get("updated",""))]
for k,v in mb:
    if not v: continue
    pa=doc.add_paragraph(); pa.paragraph_format.space_after=Pt(1)
    r=pa.add_run(f"{k}:  "); r.font.size=Pt(10); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name="Calibri"
    r2=pa.add_run(str(v)); r2.font.size=Pt(10); r2.font.color.rgb=DARK; r2.font.name="Calibri"
if meta.get("note"): p("Note: "+meta["note"], size=9.5, italic=True, color=RED, before=4, after=8)

# Status summary
allitems=[it for s in sections for it in s.get("items",[])]
cnt=Counter((it.get("status") or "Unspecified").strip() for it in allitems)
p(f"Summary: {len(allitems)} diligence items across {len(sections)} sections  —  "
  + "  ·  ".join(f"{v} {k}" for k,v in cnt.most_common()), size=9.5, bold=True, color=GREY, after=10)

def shade(cell,hex_):
    from docx.oxml.ns import qn; from docx.oxml import OxmlElement
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hex_); tcPr.append(sh)
def setcell(cell,text,size=9.5,bold=False,color=DARK):
    cell.paragraphs[0].paragraph_format.space_after=Pt(1)
    r=cell.paragraphs[0].add_run(str(text or "")); r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name="Calibri"

for s in sections:
    title=f"{s.get('id','')}.  {s.get('title','')}"
    pa=doc.add_paragraph(); pa.paragraph_format.space_before=Pt(10); pa.paragraph_format.space_after=Pt(1)
    r=pa.add_run(title); r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name="Calibri"
    if s.get("priority"):
        rp=pa.add_run("   ★ PRIORITY"); rp.font.size=Pt(9); rp.font.bold=True; rp.font.color.rgb=RED; rp.font.name="Calibri"
    if s.get("subtitle"): p(s["subtitle"], size=9, italic=True, color=GREY, after=4)
    items=s.get("items",[])
    if not items: continue
    tbl=doc.add_table(rows=1, cols=4); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit=False
    widths=[Inches(0.55),Inches(4.2),Inches(0.95),Inches(1.5)]
    hdr=tbl.rows[0].cells
    for i,h in enumerate(["ID","Diligence request","Status","Notes"]):
        setcell(hdr[i],h,size=9,bold=True,color=RGBColor(0xff,0xff,0xff)); shade(hdr[i],"0b1a3a")
    for it in items:
        cells=tbl.add_row().cells
        setcell(cells[0],it.get("id",""),size=8.5,bold=True,color=NAVY)
        setcell(cells[1],it.get("q",""),size=9.5)
        setcell(cells[2],it.get("status",""),size=9,bold=True,color=scolor(it.get("status")))
        note=it.get("notes","") or ("received "+it["received"] if it.get("received") else "")
        setcell(cells[3],note,size=8.5,color=GREY)
    for row in tbl.rows:
        for i,w in enumerate(widths): row.cells[i].width=w

p("", after=6)
p("Generated from Mission Control CTD DD tracker (ctd-dd.json). CTD identity confidential — do not use in public material.",
  size=8, italic=True, color=GREY, before=8)
doc.save(OUT); print("BUILT",OUT,"|",len(allitems),"items,",len(sections),"sections")
